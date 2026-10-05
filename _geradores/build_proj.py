# -*- coding: utf-8 -*-
"""Gera Projeto-modelo.xlsx no mesmo padrão visual das planilhas anteriores da série."""
import math
import os
import sys

_here = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, _here)
# funções de formatação compartilhadas com o gerador do PDCA (bloco inicial do arquivo)
_src = open(os.path.join(_here, "build_pdca.py"), encoding="utf-8").read()
exec(_src[: _src.index("STATUS = [")])

from openpyxl.chart import BarChart, Series  # noqa: E402
from proj_data import (ANAL, APROV, ATENDE, CHECK, EX1, EX2, LIB, NAOAT, PENDENTE, REPROV, RES_E, RESS, RESULT, TIPOS_E, TIPOS_ET, VALID, VERIF)  # noqa: E402

BLUE, TEAL, AMBER, PURPLE, REDC = "2B5C8A", "1E7B73", "A96A12", "7A4A9A", "B0413E"
S1, S2, S3 = "2A6FB0", "D19A2E", "B0413E"
S1_T, S2_T, S3_T = "DCE8F3", "F8EBCB", "F5DEDC"
YESNO = ["Sim", "Parcial", "Não"]
YESNO_CF = [("Sim", GREEN), ("Parcial", YELLOW), ("Não", RED)]
NE_, NX, NM = 20, 30, 20
PL, EN, MU = "Plano", "Entradas", "'Mudanças'"
E1, E2 = 14, 14 + NE_ - 1       # etapas, na aba Plano
X1, X2 = 7, 7 + NX - 1          # entradas
M1, M2 = 7, 7 + NM - 1          # mudanças
SEMVAL = "Liberado sem validação aprovada"
NAOAC, NAOEM = "Não atende, sem ação", "Não atende: ação em curso"


def note(ws, ref, text):
    c = Comment(text, "Modelo Projeto")
    c.width, c.height = 300, 120
    ws[ref].comment = c


def cf_warn(ws, rng_, first, ok_values=("OK",)):
    cond = ",".join(f'{first}="{v}"' for v in ok_values)
    ws.conditional_formatting.add(rng_, FormulaRule(formula=[f"OR({cond})"], fill=PatternFill("solid", bgColor=GREEN, fgColor=GREEN)))
    ws.conditional_formatting.add(rng_, FormulaRule(formula=[f"AND(LEN({first})>0,NOT(OR({cond})))"],
                                  fill=PatternFill("solid", bgColor=YELLOW, fgColor=YELLOW)))


def cf_texto(ws, rng_, first, pares, resto=None):
    for t, cor in pares:
        ws.conditional_formatting.add(rng_, FormulaRule(formula=[f'{first}="{t}"'], stopIfTrue=True, fill=PatternFill("solid", bgColor=cor, fgColor=cor)))
    if resto:
        ws.conditional_formatting.add(rng_, FormulaRule(formula=[f"LEN({first})>0"], fill=PatternFill("solid", bgColor=resto, fgColor=resto)))


def hint_row(ws, row, cells, height=30):
    for ref, text in cells:
        put(ws, f"{ref}{row}", text, f=font(9, i=True, c=MUTED), bg=GRAY, h="center")
    ws.row_dimensions[row].height = height


def alt(pares, minimo=21.75, linha=12):
    n = max(max(1, math.ceil(len(str(t)) / max(c, 1))) for t, c in pares)
    return max(minimo, n * linha + 7)


# ------------------------------------------------------------------ fórmulas compartilhadas pelas abas de entrada e pelos exemplos
def f_etapa(c, outros, tipos, results, ref):
    """Conferência de uma etapa; `c` traz as células da linha."""
    validou = f'COUNTIFS({tipos},"{VALID}",{results},"{APROV}")+COUNTIFS({tipos},"{VALID}",{results},"{RESS}")'
    return (f'IF({c["etapa"]}="",IF({outros}>0,"Falta a etapa",""),IF({c["tipo"]}="","Falta o tipo",IF({c["resp"]}="","Falta o responsável",'
            f'IF({c["prazo"]}="","Falta o prazo",IF(AND({c["concl"]}<>"",{c["result"]}=""),"Falta o resultado",IF(AND({c["result"]}<>"",{c["concl"]}=""),"Falta a data de conclusão",'
            f'IF(AND(OR({c["result"]}="{RESS}",{c["result"]}="{REPROV}"),{c["acoes"]}=""),"Falta a ação",'
            f'IF(AND({c["tipo"]}="{LIB}",{c["concl"]}<>"",{validou}=0),"{SEMVAL}",'
            f'IF(AND({c["concl"]}="",ISNUMBER({ref}),{c["prazo"]}<{ref}),"Atrasada","OK")))))))))')


def f_entrada(c, outros):
    return (f'IF({c["req"]}="",IF({outros}>0,"Falta a entrada",""),IF({c["tipo"]}="","Falta o tipo",IF({c["fonte"]}="","Falta a fonte",'
            f'IF({c["saida"]}="","Falta a saída que atende",IF({c["metodo"]}="","Falta como verificar",IF({c["result"]}="{NAOAT}",IF({c["evid"]}="","{NAOAC}","{NAOEM}"),'
            f'IF({c["result"]}="","{PENDENTE}","OK")))))))')


def f_mud(oque, outros, analise, verif, aut):
    return (f'IF({oque}="",IF({outros}>0,"Falta o que mudou",""),IF({analise}="","Falta a análise",IF({aut}="","Falta quem autorizou",'
            f'IF({verif}="","Falta o que foi verificado de novo","OK"))))')


RES_CF = [(APROV, GREEN), (RESS, S2_T), (REPROV, S3_T)]
ENT_CF = [(ATENDE, GREEN), (NAOAT, S3_T)]
CONF_ET = [("OK", GREEN), (SEMVAL, RED), ("Atrasada", RED)]
CONF_EN = [("OK", GREEN), (NAOAC, RED), (PENDENTE, S2_T), (NAOEM, S2_T)]


def chart_style(ch):
    ch.style = 2
    ch.title = None
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ch.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill="D3DBE0", w=9525))
    ch.x_axis.graphicalProperties = GraphicalProperties(ln=LineProperties(solidFill="D3DBE0", w=9525))
    ch.y_axis.graphicalProperties = GraphicalProperties(ln=LineProperties(noFill=True))


def empilhado(ws, anchor, r1, r2, c_cat, cols, width=20, height=9):
    """Barras horizontais empilhadas: entradas por tipo, por situação."""
    ch = BarChart()
    ch.type = "bar"
    ch.grouping = "stacked"
    ch.overlap = 100
    ch.gapWidth = 60
    ch.height, ch.width = height, width
    chart_style(ch)
    for (c, nome), cor in zip(cols, (S1, S3, S2)):
        s = Series(Reference(ws, min_col=c, min_row=r1, max_row=r2), title=nome)
        s.graphicalProperties.solidFill = cor
        s.graphicalProperties.line.noFill = True
        ch.append(s)
    ch.set_categories(Reference(ws, min_col=c_cat, min_row=r1, max_row=r2))
    ch.y_axis.scaling.min = 0
    ch.y_axis.majorUnit = 1
    ch.x_axis.scaling.orientation = "maxMin"
    ch.legend.position = "b"
    ws.add_chart(ch, anchor)


# ================================================================== pasta de trabalho
wb = Workbook()
wb.properties.title = "Projeto e desenvolvimento — Modelo"
wb.properties.creator = "Modelo Projeto"
wb.properties.language = "pt-BR"

# ------------------------------------------------------------------ Instruções
ws = wb.active
ws.title = "Instruções"
widths(ws, {"A": 2, "B": 26, "C": 90, "D": 2})
title(ws, "Projeto e desenvolvimento — Como usar esta planilha",
      "Modelo para planejar o projeto, ligar as entradas às saídas e às verificações, registrar os controles e as mudanças, e decidir a liberação.", "C")
r = 4


def section(text):
    global r
    band(ws, r, text, "C", sz=11)
    r += 1


def line(k, v, kbg=GRAY, vbg=None, kf=None, kh="left", height=None):
    global r
    put(ws, f"B{r}", k, f=kf or font(10, True), bg=kbg, h=kh)
    put(ws, f"C{r}", v, bg=vbg)
    ws.row_dimensions[r].height = height or (19.5 if len(v) <= 96 else (31.5 if len(v) <= 192 else 45.75))
    r += 1


section("Legenda: onde preencher")
line("Amarelo-claro", "Células de entrada. É aqui que você digita.", vbg=INPUT)
line("Cinza", "Células calculadas ou fixas (rótulos, conferências, contagens, avisos). Não altere.", vbg=GRAY)
line("Cinza em itálico", "Linha de exemplo (marcada com “Ex.”). Mostra o formato esperado e não entra nas contagens.", vbg=GRAY)
line("Listas suspensas", "Tipo da etapa, resultado, tipo da entrada, resultado da verificação e checklist aceitam apenas as opções da lista.")
line("Comentários", "Células com triângulo vermelho no canto trazem uma dica de preenchimento. Passe o mouse sobre elas.")
r += 1
section("As regras de cálculo")
for k, text in [
    ("Atrasada", "Etapa sem data de conclusão, com o prazo anterior à data da leitura, informada na aba Plano."),
    ("Falta a ação", f"Etapa com o resultado “{RESS}” ou “{REPROV}”, sem a ação registrada."),
    (SEMVAL, f"Etapa de “{LIB}” concluída sem nenhuma etapa de “{VALID}” aprovada, com ou sem ressalvas."),
    (PENDENTE, "Entrada com tudo preenchido, sem o resultado da verificação."),
    (NAOAT, "Com a ação escrita na coluna Evidência ou ação, a conferência mostra a ação em curso. Sem ela, pede a ação."),
    ("Pode liberar?", "Sim quando há validação aprovada, todas as entradas atendem, nenhuma mudança está a rever e nenhuma etapa está atrasada ou a rever."),
]:
    line(k, text)
r += 1
section("Ordem recomendada de preenchimento")
for k, text in enumerate([
    "Aba Plano: escreva o projeto, o objetivo, o cliente ou o usuário, o responsável, as datas e a data da leitura.",
    "Aba Plano: liste as etapas, com tipo, responsável, participantes e prazo. Inclua análises críticas, verificações e validações.",
    "Aba Entradas: escreva as entradas pelos sete tipos, com a fonte, a saída que atende e como verificar.",
    "Abas Plano e Entradas: conforme o projeto anda, lance as conclusões, os resultados, as ações e as evidências.",
    "Aba Mudanças: registre toda alteração em entrada ou saída já aprovada, com a análise, o que foi verificado de novo e quem autorizou.",
    "Aba Painel: leia as pendências e a resposta à pergunta “pode liberar?”. Depois, valide na aba Checklist.",
], 1):
    line(f"Passo {k}", text)
r += 1
section("Abas da planilha")
for k, text in [
    (PL, f"O projeto e até {NE_} etapas. Calcula a conferência de cada etapa e se o plano tem os três controles."),
    (EN, f"Até {NX} entradas, ligadas às saídas e às verificações. Calcula a conferência de cada entrada."),
    ("Mudanças", f"Até {NM} mudanças de projeto. Calcula a conferência de cada uma."),
    ("Painel", "Os indicadores, os controles, as entradas por tipo, o gráfico e a decisão de liberar."),
    ("Checklist", "Doze verificações do processo de projeto, com percentual de conclusão."),
    ("Exemplo 1 - Pizzaria", "A pizza vegana do cardápio de inverno, lida em 31/05/2027."),
    ("Exemplo 2 - Indústria", "O filme para vegetais congelados, lido em 30/08/2027."),
]:
    line(k, text)
r += 1
section("Regras e premissas do modelo")
for k, text in [
    ("Tamanho do controle", "Um projeto pequeno pode ter cinco etapas e seis entradas. O que não pode faltar é ao menos uma verificação, uma validação e uma análise crítica."),
    ("Resultado da etapa", f"Use “{APROV}”, “{RESS}” ou “{REPROV}” para todas as etapas concluídas, inclusive as de desenvolvimento."),
    ("Entrada nova", "A entrada que aparece durante o projeto, por exemplo na validação, entra na lista com a fonte que a revelou."),
    ("Exemplos", "Os dados das abas de exemplo são ilustrativos, criados para este material."),
]:
    line(k, text)
setup(ws, INK, f"B1:C{r}", landscape=False, fit_height=True)

# ------------------------------------------------------------------ Plano
ws = wb.create_sheet(PL)
widths(ws, {"A": 2, "B": 5, "C": 40, "D": 20, "E": 22, "F": 24, "G": 12, "H": 13, "I": 20, "J": 36, "K": 24, "L": 30, "M": 2})
title(ws, "Plano do projeto", "O projeto no alto. Abaixo, as etapas, com os três controles: análise crítica, verificação e validação.", "L")
for k, text in enumerate(["Organização", "Projeto", "Objetivo", "Cliente ou usuário", "Responsável e aprovação"]):
    label(ws, f"B{4 + k}", text, merge=f"B{4 + k}:C{4 + k}")
    inp(ws, f"D{4 + k}", merge=f"D{4 + k}:L{4 + k}", h="left")
    ws.row_dimensions[4 + k].height = 21.75
for rr, (a, b) in ((9, ("Início", "Liberação prevista")), (10, ("Data da leitura", None))):
    label(ws, f"B{rr}", a, merge=f"B{rr}:C{rr}")
    inp(ws, f"D{rr}", h="center", fmt=DATE)
    if b:
        label(ws, f"E{rr}", b)
        inp(ws, f"F{rr}", h="center", fmt=DATE)
    ws.row_dimensions[rr].height = 21.75
put(ws, "E10", "Usada para saber o que está atrasado. Atualize a cada leitura.", f=font(9, i=True, c=MUTED), bg=GRAY, merge="E10:L10")
dv_date(ws, "D9:D10")
dv_date(ws, "F9")
REF = "$D$10"
for col, text in zip("BCDEFGHIJKL", ["#", "Etapa", "Tipo", "Responsável", "Participantes", "Prazo", "Concluída em", "Resultado", "Ações e ressalvas", "Registro", "Conferência"]):
    head(ws, f"{col}12", text)
ws.row_dimensions[12].height = 33
hint_row(ws, 13, [("B", ""), ("C", "O que será feito"), ("D", "Lista"), ("E", "Uma função"), ("F", "Áreas, cliente"), ("G", "Data"), ("H", "Data"), ("I", "Lista"),
                  ("J", "Obrigatórias com ressalva ou reprovação"), ("K", "Ata, relatório"), ("L", "Calculada")])
TIP, RES = f"$D${E1}:$D${E2}", f"$I${E1}:$I${E2}"
for k in range(NE_):
    rr = E1 + k
    num(ws, f"B{rr}", k + 1)
    for col in "CEFJK":
        inp(ws, f"{col}{rr}")
    inp(ws, f"D{rr}", h="center")
    inp(ws, f"G{rr}", h="center", fmt=DATE)
    inp(ws, f"H{rr}", h="center", fmt=DATE)
    inp(ws, f"I{rr}", h="center")
    cells = dict(etapa=f"C{rr}", tipo=f"D{rr}", resp=f"E{rr}", prazo=f"G{rr}", concl=f"H{rr}", result=f"I{rr}", acoes=f"J{rr}")
    calc(ws, f"L{rr}", "=" + f_etapa(cells, f"COUNTA(D{rr}:K{rr})", TIP, RES, REF), b=False, sz=9)
    ws.row_dimensions[rr].height = 30
dv_list(ws, f"D{E1}:D{E2}", TIPOS_ET, "Planejamento, desenvolvimento, análise crítica, verificação, validação ou liberação")
dv_list(ws, f"I{E1}:I{E2}", RESULT, "Resultado da etapa concluída")
dv_date(ws, f"G{E1}:H{E2}")
cf_texto(ws, f"I{E1}:I{E2}", f"I{E1}", RES_CF)
cf_texto(ws, f"L{E1}:L{E2}", f"L{E1}", CONF_ET, resto=YELLOW)
note(ws, "D12", "Verificação: as saídas atendem às entradas? Validação: o produto serve para o uso pretendido? Análise crítica: o projeto segue?")
note(ws, "J12", "O que será feito sobre a ressalva ou a reprovação, com prazo. Na validação reprovada, diga também a entrada nova, se houver.")
s = E2 + 2
band(ws, s, "Resumo automático", "L")
PR = [("Etapas no plano", f"=COUNTA(C{E1}:C{E2})"),
      ("Controles no plano", f'=IF(D{s+1}=0,"",COUNTIF({TIP},"{ANAL}")&" análises críticas, "&COUNTIF({TIP},"{VERIF}")&" verificações, "&COUNTIF({TIP},"{VALID}")&" validações")'),
      ("Concluídas e atrasadas", f'=IF(D{s+1}=0,"",COUNTA(H{E1}:H{E2})&" concluídas, "&COUNTIF(L{E1}:L{E2},"Atrasada")&" atrasadas")'),
      ("Aviso", f'=IF(D5="","Escreva o projeto",IF(D{s+1}=0,"Escreva as etapas",IF(COUNTIF(L{E1}:L{E2},"{SEMVAL}")>0,"Há liberação sem validação aprovada",'
                f'IF(COUNTIF({TIP},"{VALID}")=0,"O plano não tem validação",IF(COUNTIF({TIP},"{VERIF}")=0,"O plano não tem verificação",'
                f'IF(COUNTIF({TIP},"{ANAL}")=0,"O plano não tem análise crítica",IF(COUNTIF(L{E1}:L{E2},"Atrasada")>0,"Há etapa atrasada",'
                f'IF(SUMPRODUCT((L{E1}:L{E2}<>"")*(L{E1}:L{E2}<>"OK"))>0,"Há etapa a rever: veja a coluna Conferência","OK"))))))))')]
for k, (text, formula) in enumerate(PR, 1):
    label(ws, f"B{s+k}", text, merge=f"B{s+k}:C{s+k}", h="right")
    calc(ws, f"D{s+k}", formula, merge=f"D{s+k}:L{s+k}", sz=9 if text == "Aviso" else 10, b=text != "Aviso", h="left")
    ws.row_dimensions[s + k].height = 21.75
PAV = s + 4
cf_warn(ws, f"D{PAV}:L{PAV}", f"D{PAV}")
ws.freeze_panes = f"D{E1}"
setup(ws, BLUE, f"B1:L{PAV}", fit_height=True)

# ------------------------------------------------------------------ Entradas
ws = wb.create_sheet(EN)
widths(ws, {"A": 2, "B": 5, "C": 9, "D": 40, "E": 22, "F": 26, "G": 30, "H": 28, "I": 13, "J": 34, "K": 28, "L": 2})
title(ws, "Entradas, saídas e verificação", "Uma linha por entrada. Cada entrada aponta a saída que a atende e como se verifica: é a matriz de rastreabilidade do projeto.", "K")
for col, text in zip("BCDEFGHIJK", ["#", "Código", "Entrada", "Tipo", "Fonte", "Saída que atende", "Como verificar", "Resultado", "Evidência ou ação", "Conferência"]):
    head(ws, f"{col}4", text)
ws.row_dimensions[4].height = 33
hint_row(ws, 5, [("B", ""), ("C", "E1, E2"), ("D", "Com número e condição"), ("E", "Lista"), ("F", "De onde veio"), ("G", "Onde o projeto atende"), ("H", "Método e critério"),
                 ("I", "Lista"), ("J", "O resultado, ou o que será feito"), ("K", "Calculada")])
ex(ws, "B6", "Ex.", h="center")
for col, v, h_ in [("C", "E3", "center"), ("D", "Queda de dardo de pelo menos 300 g a -18 °C.", "left"), ("E", TIPOS_E[1], "center"), ("F", "Uso pretendido: freezer do cliente", "left"),
                   ("G", "Resina para baixa temperatura", "left"), ("H", "Ensaio de queda de dardo a -18 °C", "left"), ("I", ATENDE, "center"), ("J", "420 g, em 16/07.", "left"),
                   ("K", "OK", "center")]:
    ex(ws, f"{col}6", v, h=h_)
ws.row_dimensions[6].height = 30
for k in range(NX):
    rr = X1 + k
    num(ws, f"B{rr}", k + 1)
    inp(ws, f"C{rr}", h="center")
    for col in "DFGHJ":
        inp(ws, f"{col}{rr}")
    inp(ws, f"E{rr}", h="center")
    inp(ws, f"I{rr}", h="center")
    cells = dict(req=f"D{rr}", tipo=f"E{rr}", fonte=f"F{rr}", saida=f"G{rr}", metodo=f"H{rr}", result=f"I{rr}", evid=f"J{rr}")
    calc(ws, f"K{rr}", "=" + f_entrada(cells, f"COUNTA(C{rr},E{rr}:J{rr})"), b=False, sz=9)
    ws.row_dimensions[rr].height = 30
dv_list(ws, f"E{X1}:E{X2}", TIPOS_E, "Cliente, desempenho, legal, norma, projeto anterior, falha potencial ou da organização")
dv_list(ws, f"I{X1}:I{X2}", RES_E, "Resultado da verificação. Deixe vazio enquanto não foi verificada.")
cf_texto(ws, f"I{X1}:I{X2}", f"I{X1}", ENT_CF)
cf_texto(ws, f"K{X1}:K{X2}", f"K{X1}", CONF_EN, resto=YELLOW)
note(ws, "E4", "Passe pelos sete tipos, um por um. O tipo esquecido é a entrada que a validação vai revelar tarde.")
note(ws, "G4", "A especificação, a receita, a instrução ou o critério em que a entrada é atendida.")
s = X2 + 2
band(ws, s, "Resumo automático", "K")
for k, (text, formula) in enumerate([
    ("Entradas escritas", f"=COUNTA(D{X1}:D{X2})"),
    ("Por resultado", f'=IF(D{s+1}=0,"",COUNTIF(I{X1}:I{X2},"{ATENDE}")&" atendem, "&COUNTIF(I{X1}:I{X2},"{NAOAT}")&" não atendem, "&COUNTIF(K{X1}:K{X2},"{PENDENTE}")&" pendentes")'),
    ("Tipos usados", f'=SUMPRODUCT((COUNTIF(E{X1}:E{X2},{{' + ",".join(f'"{t}"' for t in TIPOS_E) + f'}})>0)*1)&" de {len(TIPOS_E)}"'),
    ("Aviso", f'=IF(D{s+1}=0,"Escreva as entradas",IF(COUNTIF(K{X1}:K{X2},"{NAOAC}")>0,"Há entrada que não atende, sem ação",'
              f'IF(SUMPRODUCT((K{X1}:K{X2}<>"")*(K{X1}:K{X2}<>"OK")*(K{X1}:K{X2}<>"{PENDENTE}")*(K{X1}:K{X2}<>"{NAOEM}"))>0,"Há entrada a completar: veja a coluna Conferência",'
              f'IF(COUNTIF(K{X1}:K{X2},"{PENDENTE}")+COUNTIF(K{X1}:K{X2},"{NAOEM}")>0,"Há entrada pendente ou que não atende: não libere ainda","OK"))))'),
], 1):
    label(ws, f"B{s+k}", text, merge=f"B{s+k}:C{s+k}", h="right")
    calc(ws, f"D{s+k}", formula, merge=f"D{s+k}:K{s+k}", sz=9 if text == "Aviso" else 10, b=text != "Aviso", h="left")
    ws.row_dimensions[s + k].height = 21.75
XAV = s + 4
cf_warn(ws, f"D{XAV}:K{XAV}", f"D{XAV}")
ws.freeze_panes = f"E{X1}"
setup(ws, TEAL, f"B1:K{XAV}", fit_height=True)

# ------------------------------------------------------------------ Mudanças
ws = wb.create_sheet("Mudanças")
widths(ws, {"A": 2, "B": 5, "C": 12, "D": 36, "E": 28, "F": 36, "G": 36, "H": 22, "I": 30, "J": 2})
title(ws, "Mudanças de projeto", "Toda alteração em uma entrada ou em uma saída já aprovada. A mudança diz o que precisa ser verificado ou validado de novo.", "I")
for col, text in zip("BCDEFGHI", ["#", "Data", "O que mudou", "Motivo", "Análise do impacto", "O que foi verificado ou validado de novo", "Quem autorizou", "Conferência"]):
    head(ws, f"{col}4", text)
ws.row_dimensions[4].height = 33
hint_row(ws, 5, [("B", ""), ("C", "Data"), ("D", "De quê, para quê"), ("E", "Por que mudar"), ("F", "Efeito nas entradas, nas saídas e no uso"), ("G", "O ensaio, o teste, a validação"),
                 ("H", "Quem tem autoridade"), ("I", "Calculada")])
ex(ws, "B6", "Ex.", h="center")
for col, v, h_, fmt in [("C", date(2027, 7, 8), "center", DATE), ("D", "Resina do fornecedor B no lugar do A.", "left", None), ("E", "Prazo de entrega.", "left", None),
                        ("F", "Fichas técnicas comparadas.", "left", None), ("G", "Queda de dardo no lote de teste: 420 g.", "left", None), ("H", "Gerente de engenharia", "left", None),
                        ("I", "OK", "center", None)]:
    ex(ws, f"{col}6", v, h=h_, fmt=fmt)
ws.row_dimensions[6].height = 30
for k in range(NM):
    rr = M1 + k
    num(ws, f"B{rr}", k + 1)
    inp(ws, f"C{rr}", h="center", fmt=DATE)
    for col in "DEFGH":
        inp(ws, f"{col}{rr}")
    calc(ws, f"I{rr}", "=" + f_mud(f"D{rr}", f"COUNTA(C{rr},E{rr}:H{rr})", f"F{rr}", f"G{rr}", f"H{rr}"), b=False, sz=9)
    ws.row_dimensions[rr].height = 30
dv_date(ws, f"C{M1}:C{M2}")
cf_warn(ws, f"I{M1}:I{M2}", f"I{M1}")
note(ws, "H4", "Quem tem a autoridade para mudar o projeto. Uma mudança autorizada por quem não responde pelo projeto é uma constatação.")
s = M2 + 2
band(ws, s, "Resumo automático", "I")
for k, (text, formula) in enumerate([
    ("Mudanças registradas", f"=COUNTA(D{M1}:D{M2})"),
    ("Mudanças a rever", f'=SUMPRODUCT((I{M1}:I{M2}<>"")*(I{M1}:I{M2}<>"OK"))'),
    ("Aviso", f'=IF(D{s+1}=0,"Nenhuma mudança registrada: confirme se o projeto não mudou",IF(D{s+2}>0,"Há mudança a rever: veja a coluna Conferência","OK"))'),
], 1):
    label(ws, f"B{s+k}", text, merge=f"B{s+k}:C{s+k}", h="right")
    calc(ws, f"D{s+k}", formula, merge=f"D{s+k}:I{s+k}", sz=9 if text == "Aviso" else 10, b=text != "Aviso", h="left")
    ws.row_dimensions[s + k].height = 21.75
MAV = s + 3
cf_warn(ws, f"D{MAV}:I{MAV}", f"D{MAV}", ok_values=("OK", "Nenhuma mudança registrada: confirme se o projeto não mudou"))
ws.freeze_panes = "E7"
setup(ws, AMBER, f"B1:I{MAV}", fit_height=True)

# ------------------------------------------------------------------ Painel
ws = wb.create_sheet("Painel")
widths(ws, {"A": 2, "B": 30, "C": 14, "D": 12, "E": 12, "F": 12, "G": 34, "H": 2})
title(ws, "Painel do projeto", "Nada a preencher: tudo vem das outras abas.", "G")
label(ws, "B3", "Projeto")
calc(ws, "C3", f'=IF({PL}!D5="","",{PL}!D5)', merge="C3:G3", h="left", b=False)
for col, text, m in [("B", "Indicador", None), ("C", "Resultado", None), ("D", "Como se lê", "D5:G5")]:
    put(ws, f"{col}5", text, f=font(10, True, c=WHITE), bg=INK, h="center", merge=m)
ws.row_dimensions[5].height = 21.75
Pc, Pd, Ph, Pi, Pl = (f"{PL}!{c}{E1}:{c}{E2}" for c in "CDHIL")
Xd, Xi, Xk = (f"{EN}!{c}{X1}:{c}{X2}" for c in "DIK")
VALOK = f'COUNTIFS({Pd},"{VALID}",{Pi},"{APROV}")+COUNTIFS({Pd},"{VALID}",{Pi},"{RESS}")'
IND = [
    ("Etapas no plano", f"=COUNTA({Pc})", "Da aba Plano."),
    ("Etapas concluídas", f"=COUNTA({Ph})", "Com data de conclusão."),
    ("Etapas atrasadas", f'=COUNTIF({Pl},"Atrasada")', "Sem conclusão, com o prazo vencido na data da leitura."),
    ("Etapas a rever", f'=SUMPRODUCT(({Pl}<>"")*({Pl}<>"OK")*({Pl}<>"Atrasada"))', "Sem responsável, sem prazo, sem resultado, sem ação."),
    ("Validações aprovadas", "=" + VALOK, "Com ou sem ressalvas. Sem nenhuma, o projeto não pode ser liberado."),
    ("Entradas escritas", f"=COUNTA({Xd})", "Da aba Entradas."),
    ("Entradas que atendem", f'=COUNTIF({Xi},"{ATENDE}")', "Verificadas e conformes."),
    ("Entradas que não atendem", f'=COUNTIF({Xi},"{NAOAT}")', "Com ou sem ação em curso."),
    ("Entradas pendentes", f'=COUNTIF({Xk},"{PENDENTE}")', "Sem resultado de verificação."),
    ("Mudanças registradas", f"=COUNTA({MU}!D{M1}:D{M2})", "Da aba Mudanças."),
    ("Mudanças a rever", f'=SUMPRODUCT(({MU}!I{M1}:I{M2}<>"")*({MU}!I{M1}:I{M2}<>"OK"))', "Sem análise, sem autorização ou sem nova verificação."),
]
for k, (nome, formula, leit) in enumerate(IND):
    rr = 6 + k
    label(ws, f"B{rr}", nome)
    calc(ws, f"C{rr}", formula)
    put(ws, f"D{rr}", leit, f=font(9, c=MUTED), bg=GRAY, merge=f"D{rr}:G{rr}")
    ws.row_dimensions[rr].height = 21.75
IR = {n: 6 + k for k, (n, _, _) in enumerate(IND)}
rr = 6 + len(IND)
label(ws, f"B{rr}", "Pode liberar?")
calc(ws, f"C{rr}", f'=IF(C{IR["Etapas no plano"]}=0,"",IF(AND(C{IR["Validações aprovadas"]}>0,C{IR["Entradas escritas"]}>0,C{IR["Entradas que não atendem"]}+C{IR["Entradas pendentes"]}=0,'
     f'C{IR["Mudanças a rever"]}=0,C{IR["Etapas atrasadas"]}+C{IR["Etapas a rever"]}=0),"Sim","Ainda não"))')
calc(ws, f"D{rr}", f'=IF(C{rr}="Ainda não",IF(C{IR["Validações aprovadas"]}=0,"Falta uma validação aprovada.",IF(C{IR["Entradas que não atendem"]}+C{IR["Entradas pendentes"]}>0,'
     f'"Há entrada que não atende ou sem verificação.",IF(C{IR["Mudanças a rever"]}>0,"Há mudança sem análise, autorização ou nova verificação.","Há etapa atrasada ou a rever."))),'
     f'IF(C{rr}="Sim","As três perguntas do módulo 7 têm resposta sim.",""))', merge=f"D{rr}:G{rr}", b=False, sz=9, h="left")
cf_texto(ws, f"C{rr}", f"C{rr}", [("Sim", GREEN), ("Ainda não", RED)])
ws.row_dimensions[rr].height = 24
LIBR = rr
T0 = rr + 2
band(ws, T0, "Os três controles", "G")
for col, text in zip("BCDEFG", ["Controle", "No plano", "Concluídos", "Aprovados", "Reprovados", "Pergunta"]):
    head(ws, f"{col}{T0+1}", text)
for k, (tipo, perg) in enumerate(((ANAL, "O projeto segue, muda ou para?"), (VERIF, "As saídas atendem às entradas?"), (VALID, "O produto serve para o uso pretendido?"))):
    r2 = T0 + 2 + k
    label(ws, f"B{r2}", tipo)
    calc(ws, f"C{r2}", f'=COUNTIF({Pd},B{r2})', b=False)
    calc(ws, f"D{r2}", f'=COUNTIFS({Pd},B{r2},{Ph},"<>")', b=False)
    calc(ws, f"E{r2}", f'=COUNTIFS({Pd},B{r2},{Pi},"{APROV}")+COUNTIFS({Pd},B{r2},{Pi},"{RESS}")', b=False)
    calc(ws, f"F{r2}", f'=COUNTIFS({Pd},B{r2},{Pi},"{REPROV}")', b=False)
    put(ws, f"G{r2}", perg, f=font(9, c=MUTED), bg=GRAY)
    ws.row_dimensions[r2].height = 21.75
T1 = T0 + 6
band(ws, T1, "Entradas por tipo", "G")
for col, text in zip("BCDEFG", ["Tipo", "Entradas", "Atendem", "Não atendem", "Pendentes", "Leitura"]):
    head(ws, f"{col}{T1+1}", text)
Q1 = T1 + 2
Xe = f"{EN}!E{X1}:E{X2}"
for k, t in enumerate(TIPOS_E):
    r2 = Q1 + k
    label(ws, f"B{r2}", t)
    calc(ws, f"C{r2}", f'=COUNTIF({Xe},B{r2})', b=False)
    calc(ws, f"D{r2}", f'=COUNTIFS({Xe},B{r2},{Xi},"{ATENDE}")', b=False)
    calc(ws, f"E{r2}", f'=COUNTIFS({Xe},B{r2},{Xi},"{NAOAT}")', b=False)
    calc(ws, f"F{r2}", f'=COUNTIFS({Xe},B{r2},{Xk},"{PENDENTE}")', b=False)
    calc(ws, f"G{r2}", f'=IF(C{r2}=0,"Nenhuma entrada deste tipo: confirme se é isso mesmo",IF(E{r2}+F{r2}>0,"Há o que verificar ou resolver",""))', b=False, sz=9, h="left")
    ws.row_dimensions[r2].height = 21.75
cf_texto(ws, f"G{Q1}:G{Q1 + 6}", f"G{Q1}", [("", GRAY)], resto=YELLOW)
s = Q1 + 8
band(ws, s, "Resumo automático", "G")
label(ws, f"B{s+1}", "Aviso")
calc(ws, f"C{s+1}", f'=IF(C{IR["Etapas no plano"]}=0,"Escreva o plano na aba Plano",IF(COUNTIF({Pl},"{SEMVAL}")>0,"Há liberação sem validação aprovada",'
     f'IF(C{T0 + 4}=0,"O plano não tem validação",IF(C{IR["Etapas atrasadas"]}>0,"Há etapa atrasada",IF(C{IR["Mudanças a rever"]}>0,"Há mudança a rever",'
     f'IF(C{IR["Entradas que não atendem"]}+C{IR["Entradas pendentes"]}>0,"Há entrada que não atende ou pendente",IF(C{IR["Etapas a rever"]}>0,"Há etapa a rever","OK")))))))',
     merge=f"C{s+1}:G{s+1}", sz=9, b=False, h="left")
cf_warn(ws, f"C{s+1}:G{s+1}", f"C{s+1}")
ws.row_dimensions[s + 1].height = 21.75
NAV = s + 1
empilhado(ws, f"B{s + 3}", Q1, Q1 + 6, 2, [(4, "Atendem"), (5, "Não atendem"), (6, "Pendentes")], width=20, height=9)
setup(ws, REDC, f"B1:G{s + 22}", landscape=False, fit_height=True)

# ------------------------------------------------------------------ Checklist
ws = wb.create_sheet("Checklist")
widths(ws, {"A": 2, "B": 5, "C": 70, "D": 14, "E": 44, "F": 2})
title(ws, "Checklist do projeto e desenvolvimento", "Escolha o status de cada item na lista suspensa (coluna em amarelo-claro).", "E")
for col, text in zip("BCDE", ["#", "Item de verificação", "Status", "Observações"]):
    head(ws, f"{col}4", text)
ws.row_dimensions[4].height = 21.75
for k, text in enumerate(CHECK):
    rr = 5 + k
    num(ws, f"B{rr}", k + 1)
    put(ws, f"C{rr}", text, bg=GRAY)
    inp(ws, f"D{rr}", h="center")
    inp(ws, f"E{rr}")
    ws.row_dimensions[rr].height = 30
dv_list(ws, "D5:D16", YESNO, "Sim, Parcial ou Não")
cf_equal(ws, "D5:D16", YESNO_CF)
band(ws, 18, "Resumo automático", "E")
summary(ws, 19, "Itens atendidos (Sim)", '=COUNTIF(D5:D16,"Sim")', "C", "D")
summary(ws, 20, "Itens parciais", '=COUNTIF(D5:D16,"Parcial")', "C", "D")
summary(ws, 21, "Itens não atendidos", '=COUNTIF(D5:D16,"Não")', "C", "D")
summary(ws, 22, "Itens sem resposta", "=COUNTBLANK(D5:D16)", "C", "D")
summary(ws, 23, "Percentual atendido", "=D19/ROWS(D5:D16)", "C", "D", fmt="0%")
summary(ws, 24, "Situação", '=IF(D19=ROWS(D5:D16),"Validado",IF(D22>0,"Em andamento","Com pendências"))', "C", "D")
setup(ws, PURPLE, "B1:E24", landscape=False, fit_height=True)


# ------------------------------------------------------------------ Exemplos
def sub(ws, rr, cells, height=30):
    for a, b, t in cells:
        put(ws, f"{a}{rr}", t, f=font(10, True), bg=GRAY, h="center", merge=b and f"{a}{rr}:{b}{rr}")
    ws.row_dimensions[rr].height = height


def exemplo(ws, ex_):
    H = ex_["head"]
    widths(ws, {"A": 2, "B": 6, "C": 24, "D": 24, "E": 17, "F": 22, "G": 22, "H": 20, "I": 12, "J": 18, "K": 26, "L": 20, "M": 28, "N": 2})
    title(ws, "Projeto e desenvolvimento", "Exemplo preenchido, para consulta. Use as abas Plano, Entradas e Mudanças para o seu projeto.", "M")
    rr = 4
    for rot, val in [("Organização", H["org"]), ("Projeto", H["projeto"]), ("Cliente ou usuário", H["cliente"]), ("Objetivo", H["objetivo"]), ("Responsável e aprovação", H["resp"]),
                     ("Período", f'De {H["inicio"]:%d/%m/%Y} a {H["fim"]:%d/%m/%Y}.'), ("Origem", H["origem"])]:
        label(ws, f"B{rr}", rot, merge=f"B{rr}:C{rr}")
        inp(ws, f"D{rr}", val, merge=f"D{rr}:M{rr}", bg=WHITE, h="left")
        ws.row_dimensions[rr].height = alt([(val, 170)], minimo=19.5)
        rr += 1
    label(ws, f"B{rr}", "Data da leitura", merge=f"B{rr}:C{rr}")
    inp(ws, f"D{rr}", H["ref"], bg=WHITE, h="center", fmt=DATE)
    ref = f"$D${rr}"
    rr += 2
    band(ws, rr, "Plano do projeto", "M", color=BLUE)
    rr += 1
    sub(ws, rr, [("B", None, "#"), ("C", "D", "Etapa"), ("E", None, "Tipo"), ("F", None, "Responsável"), ("G", None, "Participantes"), ("H", None, "Prazo"), ("I", None, "Concluída"),
                 ("J", None, "Resultado"), ("K", None, "Ações e ressalvas"), ("L", None, "Registro"), ("M", None, "Conferência")])
    e1 = rr + 1
    e2 = e1 + len(ex_["etps"]) - 1
    tip, res = f"$E${e1}:$E${e2}", f"$J${e1}:$J${e2}"
    for e in ex_["etps"]:
        rr += 1
        num(ws, f"B{rr}", e["n"])
        put(ws, f"C{rr}", e["etapa"], f=font(10, True), merge=f"C{rr}:D{rr}")
        put(ws, f"E{rr}", e["tipo"], h="center")
        put(ws, f"F{rr}", e["resp"])
        put(ws, f"G{rr}", e["part"])
        put(ws, f"H{rr}", e["prazo"], h="center", fmt=DATE)
        put(ws, f"I{rr}", e["concl"], h="center", fmt=DATE)
        put(ws, f"J{rr}", e["result"] or None, h="center")
        put(ws, f"K{rr}", e["acoes"] or None)
        put(ws, f"L{rr}", e["registro"] or None)
        cells = dict(etapa=f"C{rr}", tipo=f"E{rr}", resp=f"F{rr}", prazo=f"H{rr}", concl=f"I{rr}", result=f"J{rr}", acoes=f"K{rr}")
        calc(ws, f"M{rr}", "=" + f_etapa(cells, "1", tip, res, ref), b=False, sz=9)
        ws.row_dimensions[rr].height = alt([(e["etapa"], 46), (e["acoes"], 24), (e["part"], 20), (e["registro"], 18), (e["resp"], 20)], minimo=21.75)
    cf_texto(ws, f"J{e1}:J{e2}", f"J{e1}", RES_CF)
    cf_texto(ws, f"M{e1}:M{e2}", f"M{e1}", CONF_ET, resto=YELLOW)
    rr += 2
    ws.row_breaks.append(Break(id=rr - 1))
    band(ws, rr, "Entradas, saídas e verificação", "M", color=TEAL)
    rr += 1
    sub(ws, rr, [("B", None, "#"), ("C", "D", "Entrada"), ("E", None, "Tipo"), ("F", None, "Fonte"), ("G", None, "Saída que atende"), ("H", None, "Como verificar"), ("I", None, "Resultado"),
                 ("J", "L", "Evidência ou ação"), ("M", None, "Conferência")])
    x1 = rr + 1
    for x in ex_["ents"]:
        rr += 1
        put(ws, f"B{rr}", x["id"], f=font(10, True, c=MUTED), h="center")
        put(ws, f"C{rr}", x["req"], f=font(10, True), merge=f"C{rr}:D{rr}")
        put(ws, f"E{rr}", x["tipo"], h="center")
        put(ws, f"F{rr}", x["fonte"])
        put(ws, f"G{rr}", x["saida"])
        put(ws, f"H{rr}", x["metodo"])
        put(ws, f"I{rr}", x["result"] or None, h="center")
        put(ws, f"J{rr}", x["evid"] or None, merge=f"J{rr}:L{rr}")
        cells = dict(req=f"C{rr}", tipo=f"E{rr}", fonte=f"F{rr}", saida=f"G{rr}", metodo=f"H{rr}", result=f"I{rr}", evid=f"J{rr}")
        calc(ws, f"M{rr}", "=" + f_entrada(cells, "1"), b=False, sz=9)
        ws.row_dimensions[rr].height = alt([(x["req"], 46), (x["fonte"], 20), (x["saida"], 20), (x["metodo"], 18), (x["tipo"], 16), (x["evid"], 60)], minimo=21.75)
    x2 = rr
    cf_texto(ws, f"I{x1}:I{x2}", f"I{x1}", ENT_CF)
    cf_texto(ws, f"M{x1}:M{x2}", f"M{x1}", CONF_EN, resto=YELLOW)
    rr += 2
    band(ws, rr, "Mudanças de projeto", "M", color=AMBER)
    rr += 1
    sub(ws, rr, [("B", None, "Data"), ("C", "D", "O que mudou"), ("E", "F", "Motivo"), ("G", "H", "Análise do impacto"), ("I", "K", "O que foi verificado de novo"), ("L", None, "Quem autorizou"),
                 ("M", None, "Conferência")])
    m1 = rr + 1
    for m in ex_["muds"]:
        rr += 1
        put(ws, f"B{rr}", m["data"], h="center", fmt="dd/mm")
        put(ws, f"C{rr}", m["oque"], f=font(10, True), merge=f"C{rr}:D{rr}")
        put(ws, f"E{rr}", m["motivo"], merge=f"E{rr}:F{rr}")
        put(ws, f"G{rr}", m["analise"] or None, merge=f"G{rr}:H{rr}")
        put(ws, f"I{rr}", m["verif"] or None, merge=f"I{rr}:K{rr}")
        put(ws, f"L{rr}", m["aut"])
        calc(ws, f"M{rr}", "=" + f_mud(f"C{rr}", "1", f"G{rr}", f"I{rr}", f"L{rr}"), b=False, sz=9)
        ws.row_dimensions[rr].height = alt([(m["oque"], 46), (m["motivo"], 38), (m["analise"], 40), (m["verif"], 50)], minimo=21.75)
    m2 = rr
    cf_warn(ws, f"M{m1}:M{m2}", f"M{m1}")
    rr += 2
    band(ws, rr, "Resumo automático", "M")
    vok = f'COUNTIFS({tip},"{VALID}",{res},"{APROV}")+COUNTIFS({tip},"{VALID}",{res},"{RESS}")'
    pend = f'(COUNTIF(I{x1}:I{x2},"{NAOAT}")+COUNTIF(M{x1}:M{x2},"{PENDENTE}"))'
    for k, (text, formula) in enumerate([
        ("Etapas: concluídas, atrasadas", f'=COUNTA(I{e1}:I{e2})&" de "&COUNTA(C{e1}:C{e2})&" concluídas, "&COUNTIF(M{e1}:M{e2},"Atrasada")&" atrasadas"'),
        ("Validações aprovadas", "=" + vok),
        ("Entradas: atendem, não atendem, pendentes", f'=COUNTIF(I{x1}:I{x2},"{ATENDE}")&" atendem, "&COUNTIF(I{x1}:I{x2},"{NAOAT}")&" não atendem, "&COUNTIF(M{x1}:M{x2},"{PENDENTE}")&" pendentes"'),
        ("Mudanças a rever", f'=SUMPRODUCT((M{m1}:M{m2}<>"OK")*1)'),
        ("Pode liberar?", f'=IF(AND({vok}>0,{pend}=0,SUMPRODUCT((M{m1}:M{m2}<>"OK")*1)=0,SUMPRODUCT((M{e1}:M{e2}<>"OK")*1)=0),"Sim","Ainda não")'),
    ], 1):
        label(ws, f"B{rr+k}", text, merge=f"B{rr+k}:E{rr+k}", h="right")
        calc(ws, f"F{rr+k}", formula, merge=f"F{rr+k}:M{rr+k}", h="left")
        ws.row_dimensions[rr + k].height = 21.75
    cf_texto(ws, f"F{rr + 5}:M{rr + 5}", f"F{rr + 5}", [("Sim", GREEN), ("Ainda não", RED)])
    setup(ws, MUTED, f"B1:M{rr + 5}")


exemplo(wb.create_sheet("Exemplo 1 - Pizzaria"), EX1)
exemplo(wb.create_sheet("Exemplo 2 - Indústria"), EX2)

wb.active = 1
wb.save(OUT)
print("ok", OUT, "| avisos: Plano D%d, Entradas D%d, Mudanças D%d, Painel C%d, liberar C%d" % (PAV, XAV, MAV, NAV, LIBR))
