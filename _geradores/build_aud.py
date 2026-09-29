# -*- coding: utf-8 -*-
"""Gera Auditoria-modelo.xlsx no mesmo padrão visual das planilhas anteriores da série."""
import os
import sys
from datetime import date, time

_here = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, _here)
# funções de formatação compartilhadas com o gerador do PDCA (bloco inicial do arquivo)
_src = open(os.path.join(_here, "build_pdca.py"), encoding="utf-8").read()
exec(_src[: _src.index("STATUS = [")])

from openpyxl.chart import BarChart, Reference  # noqa: E402
from openpyxl.chart.label import DataLabelList  # noqa: E402
from openpyxl.chart.series import DataPoint  # noqa: E402
from aud_data import EX1, EX2  # noqa: E402

BLUE, TEAL, AMBER, PURPLE, REDC = "2B5C8A", "1E7B73", "A96A12", "7A4A9A", "B0413E"
TEAL_T, AMBER_T, RED_T = "D9EEEB", "F6E8CF", "F5DEDC"
YESNO = ["Sim", "Parcial", "Não"]
YESNO_CF = [("Sim", GREEN), ("Parcial", YELLOW), ("Não", RED)]
PRIO_CF = [("Alta", RED), ("Média", YELLOW), ("Baixa", GREEN)]
RES = ["Conforme", "Não conforme", "Oportunidade de melhoria", "Não aplicável"]
RES_CF = [("Conforme", GREEN), ("Não conforme", RED), ("Oportunidade de melhoria", YELLOW), ("Não aplicável", "D9DEE2")]
TIPOS = ["NC maior", "NC menor", "Oportunidade de melhoria"]
TIPO_CF = [("NC maior", RED), ("NC menor", RED), ("Oportunidade de melhoria", YELLOW)]
HORA = "hh:mm"


def note(ws, ref, text):
    c = Comment(text, "Modelo Auditoria")
    c.width, c.height = 280, 110
    ws[ref].comment = c


def cf_warn(ws, rng_, first, ok_values=("OK",)):
    """Verde para os valores aceitos; amarelo para qualquer outro texto."""
    cond = ",".join(f'{first}="{v}"' for v in ok_values)
    ws.conditional_formatting.add(rng_, FormulaRule(formula=[f"OR({cond})"], fill=PatternFill("solid", bgColor=GREEN, fgColor=GREEN)))
    ws.conditional_formatting.add(rng_, FormulaRule(formula=[f"AND(LEN({first})>0,NOT(OR({cond})))"],
                                  fill=PatternFill("solid", bgColor=YELLOW, fgColor=YELLOW)))


def dv_time(ws, rng_):
    dv = DataValidation(type="time", operator="between", formula1="0", formula2="0.999988", allow_blank=True)
    dv.errorTitle, dv.error = "Horário inválido", "Digite o horário no formato hora:minuto, como 08:30."
    dv.showErrorMessage = True
    ws.add_data_validation(dv)
    dv.add(rng_)


def hint_row(ws, row, cells, height=30):
    for ref, text, merge in cells:
        put(ws, f"{ref}{row}", text, f=font(9, i=True, c=MUTED), bg=GRAY, h="center", merge=merge and f"{ref}{row}:{merge}{row}")
    ws.row_dimensions[row].height = height


# ================================================================== pasta de trabalho
wb = Workbook()
wb.properties.title = "Auditoria interna — Modelo de programa, plano e constatações"
wb.properties.creator = "Modelo Auditoria"
wb.properties.language = "pt-BR"

# ------------------------------------------------------------------ Instruções
ws = wb.active
ws.title = "Instruções"
widths(ws, {"A": 2, "B": 24, "C": 92, "D": 2})
title(ws, "Auditoria interna — Como usar esta planilha",
      "Modelo para programar as auditorias do ano, planejar cada uma e registrar as constatações.", "C")
r = 4


def section(text):
    global r
    band(ws, r, text, "C", sz=11)
    r += 1


def line(k, v, kbg=GRAY, vbg=None, kf=None, kh="left", height=19.5):
    global r
    put(ws, f"B{r}", k, f=kf or font(10, True), bg=kbg, h=kh)
    put(ws, f"C{r}", v, bg=vbg)
    ws.row_dimensions[r].height = height
    r += 1


section("Legenda: onde preencher")
line("Amarelo-claro", "Células de entrada. É aqui que você digita.", vbg=INPUT)
line("Cinza", "Células calculadas ou fixas (rótulos, contagens, avisos). Não altere.", vbg=GRAY)
line("Cinza em itálico", "Linha de exemplo (marcada com “Ex.”). Mostra o formato esperado e não entra nas contagens.", vbg=GRAY)
line("Listas suspensas", "Colunas de importância, resultado, tipo e status aceitam apenas as opções da lista. Datas e horários são conferidos na digitação.", height=31.5)
line("Comentários", "Células com triângulo vermelho no canto trazem uma dica de preenchimento. Passe o mouse sobre elas.", height=31.5)
r += 1
section("Os tipos de constatação")
line("C", "Conformidade: a evidência mostra que o requisito é atendido.", kbg=TEAL, vbg=TEAL_T, kf=font(10, True, c=WHITE), kh="center")
line("NC", "Não conformidade: a evidência mostra que um requisito não é atendido. Pode ser classificada como maior ou menor.",
     kbg=REDC, vbg=RED_T, kf=font(10, True, c=WHITE), kh="center", height=31.5)
line("OM", "Oportunidade de melhoria: o requisito é atendido, mas há uma forma melhor de atender.", kbg=AMBER, vbg=AMBER_T,
     kf=font(10, True, c=WHITE), kh="center")
line("Três partes", "Toda não conformidade tem requisito (o que foi combinado), evidência (o que foi encontrado) e declaração (o desvio, em uma frase).", height=31.5)
r += 1
section("Ordem recomendada de preenchimento")
for k, text in enumerate([
    "Aba Programa: liste os processos, avalie os três fatores de prioridade e defina a data e o auditor líder de cada auditoria.",
    "Aba Plano: para cada auditoria, registre o objetivo, o escopo, os critérios, a equipe e a agenda.",
    "Aba Lista de verificação: escreva o que será conferido em cada requisito e o tamanho da amostra.",
    "Durante a auditoria, registre na lista a evidência encontrada e o resultado de cada item.",
    "Aba Constatações: escreva cada não conformidade com as três partes e cada oportunidade de melhoria com a evidência.",
    "Aba Checklist: valide a auditoria antes de distribuir o relatório.",
    "Aba Constatações: acompanhe o status até a verificação da eficácia. Aba Programa: registre a auditoria como realizada.",
], 1):
    line(f"Passo {k}", text, height=31.5 if len(text) > 95 else 19.5)
r += 1
section("Abas da planilha")
for k, text in [
    ("Programa", "Programa anual de auditoria, com prioridade calculada, conferência de independência e cumprimento."),
    ("Plano", "Plano de uma auditoria: objetivo, escopo, critérios, equipe e agenda, com o tempo de cada atividade."),
    ("Lista de verificação", "Roteiro do auditor, com requisito, amostra, evidência encontrada e resultado."),
    ("Constatações", "Registro das constatações, com conferência das três partes, situação e gráfico por tipo."),
    ("Checklist", "Doze verificações de qualidade da auditoria, com percentual de conclusão."),
    ("Exemplo 1 - Pizzaria", "Auditoria de um processo contra os padrões da própria organização."),
    ("Exemplo 2 - Compras", "Auditoria de um processo contra um procedimento interno e um requisito da ISO 9001."),
]:
    line(k, text)
r += 1
section("Regras e premissas do modelo")
for k, text in [
    ("Prioridade", "Pontos = importância (alta 3, média 2, baixa 1) + mudanças recentes (sim 2) + resultado anterior (NC maior 3, NC menor 2, não auditado 2). Com 6 ou mais, a prioridade é alta; de 3 a 5, média. A pontuação é uma convenção deste modelo."),
    ("Frequência", "Prioridade alta: a cada 6 meses. Média: a cada 12. Baixa: a cada 18. É uma sugestão: a organização define as próprias frequências."),
    ("Independência", "A conferência compara o nome do auditor líder com o do dono do processo. Ela não identifica outros conflitos, como subordinação ou participação recente na atividade."),
    ("Uma auditoria por vez", "As abas Plano, Lista de verificação e Constatações servem a uma auditoria. Para a auditoria seguinte, salve uma cópia do arquivo ou duplique as três abas."),
    ("Situação", "Uma auditoria ou uma constatação é considerada atrasada quando a data planejada ou o prazo é anterior à data de hoje e ela ainda não foi concluída."),
    ("Exemplos", "Os dados das abas de exemplo são ilustrativos, criados para este material."),
]:
    line(k, text, height=45.75 if len(text) > 100 else 31.5)
setup(ws, INK, f"B1:C{r}", landscape=False, fit_height=True)

# ------------------------------------------------------------------ Programa
ws = wb.create_sheet("Programa")
widths(ws, {"A": 2, "B": 5, "C": 18, "D": 16, "E": 22, "F": 13, "G": 11, "H": 16, "I": 9, "J": 12, "K": 21, "L": 13, "M": 24,
            "N": 22, "O": 14, "P": 13, "Q": 20, "R": 2})
title(ws, "Programa de auditoria", "Preencha as células em amarelo-claro. As células cinza são calculadas automaticamente.", "Q")
for rr, l1, l2, f2 in [(4, "Organização ou unidade", "Gestor do programa", None), (5, "Período do programa", "Data de aprovação", DATE),
                       (6, "Elaborado por", "Versão", None)]:
    label(ws, f"B{rr}", l1, merge=f"B{rr}:C{rr}")
    inp(ws, f"D{rr}", merge=f"D{rr}:H{rr}")
    label(ws, f"I{rr}", l2, merge=f"I{rr}:K{rr}")
    inp(ws, f"L{rr}", merge=f"L{rr}:Q{rr}", fmt=f2)
    ws.row_dimensions[rr].height = 21.75
for rr, l1 in [(7, "Objetivos do programa"), (8, "Critérios de auditoria")]:
    label(ws, f"B{rr}", l1, merge=f"B{rr}:C{rr}")
    inp(ws, f"D{rr}", merge=f"D{rr}:Q{rr}")
    ws.row_dimensions[rr].height = 27
dv_date(ws, "L5")
note(ws, "D5", 'Período coberto pelo programa.\nEx.: "Janeiro a dezembro de 2027".')
note(ws, "D7", 'O que o programa quer alcançar.\nEx.: "Verificar a conformidade com a ISO 9001 e a eficácia das ações do ciclo anterior".')
note(ws, "D8", 'Normas e documentos usados como referência.\nEx.: "ISO 9001:2015 e procedimentos do sistema de gestão".')
put(ws, "B10", "Processos", f=font(10, True, c=WHITE), bg=INK, box=False, merge="B10:E10")
put(ws, "F10", "Fatores de prioridade", f=font(10, True, c=WHITE), bg=BLUE, box=False, merge="F10:H10")
put(ws, "I10", "Prioridade calculada", f=font(10, True, c=WHITE), bg=INK, box=False, merge="I10:K10")
put(ws, "L10", "Planejamento", f=font(10, True, c=WHITE), bg=BLUE, box=False, merge="L10:N10")
put(ws, "O10", "Cumprimento", f=font(10, True, c=WHITE), bg=INK, box=False, merge="O10:Q10")
ws.row_dimensions[10].height = 21.75
head(ws, "B11", "#")
put(ws, "C11", "Processo", f=font(10, True, c=WHITE), bg=INK, h="center", merge="C11:D11")
for col, text in zip("EFGHIJKLMNOPQ", ["Dono do processo", "Importância", "Mudanças recentes", "Resultado anterior", "Pontos", "Prioridade",
                                       "Frequência sugerida", "Data planejada", "Auditor líder", "Independência", "Status", "Realizada em", "Situação"]):
    head(ws, f"{col}11", text)
ws.row_dimensions[11].height = 30
hint_row(ws, 12, [("B", "", None), ("C", "Nome do processo", "D"), ("E", "Cargo de quem responde", None), ("F", "Para o cliente", None),
                  ("G", "Sim ou Não", None), ("H", "Última auditoria", None), ("I", "Soma", None), ("J", "Faixa", None), ("K", "Pela prioridade", None),
                  ("L", "Data", None), ("M", "Cargo ou nome", None), ("N", "Auditor e dono", None), ("O", "Da auditoria", None),
                  ("P", "Data", None), ("Q", "Calculada", None)])
ex(ws, "B13", "Ex.", h="center")
put(ws, "C13", "Produção", f=font(9, i=True, c=MUTED), bg=GRAY, merge="C13:D13")
for col, v, h_, f_ in [("E", "Gerente industrial", "left", None), ("F", "Alta", "center", None), ("G", "Sim", "center", None),
                       ("H", "NC maior", "center", None), ("I", 8, "center", None), ("J", "Alta", "center", None),
                       ("K", "A cada 6 meses", "center", None), ("L", date(2026, 10, 20), "center", DATE),
                       ("M", "Coordenador da Qualidade", "left", None), ("N", "OK", "center", None), ("O", "Planejada", "center", None),
                       ("P", None, "center", None), ("Q", "No prazo", "center", None)]:
    ex(ws, f"{col}13", v if v is not None else "", h=h_, fmt=f_)
ws.row_dimensions[13].height = 27
G1, G2 = 14, 25
for k in range(12):
    rr = G1 + k
    num(ws, f"B{rr}", k + 1)
    inp(ws, f"C{rr}", merge=f"C{rr}:D{rr}")
    inp(ws, f"E{rr}")
    for col in "FGH":
        inp(ws, f"{col}{rr}", h="center")
    calc(ws, f"I{rr}", f'=IF(OR(C{rr}="",F{rr}="",G{rr}="",H{rr}=""),"",IF(F{rr}="Alta",3,IF(F{rr}="Média",2,1))+IF(G{rr}="Sim",2,0)'
         f'+IF(H{rr}="NC maior",3,IF(OR(H{rr}="NC menor",H{rr}="Não auditado"),2,0)))')
    calc(ws, f"J{rr}", f'=IF(I{rr}="","",IF(I{rr}>=6,"Alta",IF(I{rr}>=3,"Média","Baixa")))', b=False)
    calc(ws, f"K{rr}", f'=IF(J{rr}="","",IF(J{rr}="Alta","A cada 6 meses",IF(J{rr}="Média","A cada 12 meses","A cada 18 meses")))', b=False, sz=9)
    inp(ws, f"L{rr}", h="center", fmt=DATE)
    inp(ws, f"M{rr}")
    calc(ws, f"N{rr}", f'=IF(OR(M{rr}="",E{rr}=""),"",IF(LOWER(TRIM(M{rr}))=LOWER(TRIM(E{rr})),"Conflito: auditor é o dono","OK"))', b=False, sz=9)
    inp(ws, f"O{rr}", h="center")
    inp(ws, f"P{rr}", h="center", fmt=DATE)
    calc(ws, f"Q{rr}", f'=IF(C{rr}="","",IF(O{rr}="Realizada",IF(AND(P{rr}<>"",L{rr}<>"",P{rr}>L{rr}+30),"Realizada com atraso","Realizada"),'
         f'IF(O{rr}="Cancelada","Cancelada",IF(L{rr}="","Sem data",IF(L{rr}<TODAY(),"Atrasada","No prazo")))))', b=False, sz=9)
    ws.row_dimensions[rr].height = 27
dv_list(ws, f"F{G1}:F{G2}", ["Alta", "Média", "Baixa"], "Alta, Média ou Baixa")
dv_list(ws, f"G{G1}:G{G2}", ["Sim", "Não"], "Sim ou Não")
dv_list(ws, f"H{G1}:H{G2}", ["NC maior", "NC menor", "Não auditado", "Sem NC"], "Resultado da última auditoria do processo")
dv_list(ws, f"O{G1}:O{G2}", ["Planejada", "Realizada", "Adiada", "Cancelada"], "Planejada, Realizada, Adiada ou Cancelada")
dv_date(ws, f"L{G1}:L{G2}")
dv_date(ws, f"P{G1}:P{G2}")
cf_equal(ws, f"F{G1}:F{G2}", PRIO_CF)
cf_equal(ws, f"J{G1}:J{G2}", PRIO_CF)
cf_warn(ws, f"N{G1}:N{G2}", f"N{G1}")
cf_equal(ws, f"Q{G1}:Q{G2}", [("Realizada", GREEN), ("No prazo", GREEN), ("Realizada com atraso", YELLOW), ("Sem data", YELLOW), ("Atrasada", RED)])
note(ws, "F11", "Quanto o processo afeta o cliente e a conformidade do produto ou do serviço.")
note(ws, "G11", "Houve mudança de pessoas, de método, de sistema ou de volume desde a última auditoria?")
note(ws, "N11", "A conferência compara o auditor líder com o dono do processo. Quem responde pelo processo não pode auditá-lo.")
note(ws, "Q11", "A auditoria aparece como realizada com atraso quando acontece mais de 30 dias depois da data planejada.")
s = G2 + 2
band(ws, s, "Resumo automático", "Q")
Q = f"Q{G1}:Q{G2}"
for k, (text, formula, fmt) in enumerate([
    ("Processos no programa", f"=COUNTA(C{G1}:C{G2})", None),
    ("Processos de prioridade alta", f'=COUNTIF(J{G1}:J{G2},"Alta")', None),
    ("Auditorias realizadas", f'=COUNTIF({Q},"Realizada")+COUNTIF({Q},"Realizada com atraso")', None),
    ("Auditorias atrasadas", f'=COUNTIF({Q},"Atrasada")', None),
    ("Cumprimento do programa", f'=IF(I{s+1}-COUNTIF({Q},"Cancelada")<=0,0,I{s+3}/(I{s+1}-COUNTIF({Q},"Cancelada")))', "0%"),
    ("Conflitos de independência", f'=COUNTIF(N{G1}:N{G2},"Conflito*")', None),
    ("Aviso", f'=IF(I{s+1}=0,"Liste os processos",IF(I{s+6}>0,"Há auditor líder que é dono do processo auditado",'
              f'IF(COUNT(I{G1}:I{G2})<I{s+1},"Há processo sem os três fatores de prioridade",IF(I{s+4}>0,"Há auditoria atrasada","OK"))))', None),
], 1):
    label(ws, f"B{s+k}", text, merge=f"B{s+k}:H{s+k}", h="right")
    calc(ws, f"I{s+k}", formula, fmt=fmt, merge=f"I{s+k}:M{s+k}", sz=9 if text == "Aviso" else 10, b=text != "Aviso")
    ws.row_dimensions[s + k].height = 21.75
cf_warn(ws, f"I{s+7}:M{s+7}", f"I{s+7}")
ws.freeze_panes = "E12"
setup(ws, BLUE, f"B1:Q{s+7}", fit_height=True)

# ------------------------------------------------------------------ Plano
ws = wb.create_sheet("Plano")
widths(ws, {"A": 2, "B": 5, "C": 11, "D": 11, "E": 40, "F": 26, "G": 24, "H": 26, "I": 20, "J": 12, "K": 2})
title(ws, "Plano de auditoria", "Um plano por auditoria. Envie ao auditado antes da data combinada.", "J")
for rr, l1, l2, f2 in [(4, "Auditoria nº", "Data da auditoria", DATE), (5, "Processo ou área", "Auditado", None),
                       (6, "Auditor líder", "Equipe auditora", None)]:
    label(ws, f"B{rr}", l1, merge=f"B{rr}:D{rr}")
    inp(ws, f"E{rr}", merge=f"E{rr}:F{rr}")
    label(ws, f"G{rr}", l2)
    inp(ws, f"H{rr}", merge=f"H{rr}:J{rr}", fmt=f2)
    ws.row_dimensions[rr].height = 21.75
for rr, l1 in [(7, "Objetivo"), (8, "Escopo"), (9, "Critérios")]:
    label(ws, f"B{rr}", l1, merge=f"B{rr}:D{rr}")
    inp(ws, f"E{rr}", merge=f"E{rr}:J{rr}")
    ws.row_dimensions[rr].height = 30
dv_date(ws, "H4")
note(ws, "E7", 'O que a auditoria quer saber.\nEx.: "Verificar a conformidade do processo com o procedimento e com o requisito 8.4".')
note(ws, "E8", 'Limites da auditoria: locais, atividades e período.\nEx.: "Requisições e pedidos de janeiro a agosto".')
note(ws, "E9", 'Documentos usados como referência.\nEx.: "ISO 9001:2015, 8.4. Procedimento PR-SUP-01 rev. 5".')
for col, text in zip("BCDEFGHIJ", ["#", "Início", "Fim", "Atividade", "Tipo", "Auditor", "Auditado", "Local", "Duração"]):
    head(ws, f"{col}11", text)
ws.row_dimensions[11].height = 21.75
ex(ws, "B12", "Ex.", h="center")
ex(ws, "C12", time(9, 0), h="center", fmt=HORA)
ex(ws, "D12", time(10, 30), h="center", fmt=HORA)
ex(ws, "E12", "Requisição e análise: amostra de 10 requisições")
ex(ws, "F12", "Coleta de evidências", h="center")
ex(ws, "G12", "Analista da Qualidade")
ex(ws, "H12", "Compradores")
ex(ws, "I12", "Suprimentos")
ex(ws, "J12", time(1, 30), h="center", fmt="[h]:mm")
ws.row_dimensions[12].height = 27
L1, L2 = 13, 24
for k in range(12):
    rr = L1 + k
    num(ws, f"B{rr}", k + 1)
    inp(ws, f"C{rr}", h="center", fmt=HORA)
    inp(ws, f"D{rr}", h="center", fmt=HORA)
    inp(ws, f"E{rr}")
    inp(ws, f"F{rr}", h="center")
    for col in "GHI":
        inp(ws, f"{col}{rr}")
    calc(ws, f"J{rr}", f'=IF(OR(C{rr}="",D{rr}=""),"",IF(D{rr}<C{rr},"Horário invertido",D{rr}-C{rr}))', b=False, fmt="[h]:mm")
    ws.row_dimensions[rr].height = 27
TP = ["Reunião de abertura", "Coleta de evidências", "Reunião da equipe", "Reunião de encerramento", "Intervalo"]
dv_list(ws, f"F{L1}:F{L2}", TP, "Tipo da atividade")
dv_time(ws, f"C{L1}:D{L2}")
cf_equal(ws, f"J{L1}:J{L2}", [("Horário invertido", RED)])
s = L2 + 2
band(ws, s, "Resumo automático", "J")
F, J = f"F{L1}:F{L2}", f"J{L1}:J{L2}"
for k, (text, formula, fmt) in enumerate([
    ("Atividades na agenda", f"=COUNTA(E{L1}:E{L2})", None),
    ("Tempo de auditoria, sem os intervalos", f'=SUM({J})-SUMIF({F},"Intervalo",{J})', "[h]:mm"),
    ("Tempo de coleta de evidências", f'=SUMIF({F},"Coleta de evidências",{J})', "[h]:mm"),
    ("Parte do tempo dedicada à coleta", f"=IF(F{s+2}=0,0,F{s+3}/F{s+2})", "0%"),
    ("Aviso: reuniões", f'=IF(F{s+1}=0,"Monte a agenda",IF(OR(COUNTIF({F},"Reunião de abertura")=0,COUNTIF({F},"Reunião de encerramento")=0),'
                        f'"Inclua as reuniões de abertura e de encerramento","OK"))', None),
    ("Aviso: coleta", f'=IF(F{s+2}=0,"Informe os horários",IF(F{s+4}<0.5,"Menos da metade do tempo está na coleta de evidências","OK"))', None),
    ("Campos do cabeçalho preenchidos", "=COUNTA(E4,H4,E5,H5,E6,H6,E7,E8,E9)", '0" de 9"'),
], 1):
    label(ws, f"B{s+k}", text, merge=f"B{s+k}:E{s+k}", h="right")
    calc(ws, f"F{s+k}", formula, fmt=fmt, merge=f"F{s+k}:H{s+k}", sz=9 if text.startswith("Aviso") else 10, b=not text.startswith("Aviso"))
    ws.row_dimensions[s + k].height = 21.75
cf_warn(ws, f"F{s+5}:H{s+6}", f"F{s+5}")
setup(ws, TEAL, f"B1:J{s+7}", fit_height=True)

# ------------------------------------------------------------------ Lista de verificação
ws = wb.create_sheet("Lista de verificação")
widths(ws, {"A": 2, "B": 5, "C": 20, "D": 36, "E": 26, "F": 40, "G": 24, "H": 13, "I": 2})
title(ws, "Lista de verificação", "Escreva os itens antes da auditoria. Durante a auditoria, registre a evidência e o resultado.", "H")
label(ws, "B4", "Auditoria", merge="B4:C4")
calc(ws, "D4", '=IF(Plano!E4="","",Plano!E4)&IF(Plano!E5="",""," · "&Plano!E5)', h="left", b=False, merge="D4:E4")
label(ws, "F4", "Data da auditoria")
calc(ws, "G4", '=IF(Plano!H4="","",Plano!H4)', h="left", b=False, fmt=DATE, merge="G4:H4")
ws.row_dimensions[4].height = 21.75
for col, text in zip("BCDEFGH", ["#", "Requisito", "O que verificar", "Amostra", "Evidência encontrada", "Resultado", "Nº da constatação"]):
    head(ws, f"{col}6", text)
ws.row_dimensions[6].height = 30
hint_row(ws, 7, [("B", "", None), ("C", "Norma ou documento, com o item", None), ("D", "Pergunta aberta", None),
                 ("E", "Quantos, de qual período", None), ("F", "Documento, data, local e números", None),
                 ("G", "Depois de comparar", None), ("H", "Aba Constatações", None)])
ex(ws, "B8", "Ex.", h="center")
ex(ws, "C8", "PR-SUP-01, item 4.2")
ex(ws, "D8", "A requisição traz a especificação técnica do item?")
ex(ws, "E8", "10 requisições, de janeiro a agosto")
ex(ws, "F8", "4 requisições sem especificação técnica: RC-0412, RC-0433, RC-0457 e RC-0461")
ex(ws, "G8", "Não conforme", h="center")
ex(ws, "H8", 1, h="center")
ws.row_dimensions[8].height = 33
V1, V2 = 9, 28
for k in range(20):
    rr = V1 + k
    num(ws, f"B{rr}", k + 1)
    for col in "CDEF":
        inp(ws, f"{col}{rr}")
    inp(ws, f"G{rr}", h="center")
    inp(ws, f"H{rr}", h="center")
    ws.row_dimensions[rr].height = 33
dv_list(ws, f"G{V1}:G{V2}", RES, "Conforme, Não conforme, Oportunidade de melhoria ou Não aplicável")
dv_number(ws, f"H{V1}:H{V2}")
cf_equal(ws, f"G{V1}:G{V2}", RES_CF)
s = V2 + 2
band(ws, s, "Resumo automático", "H")
G = f"G{V1}:G{V2}"
for k, (text, formula, fmt) in enumerate([
    ("Itens da lista", f"=COUNTA(D{V1}:D{V2})", None),
    ("Itens com resultado", f"=COUNTA({G})", None),
    ("Conformes", f'=COUNTIF({G},"Conforme")', None),
    ("Não conformes", f'=COUNTIF({G},"Não conforme")', None),
    ("Oportunidades de melhoria", f'=COUNTIF({G},"Oportunidade de melhoria")', None),
    ("Percentual verificado", f"=IF(F{s+1}=0,0,F{s+2}/F{s+1})", "0%"),
    ("Aviso: evidências", f'=IF(SUMPRODUCT(({G}<>"")*({G}<>"Não aplicável")*(F{V1}:F{V2}=""))>0,"Há item com resultado e sem evidência registrada","OK")', None),
    ("Aviso: constatações", f'=IF(SUMPRODUCT((({G}="Não conforme")+({G}="Oportunidade de melhoria"))*(H{V1}:H{V2}=""))>0,'
                            f'"Há item sem o número da constatação","OK")', None),
], 1):
    label(ws, f"B{s+k}", text, merge=f"B{s+k}:E{s+k}", h="right")
    calc(ws, f"F{s+k}", formula, fmt=fmt, sz=9 if text.startswith("Aviso") else 10, b=not text.startswith("Aviso"))
    ws.row_dimensions[s + k].height = 21.75
cf_warn(ws, f"F{s+7}:F{s+8}", f"F{s+7}")
ws.freeze_panes = "C7"
setup(ws, AMBER, f"B1:H{s+8}")
ws.print_title_rows = "6:6"

# ------------------------------------------------------------------ Constatações
ws = wb.create_sheet("Constatações")
widths(ws, {"A": 2, "B": 5, "C": 22, "D": 20, "E": 34, "F": 36, "G": 36, "H": 20, "I": 13, "J": 19, "K": 20, "L": 14, "M": 2})
title(ws, "Constatações da auditoria", "Uma linha por constatação. Toda não conformidade tem requisito, evidência e declaração.", "L")
for col, text in zip("BCDEFGHIJKL", ["Nº", "Tipo", "Processo", "Requisito", "Evidência", "Declaração", "Responsável", "Prazo", "Status",
                                      "Conferência", "Situação"]):
    head(ws, f"{col}4", text)
ws.row_dimensions[4].height = 21.75
hint_row(ws, 5, [("B", "", None), ("C", "Classificação", None), ("D", "Onde ocorreu", None), ("E", "O que foi combinado, e onde está escrito", None),
                 ("F", "O que foi encontrado, com identificação e amostra", None), ("G", "O desvio, em uma frase, sem causa e sem solução", None),
                 ("H", "Da área auditada", None), ("I", "Para a ação", None), ("J", "Do tratamento", None), ("K", "Três partes", None),
                 ("L", "Calculada", None)])
ex(ws, "B6", "Ex.", h="center")
ex(ws, "C6", "NC menor", h="center")
ex(ws, "D6", "Requisição de compra")
ex(ws, "E6", "PR-SUP-01 rev. 5, item 4.2: toda requisição deve conter a especificação técnica do item.")
ex(ws, "F6", "Requisições RC-0412, RC-0433, RC-0457 e RC-0461 sem especificação técnica, em amostra de 10.")
ex(ws, "G6", "Requisições de compra são aceitas sem a especificação técnica exigida pelo procedimento.")
ex(ws, "H6", "Gerente de Suprimentos")
ex(ws, "I6", date(2026, 10, 30), h="center", fmt=DATE)
ex(ws, "J6", "Em tratamento", h="center")
ex(ws, "K6", "Completa", h="center")
ex(ws, "L6", "No prazo", h="center")
ws.row_dimensions[6].height = 48
K1, K2 = 7, 21
for k in range(15):
    rr = K1 + k
    num(ws, f"B{rr}", k + 1)
    inp(ws, f"C{rr}", h="center")
    for col in "DEFGH":
        inp(ws, f"{col}{rr}")
    inp(ws, f"I{rr}", h="center", fmt=DATE)
    inp(ws, f"J{rr}", h="center")
    calc(ws, f"K{rr}", f'=IF(C{rr}="","",IF(LEFT(C{rr},2)="NC",IF(COUNTA(E{rr}:G{rr})=3,"Completa","Falta "&IF(E{rr}="","o requisito",'
         f'IF(F{rr}="","a evidência","a declaração"))),IF(COUNTA(F{rr}:G{rr})=2,"Completa","Falta a descrição")))', b=False, sz=9)
    calc(ws, f"L{rr}", f'=IF(C{rr}="","",IF(J{rr}="Eficácia verificada","Encerrada",IF(I{rr}="","Sem prazo",IF(I{rr}<TODAY(),"Atrasada","No prazo"))))',
         b=False, sz=9)
    ws.row_dimensions[rr].height = 48
dv_list(ws, f"C{K1}:C{K2}", TIPOS, "NC maior, NC menor ou Oportunidade de melhoria")
dv_list(ws, f"J{K1}:J{K2}", ["Aberta", "Em tratamento", "Eficácia verificada"], "Aberta, Em tratamento ou Eficácia verificada")
dv_date(ws, f"I{K1}:I{K2}")
cf_equal(ws, f"C{K1}:C{K2}", TIPO_CF)
cf_equal(ws, f"J{K1}:J{K2}", [("Eficácia verificada", GREEN), ("Em tratamento", YELLOW), ("Aberta", RED)])
cf_warn(ws, f"K{K1}:K{K2}", f"K{K1}", ok_values=("Completa",))
cf_equal(ws, f"L{K1}:L{K2}", [("Encerrada", GREEN), ("No prazo", GREEN), ("Sem prazo", YELLOW), ("Atrasada", RED)])
note(ws, "C4", "A classificação em maior e menor segue o critério definido pela organização. Oportunidade de melhoria não é não conformidade.")
note(ws, "J4", "A constatação só é encerrada depois da verificação de que a ação funcionou, e não na data em que a ação foi implantada.")
s = K2 + 2
band(ws, s, "Resumo automático", "L")
put(ws, f"B{s+1}", "Tipo", f=font(10, True), bg=GRAY, h="center", merge=f"B{s+1}:D{s+1}")
put(ws, f"E{s+1}", "Constatações", f=font(10, True), bg=GRAY, h="center")
for j, (name, color) in enumerate([("NC maior", REDC), ("NC menor", "D98B88"), ("Oportunidade de melhoria", AMBER)]):
    rr = s + 2 + j
    put(ws, f"B{rr}", name, f=font(10, True), bg=GRAY, merge=f"B{rr}:D{rr}")
    calc(ws, f"E{rr}", f'=COUNTIF(C{K1}:C{K2},"{name}")')
    ws.row_dimensions[rr].height = 21.75
for k, (text, formula) in enumerate([
    ("Total de constatações", f"=COUNTA(C{K1}:C{K2})"),
    ("Encerradas, com eficácia verificada", f'=COUNTIF(L{K1}:L{K2},"Encerrada")'),
    ("Atrasadas", f'=COUNTIF(L{K1}:L{K2},"Atrasada")'),
    ("Aviso", f'=IF(E{s+5}=0,"Registre as constatações",IF(COUNTIF(K{K1}:K{K2},"Falta*")>0,"Há constatação incompleta: veja a coluna Conferência",'
              f'IF(SUMPRODUCT((C{K1}:C{K2}<>"")*(H{K1}:H{K2}=""))>0,"Há constatação sem responsável","OK")))'),
], 5):
    rr = s + k
    put(ws, f"B{rr}", text, f=font(10, True), bg=GRAY, merge=f"B{rr}:D{rr}")
    calc(ws, f"E{rr}", formula, sz=9 if text == "Aviso" else 10, b=text != "Aviso")
    ws.row_dimensions[rr].height = 30 if text == "Aviso" else 21.75
cf_warn(ws, f"E{s+8}", f"E{s+8}")
ch = BarChart()
ch.type = "col"
ch.height, ch.width = 6.5, 15
ch.style = 2
ch.title = None
ch.legend = None
ch.gapWidth = 150
ch.add_data(Reference(ws, min_col=5, min_row=s + 1, max_row=s + 4), titles_from_data=True)
ch.set_categories(Reference(ws, min_col=2, min_row=s + 2, max_row=s + 4))
srs = ch.series[0]
srs.graphicalProperties.solidFill = MUTED
for j, color in enumerate([REDC, "D98B88", AMBER]):
    pt = DataPoint(idx=j)
    pt.graphicalProperties = GraphicalProperties(solidFill=color)
    pt.graphicalProperties.line.noFill = True
    srs.dPt.append(pt)
srs.dLbls = DataLabelList()
srs.dLbls.showVal = True
srs.dLbls.showSerName = srs.dLbls.showCatName = srs.dLbls.showLegendKey = False
ch.x_axis.delete = False
ch.y_axis.delete = False
ch.y_axis.scaling.min = 0
ch.y_axis.majorUnit = 1
ch.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill="D3DBE0", w=9525))
ch.x_axis.graphicalProperties = GraphicalProperties(ln=LineProperties(solidFill="D3DBE0", w=9525))
ch.y_axis.graphicalProperties = GraphicalProperties(ln=LineProperties(noFill=True))
ws.add_chart(ch, f"G{s+1}")
ws.freeze_panes = "D5"
setup(ws, REDC, f"B1:L{s+14}")
ws.print_title_rows = "4:4"

# ------------------------------------------------------------------ Checklist
ws = wb.create_sheet("Checklist")
widths(ws, {"A": 2, "B": 5, "C": 70, "D": 14, "E": 44, "F": 2})
title(ws, "Checklist de validação da auditoria", "Escolha o status de cada item na lista suspensa (coluna em amarelo-claro).", "E")
for col, text in zip("BCDE", ["#", "Item de verificação", "Status", "Observações"]):
    head(ws, f"{col}4", text)
ws.row_dimensions[4].height = 21.75
for k, text in enumerate([
    "A auditoria está prevista no programa, com objetivo, escopo e critérios definidos.",
    "A equipe auditora é independente da atividade auditada.",
    "O plano foi enviado ao auditado antes da auditoria.",
    "A lista de verificação cobre os critérios e indica o tamanho das amostras.",
    "Houve reunião de abertura e reunião de encerramento.",
    "As evidências foram coletadas no local, por entrevista, observação e análise de registros.",
    "As amostras foram escolhidas pelo auditor.",
    "Cada não conformidade tem requisito, evidência e declaração.",
    "As constatações não citam nomes de pessoas nem propõem soluções.",
    "O auditado conheceu as constatações antes do relatório.",
    "O relatório foi distribuído no prazo combinado.",
    "As constatações têm responsável, prazo e verificação de eficácia prevista.",
]):
    rr = 5 + k
    num(ws, f"B{rr}", k + 1)
    put(ws, f"C{rr}", text, bg=GRAY)
    inp(ws, f"D{rr}", h="center")
    inp(ws, f"E{rr}")
    ws.row_dimensions[rr].height = 30
dv_list(ws, "D5:D16", YESNO, "Sim, Parcial ou Não")
cf_equal(ws, "D5:D16", YESNO_CF)
band(ws, 18, "Resumo automático", "E")
summary(ws, 19, "Itens atendidos (Sim)", '=COUNTIF(D5:D16,"Sim")', "C", "D")
summary(ws, 20, "Itens parciais", '=COUNTIF(D5:D16,"Parcial")', "C", "D")
summary(ws, 21, "Itens não atendidos", '=COUNTIF(D5:D16,"Não")', "C", "D")
summary(ws, 22, "Itens sem resposta", "=COUNTBLANK(D5:D16)", "C", "D")
summary(ws, 23, "Percentual atendido", "=D19/ROWS(D5:D16)", "C", "D", fmt="0%")
summary(ws, 24, "Situação", '=IF(D19=ROWS(D5:D16),"Validado",IF(D22>0,"Em andamento","Com pendências"))', "C", "D")
setup(ws, PURPLE, "B1:E24", landscape=False, fit_height=True)


# ------------------------------------------------------------------ Exemplos
def example(ws, data):
    widths(ws, {"A": 2, "B": 5, "C": 22, "D": 36, "E": 28, "F": 40, "G": 24, "H": 14, "I": 2})
    title(ws, "Auditoria interna — plano, lista de verificação e constatações",
          "Exemplo preenchido, para consulta. Use as abas Plano, Lista de verificação e Constatações para a sua auditoria.", "H")
    H = data["head"]
    rr = 4
    for l1, v, fmt in [("Auditoria nº", H["num"], None), ("Data", H["data"], DATE), ("Processo", H["processo"], None),
                       ("Equipe auditora", f'{H["lider"]} (auditor líder) e {H["equipe"][0].lower() + H["equipe"][1:]}', None),
                       ("Auditado", H["auditado"], None), ("Objetivo", H["objetivo"], None), ("Escopo", H["escopo"], None),
                       ("Critérios", H["criterios"], None)]:
        label(ws, f"B{rr}", l1, merge=f"B{rr}:C{rr}")
        inp(ws, f"D{rr}", v, merge=f"D{rr}:H{rr}", bg=WHITE, fmt=fmt)
        ws.row_dimensions[rr].height = 21.75
        rr += 1
    rr += 1
    band(ws, rr, "Agenda", "H", color=TEAL)
    rr += 1
    for col, text in zip("BCDEF", ["#", "Horário", "Atividade", "Auditor", "Auditado"]):
        put(ws, f"{col}{rr}", text, f=font(10, True), bg=GRAY, h="center")
    put(ws, f"G{rr}", "Local", f=font(10, True), bg=GRAY, h="center", merge=f"G{rr}:H{rr}")
    ws.row_dimensions[rr].height = 21.75
    for k, (a, b, atv, aud, quem, local) in enumerate(data["agenda"], 1):
        rr += 1
        num(ws, f"B{rr}", k)
        put(ws, f"C{rr}", f"{a:%H:%M} a {b:%H:%M}", h="center")
        put(ws, f"D{rr}", atv)
        put(ws, f"E{rr}", aud)
        put(ws, f"F{rr}", quem)
        put(ws, f"G{rr}", local, merge=f"G{rr}:H{rr}")
        ws.row_dimensions[rr].height = 24
    rr += 2
    ws.row_breaks.append(Break(id=rr - 1))
    band(ws, rr, "Lista de verificação", "H", color=AMBER)
    rr += 1
    for col, text in zip("BCDEFGH", ["#", "Requisito", "O que verificar", "Amostra", "Evidência encontrada", "Resultado", "Constatação nº"]):
        put(ws, f"{col}{rr}", text, f=font(10, True), bg=GRAY, h="center")
    ws.row_dimensions[rr].height = 21.75
    v1 = rr + 1
    for i in data["lista"]:
        rr += 1
        num(ws, f"B{rr}", i["n"])
        put(ws, f"C{rr}", i["req"])
        put(ws, f"D{rr}", i["verificar"])
        put(ws, f"E{rr}", i["amostra"])
        put(ws, f"F{rr}", i["evid"])
        put(ws, f"G{rr}", i["res"], h="center")
        put(ws, f"H{rr}", i["ref"], h="center")
        ws.row_dimensions[rr].height = 42
    v2 = rr
    cf_equal(ws, f"G{v1}:G{v2}", RES_CF)
    rr += 2
    ws.row_breaks.append(Break(id=rr - 1))
    band(ws, rr, "Constatações", "H", color=REDC)
    rr += 1
    for col, text in zip("BCDEFGH", ["Nº", "Tipo", "Requisito", "Evidência", "Declaração", "Responsável", "Prazo"]):
        put(ws, f"{col}{rr}", text, f=font(10, True), bg=GRAY, h="center")
    ws.row_dimensions[rr].height = 21.75
    c1 = rr + 1
    for c in data["const"]:
        rr += 1
        num(ws, f"B{rr}", c["n"])
        put(ws, f"C{rr}", c["tipo"], h="center")
        put(ws, f"D{rr}", c["req"])
        put(ws, f"E{rr}", c["evid"])
        put(ws, f"F{rr}", c["decl"])
        put(ws, f"G{rr}", f'{c["resp"]}\n{c["status"]}')
        put(ws, f"H{rr}", c["prazo"], h="center", fmt=DATE)
        ws.row_dimensions[rr].height = 66
    cf_equal(ws, f"C{c1}:C{rr}", TIPO_CF)
    rr += 2
    band(ws, rr, "Resumo automático", "H")
    G = f"G{v1}:G{v2}"
    for k, (text, formula) in enumerate([
        ("Itens verificados", f"=COUNTA({G})"),
        ("Conformes", f'=COUNTIF({G},"Conforme")'),
        ("Não conformes", f'=COUNTIF({G},"Não conforme")'),
        ("Oportunidades de melhoria", f'=COUNTIF({G},"Oportunidade de melhoria")'),
        ("Constatações registradas", f"=COUNTA(C{c1}:C{c1 + len(data['const']) - 1})"),
    ], 1):
        label(ws, f"B{rr+k}", text, merge=f"B{rr+k}:E{rr+k}", h="right")
        calc(ws, f"F{rr+k}", formula)
        ws.row_dimensions[rr + k].height = 21.75
    rr += 6
    label(ws, f"B{rr}", "Conclusão da auditoria", merge=f"B{rr}:C{rr}")
    put(ws, f"D{rr}", data["conclusao"], merge=f"D{rr}:H{rr}")
    ws.row_dimensions[rr].height = 36
    setup(ws, MUTED, f"B1:H{rr}")


example(wb.create_sheet("Exemplo 1 - Pizzaria"), EX1)
example(wb.create_sheet("Exemplo 2 - Compras"), EX2)

wb.active = 1
wb.save(OUT)
print("ok", OUT)
