# -*- coding: utf-8 -*-
"""Gera Partes-Interessadas-modelo.xlsx no mesmo padrão visual das planilhas anteriores da série."""
import math
import os
import sys
from datetime import date

_here = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, _here)
# funções de formatação compartilhadas com o gerador do PDCA (bloco inicial do arquivo)
_src = open(os.path.join(_here, "build_pdca.py"), encoding="utf-8").read()
exec(_src[: _src.index("STATUS = [")])

from openpyxl.chart import BarChart, Series  # noqa: E402
from openpyxl.chart.label import DataLabelList  # noqa: E402
from pi_data import (ALTA, ATENDE, CHECK, CONTR, ESTRS, EX1, EX2, EXPEC, GERIR, GRUPOS, INFORM, LEGAL, MONIT, NAO, NAOAT, PARTE, SATISF, SIM, SITS, TIPOS,  # noqa: E402
                     estrategia, pertinente)

BLUE, TEAL, AMBER, PURPLE, REDC = "2B5C8A", "1E7B73", "A96A12", "7A4A9A", "B0413E"
BLUE_T, AMBER_T, TEAL_T, RED_T = "DCE8F3", "F6E8CF", "D9EEEB", "F5DEDC"
S1, S2, S3 = "2A6FB0", "D19A2E", "B0413E"   # atende, atende em parte, não atende
YESNO = ["Sim", "Parcial", "Não"]
YESNO_CF = [("Sim", GREEN), ("Parcial", YELLOW), ("Não", RED)]
ECOR = {GERIR: BLUE_T, SATISF: AMBER_T, INFORM: TEAL_T, MONIT: GRAY}
NP, NR = 12, 40
PT, RQ = "Partes", "Requisitos"
PA1, PA2 = 11, 11 + NP - 1     # partes, na aba Partes
R1, R2 = 7, 7 + NR - 1         # requisitos, na aba Requisitos
SEM_REQ = "Parte pertinente sem requisito"
NAO_ADOT = "Não adotado"


def note(ws, ref, text):
    c = Comment(text, "Modelo Partes interessadas")
    c.width, c.height = 300, 120
    ws[ref].comment = c


def cf_warn(ws, rng_, first, ok_values=("OK",)):
    cond = ",".join(f'{first}="{v}"' for v in ok_values)
    ws.conditional_formatting.add(rng_, FormulaRule(formula=[f"OR({cond})"], fill=PatternFill("solid", bgColor=GREEN, fgColor=GREEN)))
    ws.conditional_formatting.add(rng_, FormulaRule(formula=[f"AND(LEN({first})>0,NOT(OR({cond})))"],
                                  fill=PatternFill("solid", bgColor=YELLOW, fgColor=YELLOW)))


def cf_texto(ws, rng_, first, pares, resto=None):
    for t, cor in pares:
        ws.conditional_formatting.add(rng_, FormulaRule(formula=[f'{first}="{t}"'], stopIfTrue=True, fill=PatternFill("solid", bgColor=cor, fgColor=cor)))
    if resto:
        ws.conditional_formatting.add(rng_, FormulaRule(formula=[f"LEN({first})>0"], fill=PatternFill("solid", bgColor=resto, fgColor=resto)))


def hint_row(ws, row, cells, height=30):
    for ref, text in cells:
        put(ws, f"{ref}{row}", text, f=font(9, i=True, c=MUTED), bg=GRAY, h="center")
    ws.row_dimensions[row].height = height


def alt(pares, minimo=21.75, linha=12):
    n = max(max(1, math.ceil(len(str(t)) / max(c, 1))) for t, c in pares)
    return max(minimo, n * linha + 7)


def dv_numero(ws, rng_, minimo, maximo, prompt):
    dv = DataValidation(type="whole", operator="between", formula1=str(minimo), formula2=str(maximo), allow_blank=True)
    dv.promptTitle, dv.prompt, dv.showInputMessage = "Nota", prompt, True
    dv.errorTitle, dv.error, dv.showErrorMessage = "Nota inválida", f"Digite um número inteiro de {minimo} a {maximo}.", True
    ws.add_data_validation(dv)
    dv.add(rng_)


def f_estrategia(inf, int_):
    return f'IF({inf}>={ALTA},IF({int_}>={ALTA},"{GERIR}","{SATISF}"),IF({int_}>={ALTA},"{INFORM}","{MONIT}"))'


def chart_style(ch):
    ch.style = 2
    ch.title = None
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ch.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill="D3DBE0", w=9525))
    ch.x_axis.graphicalProperties = GraphicalProperties(ln=LineProperties(solidFill="D3DBE0", w=9525))
    ch.y_axis.graphicalProperties = GraphicalProperties(ln=LineProperties(noFill=True))


def empilhadas(ws, anchor, r1, r2, c_cat, cols, width=19, height=9.5):
    """Barras horizontais empilhadas: requisitos de cada parte, pela situação."""
    ch = BarChart()
    ch.type = "bar"
    ch.grouping = "stacked"
    ch.overlap = 100
    ch.gapWidth = 50
    ch.height, ch.width = height, width
    chart_style(ch)
    for col, tit, cor in cols:
        s = Series(Reference(ws, min_col=col, min_row=r1, max_row=r2), title=tit)
        s.graphicalProperties.solidFill = cor
        s.graphicalProperties.line.solidFill = "FFFFFF"
        ch.append(s)
    ch.set_categories(Reference(ws, min_col=c_cat, min_row=r1, max_row=r2))
    ch.y_axis.scaling.min = 0
    ch.y_axis.majorUnit = 1
    ch.y_axis.number_format = "0"
    ch.x_axis.scaling.orientation = "maxMin"
    ch.legend.position = "b"
    ws.add_chart(ch, anchor)


# ================================================================== pasta de trabalho
wb = Workbook()
wb.properties.title = "Partes interessadas — Modelo"
wb.properties.creator = "Modelo Partes interessadas"
wb.properties.language = "pt-BR"

# ------------------------------------------------------------------ Instruções
ws = wb.active
ws.title = "Instruções"
widths(ws, {"A": 2, "B": 24, "C": 92, "D": 2})
title(ws, "Partes interessadas — Como usar esta planilha",
      "Modelo para listar as partes interessadas, avaliar a influência e o interesse, registrar os requisitos e acompanhar o atendimento.", "C")
r = 4


def section(text):
    global r
    band(ws, r, text, "C", sz=11)
    r += 1


def line(k, v, kbg=GRAY, vbg=None, kf=None, kh="left", height=None):
    global r
    put(ws, f"B{r}", k, f=kf or font(10, True), bg=kbg, h=kh)
    put(ws, f"C{r}", v, bg=vbg)
    ws.row_dimensions[r].height = height or (19.5 if len(v) <= 98 else (31.5 if len(v) <= 196 else 45.75))
    r += 1


section("Legenda: onde preencher")
line("Amarelo-claro", "Células de entrada. É aqui que você digita.", vbg=INPUT)
line("Cinza", "Células calculadas ou fixas (rótulos, estratégias, conferências, contagens, avisos). Não altere.", vbg=GRAY)
line("Cinza em itálico", "Linha de exemplo (marcada com “Ex.”). Mostra o formato esperado e não entra nas contagens.", vbg=GRAY)
line("Listas suspensas", "Grupo, requisito legal, parte, tipo, adotado, situação e status aceitam apenas as opções da lista. As notas aceitam números de 1 a 5.")
line("Comentários", "Células com triângulo vermelho no canto trazem uma dica de preenchimento. Passe o mouse sobre elas.")
r += 1
section("As regras de cálculo")
for k, text in [
    ("Notas", f"Influência: o quanto a parte pode ajudar ou parar a operação. Interesse: o quanto ela acompanha e cobra. De 1 a 5. Nota {ALTA} ou 5 conta como alta."),
    ("Estratégia", "Influência alta e interesse alto: gerir de perto. Influência alta e interesse baixo: manter satisfeito. "
                   "Influência baixa e interesse alto: manter informado. As duas baixas: monitorar."),
    ("Pertinente", "É pertinente ao sistema a parte que fica fora do quadrante de monitorar, e a que tem requisito legal. Toda parte pertinente pede pelo menos um requisito."),
    ("Tipo do requisito", "Legal: vem de lei, regulamento ou licença. Contratual: vem de contrato, pedido ou promessa. Expectativa: a organização decide se adota."),
    ("Atendimento", "Pontos dos requisitos adotados: 1 para atende, 0,5 para atende em parte, 0 para não atende. Atendimento = pontos ÷ requisitos adotados."),
    ("Ações", "Todo requisito adotado que não atende por completo pede ação, responsável e prazo. A ação vencida é a que tem prazo anterior à data da análise."),
]:
    line(k, text)
r += 1
section("Ordem recomendada de preenchimento")
for k, text in enumerate([
    "Aba Partes: liste as partes, grupo por grupo, e escreva por que cada uma importa. Dê as notas e marque as que têm requisito legal.",
    "Aba Partes: para cada parte pertinente, escreva o canal de relacionamento e quem cuida dele.",
    "Aba Requisitos: escreva o que cada parte pertinente espera, com o tipo. Decida se cada expectativa é adotada.",
    "Aba Requisitos: ligue cada requisito adotado a um processo e a uma forma de monitorar. Avalie a situação e defina as ações.",
    "Abas Matriz e Painel: leia quem está em cada quadrante e que parte é atendida pior.",
    "Aba Checklist: valide a análise.",
], 1):
    line(f"Passo {k}", text)
r += 1
section("Abas da planilha")
for k, text in [
    (PT, f"Até {NP} partes interessadas, com grupo, motivo, notas, canal e quem cuida. Calcula a estratégia e a pertinência."),
    (RQ, f"Até {NR} requisitos, com tipo, decisão, processo, monitoramento, situação e ação. Confere cada linha."),
    ("Matriz", "A grade de influência e interesse, com a contagem de partes, e as partes de cada estratégia."),
    ("Painel", "Os indicadores gerais, os requisitos e o atendimento de cada parte, e o gráfico."),
    ("Checklist", "Doze verificações de qualidade da análise, com percentual de conclusão."),
    ("Exemplo 1 - Pizzaria", "As onze partes e os dezesseis requisitos da loja."),
    ("Exemplo 2 - Indústria", "As doze partes e os dezesseis requisitos da fábrica."),
]:
    line(k, text)
r += 1
section("Regras e premissas do modelo")
for k, text in [
    ("Nome da parte", "Na aba Requisitos, a parte é escolhida na lista da aba Partes. Se o nome mudar na aba Partes, escolha de novo na aba Requisitos."),
    ("Partes parecidas", f"O modelo aceita {NP} partes. Agrupe as que esperam a mesma coisa: “clientes de alimentos”, e não um cliente por linha."),
    ("Expectativa recusada", "Registre também a expectativa que não foi adotada, com o motivo na coluna “Como atende, ou por que não adota”."),
    ("Exemplos", "Os dados das abas de exemplo são ilustrativos, criados para este material."),
]:
    line(k, text)
setup(ws, INK, f"B1:C{r}", landscape=False, fit_height=True)

# ------------------------------------------------------------------ Partes
ws = wb.create_sheet(PT)
widths(ws, {"A": 2, "B": 5, "C": 30, "D": 20, "E": 40, "F": 11, "G": 11, "H": 11, "I": 18, "J": 11, "K": 36, "L": 20, "M": 28, "N": 2, "O": 24})
title(ws, "Partes interessadas", "Liste quem afeta a organização e quem é afetado por ela. A estratégia e a pertinência saem das notas.", "M")
for rr, rot in [(4, "Organização"), (5, "Data e responsáveis pela análise"), (6, "Escopo: o que a análise cobre")]:
    label(ws, f"B{rr}", rot, merge=f"B{rr}:D{rr}")
    inp(ws, f"E{rr}", merge=f"E{rr}:M{rr}", h="left")
    ws.row_dimensions[rr].height = 21.75
for col, text in zip("BCDEFGHIJKLM", ["#", "Parte interessada", "Grupo", "Por que importa", "Influência", "Interesse", "Requisito legal?", "Estratégia", "Pertinente",
                                       "Relacionamento: canal e frequência", "Quem cuida", "Conferência"]):
    head(ws, f"{col}8", text)
head(ws, "O8", "Chave: não altere", bg=MUTED)
ws.row_dimensions[8].height = 33
hint_row(ws, 9, [("B", ""), ("C", "Nome concreto"), ("D", "Escolha na lista"), ("E", "Como afeta, ou como é afetada"), ("F", "1 a 5"), ("G", "1 a 5"), ("H", "Sim ou Não"),
                 ("I", "Calculada"), ("J", "Calculada"), ("K", "Para as partes pertinentes"), ("L", "Função"), ("M", "Calculada"), ("O", "")])
ex(ws, "B10", "Ex.", h="center")
for col, v, h_ in [("C", "Vigilância sanitária", "left"), ("D", "Regulador", "left"), ("E", "Licencia e fiscaliza. Pode interditar a loja.", "left"), ("F", 5, "center"), ("G", 2, "center"),
                   ("H", "Sim", "center"), ("I", SATISF, "center"), ("J", "Sim", "center"), ("K", "Renovação anual da licença e visitas de fiscalização.", "left"),
                   ("L", "Gerente da loja", "left"), ("M", "OK", "center"), ("O", "", "left")]:
    ex(ws, f"{col}10", v, h=h_)
ws.row_dimensions[10].height = 27
for k in range(NP):
    rr = PA1 + k
    num(ws, f"B{rr}", k + 1)
    inp(ws, f"C{rr}")
    inp(ws, f"D{rr}")
    inp(ws, f"E{rr}")
    inp(ws, f"F{rr}", h="center")
    inp(ws, f"G{rr}", h="center")
    inp(ws, f"H{rr}", h="center")
    calc(ws, f"I{rr}", f'=IF(OR(F{rr}="",G{rr}=""),"",{f_estrategia(f"F{rr}", f"G{rr}")})', b=False, sz=9)
    calc(ws, f"J{rr}", f'=IF(I{rr}="","",IF(OR(I{rr}<>"{MONIT}",H{rr}="{SIM}"),"{SIM}","{NAO}"))', b=False)
    inp(ws, f"K{rr}")
    inp(ws, f"L{rr}")
    calc(ws, f"M{rr}", f'=IF(C{rr}="","",IF(COUNTIF(C${PA1}:C${PA2},C{rr})>1,"Parte repetida",IF(OR(F{rr}="",G{rr}=""),"Faltam as notas",IF(E{rr}="","Falta o motivo",'
         f'IF(J{rr}<>"{SIM}","OK",IF(K{rr}="","Falta o canal de relacionamento",IF(COUNTIF({RQ}!C${R1}:C${R2},C{rr})=0,"{SEM_REQ}","OK")))))))', b=False, sz=9)
    calc(ws, f"O{rr}", f'=IF(OR(C{rr}="",I{rr}=""),"",I{rr}&"|"&COUNTIFS(I${PA1}:I{rr},I{rr},C${PA1}:C{rr},"<>"))', b=False, sz=8, h="left")
    ws.row_dimensions[rr].height = 30
dv_list(ws, f"D{PA1}:D{PA2}", GRUPOS, "Grupo a que a parte pertence")
dv_numero(ws, f"F{PA1}:G{PA2}", 1, 5, f"De 1 a 5. Nota {ALTA} ou 5 conta como alta.")
dv_list(ws, f"H{PA1}:H{PA2}", [SIM, NAO], "A parte tem requisito de lei, de regulamento ou de licença?")
cf_texto(ws, f"I{PA1}:I{PA2}", f"I{PA1}", [(e, c) for e, c in ECOR.items() if e != MONIT])
cf_warn(ws, f"M{PA1}:M{PA2}", f"M{PA1}")
note(ws, "F8", "O quanto a parte pode ajudar ou parar a operação. 5: pode parar ou tirar a maior parte da receita. 1: sem efeito prático.")
note(ws, "G8", "O quanto a parte acompanha e cobra. 5: acompanha todos os dias. 1: não acompanha.")
note(ws, "H8", "Com “Sim”, a parte é pertinente mesmo com notas baixas, e precisa de requisito na aba Requisitos.")
note(ws, "C8", "Escreva o nome concreto: “Vigilância sanitária”, e não “Governo”. Agrupe as partes que esperam a mesma coisa.")
s = PA2 + 2
band(ws, s, "Resumo automático", "M")
EST_TXT = "&\", \"&".join(f'COUNTIF(I{PA1}:I{PA2},"{e}")&" {e.lower()}"' for e, _, _ in ESTRS)
for k, (text, formula) in enumerate([
    ("Partes listadas", f"=COUNTA(C{PA1}:C{PA2})"),
    ("Partes pertinentes ao sistema", f'=COUNTIF(J{PA1}:J{PA2},"{SIM}")'),
    ("Por estratégia", f'=IF(E{s+1}=0,"",{EST_TXT})'),
    ("Aviso", f'=IF(E{s+1}=0,"Liste as partes interessadas",IF(COUNTIF(M{PA1}:M{PA2},"OK")<E{s+1},"Há parte a rever: veja a coluna Conferência",'
              f'IF(E{s+1}<5,"Poucas partes: percorra os seis grupos","OK")))'),
], 1):
    label(ws, f"B{s+k}", text, merge=f"B{s+k}:D{s+k}", h="right")
    calc(ws, f"E{s+k}", formula, merge=f"E{s+k}:M{s+k}", sz=9 if text == "Aviso" else 10, b=text != "Aviso", h="left")
    ws.row_dimensions[s + k].height = 21.75
cf_warn(ws, f"E{s+4}:M{s+4}", f"E{s+4}")
PAV = s + 4
ws.freeze_panes = "D11"
setup(ws, BLUE, f"B1:M{s+4}", fit_height=True)

# ------------------------------------------------------------------ Requisitos
ws = wb.create_sheet(RQ)
widths(ws, {"A": 2, "B": 5, "C": 26, "D": 38, "E": 13, "F": 10, "G": 36, "H": 28, "I": 15, "J": 32, "K": 18, "L": 12, "M": 28, "N": 2})
title(ws, "Requisitos das partes interessadas", "Uma linha por necessidade ou expectativa. O que é legal ou contratual entra sempre. A expectativa entra por decisão.", "M")
for col, text in zip("BCDEFGHIJKLM", ["#", "Parte interessada", "O que espera", "Tipo", "Adotado?", "Como atende, ou por que não adota", "Como se monitora", "Situação", "Ação",
                                       "Responsável", "Prazo", "Conferência"]):
    head(ws, f"{col}4", text)
ws.row_dimensions[4].height = 33
hint_row(ws, 5, [("B", ""), ("C", "Da lista da aba Partes"), ("D", "Com verbo e critério"), ("E", "Escolha na lista"), ("F", "Sim ou Não"), ("G", "Processo ou documento, ou a justificativa"),
                 ("H", "Fonte e frequência"), ("I", "Escolha na lista"), ("J", "Se não atende por completo"), ("K", "Função"), ("L", "Data"), ("M", "Calculada")])
ex(ws, "B6", "Ex.", h="center")
for col, v, h_, fmt in [("C", "Vigilância sanitária", "left", None), ("D", "Manter a licença sanitária válida.", "left", None), ("E", LEGAL, "center", None), ("F", SIM, "center", None),
                        ("G", "Renovação protocolada em 09/10/2026.", "left", None), ("H", "Controle de vencimentos das licenças, todo mês.", "left", None), ("I", PARTE, "center", None),
                        ("J", "Acompanhar o protocolo até a emissão da licença.", "left", None), ("K", "Gerente da loja", "left", None), ("L", date(2026, 11, 13), "center", DATE),
                        ("M", "OK", "center", None)]:
    ex(ws, f"{col}6", v, h=h_, fmt=fmt)
ws.row_dimensions[6].height = 30
for k in range(NR):
    rr = R1 + k
    num(ws, f"B{rr}", k + 1)
    for col in "CDGHJK":
        inp(ws, f"{col}{rr}")
    for col in "EFI":
        inp(ws, f"{col}{rr}", h="center")
    inp(ws, f"L{rr}", h="center", fmt=DATE)
    calc(ws, f"M{rr}", f'=IF(D{rr}="","",IF(C{rr}="","Falta a parte",IF(E{rr}="","Falta o tipo",IF(F{rr}="","Decida se é adotado",'
         f'IF(AND(E{rr}<>"{EXPEC}",F{rr}="{NAO}"),"Legal ou contratual não pode ser recusado",IF(F{rr}="{NAO}",IF(G{rr}="","Falta a justificativa","{NAO_ADOT}"),'
         f'IF(G{rr}="","Falta como atende",IF(H{rr}="","Falta o monitoramento",IF(I{rr}="","Falta a situação",IF(I{rr}="{ATENDE}","OK",'
         f'IF(J{rr}="","Falta a ação",IF(OR(K{rr}="",L{rr}=""),"Ação sem responsável ou prazo","OK"))))))))))))', b=False, sz=9)
    ws.row_dimensions[rr].height = 30
dv = DataValidation(type="list", formula1=f"={PT}!$C${PA1}:$C${PA2}", allow_blank=True)
dv.promptTitle, dv.prompt, dv.showInputMessage = "Parte", "Escolha uma parte da lista da aba Partes.", True
dv.errorTitle, dv.error, dv.showErrorMessage = "Parte inválida", "Escolha uma parte cadastrada na aba Partes.", True
ws.add_data_validation(dv)
dv.add(f"C{R1}:C{R2}")
dv_list(ws, f"E{R1}:E{R2}", TIPOS, "Legal, contratual ou expectativa")
dv_list(ws, f"F{R1}:F{R2}", [SIM, NAO], "A organização adota este requisito?")
dv_list(ws, f"I{R1}:I{R2}", SITS, "Situação do requisito adotado")
dv_date(ws, f"L{R1}:L{R2}")
cf_texto(ws, f"I{R1}:I{R2}", f"I{R1}", [(ATENDE, GREEN), (PARTE, YELLOW), (NAOAT, RED)])
cf_warn(ws, f"M{R1}:M{R2}", f"M{R1}", ok_values=("OK", NAO_ADOT))
note(ws, "D4", "Escreva de modo que dê para saber se foi atendido: “Receber o pedido em até 40 minutos”, e não “Qualidade”.")
note(ws, "F4", "Requisito legal ou contratual é sempre “Sim”. A expectativa pode ser “Não”, com a justificativa na coluna ao lado.")
note(ws, "H4", "A fonte e a frequência: “Entregas em até 40 minutos, todo mês”, “Controle de vencimentos, todo mês”.")
s = R2 + 2
band(ws, s, "Resumo automático", "M")
F_ADOT = f'COUNTIF(F{R1}:F{R2},"{SIM}")'
for k, (text, formula, fmt) in enumerate([
    ("Requisitos levantados", f"=COUNTA(D{R1}:D{R2})", None),
    ("Requisitos adotados", f"={F_ADOT}", None),
    ("Expectativas não adotadas", f'=COUNTIF(F{R1}:F{R2},"{NAO}")', None),
    ("Atendimento dos requisitos adotados", f'=IF({F_ADOT}=0,"",(COUNTIFS(F{R1}:F{R2},"{SIM}",I{R1}:I{R2},"{ATENDE}")+0.5*COUNTIFS(F{R1}:F{R2},"{SIM}",I{R1}:I{R2},"{PARTE}"))/{F_ADOT})', "0%"),
    ("Ações abertas", f'=COUNTIFS(F{R1}:F{R2},"{SIM}",I{R1}:I{R2},"{PARTE}")+COUNTIFS(F{R1}:F{R2},"{SIM}",I{R1}:I{R2},"{NAOAT}")', None),
    ("Aviso", f'=IF(E{s+1}=0,"Registre os requisitos",IF(COUNTIF(M{R1}:M{R2},"OK")+COUNTIF(M{R1}:M{R2},"{NAO_ADOT}")<E{s+1},"Há requisito a rever: veja a coluna Conferência",'
              f'IF(COUNTIFS(E{R1}:E{R2},"{LEGAL}",I{R1}:I{R2},"{NAOAT}")>0,"Há requisito legal não atendido",'
              f'IF(COUNTIF({PT}!M{PA1}:M{PA2},"{SEM_REQ}")>0,"Há parte pertinente sem requisito: veja a aba Partes","OK"))))', None),
], 1):
    label(ws, f"B{s+k}", text, merge=f"B{s+k}:D{s+k}", h="right")
    calc(ws, f"E{s+k}", formula, fmt=fmt, merge=f"E{s+k}:G{s+k}", sz=9 if text == "Aviso" else 10, b=text != "Aviso", h="left")
    ws.row_dimensions[s + k].height = 21.75
cf_warn(ws, f"E{s+6}:G{s+6}", f"E{s+6}")
RS = s   # linha do resumo dos requisitos
ws.freeze_panes = "E7"
setup(ws, AMBER, f"B1:M{s+6}", fit_height=True)

# ------------------------------------------------------------------ Matriz
ws = wb.create_sheet("Matriz")
widths(ws, {"A": 2, "B": 16, "C": 8, "D": 15, "E": 15, "F": 15, "G": 15, "H": 15, "I": 2})
title(ws, "Matriz de influência e interesse", "Quantas partes há em cada célula. A cor mostra a estratégia. Tudo vem da aba Partes.", "H")
put(ws, "D4", "INTERESSE →", f=font(9, True, c=MUTED), box=False, h="center", merge="D4:H4")
ws.row_dimensions[4].height = 18
put(ws, "B5", "INFLUÊNCIA ↓", f=font(9, True, c=MUTED), bg=GRAY, h="center", merge="B5:C5")
for j in range(5):
    head(ws, f"{'DEFGH'[j]}5", str(j + 1))
ws.row_dimensions[5].height = 21.75
for i, inf in enumerate((5, 4, 3, 2, 1)):
    rr = 6 + i
    put(ws, f"B{rr}", "Alta" if inf >= ALTA else "Baixa", f=font(9, c=MUTED), bg=GRAY, h="center")
    head(ws, f"C{rr}", str(inf))
    for j in range(5):
        col = "DEFGH"[j]
        n = f"COUNTIFS({PT}!$F${PA1}:$F${PA2},{inf},{PT}!$G${PA1}:$G${PA2},{j + 1})"
        put(ws, f"{col}{rr}", f'=IF({n}=0,"",{n})', f=font(16, True), bg=ECOR[estrategia(inf, j + 1)], h="center")
    ws.row_dimensions[rr].height = 42
for j in range(5):
    put(ws, f"{'DEFGH'[j]}11", "Alto" if j + 1 >= ALTA else "Baixo", f=font(9, c=MUTED), bg=GRAY, h="center")
put(ws, "B11", "", bg=GRAY, merge="B11:C11")
ws.row_dimensions[11].height = 18
band(ws, 13, "As partes de cada estratégia", "H")
for col, text, m in [("B", "Estratégia", "B14:C14"), ("D", "Partes", None), ("E", "Quem são", "E14:H14")]:
    put(ws, f"{col}14", text, f=font(10, True), bg=GRAY, h="center", merge=m)
ws.row_dimensions[14].height = 21.75
for k, (e, quando, fazer) in enumerate(ESTRS):
    rr = 15 + k
    put(ws, f"B{rr}", e, f=font(10, True), bg=ECOR[e], merge=f"B{rr}:C{rr}")
    calc(ws, f"D{rr}", f'=COUNTIF({PT}!$I${PA1}:$I${PA2},"{e}")')
    nomes = "&".join(f'IFERROR({"" if n == 1 else chr(34) + "; " + chr(34) + "&"}INDEX({PT}!$C${PA1}:$C${PA2},MATCH("{e}|{n}",{PT}!$O${PA1}:$O${PA2},0)),"")' for n in range(1, NP + 1))
    calc(ws, f"E{rr}", f"={nomes}", h="left", b=False, sz=9, merge=f"E{rr}:H{rr}")
    ws.row_dimensions[rr].height = 45
band(ws, 20, "O que fazer em cada estratégia", "H")
for k, (e, quando, fazer) in enumerate(ESTRS):
    rr = 21 + k
    put(ws, f"B{rr}", e, f=font(10, True), bg=ECOR[e], merge=f"B{rr}:C{rr}")
    put(ws, f"D{rr}", f"{quando}. {fazer}", bg=GRAY, merge=f"D{rr}:H{rr}")
    ws.row_dimensions[rr].height = 30
label(ws, "B26", "Aviso", merge="B26:C26", h="right")
calc(ws, "D26", f'=IF(COUNT({PT}!F{PA1}:F{PA2})=0,"Dê as notas na aba Partes",IF(D15>COUNTA({PT}!C{PA1}:C{PA2})/2,"Mais da metade das partes em gerir de perto: reveja as notas",'
     f'IF(D18=0,"Nenhuma parte só monitorada: a lista pode estar incompleta","OK")))', sz=9, b=False, h="left", merge="D26:H26")
cf_warn(ws, "D26:H26", "D26")
ws.row_dimensions[26].height = 21.75
setup(ws, TEAL, "B1:H26", landscape=False, fit_height=True)

# ------------------------------------------------------------------ Painel
ws = wb.create_sheet("Painel")
widths(ws, {"A": 2, "B": 5, "C": 32, "D": 18, "E": 11, "F": 12, "G": 10, "H": 10, "I": 10, "J": 13, "K": 28, "L": 2})
title(ws, "Painel das partes interessadas", "Os indicadores gerais e o atendimento de cada parte. Informe só a data da análise.", "K")
label(ws, "B4", "Data da análise", merge="B4:C4")
inp(ws, "D4", h="center", fmt=DATE)
put(ws, "E4", "Usada para contar as ações vencidas.", f=font(9, i=True, c=MUTED), bg=GRAY, merge="E4:K4")
ws.row_dimensions[4].height = 21.75
dv_date(ws, "D4")
for col, text, m in [("B", "Indicador", "B6:C6"), ("D", "Resultado", None), ("E", "Como se lê", "E6:K6")]:
    put(ws, f"{col}6", text, f=font(10, True, c=WHITE), bg=INK, h="center", merge=m)
ws.row_dimensions[6].height = 21.75
F_, I_, L_, E_ = (f"{RQ}!{c}{R1}:{c}{R2}" for c in "FILE")
IND = [
    ("Partes listadas", f"=COUNTA({PT}!C{PA1}:C{PA2})", None, "Todas as partes da aba Partes."),
    ("Partes pertinentes ao sistema", f'=COUNTIF({PT}!J{PA1}:J{PA2},"{SIM}")', None, "Fora do quadrante de monitorar, ou com requisito legal."),
    ("Partes pertinentes sem requisito", f'=COUNTIF({PT}!M{PA1}:M{PA2},"{SEM_REQ}")', None, "Deve ser zero: toda parte pertinente pede um requisito."),
    ("Requisitos adotados", f'=COUNTIF({F_},"{SIM}")', None, "Legais, contratuais e expectativas adotadas."),
    ("Atendimento geral", f'=IF({RQ}!E{RS+4}="","",{RQ}!E{RS+4})', "0%", "Atende vale 1, atende em parte vale 0,5."),
    ("Requisitos legais atendidos por completo", f'=IF(COUNTIFS({E_},"{LEGAL}",{F_},"{SIM}")=0,"",COUNTIFS({E_},"{LEGAL}",{F_},"{SIM}",{I_},"{ATENDE}")&" de "&COUNTIFS({E_},"{LEGAL}",{F_},"{SIM}"))',
     None, "O legal em parte tem prazo correndo: acompanhe de perto."),
    ("Ações abertas", f"={RQ}!E{RS+5}", None, "Requisitos adotados que não atendem por completo."),
    ("Ações vencidas", f'=IF(D4="","",SUMPRODUCT(({F_}="{SIM}")*(({I_}="{PARTE}")+({I_}="{NAOAT}"))*ISNUMBER({L_})*({L_}<D4)))', None, "Prazo anterior à data da análise."),
]
for k, (nome, formula, fmt, leitura) in enumerate(IND):
    rr = 7 + k
    label(ws, f"B{rr}", nome, merge=f"B{rr}:C{rr}")
    calc(ws, f"D{rr}", formula, fmt=fmt)
    put(ws, f"E{rr}", leitura, f=font(9, c=MUTED), bg=GRAY, merge=f"E{rr}:K{rr}")
    ws.row_dimensions[rr].height = 21.75
I2 = 7 + len(IND) - 1
T0 = I2 + 2
band(ws, T0, "Requisitos e atendimento por parte", "K")
for col, text in zip("BCDEFGHIJK", ["#", "Parte interessada", "Estratégia", "Pertinente", "Requisitos adotados", ATENDE, "Em parte", NAOAT, "Atendimento", "Leitura"]):
    head(ws, f"{col}{T0+1}", text)
ws.row_dimensions[T0 + 1].height = 33
T1, T2 = T0 + 2, T0 + 1 + NP
C_ = f"{RQ}!$C${R1}:$C${R2}"
FA, IA = f"{RQ}!$F${R1}:$F${R2}", f"{RQ}!$I${R1}:$I${R2}"
for k in range(NP):
    rr = T1 + k
    p = PA1 + k
    num(ws, f"B{rr}", k + 1)
    calc(ws, f"C{rr}", f'=IF({PT}!C{p}="","",{PT}!C{p})', h="left", b=False)
    calc(ws, f"D{rr}", f'=IF(C{rr}="","",{PT}!I{p})', b=False, sz=9)
    calc(ws, f"E{rr}", f'=IF(C{rr}="","",{PT}!J{p})', b=False)
    calc(ws, f"F{rr}", f'=IF(C{rr}="","",COUNTIFS({C_},C{rr},{FA},"{SIM}"))', b=False)
    for col, sit in (("G", ATENDE), ("H", PARTE), ("I", NAOAT)):
        calc(ws, f"{col}{rr}", f'=IF(C{rr}="","",COUNTIFS({C_},C{rr},{FA},"{SIM}",{IA},"{sit}"))', b=False)
    calc(ws, f"J{rr}", f'=IF(OR(C{rr}="",N(F{rr})=0),"",(G{rr}+0.5*H{rr})/F{rr})', fmt="0%")
    calc(ws, f"K{rr}", f'=IF(C{rr}="","",IF(E{rr}<>"{SIM}","Só monitorada",IF(F{rr}=0,"Sem requisito",IF(G{rr}+H{rr}+I{rr}<F{rr},"Falta avaliar a situação",'
         f'IF(I{rr}>0,"Há requisito não atendido",IF(H{rr}>0,"Há pendência","Atendida"))))))', b=False, sz=9)
    ws.row_dimensions[rr].height = 21.75
cf_texto(ws, f"D{T1}:D{T2}", f"D{T1}", [(e, c) for e, c in ECOR.items() if e != MONIT])
cf_texto(ws, f"K{T1}:K{T2}", f"K{T1}", [("Atendida", GREEN), ("Há pendência", YELLOW), ("Só monitorada", GRAY)], resto=RED)
s = T2 + 2
band(ws, s, "Resumo automático", "K")
label(ws, f"B{s+1}", "Parte com o menor atendimento", merge=f"B{s+1}:C{s+1}", h="right")
calc(ws, f"D{s+1}", f'=IF(COUNT(J{T1}:J{T2})=0,"",INDEX(C{T1}:C{T2},MATCH(MIN(J{T1}:J{T2}),J{T1}:J{T2},0))&" · "&TEXT(MIN(J{T1}:J{T2}),"0%"))', h="left", merge=f"D{s+1}:K{s+1}")
label(ws, f"B{s+2}", "Aviso", merge=f"B{s+2}:C{s+2}", h="right")
calc(ws, f"D{s+2}", f'=IF(D7=0,"Liste as partes na aba Partes",IF(D10=0,"Registre os requisitos na aba Requisitos",IF(D9>0,"Há parte pertinente sem requisito",'
     f'IF(D4="","Informe a data da análise",IF(D14>0,"Há ação vencida: veja a aba Requisitos",IF(COUNTIF(K{T1}:K{T2},"Há requisito não atendido")>0,"Há requisito não atendido: leve à direção","OK"))))))',
     sz=9, b=False, h="left", merge=f"D{s+2}:K{s+2}")
cf_warn(ws, f"D{s+2}:K{s+2}", f"D{s+2}")
for k in (1, 2):
    ws.row_dimensions[s + k].height = 21.75
NAV = s + 2
empilhadas(ws, f"B{s + 4}", T1, T2, 3, [(7, ATENDE, S1), (8, PARTE, S2), (9, NAOAT, S3)], width=24, height=10)
setup(ws, REDC, f"B1:K{s + 25}", landscape=False, fit_height=True)

# ------------------------------------------------------------------ Checklist
ws = wb.create_sheet("Checklist")
widths(ws, {"A": 2, "B": 5, "C": 70, "D": 14, "E": 44, "F": 2})
title(ws, "Checklist das partes interessadas", "Escolha o status de cada item na lista suspensa (coluna em amarelo-claro).", "E")
for col, text in zip("BCDE", ["#", "Item de verificação", "Status", "Observações"]):
    head(ws, f"{col}4", text)
ws.row_dimensions[4].height = 21.75
for k, text in enumerate(CHECK):
    rr = 5 + k
    num(ws, f"B{rr}", k + 1)
    put(ws, f"C{rr}", text, bg=GRAY)
    inp(ws, f"D{rr}", h="center")
    inp(ws, f"E{rr}")
    ws.row_dimensions[rr].height = 30
dv_list(ws, "D5:D16", YESNO, "Sim, Parcial ou Não")
cf_equal(ws, "D5:D16", YESNO_CF)
band(ws, 18, "Resumo automático", "E")
summary(ws, 19, "Itens atendidos (Sim)", '=COUNTIF(D5:D16,"Sim")', "C", "D")
summary(ws, 20, "Itens parciais", '=COUNTIF(D5:D16,"Parcial")', "C", "D")
summary(ws, 21, "Itens não atendidos", '=COUNTIF(D5:D16,"Não")', "C", "D")
summary(ws, 22, "Itens sem resposta", "=COUNTBLANK(D5:D16)", "C", "D")
summary(ws, 23, "Percentual atendido", "=D19/ROWS(D5:D16)", "C", "D", fmt="0%")
summary(ws, 24, "Situação", '=IF(D19=ROWS(D5:D16),"Validado",IF(D22>0,"Em andamento","Com pendências"))', "C", "D")
setup(ws, PURPLE, "B1:E24", landscape=False, fit_height=True)


# ------------------------------------------------------------------ Exemplos
def sub(ws, rr, cells, height=30):
    for a, b, t in cells:
        put(ws, f"{a}{rr}", t, f=font(10, True), bg=GRAY, h="center", merge=b and f"{a}{rr}:{b}{rr}")
    ws.row_dimensions[rr].height = height


def exemplo(ws, ex):
    H = ex["head"]
    widths(ws, {"A": 2, "B": 5, "C": 28, "D": 36, "E": 13, "F": 10, "G": 36, "H": 28, "I": 17, "J": 30, "K": 18, "L": 12, "M": 2})
    title(ws, "Partes interessadas", "Exemplo preenchido, para consulta. Use as abas Partes e Requisitos para a sua organização.", "L")
    rr = 4
    for rot, val in [("Organização", H["org"]), ("Data e responsáveis", f'{H["data"]:%d/%m/%Y}. {H["por"]}.'), ("Escopo", H["escopo"]), ("Origem", H["origem"]), ("Revisão", H["revisao"])]:
        label(ws, f"B{rr}", rot, merge=f"B{rr}:C{rr}")
        inp(ws, f"D{rr}", val, merge=f"D{rr}:L{rr}", bg=WHITE, h="left")
        ws.row_dimensions[rr].height = alt([(val, 150)], minimo=19.5)
        rr += 1
    rr += 1
    band(ws, rr, "Partes interessadas", "L", color=BLUE)
    rr += 1
    sub(ws, rr, [("B", None, "#"), ("C", None, "Parte interessada"), ("D", None, "Por que importa"), ("E", None, "Influência"), ("F", None, "Interesse"),
                 ("G", None, "Relacionamento: canal e frequência"), ("H", None, "Quem cuida"), ("I", None, "Estratégia"), ("J", None, "Grupo"), ("K", None, "Requisito legal?"), ("L", None, "Pertinente")])
    p1 = rr + 1
    for p in ex["partes"]:
        rr += 1
        num(ws, f"B{rr}", p["n"])
        put(ws, f"C{rr}", p["nome"], f=font(10, True))
        put(ws, f"D{rr}", p["porque"])
        put(ws, f"E{rr}", p["inf"], h="center")
        put(ws, f"F{rr}", p["int"], h="center")
        put(ws, f"G{rr}", p["canal"] or "—")
        put(ws, f"H{rr}", p["quem"] or "—")
        calc(ws, f"I{rr}", f"={f_estrategia(f'E{rr}', f'F{rr}')}", b=False, sz=9)
        put(ws, f"J{rr}", p["grupo"])
        put(ws, f"K{rr}", SIM if p["legal"] else NAO, h="center")
        calc(ws, f"L{rr}", f'=IF(OR(I{rr}<>"{MONIT}",K{rr}="{SIM}"),"{SIM}","{NAO}")', b=False)
        ws.row_dimensions[rr].height = alt([(p["porque"], 38), (p["canal"], 38), (p["nome"], 28)], minimo=21.75)
    p2 = rr
    cf_texto(ws, f"I{p1}:I{p2}", f"I{p1}", [(e, c) for e, c in ECOR.items() if e != MONIT])
    rr += 2
    ws.row_breaks.append(Break(id=rr - 1))
    band(ws, rr, "Requisitos", "L", color=AMBER)
    rr += 1
    sub(ws, rr, [("B", None, "#"), ("C", None, "Parte interessada"), ("D", None, "O que espera"), ("E", None, "Tipo"), ("F", None, "Adotado?"), ("G", None, "Como atende, ou por que não adota"),
                 ("H", None, "Como se monitora"), ("I", None, "Situação"), ("J", None, "Ação"), ("K", None, "Responsável"), ("L", None, "Prazo")])
    q1 = rr + 1
    for q in ex["reqs"]:
        rr += 1
        num(ws, f"B{rr}", q["n"])
        put(ws, f"C{rr}", q["parte"], f=font(10, True))
        put(ws, f"D{rr}", q["req"])
        put(ws, f"E{rr}", q["tipo"], h="center")
        put(ws, f"F{rr}", q["adotado"], h="center")
        put(ws, f"G{rr}", q["como"])
        put(ws, f"H{rr}", q["monit"] or "—")
        put(ws, f"I{rr}", q["sit"] or "—", h="center")
        put(ws, f"J{rr}", q["acao"] or "—")
        put(ws, f"K{rr}", q["resp"] or "—")
        put(ws, f"L{rr}", q["prazo"] or "—", h="center", fmt=DATE)
        ws.row_dimensions[rr].height = alt([(q["req"], 38), (q["como"], 38), (q["monit"] or "", 30), (q["acao"] or "", 32), (q["parte"], 28)], minimo=21.75)
    q2 = rr
    cf_texto(ws, f"I{q1}:I{q2}", f"I{q1}", [(ATENDE, GREEN), (PARTE, YELLOW), (NAOAT, RED)])
    rr += 2
    band(ws, rr, "Resumo automático", "L")
    adot = f'COUNTIF(F{q1}:F{q2},"{SIM}")'
    for k, (text, formula, fmt) in enumerate([
        ("Partes listadas e pertinentes", f'=COUNTA(C{p1}:C{p2})&" listadas, "&COUNTIF(L{p1}:L{p2},"{SIM}")&" pertinentes"', None),
        ("Partes em gerir de perto", f'=COUNTIF(I{p1}:I{p2},"{GERIR}")', None),
        ("Requisitos levantados e adotados", f'=COUNTA(D{q1}:D{q2})&" levantados, "&{adot}&" adotados"', None),
        ("Atendimento dos requisitos adotados", f'=(COUNTIF(I{q1}:I{q2},"{ATENDE}")+0.5*COUNTIF(I{q1}:I{q2},"{PARTE}"))/{adot}', "0%"),
        ("Requisitos legais atendidos por completo", f'=COUNTIFS(E{q1}:E{q2},"{LEGAL}",I{q1}:I{q2},"{ATENDE}")&" de "&COUNTIF(E{q1}:E{q2},"{LEGAL}")', None),
        ("Ações abertas", f'=COUNTIF(I{q1}:I{q2},"{PARTE}")+COUNTIF(I{q1}:I{q2},"{NAOAT}")', None),
    ], 1):
        label(ws, f"B{rr+k}", text, merge=f"B{rr+k}:D{rr+k}", h="right")
        calc(ws, f"E{rr+k}", formula, fmt=fmt, merge=f"E{rr+k}:G{rr+k}", h="left")
        ws.row_dimensions[rr + k].height = 21.75
    setup(ws, MUTED, f"B1:L{rr + 6}")


exemplo(wb.create_sheet("Exemplo 1 - Pizzaria"), EX1)
exemplo(wb.create_sheet("Exemplo 2 - Indústria"), EX2)

wb.active = 1
wb.save(OUT)
print("ok", OUT, "| avisos: Partes E%d, Requisitos E%d, Painel D%d" % (PAV, RS + 6, NAV))
