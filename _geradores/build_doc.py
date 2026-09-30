# -*- coding: utf-8 -*-
"""Gera Informacao-Documentada-modelo.xlsx no mesmo padrão visual das planilhas anteriores da série."""
import math
import os
import sys
from datetime import date

_here = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, _here)
# funções de formatação compartilhadas com o gerador do PDCA (bloco inicial do arquivo)
_src = open(os.path.join(_here, "build_pdca.py"), encoding="utf-8").read()
exec(_src[: _src.index("STATUS = [")])

from openpyxl.chart import BarChart, Series  # noqa: E402
from openpyxl.chart.label import DataLabelList  # noqa: E402
from doc_data import (CHECK, EMREV, EX1, EX2, MEIOS, OBS, REV_MESES, S_APROV, S_CODIGO, S_EMREV, S_OBS, S_REV, S_VENC, S_VIG, STATUS, TIPOS, VIG)  # noqa: E402

BLUE, TEAL, AMBER, PURPLE, REDC = "2B5C8A", "1E7B73", "A96A12", "7A4A9A", "B0413E"
BLUE_T, AMBER_T, RED_T = "DCE8F3", "F8EBCB", "F5DEDC"
SIGLAS = [t[0] for t in TIPOS]
SITS = [S_VIG, S_EMREV, S_VENC, S_CODIGO, S_REV, S_APROV, S_OBS]
SIT_CF = [(S_VIG, GREEN), (S_EMREV, BLUE_T), (S_VENC, YELLOW), (S_OBS, GRAY)]
YESNO = ["Sim", "Parcial", "Não"]
YESNO_CF = [("Sim", GREEN), ("Parcial", YELLOW), ("Não", RED)]
LIGADOS = ["Sim", "Não", "Não há"]
NDOC, NREG, NEXT, NALT = 30, 20, 15, 20
L1, L2 = 9, 9 + NDOC - 1       # linhas da lista mestra
R1, R2 = 7, 7 + NREG - 1       # registros
X1, X2 = 7, 7 + NEXT - 1       # externos
A1, A2 = 7, 7 + NALT - 1       # alterações
LISTA = "'Lista mestra'"
REF = f"{LISTA}!$D$5"          # data de referência
PADRAO = f"{LISTA}!$D$6"       # intervalo padrão de revisão


def note(ws, ref, text):
    c = Comment(text, "Modelo Informação documentada")
    c.width, c.height = 300, 120
    ws[ref].comment = c


def cf_warn(ws, rng_, first, ok_values=("OK",)):
    cond = ",".join(f'{first}="{v}"' for v in ok_values)
    ws.conditional_formatting.add(rng_, FormulaRule(formula=[f"OR({cond})"], fill=PatternFill("solid", bgColor=GREEN, fgColor=GREEN)))
    ws.conditional_formatting.add(rng_, FormulaRule(formula=[f"AND(LEN({first})>0,NOT(OR({cond})))"],
                                  fill=PatternFill("solid", bgColor=YELLOW, fgColor=YELLOW)))


def cf_situacao(ws, rng_, first):
    """Cores da situação do documento: verde em vigor, azul em revisão, amarelo vencida, cinza obsoleto, vermelho para o que falta."""
    for t, cor in SIT_CF:
        ws.conditional_formatting.add(rng_, FormulaRule(formula=[f'{first}="{t}"'], stopIfTrue=True, fill=PatternFill("solid", bgColor=cor, fgColor=cor)))
    ws.conditional_formatting.add(rng_, FormulaRule(formula=[f'LEFT({first},5)="Falta"'], fill=PatternFill("solid", bgColor=RED, fgColor=RED)))


def cf_falta(ws, rng_, first, ok="Completa"):
    ws.conditional_formatting.add(rng_, FormulaRule(formula=[f'{first}="{ok}"'], stopIfTrue=True, fill=PatternFill("solid", bgColor=GREEN, fgColor=GREEN)))
    ws.conditional_formatting.add(rng_, FormulaRule(formula=[f"LEN({first})>0"], fill=PatternFill("solid", bgColor=YELLOW, fgColor=YELLOW)))


def hint_row(ws, row, cells, height=30):
    for ref, text in cells:
        put(ws, f"{ref}{row}", text, f=font(9, i=True, c=MUTED), bg=GRAY, h="center")
    ws.row_dimensions[row].height = height


def alt(pares, minimo=21.75, linha=12):
    n = max(max(1, math.ceil(len(str(t)) / max(c, 1))) for t, c in pares)
    return max(minimo, n * linha + 7)


def dv_inteiro(ws, rng_, maximo, prompt):
    dv = DataValidation(type="whole", operator="between", formula1="1", formula2=str(maximo), allow_blank=True)
    dv.promptTitle, dv.prompt, dv.showInputMessage = "Meses", prompt, True
    dv.errorTitle, dv.error, dv.showErrorMessage = "Número inválido", f"Digite um número inteiro de 1 a {maximo}.", True
    ws.add_data_validation(dv)
    dv.add(rng_)


def chart_style(ch):
    ch.style = 2
    ch.title = None
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ch.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill="D3DBE0", w=9525))
    ch.x_axis.graphicalProperties = GraphicalProperties(ln=LineProperties(solidFill="D3DBE0", w=9525))
    ch.y_axis.graphicalProperties = GraphicalProperties(ln=LineProperties(noFill=True))


def barras(ws, anchor, r1, r2, c_cat, c_val, width=16, height=6.5):
    """Barras horizontais: documentos em cada situação."""
    ch = BarChart()
    ch.type = "bar"
    ch.gapWidth = 60
    ch.height, ch.width = height, width
    chart_style(ch)
    ch.legend = None
    s = Series(Reference(ws, min_col=c_val, min_row=r1, max_row=r2), title="Documentos")
    s.graphicalProperties.solidFill = BLUE
    s.graphicalProperties.line.noFill = True
    s.dLbls = DataLabelList()
    s.dLbls.showVal = True
    s.dLbls.showSerName = s.dLbls.showCatName = s.dLbls.showLegendKey = False
    ch.append(s)
    ch.set_categories(Reference(ws, min_col=c_cat, min_row=r1, max_row=r2))
    ch.y_axis.scaling.min = 0
    ch.y_axis.majorUnit = 1
    ch.y_axis.number_format = "0"
    ch.x_axis.scaling.orientation = "maxMin"
    ws.add_chart(ch, anchor)


def f_proxima(r, data="H", meses="M", titulo="D", padrao=PADRAO):
    return f'=IF(OR({data}{r}="",{titulo}{r}=""),"",EDATE({data}{r},IF({meses}{r}="",{padrao},{meses}{r})))'


def f_situacao(r, ref, titulo="D", codigo="C", rev="G", data="H", aprovou="J", proxima="N", status="O"):
    return (f'=IF({titulo}{r}="","",IF({codigo}{r}="","{S_CODIGO}",IF({status}{r}="{OBS}","{S_OBS}",IF(OR({rev}{r}="",{data}{r}=""),"{S_REV}",'
            f'IF({aprovou}{r}="","{S_APROV}",IF({status}{r}="{EMREV}","{S_EMREV}",IF({proxima}{r}<{ref},"{S_VENC}","{S_VIG}")))))))')


# ================================================================== pasta de trabalho
wb = Workbook()
wb.properties.title = "Informação documentada — Modelo"
wb.properties.creator = "Modelo Informação documentada"
wb.properties.language = "pt-BR"

# ------------------------------------------------------------------ Instruções
ws = wb.active
ws.title = "Instruções"
widths(ws, {"A": 2, "B": 24, "C": 92, "D": 2})
title(ws, "Informação documentada — Como usar esta planilha",
      "Modelo de lista mestra, tabela de retenção de registros, lista de documentos externos e histórico de alterações.", "C")
r = 4


def section(text):
    global r
    band(ws, r, text, "C", sz=11)
    r += 1


def line(k, v, kbg=GRAY, vbg=None, kf=None, kh="left", height=None):
    global r
    put(ws, f"B{r}", k, f=kf or font(10, True), bg=kbg, h=kh)
    put(ws, f"C{r}", v, bg=vbg)
    ws.row_dimensions[r].height = height or (19.5 if len(v) <= 98 else (31.5 if len(v) <= 196 else 45.75))
    r += 1


section("Legenda: onde preencher")
line("Amarelo-claro", "Células de entrada. É aqui que você digita.", vbg=INPUT)
line("Cinza", "Células calculadas ou fixas (rótulos, próxima revisão, situação, contagens, avisos). Não altere.", vbg=GRAY)
line("Cinza em itálico", "Linha de exemplo (marcada com “Ex.”). Mostra o formato esperado e não entra nas contagens.", vbg=GRAY)
line("Listas suspensas", "Colunas de tipo, meio, status e conferência aceitam apenas as opções da lista.")
line("Comentários", "Células com triângulo vermelho no canto trazem uma dica de preenchimento. Passe o mouse sobre elas.")
r += 1
section("Os tipos de documento")
for sigla, nome, oque, exemplo_ in TIPOS:
    line(f"{sigla} · {nome}", f"{oque} Ex.: {exemplo_}")
r += 1
section("Ordem recomendada de preenchimento")
for k, text in enumerate([
    "Aba Lista mestra: escreva a organização, a data de referência e o intervalo padrão de revisão.",
    "Aba Lista mestra: cadastre cada documento, com código, tipo, processo, revisão, data, quem elaborou e quem aprovou.",
    "Aba Lista mestra: leia a coluna Situação e trate o que falta.",
    "Aba Registros: cadastre cada registro, com local de guarda, meio, prazo de retenção, acesso e disposição.",
    "Aba Externos: liste as normas, as leis e os documentos de cliente, com o responsável e o intervalo de verificação.",
    "Aba Alterações: a cada revisão de documento, registre o que mudou e confira os documentos ligados.",
    "Aba Checklist: valide o controle de documentos.",
], 1):
    line(f"Passo {k}", text)
r += 1
section("Abas da planilha")
for k, text in [
    ("Lista mestra", "Até 30 documentos, com a próxima revisão e a situação calculadas, contagem por situação e gráfico."),
    ("Registros", "Até 20 registros, com a conferência dos campos de retenção."),
    ("Externos", "Até 15 documentos de origem externa, com a próxima verificação e a situação calculadas."),
    ("Alterações", "Histórico de até 20 alterações, com a conferência dos documentos ligados."),
    ("Checklist", "Doze verificações de qualidade do controle de documentos, com percentual de conclusão."),
    ("Exemplo 1 - Pizzaria", "A lista mestra e a tabela de retenção de uma loja, depois do levantamento."),
    ("Exemplo 2 - Compras", "Os documentos e os registros do processo de aquisição de uma indústria."),
]:
    line(k, text)
r += 1
section("Regras e premissas do modelo")
for k, text in [
    ("Situação do documento", "Vigente: tem código, revisão, aprovação e a próxima revisão ainda não chegou. Revisão vencida: passou da data prevista. "
                             "Falta o código, a revisão ou a aprovação: o documento está em uso, mas fora do controle. Em revisão e Obsoleto vêm da coluna Status."),
    ("Próxima revisão", f"É a data da revisão mais o intervalo, em meses. Sem intervalo na linha, vale o intervalo padrão da aba, de {REV_MESES} meses."),
    ("Data de referência", "A situação é calculada pela data de referência da aba Lista mestra. Se ela estiver em branco, a planilha usa a data de hoje."),
    ("Documento obsoleto", "Continua na lista, marcado, para que ninguém reaproveite o código. Sai dos postos de trabalho."),
    ("Registros", "A retenção é contada em meses, a partir da data em que o registro é gerado. Doze meses valem um ano."),
    ("Tipos e códigos", "Os tipos, as siglas e a regra de código são uma convenção deste modelo. A norma não os define."),
    ("Exemplos", "Os dados das abas de exemplo são ilustrativos, criados para este material."),
]:
    line(k, text)
setup(ws, INK, f"B1:C{r}", landscape=False, fit_height=True)

# ------------------------------------------------------------------ Lista mestra
ws = wb.create_sheet("Lista mestra")
widths(ws, {"A": 2, "B": 5, "C": 13, "D": 36, "E": 8, "F": 22, "G": 10, "H": 13, "I": 20, "J": 20, "K": 12, "L": 22, "M": 10, "N": 13, "O": 12, "P": 20, "Q": 2})
title(ws, "Lista mestra de documentos", "Um documento por linha. A próxima revisão e a situação são calculadas. Os formulários em branco entram aqui; os preenchidos vão para a aba Registros.", "P")
for rr, rot, fmt, val in [(4, "Organização", None, None), (5, "Data de referência", DATE, None), (6, "Intervalo padrão de revisão, em meses", None, REV_MESES)]:
    label(ws, f"B{rr}", rot, merge=f"B{rr}:C{rr}")
    inp(ws, f"D{rr}", val, fmt=fmt, h="left")
    ws.row_dimensions[rr].height = 21.75
dv_date(ws, "D5")
dv_inteiro(ws, "D6", 120, "Intervalo usado quando a linha não informa outro.")
put(ws, "E5", "Data em que a lista é conferida. Em branco, vale a data de hoje.", f=font(9, i=True, c=MUTED), bg=GRAY, merge="E5:H5")
put(ws, "E6", "Usado nas linhas sem intervalo próprio.", f=font(9, i=True, c=MUTED), bg=GRAY, merge="E6:H6")
HEAD = ["#", "Código", "Título", "Tipo", "Processo", "Revisão", "Data da revisão", "Elaborou", "Aprovou", "Meio", "Local da versão em vigor", "Revisar a cada (meses)",
        "Próxima revisão", "Status", "Situação"]
for col, text in zip("BCDEFGHIJKLMNOP", HEAD):
    head(ws, f"{col}7", text)
ws.row_dimensions[7].height = 33
hint_row(ws, 8, [("B", ""), ("C", "Tipo-processo-nº"), ("D", "Como está no cabeçalho"), ("E", "Sigla"), ("F", "Do mapa de processos"), ("G", "Número ou data"),
                 ("H", "Da aprovação"), ("I", "Nome ou cargo"), ("J", "Nome ou cargo"), ("K", "Escolha na lista"), ("L", "Onde está a versão em vigor"),
                 ("M", "Em branco: padrão"), ("N", "Calculada"), ("O", "Escolha na lista"), ("P", "Calculada")])
assert L1 == 9
for k in range(NDOC):
    rr = L1 + k
    num(ws, f"B{rr}", k + 1)
    inp(ws, f"C{rr}", fmt="@")
    inp(ws, f"D{rr}")
    inp(ws, f"E{rr}", h="center")
    inp(ws, f"F{rr}")
    inp(ws, f"G{rr}", h="center", fmt="@")
    inp(ws, f"H{rr}", h="center", fmt=DATE)
    inp(ws, f"I{rr}")
    inp(ws, f"J{rr}")
    inp(ws, f"K{rr}", h="center")
    inp(ws, f"L{rr}")
    inp(ws, f"M{rr}", h="center")
    calc(ws, f"N{rr}", f_proxima(rr, padrao="$D$6"), fmt=DATE, b=False)
    inp(ws, f"O{rr}", h="center")
    calc(ws, f"P{rr}", f_situacao(rr, 'IF($D$5="",TODAY(),$D$5)'), b=False, sz=9)
    ws.row_dimensions[rr].height = 24
dv_list(ws, f"E{L1}:E{L2}", SIGLAS, "Sigla do tipo: " + ", ".join(SIGLAS))
dv_date(ws, f"H{L1}:H{L2}")
dv_list(ws, f"K{L1}:K{L2}", MEIOS, "Papel, Eletrônico ou Ambos")
dv_inteiro(ws, f"M{L1}:M{L2}", 120, "Meses entre uma revisão e outra. Em branco, vale o padrão.")
dv_list(ws, f"O{L1}:O{L2}", STATUS, "Vigente, Em revisão ou Obsoleto")
cf_situacao(ws, f"P{L1}:P{L2}", f"P{L1}")
ws.conditional_formatting.add(f"O{L1}:O{L2}", FormulaRule(formula=[f'O{L1}="{OBS}"'], font=Font(name="Arial", size=10, color=MUTED, strike=True)))
note(ws, "C7", "Tipo, processo e número de ordem: PR-SUP-01.\nO código não muda com a revisão e nunca é reaproveitado.")
note(ws, "E7", "\n".join(f"{s} · {n}" for s, n, _, _ in TIPOS))
note(ws, "O7", "Em revisão: a versão nova está sendo escrita, e a atual continua em vigor.\nObsoleto: não vale mais. Fica na lista para não reaproveitar o código.")
s = L2 + 2
band(ws, s, "Resumo automático", "F")
put(ws, f"H{s}", "Documentos por situação", f=font(10, True, c=WHITE), bg=INK, box=False, merge=f"H{s}:L{s}")
for k, (text, formula, fmt) in enumerate([
    ("Documentos cadastrados", f"=COUNTA(D{L1}:D{L2})", None),
    ("Em vigor, sem pendência", f'=COUNTIF(P{L1}:P{L2},"{S_VIG}")+COUNTIF(P{L1}:P{L2},"{S_EMREV}")', None),
    ("Percentual em vigor, sem os obsoletos", f'=IF(E{s+1}-COUNTIF(P{L1}:P{L2},"{S_OBS}")=0,"",E{s+2}/(E{s+1}-COUNTIF(P{L1}:P{L2},"{S_OBS}")))', "0%"),
    ("Com revisão vencida", f'=COUNTIF(P{L1}:P{L2},"{S_VENC}")', None),
    ("Sem controle: falta código, revisão ou aprovação", f'=COUNTIF(P{L1}:P{L2},"Falta*")', None),
    ("Aviso", f'=IF(E{s+1}=0,"Cadastre os documentos",IF(E{s+5}>0,"Há documento sem controle: veja a coluna Situação",'
              f'IF(E{s+4}>0,"Há documento com a revisão vencida","OK")))', None),
], 1):
    label(ws, f"B{s+k}", text, merge=f"B{s+k}:D{s+k}", h="right")
    calc(ws, f"E{s+k}", formula, fmt=fmt, merge=f"E{s+k}:F{s+k}", sz=9 if text == "Aviso" else 10, b=text != "Aviso")
    ws.row_dimensions[s + k].height = 21.75
cf_warn(ws, f"E{s+6}:F{s+6}", f"E{s+6}")
for k, t in enumerate(SITS, 1):
    label(ws, f"H{s+k}", t, merge=f"H{s+k}:K{s+k}")
    calc(ws, f"L{s+k}", f'=COUNTIF(P{L1}:P{L2},"{t}")')
    ws.row_dimensions[s + k].height = 21.75
barras(ws, f"H{s+9}", s + 1, s + 7, 8, 12, width=17, height=7.2)
ws.freeze_panes = "E9"
setup(ws, BLUE, f"B1:P{s+24}", fit_height=True)

# ------------------------------------------------------------------ Registros
ws = wb.create_sheet("Registros")
widths(ws, {"A": 2, "B": 5, "C": 30, "D": 18, "E": 20, "F": 18, "G": 24, "H": 12, "I": 10, "J": 20, "K": 22, "L": 24, "M": 2})
title(ws, "Tabela de retenção de registros", "Um registro por linha: onde é gerado, onde fica guardado, por quanto tempo, quem pode ver e o que se faz no fim do prazo.", "L")
for col, text in zip("BCDEFGHIJKL", ["#", "Registro", "Formulário ou origem", "Processo", "Gerado em", "Guardado em", "Meio", "Retenção (meses)", "Acesso",
                                      "Disposição no fim do prazo", "Conferência"]):
    head(ws, f"{col}4", text)
ws.row_dimensions[4].height = 33
hint_row(ws, 5, [("B", ""), ("C", "O que prova"), ("D", "Código do formulário ou sistema"), ("E", "Do mapa de processos"), ("F", "Posto ou sistema"),
                 ("G", "Pasta, armário ou sistema"), ("H", "Escolha na lista"), ("I", "12 = um ano"), ("J", "Quem pode ver"), ("K", "Descarte ou arquivo"), ("L", "Calculada")])
ex(ws, "B6", "Ex.", h="center")
for col, v, h_ in [("C", "Planilha de temperatura da câmara fria", "left"), ("D", "FR-02", "left"), ("E", "Comprar e armazenar insumos", "left"), ("F", "Câmara fria", "left"),
                   ("G", "Pasta da cozinha", "left"), ("H", "Papel", "center"), ("I", 12, "center"), ("J", "Pizzaiolo líder", "left"), ("K", "Picotar", "left"),
                   ("L", "Completa", "center")]:
    ex(ws, f"{col}6", v, h=h_)
ws.row_dimensions[6].height = 24
for k in range(NREG):
    rr = R1 + k
    num(ws, f"B{rr}", k + 1)
    for col in "CDEFG":
        inp(ws, f"{col}{rr}")
    inp(ws, f"H{rr}", h="center")
    inp(ws, f"I{rr}", h="center")
    inp(ws, f"J{rr}")
    inp(ws, f"K{rr}")
    calc(ws, f"L{rr}", f'=IF(C{rr}="","",IF(G{rr}="","Falta o local de guarda",IF(I{rr}="","Falta o prazo de retenção",IF(H{rr}="","Falta o meio",'
         f'IF(K{rr}="","Falta a disposição","Completa")))))', b=False, sz=9)
    ws.row_dimensions[rr].height = 24
dv_list(ws, f"H{R1}:H{R2}", MEIOS, "Papel, Eletrônico ou Ambos")
dv_inteiro(ws, f"I{R1}:I{R2}", 600, "Prazo de retenção, em meses. Doze meses valem um ano.")
cf_falta(ws, f"L{R1}:L{R2}", f"L{R1}")
note(ws, "I4", "O maior entre o prazo legal, o do cliente e o do próprio uso.\nSem exigência, um ciclo de certificação: 36 meses.")
s = R2 + 2
band(ws, s, "Resumo automático", "L")
for k, (text, formula, fmt) in enumerate([
    ("Registros cadastrados", f"=COUNTA(C{R1}:C{R2})", None),
    ("Em papel, ou nos dois meios", f'=COUNTIF(H{R1}:H{R2},"Papel")+COUNTIF(H{R1}:H{R2},"Ambos")', None),
    ("Maior prazo de retenção, em meses", f'=IF(COUNT(I{R1}:I{R2})=0,"",MAX(I{R1}:I{R2}))', None),
    ("Aviso", f'=IF(E{s+1}=0,"Cadastre os registros",IF(COUNTIF(L{R1}:L{R2},"Falta*")>0,"Há registro com campos em branco: veja a coluna Conferência","OK"))', None),
], 1):
    label(ws, f"B{s+k}", text, merge=f"B{s+k}:D{s+k}", h="right")
    calc(ws, f"E{s+k}", formula, fmt=fmt, merge=f"E{s+k}:F{s+k}", sz=9 if text == "Aviso" else 10, b=text != "Aviso")
    ws.row_dimensions[s + k].height = 21.75
cf_warn(ws, f"E{s+4}:F{s+4}", f"E{s+4}")
ws.freeze_panes = "D7"
setup(ws, AMBER, f"B1:L{s+4}", fit_height=True)

# ------------------------------------------------------------------ Externos
ws = wb.create_sheet("Externos")
widths(ws, {"A": 2, "B": 5, "C": 34, "D": 18, "E": 16, "F": 22, "G": 30, "H": 18, "I": 13, "J": 10, "K": 13, "L": 20, "M": 2})
title(ws, "Documentos de origem externa", "Normas, leis, desenhos de cliente e manuais que a organização precisa seguir. A pergunta é: existe versão nova?", "L")
for col, text in zip("BCDEFGHIJKL", ["#", "Documento", "Origem", "Versão em uso", "Onde se aplica", "Como se sabe da atualização", "Responsável", "Verificado em",
                                      "Verificar a cada (meses)", "Próxima verificação", "Situação"]):
    head(ws, f"{col}4", text)
ws.row_dimensions[4].height = 33
hint_row(ws, 5, [("B", ""), ("C", "Nome e assunto"), ("D", "Quem publica"), ("E", "Edição, revisão ou data"), ("F", "Processos que usam"), ("G", "Site, aviso, contrato"),
                 ("H", "Nome ou cargo"), ("I", "Última conferência"), ("J", "Em branco: padrão"), ("K", "Calculada"), ("L", "Calculada")])
ex(ws, "B6", "Ex.", h="center")
for col, v, h_, fmt in [("C", "Desenhos e especificações do cliente", "left", None), ("D", "Cada cliente", "left", None), ("E", "Revisão informada no desenho", "left", None),
                        ("F", "Desenvolvimento e Produção", "left", None), ("G", "O cliente envia a revisão nova", "left", None), ("H", "Gerente de engenharia", "left", None),
                        ("I", date(2027, 2, 20), "center", DATE), ("J", 3, "center", None), ("K", date(2027, 5, 20), "center", DATE), ("L", "Verificado", "center", None)]:
    ex(ws, f"{col}6", v, h=h_, fmt=fmt)
ws.row_dimensions[6].height = 24
for k in range(NEXT):
    rr = X1 + k
    num(ws, f"B{rr}", k + 1)
    for col in "CDEFGH":
        inp(ws, f"{col}{rr}")
    inp(ws, f"I{rr}", h="center", fmt=DATE)
    inp(ws, f"J{rr}", h="center")
    calc(ws, f"K{rr}", f_proxima(rr, data="I", meses="J", titulo="C"), fmt=DATE, b=False)
    calc(ws, f"L{rr}", f'=IF(C{rr}="","",IF(I{rr}="","Falta a verificação",IF(H{rr}="","Falta o responsável",IF(K{rr}<IF({REF}="",TODAY(),{REF}),"Verificação vencida","Verificado"))))',
         b=False, sz=9)
    ws.row_dimensions[rr].height = 24
dv_date(ws, f"I{X1}:I{X2}")
dv_inteiro(ws, f"J{X1}:J{X2}", 120, "Meses entre uma verificação e outra. Em branco, vale o padrão da lista mestra.")
ws.conditional_formatting.add(f"L{X1}:L{X2}", FormulaRule(formula=[f'L{X1}="Verificado"'], stopIfTrue=True, fill=PatternFill("solid", bgColor=GREEN, fgColor=GREEN)))
ws.conditional_formatting.add(f"L{X1}:L{X2}", FormulaRule(formula=[f"LEN(L{X1})>0"], fill=PatternFill("solid", bgColor=YELLOW, fgColor=YELLOW)))
note(ws, "G4", "Como a organização fica sabendo que saiu uma versão nova: consulta ao site, aviso do cliente, assinatura de serviço.")
s = X2 + 2
band(ws, s, "Resumo automático", "L")
for k, (text, formula) in enumerate([
    ("Documentos externos cadastrados", f"=COUNTA(C{X1}:C{X2})"),
    ("Verificados no prazo", f'=COUNTIF(L{X1}:L{X2},"Verificado")'),
    ("Com verificação vencida", f'=COUNTIF(L{X1}:L{X2},"Verificação vencida")'),
    ("Aviso", f'=IF(E{s+1}=0,"Cadastre os documentos externos",IF(COUNTIF(L{X1}:L{X2},"Falta*")>0,"Há documento externo com campos em branco: veja a coluna Situação",'
              f'IF(E{s+3}>0,"Há documento com a verificação vencida","OK")))'),
], 1):
    label(ws, f"B{s+k}", text, merge=f"B{s+k}:D{s+k}", h="right")
    calc(ws, f"E{s+k}", formula, merge=f"E{s+k}:F{s+k}", sz=9 if text == "Aviso" else 10, b=text != "Aviso")
    ws.row_dimensions[s + k].height = 21.75
cf_warn(ws, f"E{s+4}:F{s+4}", f"E{s+4}")
ws.freeze_panes = "D7"
setup(ws, TEAL, f"B1:L{s+4}", fit_height=True)

# ------------------------------------------------------------------ Alterações
ws = wb.create_sheet("Alterações")
widths(ws, {"A": 2, "B": 5, "C": 13, "D": 13, "E": 10, "F": 44, "G": 26, "H": 18, "I": 18, "J": 14, "K": 30, "L": 2})
title(ws, "Histórico de alterações", "Uma linha para cada revisão de documento: o que mudou, por quê, quem aprovou, e se os documentos ligados foram conferidos.", "K")
for col, text in zip("BCDEFGHIJK", ["#", "Data", "Código", "Revisão nova", "O que mudou", "Motivo", "Solicitou", "Aprovou", "Ligados conferidos?", "Conferência"]):
    head(ws, f"{col}4", text)
ws.row_dimensions[4].height = 33
hint_row(ws, 5, [("B", ""), ("C", "Da aprovação"), ("D", "Da lista mestra"), ("E", "Número ou data"), ("F", "Em uma ou duas frases"), ("G", "Ação corretiva, mudança, revisão periódica"),
                 ("H", "Nome ou cargo"), ("I", "Nome ou cargo"), ("J", "Escolha na lista"), ("K", "Calculada")])
ex(ws, "B6", "Ex.", h="center")
for col, v, h_, fmt in [("C", date(2026, 10, 16), "center", DATE), ("D", "PR-SUP-01", "left", "@"), ("E", "6", "center", "@"),
                        ("F", "Cadastro do item antes da compra, devolução da requisição incompleta e prazo contado da requisição completa.", "left", None),
                        ("G", "RNC 2026-31", "left", None), ("H", "Gerente de Suprimentos", "left", None), ("I", "Diretor geral", "left", None), ("J", "Sim", "center", None),
                        ("K", "Completa", "center", None)]:
    ex(ws, f"{col}6", v, h=h_, fmt=fmt)
ws.row_dimensions[6].height = 33
for k in range(NALT):
    rr = A1 + k
    num(ws, f"B{rr}", k + 1)
    inp(ws, f"C{rr}", h="center", fmt=DATE)
    inp(ws, f"D{rr}", fmt="@")
    inp(ws, f"E{rr}", h="center", fmt="@")
    inp(ws, f"F{rr}")
    inp(ws, f"G{rr}")
    inp(ws, f"H{rr}")
    inp(ws, f"I{rr}")
    inp(ws, f"J{rr}", h="center")
    calc(ws, f"K{rr}", f'=IF(D{rr}="","",IF(COUNTIF({LISTA}!$C${L1}:$C${L2},D{rr})=0,"Código fora da lista mestra",IF(E{rr}="","Falta a revisão",'
         f'IF(F{rr}="","Falta a descrição",IF(I{rr}="","Falta a aprovação",IF(J{rr}="","Falta conferir os ligados",IF(J{rr}="Não","Ligados não conferidos","Completa")))))))',
         b=False, sz=9)
    ws.row_dimensions[rr].height = 30
dv_date(ws, f"C{A1}:C{A2}")
dv_list(ws, f"J{A1}:J{A2}", LIGADOS, "Sim, Não ou Não há")
cf_falta(ws, f"K{A1}:K{A2}", f"K{A1}")
note(ws, "J4", "Roteiros, telas, formulários e treinamentos que dependem do documento alterado foram conferidos?\nEscolha “Não há” quando o documento não tem ligados.")
s = A2 + 2
band(ws, s, "Resumo automático", "K")
for k, (text, formula) in enumerate([
    ("Alterações registradas", f"=COUNTA(D{A1}:D{A2})"),
    ("Com os documentos ligados conferidos", f'=COUNTIF(J{A1}:J{A2},"Sim")+COUNTIF(J{A1}:J{A2},"Não há")'),
    ("Aviso", f'=IF(F{s+1}=0,"Registre as alterações",IF(COUNTIF(K{A1}:K{A2},"Código fora*")>0,"Há alteração de documento que não está na lista mestra",'
              f'IF(COUNTIF(K{A1}:K{A2},"Falta*")>0,"Há alteração com o registro incompleto: veja a coluna Conferência",'
              f'IF(COUNTIF(K{A1}:K{A2},"Ligados não conferidos")>0,"Há alteração sem a conferência dos documentos ligados","OK"))))'),
], 1):
    label(ws, f"B{s+k}", text, merge=f"B{s+k}:E{s+k}", h="right")
    calc(ws, f"F{s+k}", formula, sz=9 if text == "Aviso" else 10, b=text != "Aviso")
    ws.row_dimensions[s + k].height = 21.75
cf_warn(ws, f"F{s+3}", f"F{s+3}")
ws.freeze_panes = "E7"
setup(ws, PURPLE, f"B1:K{s+3}", fit_height=True)

# ------------------------------------------------------------------ Checklist
ws = wb.create_sheet("Checklist")
widths(ws, {"A": 2, "B": 5, "C": 70, "D": 14, "E": 44, "F": 2})
title(ws, "Checklist do controle de documentos", "Escolha o status de cada item na lista suspensa (coluna em amarelo-claro).", "E")
for col, text in zip("BCDE", ["#", "Item de verificação", "Status", "Observações"]):
    head(ws, f"{col}4", text)
ws.row_dimensions[4].height = 21.75
for k, text in enumerate(CHECK):
    rr = 5 + k
    num(ws, f"B{rr}", k + 1)
    put(ws, f"C{rr}", text, bg=GRAY)
    inp(ws, f"D{rr}", h="center")
    inp(ws, f"E{rr}")
    ws.row_dimensions[rr].height = 30
dv_list(ws, "D5:D16", YESNO, "Sim, Parcial ou Não")
cf_equal(ws, "D5:D16", YESNO_CF)
band(ws, 18, "Resumo automático", "E")
summary(ws, 19, "Itens atendidos (Sim)", '=COUNTIF(D5:D16,"Sim")', "C", "D")
summary(ws, 20, "Itens parciais", '=COUNTIF(D5:D16,"Parcial")', "C", "D")
summary(ws, 21, "Itens não atendidos", '=COUNTIF(D5:D16,"Não")', "C", "D")
summary(ws, 22, "Itens sem resposta", "=COUNTBLANK(D5:D16)", "C", "D")
summary(ws, 23, "Percentual atendido", "=D19/ROWS(D5:D16)", "C", "D", fmt="0%")
summary(ws, 24, "Situação", '=IF(D19=ROWS(D5:D16),"Validado",IF(D22>0,"Em andamento","Com pendências"))', "C", "D")
setup(ws, PURPLE, "B1:E24", landscape=False, fit_height=True)


# ------------------------------------------------------------------ Exemplos
def sub(ws, rr, cells):
    for a, b, t in cells:
        put(ws, f"{a}{rr}", t, f=font(10, True), bg=GRAY, h="center", merge=b and f"{a}{rr}:{b}{rr}")
    ws.row_dimensions[rr].height = 30


def exemplo(ws, data):
    widths(ws, {"A": 2, "B": 12, "C": 30, "D": 8, "E": 20, "F": 9, "G": 12, "H": 18, "I": 18, "J": 11, "K": 20, "L": 8, "M": 12, "N": 11, "O": 18, "P": 2})
    title(ws, "Informação documentada", "Exemplo preenchido, para consulta. Use as abas Lista mestra, Registros, Externos e Alterações para os seus documentos.", "O")
    H = data["head"]
    rr = 4
    for rot, val, fmt in [("Organização", H["org"], None), ("Data de referência", H["data"], DATE), ("Participantes", H["por"], None), ("Regra de revisão", H["regra"], None),
                          ("Origem", H["origem"], None)]:
        label(ws, f"B{rr}", rot, merge=f"B{rr}:C{rr}")
        inp(ws, f"D{rr}", val, merge=f"D{rr}:O{rr}", bg=WHITE, fmt=fmt, h="left")
        ws.row_dimensions[rr].height = 19.5
        rr += 1
    rr += 1
    band(ws, rr, "Lista mestra", "O", color=BLUE)
    rr += 1
    sub(ws, rr, [("B", None, "Código"), ("C", None, "Título"), ("D", None, "Tipo"), ("E", None, "Processo"), ("F", None, "Revisão"), ("G", None, "Data"), ("H", None, "Elaborou"),
                 ("I", None, "Aprovou"), ("J", None, "Meio"), ("K", None, "Local"), ("L", None, "Meses"), ("M", None, "Próxima"), ("N", None, "Status"), ("O", None, "Situação")])
    d1 = rr + 1
    for d in data["docs"]:
        rr += 1
        put(ws, f"B{rr}", d["codigo"] or None, f=font(10, True), fmt="@")
        put(ws, f"C{rr}", d["titulo"], f=font(10, True))
        put(ws, f"D{rr}", d["tipo"], h="center")
        put(ws, f"E{rr}", d["processo"])
        put(ws, f"F{rr}", d["rev"] or None, h="center", fmt="@")
        put(ws, f"G{rr}", d["data"], h="center", fmt=DATE)
        put(ws, f"H{rr}", d["elaborou"])
        put(ws, f"I{rr}", d["aprovou"] or None)
        put(ws, f"J{rr}", d["meio"], h="center")
        put(ws, f"K{rr}", d["onde"])
        put(ws, f"L{rr}", d["meses"], h="center")
        calc(ws, f"M{rr}", f_proxima(rr, data="G", meses="L", titulo="C", padrao=str(REV_MESES)), fmt=DATE, b=False)
        put(ws, f"N{rr}", d["status"], h="center")
        calc(ws, f"O{rr}", f_situacao(rr, "$D$5", titulo="C", codigo="B", rev="F", data="G", aprovou="I", proxima="M", status="N"), b=False, sz=9)
        ws.row_dimensions[rr].height = alt([(d["titulo"], 30), (d["processo"], 20), (d["elaborou"], 18), (d["onde"], 20)], minimo=24)
    d2 = rr
    cf_situacao(ws, f"O{d1}:O{d2}", f"O{d1}")
    rr += 2
    ws.row_breaks.append(Break(id=rr - 1))
    band(ws, rr, "Documentos por situação", "E")
    c0 = rr
    for k, t in enumerate(SITS, 1):
        label(ws, f"B{rr+k}", t, merge=f"B{rr+k}:D{rr+k}", h="right")
        calc(ws, f"E{rr+k}", f'=COUNTIF(O{d1}:O{d2},"{t}")')
        ws.row_dimensions[rr + k].height = 21.75
    label(ws, f"B{rr+8}", "Em vigor, sem os obsoletos", merge=f"B{rr+8}:D{rr+8}", h="right")
    calc(ws, f"E{rr+8}", f"=(E{rr+1}+E{rr+2})/(SUM(E{rr+1}:E{rr+7})-E{rr+7})", fmt="0%")
    ws.row_dimensions[rr + 8].height = 21.75
    barras(ws, f"F{c0 + 1}", c0 + 1, c0 + 7, 2, 5, width=17, height=6.6)
    rr += 11
    band(ws, rr, "Tabela de retenção de registros", "O", color=AMBER)
    rr += 1
    sub(ws, rr, [("B", "C", "Registro"), ("D", "E", "Formulário ou origem"), ("F", "G", "Processo"), ("H", None, "Gerado em"), ("I", None, "Guardado em"), ("J", None, "Meio"),
                 ("K", None, "Acesso"), ("L", "M", "Retenção (meses)"), ("N", "O", "Disposição")])
    r1 = rr + 1
    for g in data["regs"]:
        rr += 1
        put(ws, f"B{rr}", g["nome"], f=font(10, True), merge=f"B{rr}:C{rr}")
        put(ws, f"D{rr}", g["form"], merge=f"D{rr}:E{rr}")
        put(ws, f"F{rr}", g["processo"], merge=f"F{rr}:G{rr}")
        put(ws, f"H{rr}", g["gerado"])
        put(ws, f"I{rr}", g["guardado"])
        put(ws, f"J{rr}", g["meio"], h="center")
        put(ws, f"K{rr}", g["acesso"])
        put(ws, f"L{rr}", g["meses"], h="center", merge=f"L{rr}:M{rr}")
        put(ws, f"N{rr}", g["disposicao"], merge=f"N{rr}:O{rr}")
        ws.row_dimensions[rr].height = alt([(g["nome"], 40), (g["guardado"], 18), (g["acesso"], 20), (g["disposicao"], 28)], minimo=24)
    r2 = rr
    rr += 2
    band(ws, rr, "Resumo automático", "O")
    for k, (text, formula, fmt) in enumerate([
        ("Documentos cadastrados", f"=COUNTA(C{d1}:C{d2})", None),
        ("Documentos sem controle", f'=COUNTIF(O{d1}:O{d2},"Falta*")', None),
        ("Documentos com a revisão vencida", f'=COUNTIF(O{d1}:O{d2},"{S_VENC}")', None),
        ("Registros cadastrados", f"=COUNTA(B{r1}:B{r2})", None),
        ("Maior prazo de retenção, em meses", f"=MAX(L{r1}:L{r2})", None),
    ], 1):
        label(ws, f"B{rr+k}", text, merge=f"B{rr+k}:D{rr+k}", h="right")
        calc(ws, f"E{rr+k}", formula, fmt=fmt)
        ws.row_dimensions[rr + k].height = 21.75
    setup(ws, MUTED, f"B1:O{rr + 5}")


exemplo(wb.create_sheet("Exemplo 1 - Pizzaria"), EX1)
exemplo(wb.create_sheet("Exemplo 2 - Compras"), EX2)

wb.active = 1
wb.save(OUT)
print("ok", OUT)
