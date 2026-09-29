# -*- coding: utf-8 -*-
"""Gera 5W2H-modelo.xlsx no mesmo padrão visual das planilhas de SIPOC, PDCA, SWOT e GUT."""
import os
import sys
from datetime import date

_here = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, _here)
# funções de formatação compartilhadas com o gerador do PDCA (bloco inicial do arquivo)
_src = open(os.path.join(_here, "build_pdca.py"), encoding="utf-8").read()
exec(_src[: _src.index("STATUS = [")])

from openpyxl.chart import BarChart, Reference  # noqa: E402
from openpyxl.chart.label import DataLabelList  # noqa: E402
from openpyxl.chart.series import DataPoint  # noqa: E402
from w5_data import EX1, EX2  # noqa: E402

W, WT, HH, HT = "2B5C8A", "DCE8F3", "A96A12", "F6E8CF"
TEAL, PURPLE = "1E7B73", "7A4A9A"
MONEY = '"R$" #,##0.00'
YESNO = ["Sim", "Parcial", "Não"]
YESNO_CF = [("Sim", GREEN), ("Parcial", YELLOW), ("Não", RED)]
STATUS = ["Não iniciada", "Em andamento", "Concluída", "Cancelada"]
STATUS_CF = [("Concluída", GREEN), ("Em andamento", YELLOW), ("Não iniciada", RED)]
SIT_CF = [("Concluída", GREEN), ("No prazo", GREEN), ("Concluída com atraso", YELLOW), ("Sem prazo", YELLOW), ("Atrasada", RED)]
N = 15
R1, R2 = 15, 15 + N - 1
PLAN = "'5W2H'"


def note(ws, ref, text):
    c = Comment(text, "Modelo 5W2H")
    c.width, c.height = 280, 110
    ws[ref].comment = c


# colunas do plano: (coluna, título em inglês, título em português, dica, cor, tom)
COLS = [
    ("C", "WHAT", "O quê", "A ação: verbo no infinitivo + objeto", W, WT),
    ("E", "WHY", "Por quê", "A causa atacada ou o resultado buscado", W, WT),
    ("F", "WHERE", "Onde", "Local, área ou processo", W, WT),
    ("G", "WHEN", "Início", "Data de início", W, WT),
    ("H", None, "Prazo", "Data de conclusão", W, WT),
    ("I", "WHO", "Quem", "Uma pessoa ou um cargo", W, WT),
    ("J", "HOW", "Como", "O método, em uma frase", HH, HT),
    ("K", "HOW MUCH", "Quanto custa", "Valor, mesmo que zero", HH, HT),
]
PLAN_W = {"A": 2, "B": 5, "C": 17, "D": 17, "E": 30, "F": 18, "G": 12, "H": 12, "I": 18, "J": 34, "K": 16, "L": 16, "M": 2}


def situacao(p):
    if p["status"] == "Concluída":
        return "Concluída com atraso" if p["fim"] and p["fim"] > p["prazo"] else "Concluída"
    return p["status"]


def plan_sheet(ws, tab, data=None):
    widths(ws, PLAN_W)
    is_ex = data is not None
    n = len(data["itens"]) if is_ex else N
    r2 = R1 + n - 1
    title(ws, "5W2H — Plano de ação",
          "Exemplo preenchido, para consulta. Use a aba 5W2H para o seu plano." if is_ex
          else "Preencha as células em amarelo-claro. As células cinza são calculadas automaticamente.", "L")
    hb = WHITE if is_ex else INPUT
    H = data["head"] if is_ex else {}
    for r, l1, k1, l2, k2, f2 in [(4, "Plano", "plano", "Responsável pelo plano", "resp", None),
                                  (5, "Área / unidade", "area", "Data", "data", DATE),
                                  (6, "Elaborado por", "autor", "Versão", "versao", None),
                                  (8, "Origem do plano", "origem", "Orçamento aprovado", "orcamento", MONEY)]:
        label(ws, f"B{r}", l1, merge=f"B{r}:C{r}")
        inp(ws, f"D{r}", H.get(k1), merge=f"D{r}:F{r}", bg=hb)
        label(ws, f"G{r}", l2, merge=f"G{r}:H{r}")
        inp(ws, f"I{r}", H.get(k2), merge=f"I{r}:L{r}", bg=hb, fmt=f2)
        ws.row_dimensions[r].height = 21.75
    for r, l1, k1 in [(7, "Objetivo do plano", "objetivo"), (9, "Fora do escopo", "fora")]:
        label(ws, f"B{r}", l1, merge=f"B{r}:C{r}")
        inp(ws, f"D{r}", H.get(k1), merge=f"D{r}:L{r}", bg=hb)
        ws.row_dimensions[r].height = 30
    ws.row_dimensions[6].height = 30
    dv_date(ws, "I5")
    dv_number(ws, "I8")

    # faixa com as perguntas em inglês
    put(ws, "B11", "", bg=INK, box=False)
    big = lambda ref, text, color, merge=None: put(ws, ref, text, f=font(13, True, c=WHITE), bg=color, h="center", box=False, merge=merge)
    big("C11", "WHAT", W, "C11:D11")
    big("E11", "WHY", W)
    big("F11", "WHERE", W)
    big("G11", "WHEN", W, "G11:H11")
    big("I11", "WHO", W)
    big("J11", "HOW", HH)
    put(ws, "K11", "HOW MUCH", f=font(11, True, c=WHITE), bg=HH, h="center", box=False)
    put(ws, "L11", "", bg=INK, box=False)
    ws.row_dimensions[11].height = 30
    head(ws, "B12", "#")
    put(ws, "C12", "O quê", f=font(10, True, c=WHITE), bg=W, h="center", merge="C12:D12")
    for col, _, pt, _, color, _ in COLS[1:]:
        put(ws, f"{col}12", pt, f=font(10, True, c=WHITE), bg=color, h="center")
    head(ws, "L12", "Conferência")
    ws.row_dimensions[12].height = 21.75
    hint = lambda ref, text, bg, merge=None: put(ws, ref, text, f=font(9, i=True, c=MUTED), bg=bg, h="center", merge=merge)
    hint("B13", "", GRAY)
    hint("C13", COLS[0][3], WT, "C13:D13")
    for col, _, _, tip, _, tint in COLS[1:]:
        hint(f"{col}13", tip, tint)
    hint("L13", "Respostas da linha", GRAY)
    ws.row_dimensions[13].height = 30
    if is_ex:
        put(ws, "B14", "Leia cada linha da esquerda para a direita: a ação, o motivo, o local, as datas, o responsável, o método e o custo. "
            "Cada ação tem um único responsável e datas do calendário.", f=font(9, i=True, c=MUTED), bg=GRAY, merge="B14:L14")
    else:
        ex(ws, "B14", "Ex.", h="center")
        put(ws, "C14", "Agrupar os pedidos por bairro antes da saída", f=font(9, i=True, c=MUTED), bg=GRAY, merge="C14:D14")
        ex(ws, "E14", "Os pedidos saem um a um, e as rotas se cruzam")
        ex(ws, "F14", "Expedição da loja")
        ex(ws, "G14", date(2026, 10, 26), h="center", fmt=DATE)
        ex(ws, "H14", date(2026, 10, 30), h="center", fmt=DATE)
        ex(ws, "I14", "Líder da expedição")
        ex(ws, "J14", "Separar os pedidos em três zonas, com uma prateleira para cada uma")
        ex(ws, "K14", 350, h="center", fmt=MONEY)
        ex(ws, "L14", "Completa", h="center")
        note(ws, "D4", 'Nome curto do plano, com verbo + objeto.\nEx.: "Reduzir os atrasos nas entregas".')
        note(ws, "D7", 'O resultado que o plano busca, com indicador e meta, quando houver.\nEx.: "Aumentar de 82% para 95% as entregas em até 40 minutos".')
        note(ws, "D8", 'De onde veio a necessidade do plano.\nEx.: "Ciclo PDCA", "Auditoria interna", "Análise crítica".')
        note(ws, "I8", "Valor disponível para o plano. Digite só o número. O resumo compara com o custo previsto.")
        note(ws, "C12", "A ação, escrita com verbo no infinitivo e uma entrega que possa ser conferida.\nEvite: melhorar, conscientizar, rever.")
        note(ws, "I12", "Um único responsável por ação: nome ou cargo.\nEvite: equipe, todos, a definir.")
        note(ws, "K12", "Custo estimado, com materiais, serviços e horas. Digite 0 quando a ação não tiver custo adicional.")
    ws.row_dimensions[14].height = 33
    for k in range(n):
        r = R1 + k
        p = data["itens"][k] if is_ex and k < len(data["itens"]) else None
        put(ws, f"B{r}", f"A{k+1}", f=font(10, True, c=MUTED), bg=GRAY, h="center")
        vals = dict(C=p["oque"], E=p["porque"], F=p["onde"], G=p["inicio"], H=p["prazo"], I=p["quem"], J=p["como"], K=p["custo"]) if p else {}
        for col, _, _, _, _, tint in COLS:
            bg = (tint if p else WHITE) if is_ex else INPUT
            kw = dict(bg=bg)
            if col == "C":
                kw["merge"] = f"C{r}:D{r}"
            if col in "GH":
                kw.update(h="center", fmt=DATE)
            if col == "K":
                kw.update(h="center", fmt=MONEY)
            inp(ws, f"{col}{r}", vals.get(col), **kw)
        cnt = f"COUNTA(C{r},E{r}:K{r})"
        calc(ws, f"L{r}", f'=IF(C{r}="","",IF(AND(G{r}<>"",H{r}<>"",G{r}>H{r}),"Datas invertidas",'
             f'IF({cnt}=8,"Completa","Faltam "&(8-{cnt}))))', b=False, sz=9)
        ws.row_dimensions[r].height = 45 if is_ex else 36
    dv_date(ws, f"G{R1}:H{r2}")
    dv_number(ws, f"K{R1}:K{r2}")
    cf_equal(ws, f"L{R1}:L{r2}", [("Completa", GREEN), ("Datas invertidas", RED)])
    ws.conditional_formatting.add(f"L{R1}:L{r2}", FormulaRule(formula=[f'LEFT(L{R1},6)="Faltam"'],
                                  fill=PatternFill("solid", bgColor=YELLOW, fgColor=YELLOW)))

    s = r2 + 2
    band(ws, s, "Resumo automático", "L")

    def row(k, text, formula, fmt=None, sz=10, b=True):
        label(ws, f"B{s+k}", text, merge=f"B{s+k}:F{s+k}", h="right")
        calc(ws, f"G{s+k}", formula, fmt=fmt, merge=f"G{s+k}:J{s+k}", sz=sz, b=b)
        ws.row_dimensions[s + k].height = 21.75

    row(1, "Ações registradas", f"=COUNTA(C{R1}:C{r2})")
    row(2, "Ações completas", f'=COUNTIF(L{R1}:L{r2},"Completa")')
    row(3, "Custo total previsto", f"=SUM(K{R1}:K{r2})", fmt=MONEY)
    row(4, "Saldo do orçamento", f'=IF(I8="","Informe o orçamento",I8-G{s+3})', fmt=MONEY)
    row(5, "Primeiro início", f'=IF(COUNT(G{R1}:G{r2})=0,"",MIN(G{R1}:G{r2}))', fmt=DATE)
    row(6, "Último prazo", f'=IF(COUNT(H{R1}:H{r2})=0,"",MAX(H{R1}:H{r2}))', fmt=DATE)
    row(7, "Aviso: respostas", f'=IF(G{s+1}=0,"Liste as ações",IF(G{s+2}<G{s+1},"Há ação com resposta em branco ou com datas invertidas","OK"))', sz=9, b=False)
    row(8, "Aviso: orçamento", f'=IF(I8="","Informe o orçamento aprovado",IF(G{s+3}>I8,"Custo previsto acima do orçamento","OK"))', sz=9, b=False)
    row(9, "Campos do cabeçalho preenchidos", "=COUNTA(D4,I4,D5,I5,D6,I6,D7,D8,I8,D9)", fmt='0" de 10"')
    cf_ok(ws, f"G{s+7}:J{s+8}")
    end = s + 9

    if is_ex:
        # acompanhamento
        ws.row_breaks.append(Break(id=s - 1))  # resumo, acompanhamento e verificação na segunda página
        band(ws, end + 2, "Acompanhamento das ações", "L", color=TEAL)
        r = end + 3
        head(ws, f"B{r}", "#", bg=GRAY, fg=INK)
        put(ws, f"C{r}", "Ação", f=font(10, True), bg=GRAY, h="center", merge=f"C{r}:D{r}")
        for col, text in zip("EFGHIJK", ["Evidência de conclusão", "Status", "Prazo", "Conclusão", "Situação", "Custo previsto e realizado", "Desvio"]):
            put(ws, f"{col}{r}", text, f=font(10, True), bg=GRAY, h="center")
        ws.row_dimensions[r].height = 21.75
        a1 = r + 1
        for p in data["itens"]:
            r += 1
            put(ws, f"B{r}", p["k"], f=font(10, True, c=MUTED), bg=GRAY, h="center")
            put(ws, f"C{r}", p["curto"], merge=f"C{r}:D{r}")
            put(ws, f"E{r}", p["evidencia"])
            put(ws, f"F{r}", p["status"], h="center")
            put(ws, f"G{r}", p["prazo"], h="center", fmt=DATE)
            put(ws, f"H{r}", p["fim"], h="center", fmt=DATE)
            put(ws, f"I{r}", situacao(p), bg=GRAY, h="center")
            money = lambda v: ("R$ %s" % f"{v:,.0f}".replace(",", ".")) if v else "sem custo"
            put(ws, f"J{r}", f"Previsto: {money(p['custo'])}. Realizado: {money(p['real'])}.")
            put(ws, f"K{r}", p["real"] - p["custo"], bg=GRAY, h="center", fmt='+"R$" #,##0;-"R$" #,##0;"R$" 0')
            ws.row_dimensions[r].height = 30
        cf_equal(ws, f"F{a1}:F{r}", STATUS_CF)
        cf_equal(ws, f"I{a1}:I{r}", SIT_CF)
        r += 1
        label(ws, f"B{r}", "Custo realizado", merge=f"B{r}:I{r}", h="right")
        calc(ws, f"J{r}", sum(p["real"] for p in data["itens"]), fmt=MONEY, merge=f"J{r}:K{r}")
        ws.row_dimensions[r].height = 21.75
        # verificação
        band(ws, r + 2, "Verificação do resultado", "L", color=PURPLE)
        r += 3
        head(ws, f"B{r}", "#", bg=GRAY, fg=INK)
        put(ws, f"C{r}", "Objetivo", f=font(10, True), bg=GRAY, h="center", merge=f"C{r}:D{r}")
        for col, text in zip("EFGHI", ["Indicador", "Situação inicial", "Meta", "Resultado", "Meta atingida?"]):
            put(ws, f"{col}{r}", text, f=font(10, True), bg=GRAY, h="center")
        put(ws, f"J{r}", "Observações", f=font(10, True), bg=GRAY, h="center", merge=f"J{r}:L{r}")
        ws.row_dimensions[r].height = 21.75
        v1 = r + 1
        for k, (obj, ind, ini, meta, sentido, res, quando, obs) in enumerate(data["verif"], 1):
            r += 1
            ok = res <= meta if sentido == "Menor é melhor" else res >= meta
            num(ws, f"B{r}", k)
            put(ws, f"C{r}", obj, merge=f"C{r}:D{r}")
            put(ws, f"E{r}", ind)
            put(ws, f"F{r}", ini, h="center")
            put(ws, f"G{r}", meta, h="center")
            put(ws, f"H{r}", res, h="center")
            put(ws, f"I{r}", "Sim" if ok else "Não", bg=GRAY, h="center")
            put(ws, f"J{r}", obs or "Verificado em " + quando.strftime("%d/%m/%Y"), merge=f"J{r}:L{r}")
            ws.row_dimensions[r].height = 30
        cf_equal(ws, f"I{v1}:I{r}", [("Sim", GREEN), ("Não", RED)])
        end = r
    setup(ws, tab, f"B1:L{end}")
    if not is_ex:
        ws.print_title_rows = "11:12"


# ================================================================== pasta de trabalho
wb = Workbook()
wb.properties.title = "5W2H — Modelo de plano de ação"
wb.properties.creator = "Modelo 5W2H"
wb.properties.language = "pt-BR"

# ------------------------------------------------------------------ Instruções
ws = wb.active
ws.title = "Instruções"
widths(ws, {"A": 2, "B": 24, "C": 92, "D": 2})
title(ws, "5W2H — Como usar esta planilha",
      "Modelo para montar, acompanhar e verificar um plano de ação com as sete perguntas.", "C")
r = 4


def section(text):
    global r
    band(ws, r, text, "C", sz=11)
    r += 1


def line(k, v, kbg=GRAY, vbg=None, kf=None, kh="left", height=19.5):
    global r
    put(ws, f"B{r}", k, f=kf or font(10, True), bg=kbg, h=kh)
    put(ws, f"C{r}", v, bg=vbg)
    ws.row_dimensions[r].height = height
    r += 1


section("Legenda: onde preencher")
line("Amarelo-claro", "Células de entrada. É aqui que você digita.", vbg=INPUT)
line("Cinza", "Células calculadas ou fixas (rótulos, contagens, avisos). Não altere.", vbg=GRAY)
line("Cinza em itálico", "Linha de exemplo (marcada com “Ex.”). Mostra o formato esperado e não entra nas contagens.", vbg=GRAY)
line("Listas suspensas", "Colunas de status e de sentido aceitam apenas as opções da lista. Datas e valores são conferidos na digitação.")
line("Comentários", "Células com triângulo vermelho no canto trazem uma dica de preenchimento. Passe o mouse sobre elas.", height=31.5)
r += 1
section("As sete perguntas")
for k, text, color, tint in [
    ("What", "O quê: a ação, escrita com verbo no infinitivo e uma entrega que possa ser conferida.", W, WT),
    ("Why", "Por quê: a causa que a ação ataca ou o resultado que ela busca.", W, WT),
    ("Where", "Onde: o local, a área ou o processo em que a ação acontece.", W, WT),
    ("When", "Quando: a data de início e o prazo de conclusão, no calendário.", W, WT),
    ("Who", "Quem: a pessoa que responde pela conclusão da ação. Uma só por ação.", W, WT),
    ("How", "Como: o método, em uma frase que acrescente algo ao “o quê”.", HH, HT),
    ("How much", "Quanto custa: o custo estimado, com materiais, serviços e horas. Zero também é resposta.", HH, HT),
]:
    line(k, text, kbg=color, vbg=tint, kf=font(10, True, c=WHITE), kh="center")
r += 1
section("Ordem recomendada de preenchimento")
for k, text in enumerate([
    "Aba 5W2H: preencha o cabeçalho, com o nome, o objetivo, a origem e o orçamento do plano.",
    "Aba 5W2H: escreva as ações (o quê) e o motivo de cada uma (por quê).",
    "Aba 5W2H: descreva o método (como). Divida as ações que não cabem em uma frase.",
    "Aba 5W2H: defina o responsável (quem), as datas (quando) e o local (onde).",
    "Aba 5W2H: estime o custo (quanto) e confira o saldo do orçamento no resumo.",
    "Aba Verificação: registre os indicadores do objetivo, com a situação inicial e a meta.",
    "Aba Checklist: valide o plano com os responsáveis antes de começar a execução.",
    "Aba Acompanhamento: atualize o status, a data de conclusão, o custo e a evidência a cada reunião.",
    "Aba Verificação: registre o resultado dos indicadores e leia a conclusão sobre a eficácia.",
], 1):
    line(f"Passo {k}", text)
r += 1
section("Abas da planilha")
for k, text in [
    ("5W2H", "Modelo principal. Cabeçalho do plano, até 15 ações com as sete respostas e resumo automático."),
    ("Acompanhamento", "Status, conclusão, custo realizado e evidência de cada ação, com situação calculada e gráfico."),
    ("Verificação", "Indicadores do objetivo, com situação inicial, meta, resultado e conclusão sobre a eficácia."),
    ("Checklist", "Doze verificações de qualidade do plano, com percentual de conclusão."),
    ("Exemplo 1 - Pizzaria", "Plano preenchido com uma ação concluída com atraso e meta atingida."),
    ("Exemplo 2 - Compras", "Plano preenchido com um piloto e uma meta não atingida."),
]:
    line(k, text)
r += 1
section("Regras e premissas do modelo")
for k, text in [
    ("Conferência", "A coluna de conferência conta oito respostas por ação: o quê, por quê, onde, início, prazo, quem, como e quanto. Ela também avisa quando o início é posterior ao prazo."),
    ("Custo", "Digite o custo só com números. Quando a ação não tiver custo adicional, digite 0: a célula em branco conta como resposta que falta."),
    ("Situação", "A situação é calculada. Ação concluída depois do prazo aparece como “Concluída com atraso”. Ação não concluída com prazo anterior à data de hoje aparece como “Atrasada”."),
    ("Verificação", "O formato 5W2H não tem pergunta sobre a avaliação do resultado. A aba Verificação cobre esse ponto: compara o resultado de cada indicador com a meta, conforme o sentido informado."),
    ("Capacidade", "A planilha comporta 15 ações e 8 indicadores. Se faltar espaço, divida o plano por objetivo ou por área."),
    ("Exemplos", "Os dados das abas de exemplo são ilustrativos, criados para este material."),
]:
    line(k, text, height=45.75 if len(text) > 100 else 31.5)
setup(ws, INK, f"B1:C{r}", landscape=False, fit_height=True)

# ------------------------------------------------------------------ 5W2H
plan_sheet(wb.create_sheet("5W2H"), W)

# ------------------------------------------------------------------ Acompanhamento
ws = wb.create_sheet("Acompanhamento")
widths(ws, {"A": 2, "B": 5, "C": 40, "D": 18, "E": 12, "F": 15, "G": 13, "H": 15, "I": 34, "J": 21, "K": 15, "L": 2})
title(ws, "Acompanhamento das ações",
      "As colunas cinza vêm da aba 5W2H. Atualize as colunas em amarelo-claro a cada reunião de acompanhamento.", "K")
for col, text in zip("BCDEFGHIJK", ["#", "Ação (o quê)", "Quem", "Prazo", "Status", "Conclusão", "Custo realizado",
                                     "Evidência de conclusão", "Situação", "Desvio de custo"]):
    head(ws, f"{col}4", text)
ws.row_dimensions[4].height = 21.75
ex(ws, "B5", "Ex.", h="center")
ex(ws, "C5", "Agrupar os pedidos por bairro antes da saída")
ex(ws, "D5", "Líder da expedição")
ex(ws, "E5", date(2026, 10, 30), h="center", fmt=DATE)
ex(ws, "F5", "Concluída", h="center")
ex(ws, "G5", date(2026, 11, 4), h="center", fmt=DATE)
ex(ws, "H5", 420, h="center", fmt=MONEY)
ex(ws, "I5", "Fotos da expedição e registro das saídas")
ex(ws, "J5", "Concluída com atraso", h="center")
ex(ws, "K5", 70, h="center", fmt=MONEY)
ws.row_dimensions[5].height = 31.5
A1, A2 = 6, 6 + N - 1
for k in range(N):
    rr, src = A1 + k, R1 + k
    put(ws, f"B{rr}", f"A{k+1}", f=font(10, True, c=MUTED), bg=GRAY, h="center")
    calc(ws, f"C{rr}", f'=IF({PLAN}!C{src}="","",{PLAN}!C{src})', h="left", b=False)
    calc(ws, f"D{rr}", f'=IF({PLAN}!I{src}="","",{PLAN}!I{src})', h="left", b=False)
    calc(ws, f"E{rr}", f'=IF({PLAN}!H{src}="","",{PLAN}!H{src})', b=False, fmt=DATE)
    inp(ws, f"F{rr}", h="center")
    inp(ws, f"G{rr}", h="center", fmt=DATE)
    inp(ws, f"H{rr}", h="center", fmt=MONEY)
    inp(ws, f"I{rr}")
    calc(ws, f"J{rr}", f'=IF(C{rr}="","",IF(F{rr}="Concluída",IF(AND(G{rr}<>"",E{rr}<>"",G{rr}>E{rr}),"Concluída com atraso","Concluída"),'
         f'IF(F{rr}="Cancelada","Cancelada",IF(E{rr}="","Sem prazo",IF(E{rr}<TODAY(),"Atrasada","No prazo")))))', b=False)
    calc(ws, f"K{rr}", f'=IF(OR(C{rr}="",H{rr}=""),"",H{rr}-N({PLAN}!K{src}))', b=False, fmt='+"R$" #,##0.00;-"R$" #,##0.00;"R$" 0.00')
    ws.row_dimensions[rr].height = 30
dv_list(ws, f"F{A1}:F{A2}", STATUS, "Não iniciada, Em andamento, Concluída ou Cancelada")
dv_date(ws, f"G{A1}:G{A2}")
dv_number(ws, f"H{A1}:H{A2}")
cf_equal(ws, f"F{A1}:F{A2}", STATUS_CF)
cf_equal(ws, f"J{A1}:J{A2}", SIT_CF)
s = A2 + 2
band(ws, s, "Resumo automático", "K")
J = f"J{A1}:J{A2}"
rows = [
    ("Ações no plano", f'=SUMPRODUCT(--(C{A1}:C{A2}<>""))', None),
    ("Concluídas", f'=COUNTIF({J},"Concluída")+COUNTIF({J},"Concluída com atraso")', None),
    ("Concluídas com atraso", f'=COUNTIF({J},"Concluída com atraso")', None),
    ("No prazo", f'=COUNTIF({J},"No prazo")', None),
    ("Atrasadas", f'=COUNTIF({J},"Atrasada")', None),
    ("Canceladas", f'=COUNTIF({J},"Cancelada")', None),
    ("Percentual concluído", f"=IF(F{s+1}-F{s+6}<=0,0,F{s+2}/(F{s+1}-F{s+6}))", "0%"),
    ("Custo previsto", f"=SUM({PLAN}!K{R1}:K{R2})", MONEY),
    ("Custo realizado", f"=SUM(H{A1}:H{A2})", MONEY),
]
for k, (text, formula, fmt) in enumerate(rows, 1):
    summary(ws, s + k, text, formula, "E", "F", fmt=fmt, val_merge=f"F{s+k}:H{s+k}")
k = len(rows) + 1
summary(ws, s + k, "Aviso: conclusão",
        f'=IF(SUMPRODUCT((F{A1}:F{A2}="Concluída")*((G{A1}:G{A2}="")+(I{A1}:I{A2}="")>0))>0,"Há ação concluída sem data ou sem evidência","OK")',
        "E", "F", val_merge=f"F{s+k}:H{s+k}")
summary(ws, s + k + 1, "Aviso: status",
        f'=IF(SUMPRODUCT((C{A1}:C{A2}<>"")*(F{A1}:F{A2}=""))>0,"Há ação sem status","OK")', "E", "F", val_merge=f"F{s+k+1}:H{s+k+1}")
for rr in (s + k, s + k + 1):
    ws[f"F{rr}"].font = font(9)
cf_ok(ws, f"F{s+k}:H{s+k+1}")
t = s + k + 3
band(ws, t, "Ações por situação", "K")
put(ws, f"B{t+1}", "Situação", f=font(10, True), bg=GRAY, h="center", merge=f"B{t+1}:C{t+1}")
put(ws, f"D{t+1}", "Ações", f=font(10, True), bg=GRAY, h="center")
SITS = [("Concluída", "1F7A4D"), ("Concluída com atraso", HH), ("No prazo", W), ("Atrasada", "B0413E"), ("Sem prazo", MUTED), ("Cancelada", "9AA7B0")]
for j, (name, _) in enumerate(SITS):
    rr = t + 2 + j
    put(ws, f"B{rr}", name, f=font(10, True), bg=GRAY, merge=f"B{rr}:C{rr}")
    calc(ws, f"D{rr}", f'=COUNTIF({J},"{name}")')
    ws.row_dimensions[rr].height = 19.5
ch = BarChart()
ch.type = "col"
ch.height, ch.width = 7, 17
ch.style = 2
ch.title = None
ch.legend = None
ch.gapWidth = 120
ch.add_data(Reference(ws, min_col=4, min_row=t + 1, max_row=t + 1 + len(SITS)), titles_from_data=True)
ch.set_categories(Reference(ws, min_col=2, min_row=t + 2, max_row=t + 1 + len(SITS)))
srs = ch.series[0]
srs.graphicalProperties.solidFill = MUTED
for j, (_, color) in enumerate(SITS):
    pt = DataPoint(idx=j)
    pt.graphicalProperties = GraphicalProperties(solidFill=color)
    pt.graphicalProperties.line.noFill = True
    srs.dPt.append(pt)
srs.dLbls = DataLabelList()
srs.dLbls.showVal = True
srs.dLbls.showSerName = srs.dLbls.showCatName = srs.dLbls.showLegendKey = False
ch.x_axis.delete = False
ch.y_axis.delete = False
ch.y_axis.scaling.min = 0
ch.y_axis.majorUnit = 1
ch.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill="D3DBE0", w=9525))
ch.x_axis.graphicalProperties = GraphicalProperties(ln=LineProperties(solidFill="D3DBE0", w=9525))
ch.y_axis.graphicalProperties = GraphicalProperties(ln=LineProperties(noFill=True))
ws.add_chart(ch, f"F{t+1}")
ws.freeze_panes = "D5"
setup(ws, TEAL, f"B1:K{t+16}", fit_height=True)

# ------------------------------------------------------------------ Verificação
ws = wb.create_sheet("Verificação")
widths(ws, {"A": 2, "B": 5, "C": 34, "D": 34, "E": 15, "F": 12, "G": 18, "H": 12, "I": 13, "J": 15, "K": 34, "L": 2})
title(ws, "Verificação do resultado",
      "Registre os indicadores do objetivo do plano. A verificação é feita depois que as ações foram concluídas.", "K")
for col, text in zip("BCDEFGHIJK", ["#", "Objetivo ou causa atacada", "Indicador", "Situação inicial", "Meta", "Sentido",
                                     "Resultado", "Data", "Meta atingida?", "Observações"]):
    head(ws, f"{col}4", text)
ws.row_dimensions[4].height = 21.75
ex(ws, "B5", "Ex.", h="center")
ex(ws, "C5", "Entregas no prazo")
ex(ws, "D5", "% de entregas em até 40 minutos")
ex(ws, "E5", 82, h="center")
ex(ws, "F5", 95, h="center")
ex(ws, "G5", "Maior é melhor", h="center")
ex(ws, "H5", 95.5, h="center")
ex(ws, "I5", date(2026, 11, 30), h="center", fmt=DATE)
ex(ws, "J5", "Sim", h="center")
ex(ws, "K5", "Média das quatro semanas depois das ações")
ws.row_dimensions[5].height = 31.5
V1, V2 = 6, 13
for k in range(8):
    rr = V1 + k
    num(ws, f"B{rr}", k + 1)
    inp(ws, f"C{rr}")
    inp(ws, f"D{rr}")
    for col in "EFH":
        inp(ws, f"{col}{rr}", h="center")
    inp(ws, f"G{rr}", h="center")
    inp(ws, f"I{rr}", h="center", fmt=DATE)
    calc(ws, f"J{rr}", f'=IF(OR(H{rr}="",F{rr}=""),"",IF(G{rr}="Menor é melhor",IF(H{rr}<=F{rr},"Sim","Não"),IF(H{rr}>=F{rr},"Sim","Não")))', b=False)
    inp(ws, f"K{rr}")
    ws.row_dimensions[rr].height = 30
dv_list(ws, f"G{V1}:G{V2}", ["Maior é melhor", "Menor é melhor"], "O indicador melhora quando sobe ou quando desce?")
dv_date(ws, f"I{V1}:I{V2}")
for col in "EFH":
    dv_number(ws, f"{col}{V1}:{col}{V2}")
cf_equal(ws, f"J{V1}:J{V2}", [("Sim", GREEN), ("Não", RED)])
band(ws, 15, "Resumo automático", "K")
summary(ws, 16, "Indicadores registrados", f"=COUNTA(D{V1}:D{V2})", "D", "E", val_merge="E16:H16")
summary(ws, 17, "Indicadores verificados", f"=COUNT(H{V1}:H{V2})", "D", "E", val_merge="E17:H17")
summary(ws, 18, "Metas atingidas", f'=COUNTIF(J{V1}:J{V2},"Sim")', "D", "E", val_merge="E18:H18")
summary(ws, 19, "Metas não atingidas", f'=COUNTIF(J{V1}:J{V2},"Não")', "D", "E", val_merge="E19:H19")
summary(ws, 20, "Ações concluídas no plano", f'=Acompanhamento!F{s+2}&" de "&Acompanhamento!F{s+1}', "D", "E", val_merge="E20:H20")
summary(ws, 21, "Aviso: sentido",
        f'=IF(SUMPRODUCT((F{V1}:F{V2}<>"")*(G{V1}:G{V2}=""))>0,"Há meta sem sentido informado: a comparação supõe que maior é melhor","OK")',
        "D", "E", val_merge="E21:H21")
summary(ws, 22, "Conclusão",
        '=IF(E17=0,"Sem verificação registrada",IF(E19=0,"Plano eficaz: padronize o que funcionou","Há meta não atingida: volte à análise das causas"))',
        "D", "E", val_merge="E22:H22")
ws["E21"].font = font(9)
ws.row_dimensions[21].height = 30
cf_ok(ws, "E21:H21")
cf_equal(ws, "E22:H22", [("Plano eficaz: padronize o que funcionou", GREEN), ("Há meta não atingida: volte à análise das causas", YELLOW)])
setup(ws, PURPLE, "B1:K22", fit_height=True)

# ------------------------------------------------------------------ Checklist
ws = wb.create_sheet("Checklist")
widths(ws, {"A": 2, "B": 5, "C": 70, "D": 14, "E": 44, "F": 2})
title(ws, "Checklist de validação do plano 5W2H", "Escolha o status de cada item na lista suspensa (coluna em amarelo-claro).", "E")
for col, text in zip("BCDE", ["#", "Item de verificação", "Status", "Observações"]):
    head(ws, f"{col}4", text)
ws.row_dimensions[4].height = 21.75
for k, text in enumerate([
    "O objetivo do plano está definido e ligado a um problema, a uma causa ou a uma decisão.",
    "Cada ação começa com verbo no infinitivo e descreve uma entrega que pode ser conferida.",
    "Cada ação tem um motivo: a causa que ataca ou o resultado que busca.",
    "Cada ação tem um único responsável, identificado por nome ou cargo.",
    "Cada ação tem data de início e prazo, no calendário.",
    "O local ou o processo de cada ação está definido.",
    "O método de cada ação diz como fazer, e não só repete o que fazer.",
    "O custo de cada ação foi estimado, mesmo quando é zero.",
    "Os responsáveis participaram da elaboração e concordam com os prazos.",
    "A forma de comprovar a conclusão e de verificar o resultado está definida.",
    "O plano tem rotina de acompanhamento, com data e responsável.",
    "As ações concluídas têm evidência registrada.",
]):
    rr = 5 + k
    num(ws, f"B{rr}", k + 1)
    put(ws, f"C{rr}", text, bg=GRAY)
    inp(ws, f"D{rr}", h="center")
    inp(ws, f"E{rr}")
    ws.row_dimensions[rr].height = 30
dv_list(ws, "D5:D16", YESNO, "Sim, Parcial ou Não")
cf_equal(ws, "D5:D16", YESNO_CF)
band(ws, 18, "Resumo automático", "E")
summary(ws, 19, "Itens atendidos (Sim)", '=COUNTIF(D5:D16,"Sim")', "C", "D")
summary(ws, 20, "Itens parciais", '=COUNTIF(D5:D16,"Parcial")', "C", "D")
summary(ws, 21, "Itens não atendidos", '=COUNTIF(D5:D16,"Não")', "C", "D")
summary(ws, 22, "Itens sem resposta", "=COUNTBLANK(D5:D16)", "C", "D")
summary(ws, 23, "Percentual atendido", "=D19/ROWS(D5:D16)", "C", "D", fmt="0%")
summary(ws, 24, "Situação", '=IF(D19=ROWS(D5:D16),"Validado",IF(D22>0,"Em andamento","Com pendências"))', "C", "D")
setup(ws, HH, "B1:E24", landscape=False, fit_height=True)

# ------------------------------------------------------------------ Exemplos
plan_sheet(wb.create_sheet("Exemplo 1 - Pizzaria"), MUTED, EX1)
plan_sheet(wb.create_sheet("Exemplo 2 - Compras"), MUTED, EX2)

wb.active = 1
wb.save(OUT)
print("ok", OUT, "| resumo do acompanhamento começa na linha", s)
