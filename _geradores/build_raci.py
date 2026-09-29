# -*- coding: utf-8 -*-
"""Gera RACI-modelo.xlsx no mesmo padrão visual das planilhas anteriores da série."""
import os
import sys
from datetime import date

_here = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, _here)
# funções de formatação compartilhadas com o gerador do PDCA (bloco inicial do arquivo)
_src = open(os.path.join(_here, "build_pdca.py"), encoding="utf-8").read()
exec(_src[: _src.index("STATUS = [")])

from openpyxl.chart import BarChart, Reference  # noqa: E402
from raci_data import EX1, EX2  # noqa: E402

CR = {"R": ("1E7B73", "D9EEEB", "Responsável", "executa a atividade ou coordena a execução. Pelo menos um por atividade."),
      "A": ("7A4A9A", "E9DEF1", "Aprovador", "responde pelo resultado e tem autoridade para decidir. Um, e só um, por atividade."),
      "C": ("A96A12", "F6E8CF", "Consultado", "é ouvido antes da decisão ou da execução. Conversa de mão dupla."),
      "I": ("5B6B76", "E3E8EB", "Informado", "recebe a notícia do que foi decidido ou feito. Aviso de mão única.")}
BLUE = "2B5C8A"
YESNO = ["Sim", "Parcial", "Não"]
YESNO_CF = [("Sim", GREEN), ("Parcial", YELLOW), ("Não", RED)]
ROLES = "EFGHIJKL"          # oito colunas de papéis
NA = 15                     # linhas de atividades
R1, R2 = 15, 15 + NA - 1
HDR = 12                    # linha com os nomes dos papéis
CA = R2 + 2                 # faixa da conferência da coluna
RS = CA + 8                 # faixa do resumo
SH = "RACI"


def note(ws, ref, text):
    c = Comment(text, "Modelo RACI")
    c.width, c.height = 280, 110
    ws[ref].comment = c


def cf_letters(ws, rng_):
    for val, color in (("R", CR["R"][0]), ("A", CR["A"][0]), ("A/R", CR["A"][0]), ("C", CR["C"][0]), ("I", CR["I"][0])):
        ws.conditional_formatting.add(rng_, CellIsRule(operator="equal", formula=[f'"{val}"'], font=Font(bold=True, color=WHITE),
                                      fill=PatternFill("solid", bgColor=color, fgColor=color)))


def cf_not_ok(ws, rng_, first):
    ok = PatternFill("solid", bgColor=GREEN, fgColor=GREEN)
    warn = PatternFill("solid", bgColor=YELLOW, fgColor=YELLOW)
    ws.conditional_formatting.add(rng_, FormulaRule(formula=[f'{first}="OK"'], fill=ok))
    ws.conditional_formatting.add(rng_, FormulaRule(formula=[f'LEN({first})>2'], fill=warn))


RACI_W = {"A": 2, "B": 5, "C": 16, "D": 22, **{c: 13 for c in ROLES}, "M": 7, "N": 7, "O": 30, "P": 2}


def raci_sheet(ws, tab, data=None):
    widths(ws, RACI_W)
    is_ex = data is not None
    title(ws, "RACI — Matriz de responsabilidades",
          "Exemplo preenchido, para consulta. Use a aba RACI para a sua matriz." if is_ex
          else "Preencha as células em amarelo-claro. As células cinza são calculadas automaticamente.", "O")
    hb = WHITE if is_ex else INPUT
    H = data["head"] if is_ex else {}
    for r, l1, k1, l2, k2, f2 in [(4, "Processo ou projeto", "processo", "Dono do processo", "dono", None),
                                  (5, "Área / unidade", "area", "Data", "data", DATE),
                                  (6, "Elaborado por", "autor", "Versão", "versao", None),
                                  (8, "Origem das atividades", "origem", "Participantes", "part", None)]:
        label(ws, f"B{r}", l1, merge=f"B{r}:C{r}")
        inp(ws, f"D{r}", H.get(k1), merge=f"D{r}:G{r}", bg=hb)
        label(ws, f"H{r}", l2, merge=f"H{r}:I{r}")
        inp(ws, f"J{r}", H.get(k2), merge=f"J{r}:O{r}", bg=hb, fmt=f2)
        ws.row_dimensions[r].height = 21.75
    for r, l1, k1 in [(7, "Escopo", "escopo"), (9, "Fora do escopo", "fora")]:
        label(ws, f"B{r}", l1, merge=f"B{r}:C{r}")
        inp(ws, f"D{r}", H.get(k1), merge=f"D{r}:O{r}", bg=hb)
        ws.row_dimensions[r].height = 24
    ws.row_dimensions[8].height = 24
    dv_date(ws, "J5")

    put(ws, "B11", "Atividades", f=font(10, True, c=WHITE), bg=INK, box=False, merge="B11:D11")
    put(ws, "E11", "Papéis", f=font(10, True, c=WHITE), bg=INK, box=False, merge="E11:L11")
    put(ws, "M11", "Conferência da linha", f=font(10, True, c=WHITE), bg=INK, box=False, merge="M11:O11")
    ws.row_dimensions[11].height = 21.75
    head(ws, f"B{HDR}", "#")
    put(ws, f"C{HDR}", "Atividade", f=font(10, True, c=WHITE), bg=INK, h="center", merge=f"C{HDR}:D{HDR}")
    papeis = data["papeis"] if is_ex else []
    for j, col in enumerate(ROLES):
        put(ws, f"{col}{HDR}", papeis[j] if j < len(papeis) else None, f=font(9, True), bg=WHITE if is_ex else INPUT, h="center")
    put(ws, f"M{HDR}", "A", f=font(10, True, c=WHITE), bg=CR["A"][0], h="center")
    put(ws, f"N{HDR}", "R", f=font(10, True, c=WHITE), bg=CR["R"][0], h="center")
    head(ws, f"O{HDR}", "Aviso da linha")
    ws.row_dimensions[HDR].height = 39
    hint = lambda ref, text, merge=None: put(ws, ref, text, f=font(9, i=True, c=MUTED), bg=GRAY, h="center", merge=merge)
    hint("B13", "")
    hint("C13", "Verbo no infinitivo + objeto", "C13:D13")
    hint("E13", "Cargo ou função na linha de cima. Nas células: R, A, A/R, C ou I", "E13:L13")
    hint("M13", "Deve ser 1")
    hint("N13", "1 ou mais")
    hint("O13", "Calculado")
    ws.row_dimensions[13].height = 24
    if is_ex:
        put(ws, "B14", "Leia na horizontal para conferir cada atividade: um A e pelo menos um R. "
            "Leia na vertical para conferir a carga de cada papel.", f=font(9, i=True, c=MUTED), bg=GRAY, merge="B14:O14")
    else:
        ex(ws, "B14", "Ex.", h="center")
        put(ws, "C14", "Cotar com fornecedores", f=font(9, i=True, c=MUTED), bg=GRAY, merge="C14:D14")
        for col, v in zip(ROLES, ["C", "A/R", "", "", "", "", "", ""]):
            ex(ws, f"{col}14", v, h="center")
        ex(ws, "M14", 1, h="center")
        ex(ws, "N14", 1, h="center")
        ex(ws, "O14", "OK", h="center")
        note(ws, "D4", 'O processo ou o projeto coberto pela matriz.\nEx.: "Adquirir materiais e serviços".')
        note(ws, "D7", 'Onde o processo começa e onde termina.\nEx.: "Da requisição de compra à liberação da nota fiscal".')
        note(ws, "D8", 'De onde vieram as atividades.\nEx.: "Macroetapas do SIPOC", "Fluxograma do processo".')
        note(ws, f"E{HDR}", "Escreva um cargo ou uma função, e não o nome de uma pessoa.\nEx.: Comprador, Gerente de Suprimentos.")
        note(ws, f"M{HDR}", "Quantidade de A na linha. Cada atividade deve ter um, e só um. A célula A/R conta como A e como R.")
        note(ws, f"N{HDR}", "Quantidade de R na linha. Cada atividade deve ter pelo menos um.")
    ws.row_dimensions[14].height = 24
    linhas = data["linhas"] if is_ex else []
    for k in range(NA):
        r = R1 + k
        atv, cel = linhas[k] if k < len(linhas) else (None, [])
        bg = WHITE if is_ex else INPUT
        num(ws, f"B{r}", k + 1)
        inp(ws, f"C{r}", atv, merge=f"C{r}:D{r}", bg=bg)
        for j, col in enumerate(ROLES):
            inp(ws, f"{col}{r}", (cel[j] or None) if j < len(cel) else None, h="center", bg=bg)
        rg = f"E{r}:L{r}"
        calc(ws, f"M{r}", f'=IF(C{r}="","",COUNTIF({rg},"A")+COUNTIF({rg},"A/R"))')
        calc(ws, f"N{r}", f'=IF(C{r}="","",COUNTIF({rg},"R")+COUNTIF({rg},"A/R"))')
        calc(ws, f"O{r}", f'=IF(C{r}="","",IF(COUNTA({rg})=0,"Distribua os papéis",IF(M{r}=0,"Falta o A",IF(M{r}>1,"Mais de um A",'
             f'IF(N{r}=0,"Falta o R",IF(COUNTIF({rg},"C")>3,"Muitos consultados",IF(COUNTA({rg})=COUNTA($E${HDR}:$L${HDR}),'
             f'"Todos participam: confira","OK")))))))', b=False, sz=9)
        ws.row_dimensions[r].height = 24
        if is_ex and not atv:
            ws.row_dimensions[r].hidden = True
    grid = f"E{R1}:L{R2}"
    dv_list(ws, grid, ["R", "A", "A/R", "C", "I"], "R, A, A/R, C ou I. Deixe em branco se o papel não participa.")
    cf_letters(ws, grid)
    cf_not_ok(ws, f"O{R1}:O{R2}", f"O{R1}")
    ws.conditional_formatting.add(f"M{R1}:M{R2}", FormulaRule(formula=[f'AND(M{R1}<>"",M{R1}<>1)'],
                                  fill=PatternFill("solid", bgColor=RED, fgColor=RED)))
    ws.conditional_formatting.add(f"N{R1}:N{R2}", FormulaRule(formula=[f'AND(N{R1}<>"",N{R1}=0)'],
                                  fill=PatternFill("solid", bgColor=RED, fgColor=RED)))

    # conferência da coluna
    band(ws, CA, "Conferência da coluna: carga por papel", "O")
    rows = [("R", "Atividades em que executa (R)"), ("A", "Atividades em que responde (A)"), ("C", "Atividades em que é consultado (C)"),
            ("I", "Atividades em que é informado (I)")]
    for k, (l, text) in enumerate(rows, 1):
        r = CA + k
        put(ws, f"B{r}", l, f=font(10, True, c=WHITE), bg=CR[l][0], h="center")
        label(ws, f"C{r}", text, merge=f"C{r}:D{r}")
        for col in ROLES:
            rg = f"{col}{R1}:{col}{R2}"
            f = f'COUNTIF({rg},"{l}")' + (f'+COUNTIF({rg},"A/R")' if l in "RA" else "")
            calc(ws, f"{col}{r}", f'=IF(AND({col}${HDR}="",COUNTA({rg})=0),"",{f})')
        ws.row_dimensions[r].height = 21.75
    rt, rv = CA + 5, CA + 6
    label(ws, f"B{rt}", "Total de participações", merge=f"B{rt}:D{rt}")
    label(ws, f"B{rv}", "Aviso da coluna", merge=f"B{rv}:D{rv}")
    natv = f"$H${RS + 1}"
    for col in ROLES:
        rg = f"{col}{R1}:{col}{R2}"
        calc(ws, f"{col}{rt}", f'=IF(AND({col}${HDR}="",COUNTA({rg})=0),"",COUNTA({rg}))')
        calc(ws, f"{col}{rv}", f'=IF({col}${HDR}="",IF(COUNTA({rg})>0,"Papel sem nome",""),IF({col}{rt}=0,"Sem participação",'
             f'IF(AND({natv}>=5,{col}{CA + 2}/{natv}>0.6),"Muitos A: possível gargalo",IF(AND({natv}>=5,{col}{CA + 1}/{natv}>0.6),"Muitos R: possível sobrecarga",'
             f'IF({col}{CA + 1}+{col}{CA + 2}=0,IF({col}{CA + 3}=0,"Só informado","Só consultado ou informado"),"OK")))))', b=False, sz=9)
    ws.row_dimensions[rt].height = 21.75
    ws.row_dimensions[rv].height = 45
    cf_not_ok(ws, f"E{rv}:L{rv}", f"E{rv}")

    band(ws, RS, "Resumo automático", "O")

    def row(k, text, formula, fmt=None):
        label(ws, f"B{RS+k}", text, merge=f"B{RS+k}:G{RS+k}", h="right")
        calc(ws, f"H{RS+k}", formula, fmt=fmt, merge=f"H{RS+k}:L{RS+k}")
        ws.row_dimensions[RS + k].height = 21.75

    O = f"O{R1}:O{R2}"
    row(1, "Atividades listadas", f"=COUNTA(C{R1}:C{R2})")
    row(2, "Papéis definidos", f"=COUNTA(E{HDR}:L{HDR})")
    row(3, "Atividades sem aviso", f'=COUNTIF({O},"OK")')
    row(4, "Atividades sem A ou com mais de um A", f'=COUNTIF({O},"Falta o A")+COUNTIF({O},"Mais de um A")')
    row(5, "Atividades sem R", f'=COUNTIF({O},"Falta o R")')
    row(6, "Papéis com aviso", f'=SUMPRODUCT((E{rv}:L{rv}<>"")*(E{rv}:L{rv}<>"OK"))')
    row(7, "Situação da matriz", f'=IF(H{RS+1}=0,"Liste as atividades",IF(H{RS+2}=0,"Defina os papéis",'
                                 f'IF(H{RS+3}<H{RS+1},"Há atividade com aviso: veja a conferência da linha","Linhas conferidas")))')
    row(8, "Campos do cabeçalho preenchidos", "=COUNTA(D4,J4,D5,J5,D6,J6,D7,D8,J8,D9)", fmt='0" de 10"')
    ws[f"H{RS+7}"].font = font(9, True)
    cf_equal(ws, f"H{RS+7}:L{RS+7}", [("Linhas conferidas", GREEN)])
    ws.conditional_formatting.add(f"H{RS+7}:L{RS+7}", FormulaRule(formula=[f'H{RS+7}<>"Linhas conferidas"'],
                                  fill=PatternFill("solid", bgColor=YELLOW, fgColor=YELLOW)))
    ws.row_breaks.append(Break(id=CA - 1))
    ws.freeze_panes = f"E{HDR + 1}"
    setup(ws, tab, f"B1:O{RS + 8}")
    ws.print_title_rows = f"{HDR}:{HDR}"


# ================================================================== pasta de trabalho
wb = Workbook()
wb.properties.title = "RACI — Modelo de matriz de responsabilidades"
wb.properties.creator = "Modelo RACI"
wb.properties.language = "pt-BR"

# ------------------------------------------------------------------ Instruções
ws = wb.active
ws.title = "Instruções"
widths(ws, {"A": 2, "B": 24, "C": 92, "D": 2})
title(ws, "RACI — Como usar esta planilha",
      "Modelo para definir quem executa, quem responde, quem é consultado e quem é informado em cada atividade.", "C")
r = 4


def section(text):
    global r
    band(ws, r, text, "C", sz=11)
    r += 1


def line(k, v, kbg=GRAY, vbg=None, kf=None, kh="left", height=19.5):
    global r
    put(ws, f"B{r}", k, f=kf or font(10, True), bg=kbg, h=kh)
    put(ws, f"C{r}", v, bg=vbg)
    ws.row_dimensions[r].height = height
    r += 1


section("Legenda: onde preencher")
line("Amarelo-claro", "Células de entrada. É aqui que você digita.", vbg=INPUT)
line("Cinza", "Células calculadas ou fixas (rótulos, contagens, avisos). Não altere.", vbg=GRAY)
line("Cinza em itálico", "Linha de exemplo (marcada com “Ex.”). Mostra o formato esperado e não entra nas contagens.", vbg=GRAY)
line("Listas suspensas", "As células da matriz aceitam apenas R, A, A/R, C ou I. Deixe em branco quando o papel não participa da atividade.")
line("Comentários", "Células com triângulo vermelho no canto trazem uma dica de preenchimento. Passe o mouse sobre elas.", height=31.5)
r += 1
section("As quatro letras")
for l in "RACI":
    line(l, f"{CR[l][2]}: {CR[l][3]}", kbg=CR[l][0], vbg=CR[l][1], kf=font(10, True, c=WHITE), kh="center")
line("A/R", "O mesmo papel executa e responde pelo resultado. Conta como A e como R nas conferências.", kbg=CR["A"][0], vbg=CR["A"][1],
     kf=font(10, True, c=WHITE), kh="center")
r += 1
section("Ordem recomendada de preenchimento")
for k, text in enumerate([
    "Aba RACI: preencha o cabeçalho, com o processo e o escopo: onde começa e onde termina.",
    "Aba RACI: liste as atividades, com verbo no infinitivo e objeto, todas no mesmo nível de detalhe.",
    "Aba RACI: escreva os papéis na linha de cabeçalho da matriz. Use cargos ou funções, e não nomes.",
    "Aba RACI: marque o A de cada atividade. Um por linha.",
    "Aba RACI: marque os R. Use A/R quando o mesmo papel executa e responde.",
    "Aba RACI: marque os C e os I. Na dúvida entre os dois, escolha I.",
    "Aba RACI: confira os avisos de cada linha e de cada coluna. Aba Análise: veja a carga por papel.",
    "Aba Comunicação: registre quem ocupa cada papel, quem substitui e como a matriz foi comunicada.",
    "Aba Checklist: valide a matriz com quem ocupa os papéis.",
], 1):
    line(f"Passo {k}", text)
r += 1
section("Abas da planilha")
for k, text in [
    ("RACI", "Modelo principal. Cabeçalho, matriz com até 15 atividades e 8 papéis, conferências e resumo automático."),
    ("Análise", "Carga de cada papel, com contagem de R, A, C e I, participação, leitura da coluna e gráfico."),
    ("Comunicação", "Quem ocupa cada papel, quem substitui, e como e quando a matriz foi comunicada e entendida."),
    ("Checklist", "Doze verificações de qualidade da matriz, com percentual de conclusão."),
    ("Exemplo 1 - Pizzaria", "Matriz preenchida de um processo operacional, com cinco papéis."),
    ("Exemplo 2 - Compras", "Matriz preenchida de um processo que atravessa seis áreas, com avisos de gargalo e de sobrecarga."),
]:
    line(k, text)
r += 1
section("Regras e premissas do modelo")
for k, text in [
    ("Conferência da linha", "Cada atividade deve ter um único A e pelo menos um R. O aviso também aparece quando há mais de três consultados ou quando todos os papéis participam da atividade."),
    ("Conferência da coluna", "Com cinco atividades ou mais, o aviso aparece quando um papel tem A ou R em mais de 60% delas. O limite é uma referência do modelo: a decisão sobre redistribuir é do grupo."),
    ("Papéis, não pessoas", "As colunas da matriz trazem cargos ou funções. A relação entre o papel e as pessoas que o ocupam fica na aba Comunicação."),
    ("Célula em branco", "Deixar a célula em branco é uma resposta válida: o papel não participa da atividade."),
    ("Capacidade", "A planilha comporta 15 atividades e 8 papéis. Se faltar espaço, divida o processo em duas matrizes."),
    ("Exemplos", "Os dados das abas de exemplo são ilustrativos, criados para este material."),
]:
    line(k, text, height=45.75 if len(text) > 100 else 31.5)
setup(ws, INK, f"B1:C{r}", landscape=False, fit_height=True)

# ------------------------------------------------------------------ RACI
raci_sheet(wb.create_sheet(SH), CR["R"][0])

# ------------------------------------------------------------------ Análise
ws = wb.create_sheet("Análise")
widths(ws, {"A": 2, "B": 5, "C": 28, "D": 9, "E": 9, "F": 9, "G": 9, "H": 10, "I": 15, "J": 36, "K": 2})
title(ws, "Análise — carga por papel", "Esta aba é calculada a partir da aba RACI. Não há nada para digitar aqui.", "J")
band(ws, 4, "Participação de cada papel", "J")
put(ws, "B5", "#", f=font(10, True), bg=GRAY, h="center")
put(ws, "C5", "Papel", f=font(10, True), bg=GRAY, h="center")
for col, l in zip("DEFG", "RACI"):
    put(ws, f"{col}5", l, f=font(10, True, c=WHITE), bg=CR[l][0], h="center")
for col, text in zip("HIJ", ["Total", "Atividades de que participa", "Leitura da coluna"]):
    put(ws, f"{col}5", text, f=font(10, True), bg=GRAY, h="center")
ws.row_dimensions[5].height = 30
Q1, Q2 = 6, 13
natv = f"{SH}!$H${RS + 1}"
for j, col in enumerate(ROLES):
    rr = Q1 + j
    num(ws, f"B{rr}", j + 1)
    calc(ws, f"C{rr}", f'=IF({SH}!{col}{HDR}="","",{SH}!{col}{HDR})', h="left")
    for c2, k in zip("DEFG", (1, 2, 3, 4)):
        calc(ws, f"{c2}{rr}", f'=IF(C{rr}="","",{SH}!{col}{CA + k})', b=False)
    calc(ws, f"H{rr}", f'=IF(C{rr}="","",{SH}!{col}{CA + 5})')
    calc(ws, f"I{rr}", f'=IF(OR(C{rr}="",{natv}=0),"",H{rr}/{natv})', b=False, fmt="0%")
    calc(ws, f"J{rr}", f'=IF(C{rr}="","",{SH}!{col}{CA + 6})', b=False, sz=9)
    ws.row_dimensions[rr].height = 24
cf_not_ok(ws, f"J{Q1}:J{Q2}", f"J{Q1}")
band(ws, 15, "Resumo automático", "J")
summary(ws, 16, "Atividades na matriz", f"={natv}", "C", "D", val_merge="D16:H16")
summary(ws, 17, "Papéis na matriz", f"={SH}!$H${RS + 2}", "C", "D", val_merge="D17:H17")
summary(ws, 18, "Papel com mais R", f'=IF(MAX(D{Q1}:D{Q2})=0,"",INDEX(C{Q1}:C{Q2},MATCH(MAX(D{Q1}:D{Q2}),D{Q1}:D{Q2},0)))', "C", "D", val_merge="D18:H18")
summary(ws, 19, "Papel com mais A", f'=IF(MAX(E{Q1}:E{Q2})=0,"",INDEX(C{Q1}:C{Q2},MATCH(MAX(E{Q1}:E{Q2}),E{Q1}:E{Q2},0)))', "C", "D", val_merge="D19:H19")
summary(ws, 20, "Letras por atividade, em média", f'=IF({natv}=0,0,SUM(H{Q1}:H{Q2})/{natv})', "C", "D", fmt="0.0", val_merge="D20:H20")
summary(ws, 21, "Papéis com aviso", f"={SH}!$H${RS + 6}", "C", "D", val_merge="D21:H21")
band(ws, 23, "Gráfico: letras por papel", "J")
ch = BarChart()
ch.type = "col"
ch.grouping = "stacked"
ch.overlap = 100
ch.height, ch.width = 9, 21
ch.style = 2
ch.title = None
ch.legend.position = "b"
ch.gapWidth = 90
ch.add_data(Reference(ws, min_col=4, max_col=7, min_row=Q1 - 1, max_row=Q2), titles_from_data=True)
ch.set_categories(Reference(ws, min_col=3, min_row=Q1, max_row=Q2))
for srs, l in zip(ch.series, "RACI"):
    srs.graphicalProperties.solidFill = CR[l][0]
    srs.graphicalProperties.line.solidFill = WHITE
ch.x_axis.delete = False
ch.y_axis.delete = False
ch.y_axis.scaling.min = 0
ch.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill="D3DBE0", w=9525))
ch.x_axis.graphicalProperties = GraphicalProperties(ln=LineProperties(solidFill="D3DBE0", w=9525))
ch.y_axis.graphicalProperties = GraphicalProperties(ln=LineProperties(noFill=True))
ws.add_chart(ch, "B24")
for rr in range(24, 43):
    ws.row_dimensions[rr].height = 15
setup(ws, CR["A"][0], "B1:J42", landscape=False, fit_height=True)

# ------------------------------------------------------------------ Comunicação
ws = wb.create_sheet("Comunicação")
widths(ws, {"A": 2, "B": 5, "C": 28, "D": 28, "E": 24, "F": 28, "G": 13, "H": 15, "I": 30, "J": 2})
title(ws, "Comunicação dos papéis",
      "Os papéis vêm da aba RACI. Registre quem ocupa cada um, quem substitui e como a matriz foi comunicada.", "I")
for col, text in zip("BCDEFGHI", ["#", "Papel", "Quem ocupa", "Substituto", "Forma de comunicação", "Data", "Entendido?", "Observações"]):
    head(ws, f"{col}4", text)
ws.row_dimensions[4].height = 21.75
ex(ws, "B5", "Ex.", h="center")
ex(ws, "C5", "Comprador")
ex(ws, "D5", "Três compradores da equipe de Suprimentos")
ex(ws, "E5", "Comprador sênior")
ex(ws, "F5", "Reunião da área, com a matriz projetada")
ex(ws, "G5", date(2026, 10, 30), h="center", fmt=DATE)
ex(ws, "H5", "Sim", h="center")
ex(ws, "I5", "Matriz afixada no quadro da área")
ws.row_dimensions[5].height = 31.5
K1, K2 = 6, 13
for j, col in enumerate(ROLES):
    rr = K1 + j
    num(ws, f"B{rr}", j + 1)
    calc(ws, f"C{rr}", f'=IF({SH}!{col}{HDR}="","",{SH}!{col}{HDR})', h="left")
    for c2 in "DEFI":
        inp(ws, f"{c2}{rr}")
    inp(ws, f"G{rr}", h="center", fmt=DATE)
    inp(ws, f"H{rr}", h="center")
    ws.row_dimensions[rr].height = 30
dv_list(ws, f"H{K1}:H{K2}", YESNO, "Sim, Parcial ou Não")
dv_date(ws, f"G{K1}:G{K2}")
cf_equal(ws, f"H{K1}:H{K2}", YESNO_CF)
band(ws, 15, "Resumo automático", "I")
C = f"C{K1}:C{K2}"
summary(ws, 16, "Papéis na matriz", f'=SUMPRODUCT(--({C}<>""))', "E", "F")
summary(ws, 17, "Papéis com ocupante registrado", f'=SUMPRODUCT(({C}<>"")*(D{K1}:D{K2}<>""))', "E", "F")
summary(ws, 18, "Papéis com substituto", f'=SUMPRODUCT(({C}<>"")*(E{K1}:E{K2}<>""))', "E", "F")
summary(ws, 19, "Papéis comunicados", f'=SUMPRODUCT(({C}<>"")*(G{K1}:G{K2}<>""))', "E", "F")
summary(ws, 20, "Papéis que entenderam (Sim)", f'=COUNTIF(H{K1}:H{K2},"Sim")', "E", "F")
summary(ws, 21, "Situação",
        '=IF(F16=0,"Defina os papéis na aba RACI",IF(F19<F16,"Há papel sem comunicação registrada",IF(F20<F16,"Há papel sem entendimento confirmado","Matriz comunicada e entendida")))',
        "E", "F")
ws["F21"].font = font(9, True)
ws.row_dimensions[21].height = 30
cf_equal(ws, "F21", [("Matriz comunicada e entendida", GREEN)])
ws.conditional_formatting.add("F21", FormulaRule(formula=['F21<>"Matriz comunicada e entendida"'], fill=PatternFill("solid", bgColor=YELLOW, fgColor=YELLOW)))
setup(ws, CR["C"][0], "B1:I21", fit_height=True)

# ------------------------------------------------------------------ Checklist
ws = wb.create_sheet("Checklist")
widths(ws, {"A": 2, "B": 5, "C": 70, "D": 14, "E": 44, "F": 2})
title(ws, "Checklist de validação da Matriz RACI", "Escolha o status de cada item na lista suspensa (coluna em amarelo-claro).", "E")
for col, text in zip("BCDE", ["#", "Item de verificação", "Status", "Observações"]):
    head(ws, f"{col}4", text)
ws.row_dimensions[4].height = 21.75
for k, text in enumerate([
    "O escopo da matriz está definido: um processo ou um projeto, com início e fim.",
    "As atividades estão escritas com verbo no infinitivo e objeto.",
    "As atividades têm o mesmo nível de detalhe.",
    "As colunas trazem papéis ou cargos, e não nomes de pessoas.",
    "Cada atividade tem um único A.",
    "Cada atividade tem pelo menos um R.",
    "Os consultados de cada atividade são só os necessários.",
    "Nenhum papel concentra R ou A a ponto de virar sobrecarga ou gargalo.",
    "Quem tem o A tem autoridade para decidir sobre a atividade.",
    "A matriz foi construída com a participação de quem ocupa os papéis.",
    "A matriz foi comunicada, e as pessoas entendem o próprio papel.",
    "A data da próxima revisão está marcada.",
]):
    rr = 5 + k
    num(ws, f"B{rr}", k + 1)
    put(ws, f"C{rr}", text, bg=GRAY)
    inp(ws, f"D{rr}", h="center")
    inp(ws, f"E{rr}")
    ws.row_dimensions[rr].height = 30
dv_list(ws, "D5:D16", YESNO, "Sim, Parcial ou Não")
cf_equal(ws, "D5:D16", YESNO_CF)
band(ws, 18, "Resumo automático", "E")
summary(ws, 19, "Itens atendidos (Sim)", '=COUNTIF(D5:D16,"Sim")', "C", "D")
summary(ws, 20, "Itens parciais", '=COUNTIF(D5:D16,"Parcial")', "C", "D")
summary(ws, 21, "Itens não atendidos", '=COUNTIF(D5:D16,"Não")', "C", "D")
summary(ws, 22, "Itens sem resposta", "=COUNTBLANK(D5:D16)", "C", "D")
summary(ws, 23, "Percentual atendido", "=D19/ROWS(D5:D16)", "C", "D", fmt="0%")
summary(ws, 24, "Situação", '=IF(D19=ROWS(D5:D16),"Validado",IF(D22>0,"Em andamento","Com pendências"))', "C", "D")
setup(ws, CR["I"][0], "B1:E24", landscape=False, fit_height=True)

# ------------------------------------------------------------------ Exemplos
raci_sheet(wb.create_sheet("Exemplo 1 - Pizzaria"), MUTED, EX1)
raci_sheet(wb.create_sheet("Exemplo 2 - Compras"), MUTED, EX2)

wb.active = 1
wb.save(OUT)
print("ok", OUT, "| conferência da coluna em", CA, "| resumo em", RS)
