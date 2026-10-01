# -*- coding: utf-8 -*-
"""Gera Liberacao-modelo.xlsx no mesmo padrão visual das planilhas anteriores da série."""
import math
import os
import sys

_here = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, _here)
# funções de formatação compartilhadas com o gerador do PDCA (bloco inicial do arquivo)
_src = open(os.path.join(_here, "build_pdca.py"), encoding="utf-8").read()
exec(_src[: _src.index("STATUS = [")])

from openpyxl.chart import BarChart, Series  # noqa: E402
from openpyxl.chart.label import DataLabelList  # noqa: E402
from lib_data import (ABERTO, APOS, AUTORIZADO, CHECK, CONCES, DECISOES, DET, DISP, DISP_INFO, ENCER, EX1, EX2, L_CONF, L_DESV, L_PEND, LIBERADO, R_NAO,  # noqa: E402
                      RETIDO, RETRAB, REPETE, REVERIF, TRAT)

BLUE, TEAL, AMBER, PURPLE, REDC = "2B5C8A", "1E7B73", "A96A12", "7A4A9A", "B0413E"
S1_T, S2_T, S3_T = "DCE8F3", "F8EBCB", "F5DEDC"
YESNO = ["Sim", "Parcial", "Não"]
YESNO_CF = [("Sim", GREEN), ("Parcial", YELLOW), ("Não", RED)]
NL, NP = 60, 60
AU, LB, PN = "Autoridades", "'Liberação'", "'Produto não conforme'"
L1, L2 = 7, 7 + NL - 1          # lotes, na aba Liberação
P1, P2 = 7, 7 + NP - 1          # registros, na aba Produto não conforme
GERAL = "#,##0.##"
REAL = '"R$" #,##0'
REP = "Defeito repetido: avaliar ação corretiva"
DECIDE = [("Liberar o produto", "A função que confere as verificações e libera."),
          ("Autorizar a liberação com verificação pendente", "Quem pode correr o risco, e quando o cliente precisa saber.")] + [(n, o) for n, o, _, _, _ in DISP_INFO]


def note(ws, ref, text):
    c = Comment(text, "Modelo Liberação")
    c.width, c.height = 300, 120
    ws[ref].comment = c


def cf_warn(ws, rng_, first, ok_values=("OK",)):
    cond = ",".join(f'{first}="{v}"' for v in ok_values)
    ws.conditional_formatting.add(rng_, FormulaRule(formula=[f"OR({cond})"], fill=PatternFill("solid", bgColor=GREEN, fgColor=GREEN)))
    ws.conditional_formatting.add(rng_, FormulaRule(formula=[f"AND(LEN({first})>0,NOT(OR({cond})))"],
                                  fill=PatternFill("solid", bgColor=YELLOW, fgColor=YELLOW)))


def cf_texto(ws, rng_, first, pares, resto=None):
    for t, cor in pares:
        ws.conditional_formatting.add(rng_, FormulaRule(formula=[f'{first}="{t}"'], stopIfTrue=True, fill=PatternFill("solid", bgColor=cor, fgColor=cor)))
    if resto:
        ws.conditional_formatting.add(rng_, FormulaRule(formula=[f"LEN({first})>0"], fill=PatternFill("solid", bgColor=resto, fgColor=resto)))


def hint_row(ws, row, cells, height=30):
    for ref, text in cells:
        put(ws, f"{ref}{row}", text, f=font(9, i=True, c=MUTED), bg=GRAY, h="center")
    ws.row_dimensions[row].height = height


def alt(pares, minimo=21.75, linha=12):
    n = max(max(1, math.ceil(len(str(t)) / max(c, 1))) for t, c in pares)
    return max(minimo, n * linha + 7)


# ------------------------------------------------------------------ fórmulas compartilhadas pelas abas de entrada e pelos exemplos
def f_leitura(lote, prev, feitas, fora):
    return (f'IF({lote}="","",IF(OR({prev}="",{feitas}=""),"Falta o número de verificações",IF({feitas}>{prev},"Feitas acima das previstas",'
            f'IF({feitas}<{prev},"{L_PEND}",IF(N({fora})>0,"{L_DESV}","{L_CONF}")))))')


def f_lib_conf(lote, leit, dec, quem, aut, pnc, prev, feitas, fora):
    return (f'IF({lote}="","",IF({dec}="","Falta a decisão",IF(OR(LEFT({leit},5)="Falta",LEFT({leit},5)="Feita"),{leit},'
            f'IF(AND({feitas}<{prev},{dec}="{LIBERADO}"),"Liberado com verificação pendente",IF(AND({dec}="{AUTORIZADO}",{aut}=""),"Falta quem autorizou",'
            f'IF(AND(N({fora})>0,{pnc}=""),"Desvio sem registro de não conforme",IF(AND({dec}="{RETIDO}",{pnc}=""),"Retido sem registro de não conforme",'
            f'IF(AND({dec}<>"{RETIDO}",{quem}=""),"Falta quem liberou","OK"))))))))')


def f_sit(defe, disp, enc):
    return f'IF({defe}="","",IF({disp}="","{ABERTO}",IF({enc}="","{TRAT}","{ENCER}")))'


def f_pnc_conf(c, rep, defs, acoes, vazio):
    """Conferência de um registro; `c` traz as células da linha; `rep`, o número de registros do mesmo defeito."""
    return (f'IF({c["defe"]}="",IF({vazio}>0,"Falta o defeito",""),IF({c["data"]}="","Falta a data",IF({c["det"]}="","Falta onde foi visto",'
            f'IF({c["disp"]}="","Falta a disposição",IF({c["quem"]}="","Falta quem decidiu",IF(AND({c["seg"]}="",{c["det"]}<>"{APOS}"),"Falta a identificação e a segregação",'
            f'IF(AND({c["disp"]}="{CONCES}",{c["cli"]}=""),"Falta a concessão do cliente",IF(AND({c["det"]}="{APOS}",{c["cli"]}=""),"Falta informar o cliente",'
            f'IF(AND({c["disp"]}="{RETRAB}",{c["rev"]}=""),"Falta a reverificação",IF(AND({c["rev"]}="{R_NAO}",{c["enc"]}<>""),"Encerrado com reverificação não conforme",'
            f'IF(AND({rep}>={REPETE},COUNTIFS({defs},{c["defe"]},{acoes},"?*")=0),"{REP}","OK")))))))))))')


LEIT_CF = [(L_CONF, GREEN), (L_PEND, S2_T), (L_DESV, S3_T)]
DEC_CF = [(LIBERADO, GREEN), (AUTORIZADO, S2_T), (RETIDO, S3_T)]
SIT_CF = [(ENCER, GREEN), (TRAT, S2_T), (ABERTO, S3_T)]


def chart_style(ch):
    ch.style = 2
    ch.title = None
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ch.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill="D3DBE0", w=9525))
    ch.x_axis.graphicalProperties = GraphicalProperties(ln=LineProperties(solidFill="D3DBE0", w=9525))
    ch.y_axis.graphicalProperties = GraphicalProperties(ln=LineProperties(noFill=True))


def colunas(ws, anchor, r1, r2, c_cat, c_val, width=18, height=8.5):
    """Colunas com o custo por registro em cada ponto de detecção."""
    ch = BarChart()
    ch.type = "col"
    ch.gapWidth = 80
    ch.height, ch.width = height, width
    chart_style(ch)
    ch.legend = None
    s = Series(Reference(ws, min_col=c_val, min_row=r1, max_row=r2), title="Custo por registro")
    s.graphicalProperties.solidFill = REDC
    s.graphicalProperties.line.noFill = True
    s.dLbls = DataLabelList()
    s.dLbls.showVal = True
    s.dLbls.numFmt = REAL
    s.dLbls.showSerName = s.dLbls.showCatName = s.dLbls.showLegendKey = False
    ch.append(s)
    ch.set_categories(Reference(ws, min_col=c_cat, min_row=r1, max_row=r2))
    ch.y_axis.number_format = REAL
    ch.y_axis.scaling.min = 0
    ws.add_chart(ch, anchor)


# ================================================================== pasta de trabalho
wb = Workbook()
wb.properties.title = "Liberação e produto não conforme — Modelo"
wb.properties.creator = "Modelo Liberação"
wb.properties.language = "pt-BR"

# ------------------------------------------------------------------ Instruções
ws = wb.active
ws.title = "Instruções"
widths(ws, {"A": 2, "B": 26, "C": 90, "D": 2})
title(ws, "Liberação e produto não conforme — Como usar esta planilha",
      "Modelo para registrar a liberação, controlar o produto não conforme, decidir a disposição e ler o custo.", "C")
r = 4


def section(text):
    global r
    band(ws, r, text, "C", sz=11)
    r += 1


def line(k, v, kbg=GRAY, vbg=None, kf=None, kh="left", height=None):
    global r
    put(ws, f"B{r}", k, f=kf or font(10, True), bg=kbg, h=kh)
    put(ws, f"C{r}", v, bg=vbg)
    ws.row_dimensions[r].height = height or (19.5 if len(v) <= 96 else (31.5 if len(v) <= 192 else 45.75))
    r += 1


section("Legenda: onde preencher")
line("Amarelo-claro", "Células de entrada. É aqui que você digita.", vbg=INPUT)
line("Cinza", "Células calculadas ou fixas (rótulos, leituras, situações, avisos). Não altere.", vbg=GRAY)
line("Cinza em itálico", "Linha de exemplo (marcada com “Ex.”). Mostra o formato esperado e não entra nas contagens.", vbg=GRAY)
line("Listas suspensas", "Decisão, onde foi visto, disposição, reverificação, aceite do cliente e checklist aceitam apenas as opções da lista.")
line("Comentários", "Células com triângulo vermelho no canto trazem uma dica de preenchimento. Passe o mouse sobre elas.")
r += 1
section("As regras de cálculo")
for k, text in [
    (L_CONF, "Todas as verificações previstas foram feitas, e nenhum resultado ficou fora do critério."),
    (L_DESV, "Todas foram feitas, e algum resultado ficou fora do critério. Pede o número do registro de não conforme."),
    (L_PEND, "Foram feitas menos verificações do que as previstas. Só pode ser liberado com autorização: escolha a decisão “Liberado com autorização” e informe quem autorizou."),
    ("Situação do registro", f"{ABERTO}: sem disposição. {TRAT}: com disposição, sem data de encerramento. {ENCER}: com a data de encerramento."),
    ("Repetições", "Quantos registros têm o mesmo texto na coluna Defeito. Use sempre o mesmo nome para o mesmo defeito."),
    ("Defeito repetido", f"O mesmo defeito aparece em {REPETE} registros ou mais, e nenhum deles traz uma ação corretiva."),
    ("Custo por ponto de detecção", "A soma do custo dos registros achados no recebimento, no processo, na inspeção final e depois da entrega, dividida pelo número de registros."),
]:
    line(k, text)
r += 1
section("Ordem recomendada de preenchimento")
for k, text in enumerate([
    "Aba Autoridades: escreva quem libera, quem autoriza a liberação pendente e quem decide cada disposição, com os limites.",
    "Aba Liberação: registre cada lote ou pedido, com as verificações previstas e feitas, os resultados fora do critério, a decisão e quem liberou.",
    "Aba Produto não conforme: registre cada caso, com o defeito, onde foi visto, a segregação, a disposição e quem decidiu.",
    "Aba Produto não conforme: complete o aceite ou a comunicação ao cliente, a reverificação, o encerramento, o custo e a ação corretiva, quando houver.",
    "Aba Painel: leia os registros abertos, as liberações a rever, os defeitos repetidos e o custo por ponto de detecção.",
    "Aba Checklist: valide a rotina.",
], 1):
    line(f"Passo {k}", text)
r += 1
section("Abas da planilha")
for k, text in [
    (AU, "Quem decide cada coisa: liberar, autorizar a liberação pendente e as seis disposições. Calcula a conferência de cada linha."),
    ("Liberação", f"Até {NL} lotes ou pedidos. Calcula a leitura do lote e a conferência da decisão."),
    ("Produto não conforme", f"Até {NP} registros. Calcula a situação, as repetições do defeito e a conferência."),
    ("Painel", "Os indicadores, o custo por ponto de detecção e por disposição, o gráfico e o aviso."),
    ("Checklist", "Doze verificações da rotina, com percentual de conclusão."),
    ("Exemplo 1 - Pizzaria", "As noites de 15 a 21/03/2027, com sete liberações e oito registros."),
    ("Exemplo 2 - Indústria", "Os lotes 136 a 143 da extrusora 3, com oito liberações e nove registros."),
]:
    line(k, text)
r += 1
section("Regras e premissas do modelo")
for k, text in [
    ("Verificações", "Conte as verificações planejadas para o lote, no plano de controle e na inspeção final, e quantas foram feitas antes da decisão."),
    ("Lote com parte retida", "Quando só uma parte sai do critério, libere a parte boa e informe o número do registro de não conforme da parte separada."),
    ("Custo", "Some o que o caso custou: material perdido, retrabalho, frete, desconto, reposição. Use zero quando o custo ficou com o fornecedor."),
    ("Exemplos", "Os dados das abas de exemplo são ilustrativos, criados para este material."),
]:
    line(k, text)
setup(ws, INK, f"B1:C{r}", landscape=False, fit_height=True)

# ------------------------------------------------------------------ Autoridades
ws = wb.create_sheet(AU)
widths(ws, {"A": 2, "B": 5, "C": 36, "D": 40, "E": 26, "F": 34, "G": 14, "H": 26, "I": 30, "J": 2})
title(ws, "Quem decide", "Escreva a função, e não o nome da pessoa. Os limites dizem até onde cada função pode decidir sozinha.", "I")
label(ws, "B4", "Organização e produto", merge="B4:C4")
inp(ws, "D4", merge="D4:I4", h="left")
ws.row_dimensions[4].height = 21.75
for col, text in zip("BCDEFGHI", ["#", "Decisão", "O que é", "Quem decide", "Limite ou condição", "O cliente concorda?", "Onde se registra", "Conferência"]):
    head(ws, f"{col}6", text)
ws.row_dimensions[6].height = 33
A1 = 7
for k, (nome, oque) in enumerate(DECIDE):
    rr = A1 + k
    num(ws, f"B{rr}", k + 1)
    put(ws, f"C{rr}", nome, f=font(10, True), bg=GRAY)
    put(ws, f"D{rr}", oque, f=font(9, c=MUTED), bg=GRAY)
    inp(ws, f"E{rr}")
    inp(ws, f"F{rr}")
    inp(ws, f"G{rr}", h="center")
    inp(ws, f"H{rr}")
    regra = f'IF(AND(C{rr}="{CONCES}",G{rr}<>"Sim"),"A concessão pede o aceite do cliente",' if nome == CONCES else "IF(FALSE,\"\","
    calc(ws, f"I{rr}", f'=IF(E{rr}="","Falta quem decide",{regra}IF(H{rr}="","Falta onde se registra","OK")))', b=False, sz=9)
    ws.row_dimensions[rr].height = 33
A2 = A1 + len(DECIDE) - 1
dv_list(ws, f"G{A1}:G{A2}", ["Sim", "Não", "Quando aplicável"], "O cliente precisa concordar com esta decisão?")
cf_warn(ws, f"I{A1}:I{A2}", f"I{A1}")
note(ws, "F6", "Exemplo: “até 500 kg”, “só no turno”, “acima disso, o gerente industrial”.")
s = A2 + 2
band(ws, s, "Resumo automático", "I")
label(ws, f"B{s+1}", "Decisões com responsável", merge=f"B{s+1}:D{s+1}", h="right")
calc(ws, f"E{s+1}", f'=COUNTA(E{A1}:E{A2})&" de {len(DECIDE)}"', merge=f"E{s+1}:I{s+1}", h="left")
label(ws, f"B{s+2}", "Aviso", merge=f"B{s+2}:D{s+2}", h="right")
calc(ws, f"E{s+2}", f'=IF(D4="","Escreva a organização e o produto",IF(COUNTIF(I{A1}:I{A2},"OK")<ROWS(I{A1}:I{A2}),"Há decisão a completar: veja a coluna Conferência","OK"))',
     merge=f"E{s+2}:I{s+2}", sz=9, b=False, h="left")
cf_warn(ws, f"E{s+2}:I{s+2}", f"E{s+2}")
for k in (1, 2):
    ws.row_dimensions[s + k].height = 21.75
UAV = s + 2
setup(ws, PURPLE, f"B1:I{UAV}", fit_height=True)

# ------------------------------------------------------------------ Liberação
ws = wb.create_sheet("Liberação")
widths(ws, {"A": 2, "B": 5, "C": 12, "D": 18, "E": 24, "F": 11, "G": 11, "H": 11, "I": 20, "J": 24, "K": 20, "L": 26, "M": 16, "N": 30, "O": 30, "P": 2})
title(ws, "Registro de liberação", "Uma linha por lote, pedido ou fechamento. A decisão se apoia nas verificações previstas e feitas.", "O")
for col, text in zip("BCDEFGHIJKLMNO", ["#", "Data", "Lote ou pedido", "Produto e quantidade", "Verificações previstas", "Verificações feitas", "Resultados fora", "Leitura", "Decisão",
                                         "Quem liberou", "Quem autorizou", "Nº do registro de não conforme", "Observação", "Conferência"]):
    head(ws, f"{col}4", text)
ws.row_dimensions[4].height = 33
hint_row(ws, 5, [("B", ""), ("C", "Data"), ("D", "A identificação"), ("E", "O que e quanto"), ("F", "Número"), ("G", "Número"), ("H", "Número"), ("I", "Calculada"), ("J", "Lista"),
                 ("K", "Uma função"), ("L", "Só na liberação pendente"), ("M", "Quando houver desvio"), ("N", ""), ("O", "Calculada")])
ex(ws, "B6", "Ex.", h="center")
for col, v, h_, fmt in [("C", date(2027, 5, 20), "center", DATE), ("D", "Lote 139", "left", None), ("E", "4 bobinas, 1.190 kg", "left", None), ("F", 12, "center", None), ("G", 11, "center", None),
                        ("H", 0, "center", None), ("I", L_PEND, "center", None), ("J", AUTORIZADO, "center", None), ("K", "Analista da Qualidade", "left", None),
                        ("L", "Gerente industrial, com o aceite do cliente", "left", None), ("M", None, "center", None), ("N", "Ensaio de solda pendente.", "left", None), ("O", "OK", "center", None)]:
    ex(ws, f"{col}6", v, h=h_, fmt=fmt)
ws.row_dimensions[6].height = 30
for k in range(NL):
    rr = L1 + k
    num(ws, f"B{rr}", k + 1)
    inp(ws, f"C{rr}", h="center", fmt=DATE)
    inp(ws, f"D{rr}")
    inp(ws, f"E{rr}")
    for col in "FGH":
        inp(ws, f"{col}{rr}", h="center", fmt="0")
    calc(ws, f"I{rr}", "=" + f_leitura(f"D{rr}", f"F{rr}", f"G{rr}", f"H{rr}"), sz=9)
    inp(ws, f"J{rr}", h="center")
    for col in "KLMN":
        inp(ws, f"{col}{rr}")
    calc(ws, f"O{rr}", "=" + f_lib_conf(f"D{rr}", f"I{rr}", f"J{rr}", f"K{rr}", f"L{rr}", f"M{rr}", f"F{rr}", f"G{rr}", f"H{rr}"), b=False, sz=9)
    ws.row_dimensions[rr].height = 21.75
dv_date(ws, f"C{L1}:C{L2}")
dv_number(ws, f"F{L1}:H{L2}")
dv_list(ws, f"J{L1}:J{L2}", DECISOES, "Liberado; Liberado com autorização (verificação pendente); Retido")
cf_texto(ws, f"I{L1}:I{L2}", f"I{L1}", LEIT_CF, resto=YELLOW)
cf_texto(ws, f"J{L1}:J{L2}", f"J{L1}", DEC_CF)
cf_warn(ws, f"O{L1}:O{L2}", f"O{L1}")
note(ws, "F4", "As verificações planejadas para o lote, no plano de controle e na inspeção final.")
note(ws, "L4", "Obrigatório quando a decisão é “Liberado com autorização”: quem aprovou a liberação antes de concluir as verificações, e o aceite do cliente, se for o caso.")
s = L2 + 2
band(ws, s, "Resumo automático", "O")
LBR = [("Lotes registrados", f"=COUNTA(D{L1}:D{L2})"),
       ("Por decisão", f'=IF(E{s+1}=0,"",COUNTIF(J{L1}:J{L2},"{LIBERADO}")&" liberados, "&COUNTIF(J{L1}:J{L2},"{AUTORIZADO}")&" com autorização, "&COUNTIF(J{L1}:J{L2},"{RETIDO}")&" retidos")'),
       ("Liberações a rever", f'=SUMPRODUCT((O{L1}:O{L2}<>"")*(O{L1}:O{L2}<>"OK"))'),
       ("Aviso", f'=IF(E{s+1}=0,"Registre as liberações",IF(COUNTIF(O{L1}:O{L2},"Liberado com verificação pendente")>0,"Há lote liberado com verificação pendente, sem autorização",'
                 f'IF(E{s+3}>0,"Há liberação a rever: veja a coluna Conferência","OK")))')]
for k, (text, formula) in enumerate(LBR, 1):
    label(ws, f"B{s+k}", text, merge=f"B{s+k}:D{s+k}", h="right")
    calc(ws, f"E{s+k}", formula, merge=f"E{s+k}:I{s+k}", sz=9 if text == "Aviso" else 10, b=text != "Aviso", h="left")
    ws.row_dimensions[s + k].height = 21.75
LAV = s + 4
cf_warn(ws, f"E{LAV}:I{LAV}", f"E{LAV}")
ws.freeze_panes = f"E{L1}"
setup(ws, BLUE, f"B1:O{LAV}")

# ------------------------------------------------------------------ Produto não conforme
ws = wb.create_sheet("Produto não conforme")
widths(ws, {"A": 2, "B": 5, "C": 10, "D": 12, "E": 18, "F": 20, "G": 32, "H": 12, "I": 15, "J": 28, "K": 20, "L": 20, "M": 30, "N": 14, "O": 12, "P": 11, "Q": 20,
            "R": 14, "S": 11, "T": 30, "U": 2})
title(ws, "Registro de produto não conforme", "Uma linha por caso. O registro descreve o defeito, a segregação, a disposição, quem decidiu e o que foi feito com o cliente.", "T")
for col, text in zip("BCDEFGHIJKLMNOPQRST", ["#", "Nº do registro", "Data", "Lote ou pedido", "Defeito", "Descrição", "Quantidade", "Onde foi visto", "Identificação e segregação",
                                            "Disposição", "Quem decidiu", "Cliente: concessão ou comunicação", "Reverificação", "Encerrado em", "Custo", "Ação corretiva", "Situação",
                                            "Repetições", "Conferência"]):
    head(ws, f"{col}4", text)
ws.row_dimensions[4].height = 33
hint_row(ws, 5, [("B", ""), ("C", "Código"), ("D", "Data"), ("E", "A identificação"), ("F", "Nome curto, sempre igual"), ("G", "O que foi visto, com o critério"), ("H", "Com unidade"),
                 ("I", "Lista"), ("J", "A marca e o lugar"), ("K", "Lista"), ("L", "Uma função"), ("M", "Aceite ou aviso, com data"), ("N", "Lista"), ("O", "Data"), ("P", "R$"),
                 ("Q", "Quando houver"), ("R", "Calculada"), ("S", "Calculadas"), ("T", "Calculada")])
ex(ws, "B6", "Ex.", h="center")
for col, v, h_, fmt in [("C", "2027-21", "center", None), ("D", date(2027, 5, 20), "center", DATE), ("E", "Lote 140", "left", None), ("F", "Espessura fora", "left", None),
                        ("G", "Espessura de 43,1 e 43,4 µm, com o critério de 38 a 42 µm.", "left", None), ("H", "600 kg", "center", None), ("I", "Inspeção final", "center", None),
                        ("J", "Lote retido, etiqueta vermelha no palete.", "left", None), ("K", CONCES, "center", None), ("L", "Coordenador da Qualidade", "left", None),
                        ("M", "Cliente B aceitou por e-mail de 21/05.", "left", None), ("N", None, "center", None), ("O", date(2027, 5, 21), "center", DATE), ("P", 1350, "center", REAL),
                        ("Q", None, "left", None), ("R", ENCER, "center", None), ("S", 1, "center", None), ("T", "OK", "center", None)]:
    ex(ws, f"{col}6", v, h=h_, fmt=fmt)
ws.row_dimensions[6].height = 30
DEFS, ACOES = f"$F${P1}:$F${P2}", f"$Q${P1}:$Q${P2}"
for k in range(NP):
    rr = P1 + k
    num(ws, f"B{rr}", k + 1)
    inp(ws, f"C{rr}", h="center")
    inp(ws, f"D{rr}", h="center", fmt=DATE)
    for col in "EFGJLMQ":
        inp(ws, f"{col}{rr}")
    inp(ws, f"H{rr}", h="center")
    inp(ws, f"I{rr}", h="center")
    inp(ws, f"K{rr}", h="center")
    inp(ws, f"N{rr}", h="center")
    inp(ws, f"O{rr}", h="center", fmt=DATE)
    inp(ws, f"P{rr}", h="center", fmt=REAL)
    calc(ws, f"R{rr}", "=" + f_sit(f"F{rr}", f"K{rr}", f"O{rr}"), sz=9)
    calc(ws, f"S{rr}", f'=IF(F{rr}="","",COUNTIF({DEFS},F{rr}))', b=False)
    cells = dict(defe=f"F{rr}", data=f"D{rr}", det=f"I{rr}", disp=f"K{rr}", quem=f"L{rr}", seg=f"J{rr}", cli=f"M{rr}", rev=f"N{rr}", enc=f"O{rr}")
    calc(ws, f"T{rr}", "=" + f_pnc_conf(cells, f"S{rr}", DEFS, ACOES, f"COUNTA(C{rr}:E{rr},G{rr}:Q{rr})"), b=False, sz=9)
    ws.row_dimensions[rr].height = 30
dv_date(ws, f"D{P1}:D{P2}")
dv_date(ws, f"O{P1}:O{P2}")
dv_number(ws, f"P{P1}:P{P2}")
dv_list(ws, f"I{P1}:I{P2}", DET, "Onde o defeito foi visto")
dv_list(ws, f"K{P1}:K{P2}", DISP, "O destino decidido para o produto")
dv_list(ws, f"N{P1}:N{P2}", REVERIF, "Resultado da nova verificação, depois da correção")
cf_texto(ws, f"R{P1}:R{P2}", f"R{P1}", SIT_CF)
cf_texto(ws, f"T{P1}:T{P2}", f"T{P1}", [("OK", GREEN), ("Falta informar o cliente", RED), (REP, RED)], resto=YELLOW)
note(ws, "F4", "Use sempre o mesmo nome para o mesmo defeito: é por ele que a planilha conta as repetições.")
note(ws, "M4", "Obrigatório na concessão (o aceite do cliente, antes da entrega) e no defeito visto depois da entrega (quando e como o cliente foi informado).")
note(ws, "Q4", f"O número da ação corretiva, quando o defeito levou a uma análise de causa. A partir de {REPETE} registros do mesmo defeito, a planilha pede uma.")
s = P2 + 2
band(ws, s, "Resumo automático", "T")
PBR = [("Registros", f"=COUNTA(F{P1}:F{P2})"),
       ("Por situação", f'=IF(E{s+1}=0,"",COUNTIF(R{P1}:R{P2},"{ABERTO}")&" abertos, "&COUNTIF(R{P1}:R{P2},"{TRAT}")&" em tratamento, "&COUNTIF(R{P1}:R{P2},"{ENCER}")&" encerrados")'),
       ("Custo", f"=SUM(P{P1}:P{P2})"),
       ("Aviso", f'=IF(E{s+1}=0,"Registre os casos de produto não conforme",IF(COUNTIF(T{P1}:T{P2},"Falta informar o cliente")>0,"Há defeito visto depois da entrega sem o cliente informado",'
                 f'IF(COUNTIF(R{P1}:R{P2},"{ABERTO}")>0,"Há registro sem disposição",IF(COUNTIF(T{P1}:T{P2},"{REP}")>0,"Há defeito repetido sem ação corretiva",'
                 f'IF(SUMPRODUCT((T{P1}:T{P2}<>"")*(T{P1}:T{P2}<>"OK"))>0,"Há registro a completar: veja a coluna Conferência","OK")))))')]
for k, (text, formula) in enumerate(PBR, 1):
    label(ws, f"B{s+k}", text, merge=f"B{s+k}:D{s+k}", h="right")
    calc(ws, f"E{s+k}", formula, merge=f"E{s+k}:J{s+k}", sz=9 if text == "Aviso" else 10, b=text != "Aviso", h="left", fmt=REAL if text == "Custo" else None)
    ws.row_dimensions[s + k].height = 21.75
NAVP = s + 4
cf_warn(ws, f"E{NAVP}:J{NAVP}", f"E{NAVP}")
ws.freeze_panes = f"G{P1}"
setup(ws, REDC, f"B1:T{NAVP}")

# ------------------------------------------------------------------ Painel
ws = wb.create_sheet("Painel")
widths(ws, {"A": 2, "B": 30, "C": 14, "D": 14, "E": 16, "F": 14, "G": 30, "H": 2})
title(ws, "Painel da liberação e do produto não conforme", "Nada a preencher: tudo vem das outras abas.", "G")
for col, text, m in [("B", "Indicador", None), ("C", "Resultado", None), ("D", "Como se lê", "D4:G4")]:
    put(ws, f"{col}4", text, f=font(10, True, c=WHITE), bg=INK, h="center", merge=m)
ws.row_dimensions[4].height = 21.75
Ld, Lj, Li, Lo = (f"{LB}!{c}{L1}:{c}{L2}" for c in "DJIO")
Pf, Pi, Pk, Pp, Pr, Pt = (f"{PN}!{c}{P1}:{c}{P2}" for c in "FIKPRT")
IND = [
    ("Lotes registrados", f"=COUNTA({Ld})", None, "Da aba Liberação."),
    ("Liberados conformes", f'=COUNTIFS({Lj},"{LIBERADO}",{Li},"{L_CONF}")', None, "Todas as verificações feitas, nenhuma fora do critério."),
    ("Liberados com autorização", f'=COUNTIF({Lj},"{AUTORIZADO}")', None, "Verificação pendente, com a aprovação de quem pode."),
    ("Retidos", f'=COUNTIF({Lj},"{RETIDO}")', None, "Esperam a disposição."),
    ("Liberações a rever", f'=SUMPRODUCT(({Lo}<>"")*({Lo}<>"OK"))', None, "Liberado com pendência, sem quem liberou, desvio sem registro."),
    ("Registros de produto não conforme", f"=COUNTA({Pf})", None, "Da aba Produto não conforme."),
    ("Abertos, sem disposição", f'=COUNTIF({Pr},"{ABERTO}")', None, "Decidir na semana."),
    ("Em tratamento", f'=COUNTIF({Pr},"{TRAT}")', None, "Com disposição, sem encerramento."),
    ("Achados depois da entrega", f'=COUNTIF({Pi},"{APOS}")', None, "Os mais caros: o cliente viu o defeito."),
    ("Custo do período", f"=SUM({Pp})", REAL, "A soma do custo dos registros."),
    ("Registros a completar", f'=SUMPRODUCT(({Pt}<>"")*({Pt}<>"OK"))', None, "Veja a coluna Conferência da aba Produto não conforme."),
    ("Defeitos repetidos sem ação", f'=COUNTIF({Pt},"{REP}")', None, f"Registros de um defeito com {REPETE} casos ou mais, sem ação corretiva."),
]
for k, (nome, formula, fmt, leit) in enumerate(IND):
    rr = 5 + k
    label(ws, f"B{rr}", nome)
    calc(ws, f"C{rr}", formula, fmt=fmt)
    put(ws, f"D{rr}", leit, f=font(9, c=MUTED), bg=GRAY, merge=f"D{rr}:G{rr}")
    ws.row_dimensions[rr].height = 21.75
IR = {n: 5 + k for k, (n, _, _, _) in enumerate(IND)}
T0 = 5 + len(IND) + 1
band(ws, T0, "Onde o defeito foi visto, e quanto custou", "G")
for col, text in zip("BCDEFG", ["Ponto de detecção", "Registros", "Custo", "Custo por registro", "Parte do custo", "Leitura"]):
    head(ws, f"{col}{T0+1}", text)
ws.row_dimensions[T0 + 1].height = 30
D1 = T0 + 2
LEIT_DET = ["O defeito não entrou na operação.", "Visto onde nasce: a correção é barata.", "O produto já está pronto: o custo cresce.", "O cliente viu: o mais caro."]
for k, d in enumerate(DET):
    rr = D1 + k
    label(ws, f"B{rr}", d)
    calc(ws, f"C{rr}", f'=COUNTIF({Pi},B{rr})', b=False)
    calc(ws, f"D{rr}", f'=SUMIF({Pi},B{rr},{Pp})', fmt=REAL, b=False)
    calc(ws, f"E{rr}", f'=IF(C{rr}=0,0,D{rr}/C{rr})', fmt=REAL)
    calc(ws, f"F{rr}", f'=IF(SUM(D${D1}:D${D1 + 3})=0,"",D{rr}/SUM(D${D1}:D${D1 + 3}))', fmt="0%", b=False)
    put(ws, f"G{rr}", LEIT_DET[k], f=font(9, c=MUTED), bg=GRAY)
    ws.row_dimensions[rr].height = 21.75
T1 = D1 + 5
band(ws, T1, "Por disposição", "G")
for col, text in zip("BCDEFG", ["Disposição", "Registros", "Custo", "Custo por registro", "Parte do custo", "Cuidado"]):
    head(ws, f"{col}{T1+1}", text)
ws.row_dimensions[T1 + 1].height = 30
Q1 = T1 + 2
for k, (nome, _, _, cuidado, _) in enumerate(DISP_INFO):
    rr = Q1 + k
    label(ws, f"B{rr}", nome)
    calc(ws, f"C{rr}", f'=COUNTIF({Pk},B{rr})', b=False)
    calc(ws, f"D{rr}", f'=SUMIF({Pk},B{rr},{Pp})', fmt=REAL, b=False)
    calc(ws, f"E{rr}", f'=IF(C{rr}=0,"",D{rr}/C{rr})', fmt=REAL, b=False)
    calc(ws, f"F{rr}", f'=IF(SUM(D${Q1}:D${Q1 + 5})=0,"",D{rr}/SUM(D${Q1}:D${Q1 + 5}))', fmt="0%", b=False)
    put(ws, f"G{rr}", cuidado, f=font(9, c=MUTED), bg=GRAY)
    ws.row_dimensions[rr].height = 24
s = Q1 + 7
band(ws, s, "Resumo automático", "G")
label(ws, f"B{s+1}", "Aviso")
calc(ws, f"C{s+1}", f'=IF(AND(C{IR["Lotes registrados"]}=0,C{IR["Registros de produto não conforme"]}=0),"Registre as liberações e os casos de produto não conforme",'
     f'IF(COUNTIF({Pt},"Falta informar o cliente")>0,"Há defeito visto depois da entrega sem o cliente informado",'
     f'IF(COUNTIF({Lo},"Liberado com verificação pendente")>0,"Há lote liberado com verificação pendente, sem autorização",'
     f'IF(C{IR["Abertos, sem disposição"]}>0,"Há registro sem disposição",IF(C{IR["Defeitos repetidos sem ação"]}>0,"Há defeito repetido sem ação corretiva",'
     f'IF(C{IR["Liberações a rever"]}+C{IR["Registros a completar"]}>0,"Há liberação ou registro a completar","OK"))))))', merge=f"C{s+1}:G{s+1}", sz=9, b=False, h="left")
cf_warn(ws, f"C{s+1}:G{s+1}", f"C{s+1}")
ws.row_dimensions[s + 1].height = 21.75
NAV = s + 1
colunas(ws, f"B{s + 3}", D1, D1 + 3, 2, 5, width=20, height=8)
setup(ws, REDC, f"B1:G{s + 20}", landscape=False, fit_height=True)

# ------------------------------------------------------------------ Checklist
ws = wb.create_sheet("Checklist")
widths(ws, {"A": 2, "B": 5, "C": 70, "D": 14, "E": 44, "F": 2})
title(ws, "Checklist da liberação e do produto não conforme", "Escolha o status de cada item na lista suspensa (coluna em amarelo-claro).", "E")
for col, text in zip("BCDE", ["#", "Item de verificação", "Status", "Observações"]):
    head(ws, f"{col}4", text)
ws.row_dimensions[4].height = 21.75
for k, text in enumerate(CHECK):
    rr = 5 + k
    num(ws, f"B{rr}", k + 1)
    put(ws, f"C{rr}", text, bg=GRAY)
    inp(ws, f"D{rr}", h="center")
    inp(ws, f"E{rr}")
    ws.row_dimensions[rr].height = 30
dv_list(ws, "D5:D16", YESNO, "Sim, Parcial ou Não")
cf_equal(ws, "D5:D16", YESNO_CF)
band(ws, 18, "Resumo automático", "E")
summary(ws, 19, "Itens atendidos (Sim)", '=COUNTIF(D5:D16,"Sim")', "C", "D")
summary(ws, 20, "Itens parciais", '=COUNTIF(D5:D16,"Parcial")', "C", "D")
summary(ws, 21, "Itens não atendidos", '=COUNTIF(D5:D16,"Não")', "C", "D")
summary(ws, 22, "Itens sem resposta", "=COUNTBLANK(D5:D16)", "C", "D")
summary(ws, 23, "Percentual atendido", "=D19/ROWS(D5:D16)", "C", "D", fmt="0%")
summary(ws, 24, "Situação", '=IF(D19=ROWS(D5:D16),"Validado",IF(D22>0,"Em andamento","Com pendências"))', "C", "D")
setup(ws, PURPLE, "B1:E24", landscape=False, fit_height=True)


# ------------------------------------------------------------------ Exemplos
def sub(ws, rr, cells, height=30):
    for a, b, t in cells:
        put(ws, f"{a}{rr}", t, f=font(10, True), bg=GRAY, h="center", merge=b and f"{a}{rr}:{b}{rr}")
    ws.row_dimensions[rr].height = height


def exemplo(ws, ex_):
    H = ex_["head"]
    widths(ws, {"A": 2, "B": 10, "C": 11, "D": 17, "E": 16, "F": 10, "G": 12, "H": 18, "I": 18, "J": 16, "K": 22, "L": 13, "M": 12, "N": 11, "O": 16, "P": 14, "Q": 28, "R": 2})
    title(ws, "Liberação e produto não conforme", "Exemplo preenchido, para consulta. Use as abas Autoridades, Liberação e Produto não conforme para a sua organização.", "Q")
    rr = 4
    for rot, val in [("Organização", H["org"]), ("Processo", H["processo"]), ("Produto", H["produto"]), ("Quem libera", H["libera"]), ("Quem decide a disposição", H["decide"]),
                     ("Registros", f'{H["periodo"]}, lidos em {H["ref"]:%d/%m/%Y}.'), ("Origem", H["origem"])]:
        label(ws, f"B{rr}", rot, merge=f"B{rr}:D{rr}")
        inp(ws, f"E{rr}", val, merge=f"E{rr}:Q{rr}", bg=WHITE, h="left")
        ws.row_dimensions[rr].height = alt([(val, 170)], minimo=19.5)
        rr += 1
    rr += 1
    band(ws, rr, "Liberações", "Q", color=BLUE)
    rr += 1
    sub(ws, rr, [("B", None, "Data"), ("C", None, "Lote"), ("D", None, "Produto e quantidade"), ("E", None, "Verificações previstas"), ("F", None, "Feitas"), ("G", None, "Fora"),
                 ("H", None, "Leitura"), ("I", None, "Decisão"), ("J", None, "Quem liberou"), ("K", "L", "Quem autorizou"), ("M", None, "Registro"), ("N", "P", "Observação"), ("Q", None, "Conferência")])
    l1 = rr + 1
    for l in ex_["libs"]:
        rr += 1
        put(ws, f"B{rr}", l["data"], h="center", fmt=DATE)
        put(ws, f"C{rr}", l["lote"], f=font(10, True))
        put(ws, f"D{rr}", l["prod"])
        for col, v in zip("EFG", (l["prev"], l["feitas"], l["fora"])):
            put(ws, f"{col}{rr}", v, h="center")
        calc(ws, f"H{rr}", "=" + f_leitura(f"C{rr}", f"E{rr}", f"F{rr}", f"G{rr}"), sz=9)
        put(ws, f"I{rr}", l["dec"], h="center")
        put(ws, f"J{rr}", l["quem"] or None)
        put(ws, f"K{rr}", l["aut"] or None, merge=f"K{rr}:L{rr}")
        put(ws, f"M{rr}", l["pnc"] or None, h="center")
        put(ws, f"N{rr}", l["obs"] or None, merge=f"N{rr}:P{rr}")
        calc(ws, f"Q{rr}", "=" + f_lib_conf(f"C{rr}", f"H{rr}", f"I{rr}", f"J{rr}", f"K{rr}", f"M{rr}", f"E{rr}", f"F{rr}", f"G{rr}"), b=False, sz=9)
        ws.row_dimensions[rr].height = alt([(l["obs"], 40), (l["aut"], 30), (l["prod"], 16), (l["pnc"], 11)], minimo=21.75)
    l2 = rr
    cf_texto(ws, f"H{l1}:H{l2}", f"H{l1}", LEIT_CF)
    cf_texto(ws, f"I{l1}:I{l2}", f"I{l1}", DEC_CF)
    cf_warn(ws, f"Q{l1}:Q{l2}", f"Q{l1}")
    rr += 2
    ws.row_breaks.append(Break(id=rr - 1))
    band(ws, rr, "Produto não conforme", "Q", color=REDC)
    rr += 1
    sub(ws, rr, [("B", None, "Nº"), ("C", None, "Data"), ("D", None, "Lote ou pedido"), ("E", None, "Defeito"), ("F", None, "Quantidade"), ("G", None, "Onde foi visto"),
                 ("H", None, "Identificação e segregação"), ("I", None, "Disposição"), ("J", None, "Quem decidiu"), ("K", None, "Cliente"), ("L", None, "Reverificação"),
                 ("M", None, "Encerrado em"), ("N", None, "Custo"), ("O", None, "Ação corretiva"), ("P", None, "Situação"), ("Q", None, "Conferência")])
    p1 = rr + 1
    p2 = p1 + len(ex_["pncs"]) - 1
    defs, acoes = f"$E${p1}:$E${p2}", f"$O${p1}:$O${p2}"
    for p in ex_["pncs"]:
        rr += 1
        put(ws, f"B{rr}", p["num"], f=font(10, True, c=MUTED), h="center")
        put(ws, f"C{rr}", p["data"], h="center", fmt=DATE)
        put(ws, f"D{rr}", p["lote"])
        put(ws, f"E{rr}", p["defeito"], f=font(10, True))
        put(ws, f"F{rr}", p["qtd"], h="center")
        put(ws, f"G{rr}", p["det"], h="center")
        put(ws, f"H{rr}", p["seg"] or None)
        put(ws, f"I{rr}", p["disp"] or None, h="center")
        put(ws, f"J{rr}", p["quem"] or None)
        put(ws, f"K{rr}", p["cliente"] or None)
        put(ws, f"L{rr}", p["rev"] or None, h="center")
        put(ws, f"M{rr}", p["enc"], h="center", fmt=DATE)
        put(ws, f"N{rr}", p["custo"], h="center", fmt=REAL)
        put(ws, f"O{rr}", p["acao"] or None)
        calc(ws, f"P{rr}", "=" + f_sit(f"E{rr}", f"I{rr}", f"M{rr}"), sz=9)
        cells = dict(defe=f"E{rr}", data=f"C{rr}", det=f"G{rr}", disp=f"I{rr}", quem=f"J{rr}", seg=f"H{rr}", cli=f"K{rr}", rev=f"L{rr}", enc=f"M{rr}")
        calc(ws, f"Q{rr}", "=" + f_pnc_conf(cells, f"COUNTIF({defs},E{rr})", defs, acoes, "1"), b=False, sz=9)
        ws.row_dimensions[rr].height = alt([(p["seg"], 18), (p["cliente"], 22), (p["lote"], 16), (p["defeito"], 16), (p["quem"], 15), (p["acao"], 15), (p["disp"], 17)], minimo=21.75)
    assert rr == p2
    cf_texto(ws, f"P{p1}:P{p2}", f"P{p1}", SIT_CF)
    cf_texto(ws, f"Q{p1}:Q{p2}", f"Q{p1}", [("OK", GREEN), ("Falta informar o cliente", RED), (REP, RED)], resto=YELLOW)
    rr += 2
    band(ws, rr, "Onde o defeito foi visto, e quanto custou", "Q", color=AMBER)
    rr += 1
    sub(ws, rr, [("B", "D", "Ponto de detecção"), ("E", None, "Registros"), ("F", "G", "Custo"), ("H", None, "Custo por registro"), ("I", None, "Parte do custo")])
    d1 = rr + 1
    for d in DET:
        rr += 1
        label(ws, f"B{rr}", d, merge=f"B{rr}:D{rr}")
        calc(ws, f"E{rr}", f'=COUNTIF(G{p1}:G{p2},B{rr})', b=False)
        calc(ws, f"F{rr}", f'=SUMIF(G{p1}:G{p2},B{rr},N{p1}:N{p2})', fmt=REAL, b=False, merge=f"F{rr}:G{rr}")
        calc(ws, f"H{rr}", f'=IF(E{rr}=0,"",F{rr}/E{rr})', fmt=REAL)
        calc(ws, f"I{rr}", f'=IF(SUM(F${d1}:F${d1 + 3})=0,"",F{rr}/SUM(F${d1}:F${d1 + 3}))', fmt="0%", b=False)
        ws.row_dimensions[rr].height = 21.75
    rr += 2
    band(ws, rr, "Resumo automático", "Q")
    for k, (text, formula, fmt) in enumerate([
        ("Liberações: registradas, a rever", f'=COUNTA(C{l1}:C{l2})&" registradas, "&SUMPRODUCT((Q{l1}:Q{l2}<>"OK")*1)&" a rever"', None),
        ("Registros: abertos, em tratamento, encerrados", f'=COUNTIF(P{p1}:P{p2},"{ABERTO}")&" abertos, "&COUNTIF(P{p1}:P{p2},"{TRAT}")&" em tratamento, "&COUNTIF(P{p1}:P{p2},"{ENCER}")&" encerrados"', None),
        ("Registros a completar", f'=SUMPRODUCT((Q{p1}:Q{p2}<>"OK")*1)', None),
        ("Custo do período", f"=SUM(N{p1}:N{p2})", REAL),
        ("Custo achado depois da entrega", f'=SUMIF(G{p1}:G{p2},"{APOS}",N{p1}:N{p2})', REAL),
    ], 1):
        label(ws, f"B{rr+k}", text, merge=f"B{rr+k}:F{rr+k}", h="right")
        calc(ws, f"G{rr+k}", formula, fmt=fmt, merge=f"G{rr+k}:Q{rr+k}", h="left")
        ws.row_dimensions[rr + k].height = 21.75
    setup(ws, MUTED, f"B1:Q{rr + 5}")


exemplo(wb.create_sheet("Exemplo 1 - Pizzaria"), EX1)
exemplo(wb.create_sheet("Exemplo 2 - Indústria"), EX2)

wb.active = 1
wb.save(OUT)
print("ok", OUT, "| avisos: Autoridades E%d, Liberação E%d, Produto não conforme E%d, Painel C%d" % (UAV, LAV, NAVP, NAV))
