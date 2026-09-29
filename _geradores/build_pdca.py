# -*- coding: utf-8 -*-
"""Gera PDCA-modelo.xlsx no mesmo padrão visual de SIPOC-modelo.xlsx."""
import sys
from datetime import date

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.chart.layout import Layout, ManualLayout
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.comments import Comment
from openpyxl.drawing.line import LineProperties
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.pagebreak import Break

OUT = sys.argv[1]

INK, MUTED, GRAY, INPUT, LINE, WHITE = "16242E", "5B6B76", "EEF1F3", "FFF9DB", "C9D2D8", "FFFFFF"
GREEN, YELLOW, RED = "D8EFDD", "FBEBC8", "F6D9D7"
PH = {
    "P": ("2B5C8A", "DCE8F3", "Planejar"),
    "D": ("A96A12", "F6E8CF", "Executar"),
    "C": ("1E7B73", "D9EEEB", "Verificar"),
    "A": ("B0413E", "F5DEDC", "Agir"),
}
PURPLE = "7A4A9A"
DATE = "dd/mm/yyyy"

side = Side(style="thin", color=LINE)
BOX = Border(left=side, right=side, top=side, bottom=side)


def fill(c):
    return PatternFill("solid", fgColor=c) if c else PatternFill(fill_type=None)


def font(sz=10, b=False, i=False, c=INK):
    return Font(name="Arial", size=sz, bold=b, italic=i, color=c)


def put(ws, ref, value=None, *, f=None, bg=None, h="left", box=True, fmt=None, merge=None):
    """Escreve e formata uma célula; com merge, formata todo o intervalo."""
    cell = ws[ref]
    if value is not None:
        cell.value = value
    rng = ws[merge] if merge else ((cell,),)
    for row in rng:
        for c in row:
            c.font = f or font()
            c.fill = fill(bg)
            c.alignment = Alignment(horizontal=h, vertical="center", wrap_text=True, indent=1 if h == "left" else 0)
            if box:
                c.border = BOX
            if fmt:
                c.number_format = fmt
    if merge:
        ws.merge_cells(merge)
    return cell


def title(ws, text, sub, last_col):
    put(ws, "B1", text, f=font(16, True, c=WHITE), bg=INK, box=False, merge=f"B1:{last_col}1")
    put(ws, "B2", sub, f=font(10, i=True, c=MUTED), box=False, merge=f"B2:{last_col}2")
    ws.row_dimensions[1].height = 31.5
    ws.row_dimensions[2].height = 19.5


def band(ws, row, text, last_col, color=INK, sz=10):
    put(ws, f"B{row}", text, f=font(sz, True, c=WHITE), bg=color, box=False, merge=f"B{row}:{last_col}{row}")
    ws.row_dimensions[row].height = 21.75


def label(ws, ref, text, merge=None, h="left"):
    put(ws, ref, text, f=font(10, True), bg=GRAY, h=h, merge=merge)


def head(ws, ref, text, bg=INK, fg=WHITE):
    put(ws, ref, text, f=font(10, True, c=fg), bg=bg, h="center")


def num(ws, ref, n):
    put(ws, ref, n, f=font(10, True, c=MUTED), bg=GRAY, h="center")


def ex(ws, ref, text, h="left", fmt=None):
    put(ws, ref, text, f=font(9, i=True, c=MUTED), bg=GRAY, h=h, fmt=fmt)


def calc(ws, ref, formula, h="center", b=True, fmt=None, merge=None, sz=10):
    put(ws, ref, formula, f=font(sz, b), bg=GRAY, h=h, fmt=fmt, merge=merge)


def inp(ws, ref, value=None, h="left", fmt=None, merge=None, bg=INPUT):
    put(ws, ref, value, bg=bg, h=h, fmt=fmt, merge=merge)


def widths(ws, spec):
    for col, w in spec.items():
        ws.column_dimensions[col].width = w


def setup(ws, tab, area, landscape=True, fit_height=False):
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = tab
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.paperSize = 9
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1 if fit_height else 0
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.5
    ws.print_options.horizontalCentered = True
    ws.print_area = area


def dv_list(ws, rng, options, prompt):
    dv = DataValidation(type="list", formula1='"%s"' % ",".join(options), allow_blank=True)
    dv.promptTitle, dv.prompt = "Opções", prompt
    dv.errorTitle, dv.error = "Valor inválido", "Escolha uma das opções da lista."
    dv.showErrorMessage = dv.showInputMessage = True
    ws.add_data_validation(dv)
    dv.add(rng)


def dv_date(ws, rng):
    dv = DataValidation(type="date", operator="greaterThanOrEqual", formula1="36526", allow_blank=True)
    dv.errorTitle, dv.error = "Data inválida", "Digite uma data no formato dia/mês/ano."
    dv.showErrorMessage = True
    ws.add_data_validation(dv)
    dv.add(rng)


def dv_number(ws, rng):
    dv = DataValidation(type="decimal", operator="between", formula1="-1000000000", formula2="1000000000", allow_blank=True)
    dv.errorTitle, dv.error = "Valor inválido", "Digite apenas o número, sem texto nem unidade."
    dv.showErrorMessage = True
    ws.add_data_validation(dv)
    dv.add(rng)


def cf_equal(ws, rng, pairs):
    for text, color in pairs:
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"%s"' % text], fill=PatternFill("solid", bgColor=color, fgColor=color)))


def cf_ok(ws, rng):
    ok = PatternFill("solid", bgColor=GREEN, fgColor=GREEN)
    warn = PatternFill("solid", bgColor=YELLOW, fgColor=YELLOW)
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"OK"'], fill=ok))
    ws.conditional_formatting.add(rng, CellIsRule(operator="notEqual", formula=['"OK"'], fill=warn))


def note(ws, ref, text):
    c = Comment(text, "Modelo PDCA")
    c.width, c.height = 260, 90
    ws[ref].comment = c


def summary(ws, row, text, formula, lab_to, val, fmt=None, val_merge=None):
    label(ws, f"B{row}", text, merge=f"B{row}:{lab_to}{row}", h="right")
    calc(ws, f"{val}{row}", formula, fmt=fmt, merge=val_merge)
    ws.row_dimensions[row].height = 19.5


STATUS = ["Não iniciada", "Em andamento", "Concluída"]
STATUS_CF = [("Concluída", GREEN), ("Em andamento", YELLOW), ("Não iniciada", RED)]
YESNO = ["Sim", "Parcial", "Não"]
YESNO_CF = [("Sim", GREEN), ("Parcial", YELLOW), ("Não", RED)]
OPEN = ["Aberta", "Em andamento", "Concluída"]
OPEN_CF = [("Concluída", GREEN), ("Em andamento", YELLOW), ("Aberta", RED)]

STEPS = [
    ("P", "Identificar o problema", "Qual é o problema, quanto ele pesa e qual é a meta?"),
    ("P", "Observar", "Onde, quando e com quem o problema acontece?"),
    ("P", "Analisar as causas", "Por que acontece? Qual é a causa raiz, confirmada com dados?"),
    ("P", "Planejar a ação", "O que será feito, por quem e até quando?"),
    ("D", "Executar", "O plano foi executado como previsto? O que foi diferente?"),
    ("C", "Verificar", "A meta foi atingida? Houve efeito colateral?"),
    ("A", "Padronizar", "O que vira padrão e quem precisa ser treinado?"),
    ("A", "Concluir", "O que aprendemos e qual é o próximo ciclo?"),
]

PDCA_W = {"A": 2, "B": 5, "C": 16, "D": 24, "E": 30, "F": 44, "G": 24, "H": 16, "I": 12, "J": 14, "K": 2}


# ------------------------------------------------------------------ medições
def measurements(ws, r0, meta, sentido, data=None, example_row=True, axis=None):
    """Tabela de medições a partir da linha r0 (cabeçalho). Devolve (primeira, última, linha livre)."""
    tint_h = PH["C"][1]
    for col, text in zip("BCDEFGH", ["#", "Período", "Momento", "Valor medido", "Observações", "Meta", "Atingiu a meta?"]):
        put(ws, f"{col}{r0}", text, f=font(10, True), bg=tint_h, h="center")
    ws.row_dimensions[r0].height = 30
    r = r0 + 1
    if example_row:
        ex(ws, f"B{r}", "Ex.", h="center")
        ex(ws, f"C{r}", "Semana 1")
        ex(ws, f"D{r}", "Antes", h="center")
        ex(ws, f"E{r}", 81, h="center")
        ex(ws, f"F{r}", "Semana com feriado na sexta-feira")
        ex(ws, f"G{r}", 95, h="center")
        ex(ws, f"H{r}", "Não", h="center")
        ws.row_dimensions[r].height = 24
        r += 1
    first = r
    n = 16
    for k in range(n):
        row = first + k
        d = data[k] if data and k < len(data) else None
        bg = WHITE if data else INPUT
        num(ws, f"B{row}", k + 1)
        inp(ws, f"C{row}", d[0] if d else None, bg=bg)
        inp(ws, f"D{row}", d[1] if d else None, h="center", bg=bg)
        inp(ws, f"E{row}", d[2] if d else None, h="center", bg=bg)
        inp(ws, f"F{row}", d[3] if d else None, bg=bg)
        calc(ws, f"G{row}", f'=IF(OR({meta}="",E{row}=""),NA(),{meta})', b=False)
        calc(ws, f"H{row}", f'=IF(OR(E{row}="",{meta}=""),"",IF({sentido}="Menor é melhor",IF(E{row}<={meta},"Sim","Não"),IF(E{row}>={meta},"Sim","Não")))', b=False)
        ws.row_dimensions[row].height = 24
    last = first + n - 1
    dv_list(ws, f"D{first}:D{last}", ["Antes", "Durante", "Depois"], "Antes, Durante ou Depois das ações")
    dv_number(ws, f"E{first}:E{last}")
    cf_equal(ws, f"H{first}:H{last}", [("Sim", GREEN), ("Não", RED)])
    # esconde o #N/D das linhas sem medição (o gráfico precisa dele para não traçar zero)
    ws.conditional_formatting.add(f"G{first}:G{last}", FormulaRule(formula=[f"ISNA(G{first})"], font=Font(color=GRAY)))

    r = last + 2
    band(ws, r, "Resumo automático", "H")
    s = r + 1
    summary(ws, s, "Medições registradas", f"=COUNT(E{first}:E{last})", "E", "F")
    summary(ws, s + 1, "Média antes das ações", f'=IFERROR(AVERAGEIF(D{first}:D{last},"Antes",E{first}:E{last}),"Sem dados")', "E", "F", fmt="0.0")
    summary(ws, s + 2, "Média depois das ações", f'=IFERROR(AVERAGEIF(D{first}:D{last},"Depois",E{first}:E{last}),"Sem dados")', "E", "F", fmt="0.0")
    summary(ws, s + 3, "Variação (depois menos antes)", f'=IF(AND(ISNUMBER(F{s+1}),ISNUMBER(F{s+2})),F{s+2}-F{s+1},"Sem dados")', "E", "F", fmt="+0.0;-0.0;0.0")
    summary(ws, s + 4, "Medições de depois que atingiram a meta", f'=COUNTIFS(D{first}:D{last},"Depois",H{first}:H{last},"Sim")', "E", "F")
    summary(ws, s + 5, "Resultado do ciclo",
            f'=IF(OR(NOT(ISNUMBER(F{s+2})),{meta}=""),"Sem dados para concluir",'
            f'IF(OR(AND({sentido}="Menor é melhor",F{s+2}<={meta}),AND({sentido}<>"Menor é melhor",F{s+2}>={meta})),'
            f'"Meta atingida: padronize","Meta não atingida: volte à análise das causas"))', "E", "F")
    cf_equal(ws, f"F{s+5}", [("Meta atingida: padronize", GREEN), ("Meta não atingida: volte à análise das causas", YELLOW)])

    # gráfico
    top = s + 7
    band(ws, top, "Gráfico: indicador ao longo do tempo", "H")
    ch = LineChart()
    ch.height, ch.width = 9.5, 27
    ch.style = 2
    ch.title = None
    ch.legend.position = "b"
    ch.display_blanks = "gap"
    ch.add_data(Reference(ws, min_col=5, min_row=r0, max_row=last), titles_from_data=True)
    ch.add_data(Reference(ws, min_col=7, min_row=r0, max_row=last), titles_from_data=True)
    cats_first = first - 1 if example_row else first
    ch.set_categories(Reference(ws, min_col=3, min_row=first, max_row=last))
    if example_row:
        # a linha "Ex." fica fora do gráfico: séries começam na primeira linha de entrada
        for srs, col in zip(ch.series, ("E", "G")):
            srs.val.numRef.f = f"'{ws.title}'!${col}${first}:${col}${last}"
    m, g = ch.series
    m.graphicalProperties.line.solidFill = PH["P"][0]
    m.graphicalProperties.line.width = 25400
    m.marker.symbol = "circle"
    m.marker.size = 7
    m.marker.graphicalProperties = GraphicalProperties(solidFill=PH["P"][0])
    m.marker.graphicalProperties.line.solidFill = WHITE
    m.smooth = False
    g.graphicalProperties.line.solidFill = MUTED
    g.graphicalProperties.line.width = 19050
    g.graphicalProperties.line.dashStyle = "dash"
    g.marker.symbol = "none"
    g.smooth = False
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ch.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill="D3DBE0", w=9525))
    ch.x_axis.graphicalProperties = GraphicalProperties(ln=LineProperties(solidFill="D3DBE0", w=9525))
    ch.y_axis.graphicalProperties = GraphicalProperties(ln=LineProperties(noFill=True))
    ch.x_axis.tickLblPos = "low"
    if axis:
        ch.y_axis.scaling.min, ch.y_axis.scaling.max = axis
    ws.add_chart(ch, f"B{top + 1}")
    for rr in range(top + 1, top + 21):
        ws.row_dimensions[rr].height = 15
    return first, last, top + 21


# ------------------------------------------------------------------ aba PDCA
def pdca_sheet(ws, tab, data=None):
    widths(ws, PDCA_W)
    is_ex = data is not None
    title(ws, "PDCA — Registro do ciclo de melhoria",
          "Exemplo preenchido, para consulta. Use a aba PDCA para o seu ciclo." if is_ex
          else "Preencha as células em amarelo-claro. As células cinza são calculadas automaticamente.", "J")
    hb = WHITE if is_ex else INPUT
    H = data["head"] if is_ex else {}
    rows = [
        (4, "Tema do ciclo", "tema", "Dono do problema", "dono", None, None),
        (5, "Área / unidade", "area", "Data de início", "inicio", None, DATE),
        (6, "Elaborado por", "autor", "Versão", "versao", None, None),
        (8, "Indicador", "indicador", "Unidade de medida", "unidade", None, None),
        (9, "Situação inicial (valor)", "atual", "Meta (valor)", "meta", "General", "General"),
        (10, "Sentido do indicador", "sentido", "Prazo da meta", "prazo", None, DATE),
    ]
    for r, l1, k1, l2, k2, f1, f2 in rows:
        label(ws, f"B{r}", l1, merge=f"B{r}:C{r}")
        inp(ws, f"D{r}", H.get(k1), merge=f"D{r}:F{r}", bg=hb, fmt=f1)
        label(ws, f"G{r}", l2)
        inp(ws, f"H{r}", H.get(k2), merge=f"H{r}:J{r}", bg=hb, fmt=f2)
        ws.row_dimensions[r].height = 21.75
    for r, l1, k1 in [(7, "Descrição do problema", "problema"), (11, "Fora do escopo", "fora")]:
        label(ws, f"B{r}", l1, merge=f"B{r}:C{r}")
        inp(ws, f"D{r}", H.get(k1), merge=f"D{r}:J{r}", bg=hb)
        ws.row_dimensions[r].height = 30
    dv_list(ws, "D10", ["Maior é melhor", "Menor é melhor"], "O indicador melhora quando sobe ou quando desce?")
    dv_date(ws, "H5")
    dv_date(ws, "H10")
    dv_number(ws, "D9")
    dv_number(ws, "H9")
    if not is_ex:
        note(ws, "D4", 'Nome curto do ciclo, com verbo + objeto.\nEx.: "Reduzir os atrasos nas entregas".')
        note(ws, "H4", "Pessoa que responde pelo resultado e conduz o ciclo até o fim.")
        note(ws, "D7", 'Resultado indesejado, com número, local e período. Sem causa e sem solução.\nEx.: "18% das entregas chegam depois de 40 minutos".')
        note(ws, "D8", 'Como o problema é medido.\nEx.: "% de entregas em até 40 minutos".')
        note(ws, "D9", "Valor do indicador antes das ações. Digite só o número, na unidade do indicador.\nEx.: 82 para 82%.")
        note(ws, "H9", "Valor a alcançar. Digite só o número, na mesma unidade da situação inicial.\nEx.: 95 para 95%.")
        note(ws, "D10", "Maior é melhor: % de entregas no prazo.\nMenor é melhor: dias de prazo, número de erros.")
        note(ws, "D11", 'O que este ciclo não vai tratar.\nEx.: "Atendimento no salão; compra de ingredientes".')

    hr = 13
    for col, text in zip("BCDEFGHIJ", ["#", "Fase", "Etapa", "Pergunta-guia", "Registro: o que foi feito ou decidido",
                                        "Ferramenta ou evidência", "Responsável", "Conclusão", "Status"]):
        head(ws, f"{col}{hr}", text)
    ws.row_dimensions[hr].height = 21.75
    r = hr + 1
    if is_ex:
        put(ws, f"B{r}", "Leia de cima para baixo: cada etapa só começa quando a anterior está concluída. "
            "O registro resume o que foi feito; a evidência diz onde está a prova.",
            f=font(9, i=True, c=MUTED), bg=GRAY, merge=f"B{r}:J{r}")
    else:
        ex(ws, f"B{r}", "Ex.", h="center")
        put(ws, f"C{r}", "P", f=font(9, True, i=True, c=MUTED), bg=GRAY, h="center")
        ex(ws, f"D{r}", "Identificar o problema")
        ex(ws, f"E{r}", "Qual é o problema, quanto ele pesa e qual é a meta?")
        ex(ws, f"F{r}", "18% das entregas chegam depois de 40 minutos. Meta: passar de 82% para 95% em 12 semanas.")
        ex(ws, f"G{r}", "Relatório do aplicativo de delivery")
        ex(ws, f"H{r}", "Gerente da loja")
        ex(ws, f"I{r}", date(2026, 10, 30), h="center", fmt=DATE)
        ex(ws, f"J{r}", "Concluída", h="center")
    ws.row_dimensions[r].height = 30
    first = r + 1
    for k, (ph, step, ask) in enumerate(STEPS):
        row = first + k
        dark, tint, _ = PH[ph]
        d = data["steps"][k] if is_ex else (None, None, None, None, None)
        bg = tint if is_ex else INPUT
        num(ws, f"B{row}", k + 1)
        put(ws, f"D{row}", step, f=font(10, True), bg=tint)
        put(ws, f"E{row}", ask, f=font(9, i=True, c=MUTED), bg=tint)
        inp(ws, f"F{row}", d[0], bg=bg)
        inp(ws, f"G{row}", d[1], bg=bg)
        inp(ws, f"H{row}", d[2], bg=bg)
        inp(ws, f"I{row}", d[3], h="center", fmt=DATE, bg=bg)
        inp(ws, f"J{row}", d[4], h="center", bg=bg)
        ws.row_dimensions[row].height = 62 if is_ex else 45
    last = first + 7
    for ph, a, b in [("P", first, first + 3), ("D", first + 4, first + 4), ("C", first + 5, first + 5), ("A", first + 6, last)]:
        dark, tint, name = PH[ph]
        put(ws, f"C{a}", f"{ph}\n{name}", f=font(11, True, c=WHITE), bg=dark, h="center", merge=f"C{a}:C{b}")
    dv_list(ws, f"J{first}:J{last}", STATUS, "Não iniciada, Em andamento ou Concluída")
    dv_date(ws, f"I{first}:I{last}")
    cf_equal(ws, f"J{first}:J{last}", STATUS_CF)

    J = lambda a, b: f"J{first + a}:J{first + b}"
    s = last + 2
    band(ws, s, "Resumo automático", "J")
    summary(ws, s + 1, "Etapas concluídas", f'=COUNTIF(J{first}:J{last},"Concluída")', "E", "F", fmt='0" de 8"')
    summary(ws, s + 2, "Fase atual do ciclo",
            f'=IF(COUNTIF({J(0, 3)},"Concluída")<4,"P — Planejar",IF(J{first+4}<>"Concluída","D — Executar",'
            f'IF(J{first+5}<>"Concluída","C — Verificar",IF(COUNTIF({J(6, 7)},"Concluída")<2,"A — Agir","Ciclo concluído"))))', "E", "F")
    summary(ws, s + 3, "Aviso: sequência das etapas",
            f'=IF(SUMPRODUCT(({J(1, 7)}="Concluída")*({J(0, 6)}<>"Concluída"))>0,"Há etapa concluída antes da anterior","OK")', "E", "F")
    summary(ws, s + 4, "Aviso: registro das etapas",
            f'=IF(SUMPRODUCT((J{first}:J{last}="Concluída")*(F{first}:F{last}=""))>0,"Há etapa concluída sem registro","OK")', "E", "F")
    summary(ws, s + 5, "Aviso: meta",
            '=IF(OR(D9="",H9=""),"Informe a situação inicial e a meta",IF(D10="","Informe o sentido do indicador",'
            'IF(OR(AND(D10="Menor é melhor",H9>=D9),AND(D10="Maior é melhor",H9<=D9)),"A meta não melhora a situação inicial","OK")))', "E", "F")
    summary(ws, s + 6, "Diferença entre a meta e a situação inicial", '=IF(OR(D9="",H9=""),"",H9-D9)', "E", "F", fmt="+0.0;-0.0;0.0")
    summary(ws, s + 7, "Campos do cabeçalho preenchidos", "=COUNTA(D4,H4,D5,H5,D6,H6,D7,D8,H8,D9,H9,D10,H10,D11)", "E", "F", fmt='0" de 14"')
    for k in (3, 4, 5):
        ws[f"F{s+k}"].font = font(9)
    cf_ok(ws, f"F{s+3}:F{s+5}")
    end = s + 7

    if is_ex:
        band(ws, end + 2, "Verificação: medições do indicador", "H", color=PH["C"][0])
        _, _, free = measurements(ws, end + 3, "$H$9", "$D$10", data=data["medidas"], example_row=False, axis=data["eixo"])
        ws.row_breaks.append(Break(id=end + 1))
        setup(ws, tab, f"B1:J{free}", landscape=False)
    else:
        setup(ws, tab, f"B1:J{end}", fit_height=True)


# ------------------------------------------------------------------ dados dos exemplos
EX1 = {
    "eixo": (70, 100),
    "head": dict(tema="Reduzir os atrasos nas entregas de delivery", dono="Gerente da loja", area="Pizzaria (loja)",
                 inicio=date(2026, 6, 15), autor="Equipe da loja", versao="1.0",
                 problema="18% das entregas de delivery chegam ao cliente depois de 40 minutos. Os atrasos geram reclamações e pedidos cancelados.",
                 indicador="% de entregas em até 40 minutos", unidade="%", atual=82, meta=95,
                 sentido="Maior é melhor", prazo=date(2026, 9, 6),
                 fora="Atendimento no salão, compra de ingredientes e cardápio."),
    "steps": [
        ("18% das entregas chegam depois de 40 minutos. Meta: passar de 82% para 95% de entregas no prazo, em até 12 semanas.",
         "Relatório do aplicativo de delivery", "Gerente da loja", date(2026, 6, 26), "Concluída"),
        ("Os atrasos se concentram nas sextas e sábados, das 19h às 22h. Em 46% deles, a pizza pronta ficou esperando entregador.",
         "Estratificação por dia e horário; gráfico de Pareto", "Líder da expedição", date(2026, 7, 10), "Concluída"),
        ("Causas confirmadas: a escala de entregadores é igual em todos os dias; os pedidos saem sem agrupamento por bairro; o formulário aceita endereço sem complemento.",
         "Diagrama de Ishikawa; 5 porquês", "Equipe da loja", date(2026, 7, 17), "Concluída"),
        ("Três ações: reforçar a escala no pico, agrupar pedidos por bairro na expedição e tornar o complemento do endereço obrigatório.",
         "Plano de ação 5W2H", "Gerente da loja", date(2026, 7, 24), "Concluída"),
        ("Ações implantadas nas semanas 7 e 8, com treinamento da expedição. Desvio: o agrupamento por bairro só começou na semana 8.",
         "Plano de ação com datas reais; lista de presença", "Líder da expedição", date(2026, 8, 7), "Concluída"),
        ("Média de 95,5% nas semanas 9 a 12. Meta atingida. Efeito colateral: o custo com entregadores subiu R$ 600 por mês.",
         "Gráfico de tendência, antes e depois", "Gerente da loja", date(2026, 9, 8), "Concluída"),
        ("A nova escala virou padrão, a instrução de trabalho da expedição foi atualizada e o campo obrigatório foi mantido.",
         "Instrução de trabalho revisada; registro de treinamento", "Gerente da loja", date(2026, 9, 15), "Concluída"),
        ("Lição: medir por faixa de horário, e não só pela média do dia. Próximo ciclo: a fila do forno, segundo item do Pareto.",
         "Registro de lições aprendidas", "Equipe da loja", date(2026, 9, 18), "Concluída"),
    ],
    "medidas": [(f"Semana {i+1}", m, v, o) for i, (m, v, o) in enumerate([
        ("Antes", 81, None), ("Antes", 83, None), ("Antes", 80, "Chuva forte no sábado"), ("Antes", 84, None),
        ("Antes", 82, None), ("Antes", 82, None),
        ("Durante", 86, "Escala reforçada e campo obrigatório"), ("Durante", 90, "Início do agrupamento por bairro"),
        ("Depois", 94, None), ("Depois", 96, None), ("Depois", 95, None), ("Depois", 97, None)])],
}

EX2 = {
    "eixo": (0, 14),
    "head": dict(tema="Reduzir o prazo de compra", dono="Gerente de Suprimentos", area="Suprimentos",
                 inicio=date(2026, 6, 1), autor="Equipe de compras", versao="1.0",
                 problema="O pedido de compra leva, em média, 12 dias úteis para ser emitido depois da requisição aprovada. As áreas recebem os itens depois da data de necessidade.",
                 indicador="Dias úteis entre a requisição aprovada e o pedido emitido", unidade="dias úteis", atual=12, meta=7,
                 sentido="Menor é melhor", prazo=date(2026, 8, 31),
                 fora="Homologação de fornecedores, pagamento e gestão de estoque."),
    "steps": [
        ("O pedido leva, em média, 12 dias úteis para ser emitido. Meta: reduzir para 7 dias úteis, em 3 meses.",
         "Relatório do sistema de compras", "Gerente de Suprimentos", date(2026, 6, 5), "Concluída"),
        ("40% das requisições são devolvidas por especificação incompleta. Compras de baixo valor demoram tanto quanto as de alto valor.",
         "Estratificação por valor e por motivo de devolução", "Analista de compras", date(2026, 6, 19), "Concluída"),
        ("Causas confirmadas: o formulário de requisição não tem campos obrigatórios; a política exige três cotações para qualquer valor.",
         "5 porquês; mapa do processo", "Equipe de compras", date(2026, 6, 26), "Concluída"),
        ("Novo formulário com campos obrigatórios e catálogo de itens. Uma cotação só para compras de até R$ 2.000, com aprovação da Diretoria.",
         "Plano de ação 5W2H; política de compras revisada", "Gerente de Suprimentos", date(2026, 7, 3), "Concluída"),
        ("Piloto de 4 semanas em duas áreas requisitantes, com treinamento dos requisitantes e dos compradores.",
         "Registro do piloto", "Analista de compras", date(2026, 7, 31), "Concluída"),
        ("O prazo caiu para 8 dias úteis. Meta não atingida. As devoluções caíram de 40% para 12%. O tempo restante está na revisão de contratos.",
         "Comparação antes e depois; novo Pareto", "Gerente de Suprimentos", date(2026, 8, 28), "Concluída"),
        ("O formulário e a nova alçada foram estendidos a todas as áreas, porque funcionaram. O procedimento de compras foi revisado.",
         "Procedimento revisado; comunicado às áreas", "Gerente de Suprimentos", date(2026, 9, 11), "Concluída"),
        ("A meta de 7 dias foi mantida. Novo ciclo aberto, com foco na etapa de revisão de contratos.",
         "Registro de lições aprendidas; novo ciclo", "Equipe de compras", date(2026, 9, 18), "Concluída"),
    ],
    "medidas": [(f"Semana {i+1}", m, v, o) for i, (m, v, o) in enumerate([
        ("Antes", 12.5, None), ("Antes", 11.5, None), ("Antes", 12, None), ("Antes", 12, None), ("Antes", 12, None),
        ("Durante", 11, "Início do piloto em duas áreas"), ("Durante", 10, None), ("Durante", 9, None), ("Durante", 8.5, None),
        ("Depois", 8, None), ("Depois", 8.5, None), ("Depois", 7.5, None), ("Depois", 8, None)])],
}

# ================================================================== pasta de trabalho
wb = Workbook()
wb.properties.title = "PDCA — Modelo de ciclo de melhoria"
wb.properties.creator = "Modelo PDCA"
wb.properties.language = "pt-BR"

# ------------------------------------------------------------------ Instruções
ws = wb.active
ws.title = "Instruções"
widths(ws, {"A": 2, "B": 24, "C": 92, "D": 2})
title(ws, "PDCA — Como usar esta planilha",
      "Modelo para conduzir um ciclo de melhoria: Planejar, Executar, Verificar e Agir.", "C")
r = 4


def section(text):
    global r
    band(ws, r, text, "C", sz=11)
    r += 1


def line(k, v, kbg=GRAY, vbg=None, kf=None, kh="left", height=19.5):
    global r
    put(ws, f"B{r}", k, f=kf or font(10, True), bg=kbg, h=kh)
    put(ws, f"C{r}", v, bg=vbg)
    ws.row_dimensions[r].height = height
    r += 1


section("Legenda: onde preencher")
line("Amarelo-claro", "Células de entrada. É aqui que você digita.", vbg=INPUT)
line("Cinza", "Células calculadas ou fixas (rótulos, contagens, avisos). Não altere.", vbg=GRAY)
line("Cinza em itálico", "Linha de exemplo (marcada com “Ex.”). Mostra o formato esperado e não entra nas contagens.", vbg=GRAY)
line("Listas suspensas", "Colunas de status, tipo, momento e sentido aceitam apenas as opções da lista.")
line("Comentários", "Células com triângulo vermelho no canto trazem uma dica de preenchimento. Passe o mouse sobre elas.", height=31.5)
r += 1
section("As quatro fases")
for k, text in [("P", "Planejar: identificar o problema, observar, analisar as causas e montar o plano de ação."),
                ("D", "Executar: treinar as pessoas, executar o plano e registrar os desvios."),
                ("C", "Verificar: medir o resultado, comparar com a meta e avaliar efeitos colaterais."),
                ("A", "Agir: padronizar o que funcionou, registrar as lições e definir o próximo ciclo.")]:
    line(k, text, kbg=PH[k][0], vbg=PH[k][1], kf=font(10, True, c=WHITE), kh="center")
r += 1
section("Ordem recomendada de preenchimento")
for k, text in enumerate([
    "Aba PDCA: preencha o cabeçalho, com o tema, o dono, a descrição do problema, o indicador, a situação inicial e a meta.",
    "Aba PDCA, etapas 1 a 3: registre o problema, o que foi observado e as causas confirmadas com dados.",
    "Aba Plano de ação: detalhe as ações no formato 5W2H. Depois, conclua a etapa 4 na aba PDCA.",
    "Aba Verificação: registre as medições de antes das ações.",
    "Aba Plano de ação: atualize o status durante a execução. Registre os desvios na etapa 5 da aba PDCA.",
    "Aba Verificação: registre as medições de depois das ações e leia o resultado do ciclo.",
    "Aba Padronização: registre padrões, treinamentos, pendências e lições aprendidas.",
    "Aba PDCA, etapas 7 e 8: registre o que virou padrão e qual é o próximo ciclo.",
    "Aba Checklist: valide o ciclo com a equipe e com o dono do problema.",
], 1):
    line(f"Passo {k}", text, height=31.5 if len(text) > 95 else 19.5)
r += 1
section("Abas da planilha")
for k, text in [
    ("PDCA", "Modelo principal. Cabeçalho do ciclo, registro das oito etapas e resumo automático."),
    ("Plano de ação", "Ações no formato 5W2H, com situação calculada a partir do prazo e do status."),
    ("Verificação", "Medições do indicador, comparação com a meta e gráfico de tendência."),
    ("Checklist", "Doze verificações de qualidade do ciclo, com percentual de conclusão."),
    ("Padronização", "Registro de padrões, treinamentos, pendências e lições aprendidas."),
    ("Exemplo 1 - Pizzaria", "Ciclo preenchido em que a meta foi atingida: reduzir os atrasos nas entregas."),
    ("Exemplo 2 - Compras", "Ciclo preenchido em que a meta não foi atingida: reduzir o prazo de compra."),
]:
    line(k, text)
r += 1
section("Regras e premissas do modelo")
for k, text in [
    ("Etapas", "O ciclo tem oito etapas fixas, as mesmas do MASP. O aviso de sequência aparece quando uma etapa está concluída e a anterior não. O PDCA permite voltar etapas, mas não pular."),
    ("Indicador e meta", "Digite a situação inicial, a meta e as medições só com números, na unidade do indicador: 82 para 82%. O sentido do indicador define a comparação com a meta."),
    ("Resultado do ciclo", "A aba Verificação compara a média das medições marcadas como “Depois” com a meta. As medições marcadas como “Durante” aparecem no gráfico, mas não entram nas médias."),
    ("Ação atrasada", "Uma ação é considerada atrasada quando o prazo é anterior à data de hoje e o status não é Concluída nem Cancelada."),
    ("Capacidade", "A planilha comporta 15 ações, 16 medições e 15 itens de padronização. Se faltar espaço, o ciclo provavelmente está grande demais: divida o problema."),
    ("Exemplos", "Os dados das abas de exemplo são ilustrativos, criados para este material."),
]:
    line(k, text, height=45.75 if len(text) > 100 else 31.5)
setup(ws, INK, f"B1:C{r}", landscape=False, fit_height=True)

# ------------------------------------------------------------------ PDCA
pdca_sheet(wb.create_sheet("PDCA"), PH["P"][0])

# ------------------------------------------------------------------ Plano de ação
ws = wb.create_sheet("Plano de ação")
widths(ws, {"A": 2, "B": 5, "C": 36, "D": 30, "E": 18, "F": 16, "G": 13, "H": 34, "I": 14, "J": 15, "K": 14, "L": 2})
title(ws, "Plano de ação (5W2H)", "Uma linha por ação. Cada ação ataca uma causa confirmada e tem um único responsável.", "K")
for col, text in zip("BCDEFGHIJK", ["#", "O quê (ação)", "Por quê (causa atacada)", "Quem", "Onde", "Quando (prazo)",
                                     "Como", "Quanto custa", "Status", "Situação"]):
    head(ws, f"{col}4", text)
ws.row_dimensions[4].height = 21.75
ex(ws, "B5", "Ex.", h="center")
ex(ws, "C5", "Reforçar a escala de entregadores às sextas e sábados, das 19h às 22h")
ex(ws, "D5", "A escala não acompanha o pico de pedidos")
ex(ws, "E5", "Gerente da loja")
ex(ws, "F5", "Loja")
ex(ws, "G5", date(2026, 10, 30), h="center", fmt=DATE)
ex(ws, "H5", "Contratar dois entregadores parceiros para o horário de pico")
ex(ws, "I5", 600, h="center", fmt='"R$" #,##0.00')
ex(ws, "J5", "Em andamento", h="center")
ex(ws, "K5", "No prazo", h="center")
ws.row_dimensions[5].height = 31.5
A1, A2 = 6, 20
for k in range(15):
    row = A1 + k
    num(ws, f"B{row}", k + 1)
    for col in "CDEFH":
        inp(ws, f"{col}{row}")
    inp(ws, f"G{row}", h="center", fmt=DATE)
    inp(ws, f"I{row}", h="center", fmt='"R$" #,##0.00')
    inp(ws, f"J{row}", h="center")
    calc(ws, f"K{row}", f'=IF(C{row}="","",IF(OR(J{row}="Concluída",J{row}="Cancelada"),J{row},'
         f'IF(G{row}="","Sem prazo",IF(G{row}<TODAY(),"Atrasada","No prazo"))))', b=False)
    ws.row_dimensions[row].height = 30
PLAN = ["Não iniciada", "Em andamento", "Concluída", "Cancelada"]
dv_list(ws, f"J{A1}:J{A2}", PLAN, "Não iniciada, Em andamento, Concluída ou Cancelada")
dv_date(ws, f"G{A1}:G{A2}")
dv_number(ws, f"I{A1}:I{A2}")
cf_equal(ws, f"J{A1}:J{A2}", STATUS_CF)
cf_equal(ws, f"K{A1}:K{A2}", [("Concluída", GREEN), ("No prazo", GREEN), ("Sem prazo", YELLOW), ("Atrasada", RED)])
band(ws, 22, "Resumo automático", "K")
summary(ws, 23, "Ações registradas", f"=COUNTA(C{A1}:C{A2})", "G", "H")
summary(ws, 24, "Concluídas", f'=COUNTIF(J{A1}:J{A2},"Concluída")', "G", "H")
summary(ws, 25, "Em andamento", f'=COUNTIF(J{A1}:J{A2},"Em andamento")', "G", "H")
summary(ws, 26, "Não iniciadas", f'=COUNTIF(J{A1}:J{A2},"Não iniciada")', "G", "H")
summary(ws, 27, "Atrasadas", f'=COUNTIF(K{A1}:K{A2},"Atrasada")', "G", "H")
summary(ws, 28, "Ações sem responsável ou sem prazo", f'=SUMPRODUCT((C{A1}:C{A2}<>"")*((E{A1}:E{A2}="")+(G{A1}:G{A2}="")>0))', "G", "H")
summary(ws, 29, "Percentual concluído", "=IF(H23=0,0,H24/H23)", "G", "H", fmt="0%")
summary(ws, 30, "Custo total previsto", f"=SUM(I{A1}:I{A2})", "G", "H", fmt='"R$" #,##0.00')
ws.freeze_panes = "C5"
setup(ws, PH["D"][0], "B1:K30", fit_height=True)

# ------------------------------------------------------------------ Verificação
ws = wb.create_sheet("Verificação")
widths(ws, PDCA_W)
title(ws, "Verificação — medições do indicador",
      "As células cinza do cabeçalho vêm da aba PDCA. Preencha as colunas em amarelo-claro.", "H")
band(ws, 4, "Indicador e meta", "H", color=PH["C"][0], sz=11)
for r_, l1, c1, l2, c2, f2 in [(5, "Indicador", "D8", "Unidade", "H8", None),
                                (6, "Situação inicial", "D9", "Meta", "H9", None),
                                (7, "Sentido", "D10", "Prazo da meta", "H10", DATE)]:
    label(ws, f"B{r_}", l1, merge=f"B{r_}:C{r_}")
    calc(ws, f"D{r_}", f'=IF(PDCA!{c1}="","",PDCA!{c1})', h="left", b=False, merge=f"D{r_}:F{r_}")
    label(ws, f"G{r_}", l2)
    calc(ws, f"H{r_}", f'=IF(PDCA!{c2}="","",PDCA!{c2})', h="left", b=False, fmt=f2)
    ws.row_dimensions[r_].height = 21.75
_, _, free = measurements(ws, 9, "$H$6", "$D$7")
setup(ws, PH["C"][0], f"B1:H{free}", landscape=False, fit_height=True)

# ------------------------------------------------------------------ Checklist
ws = wb.create_sheet("Checklist")
widths(ws, {"A": 2, "B": 5, "C": 70, "D": 14, "E": 44, "F": 2})
title(ws, "Checklist de validação do ciclo PDCA", "Escolha o status de cada item na lista suspensa (coluna em amarelo-claro).", "E")
for col, text in zip("BCDE", ["#", "Item de verificação", "Status", "Observações"]):
    head(ws, f"{col}4", text)
ws.row_dimensions[4].height = 21.75
ITEMS = [
    "O problema está descrito com fatos e dados, sem causa nem solução embutida.",
    "Existe um indicador, e a situação inicial foi medida.",
    "A meta tem objetivo, valor e prazo.",
    "O problema foi observado no local e os dados foram estratificados.",
    "A causa raiz foi identificada e confirmada com dados.",
    "Cada ação do plano ataca uma causa confirmada e tem responsável e prazo.",
    "As pessoas envolvidas foram treinadas, e os desvios da execução foram registrados.",
    "O resultado foi medido com o mesmo indicador e o mesmo método da situação inicial.",
    "O resultado foi comparado com a meta, e os efeitos colaterais foram avaliados.",
    "O que funcionou virou padrão: documentos atualizados e pessoas treinadas.",
    "As lições aprendidas e as pendências foram registradas.",
    "O próximo ciclo foi definido.",
]
for k, text in enumerate(ITEMS):
    row = 5 + k
    num(ws, f"B{row}", k + 1)
    put(ws, f"C{row}", text, bg=GRAY)
    inp(ws, f"D{row}", h="center")
    inp(ws, f"E{row}")
    ws.row_dimensions[row].height = 30
dv_list(ws, "D5:D16", YESNO, "Sim, Parcial ou Não")
cf_equal(ws, "D5:D16", YESNO_CF)
band(ws, 18, "Resumo automático", "E")
summary(ws, 19, "Itens atendidos (Sim)", '=COUNTIF(D5:D16,"Sim")', "C", "D")
summary(ws, 20, "Itens parciais", '=COUNTIF(D5:D16,"Parcial")', "C", "D")
summary(ws, 21, "Itens não atendidos", '=COUNTIF(D5:D16,"Não")', "C", "D")
summary(ws, 22, "Itens sem resposta", "=COUNTBLANK(D5:D16)", "C", "D")
summary(ws, 23, "Percentual atendido", "=D19/ROWS(D5:D16)", "C", "D", fmt="0%")
summary(ws, 24, "Situação", '=IF(D19=ROWS(D5:D16),"Validado",IF(D22>0,"Em andamento","Com pendências"))', "C", "D")
setup(ws, PURPLE, "B1:E24", landscape=False, fit_height=True)

# ------------------------------------------------------------------ Padronização
ws = wb.create_sheet("Padronização")
widths(ws, {"A": 2, "B": 5, "C": 40, "D": 18, "E": 28, "F": 40, "G": 22, "H": 14, "I": 16, "J": 2})
title(ws, "Padronização e lições aprendidas",
      "Registre aqui o que vira padrão, quem precisa ser treinado, o que ficou pendente e o que se aprendeu.", "I")
for col, text in zip("BCDEFGHI", ["#", "Item", "Tipo", "Documento ou processo afetado", "Ação", "Responsável", "Prazo", "Status"]):
    head(ws, f"{col}4", text)
ws.row_dimensions[4].height = 21.75
ex(ws, "B5", "Ex.", h="center")
ex(ws, "C5", "Nova escala de entregadores para sextas e sábados")
ex(ws, "D5", "Padrão", h="center")
ex(ws, "E5", "Instrução de trabalho da expedição")
ex(ws, "F5", "Revisar a instrução e divulgar à equipe")
ex(ws, "G5", "Gerente da loja")
ex(ws, "H5", date(2026, 10, 30), h="center", fmt=DATE)
ex(ws, "I5", "Aberta", h="center")
ws.row_dimensions[5].height = 31.5
for k in range(15):
    row = 6 + k
    num(ws, f"B{row}", k + 1)
    for col in "CEFG":
        inp(ws, f"{col}{row}")
    inp(ws, f"D{row}", h="center")
    inp(ws, f"H{row}", h="center", fmt=DATE)
    inp(ws, f"I{row}", h="center")
    ws.row_dimensions[row].height = 30
dv_list(ws, "D6:D20", ["Padrão", "Treinamento", "Pendência", "Lição aprendida"], "Padrão, Treinamento, Pendência ou Lição aprendida")
dv_list(ws, "I6:I20", OPEN, "Aberta, Em andamento ou Concluída")
dv_date(ws, "H6:H20")
cf_equal(ws, "I6:I20", OPEN_CF)
band(ws, 22, "Resumo automático", "I")
summary(ws, 23, "Itens registrados", "=COUNTA(C6:C20)", "F", "G")
summary(ws, 24, "Abertos", '=COUNTIF(I6:I20,"Aberta")', "F", "G")
summary(ws, 25, "Em andamento", '=COUNTIF(I6:I20,"Em andamento")', "F", "G")
summary(ws, 26, "Concluídos", '=COUNTIF(I6:I20,"Concluída")', "F", "G")
summary(ws, 27, "Padrões e treinamentos ainda não concluídos",
        '=COUNTIFS(D6:D20,"Padrão",I6:I20,"<>Concluída")+COUNTIFS(D6:D20,"Treinamento",I6:I20,"<>Concluída")', "F", "G")
summary(ws, 28, "Lições aprendidas registradas", '=COUNTIF(D6:D20,"Lição aprendida")', "F", "G")
ws.freeze_panes = "C5"
setup(ws, PH["A"][0], "B1:I28", fit_height=True)

# ------------------------------------------------------------------ Exemplos
pdca_sheet(wb.create_sheet("Exemplo 1 - Pizzaria"), MUTED, EX1)
pdca_sheet(wb.create_sheet("Exemplo 2 - Compras"), MUTED, EX2)

wb.active = 1
wb.save(OUT)
print("ok", OUT)
