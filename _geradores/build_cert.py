# -*- coding: utf-8 -*-
"""Gera Certificacao-modelo.xlsx no mesmo padrão visual das planilhas anteriores da série."""
import math
import os
import sys

_here = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, _here)
# funções de formatação compartilhadas com o gerador do PDCA (bloco inicial do arquivo)
_src = open(os.path.join(_here, "build_pdca.py"), encoding="utf-8").read()
exec(_src[: _src.index("STATUS = [")])

from datetime import date  # noqa: E402

from openpyxl.chart import BarChart, Series  # noqa: E402
from openpyxl.chart.label import DataLabelList  # noqa: E402
from cert_data import (ABERTA, AGUARDA, ATRASADA, AUDITORIAS, AVISO, CHECK, CONSIDERAR, CRITERIOS, E2015, EDICOES, EVENTOS, EX1, EX2, F1, F2, FASES_MESES,  # noqa: E402
                       FECHADA, FIM_TRANSICAO, FORAPRAZO, MAIOR, MAIOR_DIAS, MANUT_MESES, MENOR, NA, NAOPRONTA, NOPRAZO, OM, PLANO_DIAS, PLANOATR, PONTO,
                       PREVISTA, PRONTA1, PRONTA2, REALIZADA, RECERT_ANTES, STATUS, TIPOS, TRATAR, VALIDADE_MESES, VENCE, VENCIDA)

BLUE, TEAL, AMBER, PURPLE, REDC = "2B5C8A", "1E7B73", "A96A12", "7A4A9A", "B0413E"
S2_T = "F8EBCB"
YESNO = ["Sim", "Parcial", "Não"]
YESNO_CF = [("Sim", GREEN), ("Parcial", YELLOW), ("Não", RED)]
NC_ = 40
P1, P2 = 7, 7 + len(CRITERIOS) - 1      # prontidão
V1, V2 = 13, 13 + len(EVENTOS) - 1      # eventos do ciclo
C1, C2 = 7, 7 + NC_ - 1                 # constatações
W = {"A": 2, "B": 6, "C": 44, "D": 13, "E": 12, "F": 15, "G": 30, "H": 40, "I": 16, "J": 12, "K": 12, "L": 12, "M": 12, "N": 24, "O": 28, "P": 2}
ST_CF = [("Sim", GREEN), ("Parcial", YELLOW), ("Não", RED)]
RES_CF = [(PRONTA2, GREEN), (PRONTA1, S2_T), (NAOPRONTA, RED)]
EV_CF = [(REALIZADA, GREEN), (NOPRAZO, GREEN), (FORAPRAZO, RED), (ATRASADA, RED), (VENCE, S2_T), (PREVISTA, GRAY), (NA, GRAY)]
CT_CF = [(FECHADA, GREEN), (PLANOATR, RED), (VENCIDA, RED), (AGUARDA, S2_T), (ABERTA, S2_T), (TRATAR, S2_T), (CONSIDERAR, GRAY)]


def note(ws, ref, text):
    c = Comment(text, "Modelo Certificação")
    c.width, c.height = 300, 120
    ws[ref].comment = c


def cf_warn(ws, rng_, first, ok_values=("OK",)):
    cond = ",".join(f'{first}="{v}"' for v in ok_values)
    ws.conditional_formatting.add(rng_, FormulaRule(formula=[f"OR({cond})"], fill=PatternFill("solid", bgColor=GREEN, fgColor=GREEN)))
    ws.conditional_formatting.add(rng_, FormulaRule(formula=[f"AND(LEN({first})>0,NOT(OR({cond})))"],
                                  fill=PatternFill("solid", bgColor=YELLOW, fgColor=YELLOW)))


def cf_texto(ws, rng_, first, pares, resto=None):
    for t, cor in pares:
        ws.conditional_formatting.add(rng_, FormulaRule(formula=[f'{first}="{t}"'], stopIfTrue=True, fill=PatternFill("solid", bgColor=cor, fgColor=cor)))
    if resto:
        ws.conditional_formatting.add(rng_, FormulaRule(formula=[f"LEN({first})>0"], fill=PatternFill("solid", bgColor=resto, fgColor=resto)))


def hint_row(ws, row, cells, height=30):
    for ref, text in cells:
        put(ws, f"{ref}{row}", text, f=font(9, i=True, c=MUTED), bg=GRAY, h="center")
    ws.row_dimensions[row].height = height


def alt(pares, minimo=21.75, linha=12):
    n = max(max(1, math.ceil(len(str(t or "")) / max(c, 1))) for t, c in pares)
    return max(minimo, n * linha + 7)


# ------------------------------------------------------------------ prontidão (colunas iguais na aba e nos exemplos)
def criterio_fixo(ws, rr, k, c):
    num(ws, f"B{rr}", k + 1)
    put(ws, f"C{rr}", c[2], bg=GRAY, f=font(9))
    put(ws, f"D{rr}", c[0], bg=GRAY, f=font(9), h="center")
    put(ws, f"E{rr}", "Sim" if c[1] else "—", bg=GRAY, f=font(9, c[1]), h="center")


def pront_formulas(ws, rr, merge=None):
    calc(ws, f"K{rr}", f'=IF(F{rr}="","",IF(F{rr}="Sim",1,IF(F{rr}="Parcial",0.5,0)))', b=False)
    calc(ws, f"L{rr}", f'=IF(F{rr}="","Falta o status",IF(AND(F{rr}<>"Sim",H{rr}=""),"Falta a ação",IF(AND(F{rr}<>"Sim",J{rr}=""),"Falta o prazo",'
                       f'IF(AND(F{rr}="Sim",G{rr}=""),"Falta a evidência","OK"))))', b=False, sz=9, merge=merge)


def pront_resumo(p1, p2, aba=""):
    """Fórmulas do resumo da prontidão; aba é o prefixo da aba ("'Prontidão'!") quando usadas em outra aba."""
    rg = lambda c: f"{aba}{c}{p1}:{c}{p2}"  # noqa: E731
    imp = lambda fase: f'SUMPRODUCT(({rg("D")}="{fase}")*({rg("E")}="Sim")*({rg("F")}<>"Sim"))'  # noqa: E731
    return dict(pct=f'=IF(COUNTA({rg("F")})=0,"",SUM({rg("K")})/ROWS({rg("K")}))', imp1="=" + imp(F1), imp2="=" + imp(F2),
                res=f'=IF(COUNTA({rg("F")})=0,"",IF({imp(F1)}>0,"{NAOPRONTA}",IF({imp(F2)}>0,"{PRONTA1}","{PRONTA2}")))')


# ------------------------------------------------------------------ ciclo (os parâmetros ficam em células passadas como referência)
def ciclo_eventos(ws, v1, ed, dec, ref, fim, val, fixos=None):
    """Monta as linhas dos eventos a partir de v1. fixos: datas realizadas (exemplo) ou None (entrada)."""
    lim = ['""', f'IF(E{v1}="","",EDATE(E{v1},{FASES_MESES}))', '""', f'IF({dec}="","",EDATE({dec},{MANUT_MESES}))', f'IF({dec}="","",EDATE({dec},{2 * MANUT_MESES}))',
           f'IF({dec}="","",EDATE({dec},{VALIDADE_MESES - RECERT_ANTES}))', f'IF(OR({dec}="",{ed}<>"{E2015}"),"",MIN({fim},{val}))']
    for k, ev in enumerate(EVENTOS):
        rr = v1 + k
        num(ws, f"B{rr}", k + 1)
        put(ws, f"C{rr}", ev, bg=GRAY, f=font(10, True))
        calc(ws, f"D{rr}", "=" + lim[k], b=False, fmt=DATE)
        if k == 2:
            calc(ws, f"E{rr}", f'=IF({dec}="","",{dec})', b=False, fmt=DATE)
        elif fixos is None:
            inp(ws, f"E{rr}", h="center", fmt=DATE)
        else:
            put(ws, f"E{rr}", fixos[k], h="center", fmt=DATE)
        geral = (f'IF(D{rr}="","",IF(E{rr}<>"",IF(E{rr}<=D{rr},"{NOPRAZO}","{FORAPRAZO}"),IF(N({ref})=0,"",IF({ref}>D{rr},"{ATRASADA}",'
                 f'IF(D{rr}-{ref}<={AVISO},"{VENCE}","{PREVISTA}")))))')
        if k in (0, 2):
            sit = f'IF(E{rr}="","{PREVISTA}","{REALIZADA}")'
        elif k == 6:
            sit = f'IF({ed}<>"{E2015}","{NA}",{geral})'
        else:
            sit = geral
        calc(ws, f"F{rr}", "=" + sit, sz=9)
        ws.row_dimensions[rr].height = 21.75
    cf_texto(ws, f"F{v1}:F{v1 + len(EVENTOS) - 1}", f"F{v1}", EV_CF)


# ------------------------------------------------------------------ constatações (C nº, D auditoria, E data, F tipo, G requisito, H texto, I a K datas)
def ct_formulas(ws, rr, ref, merge=None):
    calc(ws, f"L{rr}", f'=IF(OR(E{rr}="",AND(F{rr}<>"{MAIOR}",F{rr}<>"{MENOR}")),"",E{rr}+{PLANO_DIAS})', b=False, fmt=DATE)
    calc(ws, f"M{rr}", f'=IF(OR(E{rr}="",F{rr}<>"{MAIOR}"),"",E{rr}+{MAIOR_DIAS})', b=False, fmt=DATE)
    calc(ws, f"N{rr}", f'=IF(OR(F{rr}="",E{rr}=""),"",IF(OR(F{rr}="{OM}",F{rr}="{PONTO}"),IF(K{rr}<>"","{FECHADA}",IF(F{rr}="{PONTO}","{TRATAR}","{CONSIDERAR}")),'
                       f'IF(K{rr}<>"","{FECHADA}",IF(AND(I{rr}="",N({ref})>L{rr}),"{PLANOATR}",IF(AND(F{rr}="{MAIOR}",N({ref})>M{rr}),"{VENCIDA}",'
                       f'IF(AND(F{rr}="{MENOR}",I{rr}<>""),"{AGUARDA}","{ABERTA}"))))))', sz=9)
    calc(ws, f"O{rr}", f'=IF(AND(C{rr}="",COUNTA(D{rr}:K{rr})=0),"",IF(E{rr}="","Falta a data da auditoria",IF(F{rr}="","Falta o tipo",IF(G{rr}="","Falta o requisito",'
                       f'IF(H{rr}="","Falta a descrição",IF(AND(I{rr}<>"",I{rr}<E{rr}),"Plano antes da auditoria",IF(AND(F{rr}="{MAIOR}",K{rr}<>"",J{rr}=""),'
                       f'"Falta a conclusão da ação","OK")))))))', b=False, sz=9, merge=merge)


def chart_style(ch):
    ch.style = 2
    ch.title = None
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ch.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill="D3DBE0", w=9525))
    ch.x_axis.graphicalProperties = GraphicalProperties(ln=LineProperties(solidFill="D3DBE0", w=9525))
    ch.y_axis.graphicalProperties = GraphicalProperties(ln=LineProperties(noFill=True))


def colunas(ws, anchor, r1, r2, c_cat, c_val, width=16, height=7.5):
    ch = BarChart()
    ch.type = "col"
    ch.gapWidth = 70
    ch.height, ch.width = height, width
    chart_style(ch)
    ch.legend = None
    s = Series(Reference(ws, min_col=c_val, min_row=r1, max_row=r2), title="Constatações")
    s.graphicalProperties.solidFill = BLUE
    s.graphicalProperties.line.noFill = True
    s.dLbls = DataLabelList()
    s.dLbls.showVal = True
    s.dLbls.showSerName = s.dLbls.showCatName = s.dLbls.showLegendKey = False
    ch.append(s)
    ch.set_categories(Reference(ws, min_col=c_cat, min_row=r1, max_row=r2))
    ch.y_axis.scaling.min = 0
    ch.y_axis.majorUnit = 1
    ws.add_chart(ch, anchor)


def resumo_aba(ws, s, itens, last, merge_lab="D"):
    band(ws, s, "Resumo automático", last)
    col = chr(ord(merge_lab) + 1)
    for k, (text, formula, fmt) in enumerate(itens, 1):
        aviso = text == "Aviso"
        label(ws, f"B{s+k}", text, merge=f"B{s+k}:{merge_lab}{s+k}", h="right")
        calc(ws, f"{col}{s+k}", formula, merge=f"{col}{s+k}:{last}{s+k}", sz=9 if aviso else 10, b=not aviso, h="left", fmt=fmt)
        ws.row_dimensions[s + k].height = 21.75
    av = s + len(itens)
    cf_warn(ws, f"{col}{av}:{last}{av}", f"{col}{av}")
    return av


# ================================================================== pasta de trabalho
wb = Workbook()
wb.properties.title = "Processo de certificação — Modelo"
wb.properties.creator = "Modelo Certificação"
wb.properties.language = "pt-BR"

# ------------------------------------------------------------------ Instruções
ws = wb.active
ws.title = "Instruções"
widths(ws, {"A": 2, "B": 26, "C": 90, "D": 2})
title(ws, "Processo de certificação — Como usar esta planilha", "Modelo para avaliar a prontidão, acompanhar o ciclo do certificado e tratar as constatações do organismo.", "C")
r = 4


def section(text):
    global r
    band(ws, r, text, "C", sz=11)
    r += 1


def line(k, v, kbg=GRAY, vbg=None, kf=None, kh="left", height=None):
    global r
    put(ws, f"B{r}", k, f=kf or font(10, True), bg=kbg, h=kh)
    put(ws, f"C{r}", v, bg=vbg)
    ws.row_dimensions[r].height = height or (19.5 if len(v) <= 96 else (31.5 if len(v) <= 192 else 45.75))
    r += 1


section("Legenda: onde preencher")
line("Amarelo-claro", "Células de entrada. É aqui que você digita.", vbg=INPUT)
line("Cinza", "Células calculadas ou fixas (critérios, limites, prazos, situações, conferências). Não altere.", vbg=GRAY)
line("Cinza em itálico", "Linha de exemplo (marcada com “Ex.”). Mostra o formato esperado e não entra nas contagens.", vbg=GRAY)
line("Listas suspensas", "Status, edição, auditoria, tipo de constatação e checklist aceitam apenas as opções da lista.")
line("Comentários", "Células com triângulo vermelho no canto trazem uma dica de preenchimento. Passe o mouse sobre elas.")
r += 1
section("As regras de cálculo")
for k, text in [
    ("Prontidão", f"Sim vale 1, Parcial 0,5, Não 0. {NAOPRONTA}: há impeditivo da fase 1 diferente de Sim. {PRONTA1}: só faltam impeditivos da fase 2. {PRONTA2}: nenhum impeditivo pendente."),
    ("Validade", f"A data da decisão mais {VALIDADE_MESES} meses, menos um dia."),
    ("Limites do ciclo", f"Fase 2: até {FASES_MESES} meses depois da fase 1. Manutenções: {MANUT_MESES} e {2 * MANUT_MESES} meses depois da decisão. "
                         f"Recertificação: auditoria até {RECERT_ANTES} meses antes do vencimento. Transição: o fim da janela ou a validade, o que vier antes, só para a edição de 2015."),
    ("Situação do evento", f"{NOPRAZO} ou {FORAPRAZO}, se realizado. Se não: {ATRASADA} depois do limite, {VENCE} a {AVISO} dias ou menos, {PREVISTA} antes disso."),
    ("Prazos das constatações", f"Plano em até {PLANO_DIAS} dias para as não conformidades. A maior fechada, com evidência aceita, em até {MAIOR_DIAS} dias."),
    ("Situação da constatação", f"{FECHADA} com evidência aceita. {PLANOATR} sem plano depois do prazo. {VENCIDA} se a maior passou de {MAIOR_DIAS} dias. "
                                f"{AGUARDA} para a menor com plano. {TRATAR} para o ponto da fase 1 em aberto."),
]:
    line(k, text)
r += 1
section("Ordem recomendada de preenchimento")
for k, text in enumerate([
    "Aba Prontidão: a data e o status de cada critério, com a evidência, ou a ação, o responsável e o prazo.",
    "Aba Ciclo: o organismo, a edição da norma, a data da decisão (real ou prevista), a data da leitura e o fim da transição.",
    "Aba Ciclo: as datas em que cada auditoria foi realizada.",
    "Aba Constatações: depois de cada auditoria, as constatações e as datas do plano, da ação e da evidência aceita.",
    "Aba Painel: leia os avisos. Depois, valide na aba Checklist.",
], 1):
    line(f"Passo {k}", text)
r += 1
section("Abas da planilha")
for k, text in [
    ("Prontidão", f"Os {len(CRITERIOS)} critérios, com o momento e se são impeditivos. Calcula o percentual, os impeditivos e o resultado."),
    ("Ciclo", "Os dados do certificado e os sete eventos. Calcula a validade, os limites e as situações."),
    ("Constatações", f"Até {NC_} constatações do organismo. Calcula os prazos, a situação e a conferência."),
    ("Painel", "Os números das três abas, o gráfico das constatações e o aviso."),
    ("Checklist", "Doze verificações da certificação, com percentual de conclusão."),
    ("Exemplo 1 - Pizzaria", "Prontidão em 30/11/2027, certificação pela edição de 2026 em 2028."),
    ("Exemplo 2 - Indústria", "Prontidão em 30/09/2027, certificação pela edição de 2015 em dezembro de 2027."),
]:
    line(k, text)
r += 1
section("Premissas do modelo")
for k, text in [
    ("Convenções", "Os prazos de 30 e 90 dias, os 6 meses entre as fases, os 3 meses da recertificação e os critérios de prontidão são uma convenção deste material. Vale o regulamento do seu organismo."),
    ("Transição", f"A ISO 9001:2026 foi publicada em 16/09/2026. O fim da transição vem preenchido com {FIM_TRANSICAO.strftime('%d/%m/%Y')}, fixado pela Global ACI; desde 31/03/2028, certificados novos só pela edição de 2026. Confirme com o seu organismo."),
    ("Exemplos", "Os dados das abas de exemplo são ilustrativos, criados para este material."),
]:
    line(k, text)
setup(ws, INK, f"B1:C{r}", landscape=False, fit_height=True)

# ------------------------------------------------------------------ Prontidão
ws = wb.create_sheet("Prontidão")
widths(ws, {**W, "C": 52, "I": 20})
title(ws, "Prontidão para a certificação", "Um critério por linha. Os impeditivos precisam estar em Sim antes da fase indicada.", "L")
label(ws, "B4", "Data da avaliação", merge="B4:C4")
inp(ws, "D4", h="center", fmt=DATE)
dv_date(ws, "D4")
for col, text in zip("BCDEFGHIJKL", ["#", "Critério", "Antes da", "Impeditivo", "Status", "Evidência", "Ação, se não está pronto", "Responsável", "Prazo", "Pontos", "Conferência"]):
    head(ws, f"{col}6", text)
ws.row_dimensions[6].height = 33
for k, c in enumerate(CRITERIOS):
    rr = P1 + k
    criterio_fixo(ws, rr, k, c)
    inp(ws, f"F{rr}", h="center")
    for col in "GHI":
        inp(ws, f"{col}{rr}")
    inp(ws, f"J{rr}", h="center", fmt=DATE)
    pront_formulas(ws, rr)
    ws.row_dimensions[rr].height = 30
dv_list(ws, f"F{P1}:F{P2}", STATUS, "Sim, Parcial ou Não")
dv_date(ws, f"J{P1}:J{P2}")
cf_texto(ws, f"F{P1}:F{P2}", f"F{P1}", ST_CF)
ws.conditional_formatting.add(f"F{P1}:F{P2}", FormulaRule(formula=[f'AND($E{P1}="Sim",F{P1}<>"",F{P1}<>"Sim")'], fill=PatternFill("solid", bgColor=RED, fgColor=RED)))
cf_warn(ws, f"L{P1}:L{P2}", f"L{P1}")
note(ws, "E6", "Impeditivo: precisa estar em Sim antes da fase indicada. Os outros pesam no percentual, mas não impedem.")
s = P2 + 2
R = pront_resumo(P1, P2)
PAV = resumo_aba(ws, s, [
    ("Percentual atendido", R["pct"], "0%"),
    ("Impeditivos antes da fase 1", R["imp1"], None),
    ("Impeditivos antes da fase 2", R["imp2"], None),
    ("Resultado", R["res"], None),
    ("Aviso", f'=IF(COUNTA(F{P1}:F{P2})=0,"Responda os critérios",IF(E{s+2}>0,"Há impeditivo pendente antes da fase 1",IF(E{s+3}>0,"Há impeditivo pendente antes da fase 2",'
              f'IF(SUMPRODUCT((L{P1}:L{P2}<>"OK")*1)>0,"Há critério a completar: veja a coluna Conferência","OK"))))', None),
], "L")
cf_texto(ws, f"E{s+4}:L{s+4}", f"$E{s+4}", RES_CF)
ws.freeze_panes = f"D{P1}"
setup(ws, BLUE, f"B1:L{PAV}")

# ------------------------------------------------------------------ Ciclo
ws = wb.create_sheet("Ciclo")
widths(ws, {"A": 2, "B": 6, "C": 36, "D": 14, "E": 14, "F": 20, "G": 44, "H": 14, "I": 2})
title(ws, "Ciclo do certificado", "Os dados do certificado e os eventos do ciclo, com os limites calculados.", "H")
for rr, text, fmt, dica in [(4, "Organismo", None, "O organismo de certificação, e a acreditação dele."),
                            (5, "Edição da norma", None, "A edição pela qual o certificado foi, ou será, emitido."),
                            (6, "Data da decisão de certificação", DATE, "A data real, ou a prevista, para simular o ciclo."),
                            (7, "Data da leitura", DATE, "Os eventos são comparados com esta data."),
                            (8, "Fim da transição", DATE, "Prazo para os certificados de 2015 migrarem. Confirme no comunicado da acreditação.")]:
    label(ws, f"B{rr}", text, merge=f"B{rr}:C{rr}")
    inp(ws, f"D{rr}", h="center" if fmt else "left", fmt=fmt, merge=None if fmt else f"D{rr}:E{rr}")
    put(ws, f"F{rr}", dica, f=font(9, i=True, c=MUTED), bg=GRAY, merge=f"F{rr}:G{rr}")
ws["D8"] = FIM_TRANSICAO
dv_list(ws, "D5", EDICOES, "ISO 9001:2015 ou ISO 9001:2026")
dv_date(ws, "D6:D8")
label(ws, "B9", "Validade do certificado", merge="B9:C9")
calc(ws, "D9", f'=IF(D6="","",EDATE(D6,{VALIDADE_MESES})-1)', fmt=DATE)
for col, text in zip("BCDEFGH", ["#", "Evento", "Limite", "Realizado em", "Situação", "Observação", "Limite pendente"]):
    head(ws, f"{col}12", text)
ws.row_dimensions[12].height = 21.75
ciclo_eventos(ws, V1, "$D$5", "$D$6", "$D$7", "$D$8", "$D$9")
for k in range(len(EVENTOS)):
    inp(ws, f"G{V1 + k}")
    # o limite de cada evento ainda não feito; o painel tira daqui o próximo limite
    calc(ws, f"H{V1 + k}", f'=IF(AND(D{V1 + k}<>"",E{V1 + k}=""),D{V1 + k},"")', b=False, fmt=DATE)
note(ws, "H12", "Calculada: o limite dos eventos que ainda não aconteceram. O painel usa o menor deles como o próximo limite.")
dv_date(ws, f"E{V1}:E{V2}")
note(ws, "E12", "A data em que a auditoria aconteceu. A data da decisão vem do campo de cima.")
s = V2 + 2
F_ = f"F{V1}:F{V2}"
VAV = resumo_aba(ws, s, [
    ("Validade", '=IF(D9="","",D9)', DATE),
    ("Eventos", "=" + '&", "&'.join(f'COUNTIF({F_},"{x}")&" {x.lower()}"' for x in (REALIZADA, NOPRAZO, FORAPRAZO, VENCE, ATRASADA, PREVISTA)), None),
    ("Aviso", f'=IF(D6="","Informe a data da decisão, real ou prevista",IF(D7="","Informe a data da leitura",IF(COUNTIF({F_},"{ATRASADA}")>0,"Há evento atrasado: fale com o organismo",'
              f'IF(COUNTIF({F_},"{VENCE}")>0,"Há evento que vence em {AVISO} dias: confirme a data com o organismo","OK"))))', None),
], "H")
setup(ws, TEAL, f"B1:H{VAV}", landscape=False, fit_height=True)

# ------------------------------------------------------------------ Constatações
ws = wb.create_sheet("Constatações")
widths(ws, {**W, "C": 8, "D": 15, "H": 48, "I": 12})
title(ws, "Constatações do organismo", "Uma constatação por linha, com as datas do plano, da ação e da evidência aceita.", "O")
for col, text in zip("BCDEFGHIJKLMNO", ["#", "Nº", "Auditoria", "Data", "Tipo", "Requisito", "Constatação", "Plano enviado", "Ação concluída", "Evidência aceita",
                                        "Prazo do plano", "Prazo para fechar", "Situação", "Conferência"]):
    head(ws, f"{col}4", text)
ws.row_dimensions[4].height = 33
hint_row(ws, 5, [("B", ""), ("C", ""), ("D", "Lista"), ("E", "Da auditoria"), ("F", "Lista"), ("G", "Ex.: 8.5.6"), ("H", "Como o organismo escreveu"), ("I", "Data"), ("J", "Data"),
                 ("K", "Data"), ("L", "Calculado"), ("M", "Calculado"), ("N", "Calculada"), ("O", "Calculada")])
ex(ws, "B6", "Ex.", h="center")
for col, v, h_, fmt in [("C", "I-4", "center", None), ("D", "Fase 2", "center", None), ("E", date(2027, 11, 23), "center", DATE), ("F", MENOR, "center", None),
                        ("G", "7.2", "center", None), ("H", "Operador novo sem registro de treinamento", "left", None), ("I", date(2027, 12, 5), "center", DATE),
                        ("J", None, "center", None), ("K", None, "center", None), ("L", date(2027, 12, 23), "center", DATE), ("M", None, "center", None),
                        ("N", AGUARDA, "center", None), ("O", "OK", "center", None)]:
    ex(ws, f"{col}6", v, h=h_, fmt=fmt)
ws.row_dimensions[6].height = 30
REF = "Ciclo!$D$7"
for k in range(NC_):
    rr = C1 + k
    num(ws, f"B{rr}", k + 1)
    for col in "CDFG":
        inp(ws, f"{col}{rr}", h="center")
    inp(ws, f"H{rr}")
    for col in "EIJK":
        inp(ws, f"{col}{rr}", h="center", fmt=DATE)
    ct_formulas(ws, rr, REF)
    ws.row_dimensions[rr].height = 30
dv_list(ws, f"D{C1}:D{C2}", AUDITORIAS, "Em que auditoria a constatação foi feita")
dv_list(ws, f"F{C1}:F{C2}", TIPOS, "NC maior, NC menor, oportunidade de melhoria ou ponto de atenção da fase 1")
dv_date(ws, f"E{C1}:E{C2}")
dv_date(ws, f"I{C1}:K{C2}")
cf_texto(ws, f"N{C1}:N{C2}", f"N{C1}", CT_CF)
cf_texto(ws, f"F{C1}:F{C2}", f"F{C1}", [(MAIOR, RED), (MENOR, S2_T)])
cf_warn(ws, f"O{C1}:O{C2}", f"O{C1}")
note(ws, "N4", "A situação é lida na data da leitura da aba Ciclo.")
note(ws, "K4", "A data em que o organismo aceitou a evidência. Para a NC maior, é ela que permite a decisão.")
s = C2 + 2
N_ = f"N{C1}:N{C2}"
F__ = f"F{C1}:F{C2}"
CAV = resumo_aba(ws, s, [
    ("Constatações", f'=COUNTA({F__})&" constatações: "&' + '&", "&'.join(f'COUNTIF({F__},"{x}")&" {x.lower()}"' for x in TIPOS), None),
    ("Situação", "=" + '&", "&'.join(f'COUNTIF({N_},"{x}")&" {x.lower()}"' for x in (FECHADA, AGUARDA, ABERTA, PLANOATR, VENCIDA)), None),
    ("Aviso", f'=IF(COUNTA({F__})=0,"Registre as constatações depois de cada auditoria",IF(COUNTIF({N_},"{VENCIDA}")>0,"Há NC maior vencida: o certificado está em risco",'
              f'IF(COUNTIF({N_},"{PLANOATR}")>0,"Há plano atrasado: envie ao organismo",IF(SUMPRODUCT((O{C1}:O{C2}<>"")*(O{C1}:O{C2}<>"OK"))>0,'
              f'"Há constatação a completar: veja a coluna Conferência","OK"))))', None),
], "O")
ws.freeze_panes = f"E{C1}"
setup(ws, AMBER, f"B1:O{CAV}")

# ------------------------------------------------------------------ Painel
ws = wb.create_sheet("Painel")
widths(ws, {"A": 2, "B": 38, "C": 16, "D": 14, "E": 14, "F": 34, "G": 2})
title(ws, "Painel da certificação", "Nada a preencher: tudo vem das outras abas.", "F")
label(ws, "B3", "Organismo")
calc(ws, "C3", '=IF(Ciclo!D4="","",Ciclo!D4)', merge="C3:F3", h="left")
for col, text, m in [("B", "Indicador", None), ("C", "Resultado", None), ("D", "Como se lê", "D5:F5")]:
    put(ws, f"{col}5", text, f=font(10, True, c=WHITE), bg=INK, h="center", merge=m)
ws.row_dimensions[5].height = 21.75
Pf = f"'Prontidão'!F{P1}:F{P2}"
Pl = f"'Prontidão'!L{P1}:L{P2}"
Vd, Ve, Vf, Vh = (f"Ciclo!{c}{V1}:{c}{V2}" for c in "DEFH")
Cf, Cn, Co = (f"'Constatações'!{c}{C1}:{c}{C2}" for c in "FNO")
RP = pront_resumo(P1, P2, "'Prontidão'!")
IND = [
    ("PRONTIDÃO", None, None, None),
    ("Data da avaliação", "=IF('Prontidão'!D4=\"\",\"\",'Prontidão'!D4)", "Da aba Prontidão.", DATE),
    ("Percentual atendido", RP["pct"], "Sim vale 1, Parcial 0,5.", "0%"),
    ("Impeditivos antes da fase 1", RP["imp1"], "Precisam estar em Sim antes da fase 1.", None),
    ("Impeditivos antes da fase 2", RP["imp2"], "Precisam estar em Sim antes da fase 2.", None),
    ("Resultado", RP["res"], "Pronta para a fase 2, para a fase 1, ou ainda não.", None),
    ("CICLO", None, None, None),
    ("Edição", '=IF(Ciclo!D5="","",Ciclo!D5)', "Do certificado.", None),
    ("Validade", '=IF(Ciclo!D9="","",Ciclo!D9)', "Decisão mais três anos, menos um dia.", DATE),
    ("Próximo limite", f'=IF(MIN({Vh})=0,"",MIN({Vh}))', "O limite mais próximo de um evento ainda não feito.", DATE),
    ("Próximo evento", f'=IF(C{{prox}}="","",IFERROR(INDEX(Ciclo!C{V1}:C{V2},MATCH(C{{prox}},{Vh},0)),""))', "O evento desse limite.", None),
    ("Eventos atrasados", f'=COUNTIF({Vf},"{ATRASADA}")', "Falar com o organismo.", None),
    (f"Eventos que vencem em {AVISO} dias", f'=COUNTIF({Vf},"{VENCE}")', "Confirmar a data.", None),
    ("CONSTATAÇÕES", None, None, None),
    ("Constatações", f"=COUNTA({Cf})", "Da aba Constatações.", None),
    ("NC maiores", f'=COUNTIF({Cf},"{MAIOR}")', "Precisam fechar antes da decisão.", None),
    ("NC maiores abertas", f'=SUMPRODUCT(({Cf}="{MAIOR}")*({Cn}<>"{FECHADA}"))', "Sem evidência aceita.", None),
    ("Planos atrasados", f'=COUNTIF({Cn},"{PLANOATR}")', "Enviar ao organismo.", None),
    ("NC maiores vencidas", f'=COUNTIF({Cn},"{VENCIDA}")', "Risco ao certificado.", None),
    ("Aguardando verificação", f'=COUNTIF({Cn},"{AGUARDA}")', "Menores a verificar na próxima auditoria.", None),
    ("Linhas a completar", f'=SUMPRODUCT(({Pl}<>"")*({Pl}<>"OK"))+SUMPRODUCT(({Co}<>"")*({Co}<>"OK"))', "Conferências diferentes de OK.", None),
]
IR = {}
rr = 6
for nome, formula, leit, fmt in IND:
    if formula is None:
        put(ws, f"B{rr}", nome, f=font(9, True, c=MUTED), bg=None, box=False, merge=f"B{rr}:F{rr}")
        ws.row_dimensions[rr].height = 21.75
        rr += 1
        continue
    IR[nome] = rr
    if "{prox}" in formula:
        formula = formula.replace("{prox}", str(IR["Próximo limite"]))
    label(ws, f"B{rr}", nome)
    calc(ws, f"C{rr}", formula, fmt=fmt)
    put(ws, f"D{rr}", leit, f=font(9, c=MUTED), bg=GRAY, merge=f"D{rr}:F{rr}")
    ws.row_dimensions[rr].height = 21.75
    rr += 1
s = rr + 1
band(ws, s, "Resumo automático", "F")
label(ws, f"B{s+1}", "Aviso")
C = lambda n: f"C{IR[n]}"  # noqa: E731
calc(ws, f"C{s+1}", f'=IF(AND(COUNTA({Pf})=0,Ciclo!D6="",{C("Constatações")}=0),"Preencha as abas Prontidão, Ciclo e Constatações",'
     f'IF({C("NC maiores vencidas")}>0,"Há NC maior vencida: o certificado está em risco",IF({C("Planos atrasados")}>0,"Há plano atrasado: envie ao organismo",'
     f'IF({C("Eventos atrasados")}>0,"Há evento do ciclo atrasado",IF(N({C("Impeditivos antes da fase 1")})+N({C("Impeditivos antes da fase 2")})>0,"Há impeditivo pendente na prontidão",'
     f'IF({C(f"Eventos que vencem em {AVISO} dias")}>0,"Há evento que vence em {AVISO} dias",IF({C("NC maiores abertas")}>0,"Há NC maior aberta: feche antes da decisão",'
     f'IF({C("Linhas a completar")}>0,"Há linha a completar: veja as conferências","OK"))))))))', merge=f"C{s+1}:F{s+1}", sz=9, b=False, h="left")
cf_warn(ws, f"C{s+1}:F{s+1}", f"C{s+1}")
ws.row_dimensions[s + 1].height = 21.75
NAV = s + 1
GS = s + 4
put(ws, f"B{GS - 1}", "Constatações por tipo", f=font(9, True, c=MUTED), bg=None, box=False)
for k, t in enumerate(TIPOS):
    put(ws, f"B{GS + k}", t, f=font(9), bg=GRAY)
    calc(ws, f"C{GS + k}", f'=COUNTIF({Cf},B{GS + k})', b=False)
colunas(ws, f"D{GS - 1}", GS, GS + len(TIPOS) - 1, 2, 3, width=12, height=7)
setup(ws, REDC, f"B1:F{GS + 14}", landscape=False, fit_height=True)

# ------------------------------------------------------------------ Checklist
ws = wb.create_sheet("Checklist")
widths(ws, {"A": 2, "B": 5, "C": 70, "D": 14, "E": 44, "F": 2})
title(ws, "Checklist da certificação", "Escolha o status de cada item na lista suspensa (coluna em amarelo-claro).", "E")
for col, text in zip("BCDE", ["#", "Item de verificação", "Status", "Observações"]):
    head(ws, f"{col}4", text)
ws.row_dimensions[4].height = 21.75
for k, text in enumerate(CHECK):
    rr = 5 + k
    num(ws, f"B{rr}", k + 1)
    put(ws, f"C{rr}", text, bg=GRAY)
    inp(ws, f"D{rr}", h="center")
    inp(ws, f"E{rr}")
    ws.row_dimensions[rr].height = 30
dv_list(ws, "D5:D16", YESNO, "Sim, Parcial ou Não")
cf_equal(ws, "D5:D16", YESNO_CF)
band(ws, 18, "Resumo automático", "E")
summary(ws, 19, "Itens atendidos (Sim)", '=COUNTIF(D5:D16,"Sim")', "C", "D")
summary(ws, 20, "Itens parciais", '=COUNTIF(D5:D16,"Parcial")', "C", "D")
summary(ws, 21, "Itens não atendidos", '=COUNTIF(D5:D16,"Não")', "C", "D")
summary(ws, 22, "Itens sem resposta", "=COUNTBLANK(D5:D16)", "C", "D")
summary(ws, 23, "Percentual atendido", "=D19/ROWS(D5:D16)", "C", "D", fmt="0%")
summary(ws, 24, "Situação", '=IF(D19=ROWS(D5:D16),"Validado",IF(D22>0,"Em andamento","Com pendências"))', "C", "D")
setup(ws, PURPLE, "B1:E24", landscape=False, fit_height=True)


# ------------------------------------------------------------------ Exemplos
def sub(ws, rr, cells, height=30):
    for a, b, t in cells:
        put(ws, f"{a}{rr}", t, f=font(10, True), bg=GRAY, h="center", merge=b and f"{a}{rr}:{b}{rr}")
    ws.row_dimensions[rr].height = height


def exemplo(ws, ex_):
    H, P, Cc = ex_["head"], ex_["pront"], ex_["ciclo"]
    widths(ws, W)
    title(ws, "Processo de certificação", "Exemplo preenchido, para consulta. Use as abas Prontidão, Ciclo e Constatações para a sua organização.", "O")
    rr = 4
    for rot, val in [("Organização", H["org"]), ("Escopo", H["escopo"]), ("Organismo", H["organismo"]), ("Por que certificar", H["motivo"])]:
        label(ws, f"B{rr}", rot, merge=f"B{rr}:C{rr}")
        inp(ws, f"D{rr}", val, merge=f"D{rr}:O{rr}", bg=WHITE, h="left")
        ws.row_dimensions[rr].height = 19.5
        rr += 1
    pars = {}
    for rot, val, fmt in [("Data da avaliação de prontidão", P["data"], DATE), ("Edição da norma", Cc["edicao"], None), ("Data da decisão", Cc["decisao"], DATE),
                          ("Data da leitura do ciclo", Cc["ref"], DATE), ("Fim da transição", Cc["fimtrans"], DATE)]:
        label(ws, f"B{rr}", rot, merge=f"B{rr}:C{rr}")
        inp(ws, f"D{rr}", val, bg=WHITE, h="center", fmt=fmt)
        pars[rot] = f"$D${rr}"
        rr += 1
    label(ws, f"B{rr}", "Validade do certificado", merge=f"B{rr}:C{rr}")
    calc(ws, f"D{rr}", f'=IF({pars["Data da decisão"]}="","",EDATE({pars["Data da decisão"]},{VALIDADE_MESES})-1)', fmt=DATE)
    val = f"$D${rr}"

    # ---- prontidão
    rr += 2
    band(ws, rr, "Prontidão", "O", color=BLUE)
    rr += 1
    sub(ws, rr, [("B", None, "#"), ("C", None, "Critério"), ("D", None, "Antes da"), ("E", None, "Impeditivo"), ("F", None, "Status"), ("G", None, "Evidência"),
                 ("H", None, "Ação"), ("I", None, "Responsável"), ("J", None, "Prazo"), ("K", None, "Pontos"), ("L", "O", "Conferência")])
    p1 = rr + 1
    for k, (c, i) in enumerate(zip(CRITERIOS, P["itens"])):
        rr += 1
        criterio_fixo(ws, rr, k, c)
        put(ws, f"F{rr}", i["status"], h="center")
        put(ws, f"G{rr}", i["evid"] or None, f=font(9))
        put(ws, f"H{rr}", i["acao"] or None, f=font(9))
        put(ws, f"I{rr}", i["resp"] or None, f=font(9))
        put(ws, f"J{rr}", i["prazo"], h="center", fmt=DATE)
        pront_formulas(ws, rr, merge=f"L{rr}:O{rr}")
        ws.row_dimensions[rr].height = alt([(c[2], 50), (i["evid"], 34), (i["acao"], 44)], minimo=21.75)
    p2 = rr
    cf_texto(ws, f"F{p1}:F{p2}", f"F{p1}", ST_CF)
    ws.conditional_formatting.add(f"F{p1}:F{p2}", FormulaRule(formula=[f'AND($E{p1}="Sim",F{p1}<>"",F{p1}<>"Sim")'], fill=PatternFill("solid", bgColor=RED, fgColor=RED)))
    cf_warn(ws, f"L{p1}:O{p2}", f"$L{p1}")
    R = pront_resumo(p1, p2)
    rr += 1
    for rot, f_, fmt in (("Percentual atendido", R["pct"], "0%"), ("Impeditivos antes da fase 1", R["imp1"], None), ("Impeditivos antes da fase 2", R["imp2"], None),
                         ("Resultado", R["res"], None)):
        label(ws, f"B{rr}", rot, merge=f"B{rr}:C{rr}", h="right")
        calc(ws, f"D{rr}", f_, fmt=fmt, merge=f"D{rr}:F{rr}", h="left")
        ws.row_dimensions[rr].height = 21.75
        rr += 1
    res_row = rr - 1
    cf_texto(ws, f"D{res_row}:F{res_row}", f"$D{res_row}", RES_CF)

    # ---- ciclo
    rr += 1
    ws.row_breaks.append(Break(id=rr - 1))
    band(ws, rr, "Ciclo do certificado", "O", color=TEAL)
    rr += 1
    sub(ws, rr, [("B", None, "#"), ("C", None, "Evento"), ("D", None, "Limite"), ("E", None, "Realizado em"), ("F", None, "Situação")])
    v1 = rr + 1
    ciclo_eventos(ws, v1, pars["Edição da norma"], pars["Data da decisão"], pars["Data da leitura do ciclo"], pars["Fim da transição"], val, fixos=Cc["feito"])
    rr = v1 + len(EVENTOS) - 1

    # ---- constatações
    rr += 2
    band(ws, rr, "Constatações do organismo", "O", color=AMBER)
    rr += 1
    sub(ws, rr, [(c, None, t) for c, t in zip("BCDEFGHIJKLMNO", ["#", "Nº", "Auditoria", "Data", "Tipo", "Requisito", "Constatação", "Plano enviado", "Ação concluída",
                                                                  "Evidência aceita", "Prazo do plano", "Prazo para fechar", "Situação", "Conferência"])])
    c1 = rr + 1
    for k, x in enumerate(ex_["cts"]):
        rr += 1
        num(ws, f"B{rr}", k + 1)
        put(ws, f"C{rr}", x["num"], f=font(10, True, c=MUTED), h="center")
        put(ws, f"D{rr}", x["aud"], h="center", f=font(9))
        put(ws, f"E{rr}", x["data"], h="center", fmt=DATE)
        put(ws, f"F{rr}", x["tipo"], h="center", f=font(9))
        put(ws, f"G{rr}", x["req"], h="center")
        put(ws, f"H{rr}", x["desc"], f=font(9))
        for col, key in (("I", "plano"), ("J", "concl"), ("K", "evid")):
            put(ws, f"{col}{rr}", x[key], h="center", fmt=DATE)
        ct_formulas(ws, rr, pars["Data da leitura do ciclo"])
        ws.row_dimensions[rr].height = alt([(x["desc"], 46), (x["tipo"], 16)], minimo=21.75)
    c2 = rr
    cf_texto(ws, f"N{c1}:N{c2}", f"N{c1}", CT_CF)
    cf_texto(ws, f"F{c1}:F{c2}", f"F{c1}", [(MAIOR, RED), (MENOR, S2_T)])
    cf_warn(ws, f"O{c1}:O{c2}", f"O{c1}")

    # ---- resumo
    rr += 2
    band(ws, rr, "Resumo automático", "O")
    F_ = f"F{v1}:F{v1 + len(EVENTOS) - 1}"
    for k, (text, formula) in enumerate([
        ("Prontidão", f'=TEXT(D{res_row - 3},"0%")&" atendido. "&D{res_row}'),
        ("Ciclo", "=" + '&", "&'.join(f'COUNTIF({F_},"{x}")&" {x.lower()}"' for x in (REALIZADA, NOPRAZO, VENCE, ATRASADA, PREVISTA))),
        ("Constatações", "=" + '&", "&'.join(f'COUNTIF(N{c1}:N{c2},"{x}")&" {x.lower()}"' for x in (FECHADA, AGUARDA, PLANOATR, VENCIDA, CONSIDERAR))),
        ("Linhas a completar", f'=SUMPRODUCT((L{p1}:L{p2}<>"OK")*1)+SUMPRODUCT((O{c1}:O{c2}<>"OK")*1)'),
    ], 1):
        label(ws, f"B{rr+k}", text, merge=f"B{rr+k}:D{rr+k}", h="right")
        calc(ws, f"E{rr+k}", formula, merge=f"E{rr+k}:O{rr+k}", h="left")
        ws.row_dimensions[rr + k].height = 21.75
    setup(ws, MUTED, f"B1:O{rr + 4}")
    return dict(p=(p1, p2), res=res_row, v=v1, c=(c1, c2))


POS = {}
POS["ex1"] = exemplo(wb.create_sheet("Exemplo 1 - Pizzaria"), EX1)
POS["ex2"] = exemplo(wb.create_sheet("Exemplo 2 - Indústria"), EX2)

wb.active = 1
wb.save(OUT)
print("ok", OUT, "| avisos: Prontidão E%d, Ciclo E%d, Constatações E%d, Painel C%d | painel %s | exemplos %s" % (PAV, VAV, CAV, NAV, IR, POS))
