# -*- coding: utf-8 -*-
"""Gera ISO-9001-2026-modelo.xlsx no mesmo padrão visual das planilhas anteriores da série."""
import math
import os
import sys

_here = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, _here)
# funções de formatação compartilhadas com o gerador do PDCA (bloco inicial do arquivo)
_src = open(os.path.join(_here, "build_pdca.py"), encoding="utf-8").read()
exec(_src[: _src.index("STATUS = [")])

from openpyxl.chart import BarChart, Series  # noqa: E402
from openpyxl.chart.label import DataLabelList  # noqa: E402
from ed26_data import (ATENDE, ATRAS, AVISO, CHECK, CONCL, EX1, EX2, FIM_TRANSICAO, IMPACTOS, MUDANCAS, NA, NAOATENDE, NOPRAZO, P_ALTA, P_ATRAS, P_BAIXA,  # noqa: E402
                       P_FORA, P_MEDIA, P_PREV, P_REAL, P_SEM, P_VENCE, PARCIAL, PASSOS, PRIORIDADES, PUBLICACAO, SEMACAO, SITUACOES, SO_2026, TIPOS)

BLUE, TEAL, AMBER, PURPLE, REDC = "2B5C8A", "1E7B73", "A96A12", "7A4A9A", "B0413E"
S2_T = "F8EBCB"
YESNO = ["Sim", "Parcial", "Não"]
YESNO_CF = [("Sim", GREEN), ("Parcial", YELLOW), ("Não", RED)]
NLIVRE = 6
M1 = 8
M2 = M1 + len(MUDANCAS) + NLIVRE - 1
S1, S2 = 12, 12 + len(PASSOS) - 1
CERT15, SEMCERT = "Certificado pela edição de 2015", "Sem certificado"
DEPOIS = "A auditoria está depois do fim da transição"
W = {"A": 2, "B": 5, "C": 8, "D": 10, "E": 28, "F": 14, "G": 9, "H": 40, "I": 16, "J": 30, "K": 34, "L": 18, "M": 12, "N": 12, "O": 8, "P": 12, "Q": 12, "R": 26, "S": 2}
SIT_CF = [(ATENDE, GREEN), (PARCIAL, S2_T), (NAOATENDE, RED), (NA, GRAY)]
PRI_CF = [(P_ALTA, RED), (P_MEDIA, S2_T), (P_BAIXA, GRAY), (P_SEM, GREEN)]
AC_CF = [(CONCL, GREEN), (NOPRAZO, S2_T), (ATRAS, RED)]
PL_CF = [(P_REAL, GREEN), (P_FORA, RED), (P_ATRAS, RED), (P_VENCE, S2_T), (P_PREV, GRAY)]


def note(ws, ref, text):
    c = Comment(text, "Modelo ISO 9001:2026")
    c.width, c.height = 300, 120
    ws[ref].comment = c


def cf_warn(ws, rng_, first, ok_values=("OK",)):
    cond = ",".join(f'{first}="{v}"' for v in ok_values)
    ws.conditional_formatting.add(rng_, FormulaRule(formula=[f"OR({cond})"], fill=PatternFill("solid", bgColor=GREEN, fgColor=GREEN)))
    ws.conditional_formatting.add(rng_, FormulaRule(formula=[f"AND(LEN({first})>0,NOT(OR({cond})))"],
                                  fill=PatternFill("solid", bgColor=YELLOW, fgColor=YELLOW)))


def cf_texto(ws, rng_, first, pares, resto=None):
    for t, cor in pares:
        ws.conditional_formatting.add(rng_, FormulaRule(formula=[f'{first}="{t}"'], stopIfTrue=True, fill=PatternFill("solid", bgColor=cor, fgColor=cor)))
    if resto:
        ws.conditional_formatting.add(rng_, FormulaRule(formula=[f"LEN({first})>0"], fill=PatternFill("solid", bgColor=resto, fgColor=resto)))


def alt(pares, minimo=21.75, linha=12):
    n = max(max(1, math.ceil(len(str(t or "")) / max(c, 1))) for t, c in pares)
    return max(minimo, n * linha + 7)


# ------------------------------------------------------------------ fórmulas do diagnóstico (mesmas colunas na aba e nos exemplos)
def diag_formulas(ws, r, ref):
    imp = f'IF(G{r}="Alto",{IMPACTOS["Alto"]},IF(G{r}="Médio",{IMPACTOS["Médio"]},IF(G{r}="Baixo",{IMPACTOS["Baixo"]},0)))'
    lac = f'IF(I{r}="{NAOATENDE}",2,IF(I{r}="{PARCIAL}",1,0))'
    vazio = f'AND(C{r}="",E{r}="")'
    calc(ws, f"O{r}", f'=IF(OR(I{r}="",G{r}=""),"",{imp}*{lac})')
    calc(ws, f"P{r}", f'=IF(O{r}="","",IF(O{r}>=4,"{P_ALTA}",IF(O{r}>=2,"{P_MEDIA}",IF(O{r}>=1,"{P_BAIXA}","{P_SEM}"))))', sz=9)
    calc(ws, f"Q{r}", f'=IF({vazio},"",IF(K{r}="","{SEMACAO}",IF(N{r}<>"","{CONCL}",IF(AND(M{r}<>"",N({ref})>M{r}),"{ATRAS}","{NOPRAZO}"))))', sz=9)
    calc(ws, f"R{r}", f'=IF({vazio},"",IF(I{r}="","Falta a situação",IF(AND(OR(I{r}="{PARCIAL}",I{r}="{NAOATENDE}"),K{r}=""),"Lacuna sem ação",'
                      f'IF(AND(K{r}<>"",M{r}=""),"Falta o prazo",IF(AND(K{r}<>"",L{r}=""),"Falta o responsável",IF(AND(I{r}="{ATENDE}",J{r}=""),"Falta a evidência","OK"))))))',
         b=False, sz=9)


def diag_cf(ws, r1, r2):
    cf_texto(ws, f"I{r1}:I{r2}", f"I{r1}", SIT_CF)
    cf_texto(ws, f"P{r1}:P{r2}", f"P{r1}", PRI_CF)
    cf_texto(ws, f"Q{r1}:Q{r2}", f"Q{r1}", AC_CF)
    cf_warn(ws, f"R{r1}:R{r2}", f"R{r1}")


def mudanca_fixa(ws, r, k, m):
    num(ws, f"B{r}", k + 1)
    put(ws, f"C{r}", m[0], f=font(10, True, c=MUTED), bg=GRAY, h="center")
    put(ws, f"D{r}", m[1], f=font(9), bg=GRAY, h="center")
    put(ws, f"E{r}", m[2], f=font(10, True), bg=GRAY)
    put(ws, f"F{r}", m[3], f=font(9), bg=GRAY, h="center")
    put(ws, f"G{r}", m[4], f=font(9), bg=GRAY, h="center")
    put(ws, f"H{r}", m[6], f=font(9), bg=GRAY)


# ------------------------------------------------------------------ fórmulas do plano
PL_ABA = dict(etapa=("C", None), dias="D", lim=("E", None), feito="F", sit="G")
# nos exemplos, as colunas são as do diagnóstico: a etapa ocupa C:E e o limite G:H
PL_EX = dict(etapa=("C", "E"), dias="F", lim=("G", "H"), feito="I", sit="J")


def plano_passos(ws, s1, audit, ref, fixos=None, c=PL_ABA):
    (ce, ce2), cd, (cl, cl2), cf, cs = c["etapa"], c["dias"], c["lim"], c["feito"], c["sit"]
    for k, (p, dias) in enumerate(PASSOS):
        r = s1 + k
        num(ws, f"B{r}", k + 1)
        put(ws, f"{ce}{r}", p, bg=GRAY, f=font(10, True), merge=f"{ce}{r}:{ce2}{r}" if ce2 else None)
        put(ws, f"{cd}{r}", dias, bg=GRAY, h="center")
        L, F = f"{cl}{r}", f"{cf}{r}"
        calc(ws, L, f'=IF({audit}="","",{audit}-{cd}{r})', b=False, fmt=DATE, merge=f"{L}:{cl2}{r}" if cl2 else None)
        if fixos is None:
            inp(ws, F, h="center", fmt=DATE)
        else:
            put(ws, F, fixos[k], h="center", fmt=DATE)
        calc(ws, f"{cs}{r}", f'=IF({L}="","",IF({F}<>"",IF({F}<={L},"{P_REAL}","{P_FORA}"),IF(N({ref})=0,"",IF({ref}>{L},"{P_ATRAS}",'
                             f'IF({L}-{ref}<={AVISO},"{P_VENCE}","{P_PREV}")))))', sz=9)
        ws.row_dimensions[r].height = 21.75
    cf_texto(ws, f"{cs}{s1}:{cs}{s1 + len(PASSOS) - 1}", f"{cs}{s1}", PL_CF)


def f_alerta(sit, audit, fim):
    return f'=IF(OR({sit}="",{audit}=""),"",IF(AND({sit}="{CERT15}",{audit}>{fim}),"{DEPOIS}","OK"))'


def chart_style(ch):
    ch.style = 2
    ch.title = None
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ch.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill="D3DBE0", w=9525))
    ch.x_axis.graphicalProperties = GraphicalProperties(ln=LineProperties(solidFill="D3DBE0", w=9525))
    ch.y_axis.graphicalProperties = GraphicalProperties(ln=LineProperties(noFill=True))


def colunas(ws, anchor, r1, r2, c_cat, c_val, width=16, height=7.5):
    ch = BarChart()
    ch.type = "col"
    ch.gapWidth = 70
    ch.height, ch.width = height, width
    chart_style(ch)
    ch.legend = None
    s = Series(Reference(ws, min_col=c_val, min_row=r1, max_row=r2), title="Mudanças")
    s.graphicalProperties.solidFill = BLUE
    s.graphicalProperties.line.noFill = True
    s.dLbls = DataLabelList()
    s.dLbls.showVal = True
    s.dLbls.showSerName = s.dLbls.showCatName = s.dLbls.showLegendKey = False
    ch.append(s)
    ch.set_categories(Reference(ws, min_col=c_cat, min_row=r1, max_row=r2))
    ch.y_axis.scaling.min = 0
    ch.y_axis.majorUnit = 1
    ws.add_chart(ch, anchor)


def resumo_aba(ws, s, itens, last, merge_lab="D"):
    band(ws, s, "Resumo automático", last)
    col = chr(ord(merge_lab) + 1)
    for k, (text, formula) in enumerate(itens, 1):
        aviso = text == "Aviso"
        label(ws, f"B{s+k}", text, merge=f"B{s+k}:{merge_lab}{s+k}", h="right")
        calc(ws, f"{col}{s+k}", formula, merge=f"{col}{s+k}:{last}{s+k}", sz=9 if aviso else 10, b=not aviso, h="left")
        ws.row_dimensions[s + k].height = 21.75
    av = s + len(itens)
    cf_warn(ws, f"{col}{av}:{last}{av}", f"{col}{av}")
    return av


def conta_por(rng, valores):
    return "=" + '&", "&'.join(f'COUNTIF({rng},"{x}")&" {x.lower()}"' for x in valores)


# ================================================================== pasta de trabalho
wb = Workbook()
wb.properties.title = "ISO 9001:2026: o que muda — Modelo"
wb.properties.creator = "Modelo ISO 9001:2026"
wb.properties.language = "pt-BR"

# ------------------------------------------------------------------ Instruções
ws = wb.active
ws.title = "Instruções"
widths(ws, {"A": 2, "B": 26, "C": 90, "D": 2})
title(ws, "ISO 9001:2026: o que muda — Como usar esta planilha", "Modelo para diagnosticar as mudanças da edição de 2026 e planejar a transição.", "C")
r = 4


def section(text):
    global r
    band(ws, r, text, "C", sz=11)
    r += 1


def line(k, v, kbg=GRAY, vbg=None, kf=None, kh="left", height=None):
    global r
    put(ws, f"B{r}", k, f=kf or font(10, True), bg=kbg, h=kh)
    put(ws, f"C{r}", v, bg=vbg)
    ws.row_dimensions[r].height = height or (19.5 if len(v) <= 96 else (31.5 if len(v) <= 192 else 45.75))
    r += 1


section("Legenda: onde preencher")
line("Amarelo-claro", "Células de entrada. É aqui que você digita.", vbg=INPUT)
line("Cinza", "Células calculadas ou fixas (as nove mudanças, pontos, prioridade, limites, situações, conferências). Não altere.", vbg=GRAY)
line("Listas suspensas", "Situação, tipo, impacto, situação do certificado e checklist aceitam apenas as opções da lista.")
line("Comentários", "Células com triângulo vermelho no canto trazem uma dica de preenchimento. Passe o mouse sobre elas.")
r += 1
section("As regras de cálculo")
for k, text in [
    ("Pontos", "O impacto (alto 3, médio 2, baixo 1) vezes a lacuna (não atende 2, atende em parte 1, atende ou não se aplica 0). Vão de 0 a 6."),
    ("Prioridade", f"{P_ALTA}: 4 pontos ou mais. {P_MEDIA}: 2 ou 3. {P_BAIXA}: 1. {P_SEM}: 0."),
    ("Situação da ação", f"{CONCL} com data de conclusão. {ATRAS} se o prazo passou na data da leitura. {NOPRAZO} nos outros casos."),
    ("Limite de cada etapa", "A data da auditoria menos os dias da coluna. Os dias são uma convenção deste material: 180 para ler a norma, 15 para a análise crítica."),
    ("Situação da etapa", f"{P_REAL} ou {P_FORA}, se feita. Se não: {P_ATRAS} depois do limite, {P_VENCE} a {AVISO} dias ou menos, {P_PREV} antes disso."),
    ("Alerta", f"Para quem tem certificado pela edição de 2015, a auditoria precisa ser até o fim da transição."),
]:
    line(k, text)
r += 1
section("Ordem recomendada de preenchimento")
for k, text in enumerate([
    "Aba Mudanças: a data da leitura e a situação de cada mudança, com a evidência.",
    "Aba Mudanças: para cada lacuna, a ação, o responsável e o prazo. Quando concluir, a data.",
    "Aba Mudanças: se encontrar outra mudança no texto da norma, use as linhas livres.",
    "Aba Plano: a auditoria, a data, a leitura, o fim da transição e a situação do certificado.",
    "Aba Plano: as datas em que cada etapa foi feita.",
    "Aba Painel: leia os avisos. Depois, valide na aba Checklist.",
], 1):
    line(f"Passo {k}", text)
r += 1
section("Abas da planilha")
for k, text in [
    ("Mudanças", f"As {len(MUDANCAS)} mudanças da edição de 2026, com a pergunta de cada uma, e {NLIVRE} linhas livres."),
    ("Plano", f"As {len(PASSOS)} etapas da transição, com o limite calculado pela data da auditoria."),
    ("Painel", "Os números das duas abas, o gráfico das prioridades e o aviso."),
    ("Checklist", "Doze verificações da transição, com percentual de conclusão."),
    ("Exemplo 1 - Pizzaria", "A preparação para a certificação inicial, já pela edição de 2026."),
    ("Exemplo 2 - Indústria", "A transição na 1ª manutenção, depois da certificação pela edição de 2015."),
]:
    line(k, text)
r += 1
section("Premissas do modelo")
for k, text in [
    ("Datas", f"A ISO 9001:2026 foi publicada em {PUBLICACAO.strftime('%d/%m/%Y')}. Pelas regras de transição da Global ACI, certificados novos só pela edição de 2026 desde "
              f"{SO_2026.strftime('%d/%m/%Y')}, e os de 2015 valem até {FIM_TRANSICAO.strftime('%d/%m/%Y')}."),
    ("Mudanças", "Os resumos são deste material, a partir do anúncio da ISO e das análises dos organismos. Confira no texto da norma."),
    ("Exemplos", "Os dados das abas de exemplo são ilustrativos, criados para este material."),
]:
    line(k, text)
setup(ws, INK, f"B1:C{r}", landscape=False, fit_height=True)

# ------------------------------------------------------------------ Mudanças
ws = wb.create_sheet("Mudanças")
widths(ws, W)
title(ws, "Diagnóstico das mudanças", "Uma mudança por linha: a situação, a evidência e, se houver lacuna, a ação.", "R")
label(ws, "B4", "Data da leitura", merge="B4:C4")
inp(ws, "D4", h="center", fmt=DATE)
dv_date(ws, "D4")
put(ws, "E4", "As ações com prazo vencido nesta data aparecem atrasadas.", f=font(9, i=True, c=MUTED), bg=GRAY, merge="E4:R4")
for col, text in zip("BCDEFGHIJKLMNOPQR", ["#", "Código", "Onde", "Mudança", "Tipo", "Impacto", "Pergunta do diagnóstico", "Situação", "Evidência", "Ação", "Responsável",
                                           "Prazo", "Concluída em", "Pontos", "Prioridade", "Ação", "Conferência"]):
    head(ws, f"{col}6", text)
ws.row_dimensions[6].height = 33
put(ws, "B7", "As nove mudanças da edição de 2026, e linhas livres para outras que você encontrar no texto.", f=font(9, i=True, c=MUTED), bg=GRAY, merge="B7:R7")
for k in range(len(MUDANCAS) + NLIVRE):
    rr = M1 + k
    if k < len(MUDANCAS):
        mudanca_fixa(ws, rr, k, MUDANCAS[k])
    else:
        num(ws, f"B{rr}", k + 1)
        for col in "CDFG":
            inp(ws, f"{col}{rr}", h="center")
        for col in "EH":
            inp(ws, f"{col}{rr}")
    inp(ws, f"I{rr}", h="center")
    for col in "JKL":
        inp(ws, f"{col}{rr}")
    for col in "MN":
        inp(ws, f"{col}{rr}", h="center", fmt=DATE)
    diag_formulas(ws, rr, "$D$4")
    ws.row_dimensions[rr].height = 42 if k < len(MUDANCAS) else 30
dv_list(ws, f"I{M1}:I{M2}", SITUACOES, "Atende, atende em parte, não atende ou não se aplica")
dv_list(ws, f"F{M1 + len(MUDANCAS)}:F{M2}", TIPOS, "Novo, reforçado, reestruturado, incorporado ou esclarecimento")
dv_list(ws, f"G{M1 + len(MUDANCAS)}:G{M2}", list(IMPACTOS), "Alto, médio ou baixo")
dv_date(ws, f"M{M1}:N{M2}")
diag_cf(ws, M1, M2)
note(ws, "I6", "Atende: a prática e a evidência existem. Atende em parte: existe, mas falta algo. Não atende: não existe. Não se aplica: com justificativa.")
note(ws, "G6", "Convenção deste material: o quanto a mudança costuma exigir de trabalho e o quanto o auditor tende a olhar para ela.")
s = M2 + 2
I_ = f"I{M1}:I{M2}"
P_ = f"P{M1}:P{M2}"
MAV = resumo_aba(ws, s, [
    ("Situação", conta_por(I_, SITUACOES)),
    ("Prioridade", conta_por(P_, PRIORIDADES) + f'&". Total: "&SUM(O{M1}:O{M2})&" pontos"'),
    ("Aviso", f'=IF(COUNTA({I_})=0,"Diagnostique as mudanças",IF(COUNTIF(R{M1}:R{M2},"Lacuna sem ação")>0,"Há lacuna sem ação",'
              f'IF(COUNTIF(Q{M1}:Q{M2},"{ATRAS}")>0,"Há ação atrasada",IF(SUMPRODUCT((R{M1}:R{M2}<>"")*(R{M1}:R{M2}<>"OK"))>0,"Há linha a completar: veja a coluna Conferência",'
              f'IF(COUNTIF({P_},"{P_ALTA}")>0,"Há prioridade alta: leve à direção","OK")))))'),
], "R")
ws.freeze_panes = f"F{M1}"
setup(ws, BLUE, f"B1:R{MAV}")

# ------------------------------------------------------------------ Plano
ws = wb.create_sheet("Plano")
widths(ws, {"A": 2, "B": 5, "C": 40, "D": 12, "E": 14, "F": 14, "G": 20, "H": 36, "I": 14, "J": 2})
title(ws, "Plano de transição", "As etapas até a auditoria, com o limite de cada uma.", "I")
for rr, text, fmt, dica in [(4, "Auditoria", None, "Qual auditoria: manutenção, recertificação, auditoria separada, ou a fase 1 da certificação inicial."),
                            (5, "Data da auditoria", DATE, "A data combinada, ou prevista, com o organismo."),
                            (6, "Data da leitura", DATE, "As etapas são comparadas com esta data."),
                            (7, "Fim da transição", DATE, "Para certificados pela edição de 2015. Confirme com o seu organismo."),
                            (8, "Situação do certificado", None, "Certificado pela edição de 2015, ou sem certificado.")]:
    label(ws, f"B{rr}", text, merge=f"B{rr}:C{rr}")
    inp(ws, f"D{rr}", h="center" if fmt else "left", fmt=fmt, merge=f"D{rr}:E{rr}")
    put(ws, f"F{rr}", dica, f=font(9, i=True, c=MUTED), bg=GRAY, merge=f"F{rr}:I{rr}")
ws["D7"] = FIM_TRANSICAO
dv_date(ws, "D5:D7")
dv_list(ws, "D8", [CERT15, SEMCERT], "Certificado pela edição de 2015, ou sem certificado")
label(ws, "B9", "Alerta de prazo", merge="B9:C9")
calc(ws, "D9", f_alerta("D8", "D5", "D7"), merge="D9:I9", h="left", sz=9)
cf_warn(ws, "D9:I9", "$D9")
for col, text in zip("BCDEFGHI", ["#", "Etapa", "Dias antes", "Limite", "Feito em", "Situação", "Observação", "Limite pendente"]):
    head(ws, f"{col}11", text)
ws.row_dimensions[11].height = 21.75
plano_passos(ws, S1, "$D$5", "$D$6")
for k in range(len(PASSOS)):
    inp(ws, f"H{S1 + k}")
    # o limite das etapas ainda não feitas; o painel tira daqui o próximo passo
    calc(ws, f"I{S1 + k}", f'=IF(AND(E{S1 + k}<>"",F{S1 + k}=""),E{S1 + k},"")', b=False, fmt=DATE)
dv_date(ws, f"F{S1}:F{S2}")
note(ws, "D11", "Convenção deste material: quantos dias antes da auditoria cada etapa deveria estar feita.")
s = S2 + 2
G_ = f"G{S1}:G{S2}"
PAV = resumo_aba(ws, s, [
    ("Etapas", conta_por(G_, (P_REAL, P_FORA, P_ATRAS, P_VENCE, P_PREV))),
    ("Aviso", f'=IF(D5="","Informe a data da auditoria",IF(D9="{DEPOIS}","{DEPOIS}",IF(D6="","Informe a data da leitura",IF(COUNTIF({G_},"{P_ATRAS}")>0,"Há etapa atrasada",'
              f'IF(COUNTIF({G_},"{P_VENCE}")>0,"Há etapa que vence em {AVISO} dias","OK")))))'),
], "I")
setup(ws, TEAL, f"B1:I{PAV}", landscape=False, fit_height=True)

# ------------------------------------------------------------------ Painel
ws = wb.create_sheet("Painel")
widths(ws, {"A": 2, "B": 38, "C": 16, "D": 14, "E": 14, "F": 34, "G": 2})
title(ws, "Painel da transição", "Nada a preencher: tudo vem das outras abas.", "F")
label(ws, "B3", "Auditoria")
calc(ws, "C3", '=IF(Plano!D4="","",Plano!D4)', merge="C3:F3", h="left")
for col, text, m in [("B", "Indicador", None), ("C", "Resultado", None), ("D", "Como se lê", "D5:F5")]:
    put(ws, f"{col}5", text, f=font(10, True, c=WHITE), bg=INK, h="center", merge=m)
ws.row_dimensions[5].height = 21.75
Mi, Mo, Mp, Mq, Mr = (f"'Mudanças'!{c}{M1}:{c}{M2}" for c in "IOPQR")
Pg, Pi = (f"Plano!{c}{S1}:{c}{S2}" for c in "GI")
IND = [
    ("DIAGNÓSTICO", None, None, None),
    (ATENDE, f'=COUNTIF({Mi},"{ATENDE}")', "Mudanças já atendidas.", None),
    (PARCIAL, f'=COUNTIF({Mi},"{PARCIAL}")', "Lacunas parciais.", None),
    (NAOATENDE, f'=COUNTIF({Mi},"{NAOATENDE}")', "Lacunas inteiras.", None),
    ("Prioridade alta", f'=COUNTIF({Mp},"{P_ALTA}")', "Levar à direção.", None),
    ("Prioridade média", f'=COUNTIF({Mp},"{P_MEDIA}")', "Antes da auditoria interna.", None),
    ("Pontos de lacuna", f"=SUM({Mo})", "Impacto vezes lacuna, somados.", None),
    ("Ações", f'=COUNTIF({Mq},"{CONCL}")+COUNTIF({Mq},"{NOPRAZO}")+COUNTIF({Mq},"{ATRAS}")', "Ações registradas.", None),
    ("Ações concluídas", f'=COUNTIF({Mq},"{CONCL}")', "Com data de conclusão.", None),
    ("Ações atrasadas", f'=COUNTIF({Mq},"{ATRAS}")', "Prazo vencido na data da leitura.", None),
    ("Linhas a completar", f'=SUMPRODUCT(({Mr}<>"")*({Mr}<>"OK"))', "Conferências diferentes de OK.", None),
    ("PLANO", None, None, None),
    ("Data da auditoria", '=IF(Plano!D5="","",Plano!D5)', "Da aba Plano.", DATE),
    ("Alerta de prazo", '=IF(Plano!D9="","",Plano!D9)', "A auditoria cabe na transição?", None),
    ("Etapas atrasadas", f'=COUNTIF({Pg},"{P_ATRAS}")', "Passaram do limite.", None),
    (f"Etapas que vencem em {AVISO} dias", f'=COUNTIF({Pg},"{P_VENCE}")', "Confirmar a data e o responsável.", None),
    ("Próximo limite", f'=IF(MIN({Pi})=0,"",MIN({Pi}))', "O limite mais próximo de uma etapa ainda não feita.", DATE),
    ("Próxima etapa", f'=IF(C{{prox}}="","",IFERROR(INDEX(Plano!C{S1}:C{S2},MATCH(C{{prox}},{Pi},0)),""))', "A etapa desse limite.", None),
]
IR = {}
rr = 6
for nome, formula, leit, fmt in IND:
    if formula is None:
        put(ws, f"B{rr}", nome, f=font(9, True, c=MUTED), bg=None, box=False, merge=f"B{rr}:F{rr}")
        ws.row_dimensions[rr].height = 21.75
        rr += 1
        continue
    IR[nome] = rr
    if "{prox}" in formula:
        formula = formula.replace("{prox}", str(IR["Próximo limite"]))
    label(ws, f"B{rr}", nome)
    calc(ws, f"C{rr}", formula, fmt=fmt, sz=9 if nome == "Alerta de prazo" else 10)
    put(ws, f"D{rr}", leit, f=font(9, c=MUTED), bg=GRAY, merge=f"D{rr}:F{rr}")
    ws.row_dimensions[rr].height = 21.75
    rr += 1
s = rr + 1
band(ws, s, "Resumo automático", "F")
label(ws, f"B{s+1}", "Aviso")
C = lambda n: f"C{IR[n]}"  # noqa: E731
calc(ws, f"C{s+1}", f'=IF(COUNTA({Mi})=0,"Diagnostique as mudanças na aba Mudanças",IF({C("Alerta de prazo")}="{DEPOIS}","{DEPOIS}",'
     f'IF({C("Ações atrasadas")}>0,"Há ação atrasada",IF({C("Etapas atrasadas")}>0,"Há etapa do plano atrasada",IF({C("Linhas a completar")}>0,"Há linha a completar no diagnóstico",'
     f'IF({C("Prioridade alta")}>0,"Há prioridade alta: leve à direção","OK"))))))', merge=f"C{s+1}:F{s+1}", sz=9, b=False, h="left")
cf_warn(ws, f"C{s+1}:F{s+1}", f"C{s+1}")
ws.row_dimensions[s + 1].height = 21.75
NAV = s + 1
GS = s + 4
put(ws, f"B{GS - 1}", "Mudanças por prioridade", f=font(9, True, c=MUTED), bg=None, box=False)
for k, p in enumerate(PRIORIDADES):
    put(ws, f"B{GS + k}", p, f=font(9), bg=GRAY)
    calc(ws, f"C{GS + k}", f'=COUNTIF({Mp},B{GS + k})', b=False)
colunas(ws, f"D{GS - 1}", GS, GS + len(PRIORIDADES) - 1, 2, 3, width=12, height=7)
setup(ws, REDC, f"B1:F{GS + 14}", landscape=False, fit_height=True)

# ------------------------------------------------------------------ Checklist
ws = wb.create_sheet("Checklist")
widths(ws, {"A": 2, "B": 5, "C": 70, "D": 14, "E": 44, "F": 2})
title(ws, "Checklist da transição para a ISO 9001:2026", "Escolha o status de cada item na lista suspensa (coluna em amarelo-claro).", "E")
for col, text in zip("BCDE", ["#", "Item de verificação", "Status", "Observações"]):
    head(ws, f"{col}4", text)
ws.row_dimensions[4].height = 21.75
for k, text in enumerate(CHECK):
    rr = 5 + k
    num(ws, f"B{rr}", k + 1)
    put(ws, f"C{rr}", text, bg=GRAY)
    inp(ws, f"D{rr}", h="center")
    inp(ws, f"E{rr}")
    ws.row_dimensions[rr].height = 30
dv_list(ws, "D5:D16", YESNO, "Sim, Parcial ou Não")
cf_equal(ws, "D5:D16", YESNO_CF)
band(ws, 18, "Resumo automático", "E")
summary(ws, 19, "Itens atendidos (Sim)", '=COUNTIF(D5:D16,"Sim")', "C", "D")
summary(ws, 20, "Itens parciais", '=COUNTIF(D5:D16,"Parcial")', "C", "D")
summary(ws, 21, "Itens não atendidos", '=COUNTIF(D5:D16,"Não")', "C", "D")
summary(ws, 22, "Itens sem resposta", "=COUNTBLANK(D5:D16)", "C", "D")
summary(ws, 23, "Percentual atendido", "=D19/ROWS(D5:D16)", "C", "D", fmt="0%")
summary(ws, 24, "Situação", '=IF(D19=ROWS(D5:D16),"Validado",IF(D22>0,"Em andamento","Com pendências"))', "C", "D")
setup(ws, PURPLE, "B1:E24", landscape=False, fit_height=True)


# ------------------------------------------------------------------ Exemplos
def exemplo(ws, ex_, cert):
    H, pl = ex_["head"], ex_["plano"]
    widths(ws, W)
    title(ws, "ISO 9001:2026: o que muda", "Exemplo preenchido, para consulta. Use as abas Mudanças e Plano para a sua organização.", "R")
    rr = 4
    for rot, val in [("Organização", H["org"]), ("Situação", H["situacao"]), ("Auditoria", H["auditoria"]), ("Responsável", H["resp"])]:
        label(ws, f"B{rr}", rot, merge=f"B{rr}:D{rr}")
        inp(ws, f"E{rr}", val, merge=f"E{rr}:R{rr}", bg=WHITE, h="left")
        ws.row_dimensions[rr].height = 19.5
        rr += 1
    pars = {}
    for rot, val, fmt in [("Data da leitura", H["ref"], DATE), ("Data da auditoria", pl["audit"], DATE), ("Fim da transição", pl["fim"], DATE), ("Situação do certificado", cert, None)]:
        label(ws, f"B{rr}", rot, merge=f"B{rr}:D{rr}")
        inp(ws, f"E{rr}", val, bg=WHITE, h="center", fmt=fmt)
        pars[rot] = f"$E${rr}"
        rr += 1
    label(ws, f"B{rr}", "Alerta de prazo", merge=f"B{rr}:D{rr}")
    calc(ws, f"E{rr}", f_alerta(pars["Situação do certificado"], pars["Data da auditoria"], pars["Fim da transição"]), merge=f"E{rr}:H{rr}", h="left", sz=9)
    cf_warn(ws, f"E{rr}:H{rr}", f"$E{rr}")
    alerta = rr

    # ---- diagnóstico (mesmas colunas da aba Mudanças)
    rr += 2
    band(ws, rr, "Diagnóstico das mudanças", "R", color=BLUE)
    rr += 1
    for col, text in zip("BCDEFGHIJKLMNOPQR", ["#", "Código", "Onde", "Mudança", "Tipo", "Impacto", "Pergunta", "Situação", "Evidência", "Ação", "Responsável", "Prazo",
                                               "Concluída em", "Pontos", "Prioridade", "Ação", "Conferência"]):
        put(ws, f"{col}{rr}", text, f=font(10, True), bg=GRAY, h="center")
    ws.row_dimensions[rr].height = 30
    m1 = rr + 1
    for k, (m, d) in enumerate(zip(MUDANCAS, ex_["diag"])):
        rr += 1
        mudanca_fixa(ws, rr, k, m)
        put(ws, f"I{rr}", d["sit"], h="center", f=font(9))
        put(ws, f"J{rr}", d["evid"] or None, f=font(9))
        put(ws, f"K{rr}", d["acao"] or None, f=font(9))
        put(ws, f"L{rr}", d["resp"] or None, f=font(9))
        put(ws, f"M{rr}", d["prazo"], h="center", fmt=DATE)
        put(ws, f"N{rr}", d["concl"], h="center", fmt=DATE)
        diag_formulas(ws, rr, pars["Data da leitura"])
        ws.row_dimensions[rr].height = alt([(m[6], 40), (d["evid"], 30), (d["acao"], 34)], minimo=30)
    m2 = rr
    diag_cf(ws, m1, m2)

    # ---- plano
    rr += 2
    ws.row_breaks.append(Break(id=rr - 1))
    band(ws, rr, "Plano de transição", "R", color=TEAL)
    rr += 1
    for col, text, m in [("B", "#", None), ("C", "Etapa", "E"), ("F", "Dias antes", None), ("G", "Limite", "H"), ("I", "Feito em", None), ("J", "Situação", None)]:
        put(ws, f"{col}{rr}", text, f=font(10, True), bg=GRAY, h="center", merge=f"{col}{rr}:{m}{rr}" if m else None)
    ws.row_dimensions[rr].height = 30
    s1 = rr + 1
    plano_passos(ws, s1, pars["Data da auditoria"], pars["Data da leitura"], fixos=pl["feito"], c=PL_EX)
    rr = s1 + len(PASSOS) - 1

    # ---- resumo
    rr += 2
    band(ws, rr, "Resumo automático", "R")
    for k, (text, formula) in enumerate([
        ("Diagnóstico", conta_por(f"I{m1}:I{m2}", SITUACOES) + f'&". Total: "&SUM(O{m1}:O{m2})&" pontos"'),
        ("Prioridade", conta_por(f"P{m1}:P{m2}", PRIORIDADES)),
        ("Ações", conta_por(f"Q{m1}:Q{m2}", (CONCL, NOPRAZO, ATRAS))),
        ("Plano", conta_por(f"J{s1}:J{s1 + len(PASSOS) - 1}", (P_REAL, P_FORA, P_ATRAS, P_VENCE, P_PREV))),
    ], 1):
        label(ws, f"B{rr+k}", text, merge=f"B{rr+k}:D{rr+k}", h="right")
        calc(ws, f"E{rr+k}", formula, merge=f"E{rr+k}:R{rr+k}", h="left")
        ws.row_dimensions[rr + k].height = 21.75
    setup(ws, MUTED, f"B1:R{rr + 4}")
    return dict(m=(m1, m2), s=s1, alerta=alerta)


POS = {}
POS["ex1"] = exemplo(wb.create_sheet("Exemplo 1 - Pizzaria"), EX1, SEMCERT)
POS["ex2"] = exemplo(wb.create_sheet("Exemplo 2 - Indústria"), EX2, CERT15)

wb.active = 1
wb.save(OUT)
print("ok", OUT, "| avisos: Mudanças E%d, Plano E%d, Painel C%d | painel %s | exemplos %s" % (MAV, PAV, NAV, IR, POS))
