"""Lê os resultados dos testes de Processos-modelo.xlsx depois do recálculo.  Uso: proc_test_read.py <pasta com t1..t4.xlsx>"""
import openpyxl, sys
D = sys.argv[1]
for nome in ('t1', 't2', 't3', 't4'):
    w = openpyxl.load_workbook(f'{D}/{nome}.xlsx', data_only=True)
    P, I, T, E, C = w['Processos'], w['Interações'], w['Tartaruga'], w['Elementos'], w['Checklist']
    print('=====', nome)
    err = [(ws.title, c.coordinate, c.value) for ws in w.worksheets[:6] for r in ws.iter_rows() for c in r
           if isinstance(c.value, str) and ((c.value.startswith('#') and len(c.value) > 1) or 'Err:' in c.value)]
    print('erros', err)
    print('Processos  conferência', [P[f'K{r}'].value for r in range(10, 22) if P[f'K{r}'].value])
    print('           resumo', [P[f'E{r}'].value for r in range(24, 30)])
    print('Interações', [(I[f'B{r}'].value, I[f'P{r}'].value, I[f'Q{r}'].value, I[f'R{r}'].value) for r in range(6, 18) if I[f'C{r}'].value])
    print('           cabeçalho', I['D5'].value, '|', I['L5'].value, '|', I['O5'].value)
    print('           resumo', [I[f'D{r}'].value for r in range(20, 25)])
    print('Tartaruga ', [T[f'D{r}'].value for r in (4, 5, 6, 7)])
    print('           partes', [T[f'H{r}'].value for r in range(31, 39)])
    print('Elementos ', [(E[f'B{r}'].value, E[f'L{r}'].value, None if not isinstance(E[f'M{r}'].value, (int, float)) else round(100 * E[f'M{r}'].value), E[f'N{r}'].value)
                         for r in range(5, 17) if E[f'N{r}'].value])
    print('           por elemento', [None if not isinstance(E.cell(row=20, column=c).value, (int, float)) else round(100 * E.cell(row=20, column=c).value) for c in range(4, 12)])
    print('           resumo', [E[f'D{r}'].value for r in range(23, 28)])
    print('Checklist ', [C[f'D{r}'].value for r in range(19, 25)])
