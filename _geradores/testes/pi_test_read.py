"""Lê os resultados dos testes de Partes-Interessadas-modelo.xlsx depois do recálculo.  Uso: pi_test_read.py <pasta com t1..t3.xlsx>"""
import openpyxl, sys, os
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
from pi_data import EX1, EX2, ATENDE, PARTE, NAOAT, adotados, atendimento, conta, estrategia, pertinente
D = sys.argv[1]
for nome, ex in (('t1', EX1), ('t2', EX2), ('t3', None)):
    w = openpyxl.load_workbook(f'{D}/{nome}.xlsx', data_only=True)
    P, R, M, Pa, K = w['Partes'], w['Requisitos'], w['Matriz'], w['Painel'], w['Checklist']
    print('=====', nome)
    err = [(ws.title, c.coordinate, c.value) for ws in w.worksheets[:6] for r in ws.iter_rows() for c in r
           if isinstance(c.value, str) and ((c.value.startswith('#') and len(c.value) > 1) or 'Err:' in c.value)]
    print('erros', err)
    got = [(P[f'I{r}'].value, P[f'J{r}'].value, P[f'M{r}'].value) for r in range(11, 23) if P[f'C{r}'].value]
    print('Partes    ', got if ex is None else set(g[2] for g in got), '| resumo', [P[f'E{r}'].value for r in range(25, 29)])
    print('Requisitos', [R[f'M{r}'].value for r in range(7, 47) if R[f'M{r}'].value] if ex is None else set(R[f'M{r}'].value for r in range(7, 47) if R[f'M{r}'].value),
          '| resumo', [R[f'E{r}'].value for r in range(49, 55)])
    print('Matriz    ', [[M.cell(row=rr, column=c).value for c in range(4, 9)] for rr in range(6, 11)])
    print('           ', [(M[f'D{r}'].value, M[f'E{r}'].value) for r in range(15, 19)], '|', M['D26'].value)
    print('Painel    ', [Pa[f'D{r}'].value for r in range(7, 15)])
    linhas = [(Pa[f'C{r}'].value, Pa[f'F{r}'].value, Pa[f'G{r}'].value, Pa[f'H{r}'].value, Pa[f'I{r}'].value, Pa[f'J{r}'].value, Pa[f'K{r}'].value) for r in range(18, 30) if Pa[f'C{r}'].value]
    print('           ', [l[1:] for l in linhas], '|', [Pa[f'D{r}'].value for r in (32, 33)])
    if ex:
        ok1 = [(P[f'I{r}'].value, P[f'J{r}'].value) for r in range(11, 11 + len(ex['partes']))] == [(estrategia(p['inf'], p['int']), 'Sim' if pertinente(p) else 'Não') for p in ex['partes']]
        esp = [(p['nome'], len(adotados(ex, p['nome'])), conta(ex, ATENDE, p['nome']), conta(ex, PARTE, p['nome']), conta(ex, NAOAT, p['nome'])) for p in ex['partes']]
        ok2 = [l[:5] for l in linhas] == esp
        ok3 = abs(Pa['D11'].value - atendimento(ex)) < 1e-9
        print('            igual ao exemplo: estratégias', 'OK' if ok1 else 'DIFERE', '| por parte', 'OK' if ok2 else 'DIFERE', '| atendimento', 'OK' if ok3 else 'DIFERE')
    print('Checklist ', [K[f'D{r}'].value for r in range(19, 25)])
