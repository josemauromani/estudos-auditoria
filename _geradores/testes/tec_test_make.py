"""Confere Tecnica-Auditoria-modelo.xlsx e preenche cópias com dados de teste.  Uso: tec_test_make.py <modelo.xlsx> <pasta de saída>"""
import openpyxl, sys, os
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
from tec_data import (ALTO, BAIXO, CONFORME, EX1, EX2, MEDIO, NC, OM, am_conf, am_sit, aud_conf, aud_conta, aud_nivel, aud_pct, ev_conf, forca, fontes, passo,
                      tamanho)
SRC, OUTDIR = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(SRC, data_only=True)
wf0 = openpyxl.load_workbook(SRC)
for ws in wb.worksheets:
    f = wf0[ws.title]
    print('==', ws.title, ws.dimensions, 'dv', len(f.data_validations.dataValidation), 'cf', sum(len(v) for v in f.conditional_formatting._cf_rules.values()),
          'charts', len(f._charts), 'comments', sum(1 for r in f.iter_rows() for c in r if c.comment))
err = [(ws.title, c.coordinate, c.value) for ws in wb.worksheets for r in ws.iter_rows() for c in r
       if isinstance(c.value, str) and ((c.value.startswith('#') and len(c.value) > 1) or 'Err:' in c.value)]
print('erros', err)
print('vazio:', wb['Amostras']['E46'].value, '|', wb['Evidências']['E71'].value, '|', wb['Auditores']['F30'].value, '|', wb['Painel']['C28'].value)


def v(c):
    """O openpyxl lê como None o texto vazio que a fórmula devolve."""
    return '' if c.value is None else c.value


def e(x):
    return '' if x is None else x


def r4(x):
    return round(x, 4) if isinstance(x, float) else x


for aba, ex, pos in (('Exemplo 1 - Pizzaria', EX1, dict(a=14, e=22, k=36, rs=48)), ('Exemplo 2 - Indústria', EX2, dict(a=14, e=22, k=34, rs=46))):
    s = wb[aba]
    got_a = [(v(s[f'I{r}']), v(s[f'J{r}']), v(s[f'N{r}']), v(s[f'O{r}'])) for r in range(pos['a'], pos['a'] + len(ex['ams']))]
    esp_a = [(e(tamanho(a)), e(passo(a)), am_sit(a), am_conf(a)) for a in ex['ams']]
    got_e = [(v(s[f'M{r}']), v(s[f'N{r}'])) for r in range(pos['e'], pos['e'] + len(ex['evs']))]
    esp_e = [(forca(x), ev_conf(x)) for x in ex['evs']]
    got_u = [tuple(r4(v(s[f'{chr(70 + k)}{pos["rs"] + j}'])) for j in range(5)) for k in range(len(ex['auds']))]
    esp_u = [(aud_conta(a['notas'])[0], aud_conta(a['notas'])[1], r4(aud_pct(a['notas'])), aud_nivel(a['notas']), aud_conf(a['nome'], a['notas'])) for a in ex['auds']]
    print(aba, '| amostras', 'OK' if got_a == esp_a else 'DIFERE %s' % got_a, '| evidências', 'OK' if got_e == esp_e else 'DIFERE %s' % got_e,
          '| auditores', 'OK' if got_u == esp_u else 'DIFERE %s\n   esperado %s' % (got_u, esp_u))


def preencher(nome, ams=(), evs=(), auds=(), check=None):
    w = openpyxl.load_workbook(SRC)
    A, E, U, K = w['Amostras'], w['Evidências'], w['Auditores'], w['Checklist']
    for k, a in enumerate(ams):
        for col, key in zip('CDEFGHKLM', ('oque', 'req', 'periodo', 'pop', 'risco', 'inicio', 'verif', 'desv', 'nota')):
            if a[key] not in (None, ''):
                A[f'{col}{12 + k}'] = a[key]
    for k, x in enumerate(evs):
        for col, key in zip('CDEFGH', ('perg', 'req', 'ent', 'obs', 'reg', 'const')):
            if x[key] not in (None, ''):
                E[f'{col}{7 + k}'] = x[key]
    for k, a in enumerate(auds):
        c = chr(70 + k)
        if a['nome']:
            U[f'{c}5'], U[f'{c}6'] = a['nome'], a['papel']
        for j, n in enumerate(a['notas']):
            if n is not None:
                U[f'{c}{8 + j}'] = n
    for k, val in enumerate(check or ()):
        if val:
            K[f'D{5 + k}'] = val
    w.save(OUTDIR + '/%s.xlsx' % nome)


for nome, ex, check in (('t1', EX1, ['Sim'] * 12), ('t2', EX2, ['Sim'] * 6 + ['Parcial', 'Não'])):
    preencher(nome, ex['ams'], ex['evs'], ex['auds'], check=check)

# ---- t3: casos de borda
AK = ('oque', 'req', 'periodo', 'pop', 'risco', 'inicio', 'verif', 'desv', 'nota')
ams3 = [dict(zip(AK, r)) for r in [
    ('a', '8.4', 'p', 5, ALTO, 1, 5, 0, ''),          # até 5: todos
    ('b', '8.4', 'p', 6, BAIXO, 1, 3, 0, ''),         # 5 × 0,6 = 3
    ('c', '8.4', 'p', 20, ALTO, 4, 8, 2, ''),         # 7,5 → 8, passo 2: início fora
    ('d', '8.4', 'p', 21, MEDIO, 1, 8, 1, ''),        # desvio isolado
    ('e', '8.4', 'p', 100, BAIXO, 1, 6, 2, ''),       # desvio repetido
    ('f', '8.4', 'p', 501, ALTO, 1, 29, 0, ''),       # 30 planejados: incompleta
    ('g', '8.4', 'p', None, ALTO, 1, 1, 0, ''),       # falta a população
    ('h', '8.4', 'p', 30, None, 1, 1, 0, ''),         # falta o risco
    ('i', '8.4', 'p', 30, MEDIO, None, 8, 0, ''),     # falta o início
    ('j', '8.4', 'p', 30, MEDIO, 1, None, None, ''),  # falta verificados
    ('k', '8.4', 'p', 10, MEDIO, 1, 11, 0, ''),       # verificados acima da população
    ('l', '8.4', 'p', 30, MEDIO, 1, 8, None, ''),     # falta desvios
    ('m', '8.4', 'p', 30, MEDIO, 1, 8, 9, ''),        # mais desvios que verificados
    ('n', '8.4', 'p', 500, BAIXO, 1, 9, 0, ''),       # 15 × 0,6 = 9, sem erro de arredondamento
    ('o', '8.4', 'p', 50, BAIXO, 1, 5, 0, '')]]       # 8 × 0,6 = 4,8 → 5
EK = ('perg', 'req', 'ent', 'obs', 'reg', 'const')
evs3 = [dict(zip(EK, r)) for r in [
    ('p', '7.1', 'e', '', '', CONFORME),   # fraca, mas conforme: OK
    ('p', '7.1', '', 'o', '', NC),         # média: uma fonte objetiva basta, OK
    ('p', '7.1', 'e', '', '', NC),         # fraca: sem evidência objetiva
    ('p', '7.1', 'e', '', 'r', NC),        # forte: OK
    ('p', '7.1', '', '', '', CONFORME),    # falta a evidência
    ('', '7.1', 'e', '', '', CONFORME),    # falta a pergunta
    ('p', '', 'e', '', '', CONFORME),      # falta o requisito
    ('p', '7.1', 'e', 'o', '', ''),        # falta a constatação
    ('p', '7.1', 'e', 'o', 'r', OM)]]      # forte: OK
N = None
auds3 = [dict(nome='A', papel='x', notas=[3] * 12),                                   # 100 %: apto a liderar
         dict(nome='B', papel='x', notas=[2, 2, 2, 2, 2, 2, 2, 2, N, N, N, N]),       # 8 observados, 67 %: em equipe
         dict(nome='C', papel='x', notas=[3, 3, 3, 3, 3, 1, 3, 3, 3, 3, 3, 3]),       # crítico com 1: em formação
         dict(nome='D', papel='x', notas=[3, 3, 3, 3, 3, 3, 3, N, N, N, N, N])]       # 7 observados: insuficiente
preencher('t3', ams3, evs3, auds3)
print('ok')
