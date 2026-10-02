"""Confere Recursos-modelo.xlsx e preenche cópias com dados de teste.  Uso: rec_test_make.py <modelo.xlsx> <pasta de saída>"""
import openpyxl, sys, os
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
from rec_data import (ALTA, BAIXA, CORR, EX1, EX2, FIS, MEDIA, NAO, PREV, SIM, amb_conf, amb_sit, cap_conf, cap_sit, disp, infra_conf, necessarias, oc_conf,
                      paradas, prev_sit, proxima)
from datetime import date
SRC, OUTDIR = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(SRC, data_only=True)
wf0 = openpyxl.load_workbook(SRC)
for ws in wb.worksheets:
    f = wf0[ws.title]
    print('==', ws.title, ws.dimensions, 'dv', len(f.data_validations.dataValidation), 'cf', sum(len(v) for v in f.conditional_formatting._cf_rules.values()),
          'charts', len(f._charts), 'comments', sum(1 for r in f.iter_rows() for c in r if c.comment))
err = [(ws.title, c.coordinate, c.value) for ws in wb.worksheets for r in ws.iter_rows() for c in r
       if isinstance(c.value, str) and ((c.value.startswith('#') and len(c.value) > 1) or 'Err:' in c.value)]
print('erros', err)
print('vazio:', wb['Pessoas']['E41'].value, '|', wb['Infraestrutura']['E55'].value, '|', wb['Manutenções']['E91'].value, '|', wb['Ambiente']['E41'].value, '|', wb['Painel']['C29'].value)


def dd(v):
    return v.date() if hasattr(v, 'date') else v


def r4(v):
    return round(v, 4) if isinstance(v, float) else v


for aba, ex in (('Exemplo 1 - Pizzaria', EX1), ('Exemplo 2 - Indústria', EX2)):
    s = wb[aba]
    h = ex['head']
    # as linhas de cada bloco começam depois do subtítulo de cada faixa
    p1 = next(r for r in range(1, s.max_row) if s[f'B{r}'].value == 'Função') + 1
    i1 = next(r for r in range(1, s.max_row) if s[f'B{r}'].value == 'Código') + 1
    m1 = next(r for r in range(1, s.max_row) if s[f'B{r}'].value == 'Data') + 1
    a1 = next(r for r in range(1, s.max_row) if s[f'B{r}'].value == 'Fator') + 1
    got_p = [(s[f'H{r}'].value, s[f'L{r}'].value, s[f'M{r}'].value) for r in range(p1, p1 + len(ex['caps']))]
    esp_p = [(necessarias(c), cap_sit(c), cap_conf(c)) for c in ex['caps']]
    got_i = [(dd(s[f'I{r}'].value) or None, s[f'L{r}'].value, s[f'M{r}'].value, s[f'N{r}'].value, r4(s[f'O{r}'].value) if s[f'O{r}'].value != '' else None, s[f'P{r}'].value)
             for r in range(i1, i1 + len(ex['infra']))]
    esp_i = [(proxima(i), prev_sit(i, h['ref']), *paradas(i, ex['ocs'], h['ini'], h['fim']), r4(disp(i, ex['ocs'], h['ini'], h['fim'])), infra_conf(i, ex)) for i in ex['infra']]
    got_m = [s[f'Q{r}'].value for r in range(m1, m1 + len(ex['ocs']))]
    esp_m = [oc_conf(o, ex['infra']) for o in sorted(ex['ocs'], key=lambda o: o['data'])]
    got_a = [(s[f'N{r}'].value, s[f'Q{r}'].value) for r in range(a1, a1 + len(ex['amb']))]
    esp_a = [(amb_sit(a), amb_conf(a)) for a in ex['amb']]
    print(aba, '| pessoas', 'OK' if got_p == esp_p else 'DIFERE %s' % got_p, '| infra', 'OK' if got_i == esp_i else 'DIFERE %s\n   esperado %s' % (got_i, esp_i),
          '| manutenções', 'OK' if got_m == esp_m else 'DIFERE %s' % got_m, '| ambiente', 'OK' if got_a == esp_a else 'DIFERE %s' % got_a)


def preencher(nome, head=None, caps=(), infra=(), ocs=(), amb=(), check=None):
    w = openpyxl.load_workbook(SRC)
    P, I, M, A, K = w['Pessoas'], w['Infraestrutura'], w['Manutenções'], w['Ambiente'], w['Checklist']
    if head:
        I['E4'], I['E5'], I['E6'] = head['ini'], head['fim'], head['meta']
    for rows, ws, cols, r0 in ((caps, P, 'CDEFGIJ', 7), (infra, I, 'CDEFGHIJK', 10), (ocs, M, 'CDFGHIJKL', 7), (amb, A, 'CDEFGHIJKL', 7)):
        for k, row in enumerate(rows):
            for col, v in zip(cols, row):
                if v not in (None, ''):
                    ws[f'{col}{r0 + k}'] = v
    for k, v in enumerate(check or ()):
        if v:
            K[f'D{5 + k}'] = v
    w.save(OUTDIR + '/%s.xlsx' % nome)


def linhas(ex):
    caps = [(c['funcao'], c['periodo'], c['unid'], c['dem'], c['prod'], c['esc'], c['qual']) for c in ex['caps']]
    infra = [(i['cod'], i['nome'], i['tipo'], i['local'], i['crit'], i['interv'], i['ultima'], i['cont'], i['prog']) for i in ex['infra']]
    ocs = [(o['data'], o['cod'], o['tipo'], o['horas'], o['oque'], o['causa'], o['acao'], o['afetou'], o['trat']) for o in ex['ocs']]
    amb = [(a['fator'], a['tipo'], a['local'], a['porque'], a['unid'], a['min'], a['max'], a['controle'], a['valor'], a['acao']) for a in ex['amb']]
    return caps, infra, ocs, amb


for nome, ex, check in (('t1', EX1, ['Sim'] * 12), ('t2', EX2, ['Sim'] * 6 + ['Parcial', 'Não'])):
    preencher(nome, ex['head'], *linhas(ex), check=check)

# ---- t3: casos de borda
D = date
H3 = dict(ini=D(2027, 7, 1), fim=D(2027, 7, 31), ref=D(2027, 7, 31), meta=0.95)
caps3 = [('A', 'p', 'u', 10, 3, 4, 4),       # 4 necessárias: justo, OK
         ('B', 'p', 'u', 10, 3, 3, 3),       # falta 1 pessoa
         ('C', 'p', 'u', 10, 3, 2, 2),       # faltam 2 pessoas
         ('D', 'p', 'u', 10, 3, 5, 0),       # com folga, sem qualificado
         ('E', 'p', 'u', None, 3, 2, 2),     # falta a demanda
         ('F', 'p', 'u', 10, None, 2, 2),    # falta a produtividade
         ('G', 'p', 'u', 10, 3, None, None), # falta o número escalado
         ('H', 'p', 'u', 10, 3, 4, None),    # falta o número de qualificados
         ('I', 'p', 'u', 10, 3, 4, 5),       # mais qualificados que escalados
         ('J', 'p', 'u', 60, 20, 3, 3)]      # divisão exata: 3, justo
infra3 = [('X1', 'n', 'Equipamento', 'l', ALTA, 3, D(2027, 6, 1), 'c', 744),     # em dia, 40 h paradas: abaixo da meta
          ('X2', 'n', 'Equipamento', 'l', ALTA, 1, D(2027, 7, 28), '', None),     # crítico sem contingência
          ('X3', 'n', 'Equipamento', 'l', MEDIA, None, None, '', 100),           # sem plano de preventiva
          ('X4', 'n', 'Equipamento', 'l', BAIXA, None, None, '', 100),           # sem plano, mas baixa: OK
          ('X5', 'n', 'Equipamento', 'l', ALTA, 1, None, 'c', 100),              # plano sem registro: atrasada
          ('X6', 'n', 'Equipamento', 'l', MEDIA, 1, D(2027, 7, 5), '', 100),     # vence em 5 dias
          ('X7', 'n', 'Equipamento', 'l', MEDIA, 1, D(2027, 7, 7), '', 100),     # vence em 7 dias: no limite
          ('X8', 'n', 'Equipamento', 'l', MEDIA, 1, D(2027, 7, 8), '', 100),     # 8 dias: em dia
          ('X9', 'n', None, 'l', MEDIA, 1, D(2027, 7, 8), '', 100),              # falta o tipo
          ('X10', 'n', 'Equipamento', 'l', None, 1, D(2027, 7, 8), '', 100),     # falta a criticidade
          ('X11', 'n', 'Equipamento', 'l', ALTA, 3, D(2027, 6, 1), 'c', 100)]    # 5 h em 100: 95 %, na meta
ocs3 = [(D(2027, 7, 5), 'X1', CORR, 20, 'o', 'c', 'a', NAO, ''),       # OK
        (D(2027, 7, 10), 'X1', CORR, 20, 'o', 'c', 'a', SIM, 't'),     # OK
        (D(2027, 6, 30), 'X1', CORR, 50, 'o', 'c', 'a', NAO, ''),      # fora do período
        (D(2027, 8, 1), 'X1', CORR, 50, 'o', 'c', 'a', NAO, ''),       # fora do período
        (D(2027, 7, 15), 'X1', PREV, 8, 'p', '', '', '', ''),          # preventiva não conta
        (D(2027, 7, 20), 'X11', CORR, 5, 'o', 'c', 'a', NAO, ''),      # OK
        (D(2027, 7, 20), 'Z', CORR, 1, 'o', 'c', 'a', NAO, ''),        # item não cadastrado
        (None, 'X2', CORR, 1, 'o', 'c', 'a', NAO, ''),                 # falta a data
        (D(2027, 7, 20), 'X2', None, 1, 'o', 'c', 'a', NAO, ''),       # falta o tipo
        (D(2027, 7, 20), 'X2', CORR, None, 'o', 'c', 'a', NAO, ''),    # falta o tempo parado
        (D(2027, 7, 20), 'X2', CORR, 1, 'o', '', 'a', NAO, ''),        # falta a causa
        (D(2027, 7, 20), 'X2', CORR, 1, 'o', 'c', '', NAO, ''),        # falta a ação
        (D(2027, 7, 20), 'X2', CORR, 1, 'o', 'c', 'a', '', ''),        # falta dizer se afetou
        (D(2027, 7, 20), 'X2', CORR, 1, 'o', 'c', 'a', SIM, ''),       # produto afetado sem tratamento
        (D(2027, 7, 21), None, CORR, 1, 'o', 'c', 'a', NAO, '')]       # falta o item
amb3 = [('A', FIS, 'l', 'p', 'u', 0, 5, 'c', 3, ''),         # dentro
        ('B', FIS, 'l', 'p', 'u', 0, 5, 'c', 6, ''),         # fora, sem ação
        ('C', FIS, 'l', 'p', 'u', 0, 5, 'c', -1, 'a'),       # fora, com ação
        ('D', FIS, 'l', 'p', 'u', None, 5, 'c', 5, ''),      # no limite: dentro
        ('E', FIS, 'l', 'p', 'u', 10, None, 'c', None, ''),  # sem medição
        ('F', FIS, 'l', 'p', 'u', None, None, 'c', 3, ''),   # falta o limite
        ('G', FIS, 'l', '', 'u', 0, 5, 'c', 3, ''),          # falta por que afeta
        ('H', FIS, 'l', 'p', 'u', 0, 5, '', 3, '')]          # falta o controle
preencher('t3', H3, caps3, infra3, ocs3, amb3)
print('ok')
