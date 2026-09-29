import openpyxl, sys, os
S = os.path.dirname(os.path.abspath(__file__)); sys.path.insert(0, S)
from iso_data import EX1, REQ, DOCS
from datetime import date
SRC, OUTDIR = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(SRC, data_only=True)
wf0 = openpyxl.load_workbook(SRC)
for ws in wb.worksheets:
    f = wf0[ws.title]
    print('==', ws.title, ws.dimensions, 'dv', len(f.data_validations.dataValidation), 'cf', sum(len(v) for v in f.conditional_formatting._cf_rules.values()),
          'charts', len(f._charts), 'comments', sum(1 for r in f.iter_rows() for c in r if c.comment))
err = [(ws.title, c.coordinate, c.value) for ws in wb.worksheets for r in ws.iter_rows() for c in r
       if isinstance(c.value, str) and ((c.value.startswith('#') and len(c.value) > 1) or 'Err:' in c.value)]
print('erros', err)
r = wb['Resumo']; print('resumo vazio', [[r[f'{c}{k}'].value for c in 'DEFGHI'] for k in (7, 14)], [r[f'E{k}'].value for k in range(17, 23)])
print('plano vazio', [wb['Plano de ação'][f'E{k}'].value for k in range(29, 36)]); print('docs vazio', [wb['Documentos'][f'F{k}'].value for k in range(31, 38)])
for n in ('Exemplo 1 - Pizzaria', 'Exemplo 2 - Compras'):
    s = wb[n]
    rows = [[s[f'{c}{k}'].value for c in 'CDEFGHI'] for k in range(10, 90) if s[f'B{k}'].value == 'Total' or s[f'B{k}'].value in (4, 5, 6, 7, 8, 9, 10)]
    print(n); [print('   ', x) for x in rows]
    print('   ', [(s[f'B{k}'].value, s[f'E{k}'].value) for k in range(10, 95) if s[f'B{k}'].value in ('Percentual de atendimento', 'Seção com o menor atendimento', 'Estágio do sistema')])


def linhas(ws):
    return {str(ws[f'B{k}'].value): k for k in range(9, 61) if ws[f'K{k}'].value is not None}


def preencher(nome, mud=None, plano=(), docs=None):
    w = openpyxl.load_workbook(SRC); d = w['Diagnóstico']; L = linhas(d)
    assert len(L) == 45, len(L)
    d['D4'] = 'Pizzaria'; d['H4'] = date(2026, 10, 2)
    for num, (sit, ev, falta) in EX1['itens'].items():
        k = L[num]; d[f'F{k}'] = sit; d[f'G{k}'] = ev; d[f'H{k}'] = falta or None
    for num, campos in (mud or {}).items():
        for col, v in campos.items():
            d[f'{col}{L[num]}'] = v
    p = w['Plano de ação']
    for i, row in enumerate(plano):
        for col, v in zip('CFGHIJ', row):
            if v is not None:
                p[f'{col}{7 + i}'] = v
    o = w['Documentos']
    for i, v in enumerate(docs or []):
        if v:
            o[f'H{6 + i}'] = v[0]; o[f'F{6 + i}'] = v[1]
    w.save(OUTDIR + '/%s.xlsx' % nome)


preencher('t1')
preencher('t2', mud={'8.3': {'E': 'Não', 'F': None, 'G': None, 'H': None},            # NA sem justificativa
                     '7.1.5': {'E': 'Não', 'F': None, 'G': 'Não usa instrumentos', 'H': None},  # NA com justificativa
                     '4.1': {'G': None},                                           # sem evidência
                     '4.2': {'H': None},                                           # em parte, sem lacuna
                     '9.3': {'F': None, 'G': None, 'H': None}},                    # não avaliado
          plano=[('4.3', 'Escrever o escopo', 'Gerente', date(2026, 10, 30), 'Em andamento', None),
                 ('5.2', 'Escrever a política', 'Dono', date(2026, 1, 10), 'Não iniciada', None),
                 ('6.3', 'Criar rotina de mudanças', 'Gerente', date(2027, 3, 1), 'Concluída', date(2026, 10, 1)),
                 ('9.9', 'Requisito inexistente', None, None, None, None),
                 (None, 'Ação sem requisito', 'Gerente', date(2027, 3, 1), 'Cancelada', None)],
          docs=[('Existe', 'Declaração de escopo'), ('Existe', None), ('Em elaboração', 'Registros'), ('Não existe', None), ('Não aplicável', None)])
preencher('t3', mud={n: {'F': 'Atende', 'H': None} for n in EX1['itens']})          # tudo atende
preencher('t4', mud={n: {'F': 'Atende', 'H': None} for n in EX1['itens'] if n != '5.2'})   # 1 não atende, acima de 80%
