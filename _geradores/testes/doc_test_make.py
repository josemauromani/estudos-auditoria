"""Confere Informacao-Documentada-modelo.xlsx e preenche cópias com dados de teste.  Uso: doc_test_make.py <modelo.xlsx> <pasta de saída>"""
import openpyxl, sys, os
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
from doc_data import EX1, EX2, EXT, situacao, proxima
from datetime import date
SRC, OUTDIR = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(SRC, data_only=True)
wf0 = openpyxl.load_workbook(SRC)
for ws in wb.worksheets:
    f = wf0[ws.title]
    print('==', ws.title, ws.dimensions, 'dv', len(f.data_validations.dataValidation), 'cf', sum(len(v) for v in f.conditional_formatting._cf_rules.values()),
          'charts', len(f._charts), 'comments', sum(1 for r in f.iter_rows() for c in r if c.comment))
err = [(ws.title, c.coordinate, c.value) for ws in wb.worksheets for r in ws.iter_rows() for c in r
       if isinstance(c.value, str) and ((c.value.startswith('#') and len(c.value) > 1) or 'Err:' in c.value)]
print('erros', err)
print('vazio:', wb['Lista mestra']['E46'].value, '|', wb['Registros']['E32'].value, '|', wb['Externos']['E27'].value, '|', wb['Alterações']['F31'].value,
      '|', wb['Checklist']['D24'].value)
for nome, ex in (('Exemplo 1 - Pizzaria', EX1), ('Exemplo 2 - Compras', EX2)):
    s = wb[nome]
    ref = ex['head']['data']
    r0 = 12
    sit = [s[f'O{r}'].value for r in range(r0, r0 + len(ex['docs']))]
    esp = [situacao(d, ref) for d in ex['docs']]
    prox = [s[f'M{r}'].value.date() if s[f'M{r}'].value else None for r in range(r0, r0 + len(ex['docs']))]
    esp_p = [proxima(d) if d['data'] else None for d in ex['docs']]
    res = {s[f'B{r}'].value: s[f'E{r}'].value for r in range(r0 + len(ex['docs']), s.max_row + 1) if isinstance(s[f'B{r}'].value, str) and s[f'C{r}'].value is None}
    print(nome, '| situações', 'OK' if sit == esp else 'DIFERE %s' % sit, '| próximas', 'OK' if prox == esp_p else 'DIFERE %s' % prox,
          '| resumo', {k: v for k, v in res.items() if k in ('Documentos cadastrados', 'Documentos sem controle', 'Registros cadastrados', 'Em vigor, sem os obsoletos')})


def preencher(nome, ex=None, lista=None, docs=(), regs=(), ext=(), alts=(), check=None):
    w = openpyxl.load_workbook(SRC)
    L, R, X, A, C = w['Lista mestra'], w['Registros'], w['Externos'], w['Alterações'], w['Checklist']
    if ex:
        lista = dict({'D4': ex['head']['org'], 'D5': ex['head']['data']}, **(lista or {}))
        docs = docs or [(d['codigo'] or None, d['titulo'], d['tipo'], d['processo'], d['rev'] or None, d['data'], d['elaborou'], d['aprovou'] or None, d['meio'], d['onde'],
                         d['meses'], d['status']) for d in ex['docs']]
        regs = regs or [(g['nome'], g['form'], g['processo'], g['gerado'], g['guardado'], g['meio'], g['meses'], g['acesso'], g['disposicao']) for g in ex['regs']]
    for ref, v in (lista or {}).items():
        if v is not None:
            L[ref] = v
    for k, row in enumerate(docs):
        for col, v in zip('CDEFGHIJKLMO', row):
            if v is not None:
                L[f'{col}{9 + k}'] = v
    for k, row in enumerate(regs):
        for col, v in zip('CDEFGHIJK', row):
            if v is not None:
                R[f'{col}{7 + k}'] = v
    for k, row in enumerate(ext):
        for col, v in zip('CDEFGHIJ', row):
            if v is not None:
                X[f'{col}{7 + k}'] = v
    for k, row in enumerate(alts):
        for col, v in zip('CDEFGHIJ', row):
            if v is not None:
                A[f'{col}{7 + k}'] = v
    for k, v in enumerate(check or ()):
        if v:
            C[f'D{5 + k}'] = v
    w.save(OUTDIR + '/%s.xlsx' % nome)


D = date
ext2 = [(e['nome'], e['origem'], e['versao'], e['onde'], e['como'], 'Coordenador da Qualidade', e['verificado'], e['meses']) for e in EXT]
# t1: compras completo, com os externos e uma alteração
preencher('t1', EX2, ext=ext2, alts=[(D(2026, 10, 16), 'PR-SUP-01', '6', 'Cadastro do item antes da compra.', 'RNC 2026-31', 'Gerente de Suprimentos', 'Diretor geral', 'Sim')],
          check=['Sim'] * 12)
# t2: pizzaria, sem data de referência (usa a data de hoje)
preencher('t2', EX1, lista={'D5': None}, check=['Sim'] * 6 + ['Parcial', 'Não'])
# t3: erros propositais
preencher('t3', lista={'D4': 'Teste', 'D5': D(2027, 3, 5), 'D6': 12},
          docs=[('PR-01', 'Completo', 'PR', 'Proc', '1', D(2026, 6, 1), 'A', 'B', 'Papel', 'Local', None, 'Vigente'),           # vigente: 06/2026 + 12 = 06/2027
                ('PR-02', 'Vencido pelo padrão', 'PR', 'Proc', '1', D(2026, 2, 1), 'A', 'B', 'Papel', 'Local', None, 'Vigente'),  # 02/2027 < 05/03/2027
                ('PR-03', 'Vencido pelo intervalo próprio', 'PR', 'Proc', '1', D(2026, 12, 1), 'A', 'B', 'Papel', 'Local', 2, 'Vigente'),
                (None, 'Sem código', 'IT', 'Proc', '1', D(2026, 6, 1), 'A', 'B', 'Papel', 'Local', None, 'Vigente'),
                ('IT-01', 'Sem revisão', 'IT', 'Proc', None, D(2026, 6, 1), 'A', 'B', 'Papel', 'Local', None, 'Vigente'),
                ('IT-02', 'Sem data', 'IT', 'Proc', '2', None, 'A', 'B', 'Papel', 'Local', None, 'Vigente'),
                ('IT-03', 'Sem aprovação', 'IT', 'Proc', '1', D(2026, 6, 1), 'A', None, 'Papel', 'Local', None, 'Vigente'),
                ('IT-04', 'Em revisão', 'IT', 'Proc', '1', D(2026, 6, 1), 'A', 'B', 'Papel', 'Local', None, 'Em revisão'),
                ('IT-05', 'Obsoleto', 'IT', 'Proc', '1', D(2020, 6, 1), 'A', 'B', 'Papel', 'Local', None, 'Obsoleto'),
                ('IT-06', 'Sem status', 'IT', 'Proc', '1', D(2026, 6, 1), 'A', 'B', 'Papel', 'Local', None, None)],
          regs=[('Completo', 'FR-01', 'Proc', 'Posto', 'Pasta', 'Papel', 12, 'Todos', 'Picotar'),
                ('Sem local', 'FR-01', 'Proc', 'Posto', None, 'Papel', 12, 'Todos', 'Picotar'),
                ('Sem prazo', 'FR-01', 'Proc', 'Posto', 'Pasta', 'Papel', None, 'Todos', 'Picotar'),
                ('Sem meio', 'FR-01', 'Proc', 'Posto', 'Pasta', None, 12, 'Todos', 'Picotar'),
                ('Sem disposição', 'FR-01', 'Proc', 'Posto', 'Pasta', 'Papel', 60, 'Todos', None)],
          ext=[('Verificado', 'X', 'v1', 'Proc', 'Site', 'Resp', D(2027, 1, 10), 6),
               ('Vencido', 'X', 'v1', 'Proc', 'Site', 'Resp', D(2026, 1, 10), 6),
               ('Sem verificação', 'X', 'v1', 'Proc', 'Site', 'Resp', None, None),
               ('Sem responsável', 'X', 'v1', 'Proc', 'Site', None, D(2027, 1, 10), None),
               ('Vencido pelo padrão da lista', 'X', 'v1', 'Proc', 'Site', 'Resp', D(2026, 1, 10), None)],
          alts=[(D(2027, 1, 5), 'PR-01', '2', 'Mudou.', 'Revisão', 'A', 'B', 'Sim'),
                (D(2027, 1, 5), 'ZZ-99', '2', 'Código que não existe.', 'Revisão', 'A', 'B', 'Sim'),
                (D(2027, 1, 5), 'PR-02', None, 'Sem revisão.', 'Revisão', 'A', 'B', 'Sim'),
                (D(2027, 1, 5), 'PR-03', '2', None, 'Revisão', 'A', 'B', 'Sim'),
                (D(2027, 1, 5), 'IT-01', '2', 'Sem aprovação.', 'Revisão', 'A', None, 'Sim'),
                (D(2027, 1, 5), 'IT-02', '3', 'Sem conferir os ligados.', 'Revisão', 'A', 'B', None),
                (D(2027, 1, 5), 'IT-03', '2', 'Ligados não conferidos.', 'Revisão', 'A', 'B', 'Não'),
                (D(2027, 1, 5), 'IT-04', '2', 'Sem ligados.', 'Revisão', 'A', 'B', 'Não há')])
print('ok')
