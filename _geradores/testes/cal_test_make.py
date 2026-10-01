"""Confere Calibracao-modelo.xlsx e preenche cópias com dados de teste.  Uso: cal_test_make.py <modelo.xlsx> <pasta de saída>"""
import openpyxl, sys, os
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
from cal_data import CAL, CALEXT, EX1, EX2, NAO, SIM, VER, VERINT, cal_conf, inst_conf, resultado, situacao
from datetime import date
SRC, OUTDIR = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(SRC, data_only=True)
wf0 = openpyxl.load_workbook(SRC)
for ws in wb.worksheets:
    f = wf0[ws.title]
    print('==', ws.title, ws.dimensions, 'dv', len(f.data_validations.dataValidation), 'cf', sum(len(v) for v in f.conditional_formatting._cf_rules.values()),
          'charts', len(f._charts), 'comments', sum(1 for r in f.iter_rows() for c in r if c.comment))
err = [(ws.title, c.coordinate, c.value) for ws in wb.worksheets for r in ws.iter_rows() for c in r
       if isinstance(c.value, str) and ((c.value.startswith('#') and len(c.value) > 1) or 'Err:' in c.value)]
print('erros', err)
print('vazio:', wb['Instrumentos']['E53'].value, '|', wb['Calibrações']['E91'].value, '|', wb['Painel']['C19'].value)
for aba, ex in (('Exemplo 1 - Pizzaria', EX1), ('Exemplo 2 - Indústria', EX2)):
    s = wb[aba]
    ref = ex['head']['ref']
    cods = [i['cod'] for i in ex['insts']]
    ri = [r for r in range(1, s.max_row + 1) if s[f'B{r}'].value in cods]
    got_i = [(s[f'O{r}'].value, s[f'Q{r}'].value) for r in ri]
    esp_i = [(situacao(i, ref), inst_conf(i, ref)) for i in ex['insts']]
    cals = sorted(ex['cals'], key=lambda c: c['data'])
    rc = [r for r in range(1, s.max_row + 1) if isinstance(s[f'C{r}'].value, str) and ' · ' in s[f'C{r}'].value and s[f'C{r}'].value.split(' · ')[0] in cods]
    got_c = [(s[f'J{r}'].value, s[f'Q{r}'].value) for r in rc]
    esp_c = [(resultado(c, ex['insts']), cal_conf(c, ex['insts'])) for c in cals]
    print(aba, '| instrumentos', 'OK' if got_i == esp_i else 'DIFERE %s' % got_i, '| calibrações', 'OK' if got_c == esp_c else 'DIFERE %s' % got_c)


def preencher(nome, ref=None, insts=(), cals=(), check=None):
    w = openpyxl.load_workbook(SRC)
    I, C, K = w['Instrumentos'], w['Calibrações'], w['Checklist']
    if ref:
        I['E4'] = ref
    for k, row in enumerate(insts):
        for col, v in zip('CDEFGHIJKLMN', row):
            if v not in (None, ''):
                I[f'{col}{8 + k}'] = v
    for k, row in enumerate(cals):
        for col, v in zip('CDFGHIJMN', row):
            if v not in (None, ''):
                C[f'{col}{7 + k}'] = v
    for k, v in enumerate(check or ()):
        if v:
            K[f'D{5 + k}'] = v
    w.save(OUTDIR + '/%s.xlsx' % nome)


def do_exemplo(ex):
    insts = [(i['cod'], i['nome'], i['local'], i['uso'], i['unid'], i['tol'], i['res'], i['tipo'], i['interv'], i['ultima'], i['ema'], i['emuso']) for i in ex['insts']]
    cals = [(c['data'], c['cod'], c['tipo'], c['reg'], c['ponto'], c['erro'], c['inc'], c['acao'], c['impacto']) for c in ex['cals']]
    return insts, cals


for nome, ex, check in (('t1', EX1, ['Sim'] * 12), ('t2', EX2, ['Sim'] * 6 + ['Parcial', 'Não'])):
    preencher(nome, ex['head']['ref'], *do_exemplo(ex), check=check)
D = date
ref = D(2027, 6, 30)
insts3 = [('A', 'ok', 'L', 'mede', 'g', 40, 1, CALEXT, 12, D(2027, 1, 1), 2, SIM),          # em dia
          ('B', 'vence', 'L', 'mede', 'g', 40, 1, VERINT, 1, D(2027, 6, 10), 2, SIM),       # vence em 30 dias (10/07)
          ('C', 'vencido', 'L', 'mede', 'g', 40, 1, VERINT, 1, D(2027, 5, 1), 2, SIM),      # vencido
          ('D', 'sem data', 'L', 'mede', 'g', 40, 1, VERINT, 1, None, 2, SIM),             # sem calibração
          ('E', 'fora', 'L', 'mede', 'g', 40, 1, VERINT, 1, D(2027, 1, 1), 2, NAO),         # fora de uso
          ('F', 'grossa', 'L', 'mede', 'g', 40, 5, CALEXT, 12, D(2027, 1, 1), 2, SIM),     # resolução grossa
          ('G', 'no limite', 'L', 'mede', 'g', 40, 4, CALEXT, 12, D(2027, 1, 1), 2, SIM),  # 4 = 40/10: adequado
          ('H', 'sem uso', 'L', None, 'g', 40, 1, CALEXT, 12, D(2027, 1, 1), 2, SIM),      # falta o que mede
          ('I', 'sem tipo', 'L', 'mede', 'g', 40, 1, None, 12, D(2027, 1, 1), 2, SIM),     # falta o tipo
          ('J', 'sem ema', 'L', 'mede', 'g', 40, 1, CALEXT, 12, D(2027, 1, 1), None, SIM)]  # falta o erro máximo
cals3 = [(D(2027, 1, 1), 'A', CAL, 'Cert 1', '100 g', 1.5, 0.5, None, None),          # 2,0 = 2: aprovado no limite
         (D(2027, 1, 1), 'A', CAL, 'Cert 2', '100 g', -1.8, 0.3, None, None),         # 2,1: reprovado, falta a ação
         (D(2027, 1, 1), 'A', CAL, 'Cert 3', '100 g', 3, 0, 'Ajustada', None),        # reprovado, falta avaliar
         (D(2027, 1, 1), 'A', VER, 'Ver 1', '100 g', 3, 0, 'Ajustada', 'Revisto'),     # reprovado, completo
         (D(2027, 1, 1), 'Z', VER, 'Ver 2', '100 g', 1, 0, None, None),               # não cadastrado
         (D(2027, 1, 1), 'B', VER, None, '100 g', 1, 0, None, None),                  # falta o registro
         (D(2027, 1, 1), 'B', VER, 'Ver 3', '100 g', None, None, None, None),         # falta o erro
         (D(2027, 1, 1), 'J', VER, 'Ver 4', '100 g', 1, 0, None, None),               # falta o critério no cadastro
         (D(2027, 1, 1), None, VER, 'Ver 5', None, None, None, None, None)]           # falta o instrumento
preencher('t3', ref, insts3, cals3)
print('ok')
