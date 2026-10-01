"""Confere Liberacao-modelo.xlsx e preenche cópias com dados de teste.  Uso: lib_test_make.py <modelo.xlsx> <pasta de saída>"""
import openpyxl, sys, os
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
from lib_data import (APOS, AUTORIZADO, CONCES, DEVOLV, EX1, EX2, FINAL, LIBERADO, PROC, R_CONF, R_NAO, RECEB, RECLAS, REFUGO, RETIDO, RETRAB, lib_conf, lib_leitura,
                      pnc_conf, pnc_sit)
from datetime import date
SRC, OUTDIR = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(SRC, data_only=True)
wf0 = openpyxl.load_workbook(SRC)
for ws in wb.worksheets:
    f = wf0[ws.title]
    print('==', ws.title, ws.dimensions, 'dv', len(f.data_validations.dataValidation), 'cf', sum(len(v) for v in f.conditional_formatting._cf_rules.values()),
          'charts', len(f._charts), 'comments', sum(1 for r in f.iter_rows() for c in r if c.comment))
err = [(ws.title, c.coordinate, c.value) for ws in wb.worksheets for r in ws.iter_rows() for c in r
       if isinstance(c.value, str) and ((c.value.startswith('#') and len(c.value) > 1) or 'Err:' in c.value)]
print('erros', err)
print('vazio:', wb['Autoridades']['E18'].value, '|', wb['Liberação']['E72'].value, '|', wb['Produto não conforme']['E72'].value, '|', wb['Painel']['C35'].value, '|', wb['Checklist']['D24'].value)
for aba, ex in (('Exemplo 1 - Pizzaria', EX1), ('Exemplo 2 - Indústria', EX2)):
    s = wb[aba]
    lotes = [l['lote'] for l in ex['libs']]
    rl = [r for r in range(1, s.max_row + 1) if s[f'C{r}'].value in lotes and isinstance(s[f'E{r}'].value, int)]
    got_l = [(s[f'H{r}'].value, s[f'Q{r}'].value) for r in rl]
    esp_l = [(lib_leitura(l), lib_conf(l)) for l in ex['libs']]
    nums = [p['num'] for p in ex['pncs']]
    rp = [r for r in range(1, s.max_row + 1) if s[f'B{r}'].value in nums]
    got_p = [(s[f'P{r}'].value, s[f'Q{r}'].value) for r in rp]
    esp_p = [(pnc_sit(p), pnc_conf(p, ex['pncs'])) for p in ex['pncs']]
    print(aba, '| liberações', 'OK' if got_l == esp_l else 'DIFERE %s' % got_l, '| registros', 'OK' if got_p == esp_p else 'DIFERE %s' % got_p)


def preencher(nome, head=None, aut=(), libs=(), pncs=(), check=None):
    w = openpyxl.load_workbook(SRC)
    A, L, P, K = w['Autoridades'], w['Liberação'], w['Produto não conforme'], w['Checklist']
    if head:
        A['D4'] = head
    for k, row in enumerate(aut):
        for col, v in zip('EFGH', row):
            if v:
                A[f'{col}{7 + k}'] = v
    for k, row in enumerate(libs):
        for col, v in zip('CDEFGHJKLMN', row):
            if v not in (None, ''):
                L[f'{col}{7 + k}'] = v
    for k, row in enumerate(pncs):
        for col, v in zip('CDEFGHIJKLMNOPQ', row):
            if v not in (None, ''):
                P[f'{col}{7 + k}'] = v
    for k, v in enumerate(check or ()):
        if v:
            K[f'D{5 + k}'] = v
    w.save(OUTDIR + '/%s.xlsx' % nome)


def do_exemplo(ex):
    libs = [(l['data'], l['lote'], l['prod'], l['prev'], l['feitas'], l['fora'], l['dec'], l['quem'], l['aut'], l['pnc'], l['obs']) for l in ex['libs']]
    pncs = [(p['num'], p['data'], p['lote'], p['defeito'], p['desc'], p['qtd'], p['det'], p['seg'], p['disp'], p['quem'], p['cliente'], p['rev'], p['enc'], p['custo'], p['acao'])
            for p in ex['pncs']]
    return libs, pncs


AUT = [('Líder da expedição', 'Cada noite', 'Não', 'Fechamento'), ('Gerente', 'Só com o cliente avisado', 'Sim', 'Fechamento')] + [('Pizzaiolo líder', '', 'Sim', 'Registro')] * 6
D = date
for nome, ex, check in (('t1', EX1, ['Sim'] * 12), ('t2', EX2, ['Sim'] * 6 + ['Parcial', 'Não'])):
    libs, pncs = do_exemplo(ex)
    preencher(nome, head=ex['head']['org'], aut=AUT, libs=libs, pncs=pncs, check=check)
# t3: erros propositais
aut3 = [('Analista', '', '', 'Laudo'), ('', '', '', ''), ('Líder', '', '', 'Registro'), ('Coord.', '', '', 'Registro'), ('Coord.', '', 'Não', 'Registro'), ('Coord.', '', '', '')]
libs3 = [(D(2027, 6, 1), 'L1', 'P', 10, 10, 0, LIBERADO, 'Analista', None, None, None),       # OK
         (D(2027, 6, 1), 'L2', 'P', 10, 9, 0, LIBERADO, 'Analista', None, None, None),        # pendente sem autorização
         (D(2027, 6, 1), 'L3', 'P', 10, 9, 0, AUTORIZADO, 'Analista', None, None, None),      # falta quem autorizou
         (D(2027, 6, 1), 'L4', 'P', 10, 10, 1, LIBERADO, 'Analista', None, None, None),       # desvio sem registro
         (D(2027, 6, 1), 'L5', 'P', 10, 10, 2, RETIDO, None, None, '2027-1', None),           # retido OK
         (D(2027, 6, 1), 'L6', 'P', 10, 10, 0, RETIDO, None, None, None, None),               # retido sem registro
         (D(2027, 6, 1), 'L7', 'P', 10, 10, 0, LIBERADO, None, None, None, None),             # falta quem liberou
         (D(2027, 6, 1), 'L8', 'P', 10, None, 0, LIBERADO, 'Analista', None, None, None),     # falta o número de verificações
         (D(2027, 6, 1), 'L9', 'P', 10, 12, 0, LIBERADO, 'Analista', None, None, None),       # feitas acima das previstas
         (D(2027, 6, 1), 'L10', 'P', 10, 10, 0, None, 'Analista', None, None, None),          # falta a decisão
         (D(2027, 6, 1), 'L11', 'P', 10, 9, 0, AUTORIZADO, 'Analista', 'Gerente', None, None)]  # pendente autorizado: OK
pncs3 = [('1', D(2027, 6, 1), 'L', 'Defeito A', 'd', '1', PROC, 'Etiqueta', REFUGO, 'Coord.', None, None, D(2027, 6, 2), 10, None),          # OK, encerrado
         ('2', D(2027, 6, 1), 'L', 'Defeito A', 'd', '1', PROC, 'Etiqueta', REFUGO, 'Coord.', None, None, None, 10, None),                  # em tratamento
         ('3', D(2027, 6, 1), 'L', 'Defeito A', 'd', '1', PROC, 'Etiqueta', None, None, None, None, None, None, None),                       # aberto, falta disposição
         ('4', D(2027, 6, 1), 'L', 'Defeito B', 'd', '1', PROC, None, RETRAB, 'Líder', None, R_CONF, None, 5, None),                         # falta segregar
         ('5', D(2027, 6, 1), 'L', 'Defeito C', 'd', '1', FINAL, 'Etiqueta', CONCES, 'Coord.', None, None, None, 5, None),                   # falta concessão
         ('6', D(2027, 6, 1), 'L', 'Defeito D', 'd', '1', APOS, None, 'Recolher ou substituir', 'Coord.', None, None, None, 100, None),       # falta informar o cliente
         ('7', D(2027, 6, 1), 'L', 'Defeito E', 'd', '1', PROC, 'Etiqueta', RETRAB, 'Líder', None, None, None, 5, None),                     # falta reverificação
         ('8', D(2027, 6, 1), 'L', 'Defeito F', 'd', '1', PROC, 'Etiqueta', RETRAB, 'Líder', None, R_NAO, D(2027, 6, 2), 5, None),           # encerrado com reverificação NC
         ('9', D(2027, 6, 1), 'L', 'Defeito G', 'd', '1', RECEB, 'Etiqueta', DEVOLV, None, None, None, None, 0, None),                       # falta quem decidiu
         ('10', None, 'L', 'Defeito G', 'd', '1', RECEB, 'Etiqueta', DEVOLV, 'G', None, None, None, 0, None),                                # falta a data
         ('11', D(2027, 6, 1), 'L', 'Defeito H', 'd', '1', None, 'Etiqueta', RECLAS, 'G', None, None, None, 0, None),                        # falta onde foi visto
         ('12', D(2027, 6, 1), 'L', None, 'só a descrição', None, None, None, None, None, None, None, None, None, None),                     # falta o defeito
         ('13', D(2027, 6, 1), 'L', 'Defeito I', 'd', '1', PROC, 'Etiqueta', REFUGO, 'Coord.', None, None, D(2027, 6, 2), 1, None),
         ('14', D(2027, 6, 1), 'L', 'Defeito I', 'd', '1', PROC, 'Etiqueta', REFUGO, 'Coord.', None, None, D(2027, 6, 2), 1, None),
         ('15', D(2027, 6, 1), 'L', 'Defeito I', 'd', '1', PROC, 'Etiqueta', REFUGO, 'Coord.', None, None, D(2027, 6, 2), 1, 'RNC 9')]      # repetido com ação: OK
preencher('t3', head='Teste', aut=aut3, libs=libs3, pncs=pncs3)
print('ok')
