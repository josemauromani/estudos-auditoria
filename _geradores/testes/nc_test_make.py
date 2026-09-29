import openpyxl, sys, os
S = os.path.dirname(os.path.abspath(__file__)); sys.path.insert(0, S)
from nc_data import EX2
from datetime import date
SRC, OUTDIR = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(SRC, data_only=True)
err = [(ws.title, c.coordinate, c.value) for ws in wb.worksheets for r in ws.iter_rows() for c in r
       if isinstance(c.value, str) and ((c.value.startswith('#') and len(c.value) > 1) or 'Err:' in c.value)]
print('erros', err)
for n in ('Exemplo 1 - Pizzaria', 'Exemplo 2 - Compras'):
    s = wb[n]
    print(n, [(s[f'B{r}'].value, s[f'E{r}'].value) for r in range(40, 90)
              if s[f'B{r}'].value in ('Correções', 'Ações corretivas', 'Média antes das ações', 'Média depois das ações',
                                      'Períodos seguidos na meta', 'Dias entre a abertura e o encerramento', 'Conclusão')])
wf = openpyxl.load_workbook(SRC)
c = wf['Controle']
rows = [(date(2026, 9, 22), 'Auditoria interna', 'Requisição', 'NC 1', 'Sim', 'Gerente', date(2026, 10, 23), 'Encerrada', date(2027, 2, 5), 'Não'),
        (date(2026, 9, 22), 'Auditoria interna', 'Fornecedores', 'NC 2', 'Sim', 'Comprador', date(2026, 8, 1), 'Ação em implantação', None, 'Sim'),
        (date(2026, 9, 1), 'Reclamação de cliente', 'Expedição', 'NC 3', 'Não', None, date(2027, 5, 1), 'Correção feita', None, 'Não'),
        (date(2026, 7, 1), 'Fornecedor', 'Recebimento', 'NC 4', 'Não', 'Líder', None, 'Encerrada', None, 'Não'),
        (date(2026, 8, 10), 'Reclamação de cliente', 'Produção', 'NC 5', 'Sim', 'Supervisor', None, 'Causa em análise', date(2026, 8, 30), 'Não')]
for k, r in enumerate(rows):
    for col, v in zip('CDEFGHIJKL', r):
        if v is not None:
            c[f'{col}{10+k}'] = v
g = wf['Registro']
g['D4'] = '2026-31'; g['F4'] = date(2026, 9, 22); g['D5'] = 'Auditoria interna'
g['D9'] = EX2['desc']['req']; g['D10'] = EX2['desc']['evid']; g['D11'] = EX2['desc']['decl']
g['C15'] = 'Completar as 4 requisições'; g['E15'] = 'Comprador'; g['F15'] = date(2026, 9, 25); g['G15'] = 'Feita'
g['C16'] = 'Conferir os itens recebidos'; g['E16'] = 'Líder'; g['G16'] = 'Feita'          # sem data
g['D22'] = 'Sim'; g['F22'] = 'Sim'; g['D23'] = '97 de 312'; g['D24'] = 'Atinge um terço'
g['D28'] = 'Campo opcional'; g['F28'] = 'Teste'; g['D29'] = 'Itens sem cadastro'            # sem confirmação
g['F37'] = date(2027, 2, 5)
p = wf['Plano de ação']
for k, a in enumerate(EX2['acoes']):
    r = 9 + k
    for col, key in zip('CDEFGHIJ', ('tipo', 'oque', 'causa', 'resp', 'prazo', 'status', 'feito', 'evid')):
        p[f'{col}{r}'] = a[key] or None
p['E12'] = None                                        # ação corretiva sem causa
p['H13'] = 'Em andamento'; p['G13'] = date(2026, 1, 5); p['I13'] = None   # atrasada
p['J14'] = None                                        # concluída sem evidência
p['H15'] = 'Cancelada'
e = wf['Eficácia']
e['D4'] = EX2['eficacia']['indicador']; e['D6'] = 'Menor é melhor'; e['G6'] = 5; e['D7'] = 3; e['G7'] = 'Não'
for k, (per, mom, val) in enumerate(EX2['eficacia']['medicoes']):
    r = 11 + k; e[f'C{r}'] = per; e[f'D{r}'] = mom; e[f'E{r}'] = val
wf.save(OUTDIR + '/t1.xlsx')


def var(nome, med, **kw):
    w = openpyxl.load_workbook(OUTDIR + '/t1.xlsx'); e = w['Eficácia']
    for r in range(11, 23):
        for col in 'CDE':
            e[f'{col}{r}'] = None
    for k, (per, mom, val) in enumerate(med):
        r = 11 + k; e[f'C{r}'] = per; e[f'D{r}'] = mom; e[f'E{r}'] = val
    for k, v in kw.items():
        e[k] = v
    w.save(OUTDIR + '/%s.xlsx' % nome)


A = [('a', 'Antes', 30), ('b', 'Antes', 32)]
var('t2', A + [('c', 'Depois', 4), ('d', 'Depois', 3)])                                     # em observação (falta 1)
var('t3', A + [('c', 'Depois', 4), ('d', 'Depois', 12), ('e', 'Depois', 20)])                # parcial
var('t4', A + [('c', 'Depois', 31), ('d', 'Depois', 33), ('e', 'Depois', 35)])               # não eficaz
var('t5', A + [('c', 'Depois', 9), ('d', 'Depois', 4), ('e', 'Depois', 3), ('f', 'Depois', 2)])  # eficaz, primeiro fora
var('t6', A + [('c', 'Depois', 4), ('d', 'Depois', 3), ('e', 'Depois', 2)], G7='Sim')        # falha voltou
var('t7', [('a', 'Antes', 80), ('b', 'Antes', 82), ('c', 'Depois', 96), ('d', 'Depois', 97), ('e', 'Depois', 95)], D6='Maior é melhor', G6=95)
var('t8', A + [('c', 'Depois', 4), (None, None, None), ('e', 'Depois', 3)])                  # linha pulada
