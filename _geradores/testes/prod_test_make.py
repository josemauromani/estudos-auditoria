"""Confere Producao-modelo.xlsx e preenche cópias com dados de teste.  Uso: prod_test_make.py <modelo.xlsx> <pasta de saída>"""
import openpyxl, sys, os
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
from prod_data import EX1, EX2, FORA, PROCESSO, PRODUTO, leitura, resumo, resumo_k
from datetime import date
SRC, OUTDIR = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(SRC, data_only=True)
wf0 = openpyxl.load_workbook(SRC)
for ws in wb.worksheets:
    f = wf0[ws.title]
    print('==', ws.title, ws.dimensions, 'dv', len(f.data_validations.dataValidation), 'cf', sum(len(v) for v in f.conditional_formatting._cf_rules.values()),
          'charts', len(f._charts), 'comments', sum(1 for r in f.iter_rows() for c in r if c.comment))
err = [(ws.title, c.coordinate, c.value) for ws in wb.worksheets for r in ws.iter_rows() for c in r
       if isinstance(c.value, str) and ((c.value.startswith('#') and len(c.value) > 1) or 'Err:' in c.value)]
print('erros', err)
print('vazio:', wb['Plano de controle']['D35'].value, '|', wb['Registros']['E163'].value, '|', wb['Rastreabilidade']['D34'].value, '|', wb['Mudanças']['E23'].value, '|',
      wb['Painel']['D40'].value, '|', wb['Checklist']['D24'].value)
for aba, ex in (('Exemplo 1 - Pizzaria', EX1), ('Exemplo 2 - Indústria', EX2)):
    s = wb[aba]
    ids = [k['id'] for k in ex['plano']]
    ks = {k['id']: k for k in ex['plano']}
    regs = [r for r in range(1, s.max_row + 1) if s[f'E{r}'].value in ids and s[f'I{r}'].value in ('Dentro', 'Fora')]
    got = [s[f'I{r}'].value for r in regs]
    esp = [leitura(ks[r['k']], r) for r in ex['regs']]
    res = [r for r in range(1, s.max_row + 1) if s[f'B{r}'].value in ids and isinstance(s[f'E{r}'].value, (int, float)) and s[f'C{r}'].value == ks[s[f'B{r}'].value]['caract']]
    got_k = [(s[f'E{r}'].value, s[f'F{r}'].value, s[f'G{r}'].value) for r in res]
    esp_k = [(x['n'], x['fora'], x['semr']) for x in resumo_k(ex)]
    print(aba, '| leituras', 'OK' if got == esp else 'DIFERE %s' % got, '| resumo por controle', 'OK' if got_k == esp_k else 'DIFERE %s' % got_k)


def preencher(nome, head=(), plano=(), regs=(), rast=(), prop=(), mud=(), check=None):
    w = openpyxl.load_workbook(SRC)
    P, R, T, M, K = w['Plano de controle'], w['Registros'], w['Rastreabilidade'], w['Mudanças'], w['Checklist']
    for k, v in enumerate(head):
        P[f'D{4 + k}'] = v
    for k, row in enumerate(plano):
        for col, v in zip('CDEFGHIJKLM', row):
            if v not in (None, ''):
                P[f'{col}{10 + k}'] = v
    for k, row in enumerate(regs):
        for col, v in zip('CDEJKMN', row):
            if v not in (None, ''):
                R[f'{col}{7 + k}'] = v
    for k, row in enumerate(rast):
        for col, v in zip('CDEFG', row):
            if v:
                T[f'{col}{7 + k}'] = v
    for k, row in enumerate(prop):
        for col, v in zip('CDEFG', row):
            if v:
                T[f'{col}{21 + k}'] = v
    for k, row in enumerate(mud):
        for col, v in zip('CDEFGH', row):
            if v:
                M[f'{col}{7 + k}'] = v
    for k, v in enumerate(check or ()):
        if v:
            K[f'D{5 + k}'] = v
    w.save(OUTDIR + '/%s.xlsx' % nome)


def do_exemplo(ex):
    plano = [(k['etapa'], k['caract'], k['tipo'], k['min'], k['max'], k['unid'], k['metodo'], k['freq'], k['resp'], k['registro'], k['reacao']) for k in ex['plano']]
    regs = [(r['data'], r['lote'], r['k'], r['valor'], r['conf'], r['reacao'], r['quem']) for r in ex['regs']]
    prop = [(a, b, c, d, e) for a, b, c, d, e in ex['prop']]
    return plano, regs, ex['rast'], prop, ex['mud']


D = date
for nome, ex, check in (('t1', EX1, ['Sim'] * 12), ('t2', EX2, ['Sim'] * 6 + ['Parcial', 'Não'])):
    plano, regs, rast, prop, mud = do_exemplo(ex)
    h = ex['head']
    preencher(nome, head=[h['org'] + ' · ' + h['processo'], h['produto'], h['rev']], plano=plano, regs=regs, rast=rast, prop=prop, mud=mud, check=check)
# t3: erros propositais
preencher('t3', head=['Teste'],
          plano=[('Etapa 1', 'Faixa completa', PRODUTO, 10, 20, 'mm', 'Régua', 'Cada peça', 'Operador', 'Folha', 'Segregar'),
                 ('Etapa 1', 'Mínimo maior que o máximo', PRODUTO, 30, 20, 'mm', 'Régua', 'Cada peça', 'Operador', 'Folha', 'Segregar'),
                 (None, 'Sem etapa nem reação', PRODUTO, None, 5, 'g', 'Balança', 'Cada lote', 'Operador', 'Folha', None),
                 ('Etapa 2', 'Atributo sem tipo', None, None, None, None, 'Visual', 'Cada peça', 'Operador', 'Folha', 'Refazer'),
                 ('Etapa 2', 'Só mínimo, sem registro no período', PRODUTO, 65, None, '°C', 'Termômetro', 'Por hora', 'Líder', 'Planilha', 'Reaquecer')],
          regs=[(D(2027, 6, 1), 'L1', 'K1', 15, None, None, 'Operador'),            # dentro
                (D(2027, 6, 1), 'L1', 'K1', 25, None, None, 'Operador'),            # fora sem reação
                (D(2027, 6, 2), 'L2', 'K1', 9.5, None, 'Segregado', 'Operador'),    # fora com reação: desvio repetido
                (D(2027, 6, 2), 'L2', 'K1', None, None, None, 'Operador'),          # falta o valor
                (D(2027, 6, 2), 'L2', 'K3', 5, None, None, 'Operador'),             # no limite: dentro
                (D(2027, 6, 2), 'L2', 'K3', 5.1, None, 'Lote repesado', None),      # fora com reação, falta quem mediu
                (D(2027, 6, 3), 'L3', 'K4', None, 'Sim', None, 'Operador'),         # atributo dentro
                (D(2027, 6, 3), 'L3', 'K4', None, 'Não', 'Refeito', 'Operador'),    # atributo fora com reação
                (D(2027, 6, 3), 'L3', 'K4', None, None, None, 'Operador'),          # falta o resultado
                (None, None, 'K1', 12, None, None, 'Operador'),                     # falta a data
                (D(2027, 6, 3), None, 'K1', 12, None, None, 'Operador'),            # falta o lote
                (D(2027, 6, 3), 'L3', 'K9', 12, None, None, 'Operador'),            # controle sem plano
                (D(2027, 6, 3), 'L3', None, 12, None, None, 'Operador')],           # falta o controle
          rast=[('Etapa 1', 'Etiqueta', 'Carimbo', 'Folha', 'Palete'), ('Etapa 2', None, 'Carimbo', None, None), ('Etapa 3', 'Etiqueta', None, None, None), (None, 'Sem etapa', None, None, None)],
          prop=[('Molde', 'Cliente', 'Gravado', 'Estante', None), ('Palete', None, 'Marca', 'Saldo', None), ('Arte', 'Cliente', None, 'Pasta', None), ('Dado', 'Cliente', 'Cadastro', None, 'Perda')],
          mud=[(D(2027, 6, 1), 'Completa', 'Motivo', 'Teste', 'Gerente', 'Ações'), (D(2027, 6, 1), 'Sem análise', 'Motivo', None, 'Gerente', None),
               (D(2027, 6, 1), 'Sem autorização', 'Motivo', 'Teste', None, None), (None, 'Sem data', 'Motivo', 'Teste', 'Gerente', None), (D(2027, 6, 1), None, 'Sem o que mudou', None, None, None)])
print('ok')
