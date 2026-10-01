"""Confere Escopo-modelo.xlsx e preenche cópias com dados de teste.  Uso: esc_test_make.py <modelo.xlsx> <pasta de saída>"""
import openpyxl, sys, os
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
from esc_data import CAMPOS, DEM, EX1, EX2, NAO, NDEM, PARC, REQS, SIM, aplicabilidade, lider_conf
from datetime import date
SRC, OUTDIR = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(SRC, data_only=True)
wf0 = openpyxl.load_workbook(SRC)
for ws in wb.worksheets:
    f = wf0[ws.title]
    print('==', ws.title, ws.dimensions, 'dv', len(f.data_validations.dataValidation), 'cf', sum(len(v) for v in f.conditional_formatting._cf_rules.values()),
          'charts', len(f._charts), 'comments', sum(1 for r in f.iter_rows() for c in r if c.comment))
err = [(ws.title, c.coordinate, c.value) for ws in wb.worksheets for r in ws.iter_rows() for c in r
       if isinstance(c.value, str) and ((c.value.startswith('#') and len(c.value) > 1) or 'Err:' in c.value)]
print('erros', err)
print('vazio:', wb['Escopo']['C17'].value, '|', wb['Aplicabilidade']['D56'].value, '|', wb['Liderança']['D20'].value, '|', wb['Painel']['C31'].value)
for aba, ex in (('Exemplo 1 - Pizzaria', EX1), ('Exemplo 2 - Indústria', EX2)):
    s = wb[aba]
    excl = [a for a in aplicabilidade(ex) if a['aplica'] == NAO]
    ra = [r for r in range(1, s.max_row + 1) if s[f'B{r}'].value in [a['num'] for a in excl] and s[f'F{r}'].value == NAO]
    rl = [r for r in range(1, s.max_row + 1) if isinstance(s[f'B{r}'].value, int) and s[f'G{r}'].value in (DEM, PARC, NDEM)]
    ok = ([s[f'J{r}'].value for r in ra] == [a['conf'] for a in excl], [s[f'K{r}'].value for r in rl] == [lider_conf(c) for c in ex['lid']])
    print(aba, '| aplicabilidade', 'OK' if ok[0] else 'DIFERE %s' % [s[f'J{r}'].value for r in ra], '| liderança', 'OK' if ok[1] else 'DIFERE %s' % [s[f'K{r}'].value for r in rl])


def preencher(nome, org=None, escopo=None, aplic=None, lid=(), check=None):
    w = openpyxl.load_workbook(SRC)
    E, A, L, K = w['Escopo'], w['Aplicabilidade'], w['Liderança'], w['Checklist']
    if org:
        E['C4'] = org
    for k, (chave, _) in enumerate(CAMPOS):
        v = (escopo or {}).get(chave)
        if v:
            E[f'C{5 + k}'] = v
    for k, (num, _) in enumerate(REQS):
        row = (aplic or {}).get(num)
        if row:
            for col, v in zip('EFG', row):
                if v:
                    A[f'{col}{7 + k}'] = v
    for k, row in enumerate(lid):
        for col, v in zip('DEFGHIJ', row):
            if v not in (None, ''):
                L[f'{col}{6 + k}'] = v
    for k, v in enumerate(check or ()):
        if v:
            K[f'D{5 + k}'] = v
    w.save(OUTDIR + '/%s.xlsx' % nome)


def do_exemplo(ex):
    ap = {a['num']: (a['aplica'], a['just'], a['afeta']) for a in aplicabilidade(ex)}
    lid = [(c['faz'], c['freq'], c['reg'], c['sit'], c['acao'], c['resp'], c['prazo']) for c in ex['lid']]
    return ex['head']['org'], ex['escopo'], ap, lid


for nome, ex, check in (('t1', EX1, ['Sim'] * 12), ('t2', EX2, ['Sim'] * 6 + ['Parcial', 'Não'])):
    preencher(nome, *do_exemplo(ex), check=check)
D = date
aplic3 = {n: (SIM, '', '') for n, _ in REQS}
aplic3.update({'4.3': ('', '', ''), '9.3': (NAO, 'O diretor acompanha tudo', NAO), '8.3': (NAO, '', ''), '8.5.3': (NAO, 'Nada do cliente', ''), '8.5.5': (NAO, 'Sem pós-entrega', SIM),
               '7.1.5': (NAO, 'Sem medição', NAO)})
lid3 = [('Faz', 'Mensal', 'Ata', DEM, '', '', None), ('', '', '', '', '', '', None), ('Faz', '', '', '', '', '', None), ('Faz', '', '', PARC, '', '', None),
        ('Faz', '', '', NDEM, 'Ação', '', None), ('Faz', '', '', PARC, 'Ação', 'Quem', D(2027, 6, 1)), ('Faz', '', '', DEM, '', '', None)]
preencher('t3', 'Teste', dict(unidades='U', produtos='P'), aplic3, lid3)
print('ok')
