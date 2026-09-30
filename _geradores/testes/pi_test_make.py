"""Confere Partes-Interessadas-modelo.xlsx e preenche cópias com dados de teste.  Uso: pi_test_make.py <modelo.xlsx> <pasta de saída>"""
import openpyxl, sys, os
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
from pi_data import EX1, EX2, atendimento, estrategia, pertinente
from datetime import date
SRC, OUTDIR = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(SRC, data_only=True)
wf0 = openpyxl.load_workbook(SRC)
for ws in wb.worksheets:
    f = wf0[ws.title]
    print('==', ws.title, ws.dimensions, 'dv', len(f.data_validations.dataValidation), 'cf', sum(len(v) for v in f.conditional_formatting._cf_rules.values()),
          'charts', len(f._charts), 'comments', sum(1 for r in f.iter_rows() for c in r if c.comment))
err = [(ws.title, c.coordinate, c.value) for ws in wb.worksheets for r in ws.iter_rows() for c in r
       if isinstance(c.value, str) and ((c.value.startswith('#') and len(c.value) > 1) or 'Err:' in c.value)]
print('erros', err)
print('vazio:', wb['Partes']['E28'].value, '|', wb['Requisitos']['E54'].value, '|', wb['Matriz']['D26'].value, '|', wb['Painel']['D33'].value, '|', wb['Checklist']['D24'].value)
for aba, ex in (('Exemplo 1 - Pizzaria', EX1), ('Exemplo 2 - Indústria', EX2)):
    s = wb[aba]
    nomes = {p['nome']: p for p in ex['partes']}
    got = [(s[f'I{r}'].value, s[f'L{r}'].value) for r in range(1, s.max_row + 1) if s[f'C{r}'].value in nomes and isinstance(s[f'E{r}'].value, int)]
    esp = [(estrategia(p['inf'], p['int']), 'Sim' if pertinente(p) else 'Não') for p in ex['partes']]
    at = next(s[f'E{r}'].value for r in range(1, s.max_row + 1) if s[f'B{r}'].value == 'Atendimento dos requisitos adotados')
    print(aba, '| estratégias', 'OK' if got == esp else 'DIFERE %s' % got, '| atendimento', 'OK' if abs(at - atendimento(ex)) < 1e-9 else 'DIFERE %s' % at)


def preencher(nome, head=None, partes=(), reqs=(), data=None, check=None):
    w = openpyxl.load_workbook(SRC)
    P, R, Pa, K = w['Partes'], w['Requisitos'], w['Painel'], w['Checklist']
    for ref, v in (head or {}).items():
        P[ref] = v
    for k, row in enumerate(partes):
        for col, v in zip('CDEFGHKL', row):
            if v not in (None, ''):
                P[f'{col}{11 + k}'] = v
    for k, row in enumerate(reqs):
        for col, v in zip('CDEFGHIJKL', row):
            if v not in (None, ''):
                R[f'{col}{7 + k}'] = v
    if data:
        Pa['D4'] = data
    for k, v in enumerate(check or ()):
        if v:
            K[f'D{5 + k}'] = v
    w.save(OUTDIR + '/%s.xlsx' % nome)


def do_exemplo(ex):
    partes = [(p['nome'], p['grupo'], p['porque'], p['inf'], p['int'], 'Sim' if p['legal'] else 'Não', p['canal'], p['quem']) for p in ex['partes']]
    reqs = [(q['parte'], q['req'], q['tipo'], q['adotado'], q['como'], q['monit'], q['sit'], q['acao'], q['resp'], q['prazo']) for q in ex['reqs']]
    return partes, reqs


D = date
# t1: pizzaria, na data da análise; t2: indústria, em data posterior, com ações vencidas
p, q = do_exemplo(EX1)
preencher('t1', head={'E4': EX1['head']['org'], 'E5': EX1['head']['por'], 'E6': EX1['head']['escopo']}, partes=p, reqs=q, data=EX1['head']['data'], check=['Sim'] * 12)
p, q = do_exemplo(EX2)
preencher('t2', head={'E4': EX2['head']['org']}, partes=p, reqs=q, data=D(2027, 4, 15), check=['Sim'] * 6 + ['Parcial', 'Não'])
# t3: erros propositais
preencher('t3', partes=[('A', 'Cliente', 'motivo', 5, 5, 'Não', 'canal', 'quem'), ('B', 'Regulador', 'motivo', 1, 1, 'Sim', 'canal', 'quem'), ('A', 'Cliente', 'motivo', 4, 4, 'Não', 'canal', 'quem'),
                        ('C', 'Outro', 'motivo', 5, None, 'Não', None, None), ('D', 'Outro', None, 2, 2, 'Não', None, None), ('E', 'Sociedade', 'motivo', 2, 5, 'Não', None, None),
                        ('F', 'Sociedade', 'motivo', 4, 2, 'Não', 'canal', 'quem'), ('G', 'Outro', 'motivo', 3, 3, 'Não', None, None)],
          reqs=[('A', 'completo', 'Legal', 'Sim', 'como', 'monit', 'Atende', None, None, None),
                (None, 'sem parte', 'Legal', 'Sim', 'como', 'monit', 'Atende', None, None, None),
                ('A', 'sem tipo', None, 'Sim', 'como', 'monit', 'Atende', None, None, None),
                ('A', 'sem decisão', 'Expectativa', None, None, None, None, None, None, None),
                ('B', 'legal recusado', 'Legal', 'Não', 'porque', None, None, None, None, None),
                ('A', 'recusa sem justificativa', 'Expectativa', 'Não', None, None, None, None, None, None),
                ('A', 'recusa justificada', 'Expectativa', 'Não', 'porque', None, None, None, None, None),
                ('A', 'sem como', 'Contratual', 'Sim', None, 'monit', 'Atende', None, None, None),
                ('A', 'sem monitoramento', 'Contratual', 'Sim', 'como', None, 'Atende', None, None, None),
                ('A', 'sem situação', 'Contratual', 'Sim', 'como', 'monit', None, None, None, None),
                ('B', 'sem ação', 'Legal', 'Sim', 'como', 'monit', 'Não atende', None, None, None),
                ('B', 'ação sem prazo', 'Legal', 'Sim', 'como', 'monit', 'Atende em parte', 'ação', 'quem', None),
                ('B', 'ação vencida', 'Legal', 'Sim', 'como', 'monit', 'Atende em parte', 'ação', 'quem', D(2027, 1, 5))],
          data=D(2027, 2, 1))
print('ok')
