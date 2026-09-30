"""Confere Competencias-modelo.xlsx e preenche cópias com dados de teste.  Uso: comp_test_make.py <modelo.xlsx> <pasta de saída>"""
import openpyxl, sys, os
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
from comp_data import EX1, EX2, lacunas, situacao_acao, cobertura
from datetime import date
from openpyxl.utils import get_column_letter as L
SRC, OUTDIR = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(SRC, data_only=True)
wf0 = openpyxl.load_workbook(SRC)
for ws in wb.worksheets:
    f = wf0[ws.title]
    print('==', ws.title, ws.dimensions, 'dv', len(f.data_validations.dataValidation), 'cf', sum(len(v) for v in f.conditional_formatting._cf_rules.values()),
          'charts', len(f._charts), 'comments', sum(1 for r in f.iter_rows() for c in r if c.comment))
err = [(ws.title, c.coordinate, c.value) for ws in wb.worksheets for r in ws.iter_rows() for c in r
       if isinstance(c.value, str) and ((c.value.startswith('#') and len(c.value) > 1) or 'Err:' in c.value)]
print('erros', err)
print('vazio:', wb['Funções']['E24'].value, '|', wb['Matriz']['E43'].value, '|', wb['Plano']['E40'].value, '|', wb['Registros']['E37'].value, '|', wb['Checklist']['D24'].value)
for nome, ex in (('Exemplo 1 - Pizzaria', EX1), ('Exemplo 2 - Compras', EX2)):
    s = wb[nome]
    nc = len(ex['comps'])
    rows = [r for r in range(1, s.max_row + 1) if s[f'B{r}'].value in [p['nome'] for p in ex['pessoas']] and s[f'C{r}'].value in ex['funcoes']]
    lac = [s[f'{L(4 + nc)}{r}'].value for r in rows]
    prow = [r for r in range(1, s.max_row + 1) if s[f'B{r}'].value in [p['nome'] for p in ex['pessoas']] and s[f'C{r}'].value in [c[0] for c in ex['comps']]]
    sit = [s[f'{L(7 + nc)}{r}'].value for r in prow]
    cob = [s[f'{L(4 + k)}{rows[-1] + 1}'].value for k in range(nc)]
    print(nome, '| lacunas', 'OK' if lac == [len(lacunas(ex, p)) for p in ex['pessoas']] else 'DIFERE %s' % lac,
          '| plano', 'OK' if sit == [situacao_acao(a, ex['head']['data']) for a in ex['plano']] else 'DIFERE %s' % sit,
          '| cobertura', 'OK' if cob == [cobertura(ex, k) for k in range(nc)] else 'DIFERE %s' % cob)


def preencher(nome, ex=None, cab=None, comps=None, funcoes=None, pessoas=None, plano=None, regs=(), check=None):
    w = openpyxl.load_workbook(SRC)
    F, M, P, R, C = w['Funções'], w['Matriz'], w['Plano'], w['Registros'], w['Checklist']
    if ex:
        cab = dict({'D4': ex['head']['org'], 'D5': ex['head']['data']}, **(cab or {}))
        comps = comps or [c[1] for c in ex['comps']]
        funcoes = funcoes or list(ex['funcoes'].items())
        pessoas = pessoas or [(p['nome'], p['funcao'], p['niveis']) for p in ex['pessoas']]
        plano = plano if plano is not None else [(a['pessoa'], a['comp'], a['acao'], a['quem'], a['prazo'], a['feito'], a['aprend'], a['eficacia'], None, a['obs'] or None)
                                                 for a in ex['plano']]
    for ref, v in (cab or {}).items():
        if v is not None:
            F[ref] = v
    for k, c in enumerate(comps or []):
        F.cell(row=8, column=4 + k, value=c)
    for k, (fn, req) in enumerate(funcoes or []):
        F[f'C{10 + k}'] = fn
        for j, v in enumerate(req):
            F.cell(row=10 + k, column=4 + j, value=v)
    for k, (nome_, fn, nv) in enumerate(pessoas or []):
        M[f'C{10 + k}'] = nome_
        if fn is not None:
            M[f'D{10 + k}'] = fn
        for j, v in enumerate(nv):
            if v is not None:
                M.cell(row=10 + k, column=5 + j, value=v)
    for k, row in enumerate(plano or []):
        for col, v in zip('CDEFGHIJKL', row):
            if v is not None:
                P[f'{col}{7 + k}'] = v
    for k, row in enumerate(regs):
        for col, v in zip('CDEFGHI', row):
            if v is not None:
                R[f'{col}{7 + k}'] = v
    for k, v in enumerate(check or ()):
        if v:
            C[f'D{5 + k}'] = v
    w.save(OUTDIR + '/%s.xlsx' % nome)


D = date
# t1: pizzaria completa
preencher('t1', EX1, regs=[(D(2026, 10, 28), 'IT-EXP-01', 'Sérgio', 2, 1, 'Sim', 'Lista de presença'), (D(2027, 1, 20), 'Regulagem do forno', 'Rafael', 8, 1, 'Sim', 'Avaliação no posto'),
                           (D(2027, 2, 10), 'Boas práticas', 'Marina', 3, 4, None, 'Lista de presença'), (D(2027, 2, 5), 'Receitas', 'Rafael', 4, 1, 'Sim', None)],
          check=['Sim'] * 12)
# t2: compras, sem data de referência (usa hoje)
preencher('t2', EX2, cab={'D5': None}, check=['Sim'] * 6 + ['Parcial', 'Não'])
# t3: erros propositais
preencher('t3', cab={'D4': 'Teste', 'D5': D(2027, 3, 12)}, comps=['A', 'B', 'C'], funcoes=[('Operador', [2, 1, 0]), ('Líder', [3, 2, 2]), ('Sem exigência', [0, 0, 0])],
          pessoas=[('Ana', 'Operador', [2, 1, 0]), ('Beto', 'Operador', [1, 0, None]), ('Cida', 'Líder', [3, 2, 1]), ('Davi', None, [2, 2, 2]), ('Eva', 'Gerente', [3, 3, 3]),
                   ('Fabi', 'Sem exigência', [0, 0, 0])],
          plano=[('Beto', 'C1', 'Sem prazo', 'Ana', None, None, None, None, None, None),
                 ('Beto', 'C2', 'Atrasada', 'Ana', D(2027, 3, 1), None, None, None, None, None),
                 ('Beto', 'C2', 'Planejada', 'Ana', D(2027, 4, 1), None, None, None, None, None),
                 ('Cida', 'C3', 'Feita sem aprendizado', 'Ana', D(2027, 1, 1), D(2027, 1, 5), None, None, None, None),
                 ('Cida', 'C3', 'Aprendizado não', 'Ana', D(2027, 1, 1), D(2027, 1, 5), 'Não', None, None, None),
                 ('Cida', 'C3', 'Em acompanhamento', 'Ana', D(2027, 2, 1), D(2027, 2, 20), 'Sim', None, None, None),
                 ('Cida', 'C3', 'A avaliar', 'Ana', D(2027, 1, 1), D(2027, 1, 5), 'Sim', None, None, None),
                 ('Cida', 'C3', 'Eficaz', 'Ana', D(2027, 1, 1), D(2027, 1, 5), 'Sim', 'Eficaz', D(2027, 3, 10), None),
                 ('Cida', 'C3', 'Parcial', 'Ana', D(2027, 1, 1), D(2027, 1, 5), 'Sim', 'Parcial', D(2027, 3, 10), None),
                 ('Cida', 'C3', 'Não eficaz', 'Ana', D(2027, 1, 1), D(2027, 1, 5), 'Sim', 'Não eficaz', D(2027, 3, 10), None)],
          regs=[(None, 'Sem data', 'X', 1, 1, 'Sim', 'Lista'), (D(2027, 1, 1), 'Sem evidência', 'X', 1, 1, 'Sim', None), (D(2027, 1, 1), 'Sem avaliação', 'X', 1, 1, None, 'Lista'),
                (D(2027, 1, 1), 'Aprendizado não', 'X', 1, 1, 'Não', 'Lista'), (D(2027, 1, 1), 'Completo', 'X', 2, 3, 'Sim', 'Lista')])
print('ok')
