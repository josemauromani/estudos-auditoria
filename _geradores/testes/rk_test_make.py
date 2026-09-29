import openpyxl, sys, os
S = os.path.dirname(os.path.abspath(__file__)); sys.path.insert(0, S)
from risk_data import EX2, OPORT, ACEITE, ACEITAR
from datetime import date
SRC, OUTDIR = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(SRC, data_only=True)
wf0 = openpyxl.load_workbook(SRC)
for ws in wb.worksheets:
    f = wf0[ws.title]
    print('==', ws.title, ws.dimensions, 'dv', len(f.data_validations.dataValidation), 'cf', sum(len(v) for v in f.conditional_formatting._cf_rules.values()),
          'charts', len(f._charts), 'comments', sum(1 for r in f.iter_rows() for c in r if c.comment))
err = [(ws.title, c.coordinate, c.value) for ws in wb.worksheets for r in ws.iter_rows() for c in r
       if isinstance(c.value, str) and ((c.value.startswith('#') and len(c.value) > 1) or 'Err:' in c.value)]
print('erros', err)
print('riscos vazio', [wb['Riscos'][f'G{k}'].value for k in range(33, 39)]); print('mapa vazio', [wb['Mapa'][f'E{k}'].value for k in range(21, 27)])
print('oport vazio', [wb['Oportunidades'][f'E{k}'].value for k in range(19, 24)])
for n in ('Exemplo 1 - Pizzaria', 'Exemplo 2 - Compras'):
    s = wb[n]
    print(n, [(s[f'B{k}'].value, s[f'E{k}'].value, s[f'F{k}'].value) for k in range(18, 45) if s[f'B{k}'].value in ('Baixo', 'Médio', 'Alto', 'Crítico')],
          [(s[f'B{k}'].value, s[f'E{k}'].value) for k in range(18, 45) if isinstance(s[f'B{k}'].value, str) and s[f'B{k}'].value.startswith(('Riscos', 'Altos', 'Média'))])

w = openpyxl.load_workbook(SRC)
r = w['Riscos']
r['D4'] = 'Indústria'; r['K4'] = 'Adquirir materiais'; r['T4'] = date(2026, 10, 12)
for k, x in enumerate(EX2['riscos']):
    rr = 11 + k
    vals = dict(C=x['processo'], D=x['causa'], E=x['evento'], F=x['conseq'], G=x['p'], H=x['i'], K=x['resp'],
                L=ACEITE[x['id']] if x['resp'] == ACEITAR else x['acao'], M=x['quem'], N=x['prazo'], O=x['status'] or None, P=x['pr'], Q=x['ir'])
    for c, v in vals.items():
        if v is not None:
            r[f'{c}{rr}'] = v
r['N12'] = date(2026, 1, 15)                       # C2 atrasada
extra = [
    dict(E='só o evento'),
    dict(D='c', E='sem notas', F='f'),
    dict(D='c', E='sem resposta', F='f', G=3, H=3),
    dict(D='c', E='reduzir sem ação', F='f', G=3, H=3, K='Reduzir'),
    dict(D='c', E='sem responsável', F='f', G=3, H=3, K='Reduzir', L='ação'),
    dict(D='c', E='sem residual', F='f', G=3, H=3, K='Reduzir', L='ação', M='x', N=date(2027, 5, 1)),
    dict(D='c', E='residual maior', F='f', G=2, H=2, K='Reduzir', L='ação', M='x', N=date(2027, 5, 1), P=3, Q=3),
    dict(D='c', E='alto aceito', F='f', G=4, H=4, K='Aceitar', L='justificativa', P=4, Q=4),
    dict(D='c', E='residual alto', F='f', G=4, H=5, K='Reduzir', L='ação', M='x', N=date(2027, 5, 1), O='Concluída', P=3, Q=4),
]
for k, d in enumerate(extra):
    for c, v in d.items():
        r[f'{c}{19 + k}'] = v
o = w['Oportunidades']
for k, (i, txt, p, b, dec, acao, quem) in enumerate(OPORT):
    rr = 7 + k
    o[f'C{rr}'] = txt; o[f'E{rr}'] = p; o[f'F{rr}'] = b; o[f'I{rr}'] = dec; o[f'J{rr}'] = acao; o[f'K{rr}'] = quem
    if dec != 'Não agir agora':
        o[f'L{rr}'] = date(2027, 3, 1); o[f'M{rr}'] = 'Em andamento'
o['C11'] = 'alta sem ação'; o['E11'] = 5; o['F11'] = 4; o['I11'] = 'Não agir agora'
w.save(OUTDIR + '/t1.xlsx')
# só o exemplo 2, sem erros
w = openpyxl.load_workbook(OUTDIR + '/t1.xlsx'); r = w['Riscos']
for rr in range(19, 28):
    for c in 'DEFGHKLMNOPQ':
        r[f'{c}{rr}'] = None
r['N12'] = date(2026, 10, 23); r['T5'] = date(2027, 4, 12)
w['Oportunidades']['C11'] = None; w['Oportunidades']['E11'] = None; w['Oportunidades']['F11'] = None; w['Oportunidades']['I11'] = None
w.save(OUTDIR + '/t2.xlsx')
