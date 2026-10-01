"""Confere Objetivos-modelo.xlsx e preenche cópias com dados de teste.  Uso: obj_test_make.py <modelo.xlsx> <pasta de saída>"""
import openpyxl, sys, os
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
from obj_data import EX1, EX2, MAIOR, caminho, situacao
from datetime import date
SRC, OUTDIR = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(SRC, data_only=True)
wf0 = openpyxl.load_workbook(SRC)
for ws in wb.worksheets:
    f = wf0[ws.title]
    print('==', ws.title, ws.dimensions, 'dv', len(f.data_validations.dataValidation), 'cf', sum(len(v) for v in f.conditional_formatting._cf_rules.values()),
          'charts', len(f._charts), 'comments', sum(1 for r in f.iter_rows() for c in r if c.comment))
err = [(ws.title, c.coordinate, c.value) for ws in wb.worksheets for r in ws.iter_rows() for c in r
       if isinstance(c.value, str) and ((c.value.startswith('#') and len(c.value) > 1) or 'Err:' in c.value)]
print('erros', err)
print('vazio:', wb['Objetivos']['D31'].value, '|', wb['Planos']['E48'].value, '|', wb['Acompanhamento']['D23'].value, '|', wb['Painel']['D33'].value, '|', wb['Checklist']['D24'].value)
for aba, ex in (('Exemplo 1 - Pizzaria', EX1), ('Exemplo 2 - Indústria', EX2)):
    s = wb[aba]
    ref = ex['head']['ref']
    n = ex['head']['meses']
    cc, cs = openpyxl.utils.get_column_letter(5 + n), openpyxl.utils.get_column_letter(6 + n)
    rows = [r for r in range(1, s.max_row + 1) if s[f'B{r}'].value in [o['id'] for o in ex['objs']] and s[f'C{r}'].value in [o['ind'] for o in ex['objs']]]
    got = [(s[f'{cs}{r}'].value, None if s[f'{cc}{r}'].value in ('', None) else round(s[f'{cc}{r}'].value, 4)) for r in rows]
    esp = [(situacao(o, ref), None if caminho(o) is None else round(caminho(o), 4)) for o in ex['objs']]
    print(aba, '| acompanhamento', 'OK' if got == esp else 'DIFERE %s' % got)


def preencher(nome, head=None, comps=(), objs=(), acoes=(), meses=(), res=(), ref=None, check=None):
    w = openpyxl.load_workbook(SRC)
    O, P, A, K = w['Objetivos'], w['Planos'], w['Acompanhamento'], w['Checklist']
    for ref_, v in (head or {}).items():
        O[ref_] = v
    for k, c in enumerate(comps):
        O[f'D{7 + k}'] = c
    for k, row in enumerate(objs):
        for col, v in zip('CDEFGHIJKLMN', row):
            if v is not None:
                O[f'{col}{14 + k}'] = v
    for k, row in enumerate(acoes):
        for col, v in zip('CEFGHIJK', row):
            if v is not None:
                P[f'{col}{7 + k}'] = v
    for j, m in enumerate(meses):
        A.cell(row=6, column=8 + j, value=m)
    for k, row in enumerate(res):
        for j, v in enumerate(row):
            if v is not None:
                A.cell(row=7 + k, column=8 + j, value=v)
    if ref:
        A['D4'] = ref
    for k, v in enumerate(check or ()):
        if v:
            K[f'D{5 + k}'] = v
    w.save(OUTDIR + '/%s.xlsx' % nome)


def do_exemplo(ex):
    objs = [(o['texto'], o['comp'], o['origem'], o['ind'], o['unid'], o['sentido'], o['base'], o['meta'], o['prazo'], o['processo'], o['resp'], o['freq']) for o in ex['objs']]
    acoes = [(o['id'], a['oque'], a['rec'] or None, a['resp'], a['prazo'], a['status'], a['feito'], a['avalia']) for o in ex['objs'] for a in o['acoes']]
    return objs, acoes, [o['res'] for o in ex['objs']]


D = date
for nome, ex, check in (('t1', EX1, ['Sim'] * 12), ('t2', EX2, ['Sim'] * 6 + ['Parcial', 'Não'])):
    objs, acoes, res = do_exemplo(ex)
    preencher(nome, head={'D4': ex['head']['org'], 'D5': ex['head']['politica']}, comps=[t for _, t in ex['comps']], objs=objs, acoes=acoes, res=res, ref=ex['head']['ref'], check=check)
# t3: erros propositais: objetivo sem meta, sem prazo, sem ação; ação com prazo depois do objetivo; prazo vencido; resultado com lacuna
preencher('t3', head={'D5': 'Política'}, comps=['C1'],
          objs=[('Sem meta', 'C1', 'Análise crítica', 'Ind A', '%', MAIOR, 10, None, D(2027, 12, 31), 'P', 'Quem', 'Mensal'),
                ('Sem prazo nem ação', 'C1', 'Análise crítica', 'Ind B', '%', MAIOR, 10, 20, None, 'P', 'Quem', 'Mensal'),
                ('Vencido', 'C1', 'Análise crítica', 'Ind C', '%', MAIOR, 10, 20, D(2027, 3, 31), 'P', 'Quem', 'Mensal'),
                ('Manter, com lacuna', 'C1', 'Análise crítica', 'Ind D', '%', 'Menor é melhor', 5, 5, D(2027, 12, 31), 'P', 'Quem', 'Mensal'),
                ('Sem base', None, 'Análise crítica', 'Ind E', '%', MAIOR, None, 8, D(2027, 12, 31), 'P', None, 'Mensal')],
          acoes=[('O1', 'ok', None, 'Quem', D(2027, 6, 30), 'Em andamento', None, 'medida'), ('O3', 'prazo depois do objetivo', None, 'Quem', D(2027, 6, 30), 'Não iniciada', None, 'medida'),
                 ('O4', 'concluída sem data', None, 'Quem', D(2027, 2, 28), 'Concluída', None, 'medida'), ('O9', 'código sem objetivo', None, 'Quem', D(2027, 2, 28), 'Não iniciada', None, 'medida'),
                 ('O5', 'sem responsável', None, None, D(2027, 2, 28), 'Não iniciada', None, 'medida')],
          res=[[12, 14], [10, 9], [12, 15, 16], [5, None, 6], [7, 9]], ref=D(2027, 7, 1))
print('ok')
