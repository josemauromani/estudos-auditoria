"""Lê os resultados dos testes de Competencias-modelo.xlsx depois do recálculo.  Uso: comp_test_read.py <pasta com t1..t3.xlsx>"""
import openpyxl, sys
D = sys.argv[1]
for nome in ('t1', 't2', 't3'):
    w = openpyxl.load_workbook(f'{D}/{nome}.xlsx', data_only=True)
    F, M, P, R, C = w['Funções'], w['Matriz'], w['Plano'], w['Registros'], w['Checklist']
    print('=====', nome)
    err = [(ws.title, c.coordinate, c.value) for ws in w.worksheets[:6] for r in ws.iter_rows() for c in r
           if isinstance(c.value, str) and ((c.value.startswith('#') and len(c.value) > 1) or 'Err:' in c.value)]
    print('erros', err)
    print('Funções    ', [F[f'P{r}'].value for r in range(10, 20) if F[f'C{r}'].value], '| resumo', [F[f'E{r}'].value for r in range(22, 25)])
    print('Matriz     ', [(M[f'C{r}'].value, M[f'Q{r}'].value, M[f'R{r}'].value, M[f'T{r}'].value) for r in range(10, 30) if M[f'C{r}'].value])
    print('            por competência', [[M.cell(row=r, column=5 + k).value for r in range(32, 37)] for k in range(3)])
    print('            resumo', [M[f'E{r}'].value for r in range(39, 44)], '| requerido de 2ª pessoa', [M.cell(row=57, column=5 + k).value for k in range(3)])
    print('Plano      ', [(P[f'E{r}'].value, P[f'M{r}'].value) for r in range(7, 32) if P[f'M{r}'].value], '| resumo', [P[f'E{r}'].value for r in range(34, 41)])
    print('Registros  ', [R[f'J{r}'].value for r in range(7, 32) if R[f'J{r}'].value], '| resumo', [R[f'E{r}'].value for r in range(34, 38)])
    print('Checklist  ', [C[f'D{r}'].value for r in range(19, 25)])
