"""Confere Satisfacao-modelo.xlsx e preenche cópias com dados de teste.  Uso: sat_test_make.py <modelo.xlsx> <pasta de saída>"""
import openpyxl, sys, os
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
from sat_data import EX1, EX2, resumo_mes
from datetime import date
SRC, OUTDIR = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(SRC, data_only=True)
wf0 = openpyxl.load_workbook(SRC)
for ws in wb.worksheets:
    f = wf0[ws.title]
    print('==', ws.title, ws.dimensions, 'dv', len(f.data_validations.dataValidation), 'cf', sum(len(v) for v in f.conditional_formatting._cf_rules.values()),
          'charts', len(f._charts), 'comments', sum(1 for r in f.iter_rows() for c in r if c.comment))
err = [(ws.title, c.coordinate, c.value) for ws in wb.worksheets for r in ws.iter_rows() for c in r
       if isinstance(c.value, str) and ((c.value.startswith('#') and len(c.value) > 1) or 'Err:' in c.value)]
print('erros', err)
print('vazio:', wb['Pesquisa']['D29'].value, '|', wb['Respostas']['E79'].value, '|', wb['Reclamações']['E44'].value, '|', wb['Painel']['C19'].value, '|', wb['Checklist']['D24'].value)
s = wb['Exemplo 1 - Pizzaria']
rows = [r for r in range(1, s.max_row + 1) if s[f'B{r}'].value in [m[0] for m in EX1['meses']]]
got = [(round(s[f'I{r}'].value, 2), s[f'L{r}'].value, round(s[f'M{r}'].value, 1)) for r in rows]
esp = [(resumo_mes(m)['media'], resumo_mes(m)['nps'], round(resumo_mes(m)['por100'], 1)) for m in EX1['meses']]
print('Exemplo 1 | meses', 'OK' if got == esp else 'DIFERE %s' % got)
s = wb['Exemplo 2 - Indústria']
rows = [r for r in range(1, s.max_row + 1) if s[f'B{r}'].value == 'Todos']
print('Exemplo 2 | geral', round(s[f'J{rows[0]}'].value, 2), 'indicação', s[f'M{rows[0]}'].value)


def preencher(nome, pesq=None, perguntas=(), motivos=(), respostas=(), recl=(), painel=None, check=None):
    w = openpyxl.load_workbook(SRC)
    P, R, C, Pa, K = w['Pesquisa'], w['Respostas'], w['Reclamações'], w['Painel'], w['Checklist']
    for ref, v in (pesq or {}).items():
        if v is not None:
            P[ref] = v
    for k, (perg, mede) in enumerate(perguntas):
        P[f'C{9 + k}'] = perg
        if mede:
            P[f'D{9 + k}'] = mede
    for k, m in enumerate(motivos):
        P[f'C{18 + k}'] = m
    for k, row in enumerate(respostas):
        for col, v in zip('CDEFGHIJL', row):
            if v is not None:
                R[f'{col}{7 + k}'] = v
    for k, row in enumerate(recl):
        for col, v in zip('CDEFGHIKL', row):
            if v is not None:
                C[f'{col}{7 + k}'] = v
    for ref, v in (painel or {}).items():
        if v is not None:
            Pa[ref] = v
    for k, v in enumerate(check or ()):
        if v:
            K[f'D{5 + k}'] = v
    w.save(OUTDIR + '/%s.xlsx' % nome)


D = date
PERG = [(p, m) for _, p, m in EX1['perguntas']]
MOT = [m for m, _ in EX1['motivos']]
# t1: pizzaria, 20 respostas e 5 reclamações de fevereiro
resp = [(D(2027, 2, 1 + k), f'Pedido {2300 + 7 * k}', 5 if k % 4 else 3, 5 if k % 5 else 4, 4 if k % 3 else 5, None, None, 10 if k % 3 == 0 else (8 if k % 3 == 1 else 6), None) for k in range(20)]
recl = [(d, ped, canal, motivo, desc, resp_, res, rnc or None) for d, ped, canal, motivo, desc, resp_, res, rnc, _ in EX1['reclamacoes']]
preencher('t1', pesq={'D4': EX1['head']['org'], 'D5': EX1['head']['periodo'], 'D6': 5}, perguntas=PERG, motivos=MOT, respostas=resp, recl=recl, painel={'C4': 840}, check=['Sim'] * 12)
# t2: indústria, escala de 0 a 10, 8 respostas
resp2 = [(D(2026, 11, 3 + k), f'Cliente {k + 1}', 9 - k % 3, 8, 9, 7 + k % 3, 8, 9 if k % 2 else 7, None) for k in range(8)]
preencher('t2', pesq={'D4': EX2['head']['org'], 'D5': 'Novembro de 2026', 'D6': 10}, perguntas=[(p, m) for _, p, m in EX2['perguntas']], motivos=[m for m, _ in EX2['motivos']],
          respostas=resp2, painel={'C4': 58, 'D9': 8.0}, check=['Sim'] * 6 + ['Parcial', 'Não'])
# t3: erros propositais
preencher('t3', pesq={'D4': 'Teste', 'D5': 'Teste', 'D6': 5}, perguntas=[('Pergunta boa?', 'x'), ('A qualidade, o prazo e o preço foram bons?', 'x'),
                                                                        ('Uma pergunta longa demais que não cabe na tela do celular e que ninguém vai ler até o fim, nem responder?', 'x')],
          motivos=['Motivo A', 'Motivo B'],
          respostas=[(D(2027, 1, 5), 'Completa', 5, 4, 5, None, None, 9, None), (D(2027, 1, 5), 'Faltam notas', 5, None, None, None, None, 8, None),
                     (D(2027, 1, 5), 'Sem notas', None, None, None, None, None, 7, None), (D(2027, 1, 5), 'Acima da escala', 5, 6, 5, None, None, 10, None),
                     (D(2027, 1, 5), 'Detrator', 2, 1, 2, None, None, 3, 'Ruim')],
          recl=[(None, 'Sem data', 'Telefone', 'Motivo A', 'Desc', None, None, None), (D(2027, 1, 5), 'Sem motivo', 'Telefone', None, 'Desc', None, None, None),
                (D(2027, 1, 5), 'Aberta', 'Telefone', 'Motivo A', 'Desc', None, None, None), (D(2027, 1, 5), 'Respondida tarde', 'Telefone', 'Motivo B', 'Desc', D(2027, 1, 12), None, None),
                (D(2027, 1, 5), 'Resolvida', 'Telefone', 'Motivo A', 'Desc', D(2027, 1, 6), D(2027, 1, 9), 'RNC 1')],
          painel={'C4': 100})
print('ok')
