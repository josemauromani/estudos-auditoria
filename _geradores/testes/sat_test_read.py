"""Lê os resultados dos testes de Satisfacao-modelo.xlsx depois do recálculo.  Uso: sat_test_read.py <pasta com t1..t3.xlsx>"""
import openpyxl, sys
D = sys.argv[1]
for nome in ('t1', 't2', 't3'):
    w = openpyxl.load_workbook(f'{D}/{nome}.xlsx', data_only=True)
    P, R, C, Pa, K = w['Pesquisa'], w['Respostas'], w['Reclamações'], w['Painel'], w['Checklist']
    print('=====', nome)
    err = [(ws.title, c.coordinate, c.value) for ws in w.worksheets[:6] for r in ws.iter_rows() for c in r
           if isinstance(c.value, str) and ((c.value.startswith('#') and len(c.value) > 1) or 'Err:' in c.value)]
    print('erros', err)
    print('Pesquisa   ', P['D7'].value, [P[f'E{r}'].value for r in range(9, 14) if P[f'E{r}'].value], '| resumo', [P[f'D{r}'].value for r in (28, 29)])
    print('Respostas  ', [(R[f'K{r}'].value, R[f'M{r}'].value) for r in range(7, 67) if R[f'M{r}'].value][:8])
    print('            por pergunta', [[R.cell(row=rr, column=5 + k).value for rr in (69, 70, 71)] for k in range(3)], '| resumo', [R[f'E{r}'].value for r in range(74, 80)])
    print('Reclamações', [(C[f'J{r}'].value, C[f'M{r}'].value) for r in range(7, 37) if C[f'M{r}'].value], '| resumo', [C[f'E{r}'].value for r in range(39, 45)],
          '| motivos', [(C[f'I{r}'].value, C[f'M{r}'].value) for r in range(39, 47) if C[f'I{r}'].value])
    print('Painel     ', [(Pa[f'B{r}'].value[:14], Pa[f'C{r}'].value, Pa[f'E{r}'].value) for r in range(8, 15)], '|', [Pa[f'C{r}'].value for r in (17, 18, 19)])
    print('Checklist  ', [K[f'D{r}'].value for r in range(19, 25)])
