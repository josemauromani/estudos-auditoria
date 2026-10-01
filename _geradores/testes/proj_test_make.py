"""Confere Projeto-modelo.xlsx e preenche cópias com dados de teste.  Uso: proj_test_make.py <modelo.xlsx> <pasta de saída>"""
import openpyxl, sys, os
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
from proj_data import (ANAL, APROV, ATENDE, CLI, DES, EX1, EX2, FUN, LIB, NAOAT, PLAN, REPROV, RESS, VALID, VERIF, entrada_conf, etapa_conf, mud_conf)
from datetime import date
SRC, OUTDIR = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(SRC, data_only=True)
wf0 = openpyxl.load_workbook(SRC)
for ws in wb.worksheets:
    f = wf0[ws.title]
    print('==', ws.title, ws.dimensions, 'dv', len(f.data_validations.dataValidation), 'cf', sum(len(v) for v in f.conditional_formatting._cf_rules.values()),
          'charts', len(f._charts), 'comments', sum(1 for r in f.iter_rows() for c in r if c.comment))
err = [(ws.title, c.coordinate, c.value) for ws in wb.worksheets for r in ws.iter_rows() for c in r
       if isinstance(c.value, str) and ((c.value.startswith('#') and len(c.value) > 1) or 'Err:' in c.value)]
print('erros', err)
print('vazio:', wb['Plano']['D39'].value, '|', wb['Entradas']['D42'].value, '|', wb['Mudanças']['D31'].value, '|', wb['Painel']['C36'].value, '|', wb['Painel']['C17'].value)
for aba, ex in (('Exemplo 1 - Pizzaria', EX1), ('Exemplo 2 - Indústria', EX2)):
    s = wb[aba]
    ref, es = ex['head']['ref'], ex['etps']
    re_ = [r for r in range(1, s.max_row + 1) if s[f'C{r}'].value in [e['etapa'] for e in es]]
    rx = [r for r in range(1, s.max_row + 1) if s[f'C{r}'].value in [x['req'] for x in ex['ents']]]
    rm = [r for r in range(1, s.max_row + 1) if s[f'C{r}'].value in [m['oque'] for m in ex['muds']]]
    ok = ([s[f'M{r}'].value for r in re_] == [etapa_conf(e, ref, es) for e in es], [s[f'M{r}'].value for r in rx] == [entrada_conf(x) for x in ex['ents']],
          [s[f'M{r}'].value for r in rm] == [mud_conf(m) for m in ex['muds']])
    print(aba, '| etapas', 'OK' if ok[0] else 'DIFERE %s' % [s[f'M{r}'].value for r in re_], '| entradas', 'OK' if ok[1] else 'DIFERE', '| mudanças', 'OK' if ok[2] else 'DIFERE',
          '| pode liberar:', [s[f'F{r}'].value for r in range(1, s.max_row + 1) if s[f'B{r}'].value == 'Pode liberar?'])


def preencher(nome, head=(), ref=None, etps=(), ents=(), muds=(), check=None):
    w = openpyxl.load_workbook(SRC)
    P, E, M, K = w['Plano'], w['Entradas'], w['Mudanças'], w['Checklist']
    for k, v in enumerate(head):
        P[f'D{4 + k}'] = v
    if ref:
        P['D10'] = ref
    for k, row in enumerate(etps):
        for col, v in zip('CDEFGHIJK', row):
            if v not in (None, ''):
                P[f'{col}{14 + k}'] = v
    for k, row in enumerate(ents):
        for col, v in zip('CDEFGHIJ', row):
            if v not in (None, ''):
                E[f'{col}{7 + k}'] = v
    for k, row in enumerate(muds):
        for col, v in zip('CDEFGH', row):
            if v not in (None, ''):
                M[f'{col}{7 + k}'] = v
    for k, v in enumerate(check or ()):
        if v:
            K[f'D{5 + k}'] = v
    w.save(OUTDIR + '/%s.xlsx' % nome)


def do_exemplo(ex):
    h = ex['head']
    etps = [(e['etapa'], e['tipo'], e['resp'], e['part'], e['prazo'], e['concl'], e['result'], e['acoes'], e['registro']) for e in ex['etps']]
    ents = [(x['id'], x['req'], x['tipo'], x['fonte'], x['saida'], x['metodo'], x['result'], x['evid']) for x in ex['ents']]
    muds = [(m['data'], m['oque'], m['motivo'], m['analise'], m['verif'], m['aut']) for m in ex['muds']]
    return [h['org'], h['projeto'], h['objetivo'], h['cliente'], h['resp']], h['ref'], etps, ents, muds


for nome, ex, check in (('t1', EX1, ['Sim'] * 12), ('t2', EX2, ['Sim'] * 6 + ['Parcial', 'Não'])):
    head, ref, etps, ents, muds = do_exemplo(ex)
    preencher(nome, head, ref, etps, ents, muds, check)
D = date
d1, d2, d3, ref = D(2027, 6, 1), D(2027, 6, 10), D(2027, 7, 1), D(2027, 6, 20)
etps3 = [('Plano', PLAN, 'R', '', d1, d1, APROV, '', ''),                     # OK
         ('Sem tipo', None, 'R', '', d1, None, None, '', ''),                  # falta o tipo
         ('Sem responsável', DES, None, '', d1, None, None, '', ''),           # falta o responsável
         ('Sem prazo', DES, 'R', '', None, None, None, '', ''),               # falta o prazo
         ('Concluída sem resultado', VERIF, 'R', '', d1, d2, None, '', ''),   # falta o resultado
         ('Resultado sem data', VERIF, 'R', '', d1, None, APROV, '', ''),     # falta a data de conclusão
         ('Ressalva sem ação', ANAL, 'R', '', d1, d2, RESS, '', ''),          # falta a ação
         ('Reprovado com ação', VALID, 'R', '', d1, d2, REPROV, 'Refazer', ''),  # OK
         ('Liberado', LIB, 'R', '', d2, d2, APROV, '', ''),                   # liberado sem validação aprovada
         ('Atrasada', VERIF, 'R', '', d2, None, None, '', ''),                # atrasada
         ('Futura', VALID, 'R', '', d3, None, None, '', '')]                  # OK
ents3 = [('E1', 'ok', CLI, 'F', 'S', 'M', ATENDE, 'ev'), ('E2', 'sem tipo', None, 'F', 'S', 'M', ATENDE, ''), ('E3', 'sem fonte', FUN, None, 'S', 'M', '', ''),
         ('E4', 'sem saída', FUN, 'F', None, 'M', '', ''), ('E5', 'sem método', FUN, 'F', 'S', None, '', ''), ('E6', 'não atende sem ação', FUN, 'F', 'S', 'M', NAOAT, ''),
         ('E7', 'não atende com ação', FUN, 'F', 'S', 'M', NAOAT, 'ação'), ('E8', 'pendente', FUN, 'F', 'S', 'M', '', ''), ('E9', None, FUN, None, None, None, None, None)]
muds3 = [(d1, 'ok', 'm', 'a', 'v', 'Quem'), (d1, 'sem análise', 'm', None, 'v', 'Quem'), (d1, 'sem autorização', 'm', 'a', 'v', None),
         (d1, 'sem verificação', 'm', 'a', None, 'Quem'), (d1, None, 'só o motivo', None, None, None)]
preencher('t3', ['Org', 'Teste'], ref, etps3, ents3, muds3)
print('ok')
