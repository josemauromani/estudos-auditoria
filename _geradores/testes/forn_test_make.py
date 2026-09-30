"""Confere Fornecedores-modelo.xlsx e preenche cópias com dados de teste.  Uso: forn_test_make.py <modelo.xlsx> <pasta de saída>"""
import openpyxl, sys, os
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
from forn_data import EX1, EX2, HOMOLOG
from datetime import date
SRC, OUTDIR = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(SRC, data_only=True)
wf0 = openpyxl.load_workbook(SRC)
for ws in wb.worksheets:
    f = wf0[ws.title]
    print('==', ws.title, ws.dimensions, 'dv', len(f.data_validations.dataValidation), 'cf', sum(len(v) for v in f.conditional_formatting._cf_rules.values()),
          'charts', len(f._charts), 'comments', sum(1 for r in f.iter_rows() for c in r if c.comment))
err = [(ws.title, c.coordinate, c.value) for ws in wb.worksheets for r in ws.iter_rows() for c in r
       if isinstance(c.value, str) and ((c.value.startswith('#') and len(c.value) > 1) or 'Err:' in c.value)]
print('erros', err)
print('vazio:', wb['Cadastro']['E46'].value, '|', wb['Homologação']['C19'].value, '|', wb['Avaliação']['E43'].value, '|', wb['Ocorrências']['E37'].value, '|', wb['Checklist']['D24'].value)
for nome, ex in (('Exemplo 1 - Indústria', EX1), ('Exemplo 2 - Pizzaria', EX2)):
    s = wb[nome]
    rows = [r for r in range(1, s.max_row + 1) if s[f'B{r}'].value in [a['cod'] for a in ex['avals']] and isinstance(s[f'I{r}'].value, (int, float))]
    idx = [(s[f'I{r}'].value, s[f'J{r}'].value) for r in rows]
    ok = all(abs(x - y) < 0.051 and c == d for (x, c), (y, d) in zip(idx, [(a['indice'], a['classe']) for a in ex['avals']]))
    print(nome, '| índices e classes', 'OK' if ok else 'DIFERE %s' % idx)


def preencher(nome, ex=None, cab=None, forns=None, homolog=None, avals=None, ocs=(), check=None):
    w = openpyxl.load_workbook(SRC)
    C, H, A, O, K = w['Cadastro'], w['Homologação'], w['Avaliação'], w['Ocorrências'], w['Checklist']
    if ex:
        cab = dict({'E4': ex['head']['org'], 'E5': ex['head']['data']}, **(cab or {}))
        forns = forns or [(f['cod'], f['nome'], f['fornece'], f['tipo'], f['crit'], f['homologado'], None, f['docs']) for f in ex['forns']]
        avals = avals if avals is not None else [(a['cod'], ex['head']['periodo'], a['recebidos'], a['aceitos'], a['entregas'], a['no_prazo'], a['nota'], a['obs'] or None)
                                                 for a in ex['avals']]
    for ref, v in (cab or {}).items():
        if v is not None:
            C[ref] = v
    for k, row in enumerate(forns or []):
        for col, v in zip('CDEFGHIK', row):
            if v is not None:
                C[f'{col}{9 + k}'] = v
    if homolog:
        H['C4'], H['C5'], H['C6'], H['C7'] = homolog['cab']
        for k, (res, ev) in enumerate(homolog['itens']):
            if res is not None:
                H[f'D{10 + k}'] = res
            if ev is not None:
                H[f'E{10 + k}'] = ev
    for k, row in enumerate(avals or []):
        for col, v in zip('CEFGHIJP', row):
            if v is not None:
                A[f'{col}{7 + k}'] = v
    for k, row in enumerate(ocs):
        for col, v in zip('CDFGHIJ', row):
            if v is not None:
                O[f'{col}{7 + k}'] = v
    for k, v in enumerate(check or ()):
        if v:
            K[f'D{5 + k}'] = v
    w.save(OUTDIR + '/%s.xlsx' % nome)


D = date
# t1: indústria completa, com a ficha do segundo fornecedor de resina e duas ocorrências
preencher('t1', EX1, homolog=dict(cab=(f'{HOMOLOG["cod"]} · {HOMOLOG["nome"]}', HOMOLOG['fornece'], HOMOLOG['data'], HOMOLOG['por']),
                                   itens=[(res, ev) for _, res, ev in HOMOLOG['itens']]),
          ocs=[(D(2026, 10, 7), 'F-08', 'Qualidade', 'Lote de 40 paletes sem tratamento.', 'Lote devolvido.', '2026-35', 'Encerrada'),
               (D(2026, 11, 12), 'F-07', 'Prazo', 'Entrega com 6 dias de atraso.', None, None, 'Aberta')], check=['Sim'] * 12)
# t2: pizzaria, sem data de referência (usa hoje)
preencher('t2', EX2, cab={'E5': None}, check=['Sim'] * 6 + ['Parcial', 'Não'])
# t3: erros propositais
preencher('t3', cab={'E4': 'Teste', 'E5': D(2027, 3, 5), 'E6': 12},
          forns=[('X-1', 'Completo', 'Item', 'Produto', 'Crítico', D(2026, 9, 1), None, 'Sim'),          # homologado, válido até 09/2027
                 ('X-2', 'Vencido', 'Item', 'Produto', 'Crítico', D(2026, 1, 1), None, 'Sim'),           # 01/2027 < 05/03/2027
                 ('X-3', 'Vencido pela validade própria', 'Item', 'Produto', 'Crítico', D(2027, 1, 1), 1, 'Sim'),
                 ('X-4', 'Sem homologação', 'Item', 'Serviço', 'Crítico', None, None, None),
                 ('X-5', 'Sem documentos', 'Item', 'Produto', 'Crítico', D(2026, 9, 1), None, 'Não'),
                 ('X-6', 'Sem criticidade', 'Item', 'Produto', None, None, None, None),
                 ('X-7', 'Não crítico', 'Item', 'Produto', 'Não crítico', None, None, None)],
          homolog=dict(cab=('X-8 · Novo', 'Item', D(2027, 3, 1), 'Comprador'), itens=[('Atende', 'ok'), ('Atende', 'ok'), ('Não atende', 'amostra reprovada'), ('Atende', None), ('Atende', 'ok'), ('Não se aplica', None)]),
          avals=[('X-1', 'Semestre', 10, 10, 10, 10, 10, None),          # 100, A
                 ('X-2', 'Semestre', 10, 8, 10, 7, 6, None),             # 40+21+12 = 73, C sem plano
                 ('X-3', 'Semestre', 10, 4, 10, 5, 3, 'Suspenso'),       # 20+15+6 = 41, D com decisão
                 ('X-4', 'Semestre', 10, 12, 10, 10, 8, None),           # aceitos > recebidos
                 ('X-5', 'Semestre', None, None, 10, 10, 8, None),       # faltam dados
                 ('X-9', 'Semestre', 10, 10, 10, 10, 8, None),           # código fora do cadastro
                 ('X-7', 'Semestre', 0, 0, 10, 10, 8, None)],            # total zero
          ocs=[(None, 'X-1', 'Prazo', 'Sem data', None, None, 'Aberta'), (D(2027, 1, 5), None, 'Prazo', 'Sem fornecedor', None, None, 'Aberta'),
               (D(2027, 1, 5), 'X-1', None, 'Sem tipo', None, None, 'Aberta'), (D(2027, 1, 5), 'X-1', 'Prazo', 'Sem situação', None, None, None),
               (D(2027, 1, 5), 'X-1', 'Prazo', 'Tratada sem tratamento', None, None, 'Tratada'), (D(2027, 1, 5), 'X-1', 'Prazo', 'Completa', 'Reposição', None, 'Encerrada')])
print('ok')
