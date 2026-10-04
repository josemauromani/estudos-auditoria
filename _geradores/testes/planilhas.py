# Varredura de planilhas: lista células com erro e fórmulas sem valor gravado.
# Uso: planilhas.py <arquivo.xlsx>...   (sai com 1 se achar algo)
# Ficam de fora: #N/A de fórmula com NA() (de propósito, para os gráficos pularem o ponto)
# e fórmula que devolve texto vazio (o LibreOffice grava como texto, sem conteúdo).
import re, sys, zipfile
import openpyxl

ERROS = ('#DIV/0!', '#N/A', '#NAME?', '#NULL!', '#NUM!', '#REF!', '#VALUE!', 'Err:')
NS = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'


def _sem_valor(caminho):
    """Células com fórmula e sem tipo de resultado gravado (arquivo que ninguém recalculou)."""
    import xml.etree.ElementTree as ET
    achados = set()
    with zipfile.ZipFile(caminho) as z:
        nomes = [n for n in z.namelist() if re.match(r'xl/worksheets/sheet\d+\.xml$', n)]
        for n in nomes:
            for c in ET.fromstring(z.read(n)).iter(NS + 'c'):
                v = c.find(NS + 'v')
                if c.find(NS + 'f') is not None and c.get('t') != 'str' and (v is None or not (v.text or '').strip()):
                    achados.add((n, c.get('r')))
    return achados


def varrer(caminho):
    achados = []
    f = openpyxl.load_workbook(caminho)
    v = openpyxl.load_workbook(caminho, data_only=True)
    sem_valor = _sem_valor(caminho)
    for i, ws in enumerate(f.worksheets, 1):
        wv = v[ws.title]
        for linha in ws.iter_rows():
            for c in linha:
                val = wv[c.coordinate].value
                formula = c.value if isinstance(c.value, str) and c.value.startswith('=') else ''
                if isinstance(val, str) and val.startswith(ERROS):
                    if val == '#N/A' and 'NA()' in formula:
                        continue
                    achados.append((ws.title, c.coordinate, 'erro %s' % val))
                elif formula and val is None and ('xl/worksheets/sheet%d.xml' % i, c.coordinate) in sem_valor:
                    achados.append((ws.title, c.coordinate, 'fórmula sem valor gravado'))
    return achados


def main(argv):
    total = 0
    for caminho in argv:
        for aba, cel, msg in varrer(caminho):
            print('%s | %s!%s | %s' % (caminho, aba, cel, msg))
            total += 1
    print('%d célula(s) com problema em %d planilha(s)' % (total, len(argv)))
    return 1 if total else 0


if __name__ == '__main__':
    sys.exit(main(sys.argv[1:]))
