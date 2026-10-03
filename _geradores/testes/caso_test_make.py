"""Confere Caso-Integrado-modelo.xlsx e preenche cópias com dados de teste.  Uso: caso_test_make.py <modelo.xlsx> <pasta de saída>"""
import openpyxl, sys, os
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
from caso_data import EX1, EX2, cal_conf, cal_conta, cal_sit, rastro_conf, rastro_dias
from datetime import date
SRC, OUTDIR = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(SRC, data_only=True)
wf0 = openpyxl.load_workbook(SRC)
for ws in wb.worksheets:
    f = wf0[ws.title]
    print('==', ws.title, ws.dimensions, 'dv', len(f.data_validations.dataValidation), 'cf', sum(len(v) for v in f.conditional_formatting._cf_rules.values()),
          'charts', len(f._charts), 'comments', sum(1 for r in f.iter_rows() for c in r if c.comment))
err = [(ws.title, c.coordinate, c.value) for ws in wb.worksheets for r in ws.iter_rows() for c in r
       if isinstance(c.value, str) and ((c.value.startswith('#') and len(c.value) > 1) or 'Err:' in c.value)]
print('erros', err)
print('vazio:', wb['Rastro']['E42'].value, '|', wb['Calendário']['E55'].value, '|', wb['Painel']['C24'].value)


def v(c):
    """O openpyxl lê como None o texto vazio que a fórmula devolve."""
    return '' if c.value is None else c.value


def e(x):
    return '' if x is None else x


for aba, ex in (('Exemplo 1 - Pizzaria', EX1), ('Exemplo 2 - Indústria', EX2)):
    s = wb[aba]
    ref = ex['head']['ref']
    r1 = next(r for r in range(1, s.max_row) if s[f'B{r}'].value == 'Passo') + 1
    c1 = next(r for r in range(1, s.max_row) if s[f'C{r}'].value == 'Atividade') + 2
    got_r = [(v(s[f'I{r}']), v(s[f'J{r}'])) for r in range(r1, r1 + len(ex['passos']))]
    esp_r = [(e(d), c) for d, c in zip(rastro_dias(ex['passos']), rastro_conf(ex['passos']))]
    got_c = [tuple(v(s[f'{c}{r}']) for c in ('U', 'V', 'W', 'X', 'Y', 'Z')) for r in range(c1, c1 + len(ex['cal']))]
    esp_c = []
    for a in ex['cal']:
        k = cal_conta(a, ref)
        esp_c.append((k['prev'], k['feitos'], k['atras'], k['fora'], cal_sit(a, ref), cal_conf(a)))
    print(aba, '| rastro', 'OK' if got_r == esp_r else 'DIFERE %s' % got_r, '| calendário', 'OK' if got_c == esp_c else 'DIFERE %s\n   esperado %s' % (got_c, esp_c))


def preencher(nome, fio=None, ano=None, ref=None, passos=(), cal=(), check=None):
    w = openpyxl.load_workbook(SRC)
    R, C, K = w['Rastro'], w['Calendário'], w['Checklist']
    if fio:
        R['D4'] = fio
    if ano:
        C['E4'], C['E5'] = ano, ref
    for k, row in enumerate(passos):
        for col, val in zip('CDEFGH', row):
            if val not in (None, ''):
                R[f'{col}{8 + k}'] = val
    for k, (ativ, est, req, resp, freq, ini, feito) in enumerate(cal):
        for col, val in zip('CDEFGH', (ativ, est, req, resp, freq, ini)):
            if val not in (None, ''):
                C[f'{col}{10 + k}'] = val
        for m in feito:
            C.cell(row=10 + k, column=8 + m, value='X')
    for k, val in enumerate(check or ()):
        if val:
            K[f'D{5 + k}'] = val
    w.save(OUTDIR + '/%s.xlsx' % nome)


def linhas(ex):
    passos = [(p['data'], p['estudo'], p['oque'], p['reg'], p['saida'], p['prox']) for p in ex['passos']]
    cal = [(a['ativ'], a['estudo'], a['req'], a['resp'], a['freq'], a['inicio'], sorted(a['feito'])) for a in ex['cal']]
    return passos, cal


for nome, ex, check in (('t1', EX1, ['Sim'] * 12), ('t2', EX2, ['Sim'] * 6 + ['Parcial', 'Não'])):
    preencher(nome, ex['head']['fio'], ex['head']['ano'], ex['head']['ref'], *linhas(ex), check=check)

# ---- t3: casos de borda
D = date
REF3 = 6
passos3 = [(D(2027, 1, 10), 'PDCA', 'o', 'r', 's', 'Pareto'),              # OK
           (D(2027, 1, 20), 'Pareto', 'o', 'r', 's', '5W2H'),              # o passo seguinte é de outro estudo
           (D(2027, 2, 1), 'Ishikawa', 'o', '', 's', 'Matriz GUT'),         # falta o registro
           (D(2027, 1, 5), 'Matriz GUT', 'o', 'r', '', 'Indicadores'),      # data antes do passo anterior
           (D(2027, 3, 1), 'Indicadores', 'o', 'r', '', 'Análise crítica'),  # falta a saída
           (D(2027, 3, 10), 'Análise crítica', 'o', 'r', 's', ''),          # falta o próximo estudo
           (None, 'PDCA', 'o', 'r', 's', 'PDCA'),                           # falta a data
           (D(2027, 4, 1), '', 'o', 'r', 's', 'PDCA'),                      # falta o estudo
           (D(2027, 4, 2), 'PDCA', '', 'r', 's', 'PDCA'),                   # falta o que aconteceu
           (D(2027, 4, 9), 'PDCA', 'o', 'r', '', '')]                       # último passo: OK sem saída
cal3 = [('A', 'PDCA', 'r', 'p', 'Mensal', 1, [1, 2, 3, 4, 5, 6]),     # em dia
        ('B', 'PDCA', 'r', 'p', 'Mensal', 1, [1, 2, 3]),              # 3 atrasadas
        ('C', 'PDCA', 'r', 'p', 'Trimestral', 2, [2, 5]),             # em dia
        ('D', 'PDCA', 'r', 'p', 'Semestral', 7, []),                  # ainda não começou
        ('E', 'PDCA', 'r', 'p', 'Anual', 6, [6]),                     # em dia, no mês da leitura
        ('F', 'PDCA', 'r', 'p', 'Bimestral', 1, [1, 2, 3, 5]),        # em dia, 1 fora do plano
        ('G', 'PDCA', 'r', 'p', 'Semestral', 6, []),                  # atrasada no mês da leitura
        ('H', 'PDCA', 'r', '', 'Mensal', 1, [1]),                     # falta o responsável
        ('I', '', 'r', 'p', 'Mensal', 1, []),                         # falta o estudo
        ('J', 'PDCA', 'r', 'p', None, 1, []),                         # falta a frequência
        ('K', 'PDCA', 'r', 'p', 'Mensal', None, []),                  # falta o mês de início
        ('L', 'PDCA', 'r', 'p', 'Mensal', 13, []),                    # mês fora de 1 a 12
        ('M', 'PDCA', 'r', 'p', 'Trimestral', 1, [1, 4, 7])]          # marca em mês previsto futuro: não conta
preencher('t3', 'Teste', 2027, REF3, passos3, cal3)
print('ok')
