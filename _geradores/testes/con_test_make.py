"""Confere Conhecimento-modelo.xlsx e preenche cópias com dados de teste.  Uso: con_test_make.py <modelo.xlsx> <pasta de saída>"""
import openpyxl, sys, os
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
from con_data import (ALTA, BAIXA, CABECA, ESCRITO, EX1, EX2, MEDIA, NAO, SIM, TREINADO, at_conf, at_dias, at_sit, con_conf, con_sit, lic_conf, lic_dias,
                      lic_sit, pol_conf, pontos, prazo_at)
from datetime import date
SRC, OUTDIR = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(SRC, data_only=True)
wf0 = openpyxl.load_workbook(SRC)
for ws in wb.worksheets:
    f = wf0[ws.title]
    print('==', ws.title, ws.dimensions, 'dv', len(f.data_validations.dataValidation), 'cf', sum(len(v) for v in f.conditional_formatting._cf_rules.values()),
          'charts', len(f._charts), 'comments', sum(1 for r in f.iter_rows() for c in r if c.comment))
err = [(ws.title, c.coordinate, c.value) for ws in wb.worksheets for r in ws.iter_rows() for c in r
       if isinstance(c.value, str) and ((c.value.startswith('#') and len(c.value) > 1) or 'Err:' in c.value)]
print('erros', err)
print('vazio:', wb['Conhecimento']['E52'].value, '|', wb['Lições']['E71'].value, '|', wb['Pós-entrega']['E30'].value, '|', wb['Atendimentos']['E91'].value, '|',
      wb['Painel']['C28'].value)


def v(c):
    """O openpyxl lê como None o texto vazio que a fórmula devolve."""
    return '' if c.value is None else c.value


for aba, ex in (('Exemplo 1 - Pizzaria', EX1), ('Exemplo 2 - Indústria', EX2)):
    s = wb[aba]
    ref = ex['head']['ref']
    k1 = next(r for r in range(1, s.max_row) if s[f'C{r}'].value == 'Conhecimento') + 1
    l1 = next(r for r in range(1, s.max_row) if s[f'C{r}'].value == 'O que aconteceu') + 1
    p1 = next(r for r in range(1, s.max_row) if s[f'B{r}'].value == 'Produto ou serviço') + 1
    a1 = next(r for r in range(1, s.max_row) if s[f'C{r}'].value == 'Relato') + 1
    got = dict(con=[(v(s[f'M{r}']), v(s[f'N{r}']), v(s[f'O{r}'])) for r in range(k1, k1 + len(ex['cons']))],
               lic=[(v(s[f'M{r}']), v(s[f'N{r}']), v(s[f'O{r}'])) for r in range(l1, l1 + len(ex['lics']))],
               pol=[v(s[f'Q{r}']) for r in range(p1, p1 + len(ex['pols']))],
               at=[(v(s[f'M{r}']), v(s[f'N{r}']), v(s[f'O{r}']), v(s[f'P{r}'])) for r in range(a1, a1 + len(ex['ats']))])
    esp = dict(con=[(pontos(k), con_sit(k), con_conf(k, ref)) for k in ex['cons']],
               lic=[('' if lic_dias(l, ref) is None else lic_dias(l, ref), lic_sit(l), lic_conf(l, ex['cons'], ref)) for l in ex['lics']],
               pol=[pol_conf(p) for p in ex['pols']],
               at=[(prazo_at(a, ex['pols']), at_dias(a, ref), at_sit(a, ex['pols'], ref), at_conf(a, ex['pols'], ex['lics'])) for a in ex['ats']])
    print(aba, ' | '.join(f'{k} ' + ('OK' if got[k] == esp[k] else f'DIFERE {got[k]}') for k in got))


def preencher(nome, ref=None, cons=(), lics=(), pols=(), ats=(), check=None):
    w = openpyxl.load_workbook(SRC)
    K, L, P, A, C = w['Conhecimento'], w['Lições'], w['Pós-entrega'], w['Atendimentos'], w['Checklist']
    if ref:
        K['E4'] = ref
    for rows, ws, cols, r0 in ((cons, K, 'CDEFGHIJKL', 8), (lics, L, 'CDEFGHIJK', 7), (pols, P, 'CDEFGHI', 7), (ats, A, 'CDEFGHIJKL', 7)):
        for k, row in enumerate(rows):
            for col, val in zip(cols, row):
                if val not in (None, ''):
                    ws[f'{col}{r0 + k}'] = val
    for k, val in enumerate(check or ()):
        if val:
            C[f'D{5 + k}'] = val
    w.save(OUTDIR + '/%s.xlsx' % nome)


def linhas(ex):
    cons = [(k['cod'], k['nome'], k['proc'], k['crit'], k['forma'], k['onde'], k['pessoas'], k['saida'], k['acao'], k['prazo']) for k in ex['cons']]
    lics = [(l['num'], l['data'], l['origem'], l['oque'], l['aprend'], l['onde'], l['resp'], l['cod'], l['incorp']) for l in ex['lics']]
    pols = [(p['prod'], p['vida'], p['minimo'], p['aceito'], p['ativ'], p['resp'], p['conseq']) for p in ex['pols']]
    ats = [(a['num'], a['abert'], a['cliente'], a['prod'], a['tipo'], a['relato'], a['resol'], a['custo'], a['causa'], a['licao']) for a in ex['ats']]
    return cons, lics, pols, ats


for nome, ex, check in (('t1', EX1, ['Sim'] * 12), ('t2', EX2, ['Sim'] * 6 + ['Parcial', 'Não'])):
    preencher(nome, ex['head']['ref'], *linhas(ex), check=check)

# ---- t3: casos de borda
D = date
REF3 = D(2027, 6, 30)
cons3 = [('A', 'n', 'p', ALTA, CABECA, '', 1, SIM, 'a', D(2027, 7, 31)),       # 5 pontos: crítico, OK
         ('B', 'n', 'p', ALTA, CABECA, '', 1, NAO, '', None),                  # crítico sem ação
         ('C', 'n', 'p', ALTA, ESCRITO, 'doc', 2, NAO, 'a', None),             # em risco, sem prazo
         ('D', 'n', 'p', ALTA, ESCRITO, 'doc', 2, NAO, 'a', D(2027, 6, 1)),    # prazo vencido
         ('E', 'n', 'p', ALTA, ESCRITO, '', 3, NAO, '', None),                 # sob controle, falta onde
         ('F', 'n', 'p', MEDIA, CABECA, '', 3, NAO, '', None),                 # média com 2: em risco, sem ação
         ('G', 'n', 'p', BAIXA, CABECA, '', 1, SIM, '', None),                 # baixa: sob controle mesmo com 5
         ('H', 'n', 'p', MEDIA, TREINADO, 'doc', 1, SIM, 'a', D(2027, 7, 1)),  # média com 3: em risco, OK
         ('I', 'n', 'p', None, CABECA, '', 1, NAO, '', None),                  # falta a criticidade
         ('J', 'n', 'p', ALTA, None, '', 2, NAO, '', None),                    # falta a forma
         ('K', 'n', 'p', ALTA, ESCRITO, 'doc', None, NAO, '', None),           # falta quantas pessoas
         ('L', 'n', 'p', ALTA, TREINADO, 'doc', 0, NAO, '', None),             # ninguém sabe: 2 pontos, em risco
         ('M', 'n', 'p', ALTA, ESCRITO, 'doc', 2, SIM, 'a', REF3),             # prazo no dia da leitura: OK
         ('N', 'n', 'p', ALTA, ESCRITO, 'doc', 2, None, '', None)]             # saída em branco conta como não
lics3 = [('X1', D(2027, 6, 1), 'Quebra', 'o', 'a', 'd', 'r', 'A', D(2027, 6, 10)),  # incorporada
         ('X2', D(2027, 5, 1), 'Quebra', 'o', 'a', 'd', 'r', '', None),             # 60 dias: ainda OK
         ('X3', D(2027, 4, 30), 'Quebra', 'o', 'a', 'd', 'r', '', None),            # 61 dias: parada
         ('X4', None, 'Quebra', 'o', 'a', 'd', 'r', '', None),                      # falta a data
         ('X5', D(2027, 6, 1), 'Quebra', '', 'a', 'd', 'r', '', None),              # falta o que aconteceu
         ('X6', D(2027, 6, 1), 'Quebra', 'o', '', 'd', 'r', '', None),              # falta o que aprendemos
         ('X7', D(2027, 6, 1), 'Quebra', 'o', 'a', '', 'r', '', None),              # falta onde
         ('X8', D(2027, 6, 1), 'Quebra', 'o', 'a', 'd', 'r', 'ZZ', None)]           # conhecimento não cadastrado
pols3 = [('P1', 'v', 30, 30, 'a', 2, 'c'),      # OK
         ('P2', 'v', 30, 10, 'a', 2, ''),       # aceito menor que o mínimo
         ('P3', 'v', None, 10, 'a', 2, ''),     # falta o mínimo
         ('P4', 'v', 30, None, 'a', 2, ''),     # falta o aceito
         ('P5', 'v', 30, 30, '', 2, ''),        # falta o que se faz
         ('P6', 'v', 30, 30, 'a', None, ''),    # falta o prazo de resposta
         ('P7', 'v', 30, 30, 'a', 0, '')]       # resposta no mesmo dia
ats3 = [('T1', D(2027, 6, 1), 'c', 'P1', 'Reclamação', 'r', D(2027, 6, 3), 10, 'c', 'X1'),     # 2 de 2: no prazo
        ('T2', D(2027, 6, 1), 'c', 'P1', 'Reclamação', 'r', D(2027, 6, 4), 10, 'c', ''),       # 3 de 2: fora
        ('T3', D(2027, 6, 29), 'c', 'P1', 'Garantia', 'r', None, None, '', ''),                 # aberto no prazo
        ('T4', D(2027, 6, 20), 'c', 'P1', 'Devolução', 'r', None, None, '', ''),                # aberto e atrasado
        ('T5', D(2027, 6, 20), 'c', 'P7', 'Reclamação', 'r', D(2027, 6, 20), 0, 'c', ''),      # 0 de 0: no prazo
        ('T6', D(2027, 6, 20), 'c', 'PX', 'Reclamação', 'r', D(2027, 6, 20), 0, 'c', ''),      # produto sem política
        ('T7', None, 'c', 'P1', 'Reclamação', 'r', D(2027, 6, 20), 0, 'c', ''),                # falta a abertura
        ('T8', D(2027, 6, 20), 'c', 'P1', None, 'r', D(2027, 6, 21), 0, 'c', ''),              # falta o tipo
        ('T9', D(2027, 6, 20), 'c', 'P1', 'Garantia', 'r', D(2027, 6, 21), 5, '', ''),         # falta a causa
        ('T10', D(2027, 6, 20), 'c', 'P1', 'Assistência técnica', 'r', D(2027, 6, 21), 5, '', ''),  # assistência sem causa: OK
        ('T11', D(2027, 6, 20), 'c', 'P1', 'Reclamação', 'r', D(2027, 6, 21), 5, 'c', 'X99'),  # lição não encontrada
        ('T12', D(2027, 6, 20), 'c', 'P6', 'Reclamação', 'r', D(2027, 6, 21), 5, 'c', '')]     # produto sem prazo: sem situação
preencher('t3', REF3, cons3, lics3, pols3, ats3)
print('ok')
