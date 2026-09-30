"""Confere Pareto-modelo.xlsx e preenche cópias com dados de teste.  Uso: par_test_make.py <modelo.xlsx> <pasta de saída>"""
import openpyxl, sys, os
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
from par_data import EX1, EX2, EX3, linhas, pareto
SRC, OUTDIR = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(SRC, data_only=True)
wf0 = openpyxl.load_workbook(SRC)
for ws in wb.worksheets:
    f = wf0[ws.title]
    print('==', ws.title, ws.dimensions, 'dv', len(f.data_validations.dataValidation), 'cf', sum(len(v) for v in f.conditional_formatting._cf_rules.values()),
          'charts', len(f._charts), 'comments', sum(1 for r in f.iter_rows() for c in r if c.comment))
# a coluna M da aba Pareto guarda #N/D de propósito, nas posições sem categoria
err = [(ws.title, c.coordinate, c.value) for ws in wb.worksheets for r in ws.iter_rows() for c in r
       if isinstance(c.value, str) and ((c.value.startswith('#') and len(c.value) > 1) or 'Err:' in c.value) and not (ws.title == 'Pareto' and c.column_letter == 'M')]
print('erros', err)
print('vazio:', wb['Coleta']['D26'].value, '|', wb['Folha']['D24'].value, '|', wb['Pareto']['D23'].value, '|', wb['Antes e depois']['D26'].value, '|', wb['Checklist']['D24'].value)
for aba, ex in (('Exemplo 1 - Pizzaria', EX1), ('Exemplo 2 - Compras', EX2)):
    s = wb[aba]
    esp = [(r['nome'], r['v'], round(r['pct'], 4), round(r['acc'], 4), r['classe']) for r in pareto(linhas(ex))]
    r0 = next(r for r in range(1, s.max_row + 1) if s[f'B{r}'].value == 1 and s[f'C{r}'].value == esp[0][0])
    got = [(s[f'C{r}'].value, s[f'D{r}'].value, round(s[f'E{r}'].value, 4), round(s[f'F{r}'].value, 4), s[f'G{r}'].value) for r in range(r0, r0 + len(esp))]
    print(aba, '| Pareto', 'OK' if got == esp else 'DIFERE %s' % got)


def preencher(nome, coleta=None, cats=(), outros_peso=None, cols=(), folha=(), outros=(), base=None, depois=None, check=None):
    w = openpyxl.load_workbook(SRC)
    C, F, P, A, K = w['Coleta'], w['Folha'], w['Pareto'], w['Antes e depois'], w['Checklist']
    for ref, v in (coleta or {}).items():
        C[ref] = v
    for k, (cat, defin, peso) in enumerate(cats):
        for col, v in zip('CDE', (cat, defin, peso)):
            if v is not None:
                C[f'{col}{14 + k}'] = v
    if outros_peso is not None:
        C['E22'] = outros_peso
    for j, c in enumerate(cols):
        F.cell(row=6, column=4 + j, value=c)
    for k, row in enumerate(folha):
        for j, v in enumerate(row):
            if v is not None:
                F.cell(row=8 + k, column=4 + j, value=v)
    for j, v in enumerate(outros):
        F.cell(row=16, column=4 + j, value=v)
    if base:
        P['D4'] = base
    for ref, v in (depois or {}).items():
        A[ref] = v
    for k, v in enumerate(check or ()):
        if v:
            K[f'D{5 + k}'] = v
    w.save(OUTDIR + '/%s.xlsx' % nome)


H = EX1['head']
nomeadas = [f for f in EX1['folha'] if f[0] != 'Outros']
# t1: pizzaria, em outra ordem de cadastro, com o antes e o depois
ordem = [2, 0, 3, 1]
dep = dict(zip((n for n, _, _ in EX1['folha']), EX1['depois']['valores']))
preencher('t1', coleta={'D4': H['org'], 'D5': H['conta'], 'D6': H['periodo'], 'D7': H['como'], 'D8': H['base'], 'D9': H['base_nome']},
          cats=[(nomeadas[i][0], nomeadas[i][1], None) for i in ordem], cols=EX1['cols'], folha=[nomeadas[i][2] for i in ordem], outros=EX1['folha'][-1][2],
          depois=dict({'D5': EX1['depois']['periodo'], 'D7': EX1['depois']['base'], 'E19': dep['Outros']}, **{f'E{11 + k}': dep[nomeadas[i][0]] for k, i in enumerate(ordem)}),
          check=['Sim'] * 12)
# t2: indústria, com peso por ocorrência e base em frequência × peso
D3 = [d for d in EX3['defeitos'] if d[0] != 'Outros']
preencher('t2', coleta={'D4': 'Indústria', 'D5': EX3['head']['conta'], 'D6': EX3['head']['periodo']}, cats=[(n, 'Definição', kg) for n, _, kg in D3], outros_peso=EX3['defeitos'][-1][2],
          cols=['Jul', 'Ago', 'Set'], folha=[[q - 2 * (q // 3), q // 3, q // 3] for _, q, _ in D3], outros=[3, 3, 3], base='Frequência × peso',
          depois={'E11': 20}, check=['Sim'] * 6 + ['Parcial', 'Não'])
# t3: erros propositais: categoria repetida e sem definição, contagem em linha sem categoria, empate, poucos dados e Outros grande
preencher('t3', coleta={'D5': 'Teste'}, cats=[('A', 'def', None), ('B', None, None), ('A', 'def', None), ('C', 'def', None)], cols=['X', 'Y'],
          folha=[[5, 1], [3, 3], [2, 2], [6, 0], [1, 0]], outros=[4, 4], depois={'E11': 8, 'E12': 6, 'E13': 0, 'E14': 6})
# t4: Pareto plano, e base em frequência × peso sem os pesos
preencher('t4', coleta={'D5': 'Teste'}, cats=[(c, 'def', None) for c in 'ABCDEF'], cols=['X'], folha=[[12], [11], [11], [10], [10], [9]], outros=[2])
preencher('t5', coleta={'D5': 'Teste'}, cats=[(c, 'def', 2 if c != 'C' else None) for c in 'ABCD'], cols=['X'], folha=[[30], [20], [10], [5]], outros=[1], base='Frequência × peso')
print('ok')
