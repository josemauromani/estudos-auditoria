"""Confere ISO-9001-2026-modelo.xlsx e preenche cópias com dados de teste.  Uso: ed26_test_make.py <modelo.xlsx> <pasta de saída>"""
import openpyxl, sys, os
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
from ed26_data import (ATENDE, EX1, EX2, FIM_TRANSICAO, MUDANCAS, NA, NAOATENDE, PARCIAL, PASSOS, acao_sit, diag_conf, limite_passo, passo_sit, prioridade,
                       prioridade_pts)
from datetime import date
SRC, OUTDIR = sys.argv[1], sys.argv[2]
CERT15, SEMCERT = "Certificado pela edição de 2015", "Sem certificado"
wb = openpyxl.load_workbook(SRC, data_only=True)
wf0 = openpyxl.load_workbook(SRC)
for ws in wb.worksheets:
    f = wf0[ws.title]
    print('==', ws.title, ws.dimensions, 'dv', len(f.data_validations.dataValidation), 'cf', sum(len(v) for v in f.conditional_formatting._cf_rules.values()),
          'charts', len(f._charts), 'comments', sum(1 for r in f.iter_rows() for c in r if c.comment))
err = [(ws.title, c.coordinate, c.value) for ws in wb.worksheets for r in ws.iter_rows() for c in r
       if isinstance(c.value, str) and ((c.value.startswith('#') and len(c.value) > 1) or 'Err:' in c.value)]
print('erros', err)
print('vazio:', wb['Mudanças']['E27'].value, '|', wb['Plano']['E23'].value, '|', wb['Painel']['C26'].value)


def v(c):
    """O openpyxl lê como None o texto vazio que a fórmula devolve; e as datas, como datetime."""
    x = c.value
    return '' if x is None else (x.date() if hasattr(x, 'date') else x)


def e(x):
    return '' if x is None else x


for aba, ex in (('Exemplo 1 - Pizzaria', EX1), ('Exemplo 2 - Indústria', EX2)):
    s = wb[aba]
    ref, pl = ex['head']['ref'], ex['plano']
    got_d = [(v(s[f'O{r}']), v(s[f'P{r}']), v(s[f'Q{r}']), v(s[f'R{r}'])) for r in range(16, 25)]
    esp_d = [(prioridade_pts(m, d['sit']), prioridade(m, d['sit']), acao_sit(d, ref), diag_conf(d)) for m, d in zip(MUDANCAS, ex['diag'])]
    got_p = [(v(s[f'G{r}']), v(s[f'J{r}'])) for r in range(28, 28 + len(PASSOS))]
    esp_p = [(limite_passo(pl['audit'], k), passo_sit(pl, k)) for k in range(len(PASSOS))]
    print(aba, '| diagnóstico', 'OK' if got_d == esp_d else 'DIFERE %s' % got_d, '| plano', 'OK' if got_p == esp_p else 'DIFERE %s' % got_p, '| alerta', s['E12'].value)


def preencher(nome, ref=None, diag=(), livres=(), plano=None, audit_nome='Auditoria', cert=None, check=None):
    w = openpyxl.load_workbook(SRC)
    M, P, K = w['Mudanças'], w['Plano'], w['Checklist']
    if ref:
        M['D4'] = ref
    for k, d in enumerate(list(diag) + list(livres)):
        r = 8 + k
        if 'm' in d:
            for col, val in zip('CDEFGH', d['m']):
                if val not in (None, ''):
                    M[f'{col}{r}'] = val
        for col, key in zip('IJKLMN', ('sit', 'evid', 'acao', 'resp', 'prazo', 'concl')):
            if d[key] not in (None, ''):
                M[f'{col}{r}'] = d[key]
    if plano:
        P['D4'] = audit_nome
        if plano['audit']:
            P['D5'] = plano['audit']
        if plano['ref']:
            P['D6'] = plano['ref']
        P['D7'] = plano['fim']
        if cert:
            P['D8'] = cert
        for k, d in enumerate(plano['feito']):
            if d:
                P[f'F{12 + k}'] = d
    for k, val in enumerate(check or ()):
        if val:
            K[f'D{5 + k}'] = val
    w.save(OUTDIR + '/%s.xlsx' % nome)


preencher('t1', EX1['head']['ref'], EX1['diag'], plano=EX1['plano'], cert=SEMCERT, check=['Sim'] * 12)
preencher('t2', EX2['head']['ref'], EX2['diag'], plano=EX2['plano'], cert=CERT15, check=['Sim'] * 6 + ['Parcial', 'Não'])

# ---- t3: casos de borda
D = date
DK = ('sit', 'evid', 'acao', 'resp', 'prazo', 'concl')
REF3 = D(2029, 6, 1)
diag3 = [dict(zip(DK, r)) for r in [
    (PARCIAL, 'e', '', '', None, None),                          # lacuna sem ação
    (NAOATENDE, '', 'a', 'r', None, None),                       # falta o prazo
    (PARCIAL, 'e', 'a', '', D(2029, 7, 1), None),                # falta o responsável
    (ATENDE, '', '', '', None, None),                            # falta a evidência
    (None, '', '', '', None, None),                              # falta a situação
    (NA, 'e', '', '', None, None),                               # não se aplica
    (NAOATENDE, '', 'a', 'r', D(2029, 6, 1), None),              # prazo no dia da leitura: no prazo
    (NAOATENDE, '', 'a', 'r', D(2029, 5, 31), None),             # atrasada
    (ATENDE, 'e', 'a', 'r', D(2029, 1, 1), D(2029, 2, 1)),       # concluída fora do prazo: concluída
]]
livres3 = [dict(m=m, **dict(zip(DK, r))) for m, r in [
    (('X1', '8.5', 'Outra mudança', 'Novo', 'Alto', 'p'), (NAOATENDE, '', 'a', 'r', D(2029, 7, 1), None)),   # 6 pontos
    (('X2', '9.1', 'Mais uma', 'Reforçado', 'Médio', 'p'), (PARCIAL, 'e', 'a', 'r', D(2029, 7, 1), None)),   # 2 pontos
    (('X3', '10', 'Sem impacto', 'Novo', '', 'p'), (PARCIAL, 'e', 'a', 'r', D(2029, 7, 1), None)),           # sem impacto: sem pontos
]]
# auditoria depois do fim da transição; limites atrasados, vencendo e previstos; uma etapa feita no dia do limite
plano3 = dict(audit=D(2029, 10, 15), ref=REF3, fim=FIM_TRANSICAO,
              feito=[D(2029, 4, 18), D(2029, 6, 1), None, None, None, None, None, None])
preencher('t3', REF3, diag3, livres3, plano3, cert=CERT15)
print('ok')
