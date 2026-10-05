"""Confere a postura da SWOT nas duas abas de exemplo de SWOT-modelo.xlsx contra a conta feita em Python (item T29).

Uso: swot_confere_postura.py <SWOT-modelo.xlsx gerado e recalculado pelo LibreOffice>

A SWOT não tem par de testes. Este roteiro roda build_swot.py numa pasta temporária só para ler os exemplos (EX1 e EX2),
calcula a média dos pontos por fator de cada quadrante, os dois balanços e a postura, e compara com as células
recalculadas da planilha: médias na linha 43, balanços em H45 e H46, postura em H47. Confere também que a aba SWOT vazia
pede as notas. Sai com 1 se alguma célula difere.
"""
import contextlib
import io
import os
import runpy
import sys
import tempfile

import openpyxl

AQUI = os.path.dirname(os.path.abspath(__file__))
GERADORES = os.path.dirname(AQUI)
POSTURA = {(True, True): "Desenvolvimento", (True, False): "Manutenção", (False, True): "Crescimento", (False, False): "Sobrevivência"}
COLUNA_MEDIA = {"S": "E", "W": "H", "O": "J", "T": "K"}


def main(xlsx):
    argv = sys.argv
    sys.argv = ["build_swot.py", os.path.join(tempfile.mkdtemp(), "SWOT-modelo.xlsx")]
    try:
        with contextlib.redirect_stdout(io.StringIO()):
            dados = runpy.run_path(os.path.join(GERADORES, "build_swot.py"))
    finally:
        sys.argv = argv
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    falhas = 0
    for aba, ex in (("Exemplo 1 - Pizzaria", dados["EX1"]), ("Exemplo 2 - Compras", dados["EX2"])):
        ws = wb[aba]
        med = {q: sum(i * n for _, i, n in ex["fatores"][q]) / len(ex["fatores"][q]) for q in "SWOT"}
        esperado = POSTURA[(med["S"] >= med["W"], med["O"] >= med["T"])]
        cel = {q: ws[COLUNA_MEDIA[q] + "43"].value for q in "SWOT"}
        ok_med = all(isinstance(cel[q], (int, float)) and abs(cel[q] - med[q]) < 1e-9 for q in "SWOT")
        bi, be = ws["H45"].value, ws["H46"].value
        ok_bal = (isinstance(bi, (int, float)) and isinstance(be, (int, float))
                  and abs(bi - (med["S"] - med["W"])) < 1e-9 and abs(be - (med["O"] - med["T"])) < 1e-9)
        postura = ws["H47"].value
        ok_post = isinstance(postura, str) and postura.startswith(esperado + ":")
        print(aba)
        print("   médias", cel, "| Python", {q: round(med[q], 4) for q in "SWOT"}, "OK" if ok_med else "DIFERE")
        print("   balanços", bi, be, "OK" if ok_bal else "DIFERE")
        print("   postura", repr(postura), "| Python", esperado, "OK" if ok_post else "DIFERE")
        falhas += (not ok_med) + (not ok_bal) + (not ok_post)
    vazia = wb["SWOT"]["H47"].value
    ok_vazia = vazia == "Dê notas aos quatro quadrantes para ver a postura"
    print("aba SWOT vazia:", repr(vazia), "OK" if ok_vazia else "DIFERE")
    falhas += not ok_vazia
    print("falhas", falhas)
    return 1 if falhas else 0


if __name__ == "__main__":
    if len(sys.argv) != 2:
        print(__doc__)
        sys.exit(2)
    sys.exit(main(sys.argv[1]))
