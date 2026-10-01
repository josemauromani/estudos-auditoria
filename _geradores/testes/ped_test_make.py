"""Confere Pedidos-modelo.xlsx e preenche cópias com dados de teste.  Uso: ped_test_make.py <modelo.xlsx> <pasta de saída>"""
import openpyxl, sys, os
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
from ped_data import ACEITO, ALTERADO, EMANALISE, EX1, EX2, NA, NAO, RECUSADO, SIM, conf, leitura, mud_conf, oferta_conf
from datetime import date
SRC, OUTDIR = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(SRC, data_only=True)
wf0 = openpyxl.load_workbook(SRC)
for ws in wb.worksheets:
    f = wf0[ws.title]
    print('==', ws.title, ws.dimensions, 'dv', len(f.data_validations.dataValidation), 'cf', sum(len(v) for v in f.conditional_formatting._cf_rules.values()),
          'charts', len(f._charts), 'comments', sum(1 for r in f.iter_rows() for c in r if c.comment))
err = [(ws.title, c.coordinate, c.value) for ws in wb.worksheets for r in ws.iter_rows() for c in r
       if isinstance(c.value, str) and ((c.value.startswith('#') and len(c.value) > 1) or 'Err:' in c.value)]
print('erros', err)
print('vazio:', wb['Oferta']['D21'].value, '|', wb['Pedidos']['E72'].value, '|', wb['Mudanças']['E41'].value, '|', wb['Painel']['C30'].value, '|', wb['Checklist']['D24'].value)
for aba, ex in (('Exemplo 1 - Pizzaria', EX1), ('Exemplo 2 - Indústria', EX2)):
    s = wb[aba]
    nums = [p['num'] for p in ex['peds']]
    rp = [r for r in range(1, s.max_row + 1) if s[f'B{r}'].value in nums]
    got_p = [(s[f'N{r}'].value, s[f'T{r}'].value) for r in rp]
    esp_p = [(leitura(p), conf(p)) for p in ex['peds']]
    itens = [o['item'] for o in ex['ofertas']]
    ro = [r for r in range(1, s.max_row + 1) if s[f'B{r}'].value in itens]
    got_o = [s[f'T{r}'].value for r in ro]
    rm = [r for r in range(1, s.max_row + 1) if s[f'D{r}'].value in [m['oque'] for m in ex['muds']]]
    got_m = [s[f'T{r}'].value for r in rm]
    ok = (got_p == esp_p, got_o == [oferta_conf(o) for o in ex['ofertas']], got_m == [mud_conf(m) for m in ex['muds']])
    print(aba, '| pedidos', 'OK' if ok[0] else 'DIFERE %s' % got_p, '| oferta', 'OK' if ok[1] else 'DIFERE %s' % got_o, '| mudanças', 'OK' if ok[2] else 'DIFERE %s' % got_m)


def preencher(nome, ofertas=(), peds=(), muds=(), check=None):
    w = openpyxl.load_workbook(SRC)
    O, P, M, K = w['Oferta'], w['Pedidos'], w['Mudanças'], w['Checklist']
    for k, row in enumerate(ofertas):
        for col, v in zip('CDEFGH', row):
            if v:
                O[f'{col}{6 + k}'] = v
    for k, row in enumerate(peds):
        for col, v in zip('CDEFGHIJKLMNOPSTUVW', row):
            if v not in (None, ''):
                P[f'{col}{7 + k}'] = v
    for k, row in enumerate(muds):
        for col, v in zip('CDEFGHI', row):
            if v:
                M[f'{col}{7 + k}'] = v
    for k, v in enumerate(check or ()):
        if v:
            K[f'D{5 + k}'] = v
    w.save(OUTDIR + '/%s.xlsx' % nome)


def do_exemplo(ex):
    of = [(o['item'], o['espec'], o['legal'], o['naogar'], o['onde'], o['cap']) for o in ex['ofertas']]
    pe = [(p['num'], p['data'], p['cliente'], p['canal'], p['pede'], p['naodecl'], p['prazo'], *p['q'], p['dec'], p['neg'], p['quem'], p['analise'], p['confirm']) for p in ex['peds']]
    mu = [(m['data'], m['ped'], m['oque'], m['pediu'], m['analise'], m['docs'], m['inform']) for m in ex['muds']]
    return of, pe, mu


D = date
for nome, ex, check in (('t1', EX1, ['Sim'] * 12), ('t2', EX2, ['Sim'] * 6 + ['Parcial', 'Não'])):
    of, pe, mu = do_exemplo(ex)
    preencher(nome, of, pe, mu, check)
S7, d1, d2 = [SIM] * 7, D(2027, 6, 1), D(2027, 6, 2)
peds3 = [('A', d1, 'C', 'Tel', 'ok', '', d2, *S7, ACEITO, '', 'Quem', d1, d1),                                     # OK
         ('B', d1, 'C', 'Tel', 'pendência', '', d2, SIM, NAO, *[SIM] * 5, ACEITO, '', 'Quem', d1, d1),            # aceito com pendência
         ('C', d1, 'C', 'Tel', 'depois', '', d2, *S7, ACEITO, '', 'Quem', d2, d1),                                # analisado depois de aceitar
         ('D', d1, 'C', 'Tel', 'alterado sem negociação', '', d2, NAO, *[SIM] * 6, ALTERADO, '', 'Quem', d1, d1),  # falta o que foi negociado
         ('E', d1, 'C', 'Tel', 'recusado sem motivo', '', d2, NAO, *[SIM] * 6, RECUSADO, '', 'Quem', d1, None),   # falta o motivo
         ('F', d1, 'C', 'Tel', 'incompleto', '', d2, SIM, SIM, '', *[SIM] * 4, ACEITO, '', 'Quem', d1, d1),       # falta responder
         ('G', d1, 'C', 'Tel', 'sem decisão', '', d2, *S7, '', '', 'Quem', d1, None),                             # falta a decisão
         ('H', d1, 'C', 'Tel', 'sem quem', '', d2, *S7, ACEITO, '', '', d1, d1),                                  # falta quem analisou
         ('I', d1, 'C', 'Tel', 'sem data da análise', '', d2, *S7, ACEITO, '', 'Quem', None, d1),                 # falta a data da análise
         ('J', d1, 'C', 'Tel', 'sem confirmação', '', d2, *S7, ACEITO, '', 'Quem', d1, None),                     # falta a data da confirmação
         ('K', d1, 'C', 'Tel', 'em análise incompleto', '', d2, NAO, '', '', '', '', '', '', EMANALISE, 'arte', 'Quem', None, None),  # em análise: OK
         ('L', d1, 'C', 'Tel', 'não se aplica', '', d2, SIM, SIM, NA, SIM, NA, SIM, SIM, ACEITO, '', 'Quem', d1, d1)]  # OK com não se aplica
muds3 = [(d1, 'A', 'ok', 'Cliente', 'Analisada', SIM, 'Produção'), (d1, 'A', 'sem análise', 'Cliente', '', SIM, 'Produção'),
         (d1, 'A', 'sem documentos', 'Cliente', 'Analisada', NAO, 'Produção'), (d1, 'A', 'sem aviso', 'Cliente', 'Analisada', SIM, ''),
         (d1, 'Z', 'pedido inexistente', 'Cliente', 'Analisada', SIM, 'Produção'), (d1, '', 'sem pedido', 'Cliente', 'Analisada', SIM, 'Produção'),
         (d1, 'A', '', 'só quem pediu', '', '', '')]
of3 = [('Item ok', 'Espec', 'Lei', 'Nada', 'Site', SIM), ('Sem especificação', '', 'Lei', '', 'Site', SIM), ('Sem onde', 'Espec', '', '', '', SIM),
       ('Sem capacidade', 'Espec', '', '', 'Site', NAO), ('', 'só a especificação', '', '', '', '')]
preencher('t3', of3, peds3, muds3)
print('ok')
