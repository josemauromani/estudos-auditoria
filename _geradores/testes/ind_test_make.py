"""Preenche cópias de Indicadores-modelo.xlsx com dados de teste.  Uso: ind_test_make.py <modelo.xlsx> <pasta de saída>"""
import openpyxl, sys, os
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from ind_data import EX1, EX2, MESES, situacao, seguidos, tendencia, acao
from datetime import date
SRC, OUTDIR = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(SRC, data_only=True)
wf0 = openpyxl.load_workbook(SRC)
for ws in wb.worksheets:
    f = wf0[ws.title]
    print('==', ws.title, ws.dimensions, 'dv', len(f.data_validations.dataValidation), 'cf', sum(len(v) for v in f.conditional_formatting._cf_rules.values()),
          'charts', len(f._charts), 'comments', sum(1 for r in f.iter_rows() for c in r if c.comment))
err = [(ws.title, c.coordinate, c.value) for ws in wb.worksheets for r in ws.iter_rows() for c in r
       if isinstance(c.value, str) and ((c.value.startswith('#') and len(c.value) > 1 and c.value != '#N/A') or 'Err:' in c.value)]
print('erros', err)
print('vazio:', wb['Fichas']['E25'].value, '|', wb['Medições']['D37'].value, '|', wb['Painel']['D26'].value, '|', wb['Análise']['F29'].value)
for n, ex in (('Exemplo 1 - Pizzaria', EX1), ('Exemplo 2 - Compras', EX2)):
    s = wb[n]
    linhas = [k for k in range(10, 70) if s[f'B{k}'].value in [i['id'] for i in ex['inds']] and s[f'F{k}'].value in ('Na meta', 'Atenção', 'Fora da meta')]
    print(n)
    for k, i in zip(linhas, ex['inds']):
        planilha = (s[f'D{k}'].value, s[f'F{k}'].value, s[f'L{k}'].value, s[f'N{k}'].value)
        esperado = (i['valores'][-1], situacao(i['valores'][-1], i), tendencia(i), acao(i))
        print('   ', i['id'], planilha, 'OK' if (float(planilha[0]), planilha[1:]) == (float(esperado[0]), esperado[1:]) else 'DIFERE de %s' % (esperado,))
    seg = [s[f'Q{k}'].value for k in range(10, 40) if s[f'B{k}'].value in [i['id'] for i in ex['inds']] and isinstance(s[f'Q{k}'].value, (int, float))]
    print('    seguidos', seg, [seguidos(i) for i in ex['inds']])


def preencher(nome, ex, fichas=None, med=None, analise=(), extra=None):
    w = openpyxl.load_workbook(SRC)
    f, m, a, p = w['Fichas'], w['Medições'], w['Análise'], w['Painel']
    f['D4'] = ex['head']['org']; f['K5'] = ex['head']['data']
    for k, mes in enumerate(MESES):
        m.cell(row=6, column=5 + k, value=mes)
    for k, i in enumerate(ex['inds']):
        r = 10 + k
        for col, key in zip('CDEFGHIJKLM', ('nome', 'objetivo', 'processo', 'formula', 'unidade', 'sentido', 'meta', 'limite', 'fonte', 'freq', 'resp')):
            f[f'{col}{r}'] = i[key]
        for j, v in enumerate(i['valores']):
            m.cell(row=7 + k, column=5 + j, value=v)
    for ref, v in (fichas or {}).items():
        f[ref] = v
    for ref, v in (med or {}).items():
        m[ref] = v
    for k, row in enumerate(analise):
        for col, v in zip('CDFGHIJK', row):
            if v is not None:
                a[f'{col}{7 + k}'] = v
    for ref, v in (extra or {}).items():
        p[ref] = v
    w.save(OUTDIR + '/%s.xlsx' % nome)


preencher('t1', EX2, analise=[(date(2026, 10, 8), 1, 'fato', 'causa', 'decisão', 'Gerente', date(2026, 10, 23), 'Em andamento'),
                              (date(2026, 10, 8), 3, 'fato', 'causa', 'decisão', 'Comprador', date(2026, 1, 10), 'Não iniciada'),
                              (date(2026, 10, 8), 2, 'fato', 'causa', 'decisão', None, None, None)], extra={'D36': 3})
preencher('t2', EX1)
preencher('t3', EX1,
          fichas={'J10': 97,            # limite acima da meta, com maior é melhor
                  'D11': None,          # sem objetivo
                  'I12': None,          # sem meta
                  'M13': None},         # sem responsável
          med={'H8': None,              # coluna pulada no indicador 2
               'P6': None,              # último período sem nome
               'I10': None, 'J10': None, 'K10': None, 'L10': None, 'M10': None, 'N10': None, 'O10': None, 'P10': None})   # indicador 4 com 4 medições
