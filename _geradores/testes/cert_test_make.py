"""Confere Certificacao-modelo.xlsx e preenche cópias com dados de teste.  Uso: cert_test_make.py <modelo.xlsx> <pasta de saída>"""
import openpyxl, sys, os
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
from cert_data import (CRITERIOS, E2015, E2026, EVENTOS, EX1, EX2, FIM_TRANSICAO, MAIOR, MENOR, OM, PONTO, SIM, PARCIAL, NAO, ct_conf, ct_sit, evento_sit, limites,
                       prazo_fech, prazo_plano, pront_conf, pront_res, validade)
from datetime import date
SRC, OUTDIR = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(SRC, data_only=True)
wf0 = openpyxl.load_workbook(SRC)
for ws in wb.worksheets:
    f = wf0[ws.title]
    print('==', ws.title, ws.dimensions, 'dv', len(f.data_validations.dataValidation), 'cf', sum(len(v) for v in f.conditional_formatting._cf_rules.values()),
          'charts', len(f._charts), 'comments', sum(1 for r in f.iter_rows() for c in r if c.comment))
err = [(ws.title, c.coordinate, c.value) for ws in wb.worksheets for r in ws.iter_rows() for c in r
       if isinstance(c.value, str) and ((c.value.startswith('#') and len(c.value) > 1) or 'Err:' in c.value)]
print('erros', err)
print('vazio:', wb['Prontidão']['E27'].value, '|', wb['Ciclo']['E24'].value, '|', wb['Constatações']['E51'].value, '|', wb['Painel']['C29'].value)


def v(c):
    """O openpyxl lê como None o texto vazio que a fórmula devolve; e as datas, como datetime."""
    x = c.value
    return '' if x is None else (x.date() if hasattr(x, 'date') else x)


def e(x):
    return '' if x is None else x


for aba, ex, pos in (('Exemplo 1 - Pizzaria', EX1, dict(p=17, res=34, v=38, c=48)), ('Exemplo 2 - Indústria', EX2, dict(p=17, res=34, v=38, c=48))):
    s = wb[aba]
    ref = ex['ciclo']['ref']
    got_p = [v(s[f'L{r}']) for r in range(pos['p'], pos['p'] + len(CRITERIOS))]
    esp_p = [pront_conf(c, i) for c, i in zip(CRITERIOS, ex['pront']['itens'])]
    res = pront_res(ex['pront']['itens'])
    got_r = [v(s[f'D{pos["res"] - k}']) for k in (3, 2, 1, 0)]
    esp_r = [round(res['pct'], 6), res['imp1'], res['imp2'], res['res']]
    got_r[0] = round(got_r[0], 6)
    got_v = [(v(s[f'D{r}']), v(s[f'F{r}'])) for r in range(pos['v'], pos['v'] + len(EVENTOS))]
    esp_v = [(e(l), evento_sit(ex['ciclo'], k)) for k, l in enumerate(limites(ex['ciclo']))]
    got_c = [(v(s[f'L{r}']), v(s[f'M{r}']), v(s[f'N{r}']), v(s[f'O{r}'])) for r in range(pos['c'], pos['c'] + len(ex['cts']))]
    esp_c = [(e(prazo_plano(x)), e(prazo_fech(x)), ct_sit(x, ref), ct_conf(x)) for x in ex['cts']]
    print(aba, '| prontidão', 'OK' if (got_p, got_r) == (esp_p, esp_r) else 'DIFERE %s %s' % (got_p, got_r), '| ciclo', 'OK' if got_v == esp_v else 'DIFERE %s' % got_v,
          '| constatações', 'OK' if got_c == esp_c else 'DIFERE %s' % got_c)


def preencher(nome, data=None, itens=(), ciclo=None, organismo=None, cts=(), check=None):
    w = openpyxl.load_workbook(SRC)
    P, C, T, K = w['Prontidão'], w['Ciclo'], w['Constatações'], w['Checklist']
    if data:
        P['D4'] = data
    for k, i in enumerate(itens):
        for col, key in zip('FGHIJ', ('status', 'evid', 'acao', 'resp', 'prazo')):
            if i[key] not in (None, ''):
                P[f'{col}{7 + k}'] = i[key]
    if ciclo:
        C['D4'] = organismo or 'Organismo'
        C['D5'], C['D7'], C['D8'] = ciclo['edicao'], ciclo['ref'], ciclo['fimtrans']
        if ciclo['decisao']:
            C['D6'] = ciclo['decisao']
        for k, d in enumerate(ciclo['feito']):
            if d and k != 2:
                C[f'E{13 + k}'] = d
    for k, x in enumerate(cts):
        for col, key in zip('CDEFGHIJK', ('num', 'aud', 'data', 'tipo', 'req', 'desc', 'plano', 'concl', 'evid')):
            if x[key] not in (None, ''):
                T[f'{col}{7 + k}'] = x[key]
    for k, val in enumerate(check or ()):
        if val:
            K[f'D{5 + k}'] = val
    w.save(OUTDIR + '/%s.xlsx' % nome)


for nome, ex, check in (('t1', EX1, ['Sim'] * 12), ('t2', EX2, ['Sim'] * 6 + ['Parcial', 'Não'])):
    preencher(nome, ex['pront']['data'], ex['pront']['itens'], ex['ciclo'], ex['head']['organismo'], ex['cts'], check=check)

# ---- t3 e t4: casos de borda
D = date
IK = ('status', 'evid', 'acao', 'resp', 'prazo')
ok = (SIM, 'e', '', '', None)
itens3 = [dict(zip(IK, r)) for r in [(PARCIAL, 'e', '', 'r', D(2029, 9, 30)),   # falta a ação
                                     (NAO, '', 'a', 'r', None),                 # falta o prazo
                                     (SIM, '', '', '', None),                   # falta a evidência
                                     (None, '', '', '', None)] + [ok] * 10]     # falta o status
itens4 = [dict(zip(IK, ok))] * len(CRITERIOS)                                   # tudo em Sim: pronta para a fase 2
ciclo3 = dict(edicao=E2015, decisao=D(2027, 1, 31), ref=D(2029, 9, 1), fimtrans=FIM_TRANSICAO,
              feito=[D(2026, 10, 1), D(2027, 5, 1), D(2027, 1, 31), D(2028, 1, 15), None, None, None])   # fase 2 fora do prazo, 2ª atrasada, recertificação e transição vencendo
ciclo4 = dict(edicao=E2026, decisao=None, ref=D(2029, 9, 1), fimtrans=FIM_TRANSICAO, feito=[None] * 7)   # sem decisão: tudo previsto ou vazio
CK = ('num', 'aud', 'data', 'tipo', 'req', 'desc', 'plano', 'concl', 'evid')
cts3 = [dict(zip(CK, r)) for r in [
    ('a', 'Fase 2', D(2029, 8, 1), MAIOR, '8.5', 'd', None, None, None),                          # plano atrasado
    ('b', 'Fase 2', D(2029, 5, 1), MAIOR, '8.5', 'd', D(2029, 5, 10), None, None),                # maior vencida
    ('c', 'Fase 2', D(2029, 8, 20), MAIOR, '8.5', 'd', D(2029, 8, 25), None, None),               # aberta no prazo
    ('d', 'Fase 2', D(2029, 8, 20), MENOR, '8.5', 'd', D(2029, 8, 25), None, None),               # aguarda verificação
    ('e', 'Fase 2', D(2029, 8, 20), MENOR, '8.5', 'd', None, None, None),                         # aberta, plano ainda no prazo
    ('f', 'Fase 1', D(2029, 8, 20), PONTO, '8.5', 'd', None, None, None),                         # tratar antes da fase 2
    ('g', 'Fase 1', D(2029, 8, 20), PONTO, '8.5', 'd', None, None, D(2029, 8, 28)),               # fechada
    ('h', 'Fase 2', D(2029, 8, 20), OM, '8.5', 'd', None, None, None),                            # para considerar
    ('i', 'Fase 2', D(2029, 8, 20), MAIOR, '8.5', 'd', D(2029, 8, 22), None, D(2029, 8, 30)),     # fechada sem a conclusão da ação
    ('j', 'Fase 2', D(2029, 8, 20), MENOR, '8.5', 'd', D(2029, 8, 10), None, None),               # plano antes da auditoria
    ('k', 'Fase 2', None, MENOR, '8.5', 'd', None, None, None),                                   # falta a data
    ('l', 'Fase 2', D(2029, 8, 20), None, '8.5', 'd', None, None, None),                          # falta o tipo
    ('m', 'Fase 2', D(2029, 8, 20), MENOR, '', 'd', None, None, None),                            # falta o requisito
    ('n', 'Fase 2', D(2029, 8, 20), MENOR, '8.5', '', None, None, None)]]                         # falta a descrição
preencher('t3', D(2029, 9, 1), itens3, ciclo3, None, cts3)
preencher('t4', D(2029, 9, 1), itens4, ciclo4, None, ())
print('ok')
