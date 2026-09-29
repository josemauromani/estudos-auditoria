# -*- coding: utf-8 -*-
"""Gera Processos-modelo.xlsx no mesmo padrão visual das planilhas anteriores da série."""
import math
import os
import sys
from datetime import date

_here = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, _here)
# funções de formatação compartilhadas com o gerador do PDCA (bloco inicial do arquivo)
_src = open(os.path.join(_here, "build_pdca.py"), encoding="utf-8").read()
exec(_src[: _src.index("STATUS = [")])

from openpyxl.chart import BarChart, Series  # noqa: E402
from openpyxl.chart.label import DataLabelList  # noqa: E402
from openpyxl.utils import get_column_letter as L  # noqa: E402
from proc_data import (APOIO, CHECK, ELEMENTOS, EX1, EX2, GESTAO, IND, NAO, ORDEM_PARTES, PARCIAL, PARTES, PRINCIPAL, SIM, TIPOS,  # noqa: E402
                       fornece, recebe)

BLUE, TEAL, AMBER, PURPLE, REDC = "2B5C8A", "1E7B73", "A96A12", "7A4A9A", "B0413E"
BLUE_T, TEAL_T, AMBER_T = "DCE8F3", "D9EEEB", "F6E8CF"
TIPO_COR = {GESTAO: BLUE, PRINCIPAL: AMBER, APOIO: TEAL}
TIPO_TINT = {GESTAO: BLUE_T, PRINCIPAL: AMBER_T, APOIO: TEAL_T}
TIPO_CF = [(t, TIPO_TINT[t]) for t in TIPOS]
YESNO = [SIM, PARCIAL, NAO]
YESNO_CF = [(SIM, GREEN), (PARCIAL, YELLOW), (NAO, RED)]
NP = 12  # processos do modelo
TXT = "@"


def note(ws, ref, text):
    c = Comment(text, "Modelo Processos")
    c.width, c.height = 300, 120
    ws[ref].comment = c


def cf_warn(ws, rng_, first, ok_values=("OK",)):
    cond = ",".join(f'{first}="{v}"' for v in ok_values)
    ws.conditional_formatting.add(rng_, FormulaRule(formula=[f"OR({cond})"], fill=PatternFill("solid", bgColor=GREEN, fgColor=GREEN)))
    ws.conditional_formatting.add(rng_, FormulaRule(formula=[f"AND(LEN({first})>0,NOT(OR({cond})))"],
                                  fill=PatternFill("solid", bgColor=YELLOW, fgColor=YELLOW)))


def hint_row(ws, row, cells, height=30):
    for ref, text, merge in cells:
        put(ws, f"{ref}{row}", text, f=font(9, i=True, c=MUTED), bg=GRAY, h="center", merge=merge and f"{ref}{row}:{merge}{row}")
    ws.row_dimensions[row].height = height


def alt(text, chars, minimo=21.75, linha=12.75):
    n = max(1, math.ceil(len(str(text)) / max(chars, 1)))
    return max(minimo, n * linha + 9)


def dv_inteiro(ws, rng_, maximo, prompt):
    dv = DataValidation(type="whole", operator="between", formula1="1", formula2=str(maximo), allow_blank=True)
    dv.promptTitle, dv.prompt, dv.showInputMessage = "Número", prompt, True
    dv.errorTitle, dv.error, dv.showErrorMessage = "Número inválido", f"Digite um número inteiro de 1 a {maximo}.", True
    ws.add_data_validation(dv)
    dv.add(rng_)


def chart_style(ch):
    ch.style = 2
    ch.title = None
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ch.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill="D3DBE0", w=9525))
    ch.x_axis.graphicalProperties = GraphicalProperties(ln=LineProperties(solidFill="D3DBE0", w=9525))
    ch.y_axis.graphicalProperties = GraphicalProperties(ln=LineProperties(noFill=True))


def barras(ws, anchor, r1, r2, c_cat, c_val, width=24, height=8.5):
    """Barras horizontais: percentual de elementos definidos em cada processo."""
    ch = BarChart()
    ch.type = "bar"
    ch.gapWidth = 60
    ch.height, ch.width = height, width
    chart_style(ch)
    ch.legend = None
    s = Series(Reference(ws, min_col=c_val, min_row=r1, max_row=r2), title="Elementos definidos")
    s.graphicalProperties.solidFill = BLUE
    s.graphicalProperties.line.noFill = True
    s.dLbls = DataLabelList()
    s.dLbls.showVal = True
    s.dLbls.showSerName = s.dLbls.showCatName = s.dLbls.showLegendKey = False
    ch.append(s)
    ch.set_categories(Reference(ws, min_col=c_cat, min_row=r1, max_row=r2))
    ch.y_axis.scaling.min = 0
    ch.y_axis.scaling.max = 1
    ch.y_axis.majorUnit = 0.25
    ch.y_axis.number_format = "0%"
    ch.x_axis.scaling.orientation = "maxMin"
    ws.add_chart(ch, anchor)


def f_pontos(rng_):
    return f'=IF(COUNTA({rng_})=0,"",COUNTIF({rng_},"{SIM}")+0.5*COUNTIF({rng_},"{PARCIAL}"))'


# ================================================================== pasta de trabalho
wb = Workbook()
wb.properties.title = "Mapa de processos e diagrama de tartaruga — Modelo"
wb.properties.creator = "Modelo Processos"
wb.properties.language = "pt-BR"

# ------------------------------------------------------------------ Instruções
ws = wb.active
ws.title = "Instruções"
widths(ws, {"A": 2, "B": 24, "C": 92, "D": 2})
title(ws, "Mapa de processos e tartaruga — Como usar esta planilha",
      "Modelo para listar os processos, descrever as ligações, detalhar um processo e avaliar os oito elementos.", "C")
r = 4


def section(text):
    global r
    band(ws, r, text, "C", sz=11)
    r += 1


def line(k, v, kbg=GRAY, vbg=None, kf=None, kh="left", height=None):
    global r
    put(ws, f"B{r}", k, f=kf or font(10, True), bg=kbg, h=kh)
    put(ws, f"C{r}", v, bg=vbg)
    ws.row_dimensions[r].height = height or (19.5 if len(v) <= 98 else (31.5 if len(v) <= 196 else 45.75))
    r += 1


section("Legenda: onde preencher")
line("Amarelo-claro", "Células de entrada. É aqui que você digita.", vbg=INPUT)
line("Cinza", "Células calculadas ou fixas (rótulos, nomes, contagens, avisos). Não altere.", vbg=GRAY)
line("Cinza em itálico", "Linha de exemplo (marcada com “Ex.”). Mostra o formato esperado e não entra nas contagens.", vbg=GRAY)
line("Listas suspensas", "Colunas de tipo e de avaliação aceitam apenas as opções da lista.")
line("Comentários", "Células com triângulo vermelho no canto trazem uma dica de preenchimento. Passe o mouse sobre elas.")
r += 1
section("Os três tipos de processo")
line(GESTAO, "Define o rumo, os objetivos e os recursos, e confere os resultados.", kbg=BLUE, vbg=BLUE_T, kf=font(10, True, c=WHITE))
line(PRINCIPAL, "Transforma o pedido do cliente em produto ou serviço entregue.", kbg=AMBER, vbg=AMBER_T, kf=font(10, True, c=WHITE))
line(APOIO, "Fornece os recursos de que os outros processos precisam.", kbg=TEAL, vbg=TEAL_T, kf=font(10, True, c=WHITE))
r += 1
section("Ordem recomendada de preenchimento")
for k, text in enumerate([
    "Aba Processos: liste os processos, começando pelos principais, com tipo, dono e objetivo.",
    "Aba Processos: escreva as principais entradas e saídas, o cliente e o indicador de cada processo.",
    "Aba Interações: escreva, na linha de cada processo, o que ele entrega aos processos das colunas.",
    "Aba Elementos: avalie os oito elementos de cada processo.",
    "Aba Tartaruga: informe o número de um processo e preencha as sete partes.",
    "Aba Tartaruga: confira se as entradas e as saídas aparecem na aba Interações.",
    "Aba Checklist: valide o mapa com quem executa os processos.",
], 1):
    line(f"Passo {k}", text)
r += 1
section("Abas da planilha")
for k, text in [
    ("Processos", "Lista de até 12 processos, com conferência dos campos e contagem por tipo."),
    ("Interações", "Matriz com o que cada processo entrega aos outros, e situação de cada processo."),
    ("Tartaruga", "As sete partes de um processo, com conferência de cada parte."),
    ("Elementos", "Avaliação dos oito elementos de cada processo, com percentual e gráfico."),
    ("Checklist", "Doze verificações de qualidade do mapa, com percentual de conclusão."),
    ("Exemplo 1 - Pizzaria", "Lista de processos, avaliação dos elementos e tartaruga do processo de entrega."),
    ("Exemplo 2 - Compras", "Tartaruga do processo de aquisição e interações com os outros processos."),
]:
    line(k, text)
r += 1
section("Regras e premissas do modelo")
for k, text in [
    ("Número do processo", "As abas Interações, Tartaruga e Elementos usam o número da linha da aba Processos. Não troque a ordem dos processos depois de preencher as outras abas."),
    ("Interações", "A matriz mostra as ligações entre processos. O primeiro processo principal recebe do cliente, e o último entrega ao cliente: nos dois casos, a situação aparece em amarelo, e não é um erro."),
    ("Elementos", "São os oito itens do requisito 4.4.1 da ISO 9001:2015, resumidos com palavras próprias. “Sim” vale 1 ponto, e “Parcial”, meio. A pontuação é uma convenção deste modelo."),
    ("Uma tartaruga por vez", "A aba Tartaruga serve a um processo. Para o processo seguinte, duplique a aba ou salve uma cópia do arquivo."),
    ("Três tipos", "A divisão em gestão, principais e apoio é a mais usada. A norma não a exige."),
    ("Exemplos", "Os dados das abas de exemplo são ilustrativos, criados para este material."),
]:
    line(k, text)
setup(ws, INK, f"B1:C{r}", landscape=False, fit_height=True)

# ------------------------------------------------------------------ Processos
ws = wb.create_sheet("Processos")
widths(ws, {"A": 2, "B": 5, "C": 30, "D": 13, "E": 22, "F": 38, "G": 32, "H": 32, "I": 22, "J": 28, "K": 26, "L": 2})
title(ws, "Lista de processos", "Um processo por linha. Use verbo no infinitivo no nome: registrar o pedido, produzir, entregar.", "K")
for rr, l1, l2, f2 in [(4, "Organização", "Data", DATE), (5, "O que o cliente pede", "O que o cliente recebe", None)]:
    label(ws, f"B{rr}", l1, merge=f"B{rr}:C{rr}")
    inp(ws, f"D{rr}", merge=f"D{rr}:F{rr}")
    label(ws, f"G{rr}", l2)
    inp(ws, f"H{rr}", merge=f"H{rr}:K{rr}", fmt=f2, h="left")
    ws.row_dimensions[rr].height = 24
dv_date(ws, "H4")
for col, text in zip("BCDEFGHIJK", ["#", "Processo", "Tipo", "Dono", "Objetivo", "Principais entradas", "Principais saídas", "Cliente do processo",
                                     "Indicador", "Conferência"]):
    head(ws, f"{col}7", text)
ws.row_dimensions[7].height = 24
hint_row(ws, 8, [("B", "", None), ("C", "Verbo no infinitivo", None), ("D", "Escolha na lista", None), ("E", "Quem responde pelo resultado", None),
                 ("F", "O que o processo deve entregar", None), ("G", "O que recebe", None), ("H", "O que entrega", None),
                 ("I", "Quem recebe a saída", None), ("J", "Como o resultado é medido", None), ("K", "Campos da linha", None)])
ex(ws, "B9", "Ex.", h="center")
for col, v, h_ in [("C", "Entregar o pedido", "left"), ("D", PRINCIPAL, "center"), ("E", "Líder da expedição", "left"),
                   ("F", "Entregar o pedido certo, quente e no prazo.", "left"), ("G", "Pedido embalado e endereço completo", "left"),
                   ("H", "Pedido entregue e horário de chegada", "left"), ("I", "Cliente", "left"), ("J", "Entregas em até 40 minutos", "left"),
                   ("K", "Completo", "center")]:
    ex(ws, f"{col}9", v, h=h_)
ws.row_dimensions[9].height = 33
R1 = 10
R2 = R1 + NP - 1
for k in range(NP):
    rr = R1 + k
    num(ws, f"B{rr}", k + 1)
    inp(ws, f"C{rr}")
    inp(ws, f"D{rr}", h="center")
    for col in "EFGHIJ":
        inp(ws, f"{col}{rr}")
    calc(ws, f"K{rr}", f'=IF(C{rr}="","",IF(D{rr}="","Falta o tipo",IF(E{rr}="","Falta o dono",IF(F{rr}="","Falta o objetivo",'
         f'IF(OR(G{rr}="",H{rr}=""),"Faltam entradas ou saídas",IF(I{rr}="","Falta o cliente do processo",IF(J{rr}="","Falta o indicador","Completo")))))))',
         b=False, sz=9)
    ws.row_dimensions[rr].height = 36
dv_list(ws, f"D{R1}:D{R2}", TIPOS, "Gestão, Principal ou Apoio")
cf_equal(ws, f"D{R1}:D{R2}", TIPO_CF)
cf_warn(ws, f"K{R1}:K{R2}", f"K{R1}", ok_values=("Completo",))
note(ws, "C7", 'Um processo transforma uma entrada em uma saída. Setor, tarefa e documento não são processos.\nEx.: "Comprar e armazenar insumos".')
note(ws, "E7", "O dono responde pelo resultado do processo, mesmo que várias áreas participem.")
s = R2 + 2
band(ws, s, "Resumo automático", "K")
for k, (text, formula) in enumerate([
    ("Processos cadastrados", f"=COUNTA(C{R1}:C{R2})"),
    ("De gestão", f'=COUNTIF(D{R1}:D{R2},"{GESTAO}")'),
    ("Principais", f'=COUNTIF(D{R1}:D{R2},"{PRINCIPAL}")'),
    ("De apoio", f'=COUNTIF(D{R1}:D{R2},"{APOIO}")'),
    ("Sem indicador", f'=SUMPRODUCT((C{R1}:C{R2}<>"")*(J{R1}:J{R2}=""))'),
    ("Aviso", f'=IF(E{s+1}=0,"Cadastre os processos",IF(E{s+3}=0,"Não há processo principal",IF(OR(E{s+2}=0,E{s+4}=0),"Falta processo de gestão ou de apoio",'
              f'IF(COUNTIF(K{R1}:K{R2},"Falta*")>0,"Há processo com o cadastro incompleto: veja a coluna Conferência","OK"))))'),
], 1):
    label(ws, f"B{s+k}", text, merge=f"B{s+k}:D{s+k}", h="right")
    calc(ws, f"E{s+k}", formula, merge=f"E{s+k}:F{s+k}", sz=9 if text == "Aviso" else 10, b=text != "Aviso")
    ws.row_dimensions[s + k].height = 21.75
cf_warn(ws, f"E{s+6}:F{s+6}", f"E{s+6}")
ws.freeze_panes = "D8"
setup(ws, BLUE, f"B1:K{s+6}", fit_height=True)

# ------------------------------------------------------------------ Interações
ws = wb.create_sheet("Interações")
IC = [L(4 + k) for k in range(NP)]  # colunas D a O
IA, IZ = IC[0], IC[-1]
widths(ws, dict({"A": 2, "B": 5, "C": 30, "P": 11, "Q": 11, "R": 28, "S": 2}, **{c: 15 for c in IC}))
title(ws, "Matriz de interações", "Na linha de cada processo, escreva o que ele entrega ao processo de cada coluna. Deixe em branco quando não houver entrega.", "R")
head(ws, "B4", "#")
head(ws, "C4", "Quem entrega")
put(ws, f"{IA}4", "Quem recebe", f=font(10, True, c=WHITE), bg=INK, h="center", merge=f"{IA}4:{IZ}4")
head(ws, "P4", "Entrega a")
head(ws, "Q4", "Recebe de")
head(ws, "R4", "Situação")
ws.row_dimensions[4].height = 21.75
put(ws, "B5", "", bg=GRAY)
put(ws, "C5", "Os nomes vêm da aba Processos", f=font(9, i=True, c=MUTED), bg=GRAY, h="center")
for k, c in enumerate(IC):
    calc(ws, f"{c}5", f'={k+1}&IF(Processos!C{R1+k}="",""," · "&Processos!C{R1+k})', sz=9)
for c, t in (("P", "Processos"), ("Q", "Processos"), ("R", "Calculada")):
    put(ws, f"{c}5", t, f=font(9, i=True, c=MUTED), bg=GRAY, h="center")
ws.row_dimensions[5].height = 42
I1 = 6
I2 = I1 + NP - 1
for k in range(NP):
    rr = I1 + k
    num(ws, f"B{rr}", k + 1)
    calc(ws, f"C{rr}", f'=IF(Processos!C{R1+k}="","",Processos!C{R1+k})', h="left", b=False)
    for j, c in enumerate(IC):
        if j == k:
            put(ws, f"{c}{rr}", "", bg="D9DEE2")
        else:
            inp(ws, f"{c}{rr}")
            ws[f"{c}{rr}"].font = font(9)
    calc(ws, f"P{rr}", f'=IF(C{rr}="","",COUNTA({IA}{rr}:{IZ}{rr}))', b=False)
    calc(ws, f"Q{rr}", f'=IF(C{rr}="","",COUNTA(INDEX(${IA}${I1}:${IZ}${I2},0,{k+1})))', b=False)
    calc(ws, f"R{rr}", f'=IF(C{rr}="","",IF(AND(P{rr}=0,Q{rr}=0),"Sem ligação",IF(P{rr}=0,"Não entrega a outro processo",'
         f'IF(Q{rr}=0,"Não recebe de outro processo","Ligado"))))', b=False, sz=9)
    ws.row_dimensions[rr].height = 39
for cond, color in ((f'R{I1}="Ligado"', GREEN), (f'R{I1}="Sem ligação"', RED), (f"LEN(R{I1})>0", YELLOW)):
    ws.conditional_formatting.add(f"R{I1}:R{I2}", FormulaRule(formula=[cond], stopIfTrue=True, fill=PatternFill("solid", bgColor=color, fgColor=color)))
note(ws, f"{IA}4", 'Escreva o que é entregue, com o requisito, se houver.\nEx.: "Pedido embalado e conferido".')
s = I2 + 2
band(ws, s, "Resumo automático", "R")
ORF = "+".join(f'IF(Processos!C{R1+k}="",COUNTA({IA}{I1+k}:{IZ}{I1+k})+COUNTA({IC[k]}{I1}:{IC[k]}{I2}),0)' for k in range(NP))
for k, (text, formula) in enumerate([
    ("Ligações registradas", f"=COUNTA({IA}{I1}:{IZ}{I2})"),
    ("Processos ligados", f'=COUNTIF(R{I1}:R{I2},"Ligado")'),
    ("Processos sem ligação", f'=COUNTIF(R{I1}:R{I2},"Sem ligação")'),
    ("Processos que só entregam ou só recebem", f'=COUNTIF(R{I1}:R{I2},"Não*")'),
    ("Aviso", f'=IF(COUNTA(Processos!C{R1}:C{R2})=0,"Cadastre os processos na aba Processos",IF(D{s+1}=0,"Registre as ligações",'
              f'IF(({ORF})>0,"Há ligação com processo sem cadastro",IF(D{s+3}>0,"Há processo sem ligação: confira o mapa","OK"))))'),
], 1):
    label(ws, f"B{s+k}", text, merge=f"B{s+k}:C{s+k}", h="right")
    calc(ws, f"D{s+k}", formula, merge=f"D{s+k}:H{s+k}", sz=9 if text == "Aviso" else 10, b=text != "Aviso")
    ws.row_dimensions[s + k].height = 21.75
cf_warn(ws, f"D{s+5}:H{s+5}", f"D{s+5}")
ws.freeze_panes = f"{IA}6"
setup(ws, AMBER, f"B1:R{s+5}", fit_height=True)

# ------------------------------------------------------------------ Tartaruga
ws = wb.create_sheet("Tartaruga")
widths(ws, {"A": 2, "B": 5, "C": 34, "D": 28, "E": 2, "F": 5, "G": 34, "H": 28, "I": 2})
title(ws, "Diagrama de tartaruga", "Informe o número do processo e preencha as sete partes. Escreva o que existe hoje, e marque o que falta.", "H")
label(ws, "B4", "Processo nº, de 1 a 12", merge="B4:C4")
inp(ws, "D4", h="center")
label(ws, "F4", "Data", merge="F4:G4")
inp(ws, "H4", h="left", fmt=DATE)
dv_inteiro(ws, "D4", NP, "Número do processo, como na aba Processos.")
dv_date(ws, "H4")
PR = lambda col: f"INDEX(Processos!${col}${R1}:${col}${R2},$D$4)"  # noqa: E731
for rr, rot, f_ in [(5, "Processo", f'=IF($D$4="","",IF({PR("C")}="","Processo sem cadastro",{PR("C")}))'),
                    (6, "Tipo e dono", f'=IF($D$4="","",IF({PR("C")}="","",{PR("D")}&" · "&{PR("E")}))'),
                    (7, "Objetivo", f'=IF($D$4="","",IF({PR("C")}="","",{PR("F")}&""))')]:
    label(ws, f"B{rr}", rot, merge=f"B{rr}:C{rr}")
    calc(ws, f"D{rr}", f_, h="left", b=rr == 5, merge=f"D{rr}:H{rr}")
    ws.row_dimensions[rr].height = 24
ws.row_dimensions[4].height = 24
COLS = {"entradas": ("Entrada", "De onde vem"), "saidas": ("Saída", "Para onde vai"), "oque": ("Recurso", "Condição ou quantidade"),
        "quem": ("Função", "Competência necessária"), "como": ("Documento ou método", "Código, revisão ou assunto"),
        "quanto": ("Indicador", "Meta"), "riscos": ("Risco", "Resposta")}
COR_PARTE = {"entradas": AMBER, "saidas": AMBER, "oque": TEAL, "quem": TEAL, "como": TEAL, "quanto": TEAL, "riscos": REDC}
POS = {"entradas": ("B", 9), "saidas": ("F", 9), "oque": ("B", 16), "quem": ("F", 16), "como": ("B", 23), "quanto": ("F", 23), "riscos": ("B", 30)}
FAIXAS = {}
for k in ORDEM_PARTES:
    c0, r0 = POS[k]
    c1, c2 = chr(ord(c0) + 1), chr(ord(c0) + 2)
    perg, tit, desc = PARTES[k]
    put(ws, f"{c0}{r0}", f"{tit} · {perg.capitalize()}", f=font(10, True, c=WHITE), bg=COR_PARTE[k], box=False, merge=f"{c0}{r0}:{c2}{r0}")
    ws.row_dimensions[r0].height = 21.75
    put(ws, f"{c0}{r0+1}", "#", f=font(10, True), bg=GRAY, h="center")
    put(ws, f"{c1}{r0+1}", COLS[k][0], f=font(10, True), bg=GRAY, h="center")
    put(ws, f"{c2}{r0+1}", COLS[k][1], f=font(10, True), bg=GRAY, h="center")
    ws.row_dimensions[r0 + 1].height = 21.75
    for j in range(4):
        rr = r0 + 2 + j
        num(ws, f"{c0}{rr}", j + 1)
        inp(ws, f"{c1}{rr}")
        inp(ws, f"{c2}{rr}")
        ws.row_dimensions[rr].height = 30
    FAIXAS[k] = (f"{c1}{r0+2}:{c1}{r0+5}", f"{c2}{r0+2}:{c2}{r0+5}")
    note(ws, f"{c0}{r0}", desc)
put(ws, "F30", "Conferência automática", f=font(10, True, c=WHITE), bg=INK, box=False, merge="F30:H30")
for j, k in enumerate(ORDEM_PARTES):
    rr = 31 + j
    a, b = FAIXAS[k]
    label(ws, f"F{rr}", PARTES[k][1], merge=f"F{rr}:G{rr}")
    calc(ws, f"H{rr}", f'=IF(COUNTA({a})=0,"Falta preencher",IF(COUNTA({b})<COUNTA({a}),"Falta o detalhe","Completa"))', b=False, sz=9)
    ws.row_dimensions[rr].height = 21.75 if rr > 35 else 30
label(ws, "F38", "Situação", merge="F38:G38")
calc(ws, "H38", '=IF($D$4="","Informe o processo",IF(D5="Processo sem cadastro","Processo sem cadastro",IF(COUNTIF(H31:H37,"Completa")=7,'
     '"Tartaruga completa","Há partes incompletas")))', sz=9)
ws.row_dimensions[38].height = 21.75
cf_warn(ws, "H31:H37", "H31", ok_values=("Completa",))
cf_warn(ws, "H38", "H38", ok_values=("Tartaruga completa",))
setup(ws, TEAL, "B1:H38", landscape=False, fit_height=True)

# ------------------------------------------------------------------ Elementos
ws = wb.create_sheet("Elementos")
EC = [L(4 + k) for k in range(8)]  # colunas D a K
EA, EZ = EC[0], EC[-1]
widths(ws, dict({"A": 2, "B": 5, "C": 32, "L": 10, "M": 12, "N": 22, "O": 2}, **{c: 19 for c in EC}))
title(ws, "Os oito elementos de cada processo", "Avalie cada elemento com Sim, Parcial ou Não. Os nomes dos processos vêm da aba Processos.", "N")
head(ws, "B4", "#")
head(ws, "C4", "Processo")
for c, (letra, nome, perg) in zip(EC, ELEMENTOS):
    head(ws, f"{c}4", f"{letra} · {nome}")
    note(ws, f"{c}4", f"ISO 9001:2015, 4.4.1 {letra}, em resumo.\n{perg}")
head(ws, "L4", "Pontos")
head(ws, "M4", "Percentual")
head(ws, "N4", "Conferência")
ws.row_dimensions[4].height = 42
E1 = 5
E2 = E1 + NP - 1
for k in range(NP):
    rr = E1 + k
    num(ws, f"B{rr}", k + 1)
    calc(ws, f"C{rr}", f'=IF(Processos!C{R1+k}="","",Processos!C{R1+k})', h="left", b=False)
    for c in EC:
        inp(ws, f"{c}{rr}", h="center")
    calc(ws, f"L{rr}", f_pontos(f"{EA}{rr}:{EZ}{rr}"), fmt="0.0")
    calc(ws, f"M{rr}", f'=IF(L{rr}="","",L{rr}/8)', fmt="0%")
    calc(ws, f"N{rr}", f'=IF(AND(C{rr}="",COUNTA({EA}{rr}:{EZ}{rr})=0),"",IF(C{rr}="","Processo sem cadastro",IF(COUNTA({EA}{rr}:{EZ}{rr})=0,"Não avaliado",'
         f'IF(COUNTA({EA}{rr}:{EZ}{rr})<8,"Avaliação incompleta","Completa"))))', b=False, sz=9)
    ws.row_dimensions[rr].height = 27
dv_list(ws, f"{EA}{E1}:{EZ}{E2}", YESNO, "Sim, Parcial ou Não")
cf_equal(ws, f"{EA}{E1}:{EZ}{E2}", YESNO_CF)
cf_warn(ws, f"N{E1}:N{E2}", f"N{E1}", ok_values=("Completa",))
ws.conditional_formatting.add(f"M{E1}:M{E2}", FormulaRule(formula=[f"AND(ISNUMBER(M{E1}),M{E1}>=0.8)"], fill=PatternFill("solid", bgColor=GREEN, fgColor=GREEN)))
ws.conditional_formatting.add(f"M{E1}:M{E2}", FormulaRule(formula=[f"AND(ISNUMBER(M{E1}),M{E1}<0.5)"], fill=PatternFill("solid", bgColor=RED, fgColor=RED)))
ws.conditional_formatting.add(f"M{E1}:M{E2}", FormulaRule(formula=[f"ISNUMBER(M{E1})"], fill=PatternFill("solid", bgColor=YELLOW, fgColor=YELLOW)))
t = E2 + 1
for j, (rot, f_) in enumerate([(SIM, f'=COUNTIF({{c}}{E1}:{{c}}{E2},"{SIM}")'), (PARCIAL, f'=COUNTIF({{c}}{E1}:{{c}}{E2},"{PARCIAL}")'),
                               (NAO, f'=COUNTIF({{c}}{E1}:{{c}}{E2},"{NAO}")'),
                               ("Percentual do elemento", f'=IF(COUNTA({{c}}{E1}:{{c}}{E2})=0,"",({{c}}{t}+0.5*{{c}}{t+1})/COUNTA({{c}}{E1}:{{c}}{E2}))')]):
    rr = t + j
    label(ws, f"B{rr}", rot, merge=f"B{rr}:C{rr}", h="right")
    for c in EC:
        calc(ws, f"{c}{rr}", f_.replace("{c}", c), b=j == 3, fmt="0%" if j == 3 else None)
    for c in "LMN":
        put(ws, f"{c}{rr}", "", bg=GRAY)
    ws.row_dimensions[rr].height = 21.75
s = t + 5
band(ws, s, "Resumo automático", "N")
for k, (text, formula, fmt) in enumerate([
    ("Processos avaliados", f"=COUNT(L{E1}:L{E2})", None),
    ("Média dos processos", f'=IF(D{s+1}=0,"",AVERAGE(M{E1}:M{E2}))', "0%"),
    ("Processo com o menor percentual", f'=IF(D{s+1}=0,"",INDEX(C{E1}:C{E2},MATCH(MIN(M{E1}:M{E2}),M{E1}:M{E2},0)))', None),
    ("Elemento com o menor percentual", f'=IF(D{s+1}=0,"",INDEX({EA}4:{EZ}4,MATCH(MIN({EA}{t+3}:{EZ}{t+3}),{EA}{t+3}:{EZ}{t+3},0)))', None),
    ("Aviso", f'=IF(D{s+1}=0,"Avalie os elementos",IF(COUNTIF(N{E1}:N{E2},"Processo sem cadastro")>0,"Há avaliação de processo sem cadastro",'
              f'IF(COUNTIF(N{E1}:N{E2},"Avaliação incompleta")>0,"Há processo com a avaliação incompleta",'
              f'IF(COUNTIF(N{E1}:N{E2},"Não avaliado")>0,"Há processo sem avaliação","OK"))))', None),
], 1):
    label(ws, f"B{s+k}", text, merge=f"B{s+k}:C{s+k}", h="right")
    calc(ws, f"D{s+k}", formula, fmt=fmt, merge=f"D{s+k}:H{s+k}", sz=9 if text == "Aviso" else 10, b=text != "Aviso")
    ws.row_dimensions[s + k].height = 21.75
cf_warn(ws, f"D{s+5}:H{s+5}", f"D{s+5}")
barras(ws, f"B{s+7}", E1, E2, 3, 13, width=30, height=9)
ws.freeze_panes = "D5"
setup(ws, PURPLE, f"B1:N{s+26}", fit_height=True)

# ------------------------------------------------------------------ Checklist
ws = wb.create_sheet("Checklist")
widths(ws, {"A": 2, "B": 5, "C": 70, "D": 14, "E": 44, "F": 2})
title(ws, "Checklist de validação do mapa", "Escolha o status de cada item na lista suspensa (coluna em amarelo-claro).", "E")
for col, text in zip("BCDE", ["#", "Item de verificação", "Status", "Observações"]):
    head(ws, f"{col}4", text)
ws.row_dimensions[4].height = 21.75
for k, text in enumerate(CHECK):
    rr = 5 + k
    num(ws, f"B{rr}", k + 1)
    put(ws, f"C{rr}", text, bg=GRAY)
    inp(ws, f"D{rr}", h="center")
    inp(ws, f"E{rr}")
    ws.row_dimensions[rr].height = 30
dv_list(ws, "D5:D16", YESNO, "Sim, Parcial ou Não")
cf_equal(ws, "D5:D16", YESNO_CF)
band(ws, 18, "Resumo automático", "E")
summary(ws, 19, "Itens atendidos (Sim)", '=COUNTIF(D5:D16,"Sim")', "C", "D")
summary(ws, 20, "Itens parciais", '=COUNTIF(D5:D16,"Parcial")', "C", "D")
summary(ws, 21, "Itens não atendidos", '=COUNTIF(D5:D16,"Não")', "C", "D")
summary(ws, 22, "Itens sem resposta", "=COUNTBLANK(D5:D16)", "C", "D")
summary(ws, 23, "Percentual atendido", "=D19/ROWS(D5:D16)", "C", "D", fmt="0%")
summary(ws, 24, "Situação", '=IF(D19=ROWS(D5:D16),"Validado",IF(D22>0,"Em andamento","Com pendências"))', "C", "D")
setup(ws, PURPLE, "B1:E24", landscape=False, fit_height=True)


# ------------------------------------------------------------------ Exemplos
def bloco_tartaruga(ws, rr, t, ultima):
    """Tabela com as sete partes: parte, item e detalhe. Devolve a linha seguinte e as linhas de cada parte."""
    put(ws, f"B{rr}", "Parte", f=font(10, True), bg=GRAY, h="center", merge=f"B{rr}:C{rr}")
    put(ws, f"D{rr}", "Item", f=font(10, True), bg=GRAY, h="center", merge=f"D{rr}:F{rr}")
    put(ws, f"G{rr}", "De onde vem, para onde vai, condição, meta ou resposta", f=font(10, True), bg=GRAY, h="center", merge=f"G{rr}:{ultima}{rr}")
    ws.row_dimensions[rr].height = 24
    linhas = {}
    for k in ORDEM_PARTES:
        ini = rr + 1
        for a, b in t[k]:
            rr += 1
            put(ws, f"B{rr}", f"{PARTES[k][1]} · {PARTES[k][0].capitalize()}", f=font(10, True, c=WHITE), bg=COR_PARTE[k], merge=f"B{rr}:C{rr}")
            put(ws, f"D{rr}", a, merge=f"D{rr}:F{rr}")
            put(ws, f"G{rr}", b, merge=f"G{rr}:{ultima}{rr}")
            ws.row_dimensions[rr].height = 22.5
        linhas[k] = (ini, rr)
    return rr + 1, linhas


def exemplo1(ws, data):
    XC = [L(8 + k) for k in range(8)]  # colunas H a O
    widths(ws, dict({"A": 2, "B": 6, "C": 30, "D": 12, "E": 21, "F": 36, "G": 26, "P": 9, "Q": 12, "R": 2}, **{c: 15 for c in XC}))
    title(ws, "Mapa de processos e diagrama de tartaruga", "Exemplo preenchido, para consulta. Use as abas Processos, Interações, Tartaruga e Elementos para o seu mapa.", "Q")
    H = data["head"]
    rr = 4
    for rot, val, fmt in [("Organização", H["org"], None), ("Data e participantes", f'{H["data"]:%d/%m/%Y}. {H["por"]}.', None), ("Origem", H["origem"], None),
                          ("O cliente", f'Entrada: {H["cliente_in"].lower()}. Resultado: {H["cliente_out"].lower()}.', None)]:
        label(ws, f"B{rr}", rot, merge=f"B{rr}:C{rr}")
        inp(ws, f"D{rr}", val, merge=f"D{rr}:Q{rr}", bg=WHITE, fmt=fmt)
        ws.row_dimensions[rr].height = 21.75
        rr += 1
    rr += 1
    put(ws, f"B{rr}", "Processos", f=font(10, True, c=WHITE), bg=INK, box=False, merge=f"B{rr}:G{rr}")
    put(ws, f"H{rr}", "Os oito elementos de cada processo", f=font(10, True, c=WHITE), bg=PURPLE, box=False, merge=f"H{rr}:Q{rr}")
    ws.row_dimensions[rr].height = 21.75
    rr += 1
    hd = rr
    for col, text in zip("BCDEFG", ["Nº", "Processo", "Tipo", "Dono", "Objetivo", "Indicador"]):
        put(ws, f"{col}{rr}", text, f=font(10, True), bg=GRAY, h="center")
    for c, (letra, nome, _) in zip(XC, ELEMENTOS):
        put(ws, f"{c}{rr}", f"{letra} · {nome}", f=font(8, True), bg=GRAY, h="center")
    put(ws, f"P{rr}", "Pontos", f=font(10, True), bg=GRAY, h="center")
    put(ws, f"Q{rr}", "Percentual", f=font(10, True), bg=GRAY, h="center")
    ws.row_dimensions[rr].height = 42
    p1 = rr + 1
    for p in data["procs"]:
        rr += 1
        put(ws, f"B{rr}", p["id"], f=font(10, True), bg=GRAY, h="center")
        put(ws, f"C{rr}", p["nome"], f=font(10, True))
        put(ws, f"D{rr}", p["tipo"], h="center")
        put(ws, f"E{rr}", p["dono"])
        put(ws, f"F{rr}", p["objetivo"])
        put(ws, f"G{rr}", p["indicador"] or "Sem indicador", f=font(10, i=not p["indicador"], c=INK if p["indicador"] else MUTED))
        for c, v in zip(XC, p["elem"]):
            put(ws, f"{c}{rr}", v, h="center")
        calc(ws, f"P{rr}", f_pontos(f"H{rr}:O{rr}"), fmt="0.0")
        calc(ws, f"Q{rr}", f"=P{rr}/8", fmt="0%")
        ws.row_dimensions[rr].height = 30
    p2 = rr
    cf_equal(ws, f"D{p1}:D{p2}", TIPO_CF)
    cf_equal(ws, f"H{p1}:O{p2}", YESNO_CF)
    rr += 1
    label(ws, f"B{rr}", "Percentual do elemento", merge=f"B{rr}:G{rr}", h="right")
    for c in XC:
        calc(ws, f"{c}{rr}", f'=(COUNTIF({c}{p1}:{c}{p2},"{SIM}")+0.5*COUNTIF({c}{p1}:{c}{p2},"{PARCIAL}"))/ROWS({c}{p1}:{c}{p2})', fmt="0%")
    put(ws, f"P{rr}", "", bg=GRAY)
    calc(ws, f"Q{rr}", f"=AVERAGE(Q{p1}:Q{p2})", fmt="0%")
    ws.row_dimensions[rr].height = 21.75
    pe = rr
    rr += 2
    band(ws, rr, "Resumo automático", "G")
    s0 = rr
    for k, (text, formula, fmt) in enumerate([
        ("Processos", f"=COUNTA(C{p1}:C{p2})", None),
        ("De gestão, principais e de apoio", f'=COUNTIF(D{p1}:D{p2},"{GESTAO}")&", "&COUNTIF(D{p1}:D{p2},"{PRINCIPAL}")&" e "&COUNTIF(D{p1}:D{p2},"{APOIO}")', None),
        ("Processos sem indicador", f'=COUNTIF(G{p1}:G{p2},"Sem indicador")', None),
        ("Média dos elementos definidos", f"=Q{pe}", "0%"),
        ("Processo com o menor percentual", f"=INDEX(C{p1}:C{p2},MATCH(MIN(Q{p1}:Q{p2}),Q{p1}:Q{p2},0))", None),
        ("Elemento com o menor percentual", f"=INDEX(H{hd}:O{hd},MATCH(MIN(H{pe}:O{pe}),H{pe}:O{pe},0))", None),
    ], 1):
        label(ws, f"B{rr+k}", text, merge=f"B{rr+k}:D{rr+k}", h="right")
        calc(ws, f"E{rr+k}", formula, fmt=fmt, merge=f"E{rr+k}:G{rr+k}")
        ws.row_dimensions[rr + k].height = 21.75
    put(ws, f"H{s0}", "Elementos definidos em cada processo", f=font(10, True, c=WHITE), bg=INK, box=False, merge=f"H{s0}:Q{s0}")
    barras(ws, f"H{s0+1}", p1, p2, 3, 17, width=24, height=8.2)
    rr = s0 + 19
    ws.row_breaks.append(Break(id=rr - 1))
    t = data["tartaruga"]
    px = next(p for p in data["procs"] if p["id"] == t["proc"])
    band(ws, rr, f'Tartaruga do processo {px["id"]} · {px["nome"]} · dono: {px["dono"]}', "Q", color=TEAL)
    rr, _ = bloco_tartaruga(ws, rr + 1, t, "Q")
    setup(ws, MUTED, f"B1:Q{rr - 1}")


def exemplo2(ws, data):
    widths(ws, {"A": 2, "B": 8, "C": 22, "D": 20, "E": 20, "F": 20, "G": 30, "H": 30, "I": 2})
    title(ws, "Diagrama de tartaruga", "Exemplo preenchido, para consulta. Use as abas Tartaruga e Interações para o seu processo.", "H")
    H, P, t = data["head"], data["proc"], data["tartaruga"]
    rr = 4
    for rot, val in [("Organização", H["org"]), ("Processo", f'{P["id"]} · {P["nome"]}'), ("Tipo e dono", f'{P["tipo"]} · {P["dono"]}'), ("Objetivo", P["objetivo"]),
                     ("Data e participantes", f'{H["data"]:%d/%m/%Y}. {H["por"]}.'), ("Origem", H["origem"])]:
        label(ws, f"B{rr}", rot, merge=f"B{rr}:C{rr}")
        inp(ws, f"D{rr}", val, merge=f"D{rr}:H{rr}", bg=WHITE)
        ws.row_dimensions[rr].height = 21.75
        rr += 1
    rr += 1
    band(ws, rr, "As sete partes da tartaruga", "H", color=TEAL)
    rr, linhas = bloco_tartaruga(ws, rr + 1, t, "H")
    rr += 1
    ws.row_breaks.append(Break(id=rr - 1))
    band(ws, rr, f'Interações do processo {P["id"]} com os outros processos', "H", color=AMBER)
    rr += 1
    put(ws, f"B{rr}", "Sentido", f=font(10, True), bg=GRAY, h="center", merge=f"B{rr}:C{rr}")
    put(ws, f"D{rr}", "Processo", f=font(10, True), bg=GRAY, h="center", merge=f"D{rr}:F{rr}")
    put(ws, f"G{rr}", "O que é entregue", f=font(10, True), bg=GRAY, h="center", merge=f"G{rr}:H{rr}")
    ws.row_dimensions[rr].height = 24
    i1 = rr + 1
    n = int(P["id"])
    for sentido, lista in (("Recebe de", recebe(n)), ("Entrega a", fornece(n))):
        for outro, texto in lista:
            rr += 1
            put(ws, f"B{rr}", sentido, f=font(10, True), bg=GRAY if sentido == "Recebe de" else AMBER_T, merge=f"B{rr}:C{rr}")
            put(ws, f"D{rr}", f"{outro} · {IND[outro - 1]}", merge=f"D{rr}:F{rr}")
            put(ws, f"G{rr}", texto, merge=f"G{rr}:H{rr}")
            ws.row_dimensions[rr].height = 22.5
    i2 = rr
    rr += 2
    band(ws, rr, "Resumo automático", "H")
    for k, (text, formula) in enumerate(
            [(f"Itens em {PARTES[p][1].lower()}", f"=COUNTA(D{linhas[p][0]}:D{linhas[p][1]})") for p in ORDEM_PARTES]
            + [("Processos de que recebe", f'=COUNTIF(B{i1}:B{i2},"Recebe de")'), ("Processos a que entrega", f'=COUNTIF(B{i1}:B{i2},"Entrega a")')], 1):
        label(ws, f"B{rr+k}", text, merge=f"B{rr+k}:E{rr+k}", h="right")
        calc(ws, f"F{rr+k}", formula)
        ws.row_dimensions[rr + k].height = 21.75
    setup(ws, MUTED, f"B1:H{rr + 9}", landscape=False)


exemplo1(wb.create_sheet("Exemplo 1 - Pizzaria"), EX1)
exemplo2(wb.create_sheet("Exemplo 2 - Compras"), EX2)

wb.active = 1
wb.save(OUT)
print("ok", OUT)
