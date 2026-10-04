# -*- coding: utf-8 -*-
"""Gera RNC-modelo.xlsx no mesmo padrão visual das planilhas anteriores da série."""
import math
import os
import sys
from datetime import date

_here = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, _here)
# funções de formatação compartilhadas com o gerador do PDCA (bloco inicial do arquivo)
_src = open(os.path.join(_here, "build_pdca.py"), encoding="utf-8").read()
exec(_src[: _src.index("STATUS = [")])

from openpyxl.chart import BarChart, Series  # noqa: E402
from openpyxl.chart.label import DataLabelList  # noqa: E402
from nc_data import AC, CHECK, COR, EX1, EX2  # noqa: E402

BLUE, TEAL, AMBER, PURPLE, REDC = "2B5C8A", "1E7B73", "A96A12", "7A4A9A", "B0413E"
TEAL_T, AMBER_T, RED_T = "D9EEEB", "F6E8CF", "F5DEDC"
YESNO = ["Sim", "Parcial", "Não"]
YESNO_CF = [("Sim", GREEN), ("Parcial", YELLOW), ("Não", RED)]
ORIGENS = ["Auditoria interna", "Auditoria externa", "Reclamação de cliente", "Produto ou serviço não conforme", "Fornecedor",
           "Indicador fora da meta", "Outra"]
ETAPAS = ["Registrada", "Correção feita", "Causa em análise", "Ação em implantação", "Eficácia em observação", "Encerrada"]
TIPOS = [COR, AC]
TIPO_CF = [(COR, YELLOW), (AC, "CDE7E3")]
ST_ACAO = ["Não iniciada", "Em andamento", "Concluída", "Cancelada"]
ST_CF = [("Concluída", GREEN), ("Em andamento", YELLOW), ("Não iniciada", RED)]
SIT_CF = [("Encerrada", GREEN), ("Concluída", GREEN), ("No prazo", GREEN), ("Sem prazo", YELLOW), ("Falta a data", YELLOW),
          ("Confira a etapa", YELLOW), ("Atrasada", RED)]
MOMENTOS = ["Antes", "Durante", "Depois"]
SENTIDOS = ["Menor é melhor", "Maior é melhor"]
PCT = '0"%"'


def note(ws, ref, text):
    c = Comment(text, "Modelo RNC")
    c.width, c.height = 280, 110
    ws[ref].comment = c


def cf_warn(ws, rng_, first, ok_values=("OK",)):
    """Verde para os valores aceitos; amarelo para qualquer outro texto."""
    cond = ",".join(f'{first}="{v}"' for v in ok_values)
    ws.conditional_formatting.add(rng_, FormulaRule(formula=[f"OR({cond})"], fill=PatternFill("solid", bgColor=GREEN, fgColor=GREEN)))
    ws.conditional_formatting.add(rng_, FormulaRule(formula=[f"AND(LEN({first})>0,NOT(OR({cond})))"],
                                  fill=PatternFill("solid", bgColor=YELLOW, fgColor=YELLOW)))


def cf_conclusao(ws, rng_, first):
    for cond, color in ((f'LEFT({first},6)="Eficaz"', GREEN), (f'LEFT({first},3)="Não"', RED), (f"LEN({first})>0", YELLOW)):
        ws.conditional_formatting.add(rng_, FormulaRule(formula=[cond], stopIfTrue=True, fill=PatternFill("solid", bgColor=color, fgColor=color)))


def hint_row(ws, row, cells, height=30):
    for ref, text, merge in cells:
        put(ws, f"{ref}{row}", text, f=font(9, i=True, c=MUTED), bg=GRAY, h="center", merge=merge and f"{ref}{row}:{merge}{row}")
    ws.row_dimensions[row].height = height


def alt(text, chars, minimo=21.75, linha=12.75):
    """Altura da linha para um texto que quebra em uma célula com 'chars' caracteres de largura."""
    n = max(1, math.ceil(len(str(text)) / max(chars, 1)))
    return max(minimo, n * linha + 9)


def chart_style(ch):
    ch.style = 2
    ch.title = None
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ch.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill="D3DBE0", w=9525))
    ch.x_axis.graphicalProperties = GraphicalProperties(ln=LineProperties(solidFill="D3DBE0", w=9525))
    ch.y_axis.graphicalProperties = GraphicalProperties(ln=LineProperties(noFill=True))


def line_chart(ws, anchor, head_row, r1, r2, c_per, c_val, c_meta, width=17, height=7.5):
    ch = LineChart()
    ch.height, ch.width = height, width
    chart_style(ch)
    ch.legend.position = "b"
    ch.append(Series(Reference(ws, min_col=c_val, min_row=r1, max_row=r2), title="Valor"))
    ch.append(Series(Reference(ws, min_col=c_meta, min_row=r1, max_row=r2), title="Meta"))
    ch.set_categories(Reference(ws, min_col=c_per, min_row=r1, max_row=r2))
    s1, s2 = ch.series
    s1.graphicalProperties.line.solidFill = BLUE
    s1.graphicalProperties.line.width = 25400
    s1.marker.symbol = "circle"
    s1.marker.size = 7
    s1.marker.graphicalProperties = GraphicalProperties(solidFill=BLUE)
    s1.marker.graphicalProperties.line.solidFill = WHITE
    s1.smooth = False
    s2.graphicalProperties.line.solidFill = MUTED
    s2.graphicalProperties.line.width = 15875
    s2.graphicalProperties.line.dashStyle = "dash"
    s2.marker.symbol = "none"
    s2.smooth = False
    ch.y_axis.scaling.min = 0
    ws.add_chart(ch, anchor)


# ================================================================== pasta de trabalho
wb = Workbook()
wb.properties.title = "Não conformidade e ação corretiva — Modelo de registro, plano e verificação da eficácia"
wb.properties.creator = "Modelo RNC"
wb.properties.language = "pt-BR"

# ------------------------------------------------------------------ Instruções
ws = wb.active
ws.title = "Instruções"
widths(ws, {"A": 2, "B": 24, "C": 92, "D": 2})
title(ws, "Não conformidade e ação corretiva — Como usar esta planilha",
      "Modelo para registrar a não conformidade, planejar as ações e verificar a eficácia.", "C")
r = 4


def section(text):
    global r
    band(ws, r, text, "C", sz=11)
    r += 1


def line(k, v, kbg=GRAY, vbg=None, kf=None, kh="left", height=None):
    global r
    put(ws, f"B{r}", k, f=kf or font(10, True), bg=kbg, h=kh)
    put(ws, f"C{r}", v, bg=vbg)
    ws.row_dimensions[r].height = height or (19.5 if len(v) <= 98 else (31.5 if len(v) <= 196 else 45.75))
    r += 1


section("Legenda: onde preencher")
line("Amarelo-claro", "Células de entrada. É aqui que você digita.", vbg=INPUT)
line("Cinza", "Células calculadas ou fixas (rótulos, contagens, avisos). Não altere.", vbg=GRAY)
line("Cinza em itálico", "Linha de exemplo (marcada com “Ex.”). Mostra o formato esperado e não entra nas contagens.", vbg=GRAY)
line("Listas suspensas", "Colunas de origem, etapa, tipo, status e momento aceitam apenas as opções da lista. Datas e números são conferidos na digitação.")
line("Comentários", "Células com triângulo vermelho no canto trazem uma dica de preenchimento. Passe o mouse sobre elas.")
r += 1
section("Os três conceitos")
line("Não conformidade", "O não atendimento de um requisito. É descrita com requisito, evidência e declaração.", kbg=REDC, vbg=RED_T,
     kf=font(10, True, c=WHITE))
line("Correção", "Ação para eliminar a não conformidade encontrada. Trata o efeito: o que já saiu errado.", kbg=AMBER, vbg=AMBER_T,
     kf=font(10, True, c=WHITE))
line("Ação corretiva", "Ação para eliminar a causa da não conformidade e evitar que ela se repita. Depende da análise da causa.", kbg=TEAL,
     vbg=TEAL_T, kf=font(10, True, c=WHITE))
r += 1
section("Ordem recomendada de preenchimento")
for k, text in enumerate([
    "Aba Controle: inclua a não conformidade na lista, com a origem, o processo e o responsável.",
    "Aba Registro: descreva a não conformidade e registre a correção e o que foi feito pelo cliente afetado.",
    "Aba Registro: verifique a abrangência e registre a decisão sobre a ação corretiva.",
    "Aba Registro: preencha os porquês, com a forma de confirmação de cada resposta, e a causa raiz.",
    "Aba Plano de ação: escreva as correções e uma ação corretiva para cada causa confirmada.",
    "Aba Eficácia: defina o indicador, a meta e o critério antes de implantar as ações. Depois, registre as medições.",
    "Aba Checklist: valide o tratamento. Abas Registro e Controle: registre o encerramento.",
], 1):
    line(f"Passo {k}", text)
r += 1
section("Abas da planilha")
for k, text in [
    ("Controle", "Lista de todas as não conformidades, com situação, tempo de encerramento e gráfico por origem."),
    ("Registro", "Formulário de uma não conformidade: descrição, correção, abrangência, decisão, causa e encerramento."),
    ("Plano de ação", "Correções e ações corretivas, com causa tratada, responsável, prazo, status e evidência."),
    ("Eficácia", "Indicador, meta, critério e medições, com a conclusão sobre a eficácia e o gráfico."),
    ("Checklist", "Doze verificações de qualidade do tratamento, com percentual de conclusão."),
    ("Exemplo 1 - Pizzaria", "Tratamento de uma não conformidade de auditoria, com padrões da própria organização."),
    ("Exemplo 2 - Compras", "Tratamento de uma não conformidade de auditoria, com procedimento interno e indicador mensal."),
]:
    line(k, text)
r += 1
section("Regras e premissas do modelo")
for k, text in [
    ("Um registro por vez", "As abas Registro, Plano de ação e Eficácia servem a uma não conformidade. Para a seguinte, salve uma cópia do arquivo ou duplique as três abas. A aba Controle reúne todas."),
    ("Situação", "Uma não conformidade ou uma ação é considerada atrasada quando o prazo é anterior à data de hoje e ela ainda não foi concluída."),
    ("Eficácia", "A ação é considerada eficaz quando os últimos períodos medidos depois das ações ficam dentro da meta, na quantidade definida no critério, e a não conformidade não voltou."),
    ("Medições", "Preencha as medições em sequência, sem pular linhas. O modelo lê a última linha preenchida para concluir."),
    ("Conferência", "As conferências apontam campos vazios. Elas não avaliam a qualidade do texto: uma causa mal definida passa pela conferência."),
    ("Exemplos", "Os dados das abas de exemplo são ilustrativos, criados para este material."),
]:
    line(k, text)
setup(ws, INK, f"B1:C{r}", landscape=False, fit_height=True)

# ------------------------------------------------------------------ Controle
ws = wb.create_sheet("Controle")
widths(ws, {"A": 2, "B": 5, "C": 13, "D": 26, "E": 20, "F": 42, "G": 12, "H": 22, "I": 13, "J": 24, "K": 13, "L": 12, "M": 11, "N": 16,
            "O": 2})
title(ws, "Controle de não conformidades", "Uma linha por não conformidade. Preencha as células em amarelo-claro.", "N")
for rr, l1, l2, f2 in [(4, "Organização ou unidade", "Responsável pelo controle", None), (5, "Período", "Atualizado em", DATE)]:
    label(ws, f"B{rr}", l1, merge=f"B{rr}:D{rr}")
    inp(ws, f"E{rr}", merge=f"E{rr}:F{rr}")
    label(ws, f"G{rr}", l2, merge=f"G{rr}:H{rr}")
    inp(ws, f"I{rr}", merge=f"I{rr}:N{rr}", fmt=f2)
    ws.row_dimensions[rr].height = 21.75
dv_date(ws, "I5")
for col, text in zip("BCDEFGHIJKLMN", ["Nº", "Aberta em", "Origem", "Processo", "Descrição resumida", "Ação corretiva?", "Responsável", "Prazo",
                                       "Etapa atual", "Encerrada em", "Reincidente?", "Dias em aberto", "Situação"]):
    head(ws, f"{col}7", text)
ws.row_dimensions[7].height = 30
hint_row(ws, 8, [("B", "", None), ("C", "Data", None), ("D", "De onde veio", None), ("E", "Onde ocorreu", None), ("F", "A declaração, em uma frase", None),
                 ("G", "Sim ou Não", None), ("H", "Do tratamento", None), ("I", "Das ações", None), ("J", "Do tratamento", None),
                 ("K", "Data", None), ("L", "Já tinha sido tratada?", None), ("M", "Calculado", None), ("N", "Calculada", None)])
ex(ws, "B9", "Ex.", h="center")
for col, v, h_, f_ in [("C", date(2026, 9, 22), "center", DATE), ("D", "Auditoria interna", "left", None), ("E", "Requisição de compra", "left", None),
                       ("F", "Requisições de compra são aceitas sem a especificação técnica exigida pelo procedimento.", "left", None),
                       ("G", "Sim", "center", None), ("H", "Gerente de Suprimentos", "left", None), ("I", date(2026, 10, 23), "center", DATE),
                       ("J", "Ação em implantação", "center", None), ("K", "", "center", None), ("L", "Não", "center", None),
                       ("M", 7, "center", None), ("N", "No prazo", "center", None)]:
    ex(ws, f"{col}9", v, h=h_, fmt=f_)
ws.row_dimensions[9].height = 33
C1, C2 = 10, 29
for k in range(20):
    rr = C1 + k
    num(ws, f"B{rr}", k + 1)
    inp(ws, f"C{rr}", h="center", fmt=DATE)
    for col in "DEF":
        inp(ws, f"{col}{rr}")
    inp(ws, f"G{rr}", h="center")
    inp(ws, f"H{rr}")
    inp(ws, f"I{rr}", h="center", fmt=DATE)
    inp(ws, f"J{rr}", h="center")
    inp(ws, f"K{rr}", h="center", fmt=DATE)
    inp(ws, f"L{rr}", h="center")
    calc(ws, f"M{rr}", f'=IF(C{rr}="","",IF(K{rr}<>"",K{rr}-C{rr},TODAY()-C{rr}))', b=False, fmt="0")
    calc(ws, f"N{rr}", f'=IF(C{rr}="","",IF(OR(J{rr}="Encerrada",K{rr}<>""),IF(K{rr}="","Falta a data",IF(J{rr}="Encerrada","Encerrada","Confira a etapa")),'
         f'IF(I{rr}="","Sem prazo",IF(I{rr}<TODAY(),"Atrasada","No prazo"))))', b=False, sz=9)
    ws.row_dimensions[rr].height = 30
dv_list(ws, f"D{C1}:D{C2}", ORIGENS, "Escolha a origem da não conformidade")
dv_list(ws, f"G{C1}:G{C2}", ["Sim", "Não"], "A não conformidade pede ação corretiva?")
dv_list(ws, f"J{C1}:J{C2}", ETAPAS, "Escolha a etapa atual do tratamento")
dv_list(ws, f"L{C1}:L{C2}", ["Sim", "Não"], "A mesma não conformidade já tinha sido tratada antes?")
for col in "CIK":
    dv_date(ws, f"{col}{C1}:{col}{C2}")
cf_equal(ws, f"N{C1}:N{C2}", SIT_CF)
cf_equal(ws, f"L{C1}:L{C2}", [("Sim", RED)])
cf_equal(ws, f"J{C1}:J{C2}", [("Encerrada", GREEN)])
note(ws, "G7", "Toda não conformidade recebe correção. A ação corretiva depende da avaliação da abrangência e do efeito. Registre a decisão na aba Registro.")
note(ws, "L7", "Marque Sim quando a mesma falha já tinha sido tratada e encerrada. Reincidência indica que a causa não foi eliminada.")
note(ws, "M7", "Para registros encerrados, conta os dias entre a abertura e o encerramento. Para os demais, conta até hoje.")
s = C2 + 2
band(ws, s, "Resumo automático", "N")
N_, M_ = f"N{C1}:N{C2}", f"M{C1}:M{C2}"
for k, (text, formula, fmt) in enumerate([
    ("Não conformidades registradas", f"=COUNT(C{C1}:C{C2})", None),
    ("Encerradas", f'=COUNTIF({N_},"Encerrada")', None),
    ("Em aberto", f"=G{s+1}-G{s+2}", None),
    ("Atrasadas", f'=COUNTIF({N_},"Atrasada")', None),
    ("Tempo médio de encerramento, em dias", f'=IF(G{s+2}=0,"",AVERAGEIF({N_},"Encerrada",{M_}))', "0"),
    ("Com ação corretiva", f'=COUNTIF(G{C1}:G{C2},"Sim")', None),
    ("Reincidentes", f'=COUNTIF(L{C1}:L{C2},"Sim")', None),
    ("Aviso", f'=IF(G{s+1}=0,"Registre as não conformidades",IF(COUNTIF({N_},"Falta a data")>0,"Há registro encerrado sem data",IF(COUNTIF({N_},"Confira a etapa")>0,"Há registro com data de encerramento e etapa em aberto",'
              f'IF(SUMPRODUCT((C{C1}:C{C2}<>"")*(H{C1}:H{C2}=""))>0,"Há registro sem responsável",IF(G{s+4}>0,"Há não conformidade atrasada","OK")))))', None),
], 1):
    label(ws, f"B{s+k}", text, merge=f"B{s+k}:F{s+k}", h="right")
    calc(ws, f"G{s+k}", formula, fmt=fmt, merge=f"G{s+k}:H{s+k}", sz=9 if text == "Aviso" else 10, b=text != "Aviso")
    ws.row_dimensions[s + k].height = 30 if text == "Aviso" else 21.75
cf_warn(ws, f"G{s+8}:H{s+8}", f"G{s+8}")
put(ws, f"J{s+1}", "Origem", f=font(10, True), bg=GRAY, h="center", merge=f"J{s+1}:L{s+1}")
put(ws, f"M{s+1}", "Registros", f=font(10, True), bg=GRAY, h="center", merge=f"M{s+1}:N{s+1}")
for j, o in enumerate(ORIGENS):
    rr = s + 2 + j
    put(ws, f"J{rr}", o, f=font(10, True), bg=GRAY, merge=f"J{rr}:L{rr}")
    calc(ws, f"M{rr}", f'=COUNTIF(D{C1}:D{C2},"{o}")', merge=f"M{rr}:N{rr}")
ch = BarChart()
ch.type = "bar"
ch.height, ch.width = 7.5, 26
chart_style(ch)
ch.legend = None
ch.gapWidth = 60
ch.add_data(Reference(ws, min_col=13, min_row=s + 1, max_row=s + 8), titles_from_data=True)
ch.set_categories(Reference(ws, min_col=10, min_row=s + 2, max_row=s + 8))
ch.series[0].graphicalProperties.solidFill = BLUE
ch.series[0].graphicalProperties.line.noFill = True
ch.series[0].dLbls = DataLabelList()
ch.series[0].dLbls.showVal = True
ch.series[0].dLbls.showSerName = ch.series[0].dLbls.showCatName = ch.series[0].dLbls.showLegendKey = False
ch.y_axis.scaling.min = 0
ch.y_axis.majorUnit = 1
ch.x_axis.scaling.orientation = "maxMin"
ws.add_chart(ch, f"B{s+10}")
ws.freeze_panes = "D8"
setup(ws, BLUE, f"B1:N{s+25}")
ws.print_title_rows = "7:7"
ws.row_breaks.append(Break(id=s - 1))

# ------------------------------------------------------------------ Registro
ws = wb.create_sheet("Registro")
widths(ws, {"A": 2, "B": 5, "C": 24, "D": 34, "E": 24, "F": 17, "G": 17, "H": 2})
title(ws, "Registro de não conformidade", "Um registro por não conformidade. Preencha os blocos na ordem.", "G")
for rr, l1, l2, f1, f2 in [(4, "Registro nº", "Data de abertura", None, DATE), (5, "Origem", "Processo", None, None),
                           (6, "Registrado por", "Responsável", None, None)]:
    label(ws, f"B{rr}", l1, merge=f"B{rr}:C{rr}")
    inp(ws, f"D{rr}", fmt=f1)
    label(ws, f"E{rr}", l2)
    inp(ws, f"F{rr}", merge=f"F{rr}:G{rr}", fmt=f2)
    ws.row_dimensions[rr].height = 21.75
dv_date(ws, "F4")
dv_list(ws, "D5", ORIGENS, "Escolha a origem da não conformidade")

band(ws, 8, "1. Descrição da não conformidade", "G", color=REDC)
for rr, l1, tip in [(9, "Requisito", 'O que foi combinado, e onde está escrito.\nEx.: "PR-SUP-01 rev. 5, item 4.2: toda requisição deve conter a especificação técnica."'),
                    (10, "Evidência", 'O que foi encontrado, com identificação e amostra.\nEx.: "Requisições RC-0412 e RC-0433 sem especificação, em amostra de 10."'),
                    (11, "Declaração", 'O desvio, em uma frase, sem causa e sem solução.\nEx.: "Requisições são aceitas sem a especificação exigida."')]:
    label(ws, f"B{rr}", l1, merge=f"B{rr}:C{rr}")
    inp(ws, f"D{rr}", merge=f"D{rr}:G{rr}")
    note(ws, f"D{rr}", tip)
    ws.row_dimensions[rr].height = 36

band(ws, 13, "2. Correção e consequências", "G", color=AMBER)
put(ws, "B14", "#", f=font(10, True), bg=GRAY, h="center")
put(ws, "C14", "O que foi feito para corrigir e conter", f=font(10, True), bg=GRAY, h="center", merge="C14:D14")
for col, text in zip("EFG", ["Responsável", "Data", "Status"]):
    put(ws, f"{col}14", text, f=font(10, True), bg=GRAY, h="center")
ws.row_dimensions[14].height = 21.75
for k in range(4):
    rr = 15 + k
    num(ws, f"B{rr}", k + 1)
    inp(ws, f"C{rr}", merge=f"C{rr}:D{rr}")
    inp(ws, f"E{rr}")
    inp(ws, f"F{rr}", h="center", fmt=DATE)
    inp(ws, f"G{rr}", h="center")
    ws.row_dimensions[rr].height = 30
dv_date(ws, "F15:F18")
dv_list(ws, "G15:G18", ["Feita", "Pendente"], "Feita ou Pendente")
cf_equal(ws, "G15:G18", [("Feita", GREEN), ("Pendente", RED)])
label(ws, "B19", "Consequências tratadas", merge="B19:C19")
inp(ws, "D19", merge="D19:G19")
note(ws, "D19", "O que foi feito pelo cliente ou pela área afetada: contato, reposição, devolução, desconto.")
ws.row_dimensions[19].height = 36

band(ws, 21, "3. Abrangência e decisão", "G", color=TEAL)
label(ws, "B22", "A falha existe em outro lugar?", merge="B22:C22")
inp(ws, "D22", h="center")
label(ws, "E22", "Ação corretiva necessária?")
inp(ws, "F22", h="center", merge="F22:G22")
ws.row_dimensions[22].height = 30
dv_list(ws, "D22", ["Sim", "Não", "Não verificado"], "Outros produtos, locais, turnos ou períodos")
dv_list(ws, "F22", ["Sim", "Não"], "Sim ou Não")
for rr, l1, tip in [(23, "O que foi verificado", "Onde a equipe procurou a mesma falha, e o que encontrou, com números."),
                    (24, "Justificativa da decisão", "Por que a não conformidade pede, ou não pede, ação corretiva.")]:
    label(ws, f"B{rr}", l1, merge=f"B{rr}:C{rr}")
    inp(ws, f"D{rr}", merge=f"D{rr}:G{rr}")
    note(ws, f"D{rr}", tip)
    ws.row_dimensions[rr].height = 36

band(ws, 26, "4. Análise da causa", "G", color=TEAL)
put(ws, "B27", "#", f=font(10, True), bg=GRAY, h="center")
put(ws, "C27", "Pergunta", f=font(10, True), bg=GRAY, h="center")
put(ws, "D27", "Resposta", f=font(10, True), bg=GRAY, h="center", merge="D27:E27")
put(ws, "F27", "Como foi confirmada", f=font(10, True), bg=GRAY, h="center", merge="F27:G27")
ws.row_dimensions[27].height = 21.75
for k in range(5):
    rr = 28 + k
    num(ws, f"B{rr}", k + 1)
    inp(ws, f"C{rr}", "Por quê?" if k else None)
    inp(ws, f"D{rr}", merge=f"D{rr}:E{rr}")
    inp(ws, f"F{rr}", merge=f"F{rr}:G{rr}")
    ws.row_dimensions[rr].height = 36
note(ws, "C28", 'Comece pela declaração da não conformidade.\nEx.: "Por que as requisições chegam sem especificação?"')
note(ws, "F27", "O fato ou o dado que comprova a resposta: uma observação, um teste, um levantamento. Resposta sem confirmação é hipótese.")
label(ws, "B33", "Causa raiz", merge="B33:C33")
inp(ws, "D33", merge="D33:G33")
note(ws, "D33", "A causa que, eliminada, impede a repetição. Ela descreve o processo, e não a pessoa.")
ws.row_dimensions[33].height = 36

band(ws, 35, "5. Encerramento", "G")
label(ws, "B36", "Mudanças no sistema de gestão", merge="B36:C36")
inp(ws, "D36", merge="D36:G36")
note(ws, "D36", "Procedimentos, instruções, sistemas e treinamentos alterados pela ação corretiva.")
ws.row_dimensions[36].height = 36
label(ws, "B37", "Riscos atualizados?", merge="B37:C37")
inp(ws, "D37", h="center")
label(ws, "E37", "Data de encerramento")
inp(ws, "F37", h="center", fmt=DATE, merge="F37:G37")
label(ws, "B38", "Aprovado por", merge="B38:C38")
inp(ws, "D38")
label(ws, "E38", "Eficácia (aba Eficácia)")
calc(ws, "F38", '=IF(LEFT(Eficácia!E29,6)="Eficaz","Eficaz",IF(LEFT(Eficácia!E29,3)="Não","Não eficaz","Não comprovada"))', b=False, sz=9, merge="F38:G38")
for rr in (37, 38):
    ws.row_dimensions[rr].height = 21.75
dv_list(ws, "D37", ["Sim", "Não", "Não se aplica"], "Sim, Não ou Não se aplica")
dv_date(ws, "F37")

band(ws, 40, "Conferência automática", "G")
CONF = [
    ("1. Descrição", '=IF(COUNTA(D9:D11)=3,"Completa","Falta "&IF(D9="","o requisito",IF(D10="","a evidência","a declaração")))'),
    ("2. Correção", '=IF(COUNTA(C15:C18)=0,"Registre a correção",IF(SUMPRODUCT((C15:C18<>"")*(((E15:E18="")+(F15:F18=""))>0))>0,'
                    '"Há correção sem responsável ou sem data",IF(COUNTIF(G15:G18,"Pendente")>0,"Há correção pendente","Completa")))'),
    ("3. Abrangência e decisão", '=IF(D22="","Informe se a falha existe em outro lugar",IF(D22="Não verificado","Verifique a abrangência",'
                                  'IF(F22="","Registre a decisão",IF(D24="","Justifique a decisão","Completa"))))'),
    ("4. Análise da causa", '=IF(F22="Não","Não se aplica",IF(COUNTA(D28:D32)<2,"Responda a pelo menos dois porquês",'
                            'IF(COUNTA(F28:F32)<COUNTA(D28:D32),"Há resposta sem confirmação",IF(D33="","Escreva a causa raiz","Completa"))))'),
    ("5. Encerramento", '=IF(F37="","Em aberto",IF(AND(F22="Sim",F38<>"Eficaz"),"Encerrado sem eficácia comprovada","Encerrado"))'),
]
for k, (text, formula) in enumerate(CONF):
    rr = 41 + k
    label(ws, f"B{rr}", text, merge=f"B{rr}:C{rr}")
    calc(ws, f"D{rr}", formula, h="left", b=False, merge=f"D{rr}:G{rr}")
    ws.row_dimensions[rr].height = 21.75
label(ws, "B46", "Situação do registro", merge="B46:C46")
calc(ws, "D46", '=IF(COUNTA(D9:D11,C15:C18,D22,F22)=0,"Registro em branco",IF(SUMPRODUCT((D41:D44<>"Completa")*(D41:D44<>"Não se aplica"))>0,'
     '"Há blocos incompletos",IF(F22="Não","Registro completo: só correção","Pronto para o plano de ação")))', h="left", merge="D46:G46")
ws.row_dimensions[46].height = 21.75
cf_warn(ws, "D41:D44", "D41", ok_values=("Completa", "Não se aplica"))
cf_warn(ws, "D45", "D45", ok_values=("Encerrado", "Em aberto"))
cf_warn(ws, "D46", "D46", ok_values=("Pronto para o plano de ação", "Registro completo: só correção"))
setup(ws, REDC, "B1:G46", landscape=False, fit_height=True)

# ------------------------------------------------------------------ Plano de ação
ws = wb.create_sheet("Plano de ação")
widths(ws, {"A": 2, "B": 5, "C": 16, "D": 42, "E": 32, "F": 20, "G": 13, "H": 15, "I": 13, "J": 30, "K": 24, "L": 14, "M": 2})
title(ws, "Plano de ação", "Registre as correções e as ações corretivas. Cada ação corretiva trata uma causa confirmada.", "L")
label(ws, "B4", "Registro nº", merge="B4:C4")
calc(ws, "D4", '=IF(Registro!D4="","",Registro!D4)', h="left", b=False)
label(ws, "E4", "Não conformidade")
calc(ws, "F4", '=IF(Registro!D11="","",Registro!D11)', h="left", b=False, merge="F4:L4")
ws.row_dimensions[4].height = 30
for col, text in zip("BCDEFGHIJKL", ["#", "Tipo", "O que será feito", "Causa tratada", "Responsável", "Prazo", "Status", "Concluída em",
                                      "Evidência de implantação", "Conferência", "Situação"]):
    head(ws, f"{col}6", text)
ws.row_dimensions[6].height = 21.75
hint_row(ws, 7, [("B", "", None), ("C", "Efeito ou causa", None), ("D", "Verbo no infinitivo e resultado esperado", None),
                 ("E", "Só para ação corretiva", None), ("F", "Cargo ou nome", None), ("G", "Data", None), ("H", "Da ação", None),
                 ("I", "Data", None), ("J", "Documento, registro ou teste", None), ("K", "Campos da linha", None), ("L", "Calculada", None)])
ex(ws, "B8", "Ex.", h="center")
for col, v, h_, f_ in [("C", AC, "center", None), ("D", "Tornar obrigatório o campo de especificação, com bloqueio do envio da requisição incompleta.", "left", None),
                       ("E", "O campo de especificação é opcional.", "left", None), ("F", "Analista de sistemas", "left", None),
                       ("G", date(2026, 10, 16), "center", DATE), ("H", "Concluída", "center", None), ("I", date(2026, 10, 14), "center", DATE),
                       ("J", "Chamado 4471 encerrado. Teste com 5 requisições.", "left", None), ("K", "Completa", "center", None),
                       ("L", "Concluída", "center", None)]:
    ex(ws, f"{col}8", v, h=h_, fmt=f_)
ws.row_dimensions[8].height = 36
P1, P2 = 9, 20
for k in range(12):
    rr = P1 + k
    num(ws, f"B{rr}", k + 1)
    inp(ws, f"C{rr}", h="center")
    for col in "DEF":
        inp(ws, f"{col}{rr}")
    inp(ws, f"G{rr}", h="center", fmt=DATE)
    inp(ws, f"H{rr}", h="center")
    inp(ws, f"I{rr}", h="center", fmt=DATE)
    inp(ws, f"J{rr}")
    calc(ws, f"K{rr}", f'=IF(C{rr}="","",IF(D{rr}="","Falta a ação",IF(AND(C{rr}="{AC}",E{rr}=""),"Falta a causa tratada",'
         f'IF(OR(F{rr}="",G{rr}=""),"Falta responsável ou prazo",IF(AND(H{rr}="Concluída",J{rr}=""),"Falta a evidência","Completa")))))', b=False, sz=9)
    calc(ws, f"L{rr}", f'=IF(C{rr}="","",IF(H{rr}="Concluída","Concluída",IF(H{rr}="Cancelada","Cancelada",'
         f'IF(G{rr}="","Sem prazo",IF(G{rr}<TODAY(),"Atrasada","No prazo")))))', b=False, sz=9)
    ws.row_dimensions[rr].height = 36
dv_list(ws, f"C{P1}:C{P2}", TIPOS, "Correção trata o efeito. Ação corretiva trata a causa.")
dv_list(ws, f"H{P1}:H{P2}", ST_ACAO, "Não iniciada, Em andamento, Concluída ou Cancelada")
dv_date(ws, f"G{P1}:G{P2}")
dv_date(ws, f"I{P1}:I{P2}")
cf_equal(ws, f"C{P1}:C{P2}", TIPO_CF)
cf_equal(ws, f"H{P1}:H{P2}", ST_CF)
cf_warn(ws, f"K{P1}:K{P2}", f"K{P1}", ok_values=("Completa",))
cf_equal(ws, f"L{P1}:L{P2}", SIT_CF)
note(ws, "C6", "Teste: se só esta ação for feita, o problema pode voltar no mês que vem? Se sim, é uma correção.")
note(ws, "E6", "Copie a causa confirmada da aba Registro. Ação corretiva sem causa é tentativa.")
note(ws, "J6", "Evidência de que a ação foi implantada. A evidência de que ela funcionou fica na aba Eficácia.")
s = P2 + 2
band(ws, s, "Resumo automático", "L")
for k, (text, formula, fmt) in enumerate([
    ("Ações registradas", f"=COUNTA(C{P1}:C{P2})", None),
    ("Correções", f'=COUNTIF(C{P1}:C{P2},"{COR}")', None),
    ("Ações corretivas", f'=COUNTIF(C{P1}:C{P2},"{AC}")', None),
    ("Concluídas", f'=COUNTIF(L{P1}:L{P2},"Concluída")', None),
    ("Atrasadas", f'=COUNTIF(L{P1}:L{P2},"Atrasada")', None),
    ("Percentual concluído", f'=IF(E{s+1}-COUNTIF(L{P1}:L{P2},"Cancelada")<=0,0,E{s+4}/(E{s+1}-COUNTIF(L{P1}:L{P2},"Cancelada")))', "0%"),
    ("Aviso", f'=IF(E{s+1}=0,"Registre as ações",IF(AND(Registro!F22="Sim",E{s+3}=0),"A decisão pede ação corretiva, e nenhuma foi registrada",'
              f'IF(COUNTIF(K{P1}:K{P2},"Falta*")>0,"Há ação incompleta: veja a coluna Conferência",IF(E{s+5}>0,"Há ação atrasada","OK"))))', None),
], 1):
    label(ws, f"B{s+k}", text, merge=f"B{s+k}:D{s+k}", h="right")
    calc(ws, f"E{s+k}", formula, fmt=fmt, merge=f"E{s+k}:F{s+k}", sz=9 if text == "Aviso" else 10, b=text != "Aviso")
    ws.row_dimensions[s + k].height = 30 if text == "Aviso" else 21.75
cf_warn(ws, f"E{s+7}:F{s+7}", f"E{s+7}")
ws.freeze_panes = "D7"
setup(ws, AMBER, f"B1:L{s+7}", fit_height=True)

# ------------------------------------------------------------------ Eficácia
ws = wb.create_sheet("Eficácia")
widths(ws, {"A": 2, "B": 5, "C": 24, "D": 16, "E": 14, "F": 14, "G": 16, "H": 18, "I": 3, "J": 13, "K": 13, "L": 13, "M": 13, "N": 13,
            "O": 13, "P": 2})
title(ws, "Verificação da eficácia", "Defina o indicador, a meta e o critério antes de implantar as ações. Depois, registre as medições.", "O")
for rr, l1, l2 in [(4, "Indicador", "Unidade"), (5, "Método de medição", "Verificado por")]:
    label(ws, f"B{rr}", l1, merge=f"B{rr}:C{rr}")
    inp(ws, f"D{rr}", merge=f"D{rr}:H{rr}")
    label(ws, f"J{rr}", l2, merge=f"J{rr}:K{rr}")
    inp(ws, f"L{rr}", merge=f"L{rr}:O{rr}")
    ws.row_dimensions[rr].height = 27
label(ws, "B6", "Sentido", merge="B6:C6")
inp(ws, "D6", h="center", merge="D6:E6")
label(ws, "F6", "Meta", h="left")
inp(ws, "G6", h="center", fmt="General", merge="G6:H6")
label(ws, "J6", "Data da verificação", merge="J6:K6")
inp(ws, "L6", h="center", fmt=DATE, merge="L6:O6")
label(ws, "B7", "Critério: períodos seguidos na meta", merge="B7:C7")
inp(ws, "D7", h="center", merge="D7:E7")
label(ws, "F7", "A falha voltou?", h="left")
inp(ws, "G7", h="center", merge="G7:H7")
label(ws, "J7", "Registro nº", merge="J7:K7")
calc(ws, "L7", '=IF(Registro!D4="","",Registro!D4)', h="left", b=False, merge="L7:O7")
ws.row_dimensions[6].height = 21.75
ws.row_dimensions[7].height = 30
dv_list(ws, "D6", SENTIDOS, "Menor é melhor: falhas, atrasos. Maior é melhor: entregas no prazo.")
dv_number(ws, "G6")
dv_number(ws, "D7")
dv_list(ws, "G7", ["Sim", "Não"], "A mesma não conformidade ocorreu de novo depois das ações?")
dv_date(ws, "L6")
note(ws, "D4", 'O indicador mede a não conformidade, e não a ação.\nEx.: "Requisições com especificação ausente ou insuficiente".')
note(ws, "D7", "Quantos períodos seguidos o indicador precisa ficar dentro da meta para a ação ser considerada eficaz.\nEx.: 3 meses, 4 semanas.")
note(ws, "L5", "Convém que a verificação seja feita por quem não executou as ações. É uma convenção deste material.")
for col, text in zip("BCDEFGH", ["#", "Período", "Momento", "Valor", "Meta", "Dentro da meta?", "Seguidos na meta"]):
    head(ws, f"{col}9", text)
ws.row_dimensions[9].height = 30
ex(ws, "B10", "Ex.", h="center")
for col, v, h_ in [("C", "Nov/2026", "left"), ("D", "Depois", "center"), ("E", 4, "center"), ("F", 5, "center"), ("G", "Sim", "center"),
                   ("H", 1, "center")]:
    ex(ws, f"{col}10", v, h=h_)
ws.row_dimensions[10].height = 21.75
V1, V2 = 11, 22
for k in range(12):
    rr = V1 + k
    num(ws, f"B{rr}", k + 1)
    inp(ws, f"C{rr}")
    inp(ws, f"D{rr}", h="center")
    inp(ws, f"E{rr}", h="center", fmt="General")
    calc(ws, f"F{rr}", f'=IF(OR(E{rr}="",$G$6=""),"",$G$6)', b=False, fmt="General")
    calc(ws, f"G{rr}", f'=IF(OR(E{rr}="",$G$6=""),"",IF($D$6="Maior é melhor",IF(E{rr}>=$G$6,"Sim","Não"),IF(E{rr}<=$G$6,"Sim","Não")))', b=False)
    ant = f"N(H{rr-1})+1" if k else "1"
    calc(ws, f"H{rr}", f'=IF(OR(E{rr}="",D{rr}<>"Depois"),"",IF(G{rr}="Sim",{ant},0))', b=False, fmt="0")
    ws.row_dimensions[rr].height = 21.75
dv_list(ws, f"D{V1}:D{V2}", MOMENTOS, "Antes, Durante ou Depois das ações")
dv_number(ws, f"E{V1}:E{V2}")
cf_equal(ws, f"G{V1}:G{V2}", [("Sim", GREEN), ("Não", RED)])
cf_equal(ws, f"D{V1}:D{V2}", [("Depois", "CDE7E3")])
note(ws, "H9", "Conta os períodos seguidos dentro da meta, depois das ações. Volta a zero quando um período fica fora da meta.")
s = V2 + 2  # 24
band(ws, s, "Resumo automático", "H")
E_, D_ = f"E{V1}:E{V2}", f"D{V1}:D{V2}"
RES = [
    ("Medições antes das ações", f'=SUMPRODUCT(({D_}="Antes")*({E_}<>""))', None),
    ("Média antes", f'=IF(E{s+1}=0,"",AVERAGEIF({D_},"Antes",{E_}))', "0.0"),
    ("Medições depois das ações", f'=SUMPRODUCT(({D_}="Depois")*({E_}<>""))', None),
    ("Média depois", f'=IF(E{s+3}=0,"",AVERAGEIF({D_},"Depois",{E_}))', "0.0"),
    ("Conclusão", None, None),
    ("Aviso", f'=IF(COUNT({E_})=0,"Registre as medições",IF(COUNTBLANK(INDEX({E_},1):INDEX({E_},COUNT({E_})))>0,'
              f'"Preencha as medições em sequência, sem pular linhas",IF(SUMPRODUCT(({E_}<>"")*({D_}=""))>0,"Há medição sem o momento",'
              f'IF(OR($G$6="",$D$7="",$D$6=""),"Informe o sentido, a meta e o critério",IF($G$7="","Informe se a falha voltou","OK")))))', None),
]
RUN = f"N(INDEX(H{V1}:H{V2},MAX(1,COUNT({E_}))))"
LAST = f"INDEX(G{V1}:G{V2},MAX(1,COUNT({E_})))"
MELHOR = f'IF($D$6="Maior é melhor",E{s+4}>E{s+2},E{s+4}<E{s+2})'
RES[4] = ("Conclusão",
          f'=IF(E{s+3}=0,"Sem medição depois das ações",IF($G$7="Sim","Não eficaz: a falha voltou. Refaça a análise da causa",'
          f'IF(OR($G$6="",$D$7=""),"Informe a meta e o critério",IF({RUN}>=$D$7,"Eficaz: encerre e padronize",'
          f'IF(AND(E{s+3}>=$D$7,{LAST}="Não"),IF(AND(E{s+1}>0,{MELHOR}),"Parcial: melhorou, mas não atingiu a meta. Complete as ações",'
          f'"Não eficaz: meta não atingida. Refaça a análise da causa"),IF($D$7-{RUN}=1,"Em observação: falta 1 período seguido na meta","Em observação: faltam "&($D$7-{RUN})&" períodos seguidos na meta"))))))', None)
for k, (text, formula, fmt) in enumerate(RES, 1):
    label(ws, f"B{s+k}", text, merge=f"B{s+k}:D{s+k}", h="right")
    calc(ws, f"E{s+k}", formula, fmt=fmt, merge=f"E{s+k}:H{s+k}", sz=9 if text in ("Aviso", "Conclusão") else 10, b=text != "Aviso")
    ws.row_dimensions[s + k].height = 30 if text in ("Aviso", "Conclusão") else 21.75
assert s + 5 == 29, s  # a aba Registro lê a conclusão em Eficácia!E29
cf_conclusao(ws, f"E{s+5}:H{s+5}", f"E{s+5}")
cf_warn(ws, f"E{s+6}:H{s+6}", f"E{s+6}")
line_chart(ws, "J9", 9, V1, V2, 3, 5, 6, width=14.5, height=9.5)
setup(ws, TEAL, f"B1:O{s+6}", fit_height=True)

# ------------------------------------------------------------------ Checklist
ws = wb.create_sheet("Checklist")
widths(ws, {"A": 2, "B": 5, "C": 70, "D": 14, "E": 44, "F": 2})
title(ws, "Checklist de validação do tratamento", "Escolha o status de cada item na lista suspensa (coluna em amarelo-claro).", "E")
for col, text in zip("BCDE", ["#", "Item de verificação", "Status", "Observações"]):
    head(ws, f"{col}4", text)
ws.row_dimensions[4].height = 21.75
for k, text in enumerate(CHECK):
    rr = 5 + k
    num(ws, f"B{rr}", k + 1)
    put(ws, f"C{rr}", text, bg=GRAY)
    inp(ws, f"D{rr}", h="center")
    inp(ws, f"E{rr}")
    ws.row_dimensions[rr].height = 30
dv_list(ws, "D5:D16", YESNO, "Sim, Parcial ou Não")
cf_equal(ws, "D5:D16", YESNO_CF)
band(ws, 18, "Resumo automático", "E")
summary(ws, 19, "Itens atendidos (Sim)", '=COUNTIF(D5:D16,"Sim")', "C", "D")
summary(ws, 20, "Itens parciais", '=COUNTIF(D5:D16,"Parcial")', "C", "D")
summary(ws, 21, "Itens não atendidos", '=COUNTIF(D5:D16,"Não")', "C", "D")
summary(ws, 22, "Itens sem resposta", "=COUNTBLANK(D5:D16)", "C", "D")
summary(ws, 23, "Percentual atendido", "=D19/ROWS(D5:D16)", "C", "D", fmt="0%")
summary(ws, 24, "Situação", '=IF(D19=ROWS(D5:D16),"Validado",IF(D22>0,"Em andamento","Com pendências"))', "C", "D")
setup(ws, PURPLE, "B1:E24", landscape=False, fit_height=True)


# ------------------------------------------------------------------ Exemplos
def example(ws, data):
    W = {"B": 5, "C": 18, "D": 44, "E": 34, "F": 24, "G": 14, "H": 36}
    widths(ws, dict(W, A=2, I=2))
    title(ws, "Não conformidade e ação corretiva — do registro ao encerramento",
          "Exemplo preenchido, para consulta. Use as abas Registro, Plano de ação e Eficácia para o seu tratamento.", "H")
    H, d, e, f = data["head"], data["desc"], data["eficacia"], data["encerr"]
    wide = W["D"] + W["E"] + W["F"] + W["G"] + W["H"]

    def campo(rr, rot, val, fmt=None, bg=WHITE):
        label(ws, f"B{rr}", rot, merge=f"B{rr}:C{rr}")
        inp(ws, f"D{rr}", val, merge=f"D{rr}:H{rr}", bg=bg, fmt=fmt)
        ws.row_dimensions[rr].height = alt(val, wide * 1.1)

    rr = 4
    for rot, val, fmt in [("Registro nº", f'RNC {H["num"]}', None), ("Aberto em", H["aberta"], DATE), ("Origem", f'{H["origem"]}. {H["ref"]}', None),
                          ("Processo", H["processo"], None), ("Registrado por", H["por"], None), ("Responsável", H["resp"], None)]:
        campo(rr, rot, val, fmt)
        rr += 1
    rr += 1
    band(ws, rr, "1. Descrição da não conformidade", "H", color=REDC)
    for rot, val in [("Requisito", d["req"]), ("Evidência", d["evid"]), ("Declaração", d["decl"])]:
        rr += 1
        campo(rr, rot, val)
    rr += 2
    band(ws, rr, "2 e 3. Abrangência, consequências e decisão", "H", color=TEAL)
    for rot, val in [("Abrangência", data["abrang"]["texto"]), ("Consequências", data["abrang"]["conseq"]),
                     ("Decisão", "Ação corretiva necessária. " + data["decisao"]["just"])]:
        rr += 1
        campo(rr, rot, val)
    rr += 2
    band(ws, rr, "4. Análise da causa", "H", color=TEAL)
    rr += 1
    put(ws, f"B{rr}", "#", f=font(10, True), bg=GRAY, h="center")
    put(ws, f"C{rr}", "Pergunta", f=font(10, True), bg=GRAY, h="center", merge=f"C{rr}:D{rr}")
    put(ws, f"E{rr}", "Resposta", f=font(10, True), bg=GRAY, h="center", merge=f"E{rr}:F{rr}")
    put(ws, f"G{rr}", "Como foi confirmada", f=font(10, True), bg=GRAY, h="center", merge=f"G{rr}:H{rr}")
    ws.row_dimensions[rr].height = 21.75
    for p in data["porques"]:
        rr += 1
        num(ws, f"B{rr}", p["n"])
        put(ws, f"C{rr}", p["perg"], merge=f"C{rr}:D{rr}")
        put(ws, f"E{rr}", p["resp"], merge=f"E{rr}:F{rr}")
        put(ws, f"G{rr}", p["conf"], merge=f"G{rr}:H{rr}")
        ws.row_dimensions[rr].height = max(alt(p["perg"], 62 * 1.1), alt(p["resp"], 58 * 1.1), alt(p["conf"], 50 * 1.1))
    rr += 1
    campo(rr, "Causa raiz", data["raiz"], bg=TEAL_T)
    rr += 2
    ws.row_breaks.append(Break(id=rr - 1))
    band(ws, rr, "5. Plano de ação", "H", color=AMBER)
    rr += 1
    for col, text in zip("BCDEFGH", ["#", "Tipo", "O que foi feito", "Causa tratada", "Responsável", "Concluída em", "Evidência de implantação"]):
        put(ws, f"{col}{rr}", text, f=font(10, True), bg=GRAY, h="center")
    ws.row_dimensions[rr].height = 21.75
    a1 = rr + 1
    for a in data["acoes"]:
        rr += 1
        num(ws, f"B{rr}", a["n"])
        put(ws, f"C{rr}", a["tipo"], h="center")
        put(ws, f"D{rr}", a["oque"])
        put(ws, f"E{rr}", a["causa"] or "Trata o efeito.")
        put(ws, f"F{rr}", a["resp"])
        put(ws, f"G{rr}", a["feito"], h="center", fmt=DATE)
        put(ws, f"H{rr}", a["evid"])
        ws.row_dimensions[rr].height = max(alt(a["oque"], W["D"] * 1.1), alt(a["causa"], W["E"] * 1.1), alt(a["evid"], W["H"] * 1.1), 30)
    a2 = rr
    cf_equal(ws, f"C{a1}:C{a2}", TIPO_CF)
    rr += 2
    ws.row_breaks.append(Break(id=rr - 1))
    band(ws, rr, "6. Verificação da eficácia", "H", color=TEAL)
    for rot, val in [("Indicador", f'{e["indicador"]}, em percentual.'), ("Método de medição", e["metodo"]),
                     ("Verificado por", f'{e["quem"]}, em {e["data"]:%d/%m/%Y}.'), ("Observação", e["obs"])]:
        rr += 1
        campo(rr, rot, val)
    rr += 1
    label(ws, f"B{rr}", "Sentido", merge=f"B{rr}:C{rr}")
    put(ws, f"D{rr}", e["sentido"])
    label(ws, f"E{rr}", "Meta")
    put(ws, f"F{rr}", e["meta"], h="center", fmt=PCT)
    label(ws, f"G{rr}", "Critério")
    put(ws, f"H{rr}", e["criterio"], h="center", fmt=f'0" {e["periodo"]} seguidos na meta"')
    ws.row_dimensions[rr].height = 21.75
    sent, meta, crit = f"$D${rr}", f"$F${rr}", f"$H${rr}"
    rr += 2
    hd = rr
    put(ws, f"B{rr}", "#", f=font(10, True), bg=GRAY, h="center")
    put(ws, f"C{rr}", "Período", f=font(10, True), bg=GRAY, h="center")
    put(ws, f"D{rr}", "Momento", f=font(10, True), bg=GRAY, h="center")
    put(ws, f"E{rr}", "Valor", f=font(10, True), bg=GRAY, h="center")
    put(ws, f"F{rr}", "Meta", f=font(10, True), bg=GRAY, h="center")
    put(ws, f"G{rr}", "Na meta?", f=font(10, True), bg=GRAY, h="center")
    put(ws, f"H{rr}", "Períodos seguidos na meta", f=font(10, True), bg=GRAY, h="center")
    ws.row_dimensions[rr].height = 21.75
    m1 = rr + 1
    for k, (per, mom, val) in enumerate(e["medicoes"], 1):
        rr += 1
        num(ws, f"B{rr}", k)
        put(ws, f"C{rr}", per)
        put(ws, f"D{rr}", mom, h="center")
        put(ws, f"E{rr}", val, h="center", fmt=PCT)
        calc(ws, f"F{rr}", f"={meta}", b=False, fmt=PCT)
        calc(ws, f"G{rr}", f'=IF({sent}="Maior é melhor",IF(E{rr}>={meta},"Sim","Não"),IF(E{rr}<={meta},"Sim","Não"))', b=False)
        ant = f"N(H{rr-1})+1" if k > 1 else "1"
        calc(ws, f"H{rr}", f'=IF(D{rr}<>"Depois","",IF(G{rr}="Sim",{ant},0))', b=False, fmt="0")
        ws.row_dimensions[rr].height = 19.5
    m2 = rr
    cf_equal(ws, f"G{m1}:G{m2}", [("Sim", GREEN), ("Não", RED)])
    cf_equal(ws, f"D{m1}:D{m2}", [("Depois", "CDE7E3")])
    line_chart(ws, f"B{m2 + 2}", hd, m1, m2, 3, 5, 6, width=25, height=6.4)
    for k in range(1, 12):
        ws.row_dimensions[m2 + k].height = 19.5
    rr += 13
    ws.row_breaks.append(Break(id=rr - 1))
    band(ws, rr, "7. Encerramento", "H")
    for rot, val, fmt in [("Mudanças no sistema", f["mudanca"], None), ("Riscos", f["riscos"], None), ("Encerrado em", f["data"], DATE),
                          ("Aprovado por", f["aprov"], None)]:
        rr += 1
        campo(rr, rot, val, fmt)
    rr += 2
    band(ws, rr, "Resumo automático", "H")
    s0 = rr
    for k, (text, formula, fmt) in enumerate([
        ("Correções", f'=COUNTIF(C{a1}:C{a2},"{COR}")', None),
        ("Ações corretivas", f'=COUNTIF(C{a1}:C{a2},"{AC}")', None),
        ("Média antes das ações", f'=AVERAGEIF(D{m1}:D{m2},"Antes",E{m1}:E{m2})', '0.0"%"'),
        ("Média depois das ações", f'=AVERAGEIF(D{m1}:D{m2},"Depois",E{m1}:E{m2})', '0.0"%"'),
        ("Períodos seguidos na meta", f"=N(H{m2})", None),
        ("Dias entre a abertura e o encerramento", f"=D{s0-3}-D5", "0"),
        ("Conclusão", f'=IF(N(H{m2})>={crit},"Eficaz: encerre e padronize","Em observação")', None),
    ], 1):
        label(ws, f"B{rr+k}", text, merge=f"B{rr+k}:D{rr+k}", h="right")
        calc(ws, f"E{rr+k}", formula, fmt=fmt)
        ws.row_dimensions[rr + k].height = 21.75
    cf_conclusao(ws, f"E{rr+7}", f"E{rr+7}")
    setup(ws, MUTED, f"B1:H{rr + 7}")


example(wb.create_sheet("Exemplo 1 - Pizzaria"), EX1)
example(wb.create_sheet("Exemplo 2 - Compras"), EX2)

wb.active = 1
wb.save(OUT)
print("ok", OUT)
