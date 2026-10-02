# -*- coding: utf-8 -*-
"""Gera Recursos-modelo.xlsx no mesmo padrão visual das planilhas anteriores da série."""
import math
import os
import sys

_here = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, _here)
# funções de formatação compartilhadas com o gerador do PDCA (bloco inicial do arquivo)
_src = open(os.path.join(_here, "build_pdca.py"), encoding="utf-8").read()
exec(_src[: _src.index("STATUS = [")])

from openpyxl.chart import BarChart, Series  # noqa: E402
from openpyxl.chart.label import DataLabelList  # noqa: E402
from rec_data import (ALTA, ATRAS, CAP_SITS, CHECK, CORR, CRITS, DENTRO, EMDIA, EX1, EX2, FALTA, FATORES, FOLGA, FORA, JANELA, JUSTO, MEDIA, NAO, PREV,  # noqa: E402
                      PREV_SITS, SEMMED, SEMPLANO, SIM, TIPOS, TIPOS_OC, VENCE)

BLUE, TEAL, AMBER, PURPLE, REDC = "2B5C8A", "1E7B73", "A96A12", "7A4A9A", "B0413E"
S2_T = "F8EBCB"
YESNO = ["Sim", "Parcial", "Não"]
YESNO_CF = [("Sim", GREEN), ("Parcial", YELLOW), ("Não", RED)]
NP, NI, NM, NA = 30, 40, 80, 30
PS, IN, MA, AM = "Pessoas", "Infraestrutura", "'Manutenções'", "Ambiente"
P1, P2 = 7, 7 + NP - 1          # pessoas
I1, I2 = 10, 10 + NI - 1        # infraestrutura
M1, M2 = 7, 7 + NM - 1          # manutenções
A1, A2 = 7, 7 + NA - 1          # ambiente
GERAL = "#,##0.##"
PCT = "0.0%"
EPS = "0.000000001"
CAP_CF = [(FALTA, RED), (JUSTO, S2_T), (FOLGA, GREEN)]
PREV_CF = [(EMDIA, GREEN), (VENCE, S2_T), (ATRAS, RED), (SEMPLANO, GRAY)]
AMB_CF = [(DENTRO, GREEN), (FORA, RED), (SEMMED, S2_T)]
SEMQ = "Sem pessoa qualificada na escala"
SEMTRAT = "Produto afetado sem tratamento"
FORASEM = "Fora do limite: falta a ação"


def note(ws, ref, text):
    c = Comment(text, "Modelo Recursos")
    c.width, c.height = 300, 120
    ws[ref].comment = c


def cf_warn(ws, rng_, first, ok_values=("OK",)):
    cond = ",".join(f'{first}="{v}"' for v in ok_values)
    ws.conditional_formatting.add(rng_, FormulaRule(formula=[f"OR({cond})"], fill=PatternFill("solid", bgColor=GREEN, fgColor=GREEN)))
    ws.conditional_formatting.add(rng_, FormulaRule(formula=[f"AND(LEN({first})>0,NOT(OR({cond})))"],
                                  fill=PatternFill("solid", bgColor=YELLOW, fgColor=YELLOW)))


def cf_texto(ws, rng_, first, pares, resto=None):
    for t, cor in pares:
        ws.conditional_formatting.add(rng_, FormulaRule(formula=[f'{first}="{t}"'], stopIfTrue=True, fill=PatternFill("solid", bgColor=cor, fgColor=cor)))
    if resto:
        ws.conditional_formatting.add(rng_, FormulaRule(formula=[f"LEN({first})>0"], fill=PatternFill("solid", bgColor=resto, fgColor=resto)))


def hint_row(ws, row, cells, height=30):
    for ref, text in cells:
        put(ws, f"{ref}{row}", text, f=font(9, i=True, c=MUTED), bg=GRAY, h="center")
    ws.row_dimensions[row].height = height


def alt(pares, minimo=21.75, linha=12):
    n = max(max(1, math.ceil(len(str(t or "")) / max(c, 1))) for t, c in pares)
    return max(minimo, n * linha + 7)


# ------------------------------------------------------------------ fórmulas compartilhadas pelas abas de entrada e pelos exemplos
def f_nec(dem, prod):
    return f'IF(OR(N({dem})=0,N({prod})=0),"",ROUNDUP({dem}/{prod},0))'


def f_folga(nec, esc):
    return f'IF(OR({nec}="",{esc}=""),"",{esc}-{nec})'


def f_capsit(folga):
    return f'IF({folga}="","",IF({folga}<0,"{FALTA}",IF({folga}=0,"{JUSTO}","{FOLGA}")))'


def f_capconf(c):
    return (f'IF({c["func"]}="","",IF(N({c["dem"]})=0,"Falta a demanda",IF(N({c["prod"]})=0,"Falta a produtividade",IF({c["esc"]}="","Falta o número escalado",'
            f'IF({c["qual"]}="","Falta o número de qualificados",IF({c["qual"]}>{c["esc"]},"Mais qualificados que escalados",'
            f'IF({c["folga"]}<0,IF({c["folga"]}=-1,"Falta 1 pessoa","Faltam "&-{c["folga"]}&" pessoas"),IF({c["qual"]}=0,"{SEMQ}","OK"))))))))')


def f_prox(ult, interv):
    return f'IF(OR({ult}="",{interv}=""),"",EDATE({ult},{interv}))'


def f_prevsit(cod, interv, ult, prox, ref):
    return (f'IF({cod}="","",IF({interv}="","{SEMPLANO}",IF({ult}="","{ATRAS}",IF(AND(ISNUMBER({ref}),{prox}<{ref}),"{ATRAS}",'
            f'IF(AND(ISNUMBER({ref}),{prox}-{ref}<={JANELA}),"{VENCE}","{EMDIA}")))))')


def f_quebras(cod, ini, fim, r):
    """r: intervalos do registro de manutenções (data, código, tipo, horas)."""
    return (f'IF(OR({cod}="",{ini}="",{fim}=""),"",COUNTIFS({r["cod"]},{cod},{r["tipo"]},"{CORR}",{r["data"]},">="&{ini},{r["data"]},"<="&{fim}))')


def f_horas(cod, ini, fim, r):
    return (f'IF(OR({cod}="",{ini}="",{fim}=""),"",SUMIFS({r["horas"]},{r["cod"]},{cod},{r["tipo"]},"{CORR}",{r["data"]},">="&{ini},{r["data"]},"<="&{fim}))')


def f_disp(prog, horas):
    return f'IF(OR(N({prog})=0,{horas}=""),"",({prog}-{horas})/{prog})'


def f_infraconf(c, meta):
    return (f'IF({c["cod"]}="","",IF({c["tipo"]}="","Falta o tipo",IF({c["crit"]}="","Falta a criticidade",'
            f'IF(AND(OR({c["crit"]}="{ALTA}",{c["crit"]}="{MEDIA}"),{c["interv"]}=""),"Sem plano de preventiva",IF({c["sit"]}="{ATRAS}","Preventiva atrasada",'
            f'IF(AND({c["crit"]}="{ALTA}",{c["cont"]}=""),"Crítico sem contingência",IF(AND(ISNUMBER({c["disp"]}),ISNUMBER({meta})),'
            f'IF({c["disp"]}<{meta}-{EPS},"Disponibilidade abaixo da meta","OK"),"OK")))))))')


def f_occonf(c, outros, codes):
    return (f'IF({c["cod"]}="",IF({outros}>0,"Falta o item",""),IF(ISNA(MATCH({c["cod"]},{codes},0)),"Item não cadastrado",IF({c["data"]}="","Falta a data",'
            f'IF({c["tipo"]}="","Falta o tipo",IF({c["horas"]}="","Falta o tempo parado",IF({c["tipo"]}="{CORR}",IF({c["causa"]}="","Falta a causa",'
            f'IF({c["acao"]}="","Falta a ação",IF({c["afetou"]}="","Falta dizer se afetou o produto",IF(AND({c["afetou"]}="{SIM}",{c["trat"]}=""),"{SEMTRAT}","OK")))),"OK"))))))')


def f_ambsit(fator, mn, mx, val):
    return f'IF({fator}="","",IF({val}="","{SEMMED}",IF(OR(AND({mn}<>"",{val}<{mn}),AND({mx}<>"",{val}>{mx})),"{FORA}","{DENTRO}")))'


def f_ambconf(c):
    return (f'IF({c["fator"]}="","",IF({c["porque"]}="","Falta por que afeta o produto",IF({c["controle"]}="","Falta como é controlado",'
            f'IF(AND({c["min"]}="",{c["max"]}=""),"Falta o limite",IF({c["sit"]}="{SEMMED}","Falta a medição",IF(AND({c["sit"]}="{FORA}",{c["acao"]}=""),"{FORASEM}","OK"))))))')


def chart_style(ch):
    ch.style = 2
    ch.title = None
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ch.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill="D3DBE0", w=9525))
    ch.x_axis.graphicalProperties = GraphicalProperties(ln=LineProperties(solidFill="D3DBE0", w=9525))
    ch.y_axis.graphicalProperties = GraphicalProperties(ln=LineProperties(noFill=True))


def colunas(ws, anchor, r1, r2, c_cat, c_val, width=16, height=7.5):
    ch = BarChart()
    ch.type = "col"
    ch.gapWidth = 70
    ch.height, ch.width = height, width
    chart_style(ch)
    ch.legend = None
    s = Series(Reference(ws, min_col=c_val, min_row=r1, max_row=r2), title="Itens")
    s.graphicalProperties.solidFill = BLUE
    s.graphicalProperties.line.noFill = True
    s.dLbls = DataLabelList()
    s.dLbls.showVal = True
    s.dLbls.showSerName = s.dLbls.showCatName = s.dLbls.showLegendKey = False
    ch.append(s)
    ch.set_categories(Reference(ws, min_col=c_cat, min_row=r1, max_row=r2))
    ch.y_axis.scaling.min = 0
    ch.y_axis.majorUnit = 1
    ws.add_chart(ch, anchor)


def resumo_aba(ws, s, itens, last, merge_lab="D"):
    """Faixa de resumo com o aviso na última linha; devolve a linha do aviso."""
    band(ws, s, "Resumo automático", last)
    for k, (text, formula) in enumerate(itens, 1):
        aviso = text == "Aviso"
        label(ws, f"B{s+k}", text, merge=f"B{s+k}:{merge_lab}{s+k}", h="right")
        col = chr(ord(merge_lab) + 1)
        calc(ws, f"{col}{s+k}", formula, merge=f"{col}{s+k}:{last}{s+k}", sz=9 if aviso else 10, b=not aviso, h="left")
        ws.row_dimensions[s + k].height = 21.75
    av = s + len(itens)
    col = chr(ord(merge_lab) + 1)
    cf_warn(ws, f"{col}{av}:{last}{av}", f"{col}{av}")
    return av


# ================================================================== pasta de trabalho
wb = Workbook()
wb.properties.title = "Recursos, infraestrutura e ambiente — Modelo"
wb.properties.creator = "Modelo Recursos"
wb.properties.language = "pt-BR"

# ------------------------------------------------------------------ Instruções
ws = wb.active
ws.title = "Instruções"
widths(ws, {"A": 2, "B": 26, "C": 90, "D": 2})
title(ws, "Recursos, infraestrutura e ambiente — Como usar esta planilha",
      "Modelo para dimensionar as pessoas no pico, controlar a manutenção da infraestrutura e acompanhar o ambiente dos processos.", "C")
r = 4


def section(text):
    global r
    band(ws, r, text, "C", sz=11)
    r += 1


def line(k, v, kbg=GRAY, vbg=None, kf=None, kh="left", height=None):
    global r
    put(ws, f"B{r}", k, f=kf or font(10, True), bg=kbg, h=kh)
    put(ws, f"C{r}", v, bg=vbg)
    ws.row_dimensions[r].height = height or (19.5 if len(v) <= 96 else (31.5 if len(v) <= 192 else 45.75))
    r += 1


section("Legenda: onde preencher")
line("Amarelo-claro", "Células de entrada. É aqui que você digita.", vbg=INPUT)
line("Cinza", "Células calculadas ou fixas (necessárias, próxima preventiva, quebras, disponibilidade, situações, conferências). Não altere.", vbg=GRAY)
line("Cinza em itálico", "Linha de exemplo (marcada com “Ex.”). Mostra o formato esperado e não entra nas contagens.", vbg=GRAY)
line("Listas suspensas", "Tipo, criticidade, código do item, tipo de manutenção, produto afetado, tipo de fator e checklist aceitam apenas as opções da lista.")
line("Comentários", "Células com triângulo vermelho no canto trazem uma dica de preenchimento. Passe o mouse sobre elas.")
r += 1
section("As regras de cálculo")
for k, text in [
    ("Pessoas necessárias", "A demanda do pico dividida pela produtividade de uma pessoa, arredondada para cima."),
    ("Folga", f"As pessoas escaladas menos as necessárias. Negativa: {FALTA.lower()}. Zero: {JUSTO.lower()}. Positiva: {FOLGA.lower()}."),
    ("Próxima preventiva", "A última preventiva mais o intervalo, em meses."),
    (EMDIA, f"A próxima preventiva está a mais de {JANELA} dias do fim do período."),
    (VENCE, f"A próxima preventiva está a {JANELA} dias ou menos."),
    (ATRAS, "A data já passou, ou há intervalo e nenhuma preventiva registrada."),
    (SEMPLANO, "O item não tem intervalo de preventiva. Aceito só na criticidade baixa."),
    ("Quebras e horas paradas", "Só as corretivas da aba Manutenções, com data dentro do período."),
    ("Disponibilidade", "Horas programadas menos horas paradas, divididas pelas horas programadas. A preventiva é planejada e não entra na conta."),
    ("Ambiente", f"{FORA} quando a medição fica abaixo do mínimo ou acima do máximo. {SEMMED} quando não há medição."),
]:
    line(k, text)
r += 1
section("Ordem recomendada de preenchimento")
for k, text in enumerate([
    "Aba Pessoas: cada função no período de maior movimento, com a demanda do pico, a produtividade, as escaladas e as qualificadas.",
    "Aba Infraestrutura: informe o início e o fim do período e a meta de disponibilidade.",
    "Aba Infraestrutura: cadastre os itens, com tipo, criticidade, preventiva, contingência e horas programadas no período.",
    "Aba Manutenções: lance as preventivas e as corretivas, com o tempo parado, a causa, a ação e o produto afetado.",
    "Aba Ambiente: os fatores que afetam o produto, com o limite, o controle e a última medição.",
    "Aba Painel: leia os avisos. Depois, valide na aba Checklist.",
], 1):
    line(f"Passo {k}", text)
r += 1
section("Abas da planilha")
for k, text in [
    (PS, f"Até {NP} funções e períodos. Calcula as necessárias, a folga, a situação e a conferência."),
    (IN, f"Até {NI} itens. Calcula a próxima preventiva, a situação, as quebras, as horas paradas, a disponibilidade e a conferência."),
    ("Manutenções", f"Até {NM} registros de preventivas e corretivas. Traz o nome do item e calcula a conferência."),
    (AM, f"Até {NA} fatores. Calcula a situação e a conferência."),
    ("Painel", "Os números das quatro abas, o gráfico das preventivas e o aviso."),
    ("Checklist", "Doze verificações dos recursos, da infraestrutura e do ambiente, com percentual de conclusão."),
    ("Exemplo 1 - Pizzaria", "A loja em abril de 2027."),
    ("Exemplo 2 - Indústria", "A fábrica em julho de 2027."),
]:
    line(k, text)
r += 1
section("Regras e premissas do modelo")
for k, text in [
    ("Demanda do pico", "A maior demanda numa hora ou num turno, medida no sistema. A média esconde o problema."),
    ("Criticidade", "Alta: se parar, o produto ou a entrega param. Média: atrapalha, mas há como seguir. Baixa: há reserva, ou a falta não afeta o produto."),
    ("Contingência", "Obrigatória para a criticidade alta: o que fazer quando o item parar, escrito antes."),
    ("Horas programadas", "As horas em que o item deveria estar disponível no período. Deixe em branco para itens sem horário, como um sistema."),
    ("Exemplos", "Os dados das abas de exemplo são ilustrativos, criados para este material."),
]:
    line(k, text)
setup(ws, INK, f"B1:C{r}", landscape=False, fit_height=True)

# ------------------------------------------------------------------ Pessoas
ws = wb.create_sheet(PS)
widths(ws, {"A": 2, "B": 5, "C": 24, "D": 20, "E": 20, "F": 11, "G": 13, "H": 12, "I": 11, "J": 12, "K": 9, "L": 13, "M": 32, "N": 2})
title(ws, "Pessoas no pico", "Uma função por período por linha. As pessoas necessárias vêm da demanda do pico e da produtividade.", "M")
for col, text in zip("BCDEFGHIJKLM", ["#", "Função", "Período", "Unidade da demanda", "Demanda no pico", "Produtividade por pessoa", "Necessárias", "Escaladas",
                                      "Qualificadas", "Folga", "Situação", "Conferência"]):
    head(ws, f"{col}4", text)
ws.row_dimensions[4].height = 33
hint_row(ws, 5, [("B", ""), ("C", ""), ("D", "Dia e horário"), ("E", "pedidos por hora"), ("F", "Número"), ("G", "Na mesma unidade"), ("H", "Calculada"), ("I", "Número"),
                 ("J", "Das escaladas"), ("K", "Calculada"), ("L", "Calculada"), ("M", "Calculada")])
ex(ws, "B6", "Ex.", h="center")
for col, v, h_ in [("C", "Entregador", "left"), ("D", "Sexta, 20h às 22h", "left"), ("E", "entregas por hora", "left"), ("F", 29, "center"), ("G", 3, "center"), ("H", 10, "center"),
                   ("I", 8, "center"), ("J", 8, "center"), ("K", -2, "center"), ("L", FALTA, "center"), ("M", "Faltam 2 pessoas", "center")]:
    ex(ws, f"{col}6", v, h=h_)
ws.row_dimensions[6].height = 21.75
for k in range(NP):
    rr = P1 + k
    num(ws, f"B{rr}", k + 1)
    for col in "CDE":
        inp(ws, f"{col}{rr}")
    for col in "FGIJ":
        inp(ws, f"{col}{rr}", h="center", fmt=GERAL)
    calc(ws, f"H{rr}", "=" + f_nec(f"F{rr}", f"G{rr}"))
    calc(ws, f"K{rr}", "=" + f_folga(f"H{rr}", f"I{rr}"), b=False)
    calc(ws, f"L{rr}", "=" + f_capsit(f"K{rr}"), sz=9)
    calc(ws, f"M{rr}", "=" + f_capconf(dict(func=f"C{rr}", dem=f"F{rr}", prod=f"G{rr}", esc=f"I{rr}", qual=f"J{rr}", folga=f"K{rr}")), b=False, sz=9)
    ws.row_dimensions[rr].height = 21.75
dv_number(ws, f"F{P1}:G{P2}")
dv_number(ws, f"I{P1}:J{P2}")
cf_texto(ws, f"L{P1}:L{P2}", f"L{P1}", CAP_CF)
cf_warn(ws, f"M{P1}:M{P2}", f"M{P1}")
note(ws, "F4", "A maior demanda numa hora ou num turno, medida e não estimada: pedidos por hora no pico, máquinas em operação no turno.")
note(ws, "G4", "Quanto uma pessoa faz, bem feito, na mesma unidade e no mesmo tempo da demanda. Para máquinas, quantas uma pessoa opera.")
note(ws, "J4", "Das pessoas escaladas, quantas sabem fazer o trabalho sozinhas. A matriz de competências dá esse número.")
s = P2 + 2
PAV = resumo_aba(ws, s, [
    ("Funções e períodos", f"=COUNTA(C{P1}:C{P2})"),
    ("Com falta de gente", f'=COUNTIF(L{P1}:L{P2},"{FALTA}")&" linhas, "&-SUMIF(K{P1}:K{P2},"<0")&" pessoas a menos no pico"'),
    ("Aviso", f'=IF(E{s+1}=0,"Liste as funções e os períodos",IF(COUNTIF(L{P1}:L{P2},"{FALTA}")>0,"Há função com falta de gente no pico",'
              f'IF(COUNTIF(M{P1}:M{P2},"{SEMQ}")>0,"Há função sem pessoa qualificada na escala",'
              f'IF(SUMPRODUCT((M{P1}:M{P2}<>"")*(M{P1}:M{P2}<>"OK"))>0,"Há linha a completar: veja a coluna Conferência","OK"))))'),
], "M")
ws.freeze_panes = f"F{P1}"
setup(ws, BLUE, f"B1:M{PAV}")

# ------------------------------------------------------------------ Infraestrutura
ws = wb.create_sheet(IN)
widths(ws, {"A": 2, "B": 5, "C": 10, "D": 26, "E": 18, "F": 14, "G": 11, "H": 11, "I": 12, "J": 34, "K": 12, "L": 12, "M": 14, "N": 9, "O": 10, "P": 12, "Q": 30, "R": 2})
title(ws, "Infraestrutura", "Um item por linha: criticidade, preventiva, contingência e o que as quebras do período custaram.", "Q")
for rr, text, fmt, dica in [(4, "Início do período", DATE, "As quebras e as horas paradas contam só as corretivas com data entre o início e o fim."),
                            (5, "Fim do período", DATE, "Também é a data da leitura: a situação da preventiva é calculada nela."),
                            (6, "Meta de disponibilidade", "0%", "A disponibilidade mínima aceita para os itens com horas programadas.")]:
    label(ws, f"B{rr}", text, merge=f"B{rr}:D{rr}")
    inp(ws, f"E{rr}", h="center", fmt=fmt)
    put(ws, f"F{rr}", dica, f=font(9, i=True, c=MUTED), bg=GRAY, merge=f"F{rr}:Q{rr}")
dv_date(ws, "E4:E5")
dv = DataValidation(type="decimal", operator="between", formula1="0", formula2="1", allow_blank=True)
dv.errorTitle, dv.error = "Valor inválido", "Digite um percentual entre 0% e 100%."
dv.showErrorMessage = True
ws.add_data_validation(dv)
dv.add("E6")
INI, FIM, META = f"{IN}!$E$4", f"{IN}!$E$5", f"{IN}!$E$6"
for col, text in zip("BCDEFGHIJKLMNOPQ", ["#", "Código", "Item", "Tipo", "Local", "Criticidade", "Preventiva (meses)", "Última preventiva", "Contingência", "Horas programadas",
                                          "Próxima", "Situação", "Quebras", "Horas paradas", "Disponib.", "Conferência"]):
    head(ws, f"{col}8", text)
ws.row_dimensions[8].height = 33
hint_row(ws, 9, [("B", ""), ("C", "FOR-01"), ("D", ""), ("E", "Lista"), ("F", ""), ("G", "Lista"), ("H", "Número"), ("I", "Data"), ("J", "O que fazer se parar"), ("K", "No período"),
                 ("L", "Calculada"), ("M", "Calculada"), ("N", "Calculada"), ("O", "Calculada"), ("P", "Calculada"), ("Q", "Calculada")])
REG = dict(data=f"{MA}!$C${M1}:$C${M2}", cod=f"{MA}!$D${M1}:$D${M2}", tipo=f"{MA}!$F${M1}:$F${M2}", horas=f"{MA}!$G${M1}:$G${M2}")
for k in range(NI):
    rr = I1 + k
    num(ws, f"B{rr}", k + 1)
    inp(ws, f"C{rr}", h="center")
    for col in "DFJ":
        inp(ws, f"{col}{rr}")
    for col in "EG":
        inp(ws, f"{col}{rr}", h="center")
    inp(ws, f"H{rr}", h="center", fmt=GERAL)
    inp(ws, f"I{rr}", h="center", fmt=DATE)
    inp(ws, f"K{rr}", h="center", fmt=GERAL)
    calc(ws, f"L{rr}", "=" + f_prox(f"I{rr}", f"H{rr}"), b=False, fmt=DATE)
    calc(ws, f"M{rr}", "=" + f_prevsit(f"C{rr}", f"H{rr}", f"I{rr}", f"L{rr}", "$E$5"), sz=9)
    calc(ws, f"N{rr}", "=" + f_quebras(f"C{rr}", "$E$4", "$E$5", REG), b=False)
    calc(ws, f"O{rr}", "=" + f_horas(f"C{rr}", "$E$4", "$E$5", REG), b=False, fmt=GERAL)
    calc(ws, f"P{rr}", "=" + f_disp(f"K{rr}", f"O{rr}"), fmt=PCT)
    cells = dict(cod=f"C{rr}", tipo=f"E{rr}", crit=f"G{rr}", interv=f"H{rr}", sit=f"M{rr}", cont=f"J{rr}", disp=f"P{rr}")
    calc(ws, f"Q{rr}", "=" + f_infraconf(cells, "$E$6"), b=False, sz=9)
    ws.row_dimensions[rr].height = 30
dv_list(ws, f"E{I1}:E{I2}", TIPOS, "Edifício e instalações, equipamento, transporte ou TI e software")
dv_list(ws, f"G{I1}:G{I2}", CRITS, "Alta: se parar, o produto ou a entrega param. Média: atrapalha. Baixa: há reserva")
dv_number(ws, f"H{I1}:H{I2}")
dv_number(ws, f"K{I1}:K{I2}")
dv_date(ws, f"I{I1}:I{I2}")
cf_texto(ws, f"M{I1}:M{I2}", f"M{I1}", PREV_CF)
ws.conditional_formatting.add(f"P{I1}:P{I2}", FormulaRule(formula=[f'AND(ISNUMBER(P{I1}),ISNUMBER($E$6),P{I1}<$E$6-{EPS})'], fill=PatternFill("solid", bgColor=RED, fgColor=RED)))
cf_warn(ws, f"Q{I1}:Q{I2}", f"Q{I1}")
note(ws, "J8", "Para a criticidade alta: o que fazer quando o item parar, escrito antes. Equipamento reserva, caminho manual, terceiro.")
note(ws, "K8", "As horas em que o item deveria estar disponível no período: dias de trabalho × horas por dia. Em branco para itens sem horário, como um sistema.")
s = I2 + 2
M_ = f"M{I1}:M{I2}"
IAV = resumo_aba(ws, s, [
    ("Itens cadastrados", f'=COUNTA(C{I1}:C{I2})&" itens, "&COUNTIF(G{I1}:G{I2},"{ALTA}")&" críticos"'),
    ("Preventivas", "=" + '&", "&'.join(f'COUNTIF({M_},"{x}")&" {x.lower()}"' for x in PREV_SITS)),
    ("Quebras no período", f'=SUM(N{I1}:N{I2})&" corretivas, "&SUM(O{I1}:O{I2})&" horas paradas"'),
    ("Aviso", f'=IF(COUNTA(C{I1}:C{I2})=0,"Cadastre os itens",IF(OR(E4="",E5=""),"Informe o início e o fim do período",IF(COUNTIF({M_},"{ATRAS}")>0,"Há preventiva atrasada",'
              f'IF(SUMPRODUCT((G{I1}:G{I2}="{ALTA}")*(J{I1}:J{I2}=""))>0,"Há item crítico sem contingência",'
              f'IF(AND(ISNUMBER(E6),COUNTIF(P{I1}:P{I2},"<"&(E6-{EPS}))>0),"Há item abaixo da meta de disponibilidade",'
              f'IF(SUMPRODUCT((Q{I1}:Q{I2}<>"")*(Q{I1}:Q{I2}<>"OK"))>0,"Há item a completar: veja a coluna Conferência",'
              f'IF(COUNTIF({M_},"{VENCE}")>0,"Há preventiva que vence em {JANELA} dias: confirme","OK")))))))'),
], "Q")
ws.freeze_panes = f"E{I1}"
setup(ws, AMBER, f"B1:Q{IAV}")

# ------------------------------------------------------------------ Manutenções
ws = wb.create_sheet("Manutenções")
widths(ws, {"A": 2, "B": 5, "C": 12, "D": 10, "E": 22, "F": 12, "G": 10, "H": 32, "I": 30, "J": 32, "K": 10, "L": 30, "M": 30, "N": 2})
title(ws, "Manutenções", "Uma preventiva ou uma corretiva por linha. As corretivas alimentam as quebras e a disponibilidade da aba Infraestrutura.", "M")
for col, text in zip("BCDEFGHIJKLM", ["#", "Data", "Código", "Item", "Tipo", "Horas paradas", "O que aconteceu", "Causa", "Ação", "Afetou o produto?", "Tratamento do produto",
                                      "Conferência"]):
    head(ws, f"{col}4", text)
ws.row_dimensions[4].height = 33
hint_row(ws, 5, [("B", ""), ("C", "Data"), ("D", "Lista"), ("E", "Vem do cadastro"), ("F", "Lista"), ("G", "Número"), ("H", ""), ("I", "Na corretiva"), ("J", "Na corretiva"),
                 ("K", "Lista"), ("L", "Se afetou"), ("M", "Calculada")])
ex(ws, "B6", "Ex.", h="center")
for col, v, h_, fmt in [("C", date(2027, 4, 19), "center", DATE), ("D", "MOT-02", "center", None), ("E", "Moto de entrega 2", "left", None), ("F", CORR, "center", None),
                        ("G", 12, "center", None), ("H", "Corrente partiu na entrega", "left", None), ("I", "Corrente gasta", "left", None), ("J", "Corrente trocada", "left", None),
                        ("K", SIM, "center", None), ("L", "Desconto ao cliente", "left", None), ("M", "OK", "center", None)]:
    ex(ws, f"{col}6", v, h=h_, fmt=fmt)
ws.row_dimensions[6].height = 21.75
CODES = f"{IN}!$C${I1}:$C${I2}"
for k in range(NM):
    rr = M1 + k
    num(ws, f"B{rr}", k + 1)
    inp(ws, f"C{rr}", h="center", fmt=DATE)
    inp(ws, f"D{rr}", h="center")
    calc(ws, f"E{rr}", f'=IF(D{rr}="","",IFERROR(INDEX({IN}!$D${I1}:$D${I2},MATCH(D{rr},{CODES},0))&"",""))', h="left", b=False, sz=9)
    inp(ws, f"F{rr}", h="center")
    inp(ws, f"G{rr}", h="center", fmt=GERAL)
    for col in "HIJL":
        inp(ws, f"{col}{rr}")
    inp(ws, f"K{rr}", h="center")
    cells = dict(cod=f"D{rr}", data=f"C{rr}", tipo=f"F{rr}", horas=f"G{rr}", causa=f"I{rr}", acao=f"J{rr}", afetou=f"K{rr}", trat=f"L{rr}")
    calc(ws, f"M{rr}", "=" + f_occonf(cells, f"COUNTA(C{rr},F{rr}:L{rr})", CODES), b=False, sz=9)
    ws.row_dimensions[rr].height = 30
dv_date(ws, f"C{M1}:C{M2}")
dv = DataValidation(type="list", formula1=f"={CODES}", allow_blank=True)
dv.promptTitle, dv.prompt = "Opções", "Código do item, da aba Infraestrutura"
dv.showInputMessage = True
ws.add_data_validation(dv)
dv.add(f"D{M1}:D{M2}")
dv_list(ws, f"F{M1}:F{M2}", TIPOS_OC, "Corretiva, depois da quebra, ou preventiva, planejada")
dv_number(ws, f"G{M1}:G{M2}")
dv_list(ws, f"K{M1}:K{M2}", [SIM, NAO], "A quebra afetou o produto ou o serviço?")
cf_texto(ws, f"M{M1}:M{M2}", f"M{M1}", [("OK", GREEN), (SEMTRAT, RED)], resto=YELLOW)
note(ws, "G4", "As horas em que o item ficou parado dentro das horas programadas. Na preventiva, informe o tempo também, mesmo que seja zero.")
note(ws, "K4", "Na corretiva: o produto feito durante a falha, ou o que estava no equipamento, foi afetado? Se sim, escreva o que se fez com ele.")
s = M2 + 2
Mf = f"F{M1}:F{M2}"
MAV = resumo_aba(ws, s, [
    ("Registros", f"=COUNTA(D{M1}:D{M2})"),
    ("Por tipo", f'=COUNTIF({Mf},"{CORR}")&" corretivas, "&COUNTIF({Mf},"{PREV}")&" preventivas, "&SUMIF({Mf},"{CORR}",G{M1}:G{M2})&" horas paradas nas corretivas"'),
    ("Aviso", f'=IF(E{s+1}=0,"Lance as manutenções",IF(COUNTIF(M{M1}:M{M2},"{SEMTRAT}")>0,"Há produto afetado sem tratamento",'
              f'IF(SUMPRODUCT((M{M1}:M{M2}<>"")*(M{M1}:M{M2}<>"OK"))>0,"Há registro a completar: veja a coluna Conferência","OK")))'),
], "M")
ws.freeze_panes = f"F{M1}"
setup(ws, TEAL, f"B1:M{MAV}")

# ------------------------------------------------------------------ Ambiente
ws = wb.create_sheet(AM)
widths(ws, {"A": 2, "B": 5, "C": 26, "D": 12, "E": 14, "F": 34, "G": 12, "H": 9, "I": 9, "J": 30, "K": 11, "L": 30, "M": 12, "N": 28, "O": 2})
title(ws, "Ambiente dos processos", "Um fator por linha: por que afeta o produto, o limite, como é controlado e a última medição.", "N")
for col, text in zip("BCDEFGHIJKLMN", ["#", "Fator", "Tipo", "Local", "Por que afeta o produto", "Unidade", "Mínimo", "Máximo", "Como é controlado", "Última medição",
                                       "Ação, se fora", "Situação", "Conferência"]):
    head(ws, f"{col}4", text)
ws.row_dimensions[4].height = 33
hint_row(ws, 5, [("B", ""), ("C", ""), ("D", "Lista"), ("E", ""), ("F", "Numa frase"), ("G", "°C, %, lux"), ("H", "Número"), ("I", "Número"), ("J", "Como, quem, quando"),
                 ("K", "Número"), ("L", ""), ("M", "Calculada"), ("N", "Calculada")])
ex(ws, "B6", "Ex.", h="center")
for col, v, h_ in [("C", "Temperatura da cozinha", "left"), ("D", "Físico", "center"), ("E", "Cozinha", "left"), ("F", "Acima de 32 °C, a massa passa do ponto", "left"),
                   ("G", "°C", "center"), ("H", None, "center"), ("I", 32, "center"), ("J", "Termômetro de parede, às 21h", "left"), ("K", 34, "center"),
                   ("L", "Segundo ventilador", "left"), ("M", FORA, "center"), ("N", "OK", "center")]:
    ex(ws, f"{col}6", v, h=h_)
ws.row_dimensions[6].height = 30
for k in range(NA):
    rr = A1 + k
    num(ws, f"B{rr}", k + 1)
    for col in "CEFJL":
        inp(ws, f"{col}{rr}")
    for col in "DG":
        inp(ws, f"{col}{rr}", h="center")
    for col in "HIK":
        inp(ws, f"{col}{rr}", h="center", fmt=GERAL)
    calc(ws, f"M{rr}", "=" + f_ambsit(f"C{rr}", f"H{rr}", f"I{rr}", f"K{rr}"), sz=9)
    cells = dict(fator=f"C{rr}", porque=f"F{rr}", controle=f"J{rr}", min=f"H{rr}", max=f"I{rr}", sit=f"M{rr}", acao=f"L{rr}")
    calc(ws, f"N{rr}", "=" + f_ambconf(cells), b=False, sz=9)
    ws.row_dimensions[rr].height = 30
dv_list(ws, f"D{A1}:D{A2}", FATORES, "Físico, social ou psicológico")
dv_number(ws, f"H{A1}:I{A2}")
dv_number(ws, f"K{A1}:K{A2}")
cf_texto(ws, f"M{A1}:M{A2}", f"M{A1}", AMB_CF)
cf_texto(ws, f"N{A1}:N{A2}", f"N{A1}", [("OK", GREEN), (FORASEM, RED)], resto=YELLOW)
note(ws, "F4", "Só entram os fatores que afetam o produto ou o serviço. Escreva a ligação numa frase: acima de 28 °C, a tinta seca no anilox.")
note(ws, "H4", "Preencha o mínimo, o máximo ou os dois, na unidade da coluna G.")
s = A2 + 2
Ms = f"M{A1}:M{A2}"
AAV = resumo_aba(ws, s, [
    ("Fatores", f'=COUNTA(C{A1}:C{A2})&" fatores: "&COUNTIF(D{A1}:D{A2},"Físico")&" físicos, "&COUNTIF(D{A1}:D{A2},"Social")&" sociais, "&COUNTIF(D{A1}:D{A2},"Psicológico")&" psicológicos"'),
    ("Situação", "=" + '&", "&'.join(f'COUNTIF({Ms},"{x}")&" {x.lower()}"' for x in (DENTRO, FORA, SEMMED))),
    ("Aviso", f'=IF(COUNTA(C{A1}:C{A2})=0,"Liste os fatores do ambiente",IF(COUNTIF(N{A1}:N{A2},"{FORASEM}")>0,"Há fator fora do limite sem ação",'
              f'IF(SUMPRODUCT((N{A1}:N{A2}<>"")*(N{A1}:N{A2}<>"OK"))>0,"Há fator a completar: veja a coluna Conferência",'
              f'IF(COUNTIF({Ms},"{FORA}")>0,"Há fator fora do limite, com ação: acompanhe","OK"))))'),
], "N")
ws.freeze_panes = f"D{A1}"
setup(ws, PURPLE, f"B1:N{AAV}")

# ------------------------------------------------------------------ Painel
ws = wb.create_sheet("Painel")
widths(ws, {"A": 2, "B": 38, "C": 14, "D": 14, "E": 14, "F": 34, "G": 2})
title(ws, "Painel dos recursos", "Nada a preencher: tudo vem das outras abas.", "F")
label(ws, "B3", "Período: início e fim")
calc(ws, "C3", f'=IF({INI}="","",{INI})', fmt=DATE)
calc(ws, "D3", f'=IF({FIM}="","",{FIM})', fmt=DATE)
for col, text, m in [("B", "Indicador", None), ("C", "Resultado", None), ("D", "Como se lê", "D5:F5")]:
    put(ws, f"{col}5", text, f=font(10, True, c=WHITE), bg=INK, h="center", merge=m)
ws.row_dimensions[5].height = 21.75
Pl, Pk, Pm = (f"{PS}!{c}{P1}:{c}{P2}" for c in "LKM")
Ig, Ij, Im, In_, Io, Ip, Iq = (f"{IN}!{c}{I1}:{c}{I2}" for c in "GJMNOPQ")
Mm, Mf_, Mk, Mc = (f"{MA}!{c}{M1}:{c}{M2}" for c in "MFKC")
Am, An = (f"{AM}!{c}{A1}:{c}{A2}" for c in "MN")
IND = [
    ("PESSOAS", None, None),
    ("Funções e períodos", f"=COUNTA({PS}!C{P1}:C{P2})", "Da aba Pessoas."),
    ("Com falta de gente", f'=COUNTIF({Pl},"{FALTA}")', "As escaladas não cobrem o pico."),
    ("Pessoas a menos no pico", f'=-SUMIF({Pk},"<0")', "A soma das faltas."),
    ("Sem pessoa qualificada", f'=COUNTIF({Pm},"{SEMQ}")', "Ninguém na escala sabe fazer sozinho."),
    ("INFRAESTRUTURA", None, None),
    ("Itens cadastrados", f"=COUNTA({IN}!C{I1}:C{I2})", "Da aba Infraestrutura."),
    ("Itens críticos", f'=COUNTIF({Ig},"{ALTA}")', "Criticidade alta."),
    (ATRAS, f'=COUNTIF({Im},"{ATRAS}")', "Preventivas atrasadas: fazer antes de tudo."),
    (VENCE, f'=COUNTIF({Im},"{VENCE}")', "Confirmar a data, as peças e quem faz."),
    ("Críticos sem contingência", f'=SUMPRODUCT(({Ig}="{ALTA}")*({Ij}=""))', "Ninguém sabe o que fazer se parar."),
    ("Abaixo da meta de disponibilidade", f'=IF(ISNUMBER({META}),COUNTIF({Ip},"<"&({META}-{EPS})),0)', "Ler as causas das quebras."),
    ("Quebras no período", f"=SUM({In_})", "Só as corretivas dentro do período."),
    ("Horas paradas no período", f"=SUM({Io})", "Nas corretivas."),
    ("MANUTENÇÕES E AMBIENTE", None, None),
    ("Corretivas com produto afetado", f'=IF(OR({INI}="",{FIM}=""),0,COUNTIFS({Mf_},"{CORR}",{Mk},"{SIM}",{Mc},">="&{INI},{Mc},"<="&{FIM}))', "No período."),
    ("Produto afetado sem tratamento", f'=COUNTIF({Mm},"{SEMTRAT}")', "Tratar como não conforme."),
    ("Fatores do ambiente fora do limite", f'=COUNTIF({Am},"{FORA}")', "Com ou sem ação."),
    ("Fora do limite sem ação", f'=COUNTIF({An},"{FORASEM}")', "Ninguém decidiu nada."),
    ("Fatores sem medição", f'=COUNTIF({Am},"{SEMMED}")', "Limite sem medição não controla."),
    ("Linhas a completar", f'=SUMPRODUCT(({Pm}<>"")*({Pm}<>"OK"))+SUMPRODUCT(({Iq}<>"")*({Iq}<>"OK"))+SUMPRODUCT(({Mm}<>"")*({Mm}<>"OK"))+SUMPRODUCT(({An}<>"")*({An}<>"OK"))',
     "Conferências diferentes de OK nas quatro abas."),
]
IR = {}
rr = 6
for nome, formula, leit in IND:
    if formula is None:
        put(ws, f"B{rr}", nome, f=font(9, True, c=MUTED), bg=None, box=False, merge=f"B{rr}:F{rr}")
        ws.row_dimensions[rr].height = 21.75
        rr += 1
        continue
    label(ws, f"B{rr}", nome)
    calc(ws, f"C{rr}", formula)
    put(ws, f"D{rr}", leit, f=font(9, c=MUTED), bg=GRAY, merge=f"D{rr}:F{rr}")
    ws.row_dimensions[rr].height = 21.75
    IR[nome] = rr
    rr += 1
s = rr + 1
band(ws, s, "Resumo automático", "F")
label(ws, f"B{s+1}", "Aviso")
C = lambda n: f"C{IR[n]}"  # noqa: E731
calc(ws, f"C{s+1}", f'=IF({C("Funções e períodos")}+{C("Itens cadastrados")}+COUNTA({AM}!C{A1}:C{A2})=0,"Preencha as abas Pessoas, Infraestrutura e Ambiente",'
     f'IF({C("Produto afetado sem tratamento")}>0,"Há produto afetado sem tratamento",IF({C("Com falta de gente")}>0,"Há função com falta de gente no pico",'
     f'IF({C(ATRAS)}>0,"Há preventiva atrasada",IF({C("Críticos sem contingência")}>0,"Há item crítico sem contingência",'
     f'IF({C("Fora do limite sem ação")}>0,"Há fator do ambiente fora do limite sem ação",IF({C("Abaixo da meta de disponibilidade")}>0,"Há item abaixo da meta de disponibilidade",'
     f'IF({C("Sem pessoa qualificada")}>0,"Há função sem pessoa qualificada na escala",IF({C("Linhas a completar")}>0,"Há linha a completar: veja as conferências","OK")))))))))',
     merge=f"C{s+1}:F{s+1}", sz=9, b=False, h="left")
cf_warn(ws, f"C{s+1}:F{s+1}", f"C{s+1}")
ws.row_dimensions[s + 1].height = 21.75
NAV = s + 1
GS = s + 3
put(ws, f"B{GS - 1}", "Preventivas por situação", f=font(9, True, c=MUTED), bg=None, box=False)
for k, sit in enumerate(PREV_SITS):
    put(ws, f"B{GS + k}", sit, f=font(9), bg=GRAY)
    calc(ws, f"C{GS + k}", f'=COUNTIF({Im},B{GS + k})', b=False)
colunas(ws, f"D{GS - 1}", GS, GS + len(PREV_SITS) - 1, 2, 3, width=12, height=7)
setup(ws, REDC, f"B1:F{GS + 14}", landscape=False, fit_height=True)

# ------------------------------------------------------------------ Checklist
ws = wb.create_sheet("Checklist")
widths(ws, {"A": 2, "B": 5, "C": 70, "D": 14, "E": 44, "F": 2})
title(ws, "Checklist dos recursos, da infraestrutura e do ambiente", "Escolha o status de cada item na lista suspensa (coluna em amarelo-claro).", "E")
for col, text in zip("BCDE", ["#", "Item de verificação", "Status", "Observações"]):
    head(ws, f"{col}4", text)
ws.row_dimensions[4].height = 21.75
for k, text in enumerate(CHECK):
    rr = 5 + k
    num(ws, f"B{rr}", k + 1)
    put(ws, f"C{rr}", text, bg=GRAY)
    inp(ws, f"D{rr}", h="center")
    inp(ws, f"E{rr}")
    ws.row_dimensions[rr].height = 30
dv_list(ws, "D5:D16", YESNO, "Sim, Parcial ou Não")
cf_equal(ws, "D5:D16", YESNO_CF)
band(ws, 18, "Resumo automático", "E")
summary(ws, 19, "Itens atendidos (Sim)", '=COUNTIF(D5:D16,"Sim")', "C", "D")
summary(ws, 20, "Itens parciais", '=COUNTIF(D5:D16,"Parcial")', "C", "D")
summary(ws, 21, "Itens não atendidos", '=COUNTIF(D5:D16,"Não")', "C", "D")
summary(ws, 22, "Itens sem resposta", "=COUNTBLANK(D5:D16)", "C", "D")
summary(ws, 23, "Percentual atendido", "=D19/ROWS(D5:D16)", "C", "D", fmt="0%")
summary(ws, 24, "Situação", '=IF(D19=ROWS(D5:D16),"Validado",IF(D22>0,"Em andamento","Com pendências"))', "C", "D")
setup(ws, PURPLE, "B1:E24", landscape=False, fit_height=True)


# ------------------------------------------------------------------ Exemplos
def sub(ws, rr, cells, height=30):
    for a, b, t in cells:
        put(ws, f"{a}{rr}", t, f=font(10, True), bg=GRAY, h="center", merge=b and f"{a}{rr}:{b}{rr}")
    ws.row_dimensions[rr].height = height


def exemplo(ws, ex_):
    H = ex_["head"]
    widths(ws, {"A": 2, "B": 10, "C": 26, "D": 14, "E": 16, "F": 11, "G": 11, "H": 11, "I": 11, "J": 30, "K": 11, "L": 15, "M": 10, "N": 12, "O": 12, "P": 14, "Q": 28, "R": 2})
    title(ws, "Recursos, infraestrutura e ambiente", "Exemplo preenchido, para consulta. Use as abas de entrada para a sua organização.", "Q")
    rr = 4
    for rot, val in [("Organização", H["org"]), ("Responsável", H["resp"]), ("Horas programadas", H["horario"]), ("Origem", H["origem"])]:
        label(ws, f"B{rr}", rot, merge=f"B{rr}:C{rr}")
        inp(ws, f"D{rr}", val, merge=f"D{rr}:Q{rr}", bg=WHITE, h="left")
        ws.row_dimensions[rr].height = alt([(val, 170)], minimo=19.5)
        rr += 1
    pars = {}
    for rot, val, fmt in [("Início do período", H["ini"], DATE), ("Fim do período", H["fim"], DATE), ("Meta de disponibilidade", H["meta"], "0%")]:
        label(ws, f"B{rr}", rot, merge=f"B{rr}:C{rr}")
        inp(ws, f"D{rr}", val, bg=WHITE, h="center", fmt=fmt)
        pars[rot] = f"$D${rr}"
        rr += 1
    ini, fim, meta = pars["Início do período"], pars["Fim do período"], pars["Meta de disponibilidade"]

    # ---- pessoas
    rr += 1
    band(ws, rr, "Pessoas no pico", "Q", color=BLUE)
    rr += 1
    sub(ws, rr, [("B", "C", "Função"), ("D", None, "Período"), ("E", None, "Unidade"), ("F", None, "Demanda"), ("G", None, "Produtiv."), ("H", None, "Necessárias"),
                 ("I", None, "Escaladas"), ("J", None, "Qualificadas"), ("K", None, "Folga"), ("L", None, "Situação"), ("M", "Q", "Conferência")])
    p1 = rr + 1
    for c in ex_["caps"]:
        rr += 1
        put(ws, f"B{rr}", c["funcao"], f=font(10, True), merge=f"B{rr}:C{rr}")
        put(ws, f"D{rr}", c["periodo"])
        put(ws, f"E{rr}", c["unid"], f=font(9))
        for col, k in (("F", "dem"), ("G", "prod"), ("I", "esc"), ("J", "qual")):
            put(ws, f"{col}{rr}", c[k], h="center", fmt=GERAL)
        calc(ws, f"H{rr}", "=" + f_nec(f"F{rr}", f"G{rr}"))
        calc(ws, f"K{rr}", "=" + f_folga(f"H{rr}", f"I{rr}"), b=False)
        calc(ws, f"L{rr}", "=" + f_capsit(f"K{rr}"), sz=9)
        calc(ws, f"M{rr}", "=" + f_capconf(dict(func=f"B{rr}", dem=f"F{rr}", prod=f"G{rr}", esc=f"I{rr}", qual=f"J{rr}", folga=f"K{rr}")), b=False, sz=9, merge=f"M{rr}:Q{rr}")
        ws.row_dimensions[rr].height = alt([(c["unid"], 16), (c["periodo"], 14)], minimo=21.75)
    p2 = rr
    cf_texto(ws, f"L{p1}:L{p2}", f"L{p1}", CAP_CF)
    cf_warn(ws, f"M{p1}:Q{p2}", f"$M{p1}")

    # ---- infraestrutura
    rr += 2
    band(ws, rr, "Infraestrutura", "Q", color=AMBER)
    rr += 1
    sub(ws, rr, [("B", None, "Código"), ("C", None, "Item"), ("D", None, "Tipo"), ("E", None, "Local"), ("F", None, "Criticidade"), ("G", None, "Meses"), ("H", None, "Última"),
                 ("I", None, "Próxima"), ("J", None, "Contingência"), ("K", None, "Horas prog."), ("L", None, "Situação"), ("M", None, "Quebras"), ("N", None, "Horas paradas"),
                 ("O", None, "Disponib."), ("P", "Q", "Conferência")])
    i1 = rr + 1
    i2 = i1 + len(ex_["infra"]) - 1
    # o registro de manutenções fica mais abaixo, na mesma aba: as linhas são conhecidas de antemão
    m1 = i2 + 4
    m2 = m1 + len(ex_["ocs"]) - 1
    reg = dict(data=f"$B${m1}:$B${m2}", cod=f"$C${m1}:$C${m2}", tipo=f"$D${m1}:$D${m2}", horas=f"$E${m1}:$E${m2}")
    for i in ex_["infra"]:
        rr += 1
        put(ws, f"B{rr}", i["cod"], f=font(10, True, c=MUTED), h="center")
        put(ws, f"C{rr}", i["nome"], f=font(10, True))
        put(ws, f"D{rr}", i["tipo"], f=font(9))
        put(ws, f"E{rr}", i["local"])
        put(ws, f"F{rr}", i["crit"], h="center")
        put(ws, f"G{rr}", i["interv"], h="center")
        put(ws, f"H{rr}", i["ultima"], h="center", fmt=DATE)
        put(ws, f"J{rr}", i["cont"] or None, f=font(9))
        put(ws, f"K{rr}", i["prog"], h="center", fmt=GERAL)
        calc(ws, f"I{rr}", "=" + f_prox(f"H{rr}", f"G{rr}"), b=False, fmt=DATE)
        calc(ws, f"L{rr}", "=" + f_prevsit(f"B{rr}", f"G{rr}", f"H{rr}", f"I{rr}", fim), sz=9)
        calc(ws, f"M{rr}", "=" + f_quebras(f"B{rr}", ini, fim, reg), b=False)
        calc(ws, f"N{rr}", "=" + f_horas(f"B{rr}", ini, fim, reg), b=False, fmt=GERAL)
        calc(ws, f"O{rr}", "=" + f_disp(f"K{rr}", f"N{rr}"), fmt=PCT)
        cells = dict(cod=f"B{rr}", tipo=f"D{rr}", crit=f"F{rr}", interv=f"G{rr}", sit=f"L{rr}", cont=f"J{rr}", disp=f"O{rr}")
        calc(ws, f"P{rr}", "=" + f_infraconf(cells, meta), b=False, sz=9, merge=f"P{rr}:Q{rr}")
        ws.row_dimensions[rr].height = alt([(i["cont"], 34), (i["nome"], 26), (i["tipo"], 18)], minimo=21.75)
    assert rr == i2
    cf_texto(ws, f"L{i1}:L{i2}", f"L{i1}", PREV_CF)
    ws.conditional_formatting.add(f"O{i1}:O{i2}", FormulaRule(formula=[f'AND(ISNUMBER(O{i1}),O{i1}<{meta}-{EPS})'], fill=PatternFill("solid", bgColor=RED, fgColor=RED)))
    cf_warn(ws, f"P{i1}:Q{i2}", f"$P{i1}")

    # ---- manutenções
    rr += 2
    ws.row_breaks.append(Break(id=rr - 1))
    band(ws, rr, "Manutenções", "Q", color=TEAL)
    rr += 1
    sub(ws, rr, [("B", None, "Data"), ("C", None, "Código"), ("D", None, "Tipo"), ("E", None, "Horas paradas"), ("F", "I", "O que aconteceu"), ("J", None, "Causa"),
                 ("K", "M", "Ação"), ("N", None, "Afetou o produto?"), ("O", "P", "Tratamento do produto"), ("Q", None, "Conferência")])
    assert rr + 1 == m1
    codes = f"$B${i1}:$B${i2}"
    for o in sorted(ex_["ocs"], key=lambda o: o["data"]):
        rr += 1
        put(ws, f"B{rr}", o["data"], h="center", fmt=DATE)
        put(ws, f"C{rr}", o["cod"], f=font(10, True, c=MUTED), h="center")
        put(ws, f"D{rr}", o["tipo"], h="center")
        put(ws, f"E{rr}", o["horas"], h="center", fmt=GERAL)
        put(ws, f"F{rr}", o["oque"], merge=f"F{rr}:I{rr}", f=font(9))
        put(ws, f"J{rr}", o["causa"] or None, f=font(9))
        put(ws, f"K{rr}", o["acao"] or None, merge=f"K{rr}:M{rr}", f=font(9))
        put(ws, f"N{rr}", o["afetou"] or None, h="center")
        put(ws, f"O{rr}", o["trat"] or None, merge=f"O{rr}:P{rr}", f=font(9))
        cells = dict(cod=f"C{rr}", data=f"B{rr}", tipo=f"D{rr}", horas=f"E{rr}", causa=f"J{rr}", acao=f"K{rr}", afetou=f"N{rr}", trat=f"O{rr}")
        calc(ws, f"Q{rr}", "=" + f_occonf(cells, "1", codes), b=False, sz=9)
        ws.row_dimensions[rr].height = alt([(o["oque"], 48), (o["causa"], 34), (o["acao"], 40), (o["trat"], 30)], minimo=21.75)
    assert rr == m2
    cf_texto(ws, f"Q{m1}:Q{m2}", f"Q{m1}", [("OK", GREEN), (SEMTRAT, RED)], resto=YELLOW)

    # ---- ambiente
    rr += 2
    band(ws, rr, "Ambiente", "Q", color=PURPLE)
    rr += 1
    sub(ws, rr, [("B", "C", "Fator"), ("D", None, "Tipo"), ("E", None, "Local"), ("F", "H", "Por que afeta o produto"), ("I", None, "Unidade"), ("J", None, "Como é controlado"),
                 ("K", None, "Mínimo"), ("L", None, "Máximo"), ("M", None, "Medição"), ("N", None, "Situação"), ("O", "P", "Ação"), ("Q", None, "Conferência")])
    a1 = rr + 1
    for a in ex_["amb"]:
        rr += 1
        put(ws, f"B{rr}", a["fator"], f=font(10, True), merge=f"B{rr}:C{rr}")
        put(ws, f"D{rr}", a["tipo"], h="center")
        put(ws, f"E{rr}", a["local"])
        put(ws, f"F{rr}", a["porque"], merge=f"F{rr}:H{rr}", f=font(9))
        put(ws, f"I{rr}", a["unid"], h="center", f=font(9))
        put(ws, f"J{rr}", a["controle"], f=font(9))
        for col, k in (("K", "min"), ("L", "max"), ("M", "valor")):
            put(ws, f"{col}{rr}", a[k], h="center", fmt=GERAL)
        calc(ws, f"N{rr}", "=" + f_ambsit(f"B{rr}", f"K{rr}", f"L{rr}", f"M{rr}"), sz=9)
        put(ws, f"O{rr}", a["acao"] or None, merge=f"O{rr}:P{rr}", f=font(9))
        cells = dict(fator=f"B{rr}", porque=f"F{rr}", controle=f"J{rr}", min=f"K{rr}", max=f"L{rr}", sit=f"N{rr}", acao=f"O{rr}")
        calc(ws, f"Q{rr}", "=" + f_ambconf(cells), b=False, sz=9)
        ws.row_dimensions[rr].height = alt([(a["porque"], 36), (a["controle"], 32), (a["acao"], 28), (a["fator"], 36)], minimo=21.75)
    a2 = rr
    cf_texto(ws, f"N{a1}:N{a2}", f"N{a1}", AMB_CF)
    cf_texto(ws, f"Q{a1}:Q{a2}", f"Q{a1}", [("OK", GREEN), (FORASEM, RED)], resto=YELLOW)

    # ---- resumo
    rr += 2
    band(ws, rr, "Resumo automático", "Q")
    Ls = f"L{i1}:L{i2}"
    for k, (text, formula) in enumerate([
        ("Pessoas", f'=COUNTIF(L{p1}:L{p2},"{FALTA}")&" funções com falta de gente, "&-SUMIF(K{p1}:K{p2},"<0")&" pessoas a menos no pico, "'
                    f'&COUNTIF(M{p1}:M{p2},"{SEMQ}")&" sem pessoa qualificada"'),
        ("Preventivas", "=" + '&", "&'.join(f'COUNTIF({Ls},"{x}")&" {x.lower()}"' for x in PREV_SITS)),
        ("Quebras no período", f'=SUM(M{i1}:M{i2})&" corretivas, "&SUM(N{i1}:N{i2})&" horas paradas, "&COUNTIF(O{i1}:O{i2},"<"&({meta}-{EPS}))&" item abaixo da meta"'),
        ("Ambiente", "=" + '&", "&'.join(f'COUNTIF(N{a1}:N{a2},"{x}")&" {x.lower()}"' for x in (DENTRO, FORA, SEMMED))),
        ("Linhas a completar", f'=SUMPRODUCT((M{p1}:M{p2}<>"OK")*1)+SUMPRODUCT((P{i1}:P{i2}<>"OK")*1)+SUMPRODUCT((Q{m1}:Q{m2}<>"OK")*1)+SUMPRODUCT((Q{a1}:Q{a2}<>"OK")*1)'),
    ], 1):
        label(ws, f"B{rr+k}", text, merge=f"B{rr+k}:D{rr+k}", h="right")
        calc(ws, f"E{rr+k}", formula, merge=f"E{rr+k}:Q{rr+k}", h="left")
        ws.row_dimensions[rr + k].height = 21.75
    setup(ws, MUTED, f"B1:Q{rr + 5}")
    return dict(p=(p1, p2), i=(i1, i2), m=(m1, m2), a=(a1, a2), res=rr + 1)


POS = {}
POS["ex1"] = exemplo(wb.create_sheet("Exemplo 1 - Pizzaria"), EX1)
POS["ex2"] = exemplo(wb.create_sheet("Exemplo 2 - Indústria"), EX2)

wb.active = 1
wb.save(OUT)
print("ok", OUT, "| avisos: Pessoas E%d, Infraestrutura E%d, Manutenções E%d, Ambiente E%d, Painel C%d | linhas do painel %s | exemplos %s" % (PAV, IAV, MAV, AAV, NAV, IR, POS))
