# -*- coding: utf-8 -*-
"""Gera Conhecimento-modelo.xlsx no mesmo padrão visual das planilhas anteriores da série."""
import math
import os
import sys

_here = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, _here)
# funções de formatação compartilhadas com o gerador do PDCA (bloco inicial do arquivo)
_src = open(os.path.join(_here, "build_pdca.py"), encoding="utf-8").read()
exec(_src[: _src.index("STATUS = [")])

from openpyxl.chart import BarChart, Series  # noqa: E402
from openpyxl.chart.label import DataLabelList  # noqa: E402
from con_data import (ABERTO, ALTA, ATRASADO, CABECA, CHECK, CON_SITS, CONTROLE, CRITICO, CRITS, ESCRITO, EX1, EX2, FORAPRAZO, FORMAS, INCORP, MEDIA,  # noqa: E402
                      NAO, NOPRAZO, ORIGENS, PEND, PRAZO_LICAO, RISCO, SIM, TIPOS_AT, TIPOS_CAUSA)

BLUE, TEAL, AMBER, PURPLE, REDC = "2B5C8A", "1E7B73", "A96A12", "7A4A9A", "B0413E"
S2_T = "F8EBCB"
YESNO = ["Sim", "Parcial", "Não"]
YESNO_CF = [("Sim", GREEN), ("Parcial", YELLOW), ("Não", RED)]
NK, NL, NP, NA = 40, 60, 20, 80
KS, LS, PS, AS = "Conhecimento", "'Lições'", "'Pós-entrega'", "Atendimentos"
K1, K2 = 8, 8 + NK - 1          # conhecimento
L1, L2 = 7, 7 + NL - 1          # lições
P1, P2 = 7, 7 + NP - 1          # política de pós-entrega
A1, A2 = 7, 7 + NA - 1          # atendimentos
GERAL = "#,##0.##"
MOEDA = '"R$" #,##0'
CON_CF = [(CRITICO, RED), (RISCO, S2_T), (CONTROLE, GREEN)]
AT_CF = [(NOPRAZO, GREEN), (FORAPRAZO, RED), (ABERTO, S2_T), (ATRASADO, RED)]
LIC_CF = [(INCORP, GREEN), (PEND, S2_T)]
SEMACAO = "Falta a ação de transferência"
VENCIDA = "Prazo da ação vencido"
VELHA = f"Pendente há mais de {PRAZO_LICAO} dias"
MENOR = "Prazo aceito menor que o mínimo"


def note(ws, ref, text):
    c = Comment(text, "Modelo Conhecimento")
    c.width, c.height = 300, 120
    ws[ref].comment = c


def cf_warn(ws, rng_, first, ok_values=("OK",)):
    cond = ",".join(f'{first}="{v}"' for v in ok_values)
    ws.conditional_formatting.add(rng_, FormulaRule(formula=[f"OR({cond})"], fill=PatternFill("solid", bgColor=GREEN, fgColor=GREEN)))
    ws.conditional_formatting.add(rng_, FormulaRule(formula=[f"AND(LEN({first})>0,NOT(OR({cond})))"],
                                  fill=PatternFill("solid", bgColor=YELLOW, fgColor=YELLOW)))


def cf_texto(ws, rng_, first, pares, resto=None):
    for t, cor in pares:
        ws.conditional_formatting.add(rng_, FormulaRule(formula=[f'{first}="{t}"'], stopIfTrue=True, fill=PatternFill("solid", bgColor=cor, fgColor=cor)))
    if resto:
        ws.conditional_formatting.add(rng_, FormulaRule(formula=[f"LEN({first})>0"], fill=PatternFill("solid", bgColor=resto, fgColor=resto)))


def hint_row(ws, row, cells, height=30):
    for ref, text in cells:
        put(ws, f"{ref}{row}", text, f=font(9, i=True, c=MUTED), bg=GRAY, h="center")
    ws.row_dimensions[row].height = height


def alt(pares, minimo=21.75, linha=12):
    n = max(max(1, math.ceil(len(str(t or "")) / max(c, 1))) for t, c in pares)
    return max(minimo, n * linha + 7)


def lista_ref(ws, rng, ref, prompt):
    dv = DataValidation(type="list", formula1=f"={ref}", allow_blank=True)
    dv.promptTitle, dv.prompt = "Opções", prompt
    dv.showInputMessage = True
    ws.add_data_validation(dv)
    dv.add(rng)


# ------------------------------------------------------------------ fórmulas compartilhadas pelas abas de entrada e pelos exemplos
def f_pontos(forma, pessoas, saida):
    return (f'IF(OR({forma}="",{pessoas}=""),"",IF({forma}="{CABECA}",2,IF({forma}="{ESCRITO}",1,0))'
            f'+IF({pessoas}<=1,2,IF({pessoas}=2,1,0))+IF({saida}="{SIM}",1,0))')


def f_consit(crit, pts):
    return (f'IF(OR({pts}="",{crit}=""),"",IF(AND({crit}="{ALTA}",{pts}>=3),"{CRITICO}",'
            f'IF(AND(OR({crit}="{ALTA}",{crit}="{MEDIA}"),{pts}>=2),"{RISCO}","{CONTROLE}")))')


def f_conconf(c, ref):
    onde = f'IF(AND({c["forma"]}<>"{CABECA}",{c["onde"]}=""),"Falta onde está escrito","OK")'
    return (f'IF({c["cod"]}="","",IF({c["crit"]}="","Falta a criticidade",IF({c["forma"]}="","Falta a forma",IF({c["pessoas"]}="","Falta quantas pessoas sabem",'
            f'IF(OR({c["sit"]}="{CRITICO}",{c["sit"]}="{RISCO}"),IF({c["acao"]}="","{SEMACAO}",IF({c["prazo"]}="","Falta o prazo da ação",'
            f'IF(AND(ISNUMBER({ref}),{c["prazo"]}<{ref}),"{VENCIDA}",{onde}))),{onde})))))')


def f_licsit(num_, incorp):
    return f'IF({num_}="","",IF({incorp}="","{PEND}","{INCORP}"))'


def f_licdias(num_, data, incorp, ref):
    return f'IF(OR({num_}="",{incorp}<>"",{data}=""),"",IF(ISNUMBER({ref}),{ref}-{data},""))'


def f_licconf(c, codes):
    return (f'IF({c["num"]}="","",IF({c["data"]}="","Falta a data",IF({c["oque"]}="","Falta o que aconteceu",IF({c["aprend"]}="","Falta o que aprendemos",'
            f'IF({c["onde"]}="","Falta onde incorporar",IF(AND({c["cod"]}<>"",ISNA(MATCH({c["cod"]},{codes},0))),"Conhecimento não cadastrado",'
            f'IF(AND({c["sit"]}="{PEND}",N({c["dias"]})>{PRAZO_LICAO}),"{VELHA}","OK")))))))')


def f_polconf(c):
    return (f'IF({c["prod"]}="","",IF({c["min"]}="","Falta o prazo mínimo para reclamar",IF({c["aceito"]}="","Falta o prazo aceito",'
            f'IF({c["aceito"]}<{c["min"]},"{MENOR}",IF({c["ativ"]}="","Falta o que se faz",IF({c["resp"]}="","Falta o prazo de resposta","OK"))))))')


def f_prazo(prod, prods, resps):
    return f'IF({prod}="","",IFERROR(IF(INDEX({resps},MATCH({prod},{prods},0))="","",INDEX({resps},MATCH({prod},{prods},0))),""))'


def f_atdias(abert, resol, ref):
    return f'IF({abert}="","",IF({resol}<>"",{resol}-{abert},IF(ISNUMBER({ref}),{ref}-{abert},"")))'


def f_atsit(prazo, dias, resol):
    return (f'IF(OR({prazo}="",{dias}=""),"",IF({resol}<>"",IF({dias}<={prazo},"{NOPRAZO}","{FORAPRAZO}"),'
            f'IF({dias}<={prazo},"{ABERTO}","{ATRASADO}")))')


def f_atconf(c, vazio, prods, lics):
    causa = ",".join(f'{c["tipo"]}="{t}"' for t in TIPOS_CAUSA)
    return (f'IF({vazio},"",IF(ISNA(MATCH({c["prod"]},{prods},0)),"Produto sem política",IF({c["abert"]}="","Falta a data de abertura",'
            f'IF({c["tipo"]}="","Falta o tipo",IF(AND({c["resol"]}<>"",OR({causa}),{c["causa"]}=""),"Falta a causa",'
            f'IF(AND({c["licao"]}<>"",ISNA(MATCH({c["licao"]},{lics},0))),"Lição não encontrada","OK"))))))')


def chart_style(ch):
    ch.style = 2
    ch.title = None
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ch.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill="D3DBE0", w=9525))
    ch.x_axis.graphicalProperties = GraphicalProperties(ln=LineProperties(solidFill="D3DBE0", w=9525))
    ch.y_axis.graphicalProperties = GraphicalProperties(ln=LineProperties(noFill=True))


def colunas(ws, anchor, r1, r2, c_cat, c_val, width=16, height=7.5):
    ch = BarChart()
    ch.type = "col"
    ch.gapWidth = 70
    ch.height, ch.width = height, width
    chart_style(ch)
    ch.legend = None
    s = Series(Reference(ws, min_col=c_val, min_row=r1, max_row=r2), title="Conhecimentos")
    s.graphicalProperties.solidFill = BLUE
    s.graphicalProperties.line.noFill = True
    s.dLbls = DataLabelList()
    s.dLbls.showVal = True
    s.dLbls.showSerName = s.dLbls.showCatName = s.dLbls.showLegendKey = False
    ch.append(s)
    ch.set_categories(Reference(ws, min_col=c_cat, min_row=r1, max_row=r2))
    ch.y_axis.scaling.min = 0
    ch.y_axis.majorUnit = 1
    ws.add_chart(ch, anchor)


def resumo_aba(ws, s, itens, last, merge_lab="D"):
    """Faixa de resumo com o aviso na última linha; devolve a linha do aviso."""
    band(ws, s, "Resumo automático", last)
    col = chr(ord(merge_lab) + 1)
    for k, (text, formula) in enumerate(itens, 1):
        aviso = text == "Aviso"
        label(ws, f"B{s+k}", text, merge=f"B{s+k}:{merge_lab}{s+k}", h="right")
        calc(ws, f"{col}{s+k}", formula, merge=f"{col}{s+k}:{last}{s+k}", sz=9 if aviso else 10, b=not aviso, h="left")
        ws.row_dimensions[s + k].height = 21.75
    av = s + len(itens)
    cf_warn(ws, f"{col}{av}:{last}{av}", f"{col}{av}")
    return av


def conta_por(rng, valores):
    return "=" + '&", "&'.join(f'COUNTIF({rng},"{x}")&" {x.lower()}"' for x in valores)


# ================================================================== pasta de trabalho
wb = Workbook()
wb.properties.title = "Conhecimento organizacional e pós-entrega — Modelo"
wb.properties.creator = "Modelo Conhecimento"
wb.properties.language = "pt-BR"

# ------------------------------------------------------------------ Instruções
ws = wb.active
ws.title = "Instruções"
widths(ws, {"A": 2, "B": 26, "C": 90, "D": 2})
title(ws, "Conhecimento organizacional e pós-entrega — Como usar esta planilha",
      "Modelo para mapear o conhecimento, registrar as lições aprendidas, definir o pós-entrega e acompanhar os atendimentos.", "C")
r = 4


def section(text):
    global r
    band(ws, r, text, "C", sz=11)
    r += 1


def line(k, v, kbg=GRAY, vbg=None, kf=None, kh="left", height=None):
    global r
    put(ws, f"B{r}", k, f=kf or font(10, True), bg=kbg, h=kh)
    put(ws, f"C{r}", v, bg=vbg)
    ws.row_dimensions[r].height = height or (19.5 if len(v) <= 96 else (31.5 if len(v) <= 192 else 45.75))
    r += 1


section("Legenda: onde preencher")
line("Amarelo-claro", "Células de entrada. É aqui que você digita.", vbg=INPUT)
line("Cinza", "Células calculadas ou fixas (pontos, situações, dias, prazos, conferências). Não altere.", vbg=GRAY)
line("Cinza em itálico", "Linha de exemplo (marcada com “Ex.”). Mostra o formato esperado e não entra nas contagens.", vbg=GRAY)
line("Listas suspensas", "Criticidade, forma, saída prevista, origem, produto, tipo de atendimento e checklist aceitam apenas as opções da lista.")
line("Comentários", "Células com triângulo vermelho no canto trazem uma dica de preenchimento. Passe o mouse sobre elas.")
r += 1
section("As regras de cálculo")
for k, text in [
    ("Pontos de exposição", "Forma: só na cabeça 2, escrito 1, escrito e treinado 0. Pessoas que sabem fazer sozinhas: uma ou nenhuma 2, duas 1, três ou mais 0. Saída prevista: mais 1."),
    (CRITICO, "Criticidade alta com 3 pontos ou mais."),
    (RISCO, "Criticidade alta ou média com 2 pontos ou mais, quando não é crítico."),
    (CONTROLE, "Os demais, inclusive todo conhecimento de criticidade baixa."),
    ("Lição pendente", f"Sem data de incorporação. Vira aviso quando passa de {PRAZO_LICAO} dias da data do caso."),
    ("Prazo do atendimento", "O prazo de resposta do produto, na aba Pós-entrega."),
    ("Dias do atendimento", "Da abertura à resolução. Se ainda está aberto, da abertura à data da leitura."),
    ("Causa", f"Obrigatória quando o atendimento está resolvido e é {', '.join(t.lower() for t in TIPOS_CAUSA)}."),
]:
    line(k, text)
r += 1
section("Ordem recomendada de preenchimento")
for k, text in enumerate([
    "Aba Conhecimento: informe a data da leitura.",
    "Aba Conhecimento: liste o que cada processo precisa saber, com criticidade, forma, onde está, quem sabe e saída prevista. Para o que estiver crítico ou em risco, a ação e o prazo.",
    "Aba Lições: registre os casos que ensinaram alguma coisa, com o destino e o responsável.",
    "Aba Pós-entrega: a política de cada produto, com o mínimo da lei ou do contrato, o prazo aceito, o que se faz e o prazo de resposta.",
    "Aba Atendimentos: as reclamações, garantias, devoluções, assistências e orientações do período.",
    "Aba Painel: leia os avisos. Depois, valide na aba Checklist.",
], 1):
    line(f"Passo {k}", text)
r += 1
section("Abas da planilha")
for k, text in [
    (KS, f"A data da leitura e até {NK} conhecimentos. Calcula os pontos, a situação e a conferência."),
    ("Lições", f"Até {NL} lições aprendidas. Calcula a situação, os dias pendentes e a conferência."),
    ("Pós-entrega", f"Até {NP} produtos ou serviços, com a política de cada um. Calcula a conferência."),
    (AS, f"Até {NA} atendimentos. Calcula o prazo, os dias, a situação e a conferência."),
    ("Painel", "Os números das quatro abas, o gráfico do conhecimento e o aviso."),
    ("Checklist", "Doze verificações do conhecimento e do pós-entrega, com percentual de conclusão."),
    ("Exemplo 1 - Pizzaria", "A loja, lida em 31/05/2027."),
    ("Exemplo 2 - Indústria", "A fábrica, lida em 30/09/2027."),
]:
    line(k, text)
r += 1
section("Regras e premissas do modelo")
for k, text in [
    ("Pessoas que sabem", "Conte quem faz sozinho, com o resultado certo. Quem só leu o documento não conta."),
    ("Mínimo para reclamar", "O prazo da lei, para o consumidor, ou o do contrato, entre empresas. Confira a legislação e os contratos vigentes."),
    ("Prazos em dias", "Todos os prazos são em dias corridos."),
    ("Exemplos", "Os dados das abas de exemplo são ilustrativos, criados para este material."),
]:
    line(k, text)
setup(ws, INK, f"B1:C{r}", landscape=False, fit_height=True)

# ------------------------------------------------------------------ Conhecimento
ws = wb.create_sheet(KS)
widths(ws, {"A": 2, "B": 5, "C": 9, "D": 30, "E": 14, "F": 12, "G": 16, "H": 24, "I": 10, "J": 10, "K": 34, "L": 12, "M": 9, "N": 13, "O": 30, "P": 2})
title(ws, "Mapa do conhecimento", "Um conhecimento por linha: quão exposto ele está, e o que se faz para protegê-lo.", "O")
label(ws, "B4", "Data da leitura", merge="B4:D4")
inp(ws, "E4", h="center", fmt=DATE)
put(ws, "F4", "Usada para saber que ações venceram e há quanto tempo cada lição e cada atendimento estão abertos.", f=font(9, i=True, c=MUTED), bg=GRAY, merge="F4:O4")
dv_date(ws, "E4")
REF = f"{KS}!$E$4"
for col, text in zip("BCDEFGHIJKLMNO", ["#", "Código", "Conhecimento", "Processo", "Criticidade", "Forma", "Onde está", "Pessoas que sabem", "Saída prevista?",
                                        "Ação de transferência", "Prazo", "Pontos", "Situação", "Conferência"]):
    head(ws, f"{col}6", text)
ws.row_dimensions[6].height = 33
hint_row(ws, 7, [("B", ""), ("C", "K-01"), ("D", "O que é preciso saber"), ("E", ""), ("F", "Lista"), ("G", "Lista"), ("H", "Documento"), ("I", "Fazem sozinhas"),
                 ("J", "Lista"), ("K", "Se crítico ou em risco"), ("L", "Data"), ("M", "Calculado"), ("N", "Calculada"), ("O", "Calculada")])
for k in range(NK):
    rr = K1 + k
    num(ws, f"B{rr}", k + 1)
    inp(ws, f"C{rr}", h="center")
    for col in "DEHK":
        inp(ws, f"{col}{rr}")
    for col in "FGJ":
        inp(ws, f"{col}{rr}", h="center")
    inp(ws, f"I{rr}", h="center", fmt=GERAL)
    inp(ws, f"L{rr}", h="center", fmt=DATE)
    calc(ws, f"M{rr}", "=" + f_pontos(f"G{rr}", f"I{rr}", f"J{rr}"))
    calc(ws, f"N{rr}", "=" + f_consit(f"F{rr}", f"M{rr}"), sz=9)
    cells = dict(cod=f"C{rr}", crit=f"F{rr}", forma=f"G{rr}", onde=f"H{rr}", pessoas=f"I{rr}", acao=f"K{rr}", prazo=f"L{rr}", sit=f"N{rr}")
    calc(ws, f"O{rr}", "=" + f_conconf(cells, "$E$4"), b=False, sz=9)
    ws.row_dimensions[rr].height = 30
dv_list(ws, f"F{K1}:F{K2}", CRITS, "Alta: sem ele, o produto sai errado. Média: atrapalha. Baixa: pouco efeito no produto")
dv_list(ws, f"G{K1}:G{K2}", FORMAS, "Só na cabeça, escrito, ou escrito e treinado")
dv_list(ws, f"J{K1}:J{K2}", [SIM, NAO], "Alguém que sabe tem saída prevista: aposentadoria, mudança, fim de contrato?")
dv_number(ws, f"I{K1}:I{K2}")
dv_date(ws, f"L{K1}:L{K2}")
cf_texto(ws, f"N{K1}:N{K2}", f"N{K1}", CON_CF)
cf_warn(ws, f"O{K1}:O{K2}", f"O{K1}")
note(ws, "I6", "Quantas pessoas fazem a tarefa sozinhas, com o resultado certo. Quem só leu o documento não conta.")
note(ws, "G6", "Só na cabeça: não está em lugar nenhum. Escrito: há documento. Escrito e treinado: há documento, e outras pessoas já fizeram sozinhas.")
s = K2 + 2
N_ = f"N{K1}:N{K2}"
KAV = resumo_aba(ws, s, [
    ("Conhecimentos mapeados", f"=COUNTA(C{K1}:C{K2})"),
    ("Por situação", conta_por(N_, CON_SITS)),
    ("Aviso", f'=IF(E{s+1}=0,"Liste o conhecimento de cada processo",IF(E4="","Informe a data da leitura",IF(COUNTIF(O{K1}:O{K2},"{SEMACAO}")>0,"Há conhecimento exposto sem ação de transferência",'
              f'IF(COUNTIF(O{K1}:O{K2},"{VENCIDA}")>0,"Há ação de transferência vencida",'
              f'IF(SUMPRODUCT((O{K1}:O{K2}<>"")*(O{K1}:O{K2}<>"OK"))>0,"Há linha a completar: veja a coluna Conferência","OK")))))'),
], "O")
ws.freeze_panes = f"E{K1}"
setup(ws, BLUE, f"B1:O{KAV}")

# ------------------------------------------------------------------ Lições
ws = wb.create_sheet("Lições")
widths(ws, {"A": 2, "B": 5, "C": 8, "D": 12, "E": 15, "F": 30, "G": 32, "H": 26, "I": 18, "J": 12, "K": 12, "L": 12, "M": 9, "N": 30, "O": 2})
title(ws, "Lições aprendidas", "Uma lição por linha: o caso, o que se aprendeu, para onde vai e quem leva.", "N")
for col, text in zip("BCDEFGHIJKLMN", ["#", "Nº", "Data", "Origem", "O que aconteceu", "O que aprendemos", "Onde incorporar", "Responsável", "Conhecimento",
                                       "Incorporada em", "Situação", "Dias pendente", "Conferência"]):
    head(ws, f"{col}4", text)
ws.row_dimensions[4].height = 33
hint_row(ws, 5, [("B", ""), ("C", "L-01"), ("D", "Data do caso"), ("E", "Lista"), ("F", ""), ("G", "A regra nova"), ("H", "Documento ou treinamento"), ("I", ""),
                 ("J", "Código, se houver"), ("K", "Data"), ("L", "Calculada"), ("M", "Calculado"), ("N", "Calculada")])
ex(ws, "B6", "Ex.", h="center")
for col, v, h_, fmt in [("C", "L-01", "center", None), ("D", date(2027, 3, 5), "center", DATE), ("E", "Reclamação", "center", None), ("F", "Pizza chegou fria", "left", None),
                        ("G", "Bolsa térmica fechada até a entrega", "left", None), ("H", "Instrução IT-EXP-01", "left", None), ("I", "Gerente", "left", None),
                        ("J", "K-04", "center", None), ("K", date(2027, 3, 15), "center", DATE), ("L", INCORP, "center", None), ("M", None, "center", None), ("N", "OK", "center", None)]:
    ex(ws, f"{col}6", v, h=h_, fmt=fmt)
ws.row_dimensions[6].height = 30
KCODES = f"{KS}!$C${K1}:$C${K2}"
for k in range(NL):
    rr = L1 + k
    num(ws, f"B{rr}", k + 1)
    inp(ws, f"C{rr}", h="center")
    inp(ws, f"D{rr}", h="center", fmt=DATE)
    inp(ws, f"E{rr}", h="center")
    for col in "FGHI":
        inp(ws, f"{col}{rr}")
    inp(ws, f"J{rr}", h="center")
    inp(ws, f"K{rr}", h="center", fmt=DATE)
    calc(ws, f"L{rr}", "=" + f_licsit(f"C{rr}", f"K{rr}"), sz=9)
    calc(ws, f"M{rr}", "=" + f_licdias(f"C{rr}", f"D{rr}", f"K{rr}", REF), b=False)
    cells = dict(num=f"C{rr}", data=f"D{rr}", oque=f"F{rr}", aprend=f"G{rr}", onde=f"H{rr}", cod=f"J{rr}", sit=f"L{rr}", dias=f"M{rr}")
    calc(ws, f"N{rr}", "=" + f_licconf(cells, KCODES), b=False, sz=9)
    ws.row_dimensions[rr].height = 30
dv_date(ws, f"D{L1}:D{L2}")
dv_date(ws, f"K{L1}:K{L2}")
dv_list(ws, f"E{L1}:E{L2}", ORIGENS, "De onde veio a lição")
lista_ref(ws, f"J{L1}:J{L2}", KCODES, "Código do conhecimento, da aba Conhecimento (opcional)")
cf_texto(ws, f"L{L1}:L{L2}", f"L{L1}", LIC_CF)
cf_texto(ws, f"N{L1}:N{L2}", f"N{L1}", [("OK", GREEN), (VELHA, RED)], resto=YELLOW)
note(ws, "G4", "A regra nova, numa frase: o que passa a ser feito de outro jeito. Não basta contar o caso.")
note(ws, "H4", "O documento, a instrução, a embalagem ou o treinamento em que a lição vai entrar. Sem destino, a lição não muda o processo.")
s = L2 + 2
LAV = resumo_aba(ws, s, [
    ("Lições registradas", f"=COUNTA(C{L1}:C{L2})"),
    ("Por situação", conta_por(f"L{L1}:L{L2}", (INCORP, PEND))),
    ("Aviso", f'=IF(E{s+1}=0,"Registre as lições aprendidas",IF(COUNTIF(N{L1}:N{L2},"{VELHA}")>0,"Há lição pendente há mais de {PRAZO_LICAO} dias",'
              f'IF(SUMPRODUCT((N{L1}:N{L2}<>"")*(N{L1}:N{L2}<>"OK"))>0,"Há lição a completar: veja a coluna Conferência","OK")))'),
], "N")
ws.freeze_panes = f"F{L1}"
setup(ws, TEAL, f"B1:N{LAV}")

# ------------------------------------------------------------------ Pós-entrega
ws = wb.create_sheet("Pós-entrega")
widths(ws, {"A": 2, "B": 5, "C": 28, "D": 26, "E": 13, "F": 12, "G": 40, "H": 12, "I": 30, "J": 32, "K": 2})
title(ws, "Política de pós-entrega", "Um produto ou serviço por linha: o que se faz depois da entrega, e com que prazos.", "J")
for col, text in zip("BCDEFGHIJ", ["#", "Produto ou serviço", "Vida útil", "Mínimo para reclamar (dias)", "Prazo aceito (dias)", "O que se faz depois da entrega",
                                   "Prazo de resposta (dias)", "O que se quer evitar no uso", "Conferência"]):
    head(ws, f"{col}4", text)
ws.row_dimensions[4].height = 45
hint_row(ws, 5, [("B", ""), ("C", ""), ("D", "Validade, durabilidade"), ("E", "Lei ou contrato"), ("F", "O que a casa aceita"), ("G", "Troca, laudo, visita, orientação"),
                 ("H", "Dias corridos"), ("I", "Consequência indesejada"), ("J", "Calculada")])
ex(ws, "B6", "Ex.", h="center")
for col, v, h_ in [("C", "Massa congelada para levar", "left"), ("D", "60 dias no freezer", "left"), ("E", 30, "center"), ("F", 30, "center"),
                   ("G", "Troca ou devolução do valor", "left"), ("H", 2, "center"), ("I", "Descongelar e congelar de novo", "left"), ("J", "OK", "center")]:
    ex(ws, f"{col}6", v, h=h_)
ws.row_dimensions[6].height = 30
for k in range(NP):
    rr = P1 + k
    num(ws, f"B{rr}", k + 1)
    for col in "CDGI":
        inp(ws, f"{col}{rr}")
    for col in "EFH":
        inp(ws, f"{col}{rr}", h="center", fmt=GERAL)
    calc(ws, f"J{rr}", "=" + f_polconf(dict(prod=f"C{rr}", min=f"E{rr}", aceito=f"F{rr}", ativ=f"G{rr}", resp=f"H{rr}")), b=False, sz=9)
    ws.row_dimensions[rr].height = 30
dv_number(ws, f"E{P1}:F{P2}")
dv_number(ws, f"H{P1}:H{P2}")
cf_texto(ws, f"J{P1}:J{P2}", f"J{P1}", [("OK", GREEN), (MENOR, RED)], resto=YELLOW)
note(ws, "E4", "O menor prazo que o cliente tem para reclamar, pela lei (consumidor) ou pelo contrato (entre empresas). Confira a legislação e os contratos vigentes.")
note(ws, "H4", "Em quantos dias corridos a organização se compromete a resolver um atendimento deste produto.")
s = P2 + 2
PAV = resumo_aba(ws, s, [
    ("Produtos com política", f"=COUNTA(C{P1}:C{P2})"),
    ("Aviso", f'=IF(E{s+1}=0,"Escreva a política de cada produto",IF(COUNTIF(J{P1}:J{P2},"{MENOR}")>0,"Há prazo aceito menor que o mínimo da lei ou do contrato",'
              f'IF(SUMPRODUCT((J{P1}:J{P2}<>"")*(J{P1}:J{P2}<>"OK"))>0,"Há política a completar: veja a coluna Conferência","OK")))'),
], "J")
ws.freeze_panes = f"D{P1}"
setup(ws, AMBER, f"B1:J{PAV}")

# ------------------------------------------------------------------ Atendimentos
ws = wb.create_sheet(AS)
widths(ws, {"A": 2, "B": 5, "C": 8, "D": 12, "E": 18, "F": 24, "G": 16, "H": 30, "I": 12, "J": 11, "K": 30, "L": 8, "M": 8, "N": 8, "O": 15, "P": 26, "Q": 2})
title(ws, "Atendimentos depois da entrega", "Um atendimento por linha. O prazo vem da política do produto, na aba Pós-entrega.", "P")
for col, text in zip("BCDEFGHIJKLMNOP", ["#", "Nº", "Abertura", "Cliente", "Produto", "Tipo", "Relato", "Resolução", "Custo", "Causa", "Lição", "Prazo", "Dias",
                                         "Situação", "Conferência"]):
    head(ws, f"{col}4", text)
ws.row_dimensions[4].height = 33
hint_row(ws, 5, [("B", ""), ("C", "A-01"), ("D", "Data"), ("E", ""), ("F", "Lista"), ("G", "Lista"), ("H", "O que o cliente disse"), ("I", "Data, se resolvido"),
                 ("J", "R$"), ("K", "Por que aconteceu"), ("L", "Nº"), ("M", "Calculado"), ("N", "Calculado"), ("O", "Calculada"), ("P", "Calculada")])
ex(ws, "B6", "Ex.", h="center")
for col, v, h_, fmt in [("C", "A-01", "center", None), ("D", date(2027, 5, 3), "center", DATE), ("E", "Cliente do bairro A", "left", None), ("F", "Pizza entregue", "left", None),
                        ("G", "Reclamação", "center", None), ("H", "Pizza fria", "left", None), ("I", date(2027, 5, 3), "center", DATE), ("J", 45, "center", MOEDA),
                        ("K", "Entregador com 4 pedidos", "left", None), ("L", "L-01", "center", None), ("M", 1, "center", None), ("N", 0, "center", None),
                        ("O", NOPRAZO, "center", None), ("P", "OK", "center", None)]:
    ex(ws, f"{col}6", v, h=h_, fmt=fmt)
ws.row_dimensions[6].height = 30
PRODS, RESPS = f"{PS}!$C${P1}:$C${P2}", f"{PS}!$H${P1}:$H${P2}"
LNUMS = f"{LS}!$C${L1}:$C${L2}"
for k in range(NA):
    rr = A1 + k
    num(ws, f"B{rr}", k + 1)
    inp(ws, f"C{rr}", h="center")
    inp(ws, f"D{rr}", h="center", fmt=DATE)
    for col in "EFHK":
        inp(ws, f"{col}{rr}")
    inp(ws, f"G{rr}", h="center")
    inp(ws, f"I{rr}", h="center", fmt=DATE)
    inp(ws, f"J{rr}", h="center", fmt=MOEDA)
    inp(ws, f"L{rr}", h="center")
    calc(ws, f"M{rr}", "=" + f_prazo(f"F{rr}", PRODS, RESPS), b=False)
    calc(ws, f"N{rr}", "=" + f_atdias(f"D{rr}", f"I{rr}", REF))
    calc(ws, f"O{rr}", "=" + f_atsit(f"M{rr}", f"N{rr}", f"I{rr}"), sz=9)
    cells = dict(prod=f"F{rr}", abert=f"D{rr}", tipo=f"G{rr}", resol=f"I{rr}", causa=f"K{rr}", licao=f"L{rr}")
    calc(ws, f"P{rr}", "=" + f_atconf(cells, f"AND(C{rr}=\"\",COUNTA(D{rr}:L{rr})=0)", PRODS, LNUMS), b=False, sz=9)
    ws.row_dimensions[rr].height = 30
dv_date(ws, f"D{A1}:D{A2}")
dv_date(ws, f"I{A1}:I{A2}")
lista_ref(ws, f"F{A1}:F{A2}", PRODS, "Produto ou serviço, da aba Pós-entrega")
dv_list(ws, f"G{A1}:G{A2}", TIPOS_AT, "Reclamação, garantia, devolução, assistência técnica ou orientação de uso")
dv_number(ws, f"J{A1}:J{A2}")
cf_texto(ws, f"O{A1}:O{A2}", f"O{A1}", AT_CF)
cf_warn(ws, f"P{A1}:P{A2}", f"P{A1}")
note(ws, "K4", f"Obrigatória quando o atendimento está resolvido e é {', '.join(t.lower() for t in TIPOS_CAUSA)}. Sem causa, a reclamação volta.")
note(ws, "L4", "O número da lição aprendida que o atendimento gerou ou confirmou, na aba Lições.")
s = A2 + 2
O_ = f"O{A1}:O{A2}"
AAV = resumo_aba(ws, s, [
    ("Atendimentos", f'=COUNTA(D{A1}:D{A2})&" atendimentos, custo de R$ "&TEXT(SUM(J{A1}:J{A2}),"0")'),
    ("Por situação", conta_por(O_, (NOPRAZO, FORAPRAZO, ABERTO, ATRASADO))),
    ("Aviso", f'=IF(COUNTA(D{A1}:D{A2})=0,"Lance os atendimentos",IF(COUNTIF({O_},"{ATRASADO}")>0,"Há atendimento aberto e atrasado",'
              f'IF(SUMPRODUCT((P{A1}:P{A2}<>"")*(P{A1}:P{A2}<>"OK"))>0,"Há atendimento a completar: veja a coluna Conferência",'
              f'IF(COUNTIF({O_},"{FORAPRAZO}")>0,"Há atendimento resolvido fora do prazo: veja por quê","OK"))))'),
], "P")
ws.freeze_panes = f"E{A1}"
setup(ws, PURPLE, f"B1:P{AAV}")

# ------------------------------------------------------------------ Painel
ws = wb.create_sheet("Painel")
widths(ws, {"A": 2, "B": 38, "C": 14, "D": 14, "E": 14, "F": 34, "G": 2})
title(ws, "Painel do conhecimento e do pós-entrega", "Nada a preencher: tudo vem das outras abas.", "F")
label(ws, "B3", "Data da leitura")
calc(ws, "C3", f'=IF({REF}="","",{REF})', fmt=DATE)
for col, text, m in [("B", "Indicador", None), ("C", "Resultado", None), ("D", "Como se lê", "D5:F5")]:
    put(ws, f"{col}5", text, f=font(10, True, c=WHITE), bg=INK, h="center", merge=m)
ws.row_dimensions[5].height = 21.75
Kn, Ko = (f"{KS}!{c}{K1}:{c}{K2}" for c in "NO")
Ll, Ln = (f"{LS}!{c}{L1}:{c}{L2}" for c in "LN")
Pj = f"{PS}!J{P1}:J{P2}"
Ao, Ap, Aj = (f"{AS}!{c}{A1}:{c}{A2}" for c in "OPJ")
IND = [
    ("CONHECIMENTO", None, None, None),
    ("Conhecimentos mapeados", f"=COUNTA({KS}!C{K1}:C{K2})", "Da aba Conhecimento.", None),
    (CRITICO, f'=COUNTIF({Kn},"{CRITICO}")', "Alta criticidade e muito exposto.", None),
    (RISCO, f'=COUNTIF({Kn},"{RISCO}")', "Exposto, alta ou média criticidade.", None),
    (CONTROLE, f'=COUNTIF({Kn},"{CONTROLE}")', "Protegido, ou de pouco efeito.", None),
    ("Expostos sem ação de transferência", f'=COUNTIF({Ko},"{SEMACAO}")', "Ninguém está transferindo.", None),
    ("Ações de transferência vencidas", f'=COUNTIF({Ko},"{VENCIDA}")', "O prazo passou.", None),
    ("LIÇÕES APRENDIDAS", None, None, None),
    ("Lições registradas", f"=COUNTA({LS}!C{L1}:C{L2})", "Da aba Lições.", None),
    ("Lições pendentes", f'=COUNTIF({Ll},"{PEND}")', "Sem data de incorporação.", None),
    (f"Pendentes há mais de {PRAZO_LICAO} dias", f'=COUNTIF({Ln},"{VELHA}")', "Paradas: dar dono e prazo.", None),
    ("PÓS-ENTREGA", None, None, None),
    ("Produtos com política", f"=COUNTA({PS}!C{P1}:C{P2})", "Da aba Pós-entrega.", None),
    ("Prazo aceito menor que o mínimo", f'=COUNTIF({Pj},"{MENOR}")', "Abaixo da lei ou do contrato.", None),
    ("Atendimentos", f"=COUNTA({AS}!D{A1}:D{A2})", "Da aba Atendimentos.", None),
    ("Abertos", f'=COUNTIF({Ao},"{ABERTO}")+COUNTIF({Ao},"{ATRASADO}")', "Ainda sem resolução.", None),
    ("Fora do prazo, resolvidos ou não", f'=COUNTIF({Ao},"{FORAPRAZO}")+COUNTIF({Ao},"{ATRASADO}")', "Passaram do prazo de resposta.", None),
    ("Resolvidos no prazo", f'=IFERROR(COUNTIF({Ao},"{NOPRAZO}")/(COUNTIF({Ao},"{NOPRAZO}")+COUNTIF({Ao},"{FORAPRAZO}")),"")', "Entre os resolvidos.", "0%"),
    ("Custo dos atendimentos", f"=SUM({Aj})", "Reposições, descontos, fretes, horas.", MOEDA),
    ("Linhas a completar", f'=SUMPRODUCT(({Ko}<>"")*({Ko}<>"OK"))+SUMPRODUCT(({Ln}<>"")*({Ln}<>"OK"))+SUMPRODUCT(({Pj}<>"")*({Pj}<>"OK"))+SUMPRODUCT(({Ap}<>"")*({Ap}<>"OK"))',
     "Conferências diferentes de OK nas quatro abas.", None),
]
IR = {}
rr = 6
for nome, formula, leit, fmt in IND:
    if formula is None:
        put(ws, f"B{rr}", nome, f=font(9, True, c=MUTED), bg=None, box=False, merge=f"B{rr}:F{rr}")
        ws.row_dimensions[rr].height = 21.75
        rr += 1
        continue
    label(ws, f"B{rr}", nome)
    calc(ws, f"C{rr}", formula, fmt=fmt)
    put(ws, f"D{rr}", leit, f=font(9, c=MUTED), bg=GRAY, merge=f"D{rr}:F{rr}")
    ws.row_dimensions[rr].height = 21.75
    IR[nome] = rr
    rr += 1
s = rr + 1
band(ws, s, "Resumo automático", "F")
label(ws, f"B{s+1}", "Aviso")
C = lambda n: f"C{IR[n]}"  # noqa: E731
calc(ws, f"C{s+1}", f'=IF({C("Conhecimentos mapeados")}+{C("Produtos com política")}=0,"Preencha as abas Conhecimento e Pós-entrega",IF(C3="","Informe a data da leitura na aba Conhecimento",'
     f'IF(COUNTIF({Ao},"{ATRASADO}")>0,"Há atendimento aberto e atrasado",IF({C("Prazo aceito menor que o mínimo")}>0,"Há prazo aceito menor que o mínimo da lei ou do contrato",'
     f'IF({C("Expostos sem ação de transferência")}>0,"Há conhecimento exposto sem ação de transferência",IF({C(f"Pendentes há mais de {PRAZO_LICAO} dias")}>0,"Há lição pendente há mais de {PRAZO_LICAO} dias",'
     f'IF({C("Ações de transferência vencidas")}>0,"Há ação de transferência vencida",IF({C("Linhas a completar")}>0,"Há linha a completar: veja as conferências","OK"))))))))',
     merge=f"C{s+1}:F{s+1}", sz=9, b=False, h="left")
cf_warn(ws, f"C{s+1}:F{s+1}", f"C{s+1}")
ws.row_dimensions[s + 1].height = 21.75
NAV = s + 1
GS = s + 4
put(ws, f"B{GS - 1}", "Conhecimentos por situação", f=font(9, True, c=MUTED), bg=None, box=False)
for k, sit in enumerate(CON_SITS):
    put(ws, f"B{GS + k}", sit, f=font(9), bg=GRAY)
    calc(ws, f"C{GS + k}", f'=COUNTIF({Kn},B{GS + k})', b=False)
colunas(ws, f"D{GS - 1}", GS, GS + len(CON_SITS) - 1, 2, 3, width=12, height=7)
setup(ws, REDC, f"B1:F{GS + 14}", landscape=False, fit_height=True)

# ------------------------------------------------------------------ Checklist
ws = wb.create_sheet("Checklist")
widths(ws, {"A": 2, "B": 5, "C": 70, "D": 14, "E": 44, "F": 2})
title(ws, "Checklist do conhecimento e do pós-entrega", "Escolha o status de cada item na lista suspensa (coluna em amarelo-claro).", "E")
for col, text in zip("BCDE", ["#", "Item de verificação", "Status", "Observações"]):
    head(ws, f"{col}4", text)
ws.row_dimensions[4].height = 21.75
for k, text in enumerate(CHECK):
    rr = 5 + k
    num(ws, f"B{rr}", k + 1)
    put(ws, f"C{rr}", text, bg=GRAY)
    inp(ws, f"D{rr}", h="center")
    inp(ws, f"E{rr}")
    ws.row_dimensions[rr].height = 30
dv_list(ws, "D5:D16", YESNO, "Sim, Parcial ou Não")
cf_equal(ws, "D5:D16", YESNO_CF)
band(ws, 18, "Resumo automático", "E")
summary(ws, 19, "Itens atendidos (Sim)", '=COUNTIF(D5:D16,"Sim")', "C", "D")
summary(ws, 20, "Itens parciais", '=COUNTIF(D5:D16,"Parcial")', "C", "D")
summary(ws, 21, "Itens não atendidos", '=COUNTIF(D5:D16,"Não")', "C", "D")
summary(ws, 22, "Itens sem resposta", "=COUNTBLANK(D5:D16)", "C", "D")
summary(ws, 23, "Percentual atendido", "=D19/ROWS(D5:D16)", "C", "D", fmt="0%")
summary(ws, 24, "Situação", '=IF(D19=ROWS(D5:D16),"Validado",IF(D22>0,"Em andamento","Com pendências"))', "C", "D")
setup(ws, PURPLE, "B1:E24", landscape=False, fit_height=True)


# ------------------------------------------------------------------ Exemplos
def sub(ws, rr, cells, height=30):
    for a, b, t in cells:
        put(ws, f"{a}{rr}", t, f=font(10, True), bg=GRAY, h="center", merge=b and f"{a}{rr}:{b}{rr}")
    ws.row_dimensions[rr].height = height


def exemplo(ws, ex_):
    H = ex_["head"]
    widths(ws, {"A": 2, "B": 9, "C": 28, "D": 12, "E": 13, "F": 13, "G": 13, "H": 22, "I": 11, "J": 11, "K": 30, "L": 15, "M": 9, "N": 13, "O": 13, "P": 14, "Q": 28, "R": 2})
    title(ws, "Conhecimento organizacional e pós-entrega", "Exemplo preenchido, para consulta. Use as abas de entrada para a sua organização.", "Q")
    rr = 4
    for rot, val in [("Organização", H["org"]), ("Responsável", H["resp"]), ("Mudança prevista", H["mudanca"]), ("Origem", H["origem"])]:
        label(ws, f"B{rr}", rot, merge=f"B{rr}:C{rr}")
        inp(ws, f"D{rr}", val, merge=f"D{rr}:Q{rr}", bg=WHITE, h="left")
        ws.row_dimensions[rr].height = alt([(val, 170)], minimo=19.5)
        rr += 1
    label(ws, f"B{rr}", "Data da leitura", merge=f"B{rr}:C{rr}")
    inp(ws, f"D{rr}", H["ref"], bg=WHITE, h="center", fmt=DATE)
    ref = f"$D${rr}"
    # as quatro tabelas ficam uma abaixo da outra; as linhas são conhecidas de antemão, para as referências cruzadas
    k1 = rr + 4
    k2 = k1 + len(ex_["cons"]) - 1
    l1 = k2 + 4
    l2 = l1 + len(ex_["lics"]) - 1
    p1 = l2 + 4
    p2 = p1 + len(ex_["pols"]) - 1
    a1 = p2 + 4
    a2 = a1 + len(ex_["ats"]) - 1
    codes, lnums, prods, resps = f"$B${k1}:$B${k2}", f"$B${l1}:$B${l2}", f"$B${p1}:$B${p2}", f"$L${p1}:$L${p2}"

    # ---- conhecimento
    rr += 2
    band(ws, rr, "Mapa do conhecimento", "Q", color=BLUE)
    rr += 1
    sub(ws, rr, [("B", None, "Código"), ("C", None, "Conhecimento"), ("D", None, "Processo"), ("E", None, "Criticidade"), ("F", None, "Forma"), ("G", "H", "Onde está"),
                 ("I", None, "Quem sabe"), ("J", None, "Saída prevista"), ("K", None, "Ação de transferência"), ("L", None, "Prazo"), ("M", None, "Pontos"),
                 ("N", None, "Situação"), ("O", "Q", "Conferência")])
    assert rr + 1 == k1
    for k in ex_["cons"]:
        rr += 1
        put(ws, f"B{rr}", k["cod"], f=font(10, True, c=MUTED), h="center")
        put(ws, f"C{rr}", k["nome"], f=font(10, True))
        put(ws, f"D{rr}", k["proc"])
        put(ws, f"E{rr}", k["crit"], h="center")
        put(ws, f"F{rr}", k["forma"], h="center", f=font(9))
        put(ws, f"G{rr}", k["onde"] or None, merge=f"G{rr}:H{rr}", f=font(9))
        put(ws, f"I{rr}", k["pessoas"], h="center")
        put(ws, f"J{rr}", k["saida"], h="center")
        put(ws, f"K{rr}", k["acao"] or None, f=font(9))
        put(ws, f"L{rr}", k["prazo"], h="center", fmt=DATE)
        calc(ws, f"M{rr}", "=" + f_pontos(f"F{rr}", f"I{rr}", f"J{rr}"))
        calc(ws, f"N{rr}", "=" + f_consit(f"E{rr}", f"M{rr}"), sz=9)
        cells = dict(cod=f"B{rr}", crit=f"E{rr}", forma=f"F{rr}", onde=f"G{rr}", pessoas=f"I{rr}", acao=f"K{rr}", prazo=f"L{rr}", sit=f"N{rr}")
        calc(ws, f"O{rr}", "=" + f_conconf(cells, ref), b=False, sz=9, merge=f"O{rr}:Q{rr}")
        ws.row_dimensions[rr].height = alt([(k["nome"], 30), (k["acao"], 34), (k["onde"], 36)], minimo=21.75)
    assert rr == k2
    cf_texto(ws, f"N{k1}:N{k2}", f"N{k1}", CON_CF)
    cf_warn(ws, f"O{k1}:Q{k2}", f"$O{k1}")

    # ---- lições
    rr += 2
    band(ws, rr, "Lições aprendidas", "Q", color=TEAL)
    rr += 1
    sub(ws, rr, [("B", None, "Nº"), ("C", None, "O que aconteceu"), ("D", None, "Data"), ("E", None, "Origem"), ("F", "H", "O que aprendemos"), ("I", None, "Conhec."),
                 ("J", None, "Incorporada"), ("K", None, "Onde incorporar"), ("L", None, "Responsável"), ("M", None, "Dias"), ("N", None, "Situação"), ("O", "Q", "Conferência")])
    assert rr + 1 == l1
    for l in ex_["lics"]:
        rr += 1
        put(ws, f"B{rr}", l["num"], f=font(10, True, c=MUTED), h="center")
        put(ws, f"C{rr}", l["oque"], f=font(9))
        put(ws, f"D{rr}", l["data"], h="center", fmt=DATE)
        put(ws, f"E{rr}", l["origem"], h="center", f=font(9))
        put(ws, f"F{rr}", l["aprend"] or None, merge=f"F{rr}:H{rr}", f=font(9))
        put(ws, f"I{rr}", l["cod"] or None, h="center")
        put(ws, f"J{rr}", l["incorp"], h="center", fmt=DATE)
        put(ws, f"K{rr}", l["onde"] or None, f=font(9))
        put(ws, f"L{rr}", l["resp"] or None, f=font(9))
        calc(ws, f"N{rr}", "=" + f_licsit(f"B{rr}", f"J{rr}"), sz=9)
        calc(ws, f"M{rr}", "=" + f_licdias(f"B{rr}", f"D{rr}", f"J{rr}", ref), b=False)
        cells = dict(num=f"B{rr}", data=f"D{rr}", oque=f"C{rr}", aprend=f"F{rr}", onde=f"K{rr}", cod=f"I{rr}", sit=f"N{rr}", dias=f"M{rr}")
        calc(ws, f"O{rr}", "=" + f_licconf(cells, codes), b=False, sz=9, merge=f"O{rr}:Q{rr}")
        ws.row_dimensions[rr].height = alt([(l["oque"], 30), (l["aprend"], 46), (l["onde"], 32), (l["resp"], 12)], minimo=21.75)
    assert rr == l2
    cf_texto(ws, f"N{l1}:N{l2}", f"N{l1}", LIC_CF)
    cf_texto(ws, f"O{l1}:Q{l2}", f"$O{l1}", [("OK", GREEN), (VELHA, RED)], resto=YELLOW)

    # ---- política
    rr += 2
    ws.row_breaks.append(Break(id=rr - 1))
    band(ws, rr, "Política de pós-entrega", "Q", color=AMBER)
    rr += 1
    sub(ws, rr, [("B", "C", "Produto ou serviço"), ("D", "E", "Vida útil"), ("F", None, "Mínimo (dias)"), ("G", None, "Aceito (dias)"), ("H", "K", "O que se faz"),
                 ("L", None, "Resposta (dias)"), ("M", "P", "O que se quer evitar"), ("Q", None, "Conferência")])
    assert rr + 1 == p1
    for p in ex_["pols"]:
        rr += 1
        put(ws, f"B{rr}", p["prod"], f=font(10, True), merge=f"B{rr}:C{rr}")
        put(ws, f"D{rr}", p["vida"], merge=f"D{rr}:E{rr}", f=font(9))
        put(ws, f"F{rr}", p["minimo"], h="center")
        put(ws, f"G{rr}", p["aceito"], h="center")
        put(ws, f"H{rr}", p["ativ"], merge=f"H{rr}:K{rr}", f=font(9))
        put(ws, f"L{rr}", p["resp"], h="center")
        put(ws, f"M{rr}", p["conseq"] or None, merge=f"M{rr}:P{rr}", f=font(9))
        calc(ws, f"Q{rr}", "=" + f_polconf(dict(prod=f"B{rr}", min=f"F{rr}", aceito=f"G{rr}", ativ=f"H{rr}", resp=f"L{rr}")), b=False, sz=9)
        ws.row_dimensions[rr].height = alt([(p["ativ"], 66), (p["vida"], 22), (p["conseq"], 40)], minimo=21.75)
    assert rr == p2
    cf_texto(ws, f"Q{p1}:Q{p2}", f"Q{p1}", [("OK", GREEN), (MENOR, RED)], resto=YELLOW)

    # ---- atendimentos
    rr += 2
    band(ws, rr, "Atendimentos depois da entrega", "Q", color=PURPLE)
    rr += 1
    sub(ws, rr, [("B", None, "Nº"), ("C", None, "Relato"), ("D", None, "Abertura"), ("E", None, "Tipo"), ("F", "G", "Produto"), ("H", None, "Cliente"), ("I", None, "Resolução"),
                 ("J", None, "Custo"), ("K", None, "Causa"), ("L", None, "Lição"), ("M", None, "Prazo"), ("N", None, "Dias"), ("O", None, "Situação"), ("P", "Q", "Conferência")])
    assert rr + 1 == a1
    for a in ex_["ats"]:
        rr += 1
        put(ws, f"B{rr}", a["num"], f=font(10, True, c=MUTED), h="center")
        put(ws, f"C{rr}", a["relato"], f=font(9))
        put(ws, f"D{rr}", a["abert"], h="center", fmt=DATE)
        put(ws, f"E{rr}", a["tipo"], h="center", f=font(9))
        put(ws, f"F{rr}", a["prod"], merge=f"F{rr}:G{rr}", f=font(9))
        put(ws, f"H{rr}", a["cliente"], f=font(9))
        put(ws, f"I{rr}", a["resol"], h="center", fmt=DATE)
        put(ws, f"J{rr}", a["custo"], h="center", fmt=MOEDA)
        put(ws, f"K{rr}", a["causa"] or None, f=font(9))
        put(ws, f"L{rr}", a["licao"] or None, h="center")
        calc(ws, f"M{rr}", "=" + f_prazo(f"F{rr}", prods, resps), b=False)
        calc(ws, f"N{rr}", "=" + f_atdias(f"D{rr}", f"I{rr}", ref))
        calc(ws, f"O{rr}", "=" + f_atsit(f"M{rr}", f"N{rr}", f"I{rr}"), sz=9)
        cells = dict(prod=f"F{rr}", abert=f"D{rr}", tipo=f"E{rr}", resol=f"I{rr}", causa=f"K{rr}", licao=f"L{rr}")
        calc(ws, f"P{rr}", "=" + f_atconf(cells, f'B{rr}=""', prods, lnums), b=False, sz=9, merge=f"P{rr}:Q{rr}")
        ws.row_dimensions[rr].height = alt([(a["relato"], 30), (a["causa"], 32), (a["prod"], 26)], minimo=21.75)
    assert rr == a2
    cf_texto(ws, f"O{a1}:O{a2}", f"O{a1}", AT_CF)
    cf_warn(ws, f"P{a1}:Q{a2}", f"$P{a1}")

    # ---- resumo
    rr += 2
    band(ws, rr, "Resumo automático", "Q")
    Os = f"O{a1}:O{a2}"
    for k, (text, formula) in enumerate([
        ("Conhecimento", conta_por(f"N{k1}:N{k2}", CON_SITS)),
        ("Lições", conta_por(f"N{l1}:N{l2}", (INCORP, PEND)) + f'&", "&COUNTIF(O{l1}:O{l2},"{VELHA}")&" há mais de {PRAZO_LICAO} dias"'),
        ("Atendimentos", conta_por(Os, (NOPRAZO, FORAPRAZO, ABERTO, ATRASADO)) + f'&". Custo de R$ "&TEXT(SUM(J{a1}:J{a2}),"0")'),
        ("Linhas a completar", f'=SUMPRODUCT((O{k1}:O{k2}<>"OK")*1)+SUMPRODUCT((O{l1}:O{l2}<>"OK")*1)+SUMPRODUCT((Q{p1}:Q{p2}<>"OK")*1)+SUMPRODUCT((P{a1}:P{a2}<>"OK")*1)'),
    ], 1):
        label(ws, f"B{rr+k}", text, merge=f"B{rr+k}:D{rr+k}", h="right")
        calc(ws, f"E{rr+k}", formula, merge=f"E{rr+k}:Q{rr+k}", h="left")
        ws.row_dimensions[rr + k].height = 21.75
    setup(ws, MUTED, f"B1:Q{rr + 4}")
    return dict(k=(k1, k2), l=(l1, l2), p=(p1, p2), a=(a1, a2), res=rr + 1)


POS = {}
POS["ex1"] = exemplo(wb.create_sheet("Exemplo 1 - Pizzaria"), EX1)
POS["ex2"] = exemplo(wb.create_sheet("Exemplo 2 - Indústria"), EX2)

wb.active = 1
wb.save(OUT)
print("ok", OUT, "| avisos: Conhecimento E%d, Lições E%d, Pós-entrega E%d, Atendimentos E%d, Painel C%d | linhas do painel %s | exemplos %s" % (KAV, LAV, PAV, AAV, NAV, IR, POS))
