# -*- coding: utf-8 -*-
"""Gera Fornecedores-modelo.xlsx no mesmo padrão visual das planilhas anteriores da série."""
import math
import os
import sys
from datetime import date

_here = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, _here)
# funções de formatação compartilhadas com o gerador do PDCA (bloco inicial do arquivo)
_src = open(os.path.join(_here, "build_pdca.py"), encoding="utf-8").read()
exec(_src[: _src.index("STATUS = [")])

from openpyxl.chart import BarChart, Series  # noqa: E402
from openpyxl.chart.label import DataLabelList  # noqa: E402
from forn_data import (ATENDE, CHECK, CLASSES, CRITERIOS, CRITICIDADES, CRITICO, EX1, EX2, NAO, NAOATENDE, NAOCRIT, PESOS, S_DOC, S_HOM, S_VENC, SIM, TIPOS,  # noqa: E402
                       VALIDADE)

BLUE, TEAL, AMBER, PURPLE, REDC = "2B5C8A", "1E7B73", "A96A12", "7A4A9A", "B0413E"
BLUE_T, AMBER_T, RED_T, PURPLE_T = "DCE8F3", "F8EBCB", "F5DEDC", "EADFF0"
CLASSE_CF = [("A", BLUE_T), ("B", AMBER_T), ("C", RED_T), ("D", PURPLE_T)]
YESNO = ["Sim", "Parcial", "Não"]
YESNO_CF = [("Sim", GREEN), ("Parcial", YELLOW), ("Não", RED)]
RESULTADOS = [ATENDE, NAOATENDE, "Pendente", "Não se aplica"]
TIPOS_OC = ["Qualidade", "Prazo", "Documentação", "Atendimento"]
SIT_OC = ["Aberta", "Tratada", "Encerrada"]
NCAD, NAVA, NOC = 30, 30, 25
C1, C2 = 9, 9 + NCAD - 1     # cadastro
A1, A2 = 7, 7 + NAVA - 1     # avaliação
O1, O2 = 7, 7 + NOC - 1      # ocorrências
CAD = "Cadastro"
REF = f'IF({CAD}!$D$5="",TODAY(),{CAD}!$D$5)'
PESO = {n: p for n, p, _ in PESOS}


def note(ws, ref, text):
    c = Comment(text, "Modelo Fornecedores")
    c.width, c.height = 300, 120
    ws[ref].comment = c


def cf_warn(ws, rng_, first, ok_values=("OK",)):
    cond = ",".join(f'{first}="{v}"' for v in ok_values)
    ws.conditional_formatting.add(rng_, FormulaRule(formula=[f"OR({cond})"], fill=PatternFill("solid", bgColor=GREEN, fgColor=GREEN)))
    ws.conditional_formatting.add(rng_, FormulaRule(formula=[f"AND(LEN({first})>0,NOT(OR({cond})))"],
                                  fill=PatternFill("solid", bgColor=YELLOW, fgColor=YELLOW)))


def cf_texto(ws, rng_, first, pares, resto=None):
    for t, cor in pares:
        ws.conditional_formatting.add(rng_, FormulaRule(formula=[f'{first}="{t}"'], stopIfTrue=True, fill=PatternFill("solid", bgColor=cor, fgColor=cor)))
    if resto:
        ws.conditional_formatting.add(rng_, FormulaRule(formula=[f"LEN({first})>0"], fill=PatternFill("solid", bgColor=resto, fgColor=resto)))


def hint_row(ws, row, cells, height=30):
    for ref, text in cells:
        put(ws, f"{ref}{row}", text, f=font(9, i=True, c=MUTED), bg=GRAY, h="center")
    ws.row_dimensions[row].height = height


def alt(pares, minimo=21.75, linha=12):
    n = max(max(1, math.ceil(len(str(t)) / max(c, 1))) for t, c in pares)
    return max(minimo, n * linha + 7)


def dv_numero(ws, rng_, maximo, prompt, inteiro=True):
    dv = DataValidation(type="whole" if inteiro else "decimal", operator="between", formula1="0", formula2=str(maximo), allow_blank=True)
    dv.promptTitle, dv.prompt, dv.showInputMessage = "Número", prompt, True
    dv.errorTitle, dv.error, dv.showErrorMessage = "Número inválido", f"Digite um número de 0 a {maximo}.", True
    ws.add_data_validation(dv)
    dv.add(rng_)


def chart_style(ch):
    ch.style = 2
    ch.title = None
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ch.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill="D3DBE0", w=9525))
    ch.x_axis.graphicalProperties = GraphicalProperties(ln=LineProperties(solidFill="D3DBE0", w=9525))
    ch.y_axis.graphicalProperties = GraphicalProperties(ln=LineProperties(noFill=True))


def barras(ws, anchor, r1, r2, c_cat, c_val, width=16, height=6.5):
    """Barras horizontais: índice de cada fornecedor, de 0 a 100."""
    ch = BarChart()
    ch.type = "bar"
    ch.gapWidth = 50
    ch.height, ch.width = height, width
    chart_style(ch)
    ch.legend = None
    s = Series(Reference(ws, min_col=c_val, min_row=r1, max_row=r2), title="Índice")
    s.graphicalProperties.solidFill = BLUE
    s.graphicalProperties.line.noFill = True
    s.dLbls = DataLabelList()
    s.dLbls.showVal = True
    s.dLbls.numFmt = "0.0"
    s.dLbls.showSerName = s.dLbls.showCatName = s.dLbls.showLegendKey = False
    ch.append(s)
    ch.set_categories(Reference(ws, min_col=c_cat, min_row=r1, max_row=r2))
    ch.y_axis.scaling.min = 0
    ch.y_axis.scaling.max = 100
    ch.y_axis.majorUnit = 25
    ch.y_axis.number_format = "0"
    ch.x_axis.scaling.orientation = "maxMin"
    ws.add_chart(ch, anchor)


def f_indice(r, rec="F", ace="G", ent="H", pra="I", nota="J", cod="C"):
    return (f'=IF(OR({cod}{r}="",{rec}{r}="",{ent}{r}="",{nota}{r}=""),"",IF(OR({ace}{r}>{rec}{r},{pra}{r}>{ent}{r},{rec}{r}=0,{ent}{r}=0),"",'
            f'ROUND({PESO["Qualidade"]}*N({ace}{r})/{rec}{r}+{PESO["Prazo"]}*N({pra}{r})/{ent}{r}+{PESO["Atendimento"]}*{nota}{r}/10,1)))')


def f_classe(r, idx="M"):
    return f'=IF({idx}{r}="","",IF({idx}{r}>={CLASSES[0][1]},"A",IF({idx}{r}>={CLASSES[1][1]},"B",IF({idx}{r}>={CLASSES[2][1]},"C","D"))))'


def f_conduta(r, cls="N"):
    inner = '""'
    for letra, _, _, cond in reversed(CLASSES):
        inner = f'IF({cls}{r}="{letra}","{cond}",{inner})'
    return "=" + inner


def f_confere(r, rec="F", ace="G", ent="H", pra="I", nota="J", cod="C", idx="M"):
    return (f'=IF({cod}{r}="","",IF(OR({rec}{r}="",{ent}{r}="",{nota}{r}=""),"Faltam dados",IF(OR({ace}{r}>{rec}{r},{pra}{r}>{ent}{r}),"Confira os dados: aceitos ou no prazo maiores que o total",'
            f'IF(OR({rec}{r}=0,{ent}{r}=0),"Confira os dados: total zero","Calculada"))))')


# ================================================================== pasta de trabalho
wb = Workbook()
wb.properties.title = "Avaliação de fornecedores — Modelo"
wb.properties.creator = "Modelo Fornecedores"
wb.properties.language = "pt-BR"

# ------------------------------------------------------------------ Instruções
ws = wb.active
ws.title = "Instruções"
widths(ws, {"A": 2, "B": 24, "C": 92, "D": 2})
title(ws, "Avaliação de fornecedores — Como usar esta planilha",
      "Modelo para cadastrar e homologar fornecedores, avaliar o desempenho do período e registrar as ocorrências.", "C")
r = 4


def section(text):
    global r
    band(ws, r, text, "C", sz=11)
    r += 1


def line(k, v, kbg=GRAY, vbg=None, kf=None, kh="left", height=None):
    global r
    put(ws, f"B{r}", k, f=kf or font(10, True), bg=kbg, h=kh)
    put(ws, f"C{r}", v, bg=vbg)
    ws.row_dimensions[r].height = height or (19.5 if len(v) <= 98 else (31.5 if len(v) <= 196 else 45.75))
    r += 1


section("Legenda: onde preencher")
line("Amarelo-claro", "Células de entrada. É aqui que você digita.", vbg=INPUT)
line("Cinza", "Células calculadas ou fixas (rótulos, percentuais, índice, classe, conduta, avisos). Não altere.", vbg=GRAY)
line("Cinza em itálico", "Linha de exemplo (marcada com “Ex.”). Mostra o formato esperado e não entra nas contagens.", vbg=GRAY)
line("Listas suspensas", "Colunas de tipo, criticidade, documentos, resultado e situação aceitam apenas as opções da lista.")
line("Comentários", "Células com triângulo vermelho no canto trazem uma dica de preenchimento. Passe o mouse sobre elas.")
r += 1
section("O índice de desempenho e as classes")
for nome, peso, como in PESOS:
    line(f"{nome} · peso {peso}", como)
FAIXA = {"A": f"Índice de {CLASSES[0][1]} ou mais.", "B": f"Índice de {CLASSES[1][1]} a {CLASSES[0][1] - 1}.", "C": f"Índice de {CLASSES[2][1]} a {CLASSES[1][1] - 1}.",
         "D": f"Índice abaixo de {CLASSES[2][1]}."}
for letra, minimo, nome, cond in CLASSES:
    line(f"Classe {letra} · {nome}", f"{FAIXA[letra]} {cond}", kbg=[BLUE_T, AMBER_T, RED_T, PURPLE_T]["ABCD".index(letra)])
r += 1
section("Ordem recomendada de preenchimento")
for k, text in enumerate([
    "Aba Cadastro: escreva a organização, a data de referência e a validade padrão da homologação. Cadastre os fornecedores, com o que fornecem, o tipo e a criticidade.",
    "Aba Cadastro: para cada fornecedor crítico, registre a data da homologação e se os documentos estão em dia.",
    "Aba Homologação: preencha a ficha de um fornecedor novo, critério por critério, e leve a data para o Cadastro quando ele for homologado.",
    "Aba Avaliação: no fim do período, lance os dados do recebimento de cada fornecedor crítico e a nota de atendimento.",
    "Aba Avaliação: leia o índice, a classe e a conduta, e escreva o plano de ação ou a decisão.",
    "Aba Ocorrências: registre as falhas do período, com o tratamento e a situação.",
    "Aba Checklist: valide o controle de fornecedores.",
], 1):
    line(f"Passo {k}", text)
r += 1
section("Abas da planilha")
for k, text in [
    ("Cadastro", f"Até {NCAD} fornecedores, com a validade da homologação e a situação de cada um calculadas."),
    ("Homologação", "A ficha de um fornecedor novo: seis critérios, resultado, evidência e decisão calculada."),
    ("Avaliação", f"Até {NAVA} avaliações, com percentuais, índice, classe, conduta e gráfico."),
    ("Ocorrências", f"Até {NOC} ocorrências com fornecedores, com a conferência de cada uma."),
    ("Checklist", "Doze verificações de qualidade do controle de fornecedores, com percentual de conclusão."),
    ("Exemplo 1 - Indústria", "A avaliação semestral dos doze fornecedores críticos de uma indústria."),
    ("Exemplo 2 - Pizzaria", "A primeira avaliação dos fornecedores de uma loja."),
]:
    line(k, text)
r += 1
section("Regras e premissas do modelo")
for k, text in [
    ("Índice", "Índice = 50 × qualidade + 30 × prazo + 20 × atendimento, com os três em percentual. Os pesos e as classes são uma convenção deste modelo."),
    ("Dados do índice", "Lotes aceitos não podem passar dos recebidos, nem entregas no prazo das entregas. Quando isso acontece, o índice fica em branco e a conferência avisa."),
    ("Homologação", f"Vale {VALIDADE} meses, neste modelo. A situação “Homologação vencida” aparece quando a validade passa da data de referência."),
    ("Não crítico", "Fornecedor não crítico fica no cadastro, sem homologação nem avaliação obrigatórias."),
    ("Data de referência", "A situação da homologação é calculada pela data de referência da aba Cadastro. Se ela estiver em branco, a planilha usa a data de hoje."),
    ("Exemplos", "Os dados das abas de exemplo são ilustrativos, criados para este material."),
]:
    line(k, text)
setup(ws, INK, f"B1:C{r}", landscape=False, fit_height=True)

# ------------------------------------------------------------------ Cadastro
ws = wb.create_sheet(CAD)
widths(ws, {"A": 2, "B": 5, "C": 10, "D": 28, "E": 30, "F": 20, "G": 13, "H": 14, "I": 11, "J": 13, "K": 13, "L": 26, "M": 2})
title(ws, "Cadastro de fornecedores", "Um fornecedor por linha. Para os críticos, registre a homologação. A validade e a situação são calculadas.", "L")
for rr, rot, fmt, val in [(4, "Organização", None, None), (5, "Data de referência", DATE, None), (6, "Validade padrão da homologação, em meses", None, VALIDADE)]:
    label(ws, f"B{rr}", rot, merge=f"B{rr}:D{rr}")
    inp(ws, f"E{rr}", val, fmt=fmt, h="left")
    ws.row_dimensions[rr].height = 21.75
dv_date(ws, "E5")
put(ws, "F5", "Em branco, vale a data de hoje.", f=font(9, i=True, c=MUTED), bg=GRAY, merge="F5:H5")
put(ws, "F6", "Usada nas linhas sem validade própria.", f=font(9, i=True, c=MUTED), bg=GRAY, merge="F6:H6")
REF = f'IF({CAD}!$E$5="",TODAY(),{CAD}!$E$5)'
for col, text in zip("BCDEFGHIJKL", ["#", "Código", "Fornecedor", "O que fornece", "Tipo", "Criticidade", "Homologado em", "Validade (meses)", "Válida até",
                                      "Documentos em dia?", "Situação"]):
    head(ws, f"{col}8", text)
ws.row_dimensions[8].height = 33
hint_row(ws, 7, [("B", ""), ("C", "Ex.: F-01"), ("D", "Razão social ou nome"), ("E", "Item ou serviço"), ("F", "Escolha na lista"), ("G", "Escolha na lista"),
                 ("H", "Data da decisão"), ("I", "Em branco: padrão"), ("J", "Calculada"), ("K", "Sim ou Não"), ("L", "Calculada")], height=24)
for k in range(NCAD):
    rr = C1 + k
    num(ws, f"B{rr}", k + 1)
    inp(ws, f"C{rr}", fmt="@")
    inp(ws, f"D{rr}")
    inp(ws, f"E{rr}")
    inp(ws, f"F{rr}", h="center")
    inp(ws, f"G{rr}", h="center")
    inp(ws, f"H{rr}", h="center", fmt=DATE)
    inp(ws, f"I{rr}", h="center")
    calc(ws, f"J{rr}", f'=IF(OR(H{rr}="",D{rr}=""),"",EDATE(H{rr},IF(I{rr}="",$E$6,I{rr})))', fmt=DATE, b=False)
    inp(ws, f"K{rr}", h="center")
    calc(ws, f"L{rr}", f'=IF(D{rr}="","",IF(G{rr}="","Falta a criticidade",IF(G{rr}="{NAOCRIT}","Não crítico: sem homologação",IF(H{rr}="","{S_DOC}",'
         f'IF(K{rr}="Não","Documentos pendentes",IF(J{rr}<{REF},"{S_VENC}","{S_HOM}"))))))', b=False, sz=9)
    ws.row_dimensions[rr].height = 21.75
dv_list(ws, f"F{C1}:F{C2}", TIPOS, "Produto, Serviço ou Processo terceirizado")
dv_list(ws, f"G{C1}:G{C2}", CRITICIDADES, "Crítico ou Não crítico")
dv_date(ws, f"H{C1}:H{C2}")
dv_numero(ws, f"I{C1}:I{C2}", 120, "Meses de validade da homologação.")
dv_list(ws, f"K{C1}:K{C2}", [SIM, NAO], "Sim ou Não")
cf_texto(ws, f"L{C1}:L{C2}", f"L{C1}", [(S_HOM, GREEN), ("Não crítico: sem homologação", GRAY), (S_VENC, YELLOW), ("Documentos pendentes", YELLOW)], resto=RED)
note(ws, "G8", "Crítico: o fornecimento afeta o produto ou o cliente. Fornecedor único de item crítico é também um risco: leve à matriz de riscos.")
note(ws, "K8", "Certidões, licenças e contratos dentro da validade.")
s = C2 + 2
band(ws, s, "Resumo automático", "L")
for k, (text, formula) in enumerate([
    ("Fornecedores cadastrados", f"=COUNTA(D{C1}:D{C2})"),
    ("Críticos", f'=COUNTIF(G{C1}:G{C2},"{CRITICO}")'),
    ("Críticos homologados, em dia", f'=COUNTIF(L{C1}:L{C2},"{S_HOM}")'),
    ("Com homologação vencida ou documentos pendentes", f'=COUNTIF(L{C1}:L{C2},"{S_VENC}")+COUNTIF(L{C1}:L{C2},"Documentos pendentes")'),
    ("Críticos sem homologação", f'=COUNTIF(L{C1}:L{C2},"{S_DOC}")'),
    ("Aviso", f'=IF(E{s+1}=0,"Cadastre os fornecedores",IF(COUNTIF(L{C1}:L{C2},"Falta a criticidade")>0,"Há fornecedor sem criticidade",'
              f'IF(E{s+5}>0,"Há fornecedor crítico sem homologação",IF(E{s+4}>0,"Há homologação vencida ou documentos pendentes","OK"))))'),
], 1):
    label(ws, f"B{s+k}", text, merge=f"B{s+k}:D{s+k}", h="right")
    calc(ws, f"E{s+k}", formula, merge=f"E{s+k}:F{s+k}", sz=9 if text == "Aviso" else 10, b=text != "Aviso")
    ws.row_dimensions[s + k].height = 21.75
cf_warn(ws, f"E{s+6}:F{s+6}", f"E{s+6}")
ws.freeze_panes = "E9"
setup(ws, BLUE, f"B1:L{s+6}", fit_height=True)

# ------------------------------------------------------------------ Homologação
ws = wb.create_sheet("Homologação")
widths(ws, {"A": 2, "B": 24, "C": 40, "D": 16, "E": 50, "F": 2})
title(ws, "Ficha de homologação", "Um fornecedor por ficha. Para o seguinte, duplique a aba ou salve uma cópia do arquivo.", "E")
for rr, rot in [(4, "Código e fornecedor"), (5, "O que fornece"), (6, "Data da avaliação"), (7, "Quem avaliou")]:
    label(ws, f"B{rr}", rot)
    inp(ws, f"C{rr}", merge=f"C{rr}:E{rr}", fmt=DATE if rr == 6 else None, h="left")
    ws.row_dimensions[rr].height = 21.75
dv_date(ws, "C6")
for col, text in zip("BCDE", ["Critério", "O que se confere", "Resultado", "Evidência"]):
    head(ws, f"{col}9", text)
ws.row_dimensions[9].height = 21.75
H1 = 10
for k, (nome, oque, ev) in enumerate(CRITERIOS):
    rr = H1 + k
    put(ws, f"B{rr}", nome, f=font(10, True), bg=GRAY)
    put(ws, f"C{rr}", oque, f=font(9, c=MUTED), bg=GRAY)
    inp(ws, f"D{rr}", h="center")
    inp(ws, f"E{rr}")
    note(ws, f"B{rr}", f"Evidência típica: {ev}")
    ws.row_dimensions[rr].height = 33
H2 = H1 + len(CRITERIOS) - 1
dv_list(ws, f"D{H1}:D{H2}", RESULTADOS, "Atende, Não atende, Pendente ou Não se aplica")
cf_texto(ws, f"D{H1}:D{H2}", f"D{H1}", [(ATENDE, GREEN), (NAOATENDE, RED), ("Pendente", YELLOW), ("Não se aplica", GRAY)])
s = H2 + 2
band(ws, s, "Decisão", "E")
label(ws, f"B{s+1}", "Critérios atendidos")
calc(ws, f"C{s+1}", f'=COUNTIF(D{H1}:D{H2},"{ATENDE}")&" de "&(ROWS(D{H1}:D{H2})-COUNTIF(D{H1}:D{H2},"Não se aplica"))', h="left", b=False)
label(ws, f"B{s+2}", "Resultado")
calc(ws, f"C{s+2}", f'=IF(C4="","Informe o fornecedor",IF(COUNTBLANK(D{H1}:D{H2})>0,"Preencha os critérios",IF(COUNTIF(D{H1}:D{H2},"{NAOATENDE}")>0,"Não homologado",'
     f'IF(COUNTIF(D{H1}:D{H2},"Pendente")>0,"Em homologação: há critério pendente",IF(COUNTIF(E{H1}:E{H2},"")>COUNTIF(D{H1}:D{H2},"Não se aplica"),"Homologado, mas falta evidência",'
     f'"Homologado por "&{CAD}!$E$6&" meses")))))', h="left", sz=9)
label(ws, f"B{s+3}", "Válida até")
calc(ws, f"C{s+3}", f'=IF(LEFT(C{s+2},14)="Homologado por",EDATE(C6,{CAD}!$E$6),"")', fmt=DATE, h="left", b=False)
label(ws, f"B{s+4}", "Aprovado por")
inp(ws, f"C{s+4}", merge=f"C{s+4}:E{s+4}")
for k in range(1, 5):
    ws.row_dimensions[s + k].height = 21.75
    if k < 4:
        ws.merge_cells(f"C{s+k}:E{s+k}")
ws.conditional_formatting.add(f"C{s+2}", FormulaRule(formula=[f'LEFT(C{s+2},14)="Homologado por"'], stopIfTrue=True, fill=PatternFill("solid", bgColor=GREEN, fgColor=GREEN)))
ws.conditional_formatting.add(f"C{s+2}", FormulaRule(formula=[f'C{s+2}="Não homologado"'], stopIfTrue=True, fill=PatternFill("solid", bgColor=RED, fgColor=RED)))
ws.conditional_formatting.add(f"C{s+2}", FormulaRule(formula=[f"LEN(C{s+2})>0"], fill=PatternFill("solid", bgColor=YELLOW, fgColor=YELLOW)))
note(ws, "C4", "Depois de homologado, registre a data no Cadastro: é lá que a validade é acompanhada.")
setup(ws, TEAL, f"B1:E{s+4}", landscape=False, fit_height=True)

# ------------------------------------------------------------------ Avaliação
ws = wb.create_sheet("Avaliação")
widths(ws, {"A": 2, "B": 5, "C": 10, "D": 24, "E": 16, "F": 10, "G": 10, "H": 10, "I": 10, "J": 8, "K": 10, "L": 9, "M": 9, "N": 8, "O": 30, "P": 34, "Q": 24, "R": 2})
title(ws, "Avaliação de desempenho", "Uma linha por fornecedor e período. Lance os dados do recebimento e a nota de atendimento: o índice, a classe e a conduta são calculados.", "Q")
label(ws, "B4", "Período avaliado", merge="B4:C4")
inp(ws, "D4", merge="D4:F4")
put(ws, "G4", "Ex.: julho a dezembro de 2026", f=font(9, i=True, c=MUTED), bg=GRAY, merge="G4:I4")
ws.row_dimensions[4].height = 21.75
for col, text in zip("BCDEFGHIJKLMNOPQ", ["#", "Código", "Fornecedor", "Período", "Lotes recebidos", "Lotes aceitos", "Entregas", "Entregas no prazo", "Nota (0 a 10)",
                                          "Qualidade", "Prazo", "Índice", "Classe", "Conduta", "Plano de ação ou decisão", "Conferência"]):
    head(ws, f"{col}5", text)
ws.row_dimensions[5].height = 33
hint_row(ws, 6, [("B", ""), ("C", "Do Cadastro"), ("D", "Do Cadastro"), ("E", "Semestre, ano"), ("F", "Ou serviços prestados"), ("G", "Sem recusa"), ("H", "No período"),
                 ("I", "Na data combinada"), ("J", "Do comprador"), ("K", "Calculada"), ("L", "Calculada"), ("M", "Calculado"), ("N", "Calculada"), ("O", "Calculada"),
                 ("P", "O que foi combinado com o fornecedor"), ("Q", "Calculada")])
for k in range(NAVA):
    rr = A1 + k
    num(ws, f"B{rr}", k + 1)
    inp(ws, f"C{rr}", fmt="@")
    calc(ws, f"D{rr}", f'=IF(C{rr}="","",IFERROR(VLOOKUP(C{rr},{CAD}!$C${C1}:$D${C2},2,FALSE),"Não está no cadastro"))', h="left", b=False, sz=9)
    inp(ws, f"E{rr}")
    for col in "FGHI":
        inp(ws, f"{col}{rr}", h="center")
    inp(ws, f"J{rr}", h="center")
    calc(ws, f"K{rr}", f'=IF(OR(F{rr}="",F{rr}=0,G{rr}>F{rr}),"",N(G{rr})/F{rr})', fmt="0%", b=False)
    calc(ws, f"L{rr}", f'=IF(OR(H{rr}="",H{rr}=0,I{rr}>H{rr}),"",N(I{rr})/H{rr})', fmt="0%", b=False)
    calc(ws, f"M{rr}", f_indice(rr), fmt="0.0")
    calc(ws, f"N{rr}", f_classe(rr))
    calc(ws, f"O{rr}", f_conduta(rr), h="left", b=False, sz=9)
    inp(ws, f"P{rr}")
    calc(ws, f"Q{rr}", f_confere(rr), b=False, sz=9)
    ws.row_dimensions[rr].height = 30
for col in "FGHI":
    dv_numero(ws, f"{col}{A1}:{col}{A2}", 100000, "Quantidade no período.")
dv_numero(ws, f"J{A1}:J{A2}", 10, "Nota de atendimento, de 0 a 10.", inteiro=False)
cf_equal(ws, f"N{A1}:N{A2}", CLASSE_CF)
cf_texto(ws, f"Q{A1}:Q{A2}", f"Q{A1}", [("Calculada", GREEN)], resto=YELLOW)
ws.conditional_formatting.add(f"D{A1}:D{A2}", FormulaRule(formula=[f'D{A1}="Não está no cadastro"'], fill=PatternFill("solid", bgColor=YELLOW, fgColor=YELLOW)))
note(ws, "F5", "Para serviço, conte os serviços prestados no período. Para transporte, cada carga.")
note(ws, "J5", "Resposta, documentação, flexibilidade, comunicação. Uma nota por período, dada pelo comprador.")
note(ws, "P5", "Classe C: plano de ação com prazo. Classe D: suspensão ou desqualificação, com bloqueio no sistema.")
s = A2 + 2
band(ws, s, "Resumo automático", "I")
put(ws, f"K{s}", "Índice por fornecedor", f=font(10, True, c=WHITE), bg=INK, box=False, merge=f"K{s}:Q{s}")
for k, (text, formula, fmt) in enumerate([
    ("Fornecedores avaliados", f"=COUNT(M{A1}:M{A2})", None),
    ("Classe A, B, C e D", f'=COUNTIF(N{A1}:N{A2},"A")&", "&COUNTIF(N{A1}:N{A2},"B")&", "&COUNTIF(N{A1}:N{A2},"C")&" e "&COUNTIF(N{A1}:N{A2},"D")', None),
    ("Média do índice", f'=IF(E{s+1}=0,"",AVERAGE(M{A1}:M{A2}))', "0.0"),
    ("Classe C ou D sem plano ou decisão", f'=COUNTIFS(N{A1}:N{A2},"C",P{A1}:P{A2},"")+COUNTIFS(N{A1}:N{A2},"D",P{A1}:P{A2},"")', None),
    ("Aviso", f'=IF(COUNTA(C{A1}:C{A2})=0,"Registre as avaliações",IF(COUNTIF(D{A1}:D{A2},"Não está no cadastro")>0,"Há código que não está no cadastro",'
              f'IF(COUNTIF(Q{A1}:Q{A2},"Faltam dados")+COUNTIF(Q{A1}:Q{A2},"Confira*")>0,"Há avaliação com dados a conferir: veja a coluna Conferência",'
              f'IF(E{s+4}>0,"Há fornecedor de classe C ou D sem plano ou decisão","OK"))))', None),
], 1):
    label(ws, f"B{s+k}", text, merge=f"B{s+k}:D{s+k}", h="right")
    calc(ws, f"E{s+k}", formula, fmt=fmt, merge=f"E{s+k}:I{s+k}", sz=9 if text == "Aviso" else 10, b=text != "Aviso")
    ws.row_dimensions[s + k].height = 21.75
cf_warn(ws, f"E{s+5}:I{s+5}", f"E{s+5}")
barras(ws, f"K{s+1}", A1, A2, 3, 13, width=15, height=12)
ws.freeze_panes = "E7"
setup(ws, AMBER, f"B1:Q{s+24}", fit_height=True)

# ------------------------------------------------------------------ Ocorrências
ws = wb.create_sheet("Ocorrências")
widths(ws, {"A": 2, "B": 5, "C": 13, "D": 10, "E": 24, "F": 14, "G": 40, "H": 34, "I": 12, "J": 12, "K": 24, "L": 2})
title(ws, "Ocorrências com fornecedores", "Uma linha para cada falha registrada pelo recebimento: lote recusado, atraso, documento errado, atendimento.", "K")
for col, text in zip("BCDEFGHIJK", ["#", "Data", "Código", "Fornecedor", "Tipo", "Descrição", "Tratamento", "RNC nº", "Situação", "Conferência"]):
    head(ws, f"{col}4", text)
ws.row_dimensions[4].height = 33
hint_row(ws, 5, [("B", ""), ("C", "Da ocorrência"), ("D", "Do Cadastro"), ("E", "Do Cadastro"), ("F", "Escolha na lista"), ("G", "O que aconteceu, com o lote ou o pedido"),
                 ("H", "Devolução, reposição, plano do fornecedor"), ("I", "Se abriu RNC"), ("J", "Escolha na lista"), ("K", "Calculada")])
ex(ws, "B6", "Ex.", h="center")
for col, v, h_, fmt in [("C", date(2026, 10, 7), "center", DATE), ("D", "F-08", "center", "@"), ("E", "Paletes Rápido", "left", None), ("F", "Qualidade", "center", None),
                        ("G", "Lote de 40 paletes com madeira sem tratamento. Nota fiscal 1.223.", "left", None), ("H", "Lote devolvido. Reposição em 5 dias.", "left", None),
                        ("I", "2026-35", "center", "@"), ("J", "Encerrada", "center", None), ("K", "Completa", "center", None)]:
    ex(ws, f"{col}6", v, h=h_, fmt=fmt)
ws.row_dimensions[6].height = 30
for k in range(NOC):
    rr = O1 + k
    num(ws, f"B{rr}", k + 1)
    inp(ws, f"C{rr}", h="center", fmt=DATE)
    inp(ws, f"D{rr}", h="center", fmt="@")
    calc(ws, f"E{rr}", f'=IF(D{rr}="","",IFERROR(VLOOKUP(D{rr},{CAD}!$C${C1}:$D${C2},2,FALSE),"Não está no cadastro"))', h="left", b=False, sz=9)
    inp(ws, f"F{rr}", h="center")
    inp(ws, f"G{rr}")
    inp(ws, f"H{rr}")
    inp(ws, f"I{rr}", h="center", fmt="@")
    inp(ws, f"J{rr}", h="center")
    calc(ws, f"K{rr}", f'=IF(G{rr}="","",IF(C{rr}="","Falta a data",IF(D{rr}="","Falta o fornecedor",IF(F{rr}="","Falta o tipo",IF(J{rr}="","Falta a situação",'
         f'IF(AND(J{rr}<>"Aberta",H{rr}=""),"Falta o tratamento","Completa"))))))', b=False, sz=9)
    ws.row_dimensions[rr].height = 30
dv_date(ws, f"C{O1}:C{O2}")
dv_list(ws, f"F{O1}:F{O2}", TIPOS_OC, "Qualidade, Prazo, Documentação ou Atendimento")
dv_list(ws, f"J{O1}:J{O2}", SIT_OC, "Aberta, Tratada ou Encerrada")
cf_equal(ws, f"J{O1}:J{O2}", [("Aberta", RED), ("Tratada", YELLOW), ("Encerrada", GREEN)])
cf_texto(ws, f"K{O1}:K{O2}", f"K{O1}", [("Completa", GREEN)], resto=YELLOW)
note(ws, "I4", "Ocorrência grave, ou repetida, vira registro de não conformidade, com análise de causa.")
s = O2 + 2
band(ws, s, "Resumo automático", "K")
for k, (text, formula) in enumerate([
    ("Ocorrências registradas", f"=COUNTA(G{O1}:G{O2})"),
    ("Abertas", f'=COUNTIF(J{O1}:J{O2},"Aberta")'),
    ("Qualidade, prazo, documentação e atendimento", f'=COUNTIF(F{O1}:F{O2},"Qualidade")&", "&COUNTIF(F{O1}:F{O2},"Prazo")&", "&COUNTIF(F{O1}:F{O2},"Documentação")&" e "&COUNTIF(F{O1}:F{O2},"Atendimento")'),
    ("Aviso", f'=IF(E{s+1}=0,"Registre as ocorrências",IF(COUNTIF(K{O1}:K{O2},"Falta*")>0,"Há ocorrência com campos em branco: veja a coluna Conferência",'
              f'IF(E{s+2}>0,"Há ocorrência aberta","OK")))'),
], 1):
    label(ws, f"B{s+k}", text, merge=f"B{s+k}:D{s+k}", h="right")
    calc(ws, f"E{s+k}", formula, merge=f"E{s+k}:G{s+k}", sz=9 if text == "Aviso" else 10, b=text != "Aviso")
    ws.row_dimensions[s + k].height = 21.75
cf_warn(ws, f"E{s+4}:G{s+4}", f"E{s+4}")
ws.freeze_panes = "F7"
setup(ws, REDC, f"B1:K{s+4}", fit_height=True)

# ------------------------------------------------------------------ Checklist
ws = wb.create_sheet("Checklist")
widths(ws, {"A": 2, "B": 5, "C": 70, "D": 14, "E": 44, "F": 2})
title(ws, "Checklist do controle de fornecedores", "Escolha o status de cada item na lista suspensa (coluna em amarelo-claro).", "E")
for col, text in zip("BCDE", ["#", "Item de verificação", "Status", "Observações"]):
    head(ws, f"{col}4", text)
ws.row_dimensions[4].height = 21.75
for k, text in enumerate(CHECK):
    rr = 5 + k
    num(ws, f"B{rr}", k + 1)
    put(ws, f"C{rr}", text, bg=GRAY)
    inp(ws, f"D{rr}", h="center")
    inp(ws, f"E{rr}")
    ws.row_dimensions[rr].height = 30
dv_list(ws, "D5:D16", YESNO, "Sim, Parcial ou Não")
cf_equal(ws, "D5:D16", YESNO_CF)
band(ws, 18, "Resumo automático", "E")
summary(ws, 19, "Itens atendidos (Sim)", '=COUNTIF(D5:D16,"Sim")', "C", "D")
summary(ws, 20, "Itens parciais", '=COUNTIF(D5:D16,"Parcial")', "C", "D")
summary(ws, 21, "Itens não atendidos", '=COUNTIF(D5:D16,"Não")', "C", "D")
summary(ws, 22, "Itens sem resposta", "=COUNTBLANK(D5:D16)", "C", "D")
summary(ws, 23, "Percentual atendido", "=D19/ROWS(D5:D16)", "C", "D", fmt="0%")
summary(ws, 24, "Situação", '=IF(D19=ROWS(D5:D16),"Validado",IF(D22>0,"Em andamento","Com pendências"))', "C", "D")
setup(ws, PURPLE, "B1:E24", landscape=False, fit_height=True)


# ------------------------------------------------------------------ Exemplos
def sub(ws, rr, cells, height=30):
    for a, b, t in cells:
        put(ws, f"{a}{rr}", t, f=font(10, True), bg=GRAY, h="center", merge=b and f"{a}{rr}:{b}{rr}")
    ws.row_dimensions[rr].height = height


def exemplo(ws, data):
    widths(ws, {"A": 2, "B": 9, "C": 24, "D": 26, "E": 16, "F": 11, "G": 12, "H": 12, "I": 10, "J": 10, "K": 9, "L": 8, "M": 36, "N": 2})
    title(ws, "Avaliação de fornecedores", "Exemplo preenchido, para consulta. Use as abas Cadastro, Homologação, Avaliação e Ocorrências para os seus fornecedores.", "M")
    H = data["head"]
    rr = 4
    for rot, val, fmt in [("Organização", H["org"], None), ("Data de referência", H["data"], DATE), ("Período avaliado", H["periodo"], None), ("Quem avaliou", H["por"], None),
                          ("Origem", H["origem"], None)]:
        label(ws, f"B{rr}", rot, merge=f"B{rr}:C{rr}")
        inp(ws, f"D{rr}", val, merge=f"D{rr}:M{rr}", bg=WHITE, fmt=fmt, h="left")
        ws.row_dimensions[rr].height = 19.5
        rr += 1
    rr += 1
    band(ws, rr, "Cadastro e homologação", "M", color=BLUE)
    rr += 1
    sub(ws, rr, [("B", None, "Código"), ("C", None, "Fornecedor"), ("D", None, "O que fornece"), ("E", None, "Tipo"), ("F", None, "Criticidade"), ("G", None, "Homologado em"),
                 ("H", None, "Válida até"), ("I", None, "Docs. em dia?"), ("J", "M", "Situação")], height=24)
    c1 = rr + 1
    for f in data["forns"]:
        rr += 1
        put(ws, f"B{rr}", f["cod"], f=font(10, True), h="center", fmt="@")
        put(ws, f"C{rr}", f["nome"], f=font(10, True))
        put(ws, f"D{rr}", f["fornece"])
        put(ws, f"E{rr}", f["tipo"], h="center")
        put(ws, f"F{rr}", f["crit"], h="center")
        put(ws, f"G{rr}", f["homologado"], h="center", fmt=DATE)
        calc(ws, f"H{rr}", f'=IF(G{rr}="","",EDATE(G{rr},{VALIDADE}))', fmt=DATE, b=False)
        put(ws, f"I{rr}", f["docs"], h="center")
        calc(ws, f"J{rr}", f'=IF(F{rr}="{NAOCRIT}","Não crítico: sem homologação",IF(G{rr}="","{S_DOC}",IF(I{rr}="Não","Documentos pendentes",IF(H{rr}<$D$5,"{S_VENC}","{S_HOM}"))))',
             b=False, sz=9, merge=f"J{rr}:M{rr}")
        ws.row_dimensions[rr].height = alt([(f["nome"], 24), (f["fornece"], 26)], minimo=19.5)
    c2 = rr
    cf_texto(ws, f"J{c1}:J{c2}", f"J{c1}", [(S_HOM, GREEN), ("Não crítico: sem homologação", GRAY), (S_VENC, YELLOW), ("Documentos pendentes", YELLOW)], resto=RED)
    rr += 2
    ws.row_breaks.append(Break(id=rr - 1))
    band(ws, rr, f'Avaliação de desempenho · {H["periodo"].lower()}', "M", color=AMBER)
    rr += 1
    sub(ws, rr, [("B", None, "Código"), ("C", None, "Fornecedor"), ("D", None, "Lotes recebidos · aceitos"), ("E", None, "Entregas · no prazo"), ("F", None, "Nota"),
                 ("G", None, "Qualidade"), ("H", None, "Prazo"), ("I", None, "Índice"), ("J", None, "Classe"), ("K", "M", "Observação e conduta")], height=30)
    a1 = rr + 1
    nomes = {f["cod"]: f["nome"] for f in data["forns"]}
    for a in data["avals"]:
        rr += 1
        put(ws, f"B{rr}", a["cod"], f=font(10, True), h="center", fmt="@")
        put(ws, f"C{rr}", nomes[a["cod"]], f=font(10, True))
        put(ws, f"D{rr}", f'{a["recebidos"]} · {a["aceitos"]}', h="center")
        put(ws, f"E{rr}", f'{a["entregas"]} · {a["no_prazo"]}', h="center")
        put(ws, f"F{rr}", a["nota"], h="center")
        calc(ws, f"G{rr}", f"={a['aceitos']}/{a['recebidos']}", fmt="0%", b=False)
        calc(ws, f"H{rr}", f"={a['no_prazo']}/{a['entregas']}", fmt="0%", b=False)
        calc(ws, f"I{rr}", f"=ROUND({PESO['Qualidade']}*G{rr}+{PESO['Prazo']}*H{rr}+{PESO['Atendimento']}*F{rr}/10,1)", fmt="0.0")
        calc(ws, f"J{rr}", f_classe(rr, idx="I"))
        put(ws, f"K{rr}", (a["obs"] + " " if a["obs"] else ""), merge=f"K{rr}:M{rr}")
        ws[f"K{rr}"].value = (a["obs"] + " " if a["obs"] else "") + "Conduta: " + [c[3] for c in CLASSES if c[0] == a["classe"]][0]
        ws.row_dimensions[rr].height = alt([(ws[f"K{rr}"].value, 52), (nomes[a["cod"]], 24)], minimo=21.75)
    a2 = rr
    cf_equal(ws, f"J{a1}:J{a2}", CLASSE_CF)
    rr += 2
    band(ws, rr, "Resumo automático", "F")
    put(ws, f"H{rr}", "Índice por fornecedor", f=font(10, True, c=WHITE), bg=INK, box=False, merge=f"H{rr}:M{rr}")
    c0 = rr
    for k, (text, formula, fmt) in enumerate([
        ("Fornecedores avaliados", f"=COUNT(I{a1}:I{a2})", None),
        ("Classe A, B, C e D", f'=COUNTIF(J{a1}:J{a2},"A")&", "&COUNTIF(J{a1}:J{a2},"B")&", "&COUNTIF(J{a1}:J{a2},"C")&" e "&COUNTIF(J{a1}:J{a2},"D")', None),
        ("Média do índice", f"=AVERAGE(I{a1}:I{a2})", "0.0"),
        ("Menor índice", f"=INDEX(C{a1}:C{a2},MATCH(MIN(I{a1}:I{a2}),I{a1}:I{a2},0))", None),
        ("Homologações vencidas ou pendentes", f'=COUNTIF(J{c1}:J{c2},"{S_VENC}")+COUNTIF(J{c1}:J{c2},"Documentos pendentes")+COUNTIF(J{c1}:J{c2},"{S_DOC}")', None),
    ], 1):
        label(ws, f"B{rr+k}", text, merge=f"B{rr+k}:D{rr+k}", h="right")
        calc(ws, f"E{rr+k}", formula, fmt=fmt, merge=f"E{rr+k}:F{rr+k}")
        ws.row_dimensions[rr + k].height = 21.75
    barras(ws, f"H{c0+1}", a1, a2, 2, 9, width=17, height=7.5 if len(data["avals"]) > 8 else 5.5)
    setup(ws, MUTED, f"B1:M{c0 + 18}")


exemplo(wb.create_sheet("Exemplo 1 - Indústria"), EX1)
exemplo(wb.create_sheet("Exemplo 2 - Pizzaria"), EX2)

wb.active = 1
wb.save(OUT)
print("ok", OUT)
