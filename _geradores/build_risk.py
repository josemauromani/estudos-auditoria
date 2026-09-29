# -*- coding: utf-8 -*-
"""Gera Riscos-modelo.xlsx no mesmo padrão visual das planilhas anteriores da série."""
import math
import os
import sys
from datetime import date

_here = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, _here)
# funções de formatação compartilhadas com o gerador do PDCA (bloco inicial do arquivo)
_src = open(os.path.join(_here, "build_pdca.py"), encoding="utf-8").read()
exec(_src[: _src.index("STATUS = [")])

from openpyxl.chart import BarChart, Series  # noqa: E402
from openpyxl.chart.label import DataLabelList  # noqa: E402
from openpyxl.utils import get_column_letter as L  # noqa: E402
from risk_data import (ACEITAR, ACEITE, CHECK, CONDUTA, EX1, EX2, IMPACTO, NIVEIS, PROB, RESP_TXT, RESPOSTAS, nivel)  # noqa: E402

BLUE, TEAL, AMBER, PURPLE, REDC = "2B5C8A", "1E7B73", "A96A12", "7A4A9A", "B0413E"
NV_COR = {"Baixo": "2A6FB0", "Médio": "D19A2E", "Alto": "B0413E", "Crítico": "7A3E9A"}
NV_TINT = {"Baixo": "DCE8F3", "Médio": "F8EBCB", "Alto": "F5DEDC", "Crítico": "E9DEF1"}
NV_CF = [(n, NV_TINT[n]) for n in NIVEIS]
YESNO = ["Sim", "Parcial", "Não"]
YESNO_CF = [("Sim", GREEN), ("Parcial", YELLOW), ("Não", RED)]
ST_ACAO = ["Não iniciada", "Em andamento", "Concluída"]
ST_CF = [("Concluída", GREEN), ("Em andamento", YELLOW), ("Não iniciada", RED)]
SIT_CF = [("Concluída", GREEN), ("No prazo", GREEN), ("Aceito", "E3E8EB"), ("Sem prazo", YELLOW), ("Atrasada", RED)]
DECISOES = ["Aproveitar", "Testar", "Não agir agora"]
PRIO_CF = [("Alta", "DCE8F3"), ("Média", "E3E8EB"), ("Baixa", "EEF1F3")]
BENEFICIO = [(1, "Muito baixo", "Ganho pequeno, difícil de perceber."), (2, "Baixo", "Ganho interno, sem efeito para o cliente."),
             (3, "Médio", "Ganho percebido pelo cliente ou economia relevante."), (4, "Alto", "Novo cliente, nova linha ou redução grande de custo."),
             (5, "Muito alto", "Muda a posição da organização no mercado.")]


def note(ws, ref, text):
    c = Comment(text, "Modelo Riscos")
    c.width, c.height = 300, 120
    ws[ref].comment = c


def cf_warn(ws, rng_, first, ok_values=("OK",)):
    cond = ",".join(f'{first}="{v}"' for v in ok_values)
    ws.conditional_formatting.add(rng_, FormulaRule(formula=[f"OR({cond})"], fill=PatternFill("solid", bgColor=GREEN, fgColor=GREEN)))
    ws.conditional_formatting.add(rng_, FormulaRule(formula=[f"AND(LEN({first})>0,NOT(OR({cond})))"],
                                  fill=PatternFill("solid", bgColor=YELLOW, fgColor=YELLOW)))


def hint_row(ws, row, cells, height=30):
    for ref, text, merge in cells:
        put(ws, f"{ref}{row}", text, f=font(9, i=True, c=MUTED), bg=GRAY, h="center", merge=merge and f"{ref}{row}:{merge}{row}")
    ws.row_dimensions[row].height = height


def alt(text, chars, minimo=21.75, linha=12.75):
    n = max(1, math.ceil(len(str(text)) / max(chars, 1)))
    return max(minimo, n * linha + 9)


def dv_nota(ws, rng_):
    dv = DataValidation(type="whole", operator="between", formula1="1", formula2="5", allow_blank=True)
    dv.promptTitle, dv.prompt = "Nota", "Número inteiro de 1 a 5, conforme a aba Escalas."
    dv.errorTitle, dv.error = "Nota inválida", "Digite um número inteiro de 1 a 5."
    dv.showErrorMessage = dv.showInputMessage = True
    ws.add_data_validation(dv)
    dv.add(rng_)


def f_nivel(ref):
    return f'=IF({ref}="","",IF({ref}>=20,"Crítico",IF({ref}>=10,"Alto",IF({ref}>=5,"Médio","Baixo"))))'


def chart_style(ch):
    ch.style = 2
    ch.title = None
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ch.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill="D3DBE0", w=9525))
    ch.x_axis.graphicalProperties = GraphicalProperties(ln=LineProperties(solidFill="D3DBE0", w=9525))
    ch.y_axis.graphicalProperties = GraphicalProperties(ln=LineProperties(noFill=True))


def niveis_chart(ws, anchor, r1, r2, c_cat, c_a, c_b, width=16, height=7.5):
    """Colunas agrupadas: número de riscos por nível, inicial e residual."""
    ch = BarChart()
    ch.type = "col"
    ch.grouping = "clustered"
    ch.gapWidth = 80
    ch.overlap = -10
    ch.height, ch.width = height, width
    chart_style(ch)
    ch.legend.position = "b"
    for col, nome, cor in ((c_a, "Inicial", "9BABB6"), (c_b, "Residual", BLUE)):
        s = Series(Reference(ws, min_col=col, min_row=r1, max_row=r2), title=nome)
        s.graphicalProperties.solidFill = cor
        s.graphicalProperties.line.noFill = True
        s.dLbls = DataLabelList()
        s.dLbls.showVal = True
        s.dLbls.showSerName = s.dLbls.showCatName = s.dLbls.showLegendKey = False
        ch.append(s)
    ch.set_categories(Reference(ws, min_col=c_cat, min_row=r1, max_row=r2))
    ch.y_axis.scaling.min = 0
    ch.y_axis.majorUnit = 1
    ws.add_chart(ch, anchor)


def tabela_niveis(ws, top, colJ, colS, last, c0="B", merge_to="D", ca="E", cb="F"):
    """Contagem por nível, inicial e residual. colJ e colS são os intervalos com os níveis."""
    put(ws, f"{c0}{top}", "Nível", f=font(10, True, c=WHITE), bg=INK, h="center", merge=f"{c0}{top}:{merge_to}{top}")
    head(ws, f"{ca}{top}", "Inicial")
    head(ws, f"{cb}{top}", "Residual")
    ws.row_dimensions[top].height = 21.75
    for k, n in enumerate(NIVEIS, 1):
        rr = top + k
        put(ws, f"{c0}{rr}", n, f=font(10, True), bg=NV_TINT[n], merge=f"{c0}{rr}:{merge_to}{rr}")
        calc(ws, f"{ca}{rr}", f'=COUNTIF({colJ},"{n}")', b=False)
        calc(ws, f"{cb}{rr}", f'=COUNTIF({colS},"{n}")', b=False)
        ws.row_dimensions[rr].height = 21.75
    return top + 1, top + 4


# ================================================================== pasta de trabalho
wb = Workbook()
wb.properties.title = "Matriz de riscos — Modelo de riscos, mapa e oportunidades"
wb.properties.creator = "Modelo Riscos"
wb.properties.language = "pt-BR"

# ------------------------------------------------------------------ Instruções
ws = wb.active
ws.title = "Instruções"
widths(ws, {"A": 2, "B": 24, "C": 92, "D": 2})
title(ws, "Matriz de riscos — Como usar esta planilha",
      "Modelo para descrever os riscos, avaliar o nível, definir a resposta e acompanhar o risco residual.", "C")
r = 4


def section(text):
    global r
    band(ws, r, text, "C", sz=11)
    r += 1


def line(k, v, kbg=GRAY, vbg=None, kf=None, kh="left", height=None):
    global r
    put(ws, f"B{r}", k, f=kf or font(10, True), bg=kbg, h=kh)
    put(ws, f"C{r}", v, bg=vbg)
    ws.row_dimensions[r].height = height or (19.5 if len(v) <= 98 else (31.5 if len(v) <= 196 else 45.75))
    r += 1


section("Legenda: onde preencher")
line("Amarelo-claro", "Células de entrada. É aqui que você digita.", vbg=INPUT)
line("Cinza", "Células calculadas ou fixas (rótulos, pontos, níveis, avisos). Não altere.", vbg=GRAY)
line("Cinza em itálico", "Linha de exemplo (marcada com “Ex.”). Mostra o formato esperado e não entra nas contagens.", vbg=GRAY)
line("Listas suspensas", "Colunas de resposta, decisão e status aceitam apenas as opções da lista. As notas aceitam números inteiros de 1 a 5.")
line("Comentários", "Células com triângulo vermelho no canto trazem uma dica de preenchimento. Passe o mouse sobre elas.")
r += 1
section("Os quatro níveis")
for n, faixa, txt in CONDUTA:
    line(n, f"De {faixa} pontos. {txt}", kbg=NV_COR[n], vbg=NV_TINT[n], kf=font(10, True, c=INK if n == "Médio" else WHITE))
r += 1
section("Ordem recomendada de preenchimento")
for k, text in enumerate([
    "Aba Escalas: leia as descrições das notas e combine com o grupo como usá-las.",
    "Aba Riscos: preencha o cabeçalho e descreva cada risco com causa, evento e consequência.",
    "Aba Riscos: dê as notas de probabilidade e de impacto, e leia os pontos e o nível.",
    "Aba Riscos: escolha a resposta e escreva a ação, o responsável e o prazo. Se a resposta for aceitar, escreva a justificativa.",
    "Aba Riscos: estime a probabilidade e o impacto residuais, já contando com a ação.",
    "Aba Mapa: compare a matriz inicial com a residual e leia o resumo.",
    "Aba Oportunidades: registre as oportunidades do processo. Aba Checklist: valide a matriz.",
], 1):
    line(f"Passo {k}", text)
r += 1
section("Abas da planilha")
for k, text in [
    ("Riscos", "Registro dos riscos, com análise, resposta, risco residual, conferência e alertas."),
    ("Mapa", "As duas matrizes, inicial e residual, com contagem por nível, resumo e gráfico."),
    ("Oportunidades", "Registro das oportunidades, com probabilidade, benefício, prioridade e ação."),
    ("Escalas", "Descrição das notas de probabilidade, de impacto e de benefício, dos níveis e das respostas."),
    ("Checklist", "Doze verificações de qualidade da matriz, com percentual de conclusão."),
    ("Exemplo 1 - Pizzaria", "Riscos do processo de delivery de uma organização pequena."),
    ("Exemplo 2 - Compras", "Riscos do processo de aquisição de uma indústria."),
]:
    line(k, text)
r += 1
section("Regras e premissas do modelo")
for k, text in [
    ("Pontos e nível", "Pontos = probabilidade × impacto. De 1 a 4, baixo. De 5 a 9, médio. De 10 a 16, alto. De 20 a 25, crítico. As escalas e as faixas são uma convenção deste modelo."),
    ("Impacto 5", "A multiplicação trata como iguais um risco raro e grave e um risco frequente e leve. Por isso, a coluna Alerta avisa quando o impacto é 5, qualquer que seja o nível."),
    ("Risco residual", "É uma estimativa, feita ao definir a resposta. Depois que a ação estiver implantada, confira com fatos e ajuste as notas."),
    ("Risco aceito", "Na resposta Aceitar, escreva a justificativa na coluna da ação e repita as notas iniciais nas colunas residuais."),
    ("Situação", "Uma ação é considerada atrasada quando o prazo é anterior à data de hoje e ela ainda não foi concluída."),
    ("Exemplos", "Os dados das abas de exemplo são ilustrativos, criados para este material."),
]:
    line(k, text)
setup(ws, INK, f"B1:C{r}", landscape=False, fit_height=True)

# ------------------------------------------------------------------ Riscos
ws = wb.create_sheet("Riscos")
widths(ws, {"A": 2, "B": 5, "C": 18, "D": 28, "E": 28, "F": 28, "G": 7, "H": 7, "I": 9, "J": 11, "K": 15, "L": 36, "M": 18, "N": 12, "O": 14,
            "P": 7, "Q": 7, "R": 9, "S": 11, "T": 13, "U": 26, "V": 30, "W": 2})
title(ws, "Matriz de riscos", "Um risco por linha. Preencha as células em amarelo-claro. As células cinza são calculadas.", "V")
for rr, l1, l2, l3, f3 in [(4, "Organização", "Processo", "Data da análise", DATE), (5, "Objetivo do processo", "Participantes", "Data da revisão", DATE)]:
    label(ws, f"B{rr}", l1, merge=f"B{rr}:C{rr}")
    inp(ws, f"D{rr}", merge=f"D{rr}:F{rr}")
    label(ws, f"G{rr}", l2, merge=f"G{rr}:J{rr}")
    inp(ws, f"K{rr}", merge=f"K{rr}:O{rr}")
    label(ws, f"P{rr}", l3, merge=f"P{rr}:S{rr}")
    inp(ws, f"T{rr}", merge=f"T{rr}:V{rr}", fmt=f3, h="left")
    ws.row_dimensions[rr].height = 24
dv_date(ws, "T4:T5")
for rng_, text, cor in [("B7:F7", "Identificação", INK), ("G7:J7", "Análise", BLUE), ("K7:O7", "Resposta", INK), ("P7:S7", "Risco residual", BLUE),
                        ("T7:V7", "Acompanhamento", INK)]:
    put(ws, rng_.split(":")[0], text, f=font(10, True, c=WHITE), bg=cor, box=False, merge=rng_)
ws.row_dimensions[7].height = 21.75
HEADS = ["#", "Processo ou etapa", "Causa", "Evento", "Consequência", "P", "I", "Pontos", "Nível", "Resposta", "Ação ou justificativa", "Responsável",
         "Prazo", "Status", "P", "I", "Pontos", "Nível", "Situação", "Conferência", "Alerta"]
for k, text in enumerate(HEADS):
    head(ws, f"{L(2 + k)}8", text)
ws.row_dimensions[8].height = 24
hint_row(ws, 9, [("B", "", None), ("C", "Onde o risco nasce", None), ("D", "O que existe hoje", None), ("E", "O que pode acontecer", None),
                 ("F", "O efeito sobre o objetivo", None), ("G", "1 a 5", None), ("H", "1 a 5", None), ("I", "P × I", None), ("J", "Faixa", None),
                 ("K", "Escolha na lista", None), ("L", "Verbo no infinitivo", None), ("M", "Cargo ou nome", None), ("N", "Data", None),
                 ("O", "Da ação", None), ("P", "1 a 5", None), ("Q", "1 a 5", None), ("R", "P × I", None), ("S", "Faixa", None),
                 ("T", "Calculada", None), ("U", "Campos da linha", None), ("V", "Pontos de atenção", None)])
ex(ws, "B10", "Ex.", h="center")
for col, v, h_, f_ in [("C", "Planejamento de compras", "left", None), ("D", "Um só fornecedor de resina", "left", None),
                       ("E", "O fornecedor deixa de entregar", "left", None), ("F", "Parada da produção por falta de material", "left", None),
                       ("G", 4, "center", None), ("H", 5, "center", None), ("I", 20, "center", None), ("J", "Crítico", "center", None),
                       ("K", "Reduzir", "center", None), ("L", "Homologar um segundo fornecedor e manter estoque mínimo de 15 dias.", "left", None),
                       ("M", "Gerente de Suprimentos", "left", None), ("N", date(2026, 12, 18), "center", DATE), ("O", "Em andamento", "center", None),
                       ("P", 2, "center", None), ("Q", 3, "center", None), ("R", 6, "center", None), ("S", "Médio", "center", None),
                       ("T", "No prazo", "center", None), ("U", "Completo", "center", None), ("V", "Impacto 5: tenha plano de contingência", "center", None)]:
    ex(ws, f"{col}10", v, h=h_, fmt=f_)
ws.row_dimensions[10].height = 36
R1, R2 = 11, 30
for k in range(20):
    rr = R1 + k
    num(ws, f"B{rr}", k + 1)
    for col in "CDEF":
        inp(ws, f"{col}{rr}")
    inp(ws, f"G{rr}", h="center")
    inp(ws, f"H{rr}", h="center")
    calc(ws, f"I{rr}", f'=IF(OR(G{rr}="",H{rr}=""),"",G{rr}*H{rr})')
    calc(ws, f"J{rr}", f_nivel(f"I{rr}"), b=False)
    inp(ws, f"K{rr}", h="center")
    inp(ws, f"L{rr}")
    inp(ws, f"M{rr}")
    inp(ws, f"N{rr}", h="center", fmt=DATE)
    inp(ws, f"O{rr}", h="center")
    inp(ws, f"P{rr}", h="center")
    inp(ws, f"Q{rr}", h="center")
    calc(ws, f"R{rr}", f'=IF(OR(P{rr}="",Q{rr}=""),"",P{rr}*Q{rr})')
    calc(ws, f"S{rr}", f_nivel(f"R{rr}"), b=False)
    calc(ws, f"T{rr}", f'=IF(OR(E{rr}="",K{rr}=""),"",IF(K{rr}="{ACEITAR}","Aceito",IF(O{rr}="Concluída","Concluída",'
         f'IF(N{rr}="","Sem prazo",IF(N{rr}<TODAY(),"Atrasada","No prazo")))))', b=False, sz=9)
    calc(ws, f"U{rr}", f'=IF(E{rr}="","",IF(OR(D{rr}="",F{rr}=""),"Descreva a causa e a consequência",IF(I{rr}="","Faltam as notas de P e I",'
         f'IF(K{rr}="","Escolha a resposta",IF(L{rr}="",IF(K{rr}="{ACEITAR}","Justifique a decisão de aceitar","Escreva a ação"),'
         f'IF(AND(K{rr}<>"{ACEITAR}",OR(M{rr}="",N{rr}="")),"Falta responsável ou prazo",IF(R{rr}="","Avalie o risco residual",'
         f'IF(R{rr}>I{rr},"Residual maior que o inicial","Completo"))))))))', b=False, sz=9)
    calc(ws, f"V{rr}", f'=IF(I{rr}="","",IF(AND(OR(J{rr}="Alto",J{rr}="Crítico"),K{rr}="{ACEITAR}"),"Risco alto aceito: peça a decisão da direção",'
         f'IF(OR(S{rr}="Alto",S{rr}="Crítico"),"Residual ainda alto: reforce a resposta",IF(OR(H{rr}=5,Q{rr}=5),"Impacto 5: tenha plano de contingência",""))))',
         b=False, sz=9)
    ws.row_dimensions[rr].height = 39
dv_nota(ws, f"G{R1}:H{R2}")
dv_nota(ws, f"P{R1}:Q{R2}")
dv_list(ws, f"K{R1}:K{R2}", RESPOSTAS, "Evitar, Reduzir, Compartilhar ou Aceitar")
dv_list(ws, f"O{R1}:O{R2}", ST_ACAO, "Não iniciada, Em andamento ou Concluída")
dv_date(ws, f"N{R1}:N{R2}")
cf_equal(ws, f"J{R1}:J{R2}", NV_CF)
cf_equal(ws, f"S{R1}:S{R2}", NV_CF)
cf_equal(ws, f"O{R1}:O{R2}", ST_CF)
cf_equal(ws, f"T{R1}:T{R2}", SIT_CF)
cf_warn(ws, f"U{R1}:U{R2}", f"U{R1}", ok_values=("Completo",))
ws.conditional_formatting.add(f"V{R1}:V{R2}", FormulaRule(formula=[f"LEN(V{R1})>0"], fill=PatternFill("solid", bgColor=YELLOW, fgColor=YELLOW)))
note(ws, "D8", 'A condição que existe hoje e pode dar origem ao evento.\nEx.: "Um só fornecedor de resina".')
note(ws, "E8", 'O que pode acontecer. Se já aconteceu, é um problema: trate como não conformidade.\nEx.: "O fornecedor deixa de entregar".')
note(ws, "F8", 'O efeito sobre o objetivo do processo.\nEx.: "Parada da produção por falta de material".')
note(ws, "K8", "Evitar: impedir que aconteça. Reduzir: agir na causa ou na consequência. Compartilhar: dividir o efeito, por contrato ou seguro. Aceitar: não agir agora, com justificativa.")
note(ws, "P8", "Notas estimadas já contando com a ação. No risco aceito, repita as notas iniciais.")
s = R2 + 2
band(ws, s, "Resumo automático", "V")
for k, (text, formula, fmt) in enumerate([
    ("Riscos registrados", f"=COUNTA(E{R1}:E{R2})", None),
    ("Riscos com o registro completo", f'=COUNTIF(U{R1}:U{R2},"Completo")', None),
    ("Riscos aceitos", f'=COUNTIF(K{R1}:K{R2},"{ACEITAR}")', None),
    ("Ações atrasadas", f'=COUNTIF(T{R1}:T{R2},"Atrasada")', None),
    ("Riscos com alerta", f'=SUMPRODUCT(--(LEN(V{R1}:V{R2})>0))', None),
    ("Aviso", f'=IF(G{s+1}=0,"Registre os riscos",IF(G{s+2}<G{s+1},"Há risco com o registro incompleto: veja a coluna Conferência",'
              f'IF(G{s+4}>0,"Há ação atrasada",IF($T$5="","Marque a data da revisão","OK"))))', None),
], 1):
    label(ws, f"B{s+k}", text, merge=f"B{s+k}:F{s+k}", h="right")
    calc(ws, f"G{s+k}", formula, fmt=fmt, merge=f"G{s+k}:K{s+k}", sz=9 if text == "Aviso" else 10, b=text != "Aviso")
    ws.row_dimensions[s + k].height = 21.75
cf_warn(ws, f"G{s+6}:K{s+6}", f"G{s+6}")
ws.freeze_panes = "F9"
setup(ws, REDC, f"B1:V{s+6}")
ws.print_title_rows = "7:8"

# ------------------------------------------------------------------ Mapa
ws = wb.create_sheet("Mapa")
widths(ws, {"A": 2, "B": 17, "C": 11, "D": 11, "E": 11, "F": 11, "G": 11, "H": 3, "I": 17, "J": 11, "K": 11, "L": 11, "M": 11, "N": 11, "O": 2})
title(ws, "Mapa de riscos", "Número de riscos em cada célula da matriz, antes e depois das respostas. Calculado a partir da aba Riscos.", "N")
label(ws, "B4", "Processo")
calc(ws, "C4", '=IF(Riscos!K4="","",Riscos!K4)', h="left", b=False, merge="C4:G4")
label(ws, "I4", "Data da análise")
calc(ws, "J4", '=IF(Riscos!T4="","",Riscos!T4)', h="left", b=False, fmt=DATE, merge="J4:N4")
ws.row_dimensions[4].height = 21.75


def grade(c0, titulo, colp, coli, top=6):
    """Matriz 5 × 5 com contagem de riscos. c0 é o índice da coluna dos rótulos de impacto."""
    put(ws, f"{L(c0)}{top}", titulo, f=font(10, True, c=WHITE), bg=INK, box=False, merge=f"{L(c0)}{top}:{L(c0 + 5)}{top}")
    ws.row_dimensions[top].height = 21.75
    for row, (n, nome, _) in enumerate(reversed(IMPACTO)):
        rr = top + 1 + row
        put(ws, f"{L(c0)}{rr}", f"{n} · {nome}", f=font(10, True), bg=GRAY, h="right")
        for col, (p, _, _) in enumerate(PROB):
            ref = f"{L(c0 + 1 + col)}{rr}"
            put(ws, ref, f"=COUNTIFS(Riscos!${colp}${R1}:${colp}${R2},{p},Riscos!${coli}${R1}:${coli}${R2},{n})",
                f=font(14, True), bg=NV_TINT[nivel(p * n)], h="center", fmt="0;-0;;@")
        ws.row_dimensions[rr].height = 36
    rr = top + 6
    put(ws, f"{L(c0)}{rr}", "Impacto ↑  Probabilidade →", f=font(9, i=True, c=MUTED), bg=GRAY, h="right")
    for col, (p, nome, _) in enumerate(PROB):
        put(ws, f"{L(c0 + 1 + col)}{rr}", f"{p} · {nome}", f=font(9, True), bg=GRAY, h="center")
    ws.row_dimensions[rr].height = 30


grade(2, "Risco inicial", "G", "H")
grade(9, "Risco residual", "P", "Q")
J_, S_, I_, R_ = (f"Riscos!${c}${R1}:${c}${R2}" for c in "JSIR")
a, b = tabela_niveis(ws, 14, J_, S_, "G", c0="B", merge_to="D", ca="E", cb="F")
s = 20
band(ws, s, "Resumo automático", "G")
for k, (text, formula, fmt) in enumerate([
    ("Riscos avaliados", f"=COUNT({I_})", None),
    ("Altos e críticos, inicial", f"=E{a+2}+E{a+3}", None),
    ("Altos e críticos, residual", f"=F{a+2}+F{a+3}", None),
    ("Média dos pontos, inicial", f'=IF(COUNT({I_})=0,"",AVERAGE({I_}))', "0.0"),
    ("Média dos pontos, residual", f'=IF(COUNT({R_})=0,"",AVERAGE({R_}))', "0.0"),
    ("Aviso", f'=IF(E{s+1}=0,"Registre os riscos na aba Riscos",IF(COUNT({R_})<E{s+1},"Há risco sem avaliação residual",'
              f'IF(E{s+3}>0,"Há risco residual alto ou crítico: reforce a resposta","OK")))', None),
], 1):
    label(ws, f"B{s+k}", text, merge=f"B{s+k}:D{s+k}", h="right")
    calc(ws, f"E{s+k}", formula, fmt=fmt, merge=f"E{s+k}:G{s+k}", sz=9 if text == "Aviso" else 10, b=text != "Aviso")
    ws.row_dimensions[s + k].height = 30 if text == "Aviso" else 21.75
cf_warn(ws, f"E{s+6}:G{s+6}", f"E{s+6}")
niveis_chart(ws, "I14", a, b, 2, 5, 6, width=14.2, height=8.2)
setup(ws, BLUE, f"B1:N{s+7}", fit_height=True)

# ------------------------------------------------------------------ Oportunidades
ws = wb.create_sheet("Oportunidades")
widths(ws, {"A": 2, "B": 5, "C": 44, "D": 18, "E": 14, "F": 12, "G": 9, "H": 12, "I": 16, "J": 38, "K": 20, "L": 12, "M": 14, "N": 13, "O": 2})
title(ws, "Oportunidades", "Uma oportunidade por linha. A conta é a mesma do risco, com o benefício no lugar do impacto.", "N")
for col, text in zip("BCDEFGHIJKLMN", ["#", "Oportunidade", "Processo", "Probabilidade", "Benefício", "Pontos", "Prioridade", "Decisão", "Ação",
                                       "Responsável", "Prazo", "Status", "Situação"]):
    head(ws, f"{col}4", text)
ws.row_dimensions[4].height = 24
hint_row(ws, 5, [("B", "", None), ("C", "O que pode acontecer de bom", None), ("D", "Onde", None), ("E", "1 a 5", None), ("F", "1 a 5", None),
                 ("G", "Calculado", None), ("H", "Faixa", None), ("I", "Escolha na lista", None), ("J", "Verbo no infinitivo", None),
                 ("K", "Cargo ou nome", None), ("L", "Data", None), ("M", "Da ação", None), ("N", "Calculada", None)])
ex(ws, "B6", "Ex.", h="center")
for col, v, h_, f_ in [("C", "Cliente atual pede embalagem com material reciclado", "left", None), ("D", "Desenvolvimento", "left", None),
                       ("E", 4, "center", None), ("F", 4, "center", None), ("G", 16, "center", None), ("H", "Alta", "center", None),
                       ("I", "Aproveitar", "center", None), ("J", "Abrir projeto de desenvolvimento da nova linha.", "left", None),
                       ("K", "Gerente de engenharia", "left", None), ("L", date(2026, 11, 30), "center", DATE), ("M", "Em andamento", "center", None),
                       ("N", "No prazo", "center", None)]:
    ex(ws, f"{col}6", v, h=h_, fmt=f_)
ws.row_dimensions[6].height = 33
O1, O2 = 7, 16
for k in range(10):
    rr = O1 + k
    num(ws, f"B{rr}", k + 1)
    inp(ws, f"C{rr}")
    inp(ws, f"D{rr}")
    inp(ws, f"E{rr}", h="center")
    inp(ws, f"F{rr}", h="center")
    calc(ws, f"G{rr}", f'=IF(OR(E{rr}="",F{rr}=""),"",E{rr}*F{rr})')
    calc(ws, f"H{rr}", f'=IF(G{rr}="","",IF(G{rr}>=10,"Alta",IF(G{rr}>=5,"Média","Baixa")))', b=False)
    inp(ws, f"I{rr}", h="center")
    inp(ws, f"J{rr}")
    inp(ws, f"K{rr}")
    inp(ws, f"L{rr}", h="center", fmt=DATE)
    inp(ws, f"M{rr}", h="center")
    calc(ws, f"N{rr}", f'=IF(OR(C{rr}="",I{rr}=""),"",IF(I{rr}="Não agir agora","Sem ação",IF(M{rr}="Concluída","Concluída",'
         f'IF(L{rr}="","Sem prazo",IF(L{rr}<TODAY(),"Atrasada","No prazo")))))', b=False, sz=9)
    ws.row_dimensions[rr].height = 33
dv_nota(ws, f"E{O1}:F{O2}")
dv_list(ws, f"I{O1}:I{O2}", DECISOES, "Aproveitar, Testar ou Não agir agora")
dv_list(ws, f"M{O1}:M{O2}", ST_ACAO, "Não iniciada, Em andamento ou Concluída")
dv_date(ws, f"L{O1}:L{O2}")
cf_equal(ws, f"H{O1}:H{O2}", PRIO_CF)
cf_equal(ws, f"M{O1}:M{O2}", ST_CF)
cf_equal(ws, f"N{O1}:N{O2}", SIT_CF + [("Sem ação", "E3E8EB")])
s = O2 + 2
band(ws, s, "Resumo automático", "N")
for k, (text, formula) in enumerate([
    ("Oportunidades registradas", f"=COUNTA(C{O1}:C{O2})"),
    ("Prioridade alta", f'=COUNTIF(H{O1}:H{O2},"Alta")'),
    ("Com decisão de aproveitar ou testar", f'=COUNTIF(I{O1}:I{O2},"Aproveitar")+COUNTIF(I{O1}:I{O2},"Testar")'),
    ("Ações atrasadas", f'=COUNTIF(N{O1}:N{O2},"Atrasada")'),
    ("Aviso", f'=IF(E{s+1}=0,"Registre as oportunidades",IF(COUNT(G{O1}:G{O2})<E{s+1},"Há oportunidade sem as duas notas",'
              f'IF(COUNTA(I{O1}:I{O2})<E{s+1},"Há oportunidade sem decisão",IF(SUMPRODUCT((H{O1}:H{O2}="Alta")*(I{O1}:I{O2}="Não agir agora"))>0,'
              f'"Há oportunidade de prioridade alta sem ação",IF(E{s+4}>0,"Há ação atrasada","OK")))))'),
], 1):
    label(ws, f"B{s+k}", text, merge=f"B{s+k}:D{s+k}", h="right")
    calc(ws, f"E{s+k}", formula, merge=f"E{s+k}:I{s+k}", sz=9 if text == "Aviso" else 10, b=text != "Aviso")
    ws.row_dimensions[s + k].height = 21.75
cf_warn(ws, f"E{s+5}:I{s+5}", f"E{s+5}")
ws.freeze_panes = "D5"
setup(ws, TEAL, f"B1:N{s+5}", fit_height=True)

# ------------------------------------------------------------------ Escalas
ws = wb.create_sheet("Escalas")
widths(ws, {"A": 2, "B": 8, "C": 18, "D": 52, "E": 18, "F": 52, "G": 2})
title(ws, "Escalas, níveis e respostas", "Referência para as notas. Ajuste as descrições à sua organização e use as mesmas em todos os processos.", "F")
band(ws, 4, "Probabilidade e impacto", "F", color=BLUE)
for col, text in zip("BCDEF", ["Nota", "Probabilidade", "O evento acontece", "Impacto", "Se acontecer"]):
    put(ws, f"{col}5", text, f=font(10, True), bg=GRAY, h="center")
ws.row_dimensions[5].height = 21.75
for k, ((n, pn, pd_), (_, im, idesc)) in enumerate(zip(PROB, IMPACTO)):
    rr = 6 + k
    put(ws, f"B{rr}", n, f=font(12, True), bg=GRAY, h="center")
    put(ws, f"C{rr}", pn, f=font(10, True))
    put(ws, f"D{rr}", pd_)
    put(ws, f"E{rr}", im, f=font(10, True))
    put(ws, f"F{rr}", idesc)
    ws.row_dimensions[rr].height = 27
band(ws, 12, "Benefício, para as oportunidades", "F", color=TEAL)
for k, (n, nome, desc) in enumerate(BENEFICIO):
    rr = 13 + k
    put(ws, f"B{rr}", n, f=font(12, True), bg=GRAY, h="center")
    put(ws, f"C{rr}", nome, f=font(10, True))
    put(ws, f"D{rr}", desc, merge=f"D{rr}:F{rr}")
    ws.row_dimensions[rr].height = 21.75
band(ws, 19, "Níveis e conduta", "F")
for k, (n, faixa, txt) in enumerate(CONDUTA):
    rr = 20 + k
    put(ws, f"B{rr}", "", bg=NV_COR[n])
    put(ws, f"C{rr}", n, f=font(10, True), bg=NV_TINT[n])
    put(ws, f"D{rr}", f"De {faixa} pontos", bg=NV_TINT[n])
    put(ws, f"E{rr}", txt, bg=NV_TINT[n], merge=f"E{rr}:F{rr}")
    ws.row_dimensions[rr].height = 21.75
band(ws, 25, "Respostas ao risco", "F")
for col, text in zip("CDE", ["Resposta", "O que é", "Efeito no risco"]):
    put(ws, f"{col}26", text, f=font(10, True), bg=GRAY, h="center")
put(ws, "B26", "", bg=GRAY)
put(ws, "F26", "Exemplo", f=font(10, True), bg=GRAY, h="center")
ws.row_dimensions[26].height = 21.75
for k, (rsp, oque, efeito, exemplo) in enumerate(RESP_TXT):
    rr = 27 + k
    put(ws, f"B{rr}", k + 1, f=font(10, True, c=MUTED), bg=GRAY, h="center")
    put(ws, f"C{rr}", rsp, f=font(10, True))
    put(ws, f"D{rr}", oque)
    put(ws, f"E{rr}", efeito)
    put(ws, f"F{rr}", exemplo)
    ws.row_dimensions[rr].height = 42
setup(ws, PURPLE, "B1:F30", fit_height=True)

# ------------------------------------------------------------------ Checklist
ws = wb.create_sheet("Checklist")
widths(ws, {"A": 2, "B": 5, "C": 70, "D": 14, "E": 44, "F": 2})
title(ws, "Checklist de validação da matriz", "Escolha o status de cada item na lista suspensa (coluna em amarelo-claro).", "E")
for col, text in zip("BCDE", ["#", "Item de verificação", "Status", "Observações"]):
    head(ws, f"{col}4", text)
ws.row_dimensions[4].height = 21.75
for k, text in enumerate(CHECK):
    rr = 5 + k
    num(ws, f"B{rr}", k + 1)
    put(ws, f"C{rr}", text, bg=GRAY)
    inp(ws, f"D{rr}", h="center")
    inp(ws, f"E{rr}")
    ws.row_dimensions[rr].height = 30
dv_list(ws, "D5:D16", YESNO, "Sim, Parcial ou Não")
cf_equal(ws, "D5:D16", YESNO_CF)
band(ws, 18, "Resumo automático", "E")
summary(ws, 19, "Itens atendidos (Sim)", '=COUNTIF(D5:D16,"Sim")', "C", "D")
summary(ws, 20, "Itens parciais", '=COUNTIF(D5:D16,"Parcial")', "C", "D")
summary(ws, 21, "Itens não atendidos", '=COUNTIF(D5:D16,"Não")', "C", "D")
summary(ws, 22, "Itens sem resposta", "=COUNTBLANK(D5:D16)", "C", "D")
summary(ws, 23, "Percentual atendido", "=D19/ROWS(D5:D16)", "C", "D", fmt="0%")
summary(ws, 24, "Situação", '=IF(D19=ROWS(D5:D16),"Validado",IF(D22>0,"Em andamento","Com pendências"))', "C", "D")
setup(ws, PURPLE, "B1:E24", landscape=False, fit_height=True)


# ------------------------------------------------------------------ Exemplos
def example(ws, data):
    W = {"B": 6, "C": 20, "D": 24, "E": 24, "F": 26, "G": 6, "H": 6, "I": 9, "J": 11, "K": 15, "L": 40, "M": 20, "N": 12, "O": 7, "P": 7,
         "Q": 9, "R": 11}
    widths(ws, dict(W, A=2, S=2))
    title(ws, "Matriz de riscos", "Exemplo preenchido, para consulta. Use as abas Riscos, Mapa e Oportunidades para a sua matriz.", "R")
    H = data["head"]
    rr = 4
    for rot, val, fmt in [("Organização", H["org"], None), ("Processo", H["processo"], None), ("Objetivo do processo", H["objetivo"], None),
                          ("Origem dos riscos", H["origem"], None), ("Data e participantes", f'{H["data"]:%d/%m/%Y}. {H["por"]}.', None),
                          ("Data da revisão", H["revisao"], DATE)]:
        label(ws, f"B{rr}", rot, merge=f"B{rr}:C{rr}")
        inp(ws, f"D{rr}", val, merge=f"D{rr}:R{rr}", bg=WHITE, fmt=fmt)
        ws.row_dimensions[rr].height = 21.75
        rr += 1
    rr += 1
    for rng_, text, cor in [(f"B{rr}:F{rr}", "Identificação", INK), (f"G{rr}:J{rr}", "Análise", BLUE), (f"K{rr}:N{rr}", "Resposta", INK),
                            (f"O{rr}:R{rr}", "Risco residual", BLUE)]:
        put(ws, rng_.split(":")[0], text, f=font(10, True, c=WHITE), bg=cor, box=False, merge=rng_)
    ws.row_dimensions[rr].height = 21.75
    rr += 1
    hd = rr
    for k, text in enumerate(["Nº", "Processo ou etapa", "Causa", "Evento", "Consequência", "P", "I", "Pontos", "Nível", "Resposta",
                              "Ação ou justificativa", "Responsável", "Prazo", "P", "I", "Pontos", "Nível"]):
        put(ws, f"{L(2 + k)}{rr}", text, f=font(10, True), bg=GRAY, h="center")
    ws.row_dimensions[rr].height = 24
    r1 = rr + 1
    for x in data["riscos"]:
        rr += 1
        acao = ACEITE[x["id"]] if x["resp"] == ACEITAR else x["acao"]
        put(ws, f"B{rr}", x["id"], f=font(10, True), bg=GRAY, h="center")
        for col, key in zip("CDEF", ("processo", "causa", "evento", "conseq")):
            put(ws, f"{col}{rr}", x[key])
        put(ws, f"G{rr}", x["p"], h="center")
        put(ws, f"H{rr}", x["i"], h="center")
        calc(ws, f"I{rr}", f"=G{rr}*H{rr}")
        calc(ws, f"J{rr}", f_nivel(f"I{rr}"), b=False)
        put(ws, f"K{rr}", x["resp"], h="center")
        put(ws, f"L{rr}", acao)
        put(ws, f"M{rr}", x["quem"])
        put(ws, f"N{rr}", x["prazo"] if x["prazo"] else "Sem ação", h="center", fmt=DATE)
        put(ws, f"O{rr}", x["pr"], h="center")
        put(ws, f"P{rr}", x["ir"], h="center")
        calc(ws, f"Q{rr}", f"=O{rr}*P{rr}")
        calc(ws, f"R{rr}", f_nivel(f"Q{rr}"), b=False)
        ws.row_dimensions[rr].height = max(36, alt(acao, W["L"] * 1.1), alt(x["causa"], W["D"] * 1.1), alt(x["evento"], W["E"] * 1.1),
                                           alt(x["conseq"], W["F"] * 1.1))
    r2 = rr
    cf_equal(ws, f"J{r1}:J{r2}", NV_CF)
    cf_equal(ws, f"R{r1}:R{r2}", NV_CF)
    rr += 2
    ws.row_breaks.append(Break(id=rr - 1))
    band(ws, rr, "Riscos por nível", "F")
    a, b = tabela_niveis(ws, rr + 1, f"$J${r1}:$J${r2}", f"$R${r1}:$R${r2}", "F", c0="B", merge_to="D", ca="E", cb="F")
    s0 = b + 2
    band(ws, s0, "Resumo automático", "F")
    for k, (text, formula, fmt) in enumerate([
        ("Riscos avaliados", f"=COUNT(I{r1}:I{r2})", None),
        ("Altos e críticos, inicial", f"=E{a+2}+E{a+3}", None),
        ("Altos e críticos, residual", f"=F{a+2}+F{a+3}", None),
        ("Média dos pontos, inicial", f"=AVERAGE(I{r1}:I{r2})", "0.0"),
        ("Média dos pontos, residual", f"=AVERAGE(Q{r1}:Q{r2})", "0.0"),
        ("Riscos de impacto 5", f"=COUNTIF(H{r1}:H{r2},5)", None),
        ("Riscos aceitos", f'=COUNTIF(K{r1}:K{r2},"{ACEITAR}")', None),
    ], 1):
        label(ws, f"B{s0+k}", text, merge=f"B{s0+k}:D{s0+k}", h="right")
        calc(ws, f"E{s0+k}", formula, fmt=fmt, merge=f"E{s0+k}:F{s0+k}")
        ws.row_dimensions[s0 + k].height = 21.75
    niveis_chart(ws, f"H{rr}", a, b, 2, 5, 6, width=16, height=8.4)
    setup(ws, MUTED, f"B1:R{s0+8}")
    ws.print_title_rows = f"{hd}:{hd}"


example(wb.create_sheet("Exemplo 1 - Pizzaria"), EX1)
example(wb.create_sheet("Exemplo 2 - Compras"), EX2)

wb.active = 1
wb.save(OUT)
print("ok", OUT)
