# -*- coding: utf-8 -*-
"""Gera GUT-modelo.xlsx no mesmo padrão visual das planilhas de SIPOC, PDCA e SWOT."""
import os
import sys
from datetime import date

_here = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, _here)
# funções de formatação compartilhadas com o gerador do PDCA (bloco inicial do arquivo)
_src = open(os.path.join(_here, "build_pdca.py"), encoding="utf-8").read()
exec(_src[: _src.index("STATUS = [")])

from openpyxl.chart import BarChart, Reference  # noqa: E402
from gut_data import ALTA, EX1, EX2, MEDIA  # noqa: E402

CR = {
    "G": ("7A4A9A", "E9DEF1", "Gravidade", "Qual é o dano?"),
    "U": ("A96A12", "F6E8CF", "Urgência", "Qual é o prazo?"),
    "T": ("1E7B73", "D9EEEB", "Tendência", "Vai piorar?"),
}
BLUE = "2B5C8A"
YESNO = ["Sim", "Parcial", "Não"]
YESNO_CF = [("Sim", GREEN), ("Parcial", YELLOW), ("Não", RED)]
STATUS_CF = [("Concluída", GREEN), ("Em andamento", YELLOW), ("Não iniciada", RED)]
PRIO_CF = [("Alta", RED), ("Média", YELLOW), ("Baixa", GREEN)]
N = 20
R1, R2 = 15, 15 + N - 1  # linhas dos problemas na aba GUT
TH_A, TH_M = "Escala!$C$14", "Escala!$C$15"


def note(ws, ref, text):
    c = Comment(text, "Modelo GUT")
    c.width, c.height = 280, 110
    ws[ref].comment = c


def dv_score(ws, rng):
    dv = DataValidation(type="whole", operator="between", formula1="1", formula2="5", allow_blank=True)
    dv.promptTitle, dv.prompt = "Nota de 1 a 5", "Veja a descrição de cada nota na aba Escala."
    dv.errorTitle, dv.error = "Valor inválido", "Digite um número inteiro de 1 a 5."
    dv.showErrorMessage = dv.showInputMessage = True
    ws.add_data_validation(dv)
    dv.add(rng)


def cf_filled(ws, rng, first, color=RED):
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f"LEN({first})>0"], fill=PatternFill("solid", bgColor=color, fgColor=color)))


GUT_W = {"A": 2, "B": 6, "C": 18, "D": 36, "E": 16, "F": 9, "G": 9, "H": 9, "I": 10, "J": 10, "K": 13, "L": 22, "M": 2, "N": 12}


def gut_sheet(ws, tab, data=None):
    widths(ws, GUT_W)
    is_ex = data is not None
    title(ws, "GUT — Matriz de priorização",
          "Exemplo preenchido, para consulta. Use a aba GUT para a sua lista." if is_ex
          else "Preencha as células em amarelo-claro. As células cinza são calculadas automaticamente.", "L")
    hb = WHITE if is_ex else INPUT
    H = data["head"] if is_ex else {}
    for r, l1, k1, l2, k2, f2 in [(4, "Tema da priorização", "tema", "Responsável", "resp", None),
                                  (5, "Área / unidade", "area", "Data", "data", DATE),
                                  (6, "Elaborado por", "autor", "Versão", "versao", None),
                                  (8, "Origem da lista", "origem", "Participantes", "part", None)]:
        label(ws, f"B{r}", l1, merge=f"B{r}:C{r}")
        inp(ws, f"D{r}", H.get(k1), merge=f"D{r}:E{r}", bg=hb)
        label(ws, f"F{r}", l2, merge=f"F{r}:H{r}")
        inp(ws, f"I{r}", H.get(k2), merge=f"I{r}:L{r}", bg=hb, fmt=f2)
        ws.row_dimensions[r].height = 21.75
    for r, l1, k1 in [(7, "Objetivo da priorização", "objetivo"), (9, "Fora do escopo", "fora")]:
        label(ws, f"B{r}", l1, merge=f"B{r}:C{r}")
        inp(ws, f"D{r}", H.get(k1), merge=f"D{r}:L{r}", bg=hb)
        ws.row_dimensions[r].height = 30
    ws.row_dimensions[8].height = 30
    dv_date(ws, "I5")

    put(ws, "B11", "Lista de problemas", f=font(10, True, c=WHITE), bg=INK, box=False, merge="B11:E11")
    for col, q in zip("FGH", "GUT"):
        put(ws, f"{col}11", q, f=font(20, True, c=WHITE), bg=CR[q][0], h="center", box=False)
    put(ws, "I11", "Resultado", f=font(10, True, c=WHITE), bg=INK, box=False, merge="I11:L11")
    ws.row_dimensions[11].height = 33.75
    head(ws, "B12", "#")
    put(ws, "C12", "Problema", f=font(10, True, c=WHITE), bg=INK, h="center", merge="C12:D12")
    head(ws, "E12", "Origem")
    for col, q in zip("FGH", "GUT"):
        put(ws, f"{col}12", CR[q][2], f=font(9, True, c=WHITE), bg=CR[q][0], h="center")
    for col, text in zip("IJKL", ["Pontos", "Posição", "Prioridade", "Alerta"]):
        head(ws, f"{col}12", text)
    ws.row_dimensions[12].height = 21.75
    hint = lambda ref, text, bg=GRAY, merge=None: put(ws, ref, text, f=font(9, i=True, c=MUTED), bg=bg, h="center", merge=merge)
    hint("B13", "")
    hint("C13", "Resultado indesejado, com fato, local e número", merge="C13:D13")
    hint("E13", "De onde veio")
    for col, q in zip("FGH", "GUT"):
        hint(f"{col}13", CR[q][3], bg=CR[q][1])
    hint("I13", "G × U × T")
    hint("J13", "1 = primeiro")
    hint("K13", "Faixa")
    hint("L13", "Gravidade 5")
    ws.row_dimensions[13].height = 30
    if is_ex:
        put(ws, "B14", "As notas foram dadas em grupo, um critério por vez. A posição segue a pontuação; nos empates, vale a maior gravidade, "
            "depois a maior urgência e a maior tendência.", f=font(9, i=True, c=MUTED), bg=GRAY, merge="B14:L14")
    else:
        ex(ws, "B14", "Ex.", h="center")
        put(ws, "C14", "Câmara fria com falha intermitente de temperatura", f=font(9, i=True, c=MUTED), bg=GRAY, merge="C14:D14")
        ex(ws, "E14", "Reunião", h="center")
        for col, v in zip("FGHIJ", [5, 5, 4, 100, 1]):
            ex(ws, f"{col}14", v, h="center")
        ex(ws, "K14", "Alta", h="center")
        ex(ws, "L14", "Gravidade máxima", h="center")
        note(ws, "D4", 'O conjunto de problemas que está sendo comparado.\nEx.: "Problemas da operação da pizzaria".')
        note(ws, "D7", 'Para que serve esta priorização.\nEx.: "Escolher os problemas que serão tratados no trimestre".')
        note(ws, "D8", 'De onde vieram os problemas da lista.\nEx.: "Reclamações, indicadores e reunião mensal".')
        note(ws, "F12", "Gravidade: tamanho do dano, se nada for feito.\n1 = sem gravidade; 5 = extremamente grave.")
        note(ws, "G12", "Urgência: tempo disponível para agir.\n1 = pode esperar; 5 = ação imediata.")
        note(ws, "H12", "Tendência: evolução do problema, se nada for feito.\n1 = estável; 5 = piora imediata.")
        note(ws, "L12", "Todo problema com gravidade 5 deve ser avaliado pelo grupo, qualquer que seja a pontuação.")
    ws.row_dimensions[14].height = 30
    put(ws, "N14", "Apoio", f=font(9, True, c=MUTED), bg=GRAY, h="center")
    for k in range(N):
        r = R1 + k
        p = data["itens"][k] if is_ex and k < len(data["itens"]) else None
        bg = WHITE if is_ex else INPUT
        put(ws, f"B{r}", f"P{k+1}", f=font(10, True, c=MUTED), bg=GRAY, h="center")
        inp(ws, f"C{r}", p["txt"] if p else None, merge=f"C{r}:D{r}", bg=bg)
        inp(ws, f"E{r}", p["origem"] if p else None, h="center", bg=bg)
        for col, q in zip("FGH", "GUT"):
            inp(ws, f"{col}{r}", p[q.lower()] if p else None, h="center", bg=(CR[q][1] if p else WHITE) if is_ex else INPUT)
        calc(ws, f"I{r}", f'=IF(COUNT(F{r}:H{r})<3,"",F{r}*G{r}*H{r})')
        calc(ws, f"J{r}", f'=IF(N{r}="","",RANK(N{r},$N${R1}:$N${R2})+COUNTIF($N${R1 - 1}:N{r - 1},N{r}))')
        calc(ws, f"K{r}", f'=IF(I{r}="","",IF(I{r}>={TH_A},"Alta",IF(I{r}>={TH_M},"Média","Baixa")))', b=False)
        calc(ws, f"L{r}", f'=IF(F{r}=5,"Gravidade máxima","")', b=False, sz=9)
        put(ws, f"N{r}", f'=IF(I{r}="","",I{r}*1000+F{r}*100+G{r}*10+H{r})', f=font(9, c=MUTED), bg=GRAY, h="center")
        ws.row_dimensions[r].height = 30 if is_ex else 27
    dv_list(ws, f"E{R1}:E{R2}", ["Auditoria", "Reclamação", "Indicador", "SWOT", "Reunião", "Outro"], "De onde veio o problema")
    dv_score(ws, f"F{R1}:H{R2}")
    cf_equal(ws, f"K{R1}:K{R2}", PRIO_CF)
    cf_filled(ws, f"L{R1}:L{R2}", f"L{R1}")

    s = R2 + 2
    band(ws, s, "Resumo automático", "L")

    def row(k, text, formula, fmt=None, sz=10, b=True, height=21.75):
        label(ws, f"B{s+k}", text, merge=f"B{s+k}:E{s+k}", h="right")
        calc(ws, f"F{s+k}", formula, fmt=fmt, merge=f"F{s+k}:K{s+k}", sz=sz, b=b)
        ws.row_dimensions[s + k].height = height

    row(1, "Problemas listados", f"=COUNTA(C{R1}:C{R2})")
    row(2, "Problemas com as três notas", f"=COUNT(I{R1}:I{R2})")
    row(3, "Prioridade alta", f'=COUNTIF(K{R1}:K{R2},"Alta")')
    row(4, "Prioridade média", f'=COUNTIF(K{R1}:K{R2},"Média")')
    row(5, "Prioridade baixa", f'=COUNTIF(K{R1}:K{R2},"Baixa")')
    row(6, "Alertas de gravidade máxima", f"=COUNTIF(F{R1}:F{R2},5)")
    row(7, "Aviso: notas", f'=IF(F{s+1}=0,"Liste os problemas",IF(F{s+2}<F{s+1},"Há problema sem as três notas","OK"))', sz=9, b=False)
    row(8, "Aviso: distribuição",
        f'=IF(F{s+2}<4,"OK",IF(F{s+3}>F{s+2}/2,"Mais da metade em prioridade alta: revise as notas",'
        f'IF(MAX(I{R1}:I{R2})=MIN(I{R1}:I{R2}),"Todas as pontuações são iguais: revise as notas","OK")))', sz=9, b=False)
    row(9, "Campos do cabeçalho preenchidos", "=COUNTA(D4,I4,D5,I5,D6,I6,D7,D8,I8,D9)", fmt='0" de 10"')
    cf_ok(ws, f"F{s+7}:K{s+8}")
    end = s + 9

    if is_ex:
        band(ws, end + 2, "Decisões tomadas", "L")
        r = end + 3
        head(ws, f"B{r}", "#", bg=GRAY, fg=INK)
        put(ws, f"C{r}", "Problema", f=font(10, True), bg=GRAY, h="center", merge=f"C{r}:D{r}")
        put(ws, f"E{r}", "Decisão", f=font(10, True), bg=GRAY, h="center", merge=f"E{r}:I{r}")
        put(ws, f"J{r}", "Responsável", f=font(10, True), bg=GRAY, h="center", merge=f"J{r}:K{r}")
        put(ws, f"L{r}", "Prazo", f=font(10, True), bg=GRAY, h="center")
        ws.row_dimensions[r].height = 21.75
        by = {p["k"]: p for p in data["itens"]}
        for code, text, who, when in data["decisoes"]:
            r += 1
            put(ws, f"B{r}", code, f=font(10, True, c=MUTED), bg=GRAY, h="center")
            put(ws, f"C{r}", by[code]["curto"], merge=f"C{r}:D{r}")
            put(ws, f"E{r}", text, merge=f"E{r}:I{r}")
            put(ws, f"J{r}", who, merge=f"J{r}:K{r}")
            put(ws, f"L{r}", when, h="center", fmt=DATE)
            ws.row_dimensions[r].height = 33
        end = r
    setup(ws, tab, f"B1:L{end}", landscape=False, fit_height=True)


# ================================================================== pasta de trabalho
wb = Workbook()
wb.properties.title = "GUT — Modelo de matriz de priorização"
wb.properties.creator = "Modelo GUT"
wb.properties.language = "pt-BR"

# ------------------------------------------------------------------ Instruções
ws = wb.active
ws.title = "Instruções"
widths(ws, {"A": 2, "B": 24, "C": 92, "D": 2})
title(ws, "GUT — Como usar esta planilha",
      "Modelo para priorizar problemas por gravidade, urgência e tendência.", "C")
r = 4


def section(text):
    global r
    band(ws, r, text, "C", sz=11)
    r += 1


def line(k, v, kbg=GRAY, vbg=None, kf=None, kh="left", height=19.5):
    global r
    put(ws, f"B{r}", k, f=kf or font(10, True), bg=kbg, h=kh)
    put(ws, f"C{r}", v, bg=vbg)
    ws.row_dimensions[r].height = height
    r += 1


section("Legenda: onde preencher")
line("Amarelo-claro", "Células de entrada. É aqui que você digita.", vbg=INPUT)
line("Cinza", "Células calculadas ou fixas (rótulos, contagens, avisos). Não altere.", vbg=GRAY)
line("Cinza em itálico", "Linha de exemplo (marcada com “Ex.”). Mostra o formato esperado e não entra nas contagens.", vbg=GRAY)
line("Listas suspensas", "Colunas de origem, código e status aceitam apenas as opções da lista. As notas aceitam números inteiros de 1 a 5.", height=31.5)
line("Comentários", "Células com triângulo vermelho no canto trazem uma dica de preenchimento. Passe o mouse sobre elas.", height=31.5)
r += 1
section("Os três critérios")
for q, text in [("G", "Gravidade: tamanho do dano que o problema causa ou pode causar, se nada for feito."),
                ("U", "Urgência: tempo disponível para agir antes que o dano aconteça ou aumente."),
                ("T", "Tendência: evolução esperada do problema com o passar do tempo, se nada for feito.")]:
    line(q, text, kbg=CR[q][0], vbg=CR[q][1], kf=font(10, True, c=WHITE), kh="center")
line("Pontos", "Gravidade × urgência × tendência, de 1 a 125. Quanto maior a pontuação, mais cedo o problema é tratado.")
r += 1
section("Ordem recomendada de preenchimento")
for k, text in enumerate([
    "Aba GUT: preencha o cabeçalho, com o tema, o objetivo e a origem da lista.",
    "Aba GUT: liste os problemas, cada um como resultado indesejado. Retire causas e soluções.",
    "Aba Escala: leia a escala com o grupo e ajuste os textos à realidade da lista.",
    "Aba GUT: dê as notas de gravidade da lista toda. Depois as de urgência. Depois as de tendência.",
    "Aba Priorização: veja a ordem de tratamento e o gráfico. Nada para digitar.",
    "Aba GUT: avalie os alertas de gravidade máxima, qualquer que seja a pontuação.",
    "Aba Plano de ação: defina ação, responsável e prazo para os problemas escolhidos.",
    "Aba Checklist: valide a matriz com o grupo e marque a data da próxima revisão.",
], 1):
    line(f"Passo {k}", text)
r += 1
section("Abas da planilha")
for k, text in [
    ("GUT", "Modelo principal. Cabeçalho, lista de até 20 problemas com as três notas e resumo automático."),
    ("Priorização", "Problemas em ordem de tratamento, com gráfico das pontuações."),
    ("Escala", "Descrição das notas de 1 a 5 e cortes das faixas de prioridade. Os dois podem ser ajustados."),
    ("Plano de ação", "Ações para os problemas escolhidos, com situação calculada a partir do prazo e do status."),
    ("Checklist", "Doze verificações de qualidade da matriz, com percentual de conclusão."),
    ("Exemplo 1 - Pizzaria", "Matriz preenchida com oito problemas de uma operação, incluindo um alerta de gravidade máxima."),
    ("Exemplo 2 - Compras", "Matriz preenchida com sete problemas de uma área, incluindo um empate."),
]:
    line(k, text)
r += 1
section("Regras e premissas do modelo")
for k, text in [
    ("Códigos", "Cada problema tem um código fixo, de P1 a P20, que corresponde à linha da lista. A aba Plano de ação usa o código para buscar o texto e a pontuação do problema."),
    ("Faixas de prioridade", f"O modelo parte de dois cortes: {ALTA} pontos para a faixa alta, que equivale a nota 4 nos três critérios, e {MEDIA} pontos para a faixa média, que equivale a nota 3 nos três. Os cortes podem ser alterados na aba Escala."),
    ("Empates", "Problemas com a mesma pontuação são ordenados pela maior gravidade, depois pela maior urgência e depois pela maior tendência. Se o empate continuar, vale a ordem da lista."),
    ("Regra de exceção", "Todo problema com gravidade 5 recebe um alerta e deve ser avaliado pelo grupo, qualquer que seja a pontuação."),
    ("Ação atrasada", "Uma ação é considerada atrasada quando o prazo é anterior à data de hoje e o status não é Concluída nem Cancelada."),
    ("Capacidade", "A planilha comporta 20 problemas e 15 ações. Se faltar espaço, agrupe problemas parecidos ou divida a lista por área."),
    ("Exemplos", "Os dados das abas de exemplo são ilustrativos, criados para este material."),
]:
    line(k, text, height=45.75 if len(text) > 100 else 31.5)
setup(ws, INK, f"B1:C{r}", landscape=False, fit_height=True)

# ------------------------------------------------------------------ GUT
gut_sheet(wb.create_sheet("GUT"), CR["G"][0])

# ------------------------------------------------------------------ Priorização
ws = wb.create_sheet("Priorização")
widths(ws, {"A": 2, "B": 6, "C": 9, "D": 52, "E": 16, "F": 8, "G": 8, "H": 8, "I": 10, "J": 13, "K": 22, "L": 2})
title(ws, "Priorização — ordem de tratamento", "Esta aba é calculada a partir da aba GUT. Não há nada para digitar aqui.", "K")
band(ws, 4, "Problemas em ordem de pontuação", "K")
put(ws, "B5", "Pos.", f=font(10, True), bg=GRAY, h="center")
put(ws, "C5", "Código", f=font(10, True), bg=GRAY, h="center")
put(ws, "D5", "Problema", f=font(10, True), bg=GRAY, h="center")
put(ws, "E5", "Origem", f=font(10, True), bg=GRAY, h="center")
for col, q in zip("FGH", "GUT"):
    put(ws, f"{col}5", q, f=font(10, True, c=WHITE), bg=CR[q][0], h="center")
for col, text in zip("IJK", ["Pontos", "Prioridade", "Alerta"]):
    put(ws, f"{col}5", text, f=font(10, True), bg=GRAY, h="center")
ws.row_dimensions[5].height = 21.75
Q1, Q2 = 6, 6 + N - 1
pos = f"GUT!$J${R1}:$J${R2}"
for k in range(N):
    rr = Q1 + k
    num(ws, f"B{rr}", k + 1)
    for col, src, txt, h in (("C", "B", True, "center"), ("D", "C", True, "left"), ("E", "E", True, "center"),
                             ("F", "F", False, "center"), ("G", "G", False, "center"), ("H", "H", False, "center"),
                             ("I", "I", False, "center"), ("J", "K", True, "center"), ("K", "L", True, "center")):
        f = f"INDEX(GUT!${src}${R1}:${src}${R2},MATCH($B{rr},{pos},0))"
        calc(ws, f"{col}{rr}", f'=IFERROR({f}{"&" + chr(34) * 2 if txt else ""},"")', h=h, b=(col in "CI"), sz=9 if col == "K" else 10)
    ws.row_dimensions[rr].height = 24
cf_equal(ws, f"J{Q1}:J{Q2}", PRIO_CF)
cf_filled(ws, f"K{Q1}:K{Q2}", f"K{Q1}")
for col, q in zip("FGH", "GUT"):
    ws.conditional_formatting.add(f"{col}{Q1}:{col}{Q2}", FormulaRule(formula=[f'LEN($C{Q1})>0'],
                                  fill=PatternFill("solid", bgColor=CR[q][1], fgColor=CR[q][1])))
g = Q2 + 2
band(ws, g, "Gráfico: pontuação em ordem de tratamento", "K")
ch = BarChart()
ch.type = "col"
ch.height, ch.width = 8.5, 26
ch.style = 2
ch.title = None
ch.legend = None
ch.gapWidth = 80
ch.add_data(Reference(ws, min_col=9, min_row=Q1 - 1, max_row=Q2), titles_from_data=True)
ch.set_categories(Reference(ws, min_col=3, min_row=Q1, max_row=Q2))
ch.series[0].graphicalProperties.solidFill = BLUE
ch.series[0].graphicalProperties.line.noFill = True
ch.x_axis.delete = False
ch.y_axis.delete = False
ch.y_axis.scaling.min = 0
ch.y_axis.scaling.max = 125
ch.y_axis.majorUnit = 25
ch.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill="D3DBE0", w=9525))
ch.x_axis.graphicalProperties = GraphicalProperties(ln=LineProperties(solidFill="D3DBE0", w=9525))
ch.y_axis.graphicalProperties = GraphicalProperties(ln=LineProperties(noFill=True))
ws.add_chart(ch, f"B{g+1}")
for rr in range(g + 1, g + 19):
    ws.row_dimensions[rr].height = 15
setup(ws, CR["U"][0], f"B1:K{g+18}", landscape=False, fit_height=True)

# ------------------------------------------------------------------ Escala
ws = wb.create_sheet("Escala")
widths(ws, {"A": 2, "B": 12, "C": 36, "D": 36, "E": 36, "F": 2})
title(ws, "Escala das notas e faixas de prioridade",
      "Leia a escala com o grupo antes de pontuar. Os textos e os cortes em amarelo-claro podem ser ajustados.", "E")
band(ws, 4, "Escala das notas", "E")
put(ws, "B5", "Nota", f=font(10, True), bg=GRAY, h="center")
for col, q in zip("CDE", "GUT"):
    put(ws, f"{col}5", f"{q} — {CR[q][2]}", f=font(10, True, c=WHITE), bg=CR[q][0], h="center")
ws.row_dimensions[5].height = 21.75
SCALE = [
    ("Sem gravidade: dano desprezível.", "Pode esperar: não há prazo definido.", "Estável: não vai piorar."),
    ("Pouco grave: dano pequeno.", "Pouco urgente: agir em alguns meses.", "Piora lenta: a longo prazo."),
    ("Grave: dano considerável.", "Merece atenção: agir em algumas semanas.", "Piora moderada: a médio prazo."),
    ("Muito grave: dano grande.", "Muito urgente: agir em alguns dias.", "Piora rápida: a curto prazo."),
    ("Extremamente grave: dano difícil de reverter, risco à vida ou descumprimento legal.", "Ação imediata: agir hoje.", "Piora imediata, se nada for feito."),
]
for k, texts in enumerate(SCALE):
    rr = 6 + k
    put(ws, f"B{rr}", k + 1, f=font(12, True), bg=GRAY, h="center")
    for col, t in zip("CDE", texts):
        inp(ws, f"{col}{rr}", t)
    ws.row_dimensions[rr].height = 48 if k == 4 else 36
band(ws, 12, "Faixas de prioridade", "E")
for col, text in zip("BCDE", ["Faixa", "A partir de (pontos)", "Leitura", "Encaminhamento"]):
    put(ws, f"{col}13", text, f=font(10, True), bg=GRAY, h="center")
ws.row_dimensions[13].height = 21.75
for rr, name, val, read, act in [
    (14, "Alta", ALTA, "Pesa muito nos três critérios.", "Tratar agora, com responsável, prazo e acompanhamento semanal."),
    (15, "Média", MEDIA, "Pesa em pelo menos dois critérios.", "Planejar o tratamento, com data de início definida."),
    (16, "Baixa", 1, "Pesa pouco ou em um critério só.", "Registrar e reavaliar na próxima revisão da matriz."),
]:
    put(ws, f"B{rr}", name, f=font(10, True), bg=GRAY, h="center")
    if rr < 16:
        inp(ws, f"C{rr}", val, h="center")
    else:
        calc(ws, f"C{rr}", val, b=False)
    put(ws, f"D{rr}", read)
    put(ws, f"E{rr}", act)
    ws.row_dimensions[rr].height = 31.5
dv = DataValidation(type="whole", operator="between", formula1="2", formula2="125", allow_blank=False)
dv.errorTitle, dv.error = "Valor inválido", "Digite um número inteiro de 2 a 125."
dv.showErrorMessage = True
ws.add_data_validation(dv)
dv.add("C14:C15")
cf_equal(ws, "B14:B16", PRIO_CF)
label(ws, "B17", "Aviso", h="center")
calc(ws, "C17", '=IF(OR(C14="",C15=""),"Informe os dois cortes",IF(C15>=C14,"O corte da faixa média deve ser menor que o da alta","OK"))',
     b=False, sz=9, merge="C17:E17")
cf_ok(ws, "C17:E17")
ws.row_dimensions[17].height = 21.75
band(ws, 19, "Regras da priorização", "E")
for rr, k, text in [
    (20, "Exceção", "Todo problema com gravidade 5 é avaliado pelo grupo, qualquer que seja a pontuação."),
    (21, "Desempate", "Na mesma pontuação, vem primeiro a maior gravidade, depois a maior urgência, depois a maior tendência."),
    (22, "Método", "Pontue um critério por vez, para a lista toda. Registre a justificativa das notas 4 e 5."),
]:
    label(ws, f"B{rr}", k)
    put(ws, f"C{rr}", text, merge=f"C{rr}:E{rr}")
    ws.row_dimensions[rr].height = 21.75
setup(ws, CR["T"][0], "B1:E22", fit_height=True)

# ------------------------------------------------------------------ Plano de ação
ws = wb.create_sheet("Plano de ação")
widths(ws, {"A": 2, "B": 5, "C": 10, "D": 40, "E": 10, "F": 38, "G": 18, "H": 13, "I": 15, "J": 14, "K": 2})
title(ws, "Plano de ação", "Escolha o código do problema. O texto e a pontuação vêm da aba GUT.", "J")
for col, text in zip("BCDEFGHIJ", ["#", "Código", "Problema", "Pontos", "Ação", "Responsável", "Prazo", "Status", "Situação"]):
    head(ws, f"{col}4", text)
ws.row_dimensions[4].height = 21.75
ex(ws, "B5", "Ex.", h="center")
ex(ws, "C5", "P3", h="center")
ex(ws, "D5", "Câmara fria com falha intermitente de temperatura")
ex(ws, "E5", 100, h="center")
ex(ws, "F5", "Chamar a assistência técnica e registrar a temperatura duas vezes por turno")
ex(ws, "G5", "Gerente da loja")
ex(ws, "H5", date(2026, 10, 30), h="center", fmt=DATE)
ex(ws, "I5", "Em andamento", h="center")
ex(ws, "J5", "No prazo", h="center")
ws.row_dimensions[5].height = 31.5
A1, A2 = 6, 20
idx = lambda rr: f"VALUE(MID(C{rr},2,2))"
for k in range(15):
    rr = A1 + k
    num(ws, f"B{rr}", k + 1)
    inp(ws, f"C{rr}", h="center")
    calc(ws, f"D{rr}", f'=IF(C{rr}="","",INDEX(GUT!$C${R1}:$C${R2},{idx(rr)})&"")', h="left", b=False)
    calc(ws, f"E{rr}", f'=IF(C{rr}="","",IFERROR(INDEX(GUT!$I${R1}:$I${R2},{idx(rr)})+0,""))')
    inp(ws, f"F{rr}")
    inp(ws, f"G{rr}")
    inp(ws, f"H{rr}", h="center", fmt=DATE)
    inp(ws, f"I{rr}", h="center")
    calc(ws, f"J{rr}", f'=IF(F{rr}="","",IF(OR(I{rr}="Concluída",I{rr}="Cancelada"),I{rr},'
         f'IF(H{rr}="","Sem prazo",IF(H{rr}<TODAY(),"Atrasada","No prazo"))))', b=False)
    ws.row_dimensions[rr].height = 30
dv_list(ws, f"C{A1}:C{A2}", [f"P{i}" for i in range(1, N + 1)], "Código do problema na aba GUT")
dv_list(ws, f"I{A1}:I{A2}", ["Não iniciada", "Em andamento", "Concluída", "Cancelada"], "Não iniciada, Em andamento, Concluída ou Cancelada")
dv_date(ws, f"H{A1}:H{A2}")
cf_equal(ws, f"I{A1}:I{A2}", STATUS_CF)
cf_equal(ws, f"J{A1}:J{A2}", [("Concluída", GREEN), ("No prazo", GREEN), ("Sem prazo", YELLOW), ("Atrasada", RED)])
band(ws, 22, "Resumo automático", "J")
summary(ws, 23, "Ações registradas", f"=COUNTA(F{A1}:F{A2})", "F", "G", val_merge="G23:H23")
summary(ws, 24, "Concluídas", f'=COUNTIF(I{A1}:I{A2},"Concluída")', "F", "G", val_merge="G24:H24")
summary(ws, 25, "Em andamento", f'=COUNTIF(I{A1}:I{A2},"Em andamento")', "F", "G", val_merge="G25:H25")
summary(ws, 26, "Não iniciadas", f'=COUNTIF(I{A1}:I{A2},"Não iniciada")', "F", "G", val_merge="G26:H26")
summary(ws, 27, "Atrasadas", f'=COUNTIF(J{A1}:J{A2},"Atrasada")', "F", "G", val_merge="G27:H27")
summary(ws, 28, "Ações sem responsável ou sem prazo", f'=SUMPRODUCT((F{A1}:F{A2}<>"")*((G{A1}:G{A2}="")+(H{A1}:H{A2}="")>0))', "F", "G", val_merge="G28:H28")
summary(ws, 29, "Percentual concluído", "=IF(G23=0,0,G24/G23)", "F", "G", fmt="0%", val_merge="G29:H29")
summary(ws, 30, "Problemas de prioridade alta sem ação",
        f'=SUMPRODUCT((GUT!$K${R1}:$K${R2}="Alta")*(COUNTIF($C${A1}:$C${A2},GUT!$B${R1}:$B${R2})=0))', "F", "G", val_merge="G30:H30")
summary(ws, 31, "Alertas de gravidade máxima sem ação",
        f'=SUMPRODUCT((GUT!$F${R1}:$F${R2}=5)*(COUNTIF($C${A1}:$C${A2},GUT!$B${R1}:$B${R2})=0))', "F", "G", val_merge="G31:H31")
ws.freeze_panes = "C5"
setup(ws, BLUE, "B1:J31", fit_height=True)

# ------------------------------------------------------------------ Checklist
ws = wb.create_sheet("Checklist")
widths(ws, {"A": 2, "B": 5, "C": 70, "D": 14, "E": 44, "F": 2})
title(ws, "Checklist de validação da Matriz GUT", "Escolha o status de cada item na lista suspensa (coluna em amarelo-claro).", "E")
for col, text in zip("BCDE", ["#", "Item de verificação", "Status", "Observações"]):
    head(ws, f"{col}4", text)
ws.row_dimensions[4].height = 21.75
for k, text in enumerate([
    "A lista tem problemas, e não causas nem soluções.",
    "Cada problema está descrito com fato, local e número, quando houver.",
    "Os problemas da lista são comparáveis: mesmo tipo e mesmo tamanho.",
    "A escala de notas foi combinada pelo grupo antes da pontuação.",
    "As notas foram dadas em grupo, por pessoas que conhecem os problemas.",
    "Cada critério foi pontuado para a lista toda antes de passar ao seguinte.",
    "As notas se apoiam em dados, quando eles existem.",
    "Os empates foram resolvidos por um critério combinado.",
    "Os problemas de gravidade máxima foram avaliados, qualquer que seja a pontuação.",
    "Os problemas de maior prioridade têm ação, responsável e prazo.",
    "Os problemas de baixa prioridade foram registrados, e não descartados.",
    "A data da próxima revisão da matriz está marcada.",
]):
    rr = 5 + k
    num(ws, f"B{rr}", k + 1)
    put(ws, f"C{rr}", text, bg=GRAY)
    inp(ws, f"D{rr}", h="center")
    inp(ws, f"E{rr}")
    ws.row_dimensions[rr].height = 30
dv_list(ws, "D5:D16", YESNO, "Sim, Parcial ou Não")
cf_equal(ws, "D5:D16", YESNO_CF)
band(ws, 18, "Resumo automático", "E")
summary(ws, 19, "Itens atendidos (Sim)", '=COUNTIF(D5:D16,"Sim")', "C", "D")
summary(ws, 20, "Itens parciais", '=COUNTIF(D5:D16,"Parcial")', "C", "D")
summary(ws, 21, "Itens não atendidos", '=COUNTIF(D5:D16,"Não")', "C", "D")
summary(ws, 22, "Itens sem resposta", "=COUNTBLANK(D5:D16)", "C", "D")
summary(ws, 23, "Percentual atendido", "=D19/ROWS(D5:D16)", "C", "D", fmt="0%")
summary(ws, 24, "Situação", '=IF(D19=ROWS(D5:D16),"Validado",IF(D22>0,"Em andamento","Com pendências"))', "C", "D")
setup(ws, "B0413E", "B1:E24", landscape=False, fit_height=True)

# ------------------------------------------------------------------ Exemplos
gut_sheet(wb.create_sheet("Exemplo 1 - Pizzaria"), MUTED, EX1)
gut_sheet(wb.create_sheet("Exemplo 2 - Compras"), MUTED, EX2)

wb.active = 1
wb.save(OUT)
print("ok", OUT)
