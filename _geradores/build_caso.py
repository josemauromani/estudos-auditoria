# -*- coding: utf-8 -*-
"""Gera Caso-Integrado-modelo.xlsx no mesmo padrão visual das planilhas anteriores da série."""
import math
import os
import sys

_here = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, _here)
# funções de formatação compartilhadas com o gerador do PDCA (bloco inicial do arquivo)
_src = open(os.path.join(_here, "build_pdca.py"), encoding="utf-8").read()
exec(_src[: _src.index("STATUS = [")])

from openpyxl.chart import BarChart, Series  # noqa: E402
from openpyxl.chart.label import DataLabelList  # noqa: E402
from openpyxl.utils import get_column_letter as L  # noqa: E402
from caso_data import ATRASO, CAL_SITS, CHECK, EMDIA, ESTUDOS, EX1, EX2, FREQS, MESES, NAOCOMECOU  # noqa: E402

BLUE, TEAL, AMBER, PURPLE, REDC = "2B5C8A", "1E7B73", "A96A12", "7A4A9A", "B0413E"
S1_T = "DCE8F3"
YESNO = ["Sim", "Parcial", "Não"]
YESNO_CF = [("Sim", GREEN), ("Parcial", YELLOW), ("Não", RED)]
NR, NC = 30, 40
R1, R2 = 8, 8 + NR - 1          # passos do rastro
C1, C2 = 10, 10 + NC - 1        # atividades do calendário
E1, E2 = 5, 5 + len(ESTUDOS) - 1
ESTS = f"Estudos!$C${E1}:$C${E2}"
M0 = 9                          # coluna de janeiro (I)
MC = [L(M0 + k) for k in range(12)]   # I..T
CU, CV, CW, CX, CY, CZ, CS = "U", "V", "W", "X", "Y", "Z", "AA"   # previstas, feitas, atrasadas, fora, situação, conferência, passo
SIT_CF = [(EMDIA, GREEN), (ATRASO, RED), (NAOCOMECOU, GRAY)]
FREQ_LIST = list(FREQS)
OUTRO = "O passo seguinte é de outro estudo"


def note(ws, ref, text):
    c = Comment(text, "Modelo Caso integrado")
    c.width, c.height = 300, 120
    ws[ref].comment = c


def cf_warn(ws, rng_, first, ok_values=("OK",)):
    cond = ",".join(f'{first}="{v}"' for v in ok_values)
    ws.conditional_formatting.add(rng_, FormulaRule(formula=[f"OR({cond})"], fill=PatternFill("solid", bgColor=GREEN, fgColor=GREEN)))
    ws.conditional_formatting.add(rng_, FormulaRule(formula=[f"AND(LEN({first})>0,NOT(OR({cond})))"],
                                  fill=PatternFill("solid", bgColor=YELLOW, fgColor=YELLOW)))


def cf_texto(ws, rng_, first, pares, resto=None):
    for t, cor in pares:
        ws.conditional_formatting.add(rng_, FormulaRule(formula=[f'{first}="{t}"'], stopIfTrue=True, fill=PatternFill("solid", bgColor=cor, fgColor=cor)))
    if resto:
        ws.conditional_formatting.add(rng_, FormulaRule(formula=[f"LEN({first})>0"], fill=PatternFill("solid", bgColor=resto, fgColor=resto)))


def hint_row(ws, row, cells, height=30):
    for ref, text in cells:
        put(ws, f"{ref}{row}", text, f=font(9, i=True, c=MUTED), bg=GRAY, h="center")
    ws.row_dimensions[row].height = height


def alt(pares, minimo=21.75, linha=12):
    n = max(max(1, math.ceil(len(str(t or "")) / max(c, 1))) for t, c in pares)
    return max(minimo, n * linha + 7)


def lista_ref(ws, rng, ref, prompt):
    dv = DataValidation(type="list", formula1=f"={ref}", allow_blank=True)
    dv.promptTitle, dv.prompt = "Opções", prompt
    dv.showInputMessage = True
    ws.add_data_validation(dv)
    dv.add(rng)


# ------------------------------------------------------------------ fórmulas do rastro
def f_dias(r, primeira):
    return '""' if r == primeira else f'IF(OR(C{r}="",C{r - 1}=""),"",C{r}-C{r - 1})'


def f_rasconf(r, conf_col="J"):
    vazio, prox_vazio = f"COUNTA(C{r}:H{r})=0", f"COUNTA(C{r + 1}:H{r + 1})=0"
    return (f'IF({vazio},"",IF(C{r}="","Falta a data",IF(D{r}="","Falta o estudo",IF(E{r}="","Falta o que aconteceu",IF(F{r}="","Falta o registro",'
            f'IF(AND(ISNUMBER(I{r}),I{r}<0),"Data antes do passo anterior",IF({prox_vazio},"OK",IF(G{r}="","Falta a saída",IF(H{r}="","Falta o próximo estudo",'
            f'IF(H{r}<>D{r + 1},"{OUTRO}","OK"))))))))))')


# ------------------------------------------------------------------ fórmulas do calendário (a linha dos números dos meses fica em mrow)
def f_passo(r):
    return f'IF(G{r}="","",IFERROR(CHOOSE(MATCH(G{r},{{{",".join(chr(34) + f + chr(34) for f in FREQ_LIST)}}},0),{",".join(str(v) for v in FREQS.values())}),""))'


def planejado(r, mrow):
    return f"(${MC[0]}${mrow}:${MC[-1]}${mrow}>=$H{r})*(MOD(${MC[0]}${mrow}:${MC[-1]}${mrow}-$H{r},${CS}{r})=0)"


def f_prev(r, mrow, ref):
    return (f'IF(OR(C{r}="",{CS}{r}="",H{r}="",{ref}=""),"",IF(OR(H{r}<1,H{r}>12),"",'
            f'SUMPRODUCT({planejado(r, mrow)}*(${MC[0]}${mrow}:${MC[-1]}${mrow}<={ref}))))')


def f_feitas(r, mrow, ref):
    return f'IF({CU}{r}="","",SUMPRODUCT({planejado(r, mrow)}*(${MC[0]}${mrow}:${MC[-1]}${mrow}<={ref})*({MC[0]}{r}:{MC[-1]}{r}<>"")))'


def f_fora(r, mrow):
    return f'IF({CU}{r}="","",SUMPRODUCT((1-{planejado(r, mrow)})*({MC[0]}{r}:{MC[-1]}{r}<>"")))'


def f_calsit(r):
    return f'IF({CU}{r}="","",IF({CU}{r}=0,"{NAOCOMECOU}",IF({CW}{r}>0,"{ATRASO}","{EMDIA}")))'


def f_calconf(r):
    return (f'IF(C{r}="","",IF(D{r}="","Falta o estudo",IF(F{r}="","Falta o responsável",IF(G{r}="","Falta a frequência",IF(H{r}="","Falta o mês de início",'
            f'IF(OR(H{r}<1,H{r}>12),"Mês de início fora de 1 a 12","OK"))))))')


def cf_meses(ws, r1, r2, mrow, ref):
    """Cores no próprio calendário: feita, atrasada, prevista e feita fora do plano."""
    c0 = MC[0]
    plan = f"AND(ISNUMBER(${CS}{r1}),ISNUMBER($H{r1}),{c0}${mrow}>=$H{r1},MOD({c0}${mrow}-$H{r1},${CS}{r1})=0)"
    rng = f"{c0}{r1}:{MC[-1]}{r2}"
    for formula, cor in ((f'AND({plan},{c0}{r1}<>"")', S1_T), (f'AND({plan},{c0}{r1}="",ISNUMBER({ref}),{c0}${mrow}<={ref})', RED),
                         (f"{plan}", GRAY), (f'AND(NOT({plan}),{c0}{r1}<>"")', YELLOW)):
        ws.conditional_formatting.add(rng, FormulaRule(formula=[formula], stopIfTrue=True, fill=PatternFill("solid", bgColor=cor, fgColor=cor)))


def chart_style(ch):
    ch.style = 2
    ch.title = None
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ch.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill="D3DBE0", w=9525))
    ch.x_axis.graphicalProperties = GraphicalProperties(ln=LineProperties(solidFill="D3DBE0", w=9525))
    ch.y_axis.graphicalProperties = GraphicalProperties(ln=LineProperties(noFill=True))


def colunas(ws, anchor, r1, r2, c_cat, c_val, width=16, height=7.5):
    ch = BarChart()
    ch.type = "col"
    ch.gapWidth = 70
    ch.height, ch.width = height, width
    chart_style(ch)
    ch.legend = None
    s = Series(Reference(ws, min_col=c_val, min_row=r1, max_row=r2), title="Atividades")
    s.graphicalProperties.solidFill = BLUE
    s.graphicalProperties.line.noFill = True
    s.dLbls = DataLabelList()
    s.dLbls.showVal = True
    s.dLbls.showSerName = s.dLbls.showCatName = s.dLbls.showLegendKey = False
    ch.append(s)
    ch.set_categories(Reference(ws, min_col=c_cat, min_row=r1, max_row=r2))
    ch.y_axis.scaling.min = 0
    ch.y_axis.majorUnit = 1
    ws.add_chart(ch, anchor)


def resumo_aba(ws, s, itens, last, merge_lab="D"):
    band(ws, s, "Resumo automático", last)
    col = chr(ord(merge_lab) + 1)
    for k, (text, formula) in enumerate(itens, 1):
        aviso = text == "Aviso"
        label(ws, f"B{s+k}", text, merge=f"B{s+k}:{merge_lab}{s+k}", h="right")
        calc(ws, f"{col}{s+k}", formula, merge=f"{col}{s+k}:{last}{s+k}", sz=9 if aviso else 10, b=not aviso, h="left")
        ws.row_dimensions[s + k].height = 21.75
    av = s + len(itens)
    cf_warn(ws, f"{col}{av}:{last}{av}", f"{col}{av}")
    return av


# ================================================================== pasta de trabalho
wb = Workbook()
wb.properties.title = "Caso integrado — Modelo"
wb.properties.creator = "Modelo Caso integrado"
wb.properties.language = "pt-BR"

# ------------------------------------------------------------------ Instruções
ws = wb.active
ws.title = "Instruções"
widths(ws, {"A": 2, "B": 26, "C": 90, "D": 2})
title(ws, "Caso integrado — Como usar esta planilha", "Modelo para seguir o fio de um problema pelos registros e para manter o calendário anual do sistema.", "C")
r = 4


def section(text):
    global r
    band(ws, r, text, "C", sz=11)
    r += 1


def line(k, v, kbg=GRAY, vbg=None, kf=None, kh="left", height=None):
    global r
    put(ws, f"B{r}", k, f=kf or font(10, True), bg=kbg, h=kh)
    put(ws, f"C{r}", v, bg=vbg)
    ws.row_dimensions[r].height = height or (19.5 if len(v) <= 96 else (31.5 if len(v) <= 192 else 45.75))
    r += 1


section("Legenda: onde preencher")
line("Amarelo-claro", "Células de entrada. É aqui que você digita.", vbg=INPUT)
line("Cinza", "Células calculadas ou fixas (dias, previstas, feitas, atrasadas, situações, conferências). Não altere.", vbg=GRAY)
line("Cinza em itálico", "Linha de exemplo (marcada com “Ex.”). Mostra o formato esperado e não entra nas contagens.", vbg=GRAY)
line("Listas suspensas", "Estudo, próximo estudo, frequência e checklist aceitam apenas as opções da lista. Os estudos vêm da aba Estudos.")
line("Cores do calendário", "Azul-claro: feita. Vermelho: prevista e não feita até o mês da leitura. Cinza: prevista para depois. Amarelo: feita fora do plano.")
r += 1
section("As regras de cálculo")
for k, text in [
    ("Dias", "A data do passo menos a data do passo da linha de cima."),
    ("Conferência do rastro", "Cada passo, menos o último, precisa da saída e do próximo estudo, e o próximo estudo precisa ser o estudo da linha seguinte."),
    ("Meses previstos", "Do mês de início até dezembro, a cada 1, 2, 3, 6 ou 12 meses, conforme a frequência."),
    ("Previstas e feitas", "As vezes previstas até o mês da leitura, e quantas delas têm marca no mês."),
    (EMDIA, "Todas as previstas até o mês da leitura foram feitas."),
    (ATRASO, "Ao menos uma prevista até o mês da leitura não foi feita."),
    (NAOCOMECOU, "O primeiro mês previsto é depois do mês da leitura."),
    ("Fora do plano", "Marca num mês em que a atividade não estava prevista. Não entra nas feitas."),
]:
    line(k, text)
r += 1
section("Ordem recomendada de preenchimento")
for k, text in enumerate([
    "Aba Rastro: dê um nome ao fio e registre os passos do mais antigo ao mais recente, um por linha, sem linhas em branco no meio.",
    "Aba Rastro: para cada passo, o que saiu e o próximo estudo. Leia a coluna Conferência.",
    "Aba Calendário: informe o ano e o mês da leitura.",
    "Aba Calendário: liste as atividades periódicas, com estudo, requisito, responsável, frequência e mês de início.",
    "Aba Calendário: marque com X os meses em que cada atividade foi feita.",
    "Aba Painel: leia os avisos. Depois, valide na aba Checklist.",
], 1):
    line(f"Passo {k}", text)
r += 1
section("Abas da planilha")
for k, text in [
    ("Rastro", f"O nome do fio e até {NR} passos. Calcula os dias entre os passos e a conferência."),
    ("Calendário", f"O ano, o mês da leitura e até {NC} atividades. Calcula as previstas, as feitas, as atrasadas, a situação e a conferência."),
    ("Estudos", f"Os {len(ESTUDOS)} estudos da série, com o grupo e o ritmo recomendado. Alimenta as listas suspensas."),
    ("Painel", "Os números do fio e do calendário, o gráfico e o aviso."),
    ("Checklist", "Doze verificações do sistema integrado, com percentual de conclusão."),
    ("Exemplo 1 - Pizzaria", "Os atrasos das noites de pico, e o calendário de 2027 lido em junho."),
    ("Exemplo 2 - Indústria", "O filme fino da extrusora 3, e o calendário de 2027 lido em setembro."),
]:
    line(k, text)
r += 1
section("Regras e premissas do modelo")
for k, text in [
    ("Estudos", "A lista traz os estudos desta série. Para usar outro nome, acrescente uma linha no fim da aba Estudos."),
    ("Marca do mês", "Qualquer texto vale como feito. Use X, ou a data em que foi feito."),
    ("Frequências", "Mensal, bimestral, trimestral, semestral ou anual. O que é contínuo fica fora do calendário."),
    ("Exemplos", "Os dados das abas de exemplo são ilustrativos, criados para este material."),
]:
    line(k, text)
setup(ws, INK, f"B1:C{r}", landscape=False, fit_height=True)

# ------------------------------------------------------------------ Rastro
ws = wb.create_sheet("Rastro")
widths(ws, {"A": 2, "B": 7, "C": 12, "D": 26, "E": 44, "F": 24, "G": 30, "H": 26, "I": 8, "J": 32, "K": 2})
title(ws, "Rastro de um fio", "Um passo por linha, do mais antigo ao mais recente. Cada passo diz o que saiu e para qual estudo foi.", "J")
label(ws, "B4", "Fio", merge="B4:C4")
inp(ws, "D4", merge="D4:J4")
put(ws, "B5", "O problema que está sendo seguido. Ex.: reclamação de filme fino do cliente A.", f=font(9, i=True, c=MUTED), bg=GRAY, merge="B5:J5")
for col, text in zip("BCDEFGHIJ", ["Passo", "Data", "Estudo", "O que aconteceu", "Registro", "O que seguiu", "Próximo estudo", "Dias", "Conferência"]):
    head(ws, f"{col}6", text)
ws.row_dimensions[6].height = 33
hint_row(ws, 7, [("B", ""), ("C", "Data"), ("D", "Lista"), ("E", "A decisão ou a ação"), ("F", "RNC, ata, folha"), ("G", "A saída deste passo"), ("H", "Lista"),
                 ("I", "Calculado"), ("J", "Calculada")])
for k in range(NR):
    rr = R1 + k
    num(ws, f"B{rr}", k + 1)
    inp(ws, f"C{rr}", h="center", fmt=DATE)
    for col in "DEFGH":
        inp(ws, f"{col}{rr}")
    calc(ws, f"I{rr}", "=" + f_dias(rr, R1))
    calc(ws, f"J{rr}", "=" + f_rasconf(rr), b=False, sz=9)
    ws.row_dimensions[rr].height = 30
dv_date(ws, f"C{R1}:C{R2}")
lista_ref(ws, f"D{R1}:D{R2}", ESTS, "Estudo da série em que o passo aconteceu")
lista_ref(ws, f"H{R1}:H{R2}", ESTS, "Estudo que recebeu a saída deste passo")
cf_texto(ws, f"J{R1}:J{R2}", f"J{R1}", [("OK", GREEN), (OUTRO, RED)], resto=YELLOW)
note(ws, "F6", "O registro que prova o passo, com o número: RNC 2027-22, ata de 18/02/2027, folha de verificação. É ele que o auditor pede.")
note(ws, "H6", "O estudo que recebeu a saída. Precisa ser o estudo da linha seguinte; se não for, o fio pulou um passo.")
s = R2 + 2
RAV = resumo_aba(ws, s, [
    ("Passos", f'=COUNTA(D{R1}:D{R2})&" passos, "&IF(COUNTA(D{R1}:D{R2})=0,0,SUMPRODUCT((D{R1}:D{R2}<>"")/COUNTIF(D{R1}:D{R2},D{R1}:D{R2}&"")))&" estudos"'),
    ("Duração", f'=IF(COUNT(C{R1}:C{R2})=0,"",MAX(C{R1}:C{R2})-MIN(C{R1}:C{R2})&" dias do primeiro ao último passo; maior intervalo de "&MAX(I{R1}:I{R2})&" dias")'),
    ("Aviso", f'=IF(COUNTA(D{R1}:D{R2})=0,"Registre os passos do fio",IF(COUNTIF(J{R1}:J{R2},"{OUTRO}")>0,"O fio pulou um passo: veja a coluna Conferência",'
              f'IF(SUMPRODUCT((J{R1}:J{R2}<>"")*(J{R1}:J{R2}<>"OK"))>0,"Há passo a completar: veja a coluna Conferência","OK")))'),
], "J")
ws.freeze_panes = f"D{R1}"
setup(ws, BLUE, f"B1:J{RAV}")


# ------------------------------------------------------------------ Calendário
def cal_formulas(ws, rr, mrow, ref):
    calc(ws, f"{CS}{rr}", "=" + f_passo(rr), b=False, sz=9)
    calc(ws, f"{CU}{rr}", "=" + f_prev(rr, mrow, ref))
    calc(ws, f"{CV}{rr}", "=" + f_feitas(rr, mrow, ref))
    calc(ws, f"{CW}{rr}", f'=IF({CU}{rr}="","",{CU}{rr}-{CV}{rr})')
    calc(ws, f"{CX}{rr}", "=" + f_fora(rr, mrow), b=False)
    calc(ws, f"{CY}{rr}", "=" + f_calsit(rr), sz=9)
    calc(ws, f"{CZ}{rr}", "=" + f_calconf(rr), b=False, sz=9)


def cal_largura(ws):
    w = {"A": 2, "B": 6, "C": 34, "D": 24, "E": 12, "F": 22, "G": 12, "H": 9}
    w.update({c: 5.5 for c in MC})
    w.update({CU: 9, CV: 8, CW: 9, CX: 7, CY: 15, CZ: 26, CS: 7, "AB": 2})
    widths(ws, w)


ws = wb.create_sheet("Calendário")
cal_largura(ws)
title(ws, "Calendário do sistema", "Uma atividade periódica por linha. Marque com X os meses em que foi feita.", CZ)
for rr, text, dica in [(4, "Ano", "O ano do calendário."), (5, "Mês da leitura (1 a 12)", "As previstas e as atrasadas contam até este mês, inclusive.")]:
    label(ws, f"B{rr}", text, merge=f"B{rr}:D{rr}")
    inp(ws, f"E{rr}", h="center")
    put(ws, f"F{rr}", dica, f=font(9, i=True, c=MUTED), bg=GRAY, merge=f"F{rr}:{CZ}{rr}")
dv = DataValidation(type="whole", operator="between", formula1="1", formula2="12", allow_blank=True)
dv.errorTitle, dv.error = "Mês inválido", "Digite o número do mês, de 1 a 12."
dv.showErrorMessage = True
ws.add_data_validation(dv)
dv.add("E5")
REFM = "$E$5"
heads = ["#", "Atividade", "Estudo", "Requisito", "Responsável", "Frequência", "Mês de início"] + MESES + ["Previstas", "Feitas", "Atrasadas", "Fora do plano", "Situação",
                                                                                                    "Conferência", "A cada"]
for k, text in enumerate(heads):
    head(ws, f"{L(2 + k)}8", text)
ws.row_dimensions[8].height = 33
hint_row(ws, 9, [("B", ""), ("C", ""), ("D", "Lista"), ("E", "Ex.: 9.2"), ("F", ""), ("G", "Lista"), ("H", "1 a 12"), (CU, "Calculado"), (CV, "Calculado"),
                 (CW, "Calculado"), (CX, "Calculado"), (CY, "Calculada"), (CZ, "Calculada"), (CS, "Meses")])
for k, c in enumerate(MC):
    put(ws, f"{c}9", k + 1, f=font(9, i=True, c=MUTED), bg=GRAY, h="center")
for k in range(NC):
    rr = C1 + k
    num(ws, f"B{rr}", k + 1)
    for col in "CDEF":
        inp(ws, f"{col}{rr}")
    inp(ws, f"G{rr}", h="center")
    inp(ws, f"H{rr}", h="center")
    for c in MC:
        inp(ws, f"{c}{rr}", h="center")
    cal_formulas(ws, rr, 9, REFM)
    ws.row_dimensions[rr].height = 21.75
lista_ref(ws, f"D{C1}:D{C2}", ESTS, "Estudo da série a que a atividade pertence")
dv_list(ws, f"G{C1}:G{C2}", FREQ_LIST, "Mensal, bimestral, trimestral, semestral ou anual")
dv = DataValidation(type="whole", operator="between", formula1="1", formula2="12", allow_blank=True)
dv.errorTitle, dv.error = "Mês inválido", "Digite o número do mês, de 1 a 12."
dv.showErrorMessage = True
ws.add_data_validation(dv)
dv.add(f"H{C1}:H{C2}")
cf_meses(ws, C1, C2, 9, REFM)
cf_texto(ws, f"{CY}{C1}:{CY}{C2}", f"{CY}{C1}", SIT_CF)
cf_warn(ws, f"{CZ}{C1}:{CZ}{C2}", f"{CZ}{C1}")
note(ws, "H8", "O primeiro mês em que a atividade acontece no ano. Uma atividade semestral com início em 2 acontece em fevereiro e em agosto.")
note(ws, f"{MC[0]}8", "Marque com X, ou com a data, o mês em que a atividade foi feita. As cores mostram o que estava previsto.")
s = C2 + 2
Y_ = f"{CY}{C1}:{CY}{C2}"
CAV = resumo_aba(ws, s, [
    ("Atividades", f"=COUNTA(C{C1}:C{C2})"),
    ("Até o mês da leitura", f'=IF(SUM({CU}{C1}:{CU}{C2})=0,"",SUM({CV}{C1}:{CV}{C2})&" de "&SUM({CU}{C1}:{CU}{C2})&" feitas, "&SUM({CW}{C1}:{CW}{C2})&" atrasadas, "'
                             f'&SUM({CX}{C1}:{CX}{C2})&" fora do plano")'),
    ("Por situação", "=" + '&", "&'.join(f'COUNTIF({Y_},"{x}")&" {x.lower()}"' for x in CAL_SITS)),
    ("Aviso", f'=IF(E{s+1}=0,"Liste as atividades do calendário",IF(E5="","Informe o mês da leitura",IF(COUNTIF({Y_},"{ATRASO}")>0,"Há atividade atrasada: marque a data nova",'
              f'IF(SUMPRODUCT(({CZ}{C1}:{CZ}{C2}<>"")*({CZ}{C1}:{CZ}{C2}<>"OK"))>0,"Há atividade a completar: veja a coluna Conferência","OK"))))'),
], CZ)
ws.freeze_panes = f"D{C1}"
setup(ws, TEAL, f"B1:{CZ}{CAV}")

# ------------------------------------------------------------------ Estudos
ws = wb.create_sheet("Estudos")
widths(ws, {"A": 2, "B": 5, "C": 32, "D": 44, "E": 36, "F": 18, "G": 50, "H": 2})
title(ws, "Os estudos da série", "A lista usada nas listas suspensas, com o ritmo recomendado por este material.", "G")
for col, text in zip("BCDEFG", ["#", "Estudo", "Título", "Grupo", "Ritmo", "O que se faz"]):
    head(ws, f"{col}4", text)
for k, e in enumerate(ESTUDOS):
    rr = E1 + k
    num(ws, f"B{rr}", k + 1)
    put(ws, f"C{rr}", e[0], f=font(10, True), bg=GRAY)
    for col, v in zip("DEFG", (e[1], e[2], e[5], e[4])):
        put(ws, f"{col}{rr}", v, bg=GRAY, f=font(9))
    ws.row_dimensions[rr].height = 21.75
setup(ws, MUTED, f"B1:G{E2}")

# ------------------------------------------------------------------ Painel
ws = wb.create_sheet("Painel")
widths(ws, {"A": 2, "B": 38, "C": 14, "D": 14, "E": 14, "F": 34, "G": 2})
title(ws, "Painel do sistema integrado", "Nada a preencher: tudo vem das outras abas.", "F")
label(ws, "B3", "Fio")
calc(ws, "C3", '=IF(Rastro!D4="","",Rastro!D4)', merge="C3:F3", h="left")
for col, text, m in [("B", "Indicador", None), ("C", "Resultado", None), ("D", "Como se lê", "D5:F5")]:
    put(ws, f"{col}5", text, f=font(10, True, c=WHITE), bg=INK, h="center", merge=m)
ws.row_dimensions[5].height = 21.75
Rd, Rc, Ri, Rj = (f"Rastro!{c}{R1}:{c}{R2}" for c in "DCIJ")
Cal = lambda c: f"'Calendário'!{c}{C1}:{c}{C2}"  # noqa: E731
IND = [
    ("O FIO", None, None, None),
    ("Passos", f"=COUNTA({Rd})", "Da aba Rastro.", None),
    ("Estudos diferentes", f'=IF(COUNTA({Rd})=0,0,SUMPRODUCT(({Rd}<>"")/COUNTIF({Rd},{Rd}&"")))', "Por quantos estudos o problema passou.", None),
    ("Dias do primeiro ao último passo", f'=IF(COUNT({Rc})=0,"",MAX({Rc})-MIN({Rc}))', "A duração do fio.", None),
    ("Maior intervalo entre passos", f'=IF(COUNT({Ri})=0,"",MAX({Ri}))', "Onde o fio esperou mais.", None),
    ("Passos a completar", f'=SUMPRODUCT(({Rj}<>"")*({Rj}<>"OK"))', "Conferências diferentes de OK.", None),
    ("O CALENDÁRIO", None, None, None),
    ("Mês da leitura", "=IF('Calendário'!E5=\"\",\"\",'Calendário'!E5)", "Da aba Calendário.", None),
    ("Atividades", f"=COUNTA({Cal('C')})", "Atividades periódicas do sistema.", None),
    ("Previstas até o mês", f"=SUM({Cal(CU)})", "Vezes previstas até o mês da leitura.", None),
    ("Feitas", f"=SUM({Cal(CV)})", "Das previstas, as que têm marca.", None),
    ("Cumprimento", f'=IFERROR(SUM({Cal(CV)})/SUM({Cal(CU)}),"")', "Feitas ÷ previstas.", "0%"),
    ("Atrasadas", f"=SUM({Cal(CW)})", "Previstas e não feitas.", None),
    ("Atividades com atraso", f'=COUNTIF({Cal(CY)},"{ATRASO}")', "Levar à reunião mensal.", None),
    ("Feitas fora do plano", f"=SUM({Cal(CX)})", "Recuperação de atraso, ou plano a rever.", None),
    ("Atividades a completar", f'=SUMPRODUCT(({Cal(CZ)}<>"")*({Cal(CZ)}<>"OK"))', "Conferências diferentes de OK.", None),
]
IR = {}
rr = 6
for nome, formula, leit, fmt in IND:
    if formula is None:
        put(ws, f"B{rr}", nome, f=font(9, True, c=MUTED), bg=None, box=False, merge=f"B{rr}:F{rr}")
        ws.row_dimensions[rr].height = 21.75
        rr += 1
        continue
    label(ws, f"B{rr}", nome)
    calc(ws, f"C{rr}", formula, fmt=fmt)
    put(ws, f"D{rr}", leit, f=font(9, c=MUTED), bg=GRAY, merge=f"D{rr}:F{rr}")
    ws.row_dimensions[rr].height = 21.75
    IR[nome] = rr
    rr += 1
s = rr + 1
band(ws, s, "Resumo automático", "F")
label(ws, f"B{s+1}", "Aviso")
C = lambda n: f"C{IR[n]}"  # noqa: E731
calc(ws, f"C{s+1}", f'=IF({C("Passos")}+{C("Atividades")}=0,"Preencha as abas Rastro e Calendário",IF({C("Atividades com atraso")}>0,"Há atividade atrasada no calendário",'
     f'IF(COUNTIF({Rj},"{OUTRO}")>0,"O fio pulou um passo: veja a aba Rastro",IF({C("Passos a completar")}>0,"Há passo do fio a completar",'
     f'IF({C("Atividades a completar")}>0,"Há atividade do calendário a completar","OK")))))', merge=f"C{s+1}:F{s+1}", sz=9, b=False, h="left")
cf_warn(ws, f"C{s+1}:F{s+1}", f"C{s+1}")
ws.row_dimensions[s + 1].height = 21.75
NAV = s + 1
GS = s + 4
put(ws, f"B{GS - 1}", "Atividades por situação", f=font(9, True, c=MUTED), bg=None, box=False)
for k, sit in enumerate(CAL_SITS):
    put(ws, f"B{GS + k}", sit, f=font(9), bg=GRAY)
    calc(ws, f"C{GS + k}", f'=COUNTIF({Cal(CY)},B{GS + k})', b=False)
colunas(ws, f"D{GS - 1}", GS, GS + len(CAL_SITS) - 1, 2, 3, width=12, height=7)
setup(ws, REDC, f"B1:F{GS + 14}", landscape=False, fit_height=True)

# ------------------------------------------------------------------ Checklist
ws = wb.create_sheet("Checklist")
widths(ws, {"A": 2, "B": 5, "C": 70, "D": 14, "E": 44, "F": 2})
title(ws, "Checklist do sistema integrado", "Escolha o status de cada item na lista suspensa (coluna em amarelo-claro).", "E")
for col, text in zip("BCDE", ["#", "Item de verificação", "Status", "Observações"]):
    head(ws, f"{col}4", text)
ws.row_dimensions[4].height = 21.75
for k, text in enumerate(CHECK):
    rr = 5 + k
    num(ws, f"B{rr}", k + 1)
    put(ws, f"C{rr}", text, bg=GRAY)
    inp(ws, f"D{rr}", h="center")
    inp(ws, f"E{rr}")
    ws.row_dimensions[rr].height = 30
dv_list(ws, "D5:D16", YESNO, "Sim, Parcial ou Não")
cf_equal(ws, "D5:D16", YESNO_CF)
band(ws, 18, "Resumo automático", "E")
summary(ws, 19, "Itens atendidos (Sim)", '=COUNTIF(D5:D16,"Sim")', "C", "D")
summary(ws, 20, "Itens parciais", '=COUNTIF(D5:D16,"Parcial")', "C", "D")
summary(ws, 21, "Itens não atendidos", '=COUNTIF(D5:D16,"Não")', "C", "D")
summary(ws, 22, "Itens sem resposta", "=COUNTBLANK(D5:D16)", "C", "D")
summary(ws, 23, "Percentual atendido", "=D19/ROWS(D5:D16)", "C", "D", fmt="0%")
summary(ws, 24, "Situação", '=IF(D19=ROWS(D5:D16),"Validado",IF(D22>0,"Em andamento","Com pendências"))', "C", "D")
setup(ws, PURPLE, "B1:E24", landscape=False, fit_height=True)


# ------------------------------------------------------------------ Exemplos
def exemplo(ws, ex_):
    H = ex_["head"]
    cal_largura(ws)
    ws.column_dimensions["E"].width = 40
    ws.column_dimensions["G"].width = 22
    ws.column_dimensions["H"].width = 22
    title(ws, "Caso integrado", "Exemplo preenchido, para consulta. Use as abas Rastro e Calendário para a sua organização.", CZ)
    rr = 4
    for rot, val in [("Organização", H["org"]), ("Fio", H["fio"]), ("Origem", H["origem"]), ("Responsável pelo calendário", H["resp"])]:
        label(ws, f"B{rr}", rot, merge=f"B{rr}:C{rr}")
        inp(ws, f"D{rr}", val, merge=f"D{rr}:{CZ}{rr}", bg=WHITE, h="left")
        ws.row_dimensions[rr].height = 19.5
        rr += 1
    label(ws, f"B{rr}", "Ano e mês da leitura", merge=f"B{rr}:C{rr}")
    inp(ws, f"D{rr}", H["ano"], bg=WHITE, h="center")
    inp(ws, f"E{rr}", H["ref"], bg=WHITE, h="center")
    ref = f"$E${rr}"

    # ---- rastro (as colunas C a J seguem a aba Rastro: data, estudo, o que aconteceu, registro, saída, próximo, dias, conferência)
    rr += 2
    band(ws, rr, "Rastro do fio", CZ, color=BLUE)
    rr += 1
    for col, text in zip("BCDEFGHIJ", ["Passo", "Data", "Estudo", "O que aconteceu", "Registro", "O que seguiu", "Próximo estudo", "Dias", "Conferência"]):
        put(ws, f"{col}{rr}", text, f=font(10, True), bg=GRAY, h="center", merge=f"J{rr}:N{rr}" if col == "J" else None)
    ws.row_dimensions[rr].height = 30
    r1 = rr + 1
    for k, p in enumerate(ex_["passos"]):
        rr += 1
        num(ws, f"B{rr}", k + 1)
        put(ws, f"C{rr}", p["data"], h="center", fmt=DATE)
        put(ws, f"D{rr}", p["estudo"], f=font(10, True))
        put(ws, f"E{rr}", p["oque"], f=font(9))
        put(ws, f"F{rr}", p["reg"], f=font(9))
        put(ws, f"G{rr}", p["saida"] or None, f=font(9))
        put(ws, f"H{rr}", p["prox"] or None, f=font(9))
        calc(ws, f"I{rr}", "=" + f_dias(rr, r1))
        calc(ws, f"J{rr}", "=" + f_rasconf(rr), b=False, sz=9, merge=f"J{rr}:N{rr}")
        ws.row_dimensions[rr].height = alt([(p["oque"], 44), (p["reg"], 24), (p["saida"], 24), (p["estudo"], 22)], minimo=21.75)
    r2 = rr
    cf_texto(ws, f"J{r1}:N{r2}", f"$J{r1}", [("OK", GREEN), (OUTRO, RED)], resto=YELLOW)

    # ---- calendário
    rr += 2
    ws.row_breaks.append(Break(id=rr - 1))
    band(ws, rr, "Calendário do sistema", CZ, color=TEAL)
    rr += 1
    for k, text in enumerate(["#", "Atividade", "Estudo", "Requisito", "Responsável", "Frequência", "Início"] + MESES + ["Previstas", "Feitas", "Atrasadas", "Fora", "Situação",
                                                                                                                 "Conferência", "A cada"]):
        put(ws, f"{L(2 + k)}{rr}", text, f=font(10, True), bg=GRAY, h="center")
    ws.row_dimensions[rr].height = 30
    mrow = rr + 1
    for k, c in enumerate(MC):
        put(ws, f"{c}{mrow}", k + 1, f=font(9, i=True, c=MUTED), bg=GRAY, h="center")
    rr = mrow
    c1 = rr + 1
    for k, a in enumerate(ex_["cal"]):
        rr += 1
        num(ws, f"B{rr}", k + 1)
        put(ws, f"C{rr}", a["ativ"], f=font(10, True))
        put(ws, f"D{rr}", a["estudo"], f=font(9))
        put(ws, f"E{rr}", a["req"], f=font(9), h="center")
        put(ws, f"F{rr}", a["resp"], f=font(9))
        put(ws, f"G{rr}", a["freq"], h="center", f=font(9))
        put(ws, f"H{rr}", a["inicio"], h="center")
        for m in range(1, 13):
            put(ws, f"{MC[m - 1]}{rr}", "X" if m in a["feito"] else None, h="center")
        cal_formulas(ws, rr, mrow, ref)
        ws.row_dimensions[rr].height = alt([(a["ativ"], 34), (a["resp"], 22), (a["estudo"], 24)], minimo=21.75)
    c2 = rr
    cf_meses(ws, c1, c2, mrow, ref)
    cf_texto(ws, f"{CY}{c1}:{CY}{c2}", f"{CY}{c1}", SIT_CF)
    cf_warn(ws, f"{CZ}{c1}:{CZ}{c2}", f"{CZ}{c1}")

    # ---- resumo
    rr += 2
    band(ws, rr, "Resumo automático", CZ)
    for k, (text, formula) in enumerate([
        ("Fio", f'=COUNTA(D{r1}:D{r2})&" passos, "&SUMPRODUCT((D{r1}:D{r2}<>"")/COUNTIF(D{r1}:D{r2},D{r1}:D{r2}&""))&" estudos, "&(MAX(C{r1}:C{r2})-MIN(C{r1}:C{r2}))'
                f'&" dias; maior intervalo de "&MAX(I{r1}:I{r2})&" dias"'),
        ("Calendário", f'=SUM({CV}{c1}:{CV}{c2})&" de "&SUM({CU}{c1}:{CU}{c2})&" feitas até o mês da leitura, "&SUM({CW}{c1}:{CW}{c2})&" atrasadas, "&SUM({CX}{c1}:{CX}{c2})&" fora do plano"'),
        ("Por situação", "=" + '&", "&'.join(f'COUNTIF({CY}{c1}:{CY}{c2},"{x}")&" {x.lower()}"' for x in CAL_SITS)),
        ("Linhas a completar", f'=SUMPRODUCT((J{r1}:J{r2}<>"OK")*1)+SUMPRODUCT(({CZ}{c1}:{CZ}{c2}<>"OK")*1)'),
    ], 1):
        label(ws, f"B{rr+k}", text, merge=f"B{rr+k}:D{rr+k}", h="right")
        calc(ws, f"E{rr+k}", formula, merge=f"E{rr+k}:{CZ}{rr+k}", h="left")
        ws.row_dimensions[rr + k].height = 21.75
    setup(ws, MUTED, f"B1:{CZ}{rr + 4}")
    return dict(r=(r1, r2), c=(c1, c2), res=rr + 1)


POS = {}
POS["ex1"] = exemplo(wb.create_sheet("Exemplo 1 - Pizzaria"), EX1)
POS["ex2"] = exemplo(wb.create_sheet("Exemplo 2 - Indústria"), EX2)

wb.active = 1
wb.save(OUT)
print("ok", OUT, "| avisos: Rastro E%d, Calendário E%d, Painel C%d | linhas do painel %s | exemplos %s" % (RAV, CAV, NAV, IR, POS))
