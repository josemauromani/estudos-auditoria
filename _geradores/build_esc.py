# -*- coding: utf-8 -*-
"""Gera Escopo-modelo.xlsx no mesmo padrão visual das planilhas anteriores da série."""
import math
import os
import sys

_here = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, _here)
# funções de formatação compartilhadas com o gerador do PDCA (bloco inicial do arquivo)
_src = open(os.path.join(_here, "build_pdca.py"), encoding="utf-8").read()
exec(_src[: _src.index("STATUS = [")])

from openpyxl.chart import BarChart, Series  # noqa: E402
from openpyxl.chart.label import DataLabelList  # noqa: E402
from esc_data import (CAMPOS, CHECK, COMPROMISSOS, DEM, EX1, EX2, EXCLUIVEL, NAO, NC, NDEM, OBRIG, PARC, REQS, SIM, SITS, aplicabilidade)  # noqa: E402

BLUE, TEAL, AMBER, PURPLE, REDC = "2B5C8A", "1E7B73", "A96A12", "7A4A9A", "B0413E"
S1_T, S2_T, S3_T = "DCE8F3", "F8EBCB", "F5DEDC"
YESNO = ["Sim", "Parcial", "Não"]
YESNO_CF = [("Sim", GREEN), ("Parcial", YELLOW), ("Não", RED)]
PODE, SEMPRE = "Pode não se aplicar", "Aplica-se sempre"
INDEV, SEMPRE_C = "Exclusão indevida", "Requisito que se aplica sempre"
ES, AP, LI = "Escopo", "Aplicabilidade", "'Liderança'"
F1 = 5                          # primeiro campo do escopo, na aba Escopo
A1, A2 = 7, 7 + len(REQS) - 1   # requisitos, na aba Aplicabilidade
L1, L2 = 6, 6 + NC - 1          # compromissos, na aba Liderança
SIT_CF = [(DEM, GREEN), (PARC, S2_T), (NDEM, S3_T)]


def note(ws, ref, text):
    c = Comment(text, "Modelo Escopo")
    c.width, c.height = 300, 120
    ws[ref].comment = c


def cf_warn(ws, rng_, first, ok_values=("OK",)):
    cond = ",".join(f'{first}="{v}"' for v in ok_values)
    ws.conditional_formatting.add(rng_, FormulaRule(formula=[f"OR({cond})"], fill=PatternFill("solid", bgColor=GREEN, fgColor=GREEN)))
    ws.conditional_formatting.add(rng_, FormulaRule(formula=[f"AND(LEN({first})>0,NOT(OR({cond})))"],
                                  fill=PatternFill("solid", bgColor=YELLOW, fgColor=YELLOW)))


def cf_texto(ws, rng_, first, pares, resto=None):
    for t, cor in pares:
        ws.conditional_formatting.add(rng_, FormulaRule(formula=[f'{first}="{t}"'], stopIfTrue=True, fill=PatternFill("solid", bgColor=cor, fgColor=cor)))
    if resto:
        ws.conditional_formatting.add(rng_, FormulaRule(formula=[f"LEN({first})>0"], fill=PatternFill("solid", bgColor=resto, fgColor=resto)))


def hint_row(ws, row, cells, height=30):
    for ref, text in cells:
        put(ws, f"{ref}{row}", text, f=font(9, i=True, c=MUTED), bg=GRAY, h="center")
    ws.row_dimensions[row].height = height


def alt(pares, minimo=21.75, linha=12):
    n = max(max(1, math.ceil(len(str(t)) / max(c, 1))) for t, c in pares)
    return max(minimo, n * linha + 7)


# ------------------------------------------------------------------ fórmulas compartilhadas pelas abas de entrada e pelos exemplos
def f_aplic(pode, aplica, just, afeta):
    return (f'IF({aplica}="","Falta decidir",IF({aplica}="{SIM}","OK",IF({pode}="{SEMPRE}","{SEMPRE_C}",IF({just}="","Falta a justificativa",'
            f'IF({afeta}="","Falta avaliar o efeito",IF({afeta}="{SIM}","{INDEV}","OK"))))))')


def f_lider(faz, sit, acao, resp, prazo, reg):
    return (f'IF({faz}="","Falta a evidência",IF({sit}="","Falta a situação",IF(AND(OR({sit}="{PARC}",{sit}="{NDEM}"),{acao}=""),"Falta a ação",'
            f'IF(AND({acao}<>"",OR({resp}="",{prazo}="")),"Falta o responsável ou o prazo",IF(AND({sit}="{DEM}",{reg}=""),"Falta o registro","OK")))))')


def chart_style(ch):
    ch.style = 2
    ch.title = None
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ch.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill="D3DBE0", w=9525))
    ch.x_axis.graphicalProperties = GraphicalProperties(ln=LineProperties(solidFill="D3DBE0", w=9525))
    ch.y_axis.graphicalProperties = GraphicalProperties(ln=LineProperties(noFill=True))


def colunas(ws, anchor, r1, r2, c_cat, c_val, width=14, height=7.5):
    ch = BarChart()
    ch.type = "col"
    ch.gapWidth = 80
    ch.height, ch.width = height, width
    chart_style(ch)
    ch.legend = None
    s = Series(Reference(ws, min_col=c_val, min_row=r1, max_row=r2), title="Compromissos")
    s.graphicalProperties.solidFill = BLUE
    s.graphicalProperties.line.noFill = True
    s.dLbls = DataLabelList()
    s.dLbls.showVal = True
    s.dLbls.showSerName = s.dLbls.showCatName = s.dLbls.showLegendKey = False
    ch.append(s)
    ch.set_categories(Reference(ws, min_col=c_cat, min_row=r1, max_row=r2))
    ch.y_axis.scaling.min = 0
    ch.y_axis.majorUnit = 2
    ws.add_chart(ch, anchor)


# ================================================================== pasta de trabalho
wb = Workbook()
wb.properties.title = "Escopo e liderança — Modelo"
wb.properties.creator = "Modelo Escopo"
wb.properties.language = "pt-BR"

# ------------------------------------------------------------------ Instruções
ws = wb.active
ws.title = "Instruções"
widths(ws, {"A": 2, "B": 26, "C": 90, "D": 2})
title(ws, "Escopo e liderança — Como usar esta planilha",
      "Modelo para escrever o escopo, decidir a aplicabilidade de cada requisito e registrar como a direção demonstra a liderança.", "C")
r = 4


def section(text):
    global r
    band(ws, r, text, "C", sz=11)
    r += 1


def line(k, v, kbg=GRAY, vbg=None, kf=None, kh="left", height=None):
    global r
    put(ws, f"B{r}", k, f=kf or font(10, True), bg=kbg, h=kh)
    put(ws, f"C{r}", v, bg=vbg)
    ws.row_dimensions[r].height = height or (19.5 if len(v) <= 96 else (31.5 if len(v) <= 192 else 45.75))
    r += 1


section("Legenda: onde preencher")
line("Amarelo-claro", "Células de entrada. É aqui que você digita.", vbg=INPUT)
line("Cinza", "Células calculadas ou fixas (requisitos, compromissos, conferências, avisos). Não altere.", vbg=GRAY)
line("Listas suspensas", "Aplica-se, efeito no cliente, situação do compromisso e checklist aceitam apenas as opções da lista.")
line("Comentários", "Células com triângulo vermelho no canto trazem uma dica de preenchimento. Passe o mouse sobre elas.")
r += 1
section("As regras de cálculo")
for k, text in [
    (PODE, "Os requisitos 7.1.5, 8.3, 8.5.3 e 8.5.5. Podem ser declarados não aplicáveis, com justificativa, se a ausência não afetar o cliente."),
    (SEMPRE, "Todos os outros. Marcar “Não” neles gera o aviso “requisito que se aplica sempre”."),
    (INDEV, "Requisito declarado não aplicável cuja ausência afeta a conformidade do produto ou a satisfação do cliente."),
    ("Falta a ação", f"Compromisso “{PARC}” ou “{NDEM}” sem ação."),
    ("Falta o registro", f"Compromisso “{DEM}” sem dizer onde fica a evidência."),
]:
    line(k, text)
r += 1
section("Ordem recomendada de preenchimento")
for k, text in enumerate([
    "Aba Escopo: escreva a organização, as unidades, os produtos e serviços, os processos e os processos terceirizados.",
    "Aba Escopo: escreva o que fica fora, se houver, e por quê.",
    "Aba Aplicabilidade: para cada requisito, diga se se aplica. Para os que não se aplicam, escreva a justificativa e avalie o efeito no cliente.",
    "Aba Escopo: escreva a declaração, a revisão e quem aprovou.",
    "Aba Liderança: para cada compromisso, o que a direção faz, a frequência, o registro, a situação e, se preciso, a ação.",
    "Aba Painel: leia os avisos. Depois, valide na aba Checklist.",
], 1):
    line(f"Passo {k}", text)
r += 1
section("Abas da planilha")
for k, text in [
    (ES, "Os campos do escopo, com a conferência dos obrigatórios: unidades, produtos e serviços, processos, declaração e aprovação."),
    (AP, f"Os {len(REQS)} requisitos do estudo da ISO 9001, com a decisão e a conferência de cada um."),
    ("Liderança", f"Os {NC} compromissos da direção, com evidência, situação e ação."),
    ("Painel", "Os indicadores, os compromissos por situação, o gráfico e o aviso."),
    ("Checklist", "Doze verificações do escopo e da liderança, com percentual de conclusão."),
    ("Exemplo 1 - Pizzaria", "O primeiro escopo da loja, em rascunho, e a liderança do dono."),
    ("Exemplo 2 - Indústria", "A revisão do escopo antes da certificação, e a liderança da diretoria."),
]:
    line(k, text)
r += 1
section("Regras e premissas do modelo")
for k, text in [
    ("Lista de requisitos", "A mesma do estudo ISO 9001 requisito a requisito: 45 itens, com alguns subitens agrupados."),
    ("Requisitos que podem não se aplicar", "Convenção deste material, a partir da prática. A norma não traz uma lista; ela pede justificativa e que a ausência não afete o cliente."),
    ("Exemplos", "Os dados das abas de exemplo são ilustrativos, criados para este material."),
]:
    line(k, text)
setup(ws, INK, f"B1:C{r}", landscape=False, fit_height=True)

# ------------------------------------------------------------------ Escopo
ws = wb.create_sheet(ES)
widths(ws, {"A": 2, "B": 26, "C": 80, "D": 26, "E": 2})
title(ws, "Escopo do sistema de gestão da qualidade", "O que, onde e como o sistema cobre. Os requisitos não aplicáveis ficam na aba Aplicabilidade.", "D")
label(ws, "B4", "Organização")
inp(ws, "C4", h="left")
head(ws, "D4", "Conferência")
ws.row_dimensions[4].height = 21.75
CR = {}
for k, (chave, rot) in enumerate(CAMPOS):
    rr = F1 + k
    CR[chave] = rr
    label(ws, f"B{rr}", rot)
    inp(ws, f"C{rr}", h="left")
    if chave in OBRIG:
        calc(ws, f"D{rr}", f'=IF(C{rr}="","Falta: {rot.lower()}","OK")', b=False, sz=9)
    else:
        put(ws, f"D{rr}", "Opcional", f=font(9, i=True, c=MUTED), bg=GRAY, h="center")
    ws.row_dimensions[rr].height = 45 if chave in ("produtos", "processos", "declaracao") else 30
F2 = F1 + len(CAMPOS) - 1
cf_warn(ws, f"D{F1}:D{F2}", f"D{F1}", ok_values=("OK", "Opcional"))
note(ws, f"B{CR['terceiros']}", "Atividades do sistema feitas por outras empresas: transporte, armazenagem, laboratório, calibração. Continuam no escopo.")
note(ws, f"B{CR['fora']}", "Unidades, linhas de produto ou serviços que ficam fora do sistema, e o motivo. É uma decisão de fronteira, e não de aplicabilidade.")
note(ws, f"B{CR['declaracao']}", "O que, onde e como, em uma ou duas frases. É o texto que costuma ir para o certificado.")
s = F2 + 2
band(ws, s, "Resumo automático", "D")
for k, (text, formula) in enumerate([
    ("Campos obrigatórios preenchidos", f'=COUNTIF(D{F1}:D{F2},"OK")&" de {len(OBRIG)}"'),
    ("Requisitos não aplicáveis", f'=COUNTIF({AP}!E{A1}:E{A2},"{NAO}")&" de {len(REQS)}: veja a aba Aplicabilidade"'),
    ("Aviso", f'=IF(C4="","Escreva a organização",IF(COUNTIF(D{F1}:D{F2},"OK")<{len(OBRIG)},"Há campo obrigatório em branco",'
              f'IF(COUNTIF({AP}!H{A1}:H{A2},"{INDEV}")+COUNTIF({AP}!H{A1}:H{A2},"{SEMPRE_C}")>0,"Há requisito declarado não aplicável que se aplica",'
              f'IF(SUMPRODUCT(({AP}!H{A1}:H{A2}<>"OK")*1)>0,"Há requisito a decidir ou a completar na aba Aplicabilidade","OK"))))'),
], 1):
    label(ws, f"B{s+k}", text, h="right")
    calc(ws, f"C{s+k}", formula, merge=f"C{s+k}:D{s+k}", sz=9 if text == "Aviso" else 10, b=text != "Aviso", h="left")
    ws.row_dimensions[s + k].height = 21.75
EAV = s + 3
cf_warn(ws, f"C{EAV}:D{EAV}", f"C{EAV}")
setup(ws, BLUE, f"B1:D{EAV}", landscape=False, fit_height=True)

# ------------------------------------------------------------------ Aplicabilidade
ws = wb.create_sheet(AP)
widths(ws, {"A": 2, "B": 8, "C": 50, "D": 18, "E": 12, "F": 46, "G": 14, "H": 30, "I": 2})
title(ws, "Aplicabilidade dos requisitos", "Requisito por requisito: aplica-se? Para os que não se aplicam, a justificativa e o efeito no cliente.", "H")
for col, text in zip("BCDEFGH", ["Requisito", "Título", "Pode não se aplicar?", "Aplica-se?", "Justificativa, se não se aplica", "Afeta o cliente?", "Conferência"]):
    head(ws, f"{col}4", text)
ws.row_dimensions[4].height = 33
hint_row(ws, 5, [("B", ""), ("C", "Do estudo da ISO 9001"), ("D", "Convenção do material"), ("E", "Lista"), ("F", "Por que nenhuma atividade se liga a ele"), ("G", "Lista"), ("H", "Calculada")])
ex(ws, "B6", "Ex.", h="center")
for col, v, h_ in [("C", "Projeto e desenvolvimento de produtos e serviços", "left"), ("D", PODE, "center"), ("E", NAO, "center"),
                   ("F", "Produzimos só pelo desenho e pela especificação do cliente, sem alterar nada.", "left"), ("G", NAO, "center"), ("H", "OK", "center")]:
    ex(ws, f"{col}6", v, h=h_)
ws.row_dimensions[6].height = 30
for k, (n, t) in enumerate(REQS):
    rr = A1 + k
    num(ws, f"B{rr}", n)
    put(ws, f"C{rr}", t, f=font(9), bg=GRAY)
    put(ws, f"D{rr}", PODE if n in EXCLUIVEL else SEMPRE, f=font(9, c=MUTED), bg=GRAY, h="center")
    if n in EXCLUIVEL:
        note(ws, f"D{rr}", EXCLUIVEL[n])
    inp(ws, f"E{rr}", h="center")
    inp(ws, f"F{rr}")
    inp(ws, f"G{rr}", h="center")
    calc(ws, f"H{rr}", "=" + f_aplic(f"D{rr}", f"E{rr}", f"F{rr}", f"G{rr}"), b=False, sz=9)
    ws.row_dimensions[rr].height = 24
dv_list(ws, f"E{A1}:E{A2}", [SIM, NAO], "O requisito se aplica ao sistema?")
dv_list(ws, f"G{A1}:G{A2}", [SIM, NAO], "Deixar de aplicá-lo afeta a conformidade do produto ou a satisfação do cliente?")
cf_texto(ws, f"E{A1}:E{A2}", f"E{A1}", [(SIM, GREEN), (NAO, S2_T)])
cf_texto(ws, f"H{A1}:H{A2}", f"H{A1}", [("OK", GREEN), (INDEV, RED), (SEMPRE_C, RED)], resto=YELLOW)
s = A2 + 2
band(ws, s, "Resumo automático", "H")
for k, (text, formula) in enumerate([
    ("Requisitos decididos", f'=COUNTA(E{A1}:E{A2})&" de {len(REQS)}"'),
    ("Aplicáveis e não aplicáveis", f'=COUNTIF(E{A1}:E{A2},"{SIM}")&" aplicáveis, "&COUNTIF(E{A1}:E{A2},"{NAO}")&" não aplicáveis"'),
    ("Aviso", f'=IF(COUNTA(E{A1}:E{A2})=0,"Decida a aplicabilidade de cada requisito",IF(COUNTIF(H{A1}:H{A2},"{INDEV}")+COUNTIF(H{A1}:H{A2},"{SEMPRE_C}")>0,'
              f'"Há requisito declarado não aplicável que se aplica",IF(SUMPRODUCT((H{A1}:H{A2}<>"OK")*1)>0,"Há requisito a decidir ou a completar","OK")))'),
], 1):
    label(ws, f"B{s+k}", text, merge=f"B{s+k}:C{s+k}", h="right")
    calc(ws, f"D{s+k}", formula, merge=f"D{s+k}:H{s+k}", sz=9 if text == "Aviso" else 10, b=text != "Aviso", h="left")
    ws.row_dimensions[s + k].height = 21.75
AAV = s + 3
cf_warn(ws, f"D{AAV}:H{AAV}", f"D{AAV}")
ws.freeze_panes = f"D{A1}"
setup(ws, TEAL, f"B1:H{AAV}", landscape=False, fit_height=False)

# ------------------------------------------------------------------ Liderança
ws = wb.create_sheet("Liderança")
widths(ws, {"A": 2, "B": 5, "C": 34, "D": 42, "E": 16, "F": 26, "G": 16, "H": 36, "I": 20, "J": 12, "K": 28, "L": 2})
title(ws, "Liderança e comprometimento da direção", "Para cada compromisso: o que a direção faz, com que frequência e onde fica o registro. O que é parcial pede ação.", "K")
for col, text in zip("BCDEFGHIJK", ["#", "Compromisso", "O que a direção faz", "Frequência", "Registro", "Situação", "Ação", "Responsável", "Prazo", "Conferência"]):
    head(ws, f"{col}4", text)
ws.row_dimensions[4].height = 33
hint_row(ws, 5, [("B", ""), ("C", "Fixo"), ("D", "Um fato, e não uma intenção"), ("E", "Mensal, anual"), ("F", "Onde a evidência fica"), ("G", "Lista"), ("H", "Se parcial ou não demonstrado"),
                 ("I", "Uma função"), ("J", "Data"), ("K", "Calculada")])
for k, (tit, desc) in enumerate(COMPROMISSOS):
    rr = L1 + k
    num(ws, f"B{rr}", k + 1)
    put(ws, f"C{rr}", desc, f=font(9), bg=GRAY)
    for col in "DEFHI":
        inp(ws, f"{col}{rr}")
    inp(ws, f"G{rr}", h="center")
    inp(ws, f"J{rr}", h="center", fmt=DATE)
    calc(ws, f"K{rr}", "=" + f_lider(f"D{rr}", f"G{rr}", f"H{rr}", f"I{rr}", f"J{rr}", f"F{rr}"), b=False, sz=9)
    ws.row_dimensions[rr].height = 36
dv_list(ws, f"G{L1}:G{L2}", SITS, "Demonstrado, Parcial ou Não demonstrado")
dv_date(ws, f"J{L1}:J{L2}")
cf_texto(ws, f"G{L1}:G{L2}", f"G{L1}", SIT_CF)
cf_warn(ws, f"K{L1}:K{L2}", f"K{L1}")
note(ws, "D4", "O que a direção faz, de fato: “conduz a análise crítica”, “aprovou a compra do medidor”. Evite “apoia a qualidade”.")
s = L2 + 2
band(ws, s, "Resumo automático", "K")
for k, (text, formula) in enumerate([
    ("Por situação", f'=COUNTIF(G{L1}:G{L2},"{DEM}")&" demonstrados, "&COUNTIF(G{L1}:G{L2},"{PARC}")&" parciais, "&COUNTIF(G{L1}:G{L2},"{NDEM}")&" não demonstrados"'),
    ("Ações", f'=COUNTA(H{L1}:H{L2})'),
    ("Aviso", f'=IF(COUNTA(D{L1}:D{L2})=0,"Escreva o que a direção faz em cada compromisso",IF(COUNTIF(K{L1}:K{L2},"OK")<{NC},"Há compromisso a completar: veja a coluna Conferência",'
              f'IF(COUNTIF(G{L1}:G{L2},"{NDEM}")>0,"Há compromisso não demonstrado: leve à análise crítica","OK")))'),
], 1):
    label(ws, f"B{s+k}", text, merge=f"B{s+k}:C{s+k}", h="right")
    calc(ws, f"D{s+k}", formula, merge=f"D{s+k}:K{s+k}", sz=9 if text == "Aviso" else 10, b=text != "Aviso", h="left")
    ws.row_dimensions[s + k].height = 21.75
LAV = s + 3
cf_warn(ws, f"D{LAV}:K{LAV}", f"D{LAV}")
ws.freeze_panes = f"D{L1}"
setup(ws, AMBER, f"B1:K{LAV}", fit_height=True)

# ------------------------------------------------------------------ Painel
ws = wb.create_sheet("Painel")
widths(ws, {"A": 2, "B": 34, "C": 16, "D": 14, "E": 14, "F": 34, "G": 2})
title(ws, "Painel do escopo e da liderança", "Nada a preencher: tudo vem das outras abas.", "F")
label(ws, "B3", "Organização")
calc(ws, "C3", f'=IF({ES}!C4="","",{ES}!C4)', merge="C3:F3", h="left", b=False)
for col, text, m in [("B", "Indicador", None), ("C", "Resultado", None), ("D", "Como se lê", "D5:F5")]:
    put(ws, f"{col}5", text, f=font(10, True, c=WHITE), bg=INK, h="center", merge=m)
ws.row_dimensions[5].height = 21.75
He, Ee, Gl, Kl = f"{AP}!H{A1}:H{A2}", f"{AP}!E{A1}:E{A2}", f"{LI}!G{L1}:G{L2}", f"{LI}!K{L1}:K{L2}"
IND = [
    ("Campos obrigatórios do escopo", f'=COUNTIF({ES}!D{F1}:D{F2},"OK")&" de {len(OBRIG)}"', "Unidades, produtos e serviços, processos, declaração e aprovação."),
    ("Escopo aprovado", f'=IF({ES}!C{CR["aprov"]}="","Não","Sim")', "A direção aprova o escopo."),
    ("Requisitos decididos", f'=COUNTA({Ee})&" de {len(REQS)}"', "Da aba Aplicabilidade."),
    ("Requisitos não aplicáveis", f'=COUNTIF({Ee},"{NAO}")', "Cada um com justificativa, e sem efeito no cliente."),
    ("Não aplicáveis que se aplicam", f'=COUNTIF({He},"{INDEV}")+COUNTIF({He},"{SEMPRE_C}")', "Exclusão indevida, ou requisito que se aplica sempre."),
    ("Requisitos a decidir ou completar", f'=SUMPRODUCT(({He}<>"OK")*1)-COUNTIF({He},"{INDEV}")-COUNTIF({He},"{SEMPRE_C}")', "Sem decisão, sem justificativa ou sem o efeito avaliado."),
    (DEM, f'=COUNTIF({Gl},"{DEM}")', f"Compromissos da direção, de {NC}."),
    (PARC, f'=COUNTIF({Gl},"{PARC}")', "Pedem ação, com responsável e prazo."),
    (NDEM, f'=COUNTIF({Gl},"{NDEM}")', "Levar à análise crítica."),
    ("Compromissos a completar", f'=SUMPRODUCT(({Kl}<>"OK")*1)', "Sem evidência, sem ação ou sem registro."),
]
for k, (nome, formula, leit) in enumerate(IND):
    rr = 6 + k
    label(ws, f"B{rr}", nome)
    calc(ws, f"C{rr}", formula)
    put(ws, f"D{rr}", leit, f=font(9, c=MUTED), bg=GRAY, merge=f"D{rr}:F{rr}")
    ws.row_dimensions[rr].height = 21.75
IR = {n: 6 + k for k, (n, _, _) in enumerate(IND)}
cf_texto(ws, f"C{IR['Escopo aprovado']}", f"C{IR['Escopo aprovado']}", [("Sim", GREEN), ("Não", RED)])
T0 = 6 + len(IND) + 1
band(ws, T0, "Os compromissos da direção", "F")
for col, text in zip("BCDEF", ["Compromisso", "Situação", "Prazo da ação", "Responsável", "Conferência"]):
    head(ws, f"{col}{T0+1}", text)
for k, (tit, _) in enumerate(COMPROMISSOS):
    rr = T0 + 2 + k
    lr = L1 + k
    label(ws, f"B{rr}", f"{k + 1} · {tit}")
    calc(ws, f"C{rr}", f'=IF({LI}!G{lr}="","",{LI}!G{lr})', b=False, sz=9)
    calc(ws, f"D{rr}", f'=IF({LI}!J{lr}="","",{LI}!J{lr})', b=False, fmt=DATE)
    calc(ws, f"E{rr}", f'=IF({LI}!I{lr}="","",{LI}!I{lr})', b=False, sz=9)
    calc(ws, f"F{rr}", f'={LI}!K{lr}', b=False, sz=9)
    ws.row_dimensions[rr].height = 21.75
C1, C2 = T0 + 2, T0 + 1 + NC
cf_texto(ws, f"C{C1}:C{C2}", f"C{C1}", SIT_CF)
cf_warn(ws, f"F{C1}:F{C2}", f"F{C1}")
s = C2 + 2
band(ws, s, "Resumo automático", "F")
label(ws, f"B{s+1}", "Aviso")
calc(ws, f"C{s+1}", f'=IF({ES}!C4="","Escreva o escopo na aba Escopo",IF(C{IR["Não aplicáveis que se aplicam"]}>0,"Há requisito declarado não aplicável que se aplica",'
     f'IF(COUNTIF({ES}!D{F1}:D{F2},"OK")<{len(OBRIG)},"Há campo obrigatório do escopo em branco",IF(C{IR["Requisitos a decidir ou completar"]}>0,"Há requisito a decidir na aba Aplicabilidade",'
     f'IF(C{IR[NDEM]}>0,"Há compromisso da direção não demonstrado",IF(C{IR["Compromissos a completar"]}>0,"Há compromisso a completar","OK"))))))',
     merge=f"C{s+1}:F{s+1}", sz=9, b=False, h="left")
cf_warn(ws, f"C{s+1}:F{s+1}", f"C{s+1}")
ws.row_dimensions[s + 1].height = 21.75
NAV = s + 1
GS = s + 3
for k, sit in enumerate(SITS):
    put(ws, f"B{GS + k}", sit, f=font(9), bg=GRAY)
    calc(ws, f"C{GS + k}", f'=COUNTIF({Gl},B{GS + k})', b=False)
colunas(ws, f"D{GS}", GS, GS + 2, 2, 3, width=12, height=6.5)
setup(ws, REDC, f"B1:F{GS + 14}", landscape=False, fit_height=True)

# ------------------------------------------------------------------ Checklist
ws = wb.create_sheet("Checklist")
widths(ws, {"A": 2, "B": 5, "C": 70, "D": 14, "E": 44, "F": 2})
title(ws, "Checklist do escopo e da liderança", "Escolha o status de cada item na lista suspensa (coluna em amarelo-claro).", "E")
for col, text in zip("BCDE", ["#", "Item de verificação", "Status", "Observações"]):
    head(ws, f"{col}4", text)
ws.row_dimensions[4].height = 21.75
for k, text in enumerate(CHECK):
    rr = 5 + k
    num(ws, f"B{rr}", k + 1)
    put(ws, f"C{rr}", text, bg=GRAY)
    inp(ws, f"D{rr}", h="center")
    inp(ws, f"E{rr}")
    ws.row_dimensions[rr].height = 30
dv_list(ws, "D5:D16", YESNO, "Sim, Parcial ou Não")
cf_equal(ws, "D5:D16", YESNO_CF)
band(ws, 18, "Resumo automático", "E")
summary(ws, 19, "Itens atendidos (Sim)", '=COUNTIF(D5:D16,"Sim")', "C", "D")
summary(ws, 20, "Itens parciais", '=COUNTIF(D5:D16,"Parcial")', "C", "D")
summary(ws, 21, "Itens não atendidos", '=COUNTIF(D5:D16,"Não")', "C", "D")
summary(ws, 22, "Itens sem resposta", "=COUNTBLANK(D5:D16)", "C", "D")
summary(ws, 23, "Percentual atendido", "=D19/ROWS(D5:D16)", "C", "D", fmt="0%")
summary(ws, 24, "Situação", '=IF(D19=ROWS(D5:D16),"Validado",IF(D22>0,"Em andamento","Com pendências"))', "C", "D")
setup(ws, PURPLE, "B1:E24", landscape=False, fit_height=True)


# ------------------------------------------------------------------ Exemplos
def sub(ws, rr, cells, height=30):
    for a, b, t in cells:
        put(ws, f"{a}{rr}", t, f=font(10, True), bg=GRAY, h="center", merge=b and f"{a}{rr}:{b}{rr}")
    ws.row_dimensions[rr].height = height


def exemplo(ws, ex_):
    H = ex_["head"]
    widths(ws, {"A": 2, "B": 6, "C": 26, "D": 40, "E": 14, "F": 22, "G": 16, "H": 34, "I": 18, "J": 12, "K": 28, "L": 2})
    title(ws, "Escopo e liderança", "Exemplo preenchido, para consulta. Use as abas Escopo, Aplicabilidade e Liderança para a sua organização.", "K")
    rr = 4
    for rot, val in [("Organização", H["org"]), ("Escopo escrito em", f'{H["data"]:%d/%m/%Y}. {H["por"]}.'), ("Leitura", f'{H["ref"]:%d/%m/%Y}.'), ("Origem", H["origem"])]:
        label(ws, f"B{rr}", rot, merge=f"B{rr}:C{rr}")
        inp(ws, f"D{rr}", val, merge=f"D{rr}:K{rr}", bg=WHITE, h="left")
        ws.row_dimensions[rr].height = alt([(val, 150)], minimo=19.5)
        rr += 1
    rr += 1
    band(ws, rr, "Escopo", "K", color=BLUE)
    rr += 1
    sub(ws, rr, [("B", "C", "Campo"), ("D", "I", "Conteúdo"), ("J", "K", "Conferência")])
    e = ex_["escopo"]
    c1 = rr + 1
    for k, rot in CAMPOS:
        rr += 1
        label(ws, f"B{rr}", rot, merge=f"B{rr}:C{rr}")
        put(ws, f"D{rr}", e[k] or None, merge=f"D{rr}:I{rr}")
        if k in OBRIG:
            calc(ws, f"J{rr}", f'=IF(D{rr}="","Falta: {rot.lower()}","OK")', b=False, sz=9, merge=f"J{rr}:K{rr}")
        else:
            put(ws, f"J{rr}", "Opcional", f=font(9, i=True, c=MUTED), bg=GRAY, h="center", merge=f"J{rr}:K{rr}")
        ws.row_dimensions[rr].height = alt([(e[k], 120)], minimo=21.75)
    cf_warn(ws, f"J{c1}:K{rr}", f"$J{c1}", ok_values=("OK", "Opcional"))
    rr += 2
    ap = aplicabilidade(ex_)
    band(ws, rr, f"Aplicabilidade: {sum(1 for a in ap if a['aplica'] == SIM)} de {len(ap)} requisitos aplicáveis. Os declarados não aplicáveis:", "K", color=TEAL)
    rr += 1
    sub(ws, rr, [("B", None, "Req."), ("C", None, "Título"), ("D", None, "Justificativa"), ("E", None, "Pode não se aplicar?"), ("F", None, "Aplica-se?"), ("G", None, "Afeta o cliente?"),
                 ("H", "I", "Por que pode não se aplicar"), ("J", "K", "Conferência")])
    for a in ap:
        if a["aplica"] != NAO:
            continue
        rr += 1
        num(ws, f"B{rr}", a["num"])
        put(ws, f"C{rr}", a["tit"], f=font(10, True))
        put(ws, f"D{rr}", a["just"])
        put(ws, f"E{rr}", PODE if a["num"] in EXCLUIVEL else SEMPRE, h="center", f=font(9))
        put(ws, f"F{rr}", a["aplica"], h="center")
        put(ws, f"G{rr}", a["afeta"], h="center")
        put(ws, f"H{rr}", EXCLUIVEL.get(a["num"], ""), merge=f"H{rr}:I{rr}", f=font(9, c=MUTED))
        calc(ws, f"J{rr}", "=" + f_aplic(f"E{rr}", f"F{rr}", f"D{rr}", f"G{rr}"), b=False, sz=9, merge=f"J{rr}:K{rr}")
        ws.row_dimensions[rr].height = alt([(a["tit"], 26), (a["just"], 40), (EXCLUIVEL.get(a["num"], ""), 50)], minimo=21.75)
        cf_texto(ws, f"J{rr}:K{rr}", f"$J{rr}", [("OK", GREEN), (INDEV, RED), (SEMPRE_C, RED)], resto=YELLOW)
    rr += 2
    ws.row_breaks.append(Break(id=rr - 1))
    band(ws, rr, "Liderança", "K", color=AMBER)
    rr += 1
    sub(ws, rr, [("B", None, "#"), ("C", None, "Compromisso"), ("D", None, "O que a direção faz"), ("E", None, "Frequência"), ("F", None, "Registro"), ("G", None, "Situação"), ("H", None, "Ação"),
                 ("I", None, "Responsável"), ("J", None, "Prazo"), ("K", None, "Conferência")])
    l1 = rr + 1
    for k, ((tit, _), c) in enumerate(zip(COMPROMISSOS, ex_["lid"]), 1):
        rr += 1
        num(ws, f"B{rr}", k)
        put(ws, f"C{rr}", tit, f=font(10, True))
        put(ws, f"D{rr}", c["faz"])
        put(ws, f"E{rr}", c["freq"] or None, h="center")
        put(ws, f"F{rr}", c["reg"] or None)
        put(ws, f"G{rr}", c["sit"], h="center")
        put(ws, f"H{rr}", c["acao"] or None)
        put(ws, f"I{rr}", c["resp"] or None)
        put(ws, f"J{rr}", c["prazo"], h="center", fmt=DATE)
        calc(ws, f"K{rr}", "=" + f_lider(f"D{rr}", f"G{rr}", f"H{rr}", f"I{rr}", f"J{rr}", f"F{rr}"), b=False, sz=9)
        ws.row_dimensions[rr].height = alt([(c["faz"], 38), (c["reg"], 22), (c["acao"], 32), (c["resp"], 17)], minimo=21.75)
    l2 = rr
    cf_texto(ws, f"G{l1}:G{l2}", f"G{l1}", SIT_CF)
    cf_warn(ws, f"K{l1}:K{l2}", f"K{l1}")
    rr += 2
    band(ws, rr, "Resumo automático", "K")
    for k, (text, formula) in enumerate([
        ("Campos obrigatórios do escopo", f'=COUNTIF(J{c1}:J{c1 + len(CAMPOS) - 1},"OK")&" de {len(OBRIG)}"'),
        ("Compromissos por situação", f'=COUNTIF(G{l1}:G{l2},"{DEM}")&" demonstrados, "&COUNTIF(G{l1}:G{l2},"{PARC}")&" parciais, "&COUNTIF(G{l1}:G{l2},"{NDEM}")&" não demonstrados"'),
        ("Compromissos a completar", f'=SUMPRODUCT((K{l1}:K{l2}<>"OK")*1)'),
    ], 1):
        label(ws, f"B{rr+k}", text, merge=f"B{rr+k}:D{rr+k}", h="right")
        calc(ws, f"E{rr+k}", formula, merge=f"E{rr+k}:K{rr+k}", h="left")
        ws.row_dimensions[rr + k].height = 21.75
    setup(ws, MUTED, f"B1:K{rr + 3}")


exemplo(wb.create_sheet("Exemplo 1 - Pizzaria"), EX1)
exemplo(wb.create_sheet("Exemplo 2 - Indústria"), EX2)

wb.active = 1
wb.save(OUT)
print("ok", OUT, "| avisos: Escopo C%d, Aplicabilidade D%d, Liderança D%d, Painel C%d" % (EAV, AAV, LAV, NAV))
