# -*- coding: utf-8 -*-
"""Gera Pareto-modelo.xlsx no mesmo padrão visual das planilhas anteriores da série."""
import math
import os
import sys

_here = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, _here)
# funções de formatação compartilhadas com o gerador do PDCA (bloco inicial do arquivo)
_src = open(os.path.join(_here, "build_pdca.py"), encoding="utf-8").read()
exec(_src[: _src.index("STATUS = [")])

from openpyxl.chart import BarChart, Series  # noqa: E402
from openpyxl.chart.label import DataLabelList  # noqa: E402
from openpyxl.chart.marker import Marker  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402
from par_data import CHECK, CORTE, DEPOIS, EX1, EX2, MIN_OBS, OUTROS, OUTROS_MAX, PLANO, PRIOR, colunas, linhas, ordenar  # noqa: E402

BLUE, TEAL, AMBER, PURPLE, REDC = "2B5C8A", "1E7B73", "A96A12", "7A4A9A", "B0413E"
BLUE_T = "DCE8F3"
YESNO = ["Sim", "Parcial", "Não"]
YESNO_CF = [("Sim", GREEN), ("Parcial", YELLOW), ("Não", RED)]
FREQ, PESO = "Frequência", "Frequência × peso"
TAXA = "Taxa por 100"
NCAT, NCOL = 8, 7
CO, FO, PA, AD = "Coleta", "Folha", "Pareto", "'Antes e depois'"
K1, K2 = 14, 14 + NCAT - 1      # categorias, na aba Coleta
KO = K2 + 1                     # linha de "Outros", na aba Coleta
F1, F2 = 8, 8 + NCAT - 1        # categorias, na aba Folha
FOU, FT = F2 + 1, F2 + 2        # "Outros" e total, na aba Folha
FC = [get_column_letter(4 + j) for j in range(NCOL)]   # colunas D a J das contagens
P1, P2 = 8, 8 + NCAT            # as nove posições do Pareto: oito categorias e "Outros"
A1, A2 = 11, 11 + NCAT - 1      # categorias, na aba Antes e depois
AO, AT = A2 + 1, A2 + 2
PCT = "0%"


def note(ws, ref, text):
    c = Comment(text, "Modelo Pareto")
    c.width, c.height = 300, 120
    ws[ref].comment = c


def cf_warn(ws, rng_, first, ok_values=("OK",)):
    cond = ",".join(f'{first}="{v}"' for v in ok_values)
    ws.conditional_formatting.add(rng_, FormulaRule(formula=[f"OR({cond})"], fill=PatternFill("solid", bgColor=GREEN, fgColor=GREEN)))
    ws.conditional_formatting.add(rng_, FormulaRule(formula=[f"AND(LEN({first})>0,NOT(OR({cond})))"],
                                  fill=PatternFill("solid", bgColor=YELLOW, fgColor=YELLOW)))


def cf_texto(ws, rng_, first, pares, resto=None):
    for t, cor in pares:
        ws.conditional_formatting.add(rng_, FormulaRule(formula=[f'{first}="{t}"'], stopIfTrue=True, fill=PatternFill("solid", bgColor=cor, fgColor=cor)))
    if resto:
        ws.conditional_formatting.add(rng_, FormulaRule(formula=[f"LEN({first})>0"], fill=PatternFill("solid", bgColor=resto, fgColor=resto)))


def hint_row(ws, row, cells, height=30):
    for ref, text in cells:
        put(ws, f"{ref}{row}", text, f=font(9, i=True, c=MUTED), bg=GRAY, h="center")
    ws.row_dimensions[row].height = height


def mil(v):
    return f"{v:,}".replace(",", ".")


def alt(pares, minimo=21.75, linha=12):
    n = max(max(1, math.ceil(len(str(t)) / max(c, 1))) for t, c in pares)
    return max(minimo, n * linha + 7)


def dv_numero(ws, rng_, minimo, maximo, prompt, inteiro=True):
    dv = DataValidation(type="whole" if inteiro else "decimal", operator="between", formula1=str(minimo), formula2=str(maximo), allow_blank=True)
    dv.promptTitle, dv.prompt, dv.showInputMessage = "Número", prompt, True
    dv.errorTitle, dv.error, dv.showErrorMessage = "Número inválido", f"Digite um número de {minimo} a {maximo}.", True
    ws.add_data_validation(dv)
    dv.add(rng_)


def chart_style(ch):
    ch.style = 2
    ch.title = None
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ch.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill="D3DBE0", w=9525))
    ch.x_axis.graphicalProperties = GraphicalProperties(ln=LineProperties(solidFill="D3DBE0", w=9525))
    ch.y_axis.graphicalProperties = GraphicalProperties(ln=LineProperties(noFill=True))


def pareto_chart(ws, anchor, r1, r2, c_cat, c_pct, c_acc, c_corte, width=19, height=9):
    """Barras com a parte do total e linha com o acumulado, na mesma régua de 0 a 100%."""
    bar = BarChart()
    bar.type = "col"
    bar.gapWidth = 50
    bar.height, bar.width = height, width
    chart_style(bar)
    s = Series(Reference(ws, min_col=c_pct, min_row=r1, max_row=r2), title="Parte do total")
    s.graphicalProperties.solidFill = BLUE
    s.graphicalProperties.line.noFill = True
    s.dLbls = DataLabelList()
    s.dLbls.showVal = True
    s.dLbls.numFmt = PCT
    s.dLbls.showSerName = s.dLbls.showCatName = s.dLbls.showLegendKey = False
    bar.append(s)
    bar.set_categories(Reference(ws, min_col=c_cat, min_row=r1, max_row=r2))
    bar.y_axis.scaling.min, bar.y_axis.scaling.max = 0, 1
    bar.y_axis.majorUnit = 0.2
    bar.y_axis.number_format = PCT
    line = LineChart()
    a = Series(Reference(ws, min_col=c_acc, min_row=r1, max_row=r2), title="Acumulado")
    a.graphicalProperties.line.solidFill = INK
    a.graphicalProperties.line.width = 22000
    a.marker = Marker(symbol="circle", size=7)
    a.marker.graphicalProperties = GraphicalProperties(solidFill=INK, ln=LineProperties(solidFill="FFFFFF", w=12700))
    a.smooth = False
    line.append(a)
    c = Series(Reference(ws, min_col=c_corte, min_row=r1, max_row=r2), title=f"Corte de {round(100 * CORTE)}%")
    c.graphicalProperties.line.solidFill = MUTED
    c.graphicalProperties.line.width = 12700
    c.graphicalProperties.line.dashStyle = "dash"
    c.marker = Marker(symbol="none")
    c.smooth = False
    line.append(c)
    bar += line
    bar.legend.position = "b"
    ws.add_chart(bar, anchor)


def par_chart(ws, anchor, r1, r2, c_cat, c_a, c_b, width=19, height=8.5):
    """Barras lado a lado, antes e depois, para cada categoria."""
    ch = BarChart()
    ch.type = "col"
    ch.gapWidth = 60
    ch.overlap = -10
    ch.height, ch.width = height, width
    chart_style(ch)
    for col, tit, cor in ((c_a, "Antes", "9BABB6"), (c_b, "Depois", BLUE)):
        s = Series(Reference(ws, min_col=col, min_row=r1, max_row=r2), title=tit)
        s.graphicalProperties.solidFill = cor
        s.graphicalProperties.line.noFill = True
        s.dLbls = DataLabelList()
        s.dLbls.showVal = True
        s.dLbls.numFmt = "0.0"
        s.dLbls.showSerName = s.dLbls.showCatName = s.dLbls.showLegendKey = False
        ch.append(s)
    ch.set_categories(Reference(ws, min_col=c_cat, min_row=r1, max_row=r2))
    ch.y_axis.scaling.min = 0
    ch.legend.position = "b"
    ws.add_chart(ch, anchor)


# ================================================================== pasta de trabalho
wb = Workbook()
wb.properties.title = "Pareto e folha de verificação — Modelo"
wb.properties.creator = "Modelo Pareto"
wb.properties.language = "pt-BR"

# ------------------------------------------------------------------ Instruções
ws = wb.active
ws.title = "Instruções"
widths(ws, {"A": 2, "B": 24, "C": 92, "D": 2})
title(ws, "Pareto e folha de verificação — Como usar esta planilha",
      "Modelo para planejar a coleta, lançar as contagens, ordenar as categorias e comparar o antes com o depois.", "C")
r = 4


def section(text):
    global r
    band(ws, r, text, "C", sz=11)
    r += 1


def line(k, v, kbg=GRAY, vbg=None, kf=None, kh="left", height=None):
    global r
    put(ws, f"B{r}", k, f=kf or font(10, True), bg=kbg, h=kh)
    put(ws, f"C{r}", v, bg=vbg)
    ws.row_dimensions[r].height = height or (19.5 if len(v) <= 98 else (31.5 if len(v) <= 196 else 45.75))
    r += 1


section("Legenda: onde preencher")
line("Amarelo-claro", "Células de entrada. É aqui que você digita.", vbg=INPUT)
line("Cinza", "Células calculadas ou fixas (rótulos, totais, ordem, acumulado, classes, avisos). Não altere.", vbg=GRAY)
line("Cinza em itálico", "Linha de exemplo (marcada com “Ex.”). Mostra o formato esperado e não entra nas contagens.", vbg=GRAY)
line("Listas suspensas", "A base do Pareto e o status do checklist aceitam apenas as opções da lista. As contagens aceitam números inteiros.")
line("Comentários", "Células com triângulo vermelho no canto trazem uma dica de preenchimento. Passe o mouse sobre elas.")
r += 1
section("As regras de cálculo")
for k, text in [
    ("Ordem", "As categorias são ordenadas do maior valor para o menor. No empate, vale a ordem da aba Coleta. “Outros” fica sempre por último."),
    ("Parte e acumulado", "Parte do total = valor da categoria ÷ total. Acumulado = soma das partes, da primeira categoria até a atual."),
    ("Classe", f"Prioridade: a categoria começa antes de {round(100 * CORTE)}% do acumulado, e por isso entra a que cruza o corte. Depois: começa em {round(100 * CORTE)}% ou mais."),
    ("Base do Pareto", "Frequência: o valor é a contagem. Frequência × peso: o valor é a contagem vezes o peso por ocorrência da aba Coleta (custo, tempo, quilos)."),
    ("Taxa", "Ocorrências ÷ base × 100. A base é o número de oportunidades do período: entregas, requisições, peças."),
    ("Avisos", f"Menos de {MIN_OBS} ocorrências: poucos dados. “Outros” acima de {round(100 * OUTROS_MAX)}%: abrir em categorias. "
               f"Duas primeiras barras abaixo de {round(100 * PLANO)}%, com quatro categorias ou mais: Pareto plano."),
]:
    line(k, text)
r += 1
section("Ordem recomendada de preenchimento")
for k, text in enumerate([
    "Aba Coleta: escreva o que se conta, o período, quem registra e a base. Cadastre as categorias, cada uma com a sua definição.",
    "Aba Folha: escreva o fator das colunas (dia, turno, área) e lance as contagens do período.",
    "Aba Pareto: leia a ordem, o acumulado, a classe e o aviso. Troque a base para frequência × peso, se as ocorrências não custam o mesmo.",
    "Leve as categorias da prioridade para a análise de causa, e registre a decisão.",
    "Aba Antes e depois: depois da ação, lance a nova contagem e a nova base, e leia a variação.",
    "Aba Checklist: valide a análise.",
], 1):
    line(f"Passo {k}", text)
r += 1
section("Abas da planilha")
for k, text in [
    ("Coleta", f"O plano da coleta e até {NCAT} categorias, com definição e peso. “Outros” já está na lista."),
    ("Folha", f"As contagens por categoria, em até {NCOL} colunas do segundo fator. Calcula os totais e a coluna com mais ocorrências."),
    ("Pareto", "A ordem, a parte do total, o acumulado, a classe, o aviso e o gráfico."),
    ("Antes e depois", "A contagem feita depois da ação. Calcula a taxa, a variação e a primeira categoria de agora."),
    ("Checklist", "Doze verificações de qualidade da análise, com percentual de conclusão."),
    ("Exemplo 1 - Pizzaria", "Os atrasos nas entregas, por motivo e por dia da semana, com o antes e o depois."),
    ("Exemplo 2 - Compras", "As requisições devolvidas, por motivo e por área requisitante."),
]:
    line(k, text)
r += 1
section("Regras e premissas do modelo")
for k, text in [
    ("Categorias", "De quatro a oito, mais “Outros”. Com mais do que isso, junte as menores em “Outros” ou faça um segundo Pareto da primeira barra."),
    ("Folha impressa", "A contagem é feita no papel, na hora, com um traço por ocorrência. A planilha recebe os totais do período."),
    ("Gráfico", "As barras mostram a parte de cada categoria, e a linha, o acumulado, na mesma régua de 0 a 100%. A coluna “Linha do gráfico” fica vazia nas posições sem categoria."),
    ("Exemplos", "Os dados das abas de exemplo são ilustrativos, criados para este material."),
]:
    line(k, text)
setup(ws, INK, f"B1:C{r}", landscape=False, fit_height=True)

# ------------------------------------------------------------------ Coleta
ws = wb.create_sheet(CO)
widths(ws, {"A": 2, "B": 5, "C": 36, "D": 58, "E": 14, "F": 26, "G": 2})
title(ws, "O plano da coleta", "Defina o que será contado, onde, quando e por quem. Cada categoria precisa de uma definição escrita.", "F")
for rr, rot in [(4, "Organização e processo"), (5, "O que se conta (unidade e critério)"), (6, "Onde e em que período"), (7, "Quem registra, e como"),
                (8, "Base: oportunidades no período"), (9, "Nome da base (entregas, requisições, peças)")]:
    label(ws, f"B{rr}", rot, merge=f"B{rr}:C{rr}")
    inp(ws, f"D{rr}", merge=f"D{rr}:F{rr}", h="left")
    ws.row_dimensions[rr].height = 21.75
dv_numero(ws, "D8", 1, 1000000000, "Quantas oportunidades houve no período: entregas, requisições, peças produzidas.")
note(ws, "B5", "Uma unidade e um critério. Ex.: entregas de delivery que chegaram depois de 40 minutos, pelo motivo principal.")
note(ws, "B8", "Opcional. Com a base, a planilha calcula a taxa por 100 e permite comparar períodos de tamanhos diferentes.")
for col, text in zip("BCDEF", ["#", "Categoria", "Definição: conta quando…", "Peso por ocorrência", "Conferência"]):
    head(ws, f"{col}11", text)
ws.row_dimensions[11].height = 33
hint_row(ws, 12, [("B", ""), ("C", "Nome curto, com as mesmas palavras sempre"), ("D", "A regra que decide se a ocorrência entra nesta categoria"),
                  ("E", "Opcional: custo, tempo ou quilos"), ("F", "Calculada")])
ex(ws, "B13", "Ex.", h="center")
for col, v, h_ in [("C", "Fila no forno", "left"), ("D", "O pedido esperou mais de 10 minutos para entrar no forno.", "left"), ("E", "", "center"), ("F", "OK", "center")]:
    ex(ws, f"{col}13", v, h=h_)
ws.row_dimensions[13].height = 21.75
for k in range(NCAT):
    rr = K1 + k
    num(ws, f"B{rr}", k + 1)
    inp(ws, f"C{rr}")
    inp(ws, f"D{rr}")
    inp(ws, f"E{rr}", h="center")
    calc(ws, f"F{rr}", f'=IF(C{rr}="","",IF(COUNTIF(C${K1}:C${K2},C{rr})>1,"Categoria repetida",IF(D{rr}="","Falta a definição","OK")))', b=False, sz=9)
    ws.row_dimensions[rr].height = 27
num(ws, f"B{KO}", "—")
put(ws, f"C{KO}", OUTROS, f=font(10, True), bg=GRAY)
inp(ws, f"D{KO}", "O que não cabe nas categorias acima. Anotar o motivo ao lado do traço.")
inp(ws, f"E{KO}", h="center")
put(ws, f"F{KO}", "Fica sempre por último", f=font(9, i=True, c=MUTED), bg=GRAY, h="center")
ws.row_dimensions[KO].height = 27
dv_numero(ws, f"E{K1}:E{KO}", 0, 1000000000, "Custo, tempo ou quantidade perdida em cada ocorrência desta categoria.", inteiro=False)
cf_warn(ws, f"F{K1}:F{K2}", f"F{K1}")
note(ws, "D11", "Teste a definição: duas pessoas, lendo o mesmo caso, precisam pôr o traço na mesma linha.")
note(ws, "E11", "Preencha quando as ocorrências não custam o mesmo. Na aba Pareto, troque a base para “Frequência × peso”.")
s = KO + 2
band(ws, s, "Resumo automático", "F")
label(ws, f"B{s+1}", "Categorias cadastradas", merge=f"B{s+1}:C{s+1}", h="right")
calc(ws, f"D{s+1}", f"=COUNTA(C{K1}:C{K2})", h="left")
label(ws, f"B{s+2}", "Aviso", merge=f"B{s+2}:C{s+2}", h="right")
calc(ws, f"D{s+2}", f'=IF(D5="","Descreva o que se conta",IF(D{s+1}=0,"Cadastre as categorias",IF(COUNTIF(F{K1}:F{K2},"OK")<D{s+1},"Há categoria a rever: veja a coluna Conferência",'
     f'IF(D{s+1}<4,"Poucas categorias: o usual é de quatro a oito","OK"))))', merge=f"D{s+2}:F{s+2}", sz=9, b=False, h="left")
cf_warn(ws, f"D{s+2}:F{s+2}", f"D{s+2}")
for k in (1, 2):
    ws.row_dimensions[s + k].height = 21.75
setup(ws, BLUE, f"B1:F{s+2}", fit_height=True)

# ------------------------------------------------------------------ Folha
ws = wb.create_sheet(FO)
LT, LP = get_column_letter(4 + NCOL), get_column_letter(5 + NCOL)   # colunas do total e da parte
widths(ws, dict({"A": 2, "B": 5, "C": 36, LT: 10, LP: 10, get_column_letter(6 + NCOL): 2}, **{c: 10 for c in FC}))
title(ws, "Folha de verificação", "Lance os totais do período, por categoria e por coluna. As categorias vêm da aba Coleta.", LP)
label(ws, "B4", "Fator das colunas", merge="B4:C4")
inp(ws, "D4", merge="D4:G4", h="left")
put(ws, "H4", "Dia da semana, turno, área, máquina, produto.", f=font(9, i=True, c=MUTED), bg=GRAY, merge=f"H4:{LP}4")
ws.row_dimensions[4].height = 21.75
head(ws, "B6", "#")
head(ws, "C6", "Categoria")
for c in FC:
    inp(ws, f"{c}6", h="center")
    ws[f"{c}6"].font = font(10, True)
head(ws, f"{LT}6", "Total")
head(ws, f"{LP}6", "Parte do total")
ws.row_dimensions[6].height = 33
hint_row(ws, 7, [("B", ""), ("C", "Vem da aba Coleta"), (LT, "Calculado"), (LP, "Calculada")] + [(c, "Contagem") for c in FC], height=21.75)
note(ws, "D6", "Escreva aqui o nome de cada coluna: Seg, Ter, Qua… ou Manhã, Tarde, Noite… ou o nome de cada área.")
for k in range(NCAT):
    rr = F1 + k
    num(ws, f"B{rr}", k + 1)
    calc(ws, f"C{rr}", f'=IF({CO}!C{K1 + k}="","",{CO}!C{K1 + k})', h="left", b=False)
    for c in FC:
        inp(ws, f"{c}{rr}", h="center")
    calc(ws, f"{LT}{rr}", f'=IF(AND(C{rr}="",COUNT({FC[0]}{rr}:{FC[-1]}{rr})=0),"",SUM({FC[0]}{rr}:{FC[-1]}{rr}))')
    calc(ws, f"{LP}{rr}", f'=IF(OR({LT}{rr}="",${LT}${FT}=0),"",{LT}{rr}/${LT}${FT})', fmt=PCT, b=False)
    ws.row_dimensions[rr].height = 21.75
num(ws, f"B{FOU}", "—")
put(ws, f"C{FOU}", OUTROS, f=font(10, True), bg=GRAY)
for c in FC:
    inp(ws, f"{c}{FOU}", h="center")
calc(ws, f"{LT}{FOU}", f"=SUM({FC[0]}{FOU}:{FC[-1]}{FOU})")
calc(ws, f"{LP}{FOU}", f'=IF(${LT}${FT}=0,"",{LT}{FOU}/${LT}${FT})', fmt=PCT, b=False)
ws.row_dimensions[FOU].height = 21.75
label(ws, f"B{FT}", "Total", merge=f"B{FT}:C{FT}", h="right")
label(ws, f"B{FT+1}", "Parte do total", merge=f"B{FT+1}:C{FT+1}", h="right")
for c in FC + [LT]:
    calc(ws, f"{c}{FT}", f"=SUM({c}{F1}:{c}{FOU})")
    calc(ws, f"{c}{FT+1}", f'=IF(${LT}${FT}=0,"",{c}{FT}/${LT}${FT})', fmt=PCT, b=False)
put(ws, f"{LP}{FT}", "", bg=GRAY)
put(ws, f"{LP}{FT+1}", "", bg=GRAY)
for rr in (FT, FT + 1):
    ws.row_dimensions[rr].height = 21.75
dv_numero(ws, f"{FC[0]}{F1}:{FC[-1]}{FOU}", 0, 1000000, "Quantas ocorrências desta categoria, nesta coluna.")
ws.conditional_formatting.add(f"{FC[0]}{FT}:{FC[-1]}{FT}", FormulaRule(formula=[f"AND({FC[0]}{FT}>0,{FC[0]}{FT}=MAX(${FC[0]}${FT}:${FC[-1]}${FT}))"],
                              fill=PatternFill("solid", bgColor=BLUE_T, fgColor=BLUE_T)))
s = FT + 3
band(ws, s, "Resumo automático", LP)
for k, (text, formula, fmt) in enumerate([
    ("Ocorrências registradas", f"={LT}{FT}", None),
    ("Coluna com mais ocorrências", f'=IF({LT}{FT}=0,"",INDEX({FC[0]}6:{FC[-1]}6,MATCH(MAX({FC[0]}{FT}:{FC[-1]}{FT}),{FC[0]}{FT}:{FC[-1]}{FT},0))&" · "&TEXT(MAX({FC[0]}{FT}:{FC[-1]}{FT})/{LT}{FT},"0%"))', None),
    ("Ocorrências por 100 da base", f'=IF(OR({CO}!D8="",{LT}{FT}=0),"",100*{LT}{FT}/{CO}!D8)', "0.0"),
    ("Aviso", f'=IF({LT}{FT}=0,"Lance as contagens",IF(SUMPRODUCT((C{F1}:C{F2}="")*ISNUMBER({LT}{F1}:{LT}{F2}))>0,"Há contagem em linha sem categoria: cadastre na aba Coleta",'
              f'IF({LT}{FT}<{MIN_OBS},"Poucos dados: menos de {MIN_OBS} ocorrências","OK")))', None),
], 1):
    label(ws, f"B{s+k}", text, merge=f"B{s+k}:C{s+k}", h="right")
    calc(ws, f"D{s+k}", formula, fmt=fmt, merge=f"D{s+k}:{LP}{s+k}", sz=9 if text == "Aviso" else 10, b=text != "Aviso", h="left")
    ws.row_dimensions[s + k].height = 21.75
cf_warn(ws, f"D{s+4}:{LP}{s+4}", f"D{s+4}")
FAV = s + 4
ws.freeze_panes = "D8"
setup(ws, AMBER, f"B1:{LP}{s+4}", fit_height=True)

# ------------------------------------------------------------------ Pareto
ws = wb.create_sheet(PA)
widths(ws, {"A": 2, "B": 8, "C": 36, "D": 12, "E": 13, "F": 13, "G": 14, "H": 2, "I": 34, "J": 12, "K": 14, "L": 10, "M": 14, "N": 10, "O": 2})
title(ws, "Gráfico de Pareto", "A ordem, a parte do total e o acumulado são calculados a partir da aba Folha. Escolha só a base.", "G")
label(ws, "B4", "Base do Pareto", merge="B4:C4")
inp(ws, "D4", FREQ, merge="D4:E4", h="left")
put(ws, "F4", "Escolha na lista.", f=font(9, i=True, c=MUTED), bg=GRAY, merge="F4:G4")
dv_list(ws, "D4", [FREQ, PESO], "Frequência: a contagem. Frequência × peso: a contagem vezes o peso da aba Coleta.")
label(ws, "B5", "O que se conta", merge="B5:C5")
calc(ws, "D5", f'=IF({CO}!D5="","",{CO}!D5)', h="left", b=False, merge="D5:G5")
for rr in (4, 5):
    ws.row_dimensions[rr].height = 21.75
for col, text in zip("BCDEFG", ["Ordem", "Categoria", "Valor", "Parte do total", "Acumulado", "Classe"]):
    head(ws, f"{col}7", text)
put(ws, "I6", "Base de cálculo: não altere", f=font(9, i=True, c=MUTED), box=False, merge="I6:N6")
for col, text in zip("IJKLMN", ["Categoria, na ordem da folha", "Valor", "Chave de ordem", "Posição", "Linha do gráfico", "Corte"]):
    head(ws, f"{col}7", text, bg=MUTED)
ws.row_dimensions[7].height = 33
NV, TV = f"$J${P2 + 2}", f"$J${P2 + 3}"   # categorias com valor e total
for k in range(NCAT + 1):
    rr = P1 + k
    if k < NCAT:
        calc(ws, f"I{rr}", f'=IF({FO}!C{F1 + k}="","",{FO}!C{F1 + k})', h="left", b=False, sz=9)
        calc(ws, f"J{rr}", f'=IF({FO}!{LT}{F1 + k}="","",IF($D$4="{PESO}",{FO}!{LT}{F1 + k}*N({CO}!E{K1 + k}),{FO}!{LT}{F1 + k}))', b=False, sz=9)
        calc(ws, f"K{rr}", f'=IF(OR(I{rr}="",J{rr}="",J{rr}=0),"",J{rr}-ROW()/1000000)', b=False, sz=9, fmt="0.000000")
        calc(ws, f"L{rr}", f'=IF(K{rr}="","",RANK(K{rr},K${P1}:K${P2 - 1}))', b=False, sz=9)
    else:
        put(ws, f"I{rr}", OUTROS, f=font(9), bg=GRAY)
        calc(ws, f"J{rr}", f'=IF($D$4="{PESO}",{FO}!{LT}{FOU}*N({CO}!E{KO}),{FO}!{LT}{FOU})', b=False, sz=9)
        put(ws, f"K{rr}", "", bg=GRAY)
        put(ws, f"L{rr}", "", bg=GRAY)
    n = k + 1
    num(ws, f"B{rr}", n)
    calc(ws, f"C{rr}", f'=IF({n}<={NV},INDEX($I${P1}:$I${P2 - 1},MATCH({n},$L${P1}:$L${P2 - 1},0)),IF(AND({n}={NV}+1,$J${P2}>0),"{OUTROS}",""))', h="left")
    calc(ws, f"D{rr}", f'=IF(C{rr}="","",IF({n}<={NV},INDEX($J${P1}:$J${P2 - 1},MATCH({n},$L${P1}:$L${P2 - 1},0)),$J${P2}))', fmt="#,##0.##")
    calc(ws, f"E{rr}", f'=IF(D{rr}="","",D{rr}/{TV})', fmt=PCT, b=False)
    calc(ws, f"F{rr}", f'=IF(D{rr}="","",SUM(D${P1}:D{rr})/{TV})', fmt=PCT, b=False)
    calc(ws, f"G{rr}", f'=IF(D{rr}="","",IF({n}>{NV},"{OUTROS}",IF(ROUND(F{rr}-E{rr},6)<{CORTE},"{PRIOR}","{DEPOIS}")))', b=False, sz=9)
    # nas posições sem categoria, a linha do gráfico fica sem ponto: #N/D, escrito em cinza sobre cinza
    put(ws, f"M{rr}", f'=IF(D{rr}="",NA(),F{rr})', f=font(9, c=GRAY), bg=GRAY, h="center", fmt=PCT)
    put(ws, f"N{rr}", CORTE, f=font(9, c=MUTED), bg=GRAY, h="center", fmt=PCT)
    ws.row_dimensions[rr].height = 21.75
ws.conditional_formatting.add(f"M{P1}:M{P2}", FormulaRule(formula=[f"ISNUMBER(M{P1})"], font=Font(name="Arial", size=9, color=MUTED)))
label(ws, f"I{P2 + 2}", "Categorias com valor")
calc(ws, f"J{P2 + 2}", f"=COUNT(L{P1}:L{P2 - 1})", b=False, sz=9)
label(ws, f"I{P2 + 3}", "Total")
calc(ws, f"J{P2 + 3}", f"=SUM(J{P1}:J{P2})", b=False, sz=9, fmt="#,##0.##")
cf_texto(ws, f"G{P1}:G{P2}", f"G{P1}", [(PRIOR, BLUE_T)])
note(ws, "G7", f"Prioridade: a categoria começa antes de {round(100 * CORTE)}% do acumulado. A decisão é da equipe, e considera também o custo e a facilidade de agir.")
note(ws, "B4", "Use “Frequência × peso” quando as ocorrências não custam o mesmo. O peso de cada categoria fica na aba Coleta.")
s = P2 + 2
band(ws, s, "Resumo automático", "G")
NPR = f'COUNTIF(G{P1}:G{P2},"{PRIOR}")'
AVISO = (f'=IF({FO}!{LT}{FT}=0,"Lance as contagens na aba Folha",'
         f'IF(AND(D4="{PESO}",SUMPRODUCT(ISNUMBER({FO}!{LT}{F1}:{LT}{FOU})*({FO}!{LT}{F1}:{LT}{FOU}<>0)*({CO}!E{K1}:E{KO}=""))>0),"Falta o peso de uma categoria: veja a aba Coleta",'
         f'IF({FO}!{LT}{FT}<{MIN_OBS},"Poucos dados: menos de {MIN_OBS} ocorrências",'
         f'IF(D{s+4}>{OUTROS_MAX},"Outros acima de {round(100 * OUTROS_MAX)}%: abra em categorias",'
         f'IF(AND({NV}>=4,N(F{P1 + 1})<{PLANO}),"Pareto plano: estratifique por outro fator","OK")))))')
for k, (text, formula, fmt) in enumerate([
    ("Total", f'=IF({TV}=0,"",{TV})', "#,##0.##"),
    ("Primeira categoria", f'=IF(C{P1}="","",C{P1}&" · "&TEXT(E{P1},"0%"))', None),
    ("Prioridade", f'=IF({NPR}=0,"",{NPR}&" de "&{NV}&" categorias, com "&TEXT(INDEX(F{P1}:F{P2},{NPR}),"0%")&" do total")', None),
    ("Outros", f'=IF({TV}=0,"",$J${P2}/{TV})', PCT),
    ("Aviso", AVISO, None),
], 1):
    label(ws, f"B{s+k}", text, merge=f"B{s+k}:C{s+k}", h="right")
    calc(ws, f"D{s+k}", formula, fmt=fmt, merge=f"D{s+k}:G{s+k}", sz=9 if text == "Aviso" else 10, b=text != "Aviso", h="left")
    ws.row_dimensions[s + k].height = 21.75
cf_warn(ws, f"D{s+5}:G{s+5}", f"D{s+5}")
PAV = s + 5
pareto_chart(ws, f"B{s + 7}", P1, P2, 3, 5, 13, 14, width=18.6, height=9)
setup(ws, TEAL, f"B1:G{s + 26}", landscape=False, fit_height=True)

# ------------------------------------------------------------------ Antes e depois
ws = wb.create_sheet("Antes e depois")
widths(ws, {"A": 2, "B": 5, "C": 36, "D": 11, "E": 11, "F": 13, "G": 13, "H": 12, "I": 14, "J": 2})
title(ws, "Antes e depois", "Depois da ação, conte de novo com as mesmas categorias. O antes vem da aba Folha.", "I")
label(ws, "B4", "Período do antes", merge="B4:C4")
calc(ws, "D4", f'=IF({CO}!D6="","",{CO}!D6)', h="left", b=False, merge="D4:I4")
label(ws, "B5", "Período do depois", merge="B5:C5")
inp(ws, "D5", merge="D5:I5", h="left")
label(ws, "B6", "Base do antes", merge="B6:C6")
calc(ws, "D6", f'=IF({CO}!D8="","",{CO}!D8)', fmt="#,##0")
put(ws, "E6", "Vem da aba Coleta.", f=font(9, i=True, c=MUTED), bg=GRAY, merge="E6:I6")
label(ws, "B7", "Base do depois", merge="B7:C7")
inp(ws, "D7", h="center", fmt="#,##0")
calc(ws, "E7", f'=IF(AND(ISNUMBER(D6),ISNUMBER(D7)),IF(AND(D6>0,D7>0),"{TAXA}","Contagem"),"Contagem")', b=False, sz=9)
put(ws, "F7", "Como a comparação é feita. Com as duas bases, pela taxa.", f=font(9, i=True, c=MUTED), bg=GRAY, merge="F7:I7")
for rr in (4, 5, 6, 7):
    ws.row_dimensions[rr].height = 21.75
dv_numero(ws, "D7", 1, 1000000000, "Quantas oportunidades houve no período do depois.")
for col, text in zip("BCDEFGHI", ["#", "Categoria", "Antes", "Depois", "Antes, comparável", "Depois, comparável", "Variação", "Leitura"]):
    head(ws, f"{col}9", text)
ws.row_dimensions[9].height = 33
hint_row(ws, 10, [("B", ""), ("C", "Vem da aba Coleta"), ("D", "Da aba Folha"), ("E", "Nova contagem"), ("F", "Por 100, ou contagem"), ("G", "Por 100, ou contagem"),
                  ("H", "Calculada"), ("I", "Calculada")])
note(ws, "F9", "Com as duas bases informadas, esta coluna mostra as ocorrências por 100 da base. Sem elas, repete a contagem.")


def linha_ad(rr, total=False):
    nome = "B" if total else "C"   # na linha do total, o rótulo fica na coluna B
    calc(ws, f"F{rr}", f'=IF(OR({nome}{rr}="",D{rr}=""),"",IF($E$7="{TAXA}",100*D{rr}/$D$6,D{rr}))', fmt="0.0", b=total)
    calc(ws, f"G{rr}", f'=IF(OR({nome}{rr}="",E{rr}=""),"",IF($E$7="{TAXA}",100*E{rr}/$D$7,E{rr}))', fmt="0.0", b=total)
    calc(ws, f"H{rr}", f'=IF(OR(F{rr}="",G{rr}=""),"",IF(F{rr}=0,"",(G{rr}-F{rr})/F{rr}))', fmt='+0%;-0%;0%', b=total)
    calc(ws, f"I{rr}", f'=IF(OR(F{rr}="",G{rr}=""),"",IF(AND(F{rr}=0,G{rr}>0),"Nova",IF(G{rr}<F{rr},"Caiu",IF(G{rr}>F{rr},"Subiu","Igual"))))', b=False, sz=9)
    ws.row_dimensions[rr].height = 21.75


for k in range(NCAT):
    rr = A1 + k
    num(ws, f"B{rr}", k + 1)
    calc(ws, f"C{rr}", f'=IF({FO}!C{F1 + k}="","",{FO}!C{F1 + k})', h="left", b=False)
    calc(ws, f"D{rr}", f'=IF({FO}!{LT}{F1 + k}="","",{FO}!{LT}{F1 + k})', b=False)
    inp(ws, f"E{rr}", h="center")
    linha_ad(rr)
num(ws, f"B{AO}", "—")
put(ws, f"C{AO}", OUTROS, f=font(10, True), bg=GRAY)
calc(ws, f"D{AO}", f"={FO}!{LT}{FOU}", b=False)
inp(ws, f"E{AO}", h="center")
linha_ad(AO)
label(ws, f"B{AT}", "Total", merge=f"B{AT}:C{AT}", h="right")
calc(ws, f"D{AT}", f"=SUM(D{A1}:D{AO})")
calc(ws, f"E{AT}", f'=IF(COUNT(E{A1}:E{AO})=0,"",SUM(E{A1}:E{AO}))')
linha_ad(AT, total=True)
dv_numero(ws, f"E{A1}:E{AO}", 0, 1000000, "Quantas ocorrências desta categoria, no período do depois.")
cf_texto(ws, f"I{A1}:I{AT}", f"I{A1}", [("Caiu", GREEN), ("Subiu", RED), ("Nova", RED), ("Igual", YELLOW)])
s = AT + 2
band(ws, s, "Resumo automático", "I")
for k, (text, formula, fmt) in enumerate([
    ("Variação do total", f'=IF(H{AT}="","",H{AT})', '+0%;-0%;0%'),
    ("Primeira categoria do antes", f'=IF(OR({PA}!C{P1}="",COUNT(E{A1}:E{AO})=0),"",{PA}!C{P1}&IF(IFERROR(INDEX(I{A1}:I{AO},MATCH({PA}!C{P1},C{A1}:C{AO},0)),"")="","",": "&LOWER(IFERROR(INDEX(I{A1}:I{AO},MATCH({PA}!C{P1},C{A1}:C{AO},0)),""))))', None),
    ("Primeira categoria de agora", f'=IF(COUNT(E{A1}:E{A2})=0,"",INDEX(C{A1}:C{A2},MATCH(MAX(E{A1}:E{A2}),E{A1}:E{A2},0)))', None),
    ("Aviso", f'=IF({FO}!{LT}{FT}=0,"Lance as contagens na aba Folha",IF(COUNT(E{A1}:E{AO})=0,"Lance as contagens do depois",'
              f'IF(COUNT(E{A1}:E{AO})<COUNT(D{A1}:D{AO}),"Faltam contagens do depois: escreva 0 onde não houve ocorrência",'
              f'IF(E7<>"{TAXA}","Informe as duas bases para comparar a taxa",'
              f'IF(COUNTIF(I{A1}:I{AO},"Subiu")+COUNTIF(I{A1}:I{AO},"Nova")>0,"Há categoria que subiu: veja a coluna Leitura","OK")))))', None),
], 1):
    label(ws, f"B{s+k}", text, merge=f"B{s+k}:C{s+k}", h="right")
    calc(ws, f"D{s+k}", formula, fmt=fmt, merge=f"D{s+k}:I{s+k}", sz=9 if text == "Aviso" else 10, b=text != "Aviso", h="left")
    ws.row_dimensions[s + k].height = 21.75
cf_warn(ws, f"D{s+4}:I{s+4}", f"D{s+4}")
par_chart(ws, f"B{s + 6}", A1, AO, 3, 6, 7, width=19.5, height=8.5)
setup(ws, REDC, f"B1:I{s + 24}", landscape=False, fit_height=True)

# ------------------------------------------------------------------ Checklist
ws = wb.create_sheet("Checklist")
widths(ws, {"A": 2, "B": 5, "C": 70, "D": 14, "E": 44, "F": 2})
title(ws, "Checklist do Pareto e da folha de verificação", "Escolha o status de cada item na lista suspensa (coluna em amarelo-claro).", "E")
for col, text in zip("BCDE", ["#", "Item de verificação", "Status", "Observações"]):
    head(ws, f"{col}4", text)
ws.row_dimensions[4].height = 21.75
for k, text in enumerate(CHECK):
    rr = 5 + k
    num(ws, f"B{rr}", k + 1)
    put(ws, f"C{rr}", text, bg=GRAY)
    inp(ws, f"D{rr}", h="center")
    inp(ws, f"E{rr}")
    ws.row_dimensions[rr].height = 30
dv_list(ws, "D5:D16", YESNO, "Sim, Parcial ou Não")
cf_equal(ws, "D5:D16", YESNO_CF)
band(ws, 18, "Resumo automático", "E")
summary(ws, 19, "Itens atendidos (Sim)", '=COUNTIF(D5:D16,"Sim")', "C", "D")
summary(ws, 20, "Itens parciais", '=COUNTIF(D5:D16,"Parcial")', "C", "D")
summary(ws, 21, "Itens não atendidos", '=COUNTIF(D5:D16,"Não")', "C", "D")
summary(ws, 22, "Itens sem resposta", "=COUNTBLANK(D5:D16)", "C", "D")
summary(ws, 23, "Percentual atendido", "=D19/ROWS(D5:D16)", "C", "D", fmt="0%")
summary(ws, 24, "Situação", '=IF(D19=ROWS(D5:D16),"Validado",IF(D22>0,"Em andamento","Com pendências"))', "C", "D")
setup(ws, PURPLE, "B1:E24", landscape=False, fit_height=True)


# ------------------------------------------------------------------ Exemplos
LAST = "L"


def sub(ws, rr, cells, height=24):
    for a, b, t in cells:
        put(ws, f"{a}{rr}", t, f=font(10, True), bg=GRAY, h="center", merge=b and f"{a}{rr}:{b}{rr}")
    ws.row_dimensions[rr].height = height


def cabecalho(ws, H):
    rr = 4
    for rot, val in [("Organização", H["org"]), ("O que se conta", H["conta"]), ("Período e fonte", H["periodo"]), ("Como se registra", H["como"]),
                     ("Análise", f'{H["data"]:%d/%m/%Y}, por {H["por"][0].lower() + H["por"][1:]}. Base: {mil(H["base"])} {H["base_nome"]}.'), ("Origem", H["origem"])]:
        label(ws, f"B{rr}", rot, merge=f"B{rr}:C{rr}")
        inp(ws, f"D{rr}", val, merge=f"D{rr}:{LAST}{rr}", bg=WHITE, h="left")
        ws.row_dimensions[rr].height = alt([(val, 95)], minimo=19.5)
        rr += 1
    return rr + 1


def folha(ws, rr, ex):
    """A folha do período: categorias nas linhas, o segundo fator nas colunas."""
    cols = ex["cols"]
    cl = [get_column_letter(4 + j) for j in range(len(cols))]
    lt = get_column_letter(4 + len(cols))
    band(ws, rr, f'Folha do período · {ex["head"]["fator"].lower()} nas colunas', LAST, color=AMBER)
    rr += 1
    sub(ws, rr, [("B", "C", "Categoria")] + [(c, None, n) for c, n in zip(cl, cols)] + [(lt, None, "Total")], height=30 if max(len(c) for c in cols) > 10 else 24)
    f1 = rr + 1
    for nome, _, v in ex["folha"]:
        rr += 1
        put(ws, f"B{rr}", nome, f=font(10, True), merge=f"B{rr}:C{rr}")
        for c, x in zip(cl, v):
            put(ws, f"{c}{rr}", x, h="center")
        calc(ws, f"{lt}{rr}", f"=SUM({cl[0]}{rr}:{cl[-1]}{rr})")
        ws.row_dimensions[rr].height = 21.75
    f2 = rr
    rr += 1
    label(ws, f"B{rr}", "Total", merge=f"B{rr}:C{rr}", h="right")
    for c in cl + [lt]:
        calc(ws, f"{c}{rr}", f"=SUM({c}{f1}:{c}{f2})")
    ws.row_dimensions[rr].height = 21.75
    ws.conditional_formatting.add(f"{cl[0]}{rr}:{cl[-1]}{rr}", FormulaRule(formula=[f"{cl[0]}{rr}=MAX(${cl[0]}${rr}:${cl[-1]}${rr})"], fill=PatternFill("solid", bgColor=BLUE_T, fgColor=BLUE_T)))
    rr += 1
    label(ws, f"B{rr}", "Parte do total", merge=f"B{rr}:C{rr}", h="right")
    for c in cl + [lt]:
        calc(ws, f"{c}{rr}", f"={c}{rr - 1}/${lt}${rr - 1}", fmt=PCT, b=False)
    ws.row_dimensions[rr].height = 21.75
    return rr + 2


def tabela_pareto(ws, rr, rows, titulo, col1, unidade, base=None, base_nome=None, bases=None):
    """Pareto já ordenado, com fórmulas para a parte, o acumulado e a classe. Devolve a linha seguinte e as linhas da tabela.
    Com `bases` (nome -> base de cada linha), a coluna J traz a base da linha, e a K, a taxa por 100."""
    rows = ordenar(rows)
    band(ws, rr, titulo, LAST, color=TEAL)
    rr += 1
    sub(ws, rr, [("B", None, "Ordem"), ("C", None, col1), ("D", None, unidade), ("E", None, "Parte"), ("F", None, "Acumulado"), ("G", "H", "Classe"), ("I", None, "Corte")]
        + ([("J", None, base_nome[0].upper() + base_nome[1:]), ("K", None, f"Por 100 {base_nome}")] if bases else [])
        + ([("J", "K", f"Por 100 {base_nome}")] if base else []), height=30)
    p1, p2 = rr + 1, rr + len(rows)
    for k, (nome, v) in enumerate(rows, 1):
        rr += 1
        num(ws, f"B{rr}", k)
        put(ws, f"C{rr}", nome, f=font(10, True))
        put(ws, f"D{rr}", v, h="center")
        calc(ws, f"E{rr}", f"=D{rr}/SUM(D${p1}:D${p2})", fmt=PCT, b=False)
        calc(ws, f"F{rr}", f"=SUM(D${p1}:D{rr})/SUM(D${p1}:D${p2})", fmt=PCT, b=False)
        calc(ws, f"G{rr}", f'=IF(C{rr}="{OUTROS}","{OUTROS}",IF(ROUND(F{rr}-E{rr},6)<{CORTE},"{PRIOR}","{DEPOIS}"))', b=False, sz=9, merge=f"G{rr}:H{rr}")
        put(ws, f"I{rr}", CORTE, f=font(9, c=MUTED), bg=GRAY, h="center", fmt=PCT)
        if base:
            calc(ws, f"J{rr}", f"=100*D{rr}/{base}", fmt="0.0", b=False, merge=f"J{rr}:K{rr}")
        if bases:
            put(ws, f"J{rr}", bases[nome], h="center")
            calc(ws, f"K{rr}", f"=100*D{rr}/J{rr}", fmt="0.0", b=False)
        ws.row_dimensions[rr].height = 21.75
    cf_texto(ws, f"G{p1}:H{p2}", f"G{p1}", [(PRIOR, BLUE_T)])
    rr += 1
    label(ws, f"B{rr}", "Total", merge=f"B{rr}:C{rr}", h="right")
    calc(ws, f"D{rr}", f"=SUM(D{p1}:D{p2})")
    calc(ws, f"E{rr}", f"=SUM(E{p1}:E{p2})", fmt=PCT)
    for c in "FI":
        put(ws, f"{c}{rr}", "", bg=GRAY)
    put(ws, f"G{rr}", "", bg=GRAY, merge=f"G{rr}:H{rr}")
    if base:
        calc(ws, f"J{rr}", f"=100*D{rr}/{base}", fmt="0.0", merge=f"J{rr}:K{rr}")
    if bases:
        calc(ws, f"J{rr}", f"=SUM(J{p1}:J{p2})")
        calc(ws, f"K{rr}", f"=100*D{rr}/J{rr}", fmt="0.0")
    ws.row_dimensions[rr].height = 21.75
    return rr + 2, p1, p2


def resumo(ws, rr, itens):
    band(ws, rr, "Resumo automático", LAST)
    for k, (text, formula, fmt) in enumerate(itens, 1):
        label(ws, f"B{rr+k}", text, merge=f"B{rr+k}:E{rr+k}", h="right")
        calc(ws, f"F{rr+k}", formula, fmt=fmt, merge=f"F{rr+k}:{LAST}{rr+k}", h="left")
        ws.row_dimensions[rr + k].height = 21.75
    return rr + len(itens)


def base_ex(ws, ex):
    widths(ws, {"A": 2, "B": 8, "C": 36, "D": 12, "E": 12, "F": 12, "G": 12, "H": 12, "I": 10, "J": 10, "K": 10, "L": 10, "M": 2})
    title(ws, "Pareto e folha de verificação", "Exemplo preenchido, para consulta. Use as abas Coleta, Folha, Pareto e Antes e depois para a sua organização.", LAST)
    return cabecalho(ws, ex["head"])


def exemplo1(ws, ex):
    H = ex["head"]
    rr = base_ex(ws, ex)
    rr = folha(ws, rr, ex)
    ws.row_breaks.append(Break(id=rr - 1))
    rr, p1, p2 = tabela_pareto(ws, rr, linhas(ex), "Pareto por motivo do atraso", "Motivo principal", "Atrasos", H["base"], H["base_nome"])
    pareto_chart(ws, f"B{rr}", p1, p2, 3, 5, 6, 9, width=22, height=8.5)
    rr += 19
    ws.row_breaks.append(Break(id=rr - 1))
    dp = ex["depois"]
    band(ws, rr, f'Antes e depois · depois: {dp["periodo"].lower()}, com {mil(dp["base"])} {H["base_nome"]}', LAST, color=REDC)
    rr += 1
    sub(ws, rr, [("B", "C", "Motivo principal"), ("D", None, "Antes"), ("E", None, "Depois"), ("F", None, "Antes por 100"), ("G", None, "Depois por 100"), ("H", None, "Variação"),
                 ("I", "J", "Leitura")], height=30)
    a1 = rr + 1
    depois = dict(zip((n for n, _, _ in ex["folha"]), dp["valores"]))
    for nome, v in ordenar(linhas(ex)):
        rr += 1
        put(ws, f"B{rr}", nome, f=font(10, True), merge=f"B{rr}:C{rr}")
        put(ws, f"D{rr}", v, h="center")
        put(ws, f"E{rr}", depois[nome], h="center")
        calc(ws, f"F{rr}", f"=100*D{rr}/{H['base']}", fmt="0.0", b=False)
        calc(ws, f"G{rr}", f"=100*E{rr}/{dp['base']}", fmt="0.0", b=False)
        calc(ws, f"H{rr}", f"=(G{rr}-F{rr})/F{rr}", fmt='+0%;-0%;0%')
        calc(ws, f"I{rr}", f'=IF(G{rr}<F{rr},"Caiu",IF(G{rr}>F{rr},"Subiu","Igual"))', b=False, sz=9, merge=f"I{rr}:J{rr}")
        ws.row_dimensions[rr].height = 21.75
    a2 = rr
    cf_texto(ws, f"I{a1}:J{a2}", f"I{a1}", [("Caiu", GREEN), ("Subiu", RED), ("Igual", YELLOW)])
    rr += 1
    label(ws, f"B{rr}", "Total", merge=f"B{rr}:C{rr}", h="right")
    calc(ws, f"D{rr}", f"=SUM(D{a1}:D{a2})")
    calc(ws, f"E{rr}", f"=SUM(E{a1}:E{a2})")
    calc(ws, f"F{rr}", f"=100*D{rr}/{H['base']}", fmt="0.0")
    calc(ws, f"G{rr}", f"=100*E{rr}/{dp['base']}", fmt="0.0")
    calc(ws, f"H{rr}", f"=(G{rr}-F{rr})/F{rr}", fmt='+0%;-0%;0%')
    put(ws, f"I{rr}", "", bg=GRAY, merge=f"I{rr}:J{rr}")
    ws.row_dimensions[rr].height = 21.75
    at = rr
    rr += 2
    npr = f'COUNTIF(G{p1}:H{p2},"{PRIOR}")'
    rr = resumo(ws, rr, [
        ("Atrasos no período do antes", f"=D{p2 + 1}", None),
        ("Primeira categoria", f'=C{p1}&" · "&TEXT(E{p1},"0%")', None),
        ("Prioridade", f'={npr}&" categorias, com "&TEXT(INDEX(F{p1}:F{p2},{npr}),"0%")&" do total"', None),
        ("Atrasos por 100 entregas, antes e depois", f'=TEXT(F{at},"0,0")&" → "&TEXT(G{at},"0,0")', None),
        ("Primeira categoria de agora", f"=INDEX(B{a1}:B{a2},MATCH(MAX(E{a1}:E{a2 - 1}),E{a1}:E{a2 - 1},0))", None),
    ])
    setup(ws, MUTED, f"B1:{LAST}{rr}")


def exemplo2(ws, ex):
    H = ex["head"]
    rr = base_ex(ws, ex)
    rr = folha(ws, rr, ex)
    ws.row_breaks.append(Break(id=rr - 1))
    rr, p1, p2 = tabela_pareto(ws, rr, linhas(ex), "Pareto por motivo da devolução", "Motivo", "Devoluções", H["base"], H["base_nome"])
    pareto_chart(ws, f"B{rr}", p1, p2, 3, 5, 6, 9, width=22, height=8.5)
    rr += 19
    ws.row_breaks.append(Break(id=rr - 1))
    rr, q1, q2 = tabela_pareto(ws, rr, colunas(ex), "Pareto por área requisitante · as colunas da mesma folha, com as requisições de cada área",
                               "Área requisitante", "Devoluções", base_nome=H["base_nome"], bases=dict(zip(ex["cols"], ex["req"])))
    npr = f'COUNTIF(G{p1}:H{p2},"{PRIOR}")'
    rr = resumo(ws, rr, [
        ("Devoluções no período", f"=D{p2 + 1}", None),
        ("Primeira categoria", f'=C{p1}&" · "&TEXT(E{p1},"0%")', None),
        ("Prioridade", f'={npr}&" categorias, com "&TEXT(INDEX(F{p1}:F{p2},{npr}),"0%")&" do total"', None),
        ("Campos em branco: as três primeiras categorias", f"=F{p1 + 2}", PCT),
        ("Área com mais devoluções", f'=C{q1}&" · "&TEXT(E{q1},"0%")', None),
        ("Área com a maior taxa de devolução", f'=INDEX(C{q1}:C{q2},MATCH(MAX(K{q1}:K{q2}),K{q1}:K{q2},0))&" · "&TEXT(MAX(K{q1}:K{q2}),"0,0")&" por 100 requisições"', None),
    ])
    setup(ws, MUTED, f"B1:{LAST}{rr}")


exemplo1(wb.create_sheet("Exemplo 1 - Pizzaria"), EX1)
exemplo2(wb.create_sheet("Exemplo 2 - Compras"), EX2)

wb.active = 1
wb.save(OUT)
print("ok", OUT, "| aviso da folha: D%d | aviso do Pareto: D%d" % (FAV, PAV))
