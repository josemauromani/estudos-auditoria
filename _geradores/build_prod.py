# -*- coding: utf-8 -*-
"""Gera Producao-modelo.xlsx no mesmo padrão visual das planilhas anteriores da série."""
import math
import os
import sys

_here = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, _here)
# funções de formatação compartilhadas com o gerador do PDCA (bloco inicial do arquivo)
_src = open(os.path.join(_here, "build_pdca.py"), encoding="utf-8").read()
exec(_src[: _src.index("STATUS = [")])

from openpyxl.chart import BarChart, Series  # noqa: E402
from openpyxl.chart.label import DataLabelList  # noqa: E402
from prod_data import CHECK, DENTRO, EX1, EX2, FORA, NAO, SEMREACAO, SIM, TIPOS, atributo  # noqa: E402

BLUE, TEAL, AMBER, PURPLE, REDC = "2B5C8A", "1E7B73", "A96A12", "7A4A9A", "B0413E"
S1_T, S2_T, S3_T = "DCE8F3", "F8EBCB", "F5DEDC"
YESNO = ["Sim", "Parcial", "Não"]
YESNO_CF = [("Sim", GREEN), ("Parcial", YELLOW), ("Não", RED)]
NK, NR, NE, NP, NM = 20, 150, 10, 8, 12
PC, RG, RT, MU = "'Plano de controle'", "Registros", "Rastreabilidade", "'Mudanças'"
K1, K2 = 10, 10 + NK - 1        # controles, na aba Plano de controle
R1, R2 = 7, 7 + NR - 1          # verificações, na aba Registros
E1, E2 = 7, 7 + NE - 1          # etapas, na aba Rastreabilidade
P1, P2 = E2 + 5, E2 + 5 + NP - 1  # itens de terceiros, na aba Rastreabilidade
M1, M2 = 7, 7 + NM - 1          # mudanças, na aba Mudanças
IDS = [f"K{k + 1}" for k in range(NK)]
GERAL = "#,##0.##"
SEMPLANO, ATRIB = "Controle sem plano", "Conforme ou não"
L_SEM, L_COBRAR, L_REP, L_UM, L_OK = "Sem registro", "Fora sem reação: cobrar", "Desvio repetido: analisar a causa", "Um desvio, com reação", "Dentro do critério"


def note(ws, ref, text):
    c = Comment(text, "Modelo Produção")
    c.width, c.height = 300, 120
    ws[ref].comment = c


def cf_warn(ws, rng_, first, ok_values=("OK",)):
    cond = ",".join(f'{first}="{v}"' for v in ok_values)
    ws.conditional_formatting.add(rng_, FormulaRule(formula=[f"OR({cond})"], fill=PatternFill("solid", bgColor=GREEN, fgColor=GREEN)))
    ws.conditional_formatting.add(rng_, FormulaRule(formula=[f"AND(LEN({first})>0,NOT(OR({cond})))"],
                                  fill=PatternFill("solid", bgColor=YELLOW, fgColor=YELLOW)))


def cf_texto(ws, rng_, first, pares, resto=None):
    for t, cor in pares:
        ws.conditional_formatting.add(rng_, FormulaRule(formula=[f'{first}="{t}"'], stopIfTrue=True, fill=PatternFill("solid", bgColor=cor, fgColor=cor)))
    if resto:
        ws.conditional_formatting.add(rng_, FormulaRule(formula=[f"LEN({first})>0"], fill=PatternFill("solid", bgColor=resto, fgColor=resto)))


def hint_row(ws, row, cells, height=30):
    for ref, text in cells:
        put(ws, f"{ref}{row}", text, f=font(9, i=True, c=MUTED), bg=GRAY, h="center")
    ws.row_dimensions[row].height = height


def alt(pares, minimo=21.75, linha=12):
    n = max(max(1, math.ceil(len(str(t)) / max(c, 1))) for t, c in pares)
    return max(minimo, n * linha + 7)


def f_crit(d, mn, mx, un):
    """Texto do critério de um controle: faixa, mínimo, máximo ou atributo."""
    return (f'IF({d}="","",IF(AND({mn}="",{mx}=""),"{ATRIB}",TRIM(IF({mx}="","mín. "&{mn},IF({mn}="","máx. "&{mx},{mn}&" a "&{mx}))&" "&{un})))')


def f_leitura(mn, mx, val, conf):
    """Leitura de uma verificação: a regra do módulo 7 do treinamento."""
    return (f'IF(AND({mn}="",{mx}=""),IF({conf}="","Falta o resultado",IF({conf}="{SIM}","{DENTRO}","{FORA}")),'
            f'IF({val}="","Falta o valor",IF(AND(OR({mn}="",{val}>={mn}),OR({mx}="",{val}<={mx})),"{DENTRO}","{FORA}")))')


def busca(col, key, r1, r2, aba=PC, texto=True):
    """Valor da coluna `col` da tabela de controles, pela chave em `key`. Célula vazia volta como texto vazio."""
    pre = f"{aba}!" if aba else ""
    ix = f"INDEX({pre}${col}${r1}:${col}${r2},MATCH({key},{pre}$B${r1}:$B${r2},0))"
    return f'IFERROR({ix}&"","")' if texto else f'IFERROR(IF(ISNUMBER({ix}),{ix},""),"")'


def chart_style(ch):
    ch.style = 2
    ch.title = None
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ch.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill="D3DBE0", w=9525))
    ch.x_axis.graphicalProperties = GraphicalProperties(ln=LineProperties(solidFill="D3DBE0", w=9525))
    ch.y_axis.graphicalProperties = GraphicalProperties(ln=LineProperties(noFill=True))


def barras(ws, anchor, r1, r2, c_cat, c_val, width=19, height=9.5):
    """Barras horizontais com a parte das verificações de cada controle que ficou dentro do critério."""
    ch = BarChart()
    ch.type = "bar"
    ch.gapWidth = 60
    ch.height, ch.width = height, width
    chart_style(ch)
    ch.legend = None
    s = Series(Reference(ws, min_col=c_val, min_row=r1, max_row=r2), title="Dentro do critério")
    s.graphicalProperties.solidFill = BLUE
    s.graphicalProperties.line.noFill = True
    s.dLbls = DataLabelList()
    s.dLbls.showVal = True
    s.dLbls.numFmt = "0%"
    s.dLbls.showSerName = s.dLbls.showCatName = s.dLbls.showLegendKey = False
    ch.append(s)
    ch.set_categories(Reference(ws, min_col=c_cat, min_row=r1, max_row=r2))
    ch.y_axis.scaling.min, ch.y_axis.scaling.max = 0, 1
    ch.y_axis.majorUnit = 0.25
    ch.y_axis.number_format = "0%"
    ch.x_axis.scaling.orientation = "maxMin"
    ws.add_chart(ch, anchor)


def resumo_bloco(ws, s, last, itens, lab_to, val, aviso, ok_values=("OK",)):
    """Faixa de resumo com rótulo à esquerda e valor mesclado até `last`. Devolve a linha do aviso."""
    band(ws, s, "Resumo automático", last)
    for k, (text, formula, fmt) in enumerate(itens, 1):
        label(ws, f"B{s+k}", text, merge=f"B{s+k}:{lab_to}{s+k}", h="right")
        calc(ws, f"{val}{s+k}", formula, fmt=fmt, merge=f"{val}{s+k}:{last}{s+k}", sz=9 if text == aviso else 10, b=text != aviso, h="left")
        ws.row_dimensions[s + k].height = 21.75
    r = s + len(itens)
    cf_warn(ws, f"{val}{r}:{last}{r}", f"{val}{r}", ok_values=ok_values)
    return r


# ================================================================== pasta de trabalho
wb = Workbook()
wb.properties.title = "Controle de produção e de serviço — Modelo"
wb.properties.creator = "Modelo Produção"
wb.properties.language = "pt-BR"

# ------------------------------------------------------------------ Instruções
ws = wb.active
ws.title = "Instruções"
widths(ws, {"A": 2, "B": 24, "C": 92, "D": 2})
title(ws, "Controle de produção e de serviço — Como usar esta planilha",
      "Modelo para escrever o plano de controle, registrar as verificações, cuidar da identificação e controlar as mudanças.", "C")
r = 4


def section(text):
    global r
    band(ws, r, text, "C", sz=11)
    r += 1


def line(k, v, kbg=GRAY, vbg=None, kf=None, kh="left", height=None):
    global r
    put(ws, f"B{r}", k, f=kf or font(10, True), bg=kbg, h=kh)
    put(ws, f"C{r}", v, bg=vbg)
    ws.row_dimensions[r].height = height or (19.5 if len(v) <= 98 else (31.5 if len(v) <= 196 else 45.75))
    r += 1


section("Legenda: onde preencher")
line("Amarelo-claro", "Células de entrada. É aqui que você digita.", vbg=INPUT)
line("Cinza", "Células calculadas ou fixas (rótulos, critérios, leituras, avisos). Não altere.", vbg=GRAY)
line("Cinza em itálico", "Linha de exemplo (marcada com “Ex.”). Mostra o formato esperado e não entra nas contagens.", vbg=GRAY)
line("Listas suspensas", "Tipo do controle, código do controle, conforme ou não, dono do item e checklist aceitam apenas as opções da lista.")
line("Comentários", "Células com triângulo vermelho no canto trazem uma dica de preenchimento. Passe o mouse sobre elas.")
r += 1
section("As regras de cálculo")
for k, text in [
    ("Critério", "Com mínimo e máximo: faixa. Só com um deles: mínimo ou máximo. Sem nenhum: o controle é por atributo, e o resultado é conforme ou não."),
    (DENTRO, "Controle por medida: o valor está entre o mínimo e o máximo. Controle por atributo: a resposta é Sim."),
    (FORA, "O valor está abaixo do mínimo ou acima do máximo, ou a resposta é Não."),
    (SEMREACAO, "A leitura é Fora, e a coluna Reação tomada está vazia. É o aviso mais importante da planilha."),
    (L_SEM, "O controle está no plano, e não há verificação lançada na aba Registros."),
    (L_REP, "O mesmo controle tem mais de um resultado fora do critério. A reação não basta: abrir a análise de causa."),
    ("Mudança a rever", "A mudança foi escrita sem a análise feita antes de mudar, ou sem quem autorizou."),
]:
    line(k, text)
r += 1
section("Ordem recomendada de preenchimento")
for k, text in enumerate([
    "Aba Plano de controle: escreva a organização, o processo, o produto e a revisão do plano.",
    "Aba Plano de controle: para cada controle, a etapa, o que controlar, o tipo, o mínimo e o máximo, como medir, a frequência, quem mede, o registro e a reação.",
    "Aba Registros: lance cada verificação, com data, lote ou pedido, código do controle e resultado. Se o resultado sair do critério, escreva a reação tomada.",
    "Aba Rastreabilidade: descreva a identificação, a situação e a preservação em cada etapa, e liste o que pertence a clientes e a fornecedores.",
    "Aba Mudanças: registre cada mudança, com o motivo, a análise, quem autorizou e as ações decorrentes.",
    "Aba Painel: leia os desvios por controle, os desvios sem reação e os controles sem registro. Depois, valide na aba Checklist.",
], 1):
    line(f"Passo {k}", text)
r += 1
section("Abas da planilha")
for k, text in [
    ("Plano de controle", f"Até {NK} controles. Calcula o texto do critério e a conferência de cada um."),
    (RG, f"Até {NR} verificações, ligadas aos controles pelo código. Calcula a leitura e a conferência."),
    (RT, f"Até {NE} etapas, com identificação, situação e preservação, e até {NP} itens de clientes e de fornecedores."),
    ("Mudanças", f"Até {NM} mudanças. Calcula a conferência: análise e autorização."),
    ("Painel", "Os indicadores gerais, a leitura de cada controle e o gráfico das verificações dentro do critério."),
    ("Checklist", "Doze verificações de qualidade do controle, com percentual de conclusão."),
    ("Exemplo 1 - Pizzaria", "O plano da cozinha e da expedição, com sete controles e os registros de uma semana."),
    ("Exemplo 2 - Indústria", "O plano da extrusora 3, com sete controles e os registros de cinco lotes."),
]:
    line(k, text)
r += 1
section("Regras e premissas do modelo")
for k, text in [
    ("Poucos controles", f"O modelo aceita {NK}. O usual é de cinco a dez por processo: os que pesam para o cliente, para a lei e para a etapa seguinte."),
    ("Código do controle", f"Os códigos K1 a K{NK} são fixos. Na aba Registros, cada verificação é ligada ao controle pelo código."),
    ("Mínimo e máximo", "Digite só o número, na unidade informada. Para um controle por atributo, deixe os dois vazios e responda Sim ou Não no registro."),
    ("Exemplos", "Os dados das abas de exemplo são ilustrativos, criados para este material."),
]:
    line(k, text)
setup(ws, INK, f"B1:C{r}", landscape=False, fit_height=True)

# ------------------------------------------------------------------ Plano de controle
ws = wb.create_sheet("Plano de controle")
widths(ws, {"A": 2, "B": 6, "C": 22, "D": 32, "E": 11, "F": 9, "G": 9, "H": 10, "I": 30, "J": 20, "K": 18, "L": 20, "M": 40, "N": 18, "O": 24, "P": 2})
title(ws, "Plano de controle", "Uma linha por controle: o que se confere, com que critério, como, quando, por quem, onde se registra e o que fazer se sair do critério.", "O")
for k, text in enumerate(["Organização e processo", "Produto ou serviço", "Revisão do plano e data"]):
    label(ws, f"B{4 + k}", text, merge=f"B{4 + k}:C{4 + k}")
    inp(ws, f"D{4 + k}", merge=f"D{4 + k}:O{4 + k}", h="left")
    ws.row_dimensions[4 + k].height = 21.75
for col, text in zip("BCDEFGHIJKLMNO", ["#", "Etapa", "O que controlar", "Tipo", "Mínimo", "Máximo", "Unidade", "Como medir", "Frequência", "Quem mede", "Registro",
                                         "Reação se sair do critério", "Critério", "Conferência"]):
    head(ws, f"{col}8", text)
ws.row_dimensions[8].height = 33
hint_row(ws, 9, [("B", ""), ("C", "Do recebimento à entrega"), ("D", "Característica ou parâmetro"), ("E", "Lista"), ("F", "Número"), ("G", "Número"), ("H", "°C, g, mm"),
                 ("I", "Método e instrumento"), ("J", "Quando e quantos"), ("K", "Uma função"), ("L", "Onde fica o resultado"), ("M", "Com o produto e com o processo"), ("N", "Calculado"),
                 ("O", "Calculada")])
for k in range(NK):
    rr = K1 + k
    num(ws, f"B{rr}", IDS[k])
    for col in "CDIJKLM":
        inp(ws, f"{col}{rr}")
    inp(ws, f"E{rr}", h="center")
    inp(ws, f"F{rr}", h="center", fmt=GERAL)
    inp(ws, f"G{rr}", h="center", fmt=GERAL)
    inp(ws, f"H{rr}", h="center")
    calc(ws, f"N{rr}", "=" + f_crit(f"D{rr}", f"F{rr}", f"G{rr}", f"H{rr}"), b=False, sz=9)
    calc(ws, f"O{rr}", f'=IF(D{rr}="","",IF(C{rr}="","Falta a etapa",IF(E{rr}="","Falta o tipo",IF(AND(F{rr}<>"",G{rr}<>"",F{rr}>G{rr}),"Mínimo maior que o máximo",'
         f'IF(I{rr}="","Falta como medir",IF(J{rr}="","Falta a frequência",IF(K{rr}="","Falta quem mede",IF(L{rr}="","Falta o registro",IF(M{rr}="","Falta a reação","OK")))))))))', b=False, sz=9)
    ws.row_dimensions[rr].height = 30
dv_list(ws, f"E{K1}:E{K2}", TIPOS, "Produto: uma característica do que é feito. Processo: uma condição que faz o produto sair certo.")
dv_number(ws, f"F{K1}:G{K2}")
cf_warn(ws, f"O{K1}:O{K2}", f"O{K1}")
note(ws, "D8", "A característica do produto ou o parâmetro do processo: “Temperatura da pizza na saída”, “Espessura do filme”.")
note(ws, "F8", "Deixe o mínimo e o máximo vazios para um controle por atributo: o resultado será conforme ou não.")
note(ws, "M8", "O que fazer com o produto feito desde a última verificação boa, e com o processo. “Avisar o responsável” não é reação.")
s = K2 + 2
PAV = resumo_bloco(ws, s, "O", [
    ("Controles escritos", f"=COUNTA(D{K1}:D{K2})", None),
    ("Por tipo", f'=IF(D{s+1}=0,"",COUNTIF(E{K1}:E{K2},"{TIPOS[0]}")&" de produto, "&COUNTIF(E{K1}:E{K2},"{TIPOS[1]}")&" de processo")', None),
    ("Por critério", f'=IF(D{s+1}=0,"",(D{s+1}-COUNTIF(N{K1}:N{K2},"{ATRIB}"))&" por medida, "&COUNTIF(N{K1}:N{K2},"{ATRIB}")&" por atributo")', None),
    ("Aviso", f'=IF(D4="","Escreva a organização e o processo",IF(D{s+1}=0,"Escreva os controles",IF(COUNTIF(O{K1}:O{K2},"OK")<D{s+1},"Há controle a rever: veja a coluna Conferência",'
              f'IF(COUNTIF(E{K1}:E{K2},"{TIPOS[1]}")=0,"Só há controles de produto: avalie controlar também o processo","OK"))))', None),
], "C", "D", "Aviso")
ws.freeze_panes = f"E{K1}"
setup(ws, BLUE, f"B1:O{PAV}", fit_height=True)

# ------------------------------------------------------------------ Registros
ws = wb.create_sheet(RG)
widths(ws, {"A": 2, "B": 5, "C": 12, "D": 20, "E": 10, "F": 30, "G": 18, "H": 8, "I": 8, "J": 11, "K": 11, "L": 12, "M": 40, "N": 18, "O": 26, "P": 2})
title(ws, "Registros das verificações", "Uma linha por verificação. Para um controle por medida, digite o valor. Para um controle por atributo, responda Sim ou Não.", "O")
for col, text in zip("BCDEFGHIJKLMNO", ["#", "Data", "Lote ou pedido", "Controle", "O que se controla", "Critério", "Mín.", "Máx.", "Valor medido", "Conforme?", "Leitura",
                                         "Reação tomada", "Quem mediu", "Conferência"]):
    head(ws, f"{col}4", text)
ws.row_dimensions[4].height = 33
hint_row(ws, 5, [("B", ""), ("C", "Data"), ("D", "A identificação"), ("E", "Código"), ("F", "Vem do plano"), ("G", "Vem do plano"), ("H", ""), ("I", ""), ("J", "Só o número"),
                 ("K", "Lista"), ("L", "Calculada"), ("M", "Só quando sair do critério"), ("N", "Uma função"), ("O", "Calculada")])
ex(ws, "B6", "Ex.", h="center")
for col, v, h_, fmt in [("C", date(2027, 3, 11), "center", DATE), ("D", "Câmara, 2ª leitura", "left", None), ("E", "K1", "center", None), ("F", "Temperatura da câmara fria", "left", None),
                        ("G", "0 a 5 °C", "center", None), ("H", 0, "center", None), ("I", 5, "center", None), ("J", 7, "center", None), ("K", None, "center", None), ("L", FORA, "center", None),
                        ("M", "Insumos passados para o refrigerador reserva às 21h40.", "left", None), ("N", "Pizzaiolo do turno", "left", None), ("O", "OK", "center", None)]:
    ex(ws, f"{col}6", v, h=h_, fmt=fmt)
ws.row_dimensions[6].height = 30
for k in range(NR):
    rr = R1 + k
    num(ws, f"B{rr}", k + 1)
    inp(ws, f"C{rr}", h="center", fmt=DATE)
    inp(ws, f"D{rr}")
    inp(ws, f"E{rr}", h="center")
    calc(ws, f"F{rr}", f'=IF(E{rr}="","",{busca("D", f"E{rr}", K1, K2)})', h="left", b=False, sz=9)
    calc(ws, f"G{rr}", f'=IF(E{rr}="","",{busca("N", f"E{rr}", K1, K2)})', b=False, sz=9)
    calc(ws, f"H{rr}", f'=IF(E{rr}="","",{busca("F", f"E{rr}", K1, K2, texto=False)})', b=False, sz=9, fmt=GERAL)
    calc(ws, f"I{rr}", f'=IF(E{rr}="","",{busca("G", f"E{rr}", K1, K2, texto=False)})', b=False, sz=9, fmt=GERAL)
    inp(ws, f"J{rr}", h="center", fmt=GERAL)
    inp(ws, f"K{rr}", h="center")
    calc(ws, f"L{rr}", f'=IF(OR(E{rr}="",F{rr}=""),"",{f_leitura(f"H{rr}", f"I{rr}", f"J{rr}", f"K{rr}")})', sz=9)
    inp(ws, f"M{rr}")
    inp(ws, f"N{rr}")
    calc(ws, f"O{rr}", f'=IF(E{rr}="",IF(COUNTA(C{rr}:D{rr},J{rr}:K{rr},M{rr})>0,"Falta o controle",""),IF(F{rr}="","{SEMPLANO}",IF(C{rr}="","Falta a data",IF(D{rr}="","Falta o lote ou o pedido",'
         f'IF(LEFT(L{rr},5)="Falta",L{rr},IF(AND(L{rr}="{FORA}",M{rr}=""),"{SEMREACAO}",IF(N{rr}="","Falta quem mediu","OK")))))))', b=False, sz=9)
    ws.row_dimensions[rr].height = 21.75
dv_date(ws, f"C{R1}:C{R2}")
dv_list(ws, f"E{R1}:E{R2}", IDS, "Código do controle, da aba Plano de controle")
dv_number(ws, f"J{R1}:J{R2}")
dv_list(ws, f"K{R1}:K{R2}", [SIM, NAO], "Só para controle por atributo: o resultado está conforme?")
cf_texto(ws, f"L{R1}:L{R2}", f"L{R1}", [(DENTRO, S1_T), (FORA, S3_T)], resto=YELLOW)
cf_texto(ws, f"O{R1}:O{R2}", f"O{R1}", [("OK", GREEN), (SEMREACAO, RED)], resto=YELLOW)
note(ws, "D4", "O que identifica o produto verificado: o lote, o pedido, a bobina, o turno. É por aqui que se rastreia.")
note(ws, "M4", "O que foi feito com o produto e com o processo. Um resultado fora do critério sem reação é o aviso mais grave.")
s = R2 + 2
RAV = resumo_bloco(ws, s, "G", [
    ("Verificações lançadas", f"=COUNTA(E{R1}:E{R2})", None),
    ("Dentro do critério", f'=COUNTIF(L{R1}:L{R2},"{DENTRO}")', None),
    ("Fora do critério", f'=COUNTIF(L{R1}:L{R2},"{FORA}")', None),
    ("Fora sem reação registrada", f'=COUNTIF(O{R1}:O{R2},"{SEMREACAO}")', None),
    ("Aviso", f'=IF(E{s+1}=0,"Lance as verificações",IF(E{s+4}>0,"Há resultado fora do critério sem reação registrada",'
              f'IF(COUNTIF(O{R1}:O{R2},"<>OK")-COUNTBLANK(O{R1}:O{R2})>0,"Há registro a rever: veja a coluna Conferência","OK")))', None),
], "D", "E", "Aviso")
ws.freeze_panes = f"F{R1}"
setup(ws, TEAL, f"B1:O{RAV}")

# ------------------------------------------------------------------ Rastreabilidade
ws = wb.create_sheet(RT)
widths(ws, {"A": 2, "B": 5, "C": 26, "D": 34, "E": 34, "F": 30, "G": 34, "H": 24, "I": 2})
title(ws, "Identificação, rastreabilidade e preservação", "Em cada etapa: como o produto é identificado, como se sabe se já foi verificado, o que se registra e como ele é protegido.", "H")
band(ws, 4, "Identificação, situação e preservação, por etapa", "H", color=AMBER)
for col, text in zip("BCDEFGH", ["#", "Etapa", "Como o produto é identificado", "Como se sabe a situação", "O que se registra para rastrear", "Como se preserva", "Conferência"]):
    head(ws, f"{col}5", text)
ws.row_dimensions[5].height = 33
hint_row(ws, 6, [("B", ""), ("C", "A mesma do plano"), ("D", "Etiqueta, comanda, lote"), ("E", "Aprovado, aguardando, reprovado"), ("F", "O registro que liga à etapa anterior"),
                 ("G", "Manuseio, embalagem, armazenagem"), ("H", "Calculada")])
for k in range(NE):
    rr = E1 + k
    num(ws, f"B{rr}", k + 1)
    for col in "CDEFG":
        inp(ws, f"{col}{rr}")
    calc(ws, f"H{rr}", f'=IF(C{rr}="",IF(COUNTA(D{rr}:G{rr})>0,"Falta a etapa",""),IF(D{rr}="","Falta a identificação",IF(E{rr}="","Falta a situação","OK")))', b=False, sz=9)
    ws.row_dimensions[rr].height = 30
cf_warn(ws, f"H{E1}:H{E2}", f"H{E1}")
band(ws, P1 - 3, "Propriedade de clientes e de fornecedores", "H", color=AMBER)
for col, text in zip("BCDEFGH", ["#", "Item", "De quem é", "Como é identificado", "Como é cuidado", "Ocorrência e comunicação ao dono", "Conferência"]):
    head(ws, f"{col}{P1 - 2}", text)
ws.row_dimensions[P1 - 2].height = 33
hint_row(ws, P1 - 1, [("B", ""), ("C", "Material, ferramenta, embalagem, dado"), ("D", "Lista"), ("E", "Marca, código, cadastro"), ("F", "Guarda, inspeção, acesso"),
                      ("G", "Perda ou dano, e quando o dono foi avisado"), ("H", "Calculada")])
for k in range(NP):
    rr = P1 + k
    num(ws, f"B{rr}", k + 1)
    for col in "CEFG":
        inp(ws, f"{col}{rr}")
    inp(ws, f"D{rr}", h="center")
    calc(ws, f"H{rr}", f'=IF(C{rr}="",IF(COUNTA(D{rr}:G{rr})>0,"Falta o item",""),IF(D{rr}="","Falta o dono",IF(E{rr}="","Falta a identificação",IF(F{rr}="","Falta o cuidado","OK"))))', b=False, sz=9)
    ws.row_dimensions[rr].height = 30
dv_list(ws, f"D{P1}:D{P2}", ["Cliente", "Fornecedor"], "A quem o item pertence")
cf_warn(ws, f"H{P1}:H{P2}", f"H{P1}")
note(ws, f"G{P1 - 2}", "A norma pede que a perda, o dano ou a inadequação do item seja relatada ao dono, com registro.")
s = P2 + 2
TAV = resumo_bloco(ws, s, "H", [
    ("Etapas descritas", f"=COUNTA(C{E1}:C{E2})", None),
    ("Itens de clientes e de fornecedores", f"=COUNTA(C{P1}:C{P2})", None),
    ("Ocorrências registradas", f"=COUNTA(G{P1}:G{P2})", None),
    ("Aviso", f'=IF(D{s+1}=0,"Descreva as etapas",IF(COUNTIF(H{E1}:H{E2},"<>OK")-COUNTBLANK(H{E1}:H{E2})+COUNTIF(H{P1}:H{P2},"<>OK")-COUNTBLANK(H{P1}:H{P2})>0,'
              f'"Há linha a rever: veja a coluna Conferência",IF(D{s+2}=0,"Nenhum item de terceiros: confirme se é isso mesmo","OK")))', None),
], "C", "D", "Aviso", ok_values=("OK", "Nenhum item de terceiros: confirme se é isso mesmo"))
setup(ws, AMBER, f"B1:H{TAV}", fit_height=True)

# ------------------------------------------------------------------ Mudanças
ws = wb.create_sheet("Mudanças")
widths(ws, {"A": 2, "B": 5, "C": 12, "D": 36, "E": 28, "F": 36, "G": 20, "H": 36, "I": 24, "J": 2})
title(ws, "Mudanças na produção e no serviço", "Uma linha por mudança de insumo, método, equipamento, parâmetro ou pessoas. A análise vem antes de a mudança entrar em uso.", "I")
for col, text in zip("BCDEFGHI", ["#", "Data", "O que mudou", "Motivo", "Análise antes de mudar", "Quem autorizou", "Ações decorrentes", "Conferência"]):
    head(ws, f"{col}4", text)
ws.row_dimensions[4].height = 33
hint_row(ws, 5, [("B", ""), ("C", "Data"), ("D", "De quê, para quê"), ("E", "Por que mudar"), ("F", "O que foi conferido ou testado"), ("G", "Uma função"),
                 ("H", "Documentos, treinamento, plano de controle"), ("I", "Calculada")])
ex(ws, "B6", "Ex.", h="center")
for col, v, h_, fmt in [("C", date(2027, 3, 2), "center", DATE), ("D", "Farinha da marca B no lugar da marca A.", "left", None), ("E", "Falta no fornecedor.", "left", None),
                        ("F", "Um lote de massa de teste: peso e ponto conferidos.", "left", None), ("G", "Pizzaiolo líder", "left", None),
                        ("H", "Receita ajustada. Ficha da massa atualizada.", "left", None), ("I", "OK", "center", None)]:
    ex(ws, f"{col}6", v, h=h_, fmt=fmt)
ws.row_dimensions[6].height = 30
for k in range(NM):
    rr = M1 + k
    num(ws, f"B{rr}", k + 1)
    inp(ws, f"C{rr}", h="center", fmt=DATE)
    for col in "DEFGH":
        inp(ws, f"{col}{rr}")
    calc(ws, f"I{rr}", f'=IF(D{rr}="",IF(COUNTA(C{rr},E{rr}:H{rr})>0,"Falta o que mudou",""),IF(C{rr}="","Falta a data",IF(F{rr}="","Falta a análise",IF(G{rr}="","Falta quem autorizou","OK"))))', b=False, sz=9)
    ws.row_dimensions[rr].height = 30
dv_date(ws, f"C{M1}:C{M2}")
cf_warn(ws, f"I{M1}:I{M2}", f"I{M1}")
note(ws, "F4", "O que foi conferido, testado ou medido antes de a mudança entrar em uso, e o resultado.")
s = M2 + 2
MAV = resumo_bloco(ws, s, "I", [
    ("Mudanças registradas", f"=COUNTA(D{M1}:D{M2})", None),
    ("Com análise e autorização", f'=COUNTIF(I{M1}:I{M2},"OK")', None),
    ("Aviso", f'=IF(E{s+1}=0,"Nenhuma mudança registrada: confirme se o processo não mudou",IF(E{s+2}<E{s+1},"Há mudança sem análise ou sem autorização","OK"))', None),
], "D", "E", "Aviso", ok_values=("OK", "Nenhuma mudança registrada: confirme se o processo não mudou"))
ws.freeze_panes = "E7"
setup(ws, PURPLE, f"B1:I{MAV}", fit_height=True)

# ------------------------------------------------------------------ Painel
ws = wb.create_sheet("Painel")
widths(ws, {"A": 2, "B": 6, "C": 36, "D": 26, "E": 13, "F": 9, "G": 12, "H": 10, "I": 34, "J": 2})
title(ws, "Painel do controle", "As verificações e os desvios por controle. Nada a preencher: tudo vem das outras abas.", "I")
label(ws, "B4", "Organização e processo", merge="B4:C4")
calc(ws, "D4", f'=IF({PC}!D4="","",{PC}!D4)', merge="D4:I4", h="left", b=False)
ws.row_dimensions[4].height = 21.75
for col, text, m in [("B", "Indicador", "B6:C6"), ("D", "Resultado", None), ("E", "Como se lê", "E6:I6")]:
    put(ws, f"{col}6", text, f=font(10, True, c=WHITE), bg=INK, h="center", merge=m)
ws.row_dimensions[6].height = 21.75
L_, O_, E_ = (f"{RG}!{c}{R1}:{c}{R2}" for c in "LOE")
T0 = 16
T1, T2 = T0 + 2, T0 + 1 + NK
IND = [
    ("Controles no plano", f"=COUNTA({PC}!D{K1}:D{K2})", None, "Da aba Plano de controle. O usual é de cinco a dez por processo."),
    ("Verificações lançadas", f"=COUNTA({E_})", None, "Da aba Registros."),
    ("Dentro do critério", f'=IF(COUNTIF({L_},"{DENTRO}")+COUNTIF({L_},"{FORA}")=0,"",COUNTIF({L_},"{DENTRO}")/(COUNTIF({L_},"{DENTRO}")+COUNTIF({L_},"{FORA}")))', "0%",
     "Parte das verificações com leitura que atendeu ao critério."),
    ("Fora do critério", f'=COUNTIF({L_},"{FORA}")', None, "Cada uma pede uma reação, com o produto e com o processo."),
    ("Fora sem reação registrada", f'=COUNTIF({O_},"{SEMREACAO}")', None, "O aviso mais grave: o que foi feito com o produto?"),
    ("Controles sem registro", f'=COUNTIF(I{T1}:I{T2},"{L_SEM}")', None, "Estão no plano, e não há verificação lançada."),
    ("Controles com desvio repetido", f'=COUNTIF(I{T1}:I{T2},"{L_REP}")', None, "Mais de um resultado fora: a reação não basta."),
    ("Mudanças: registradas, a rever", f'=COUNTA({MU}!D{M1}:D{M2})&" registradas, "&(COUNTA({MU}!D{M1}:D{M2})-COUNTIF({MU}!I{M1}:I{M2},"OK"))&" a rever"', None,
     "A rever: sem a análise feita antes, ou sem quem autorizou."),
]
for k, (nome, formula, fmt, leit) in enumerate(IND):
    rr = 7 + k
    label(ws, f"B{rr}", nome, merge=f"B{rr}:C{rr}")
    calc(ws, f"D{rr}", formula, fmt=fmt)
    put(ws, f"E{rr}", leit, f=font(9, c=MUTED), bg=GRAY, merge=f"E{rr}:I{rr}")
    ws.row_dimensions[rr].height = 21.75
assert 7 + len(IND) - 1 == T0 - 2
band(ws, T0, "Leitura de cada controle", "I")
for col, text in zip("BCDEFGHI", ["#", "O que se controla", "Critério", "Verificações", "Fora", "Sem reação", "Dentro", "Leitura"]):
    head(ws, f"{col}{T0+1}", text)
ws.row_dimensions[T0 + 1].height = 33
for k in range(NK):
    rr = T1 + k
    prow = K1 + k
    num(ws, f"B{rr}", IDS[k])
    calc(ws, f"C{rr}", f'=IF({PC}!D{prow}="","",{PC}!D{prow})', h="left", b=False, sz=9)
    calc(ws, f"D{rr}", f'=IF(C{rr}="","",{PC}!N{prow})', b=False, sz=9)
    calc(ws, f"E{rr}", f'=IF(C{rr}="","",COUNTIF({E_},B{rr}))', b=False)
    calc(ws, f"F{rr}", f'=IF(C{rr}="","",SUMPRODUCT(({E_}=B{rr})*({L_}="{FORA}")))')
    calc(ws, f"G{rr}", f'=IF(C{rr}="","",SUMPRODUCT(({E_}=B{rr})*({O_}="{SEMREACAO}")))', b=False)
    calc(ws, f"H{rr}", f'=IF(OR(C{rr}="",N(E{rr})=0),"",1-F{rr}/E{rr})', fmt="0%", b=False)
    calc(ws, f"I{rr}", f'=IF(C{rr}="","",IF(E{rr}=0,"{L_SEM}",IF(G{rr}>0,"{L_COBRAR}",IF(F{rr}>1,"{L_REP}",IF(F{rr}=1,"{L_UM}","{L_OK}")))))', b=False, sz=9, h="left")
    ws.row_dimensions[rr].height = 24
cf_texto(ws, f"I{T1}:I{T2}", f"I{T1}", [(L_OK, GREEN), (L_UM, S1_T), (L_COBRAR, RED), (L_SEM, RED)], resto=YELLOW)
s = T2 + 2
band(ws, s, "Resumo automático", "I")
label(ws, f"B{s+1}", "Aviso", merge=f"B{s+1}:C{s+1}", h="right")
calc(ws, f"D{s+1}", f'=IF(D7=0,"Escreva os controles na aba Plano de controle",IF(D8=0,"Lance as verificações na aba Registros",IF(D11>0,"Há resultado fora do critério sem reação registrada",'
     f'IF(D12>0,"Há controle sem registro",IF(D13>0,"Há desvio que se repete: abra a análise de causa",'
     f'IF(COUNTA({MU}!D{M1}:D{M2})>COUNTIF({MU}!I{M1}:I{M2},"OK"),"Há mudança sem análise ou sem autorização","OK"))))))', sz=9, b=False, h="left", merge=f"D{s+1}:I{s+1}")
cf_warn(ws, f"D{s+1}:I{s+1}", f"D{s+1}")
ws.row_dimensions[s + 1].height = 21.75
NAV = s + 1
barras(ws, f"B{s + 3}", T1, T2, 2, 8, width=24, height=11)
setup(ws, REDC, f"B1:I{s + 26}", landscape=False, fit_height=True)

# ------------------------------------------------------------------ Checklist
ws = wb.create_sheet("Checklist")
widths(ws, {"A": 2, "B": 5, "C": 70, "D": 14, "E": 44, "F": 2})
title(ws, "Checklist do controle de produção e de serviço", "Escolha o status de cada item na lista suspensa (coluna em amarelo-claro).", "E")
for col, text in zip("BCDE", ["#", "Item de verificação", "Status", "Observações"]):
    head(ws, f"{col}4", text)
ws.row_dimensions[4].height = 21.75
for k, text in enumerate(CHECK):
    rr = 5 + k
    num(ws, f"B{rr}", k + 1)
    put(ws, f"C{rr}", text, bg=GRAY)
    inp(ws, f"D{rr}", h="center")
    inp(ws, f"E{rr}")
    ws.row_dimensions[rr].height = 30
dv_list(ws, "D5:D16", YESNO, "Sim, Parcial ou Não")
cf_equal(ws, "D5:D16", YESNO_CF)
band(ws, 18, "Resumo automático", "E")
summary(ws, 19, "Itens atendidos (Sim)", '=COUNTIF(D5:D16,"Sim")', "C", "D")
summary(ws, 20, "Itens parciais", '=COUNTIF(D5:D16,"Parcial")', "C", "D")
summary(ws, 21, "Itens não atendidos", '=COUNTIF(D5:D16,"Não")', "C", "D")
summary(ws, 22, "Itens sem resposta", "=COUNTBLANK(D5:D16)", "C", "D")
summary(ws, 23, "Percentual atendido", "=D19/ROWS(D5:D16)", "C", "D", fmt="0%")
summary(ws, 24, "Situação", '=IF(D19=ROWS(D5:D16),"Validado",IF(D22>0,"Em andamento","Com pendências"))', "C", "D")
setup(ws, PURPLE, "B1:E24", landscape=False, fit_height=True)


# ------------------------------------------------------------------ Exemplos
def sub(ws, rr, cells, height=30):
    for a, b, t in cells:
        put(ws, f"{a}{rr}", t, f=font(10, True), bg=GRAY, h="center", merge=b and f"{a}{rr}:{b}{rr}")
    ws.row_dimensions[rr].height = height


def exemplo(ws, ex_):
    H = ex_["head"]
    widths(ws, {"A": 2, "B": 11, "C": 26, "D": 30, "E": 11, "F": 9, "G": 9, "H": 11, "I": 26, "J": 18, "K": 18, "L": 30, "M": 2})
    title(ws, "Controle de produção e de serviço", "Exemplo preenchido, para consulta. Use as abas Plano de controle, Registros, Rastreabilidade e Mudanças para o seu processo.", "L")
    rr = 4
    for rot, val in [("Organização", H["org"]), ("Processo", H["processo"]), ("Produto ou serviço", H["produto"]), ("Plano de controle", f'{H["rev"]}. {H["por"]}.'),
                     ("Registros", f'{H["periodo"]}, lidos em {H["ref"]:%d/%m/%Y}.'), ("Origem", H["origem"])]:
        label(ws, f"B{rr}", rot, merge=f"B{rr}:C{rr}")
        inp(ws, f"D{rr}", val, merge=f"D{rr}:L{rr}", bg=WHITE, h="left")
        ws.row_dimensions[rr].height = alt([(val, 130)], minimo=19.5)
        rr += 1
    rr += 1
    band(ws, rr, "Plano de controle", "L", color=BLUE)
    rr += 1
    sub(ws, rr, [("B", None, "#"), ("C", None, "Etapa"), ("D", None, "O que controlar"), ("E", None, "Tipo"), ("F", None, "Mín."), ("G", None, "Máx."), ("H", None, "Unidade"),
                 ("I", None, "Como medir"), ("J", None, "Frequência"), ("K", None, "Quem · registro"), ("L", None, "Reação se sair do critério")])
    p1 = rr + 1
    for k in ex_["plano"]:
        rr += 1
        num(ws, f"B{rr}", k["id"])
        put(ws, f"C{rr}", k["etapa"])
        put(ws, f"D{rr}", k["caract"], f=font(10, True))
        put(ws, f"E{rr}", k["tipo"], h="center")
        put(ws, f"F{rr}", k["min"] if k["min"] is not None else None, h="center", fmt=GERAL)
        put(ws, f"G{rr}", k["max"] if k["max"] is not None else None, h="center", fmt=GERAL)
        put(ws, f"H{rr}", k["unid"] or ("atributo" if atributo(k) else ""), h="center")
        put(ws, f"I{rr}", k["metodo"])
        put(ws, f"J{rr}", k["freq"])
        put(ws, f"K{rr}", f'{k["resp"]} · {k["registro"]}')
        put(ws, f"L{rr}", k["reacao"])
        ws.row_dimensions[rr].height = alt([(k["caract"], 30), (k["metodo"], 26), (k["reacao"], 30), (k["resp"] + k["registro"], 16), (k["etapa"], 24), (k["freq"], 16)], minimo=21.75)
    p2 = rr
    rr += 2
    ws.row_breaks.append(Break(id=rr - 1))
    band(ws, rr, "Resumo por controle", "L", color=AMBER)
    rr += 1
    sub(ws, rr, [("B", None, "#"), ("C", "D", "O que se controla"), ("E", None, "Verificações"), ("F", None, "Fora"), ("G", "H", "Sem reação"), ("I", None, "Dentro do critério"), ("J", "L", "Leitura")])
    q1 = rr + 1
    q2 = rr + len(ex_["plano"])
    g1 = q2 + 4                          # primeira linha dos registros, abaixo deste resumo
    g2 = g1 + len(ex_["regs"]) - 1
    E_, I_, J_ = (f"${c}${g1}:${c}${g2}" for c in "EIJ")
    for k in ex_["plano"]:
        rr += 1
        num(ws, f"B{rr}", k["id"])
        put(ws, f"C{rr}", k["caract"], merge=f"C{rr}:D{rr}")
        calc(ws, f"E{rr}", f"=COUNTIF({E_},B{rr})", b=False)
        calc(ws, f"F{rr}", f'=SUMPRODUCT(({E_}=B{rr})*({I_}="{FORA}"))')
        calc(ws, f"G{rr}", f'=SUMPRODUCT(({E_}=B{rr})*({I_}="{FORA}")*({J_}=""))', b=False, merge=f"G{rr}:H{rr}")
        calc(ws, f"I{rr}", f'=IF(E{rr}=0,"",1-F{rr}/E{rr})', fmt="0%", b=False)
        calc(ws, f"J{rr}", f'=IF(E{rr}=0,"{L_SEM}",IF(G{rr}>0,"{L_COBRAR}",IF(F{rr}>1,"{L_REP}",IF(F{rr}=1,"{L_UM}","{L_OK}"))))', b=False, sz=9, h="left", merge=f"J{rr}:L{rr}")
        ws.row_dimensions[rr].height = 21.75
    cf_texto(ws, f"J{q1}:L{q2}", f"$J{q1}", [(L_OK, GREEN), (L_UM, S1_T), (L_COBRAR, RED), (L_SEM, RED)], resto=YELLOW)
    rr += 2
    band(ws, rr, "Registros", "L", color=TEAL)
    rr += 1
    sub(ws, rr, [("B", None, "Data"), ("C", None, "Lote ou pedido"), ("D", None, "O que se controla"), ("E", None, "Controle"), ("F", None, "Mín."), ("G", None, "Máx."), ("H", None, "Resultado"),
                 ("I", None, "Leitura"), ("J", "L", "Reação tomada")])
    assert rr + 1 == g1
    ks = {k["id"]: k for k in ex_["plano"]}
    for reg in ex_["regs"]:
        rr += 1
        k = ks[reg["k"]]
        put(ws, f"B{rr}", reg["data"], h="center", fmt=DATE)
        put(ws, f"C{rr}", reg["lote"])
        put(ws, f"D{rr}", k["caract"], f=font(9))
        put(ws, f"E{rr}", reg["k"], h="center")
        calc(ws, f"F{rr}", "=" + busca("F", f"E{rr}", p1, p2, aba=None, texto=False), b=False, sz=9, fmt=GERAL)
        calc(ws, f"G{rr}", "=" + busca("G", f"E{rr}", p1, p2, aba=None, texto=False), b=False, sz=9, fmt=GERAL)
        put(ws, f"H{rr}", reg["conf"] if atributo(k) else reg["valor"], h="center", fmt=GERAL)
        calc(ws, f"I{rr}", "=" + f_leitura(f"F{rr}", f"G{rr}", f"H{rr}", f"H{rr}"), sz=9)
        put(ws, f"J{rr}", reg["reacao"] or None, merge=f"J{rr}:L{rr}")
        ws.row_dimensions[rr].height = alt([(reg["reacao"], 70)], minimo=19.5)
    assert rr == g2
    cf_texto(ws, f"I{g1}:I{g2}", f"I{g1}", [(DENTRO, S1_T), (FORA, S3_T)])
    rr += 2
    ws.row_breaks.append(Break(id=rr - 1))
    band(ws, rr, "Identificação, situação e preservação", "L", color=AMBER)
    rr += 1
    sub(ws, rr, [("B", None, "#"), ("C", None, "Etapa"), ("D", None, "Como o produto é identificado"), ("E", "H", "Como se sabe a situação"), ("I", "J", "O que se registra para rastrear"),
                 ("K", "L", "Como se preserva")])
    for j, (e, ident, sit, reg_, pres) in enumerate(ex_["rast"], 1):
        rr += 1
        num(ws, f"B{rr}", j)
        put(ws, f"C{rr}", e, f=font(10, True))
        put(ws, f"D{rr}", ident)
        put(ws, f"E{rr}", sit, merge=f"E{rr}:H{rr}")
        put(ws, f"I{rr}", reg_, merge=f"I{rr}:J{rr}")
        put(ws, f"K{rr}", pres or "—", merge=f"K{rr}:L{rr}")
        ws.row_dimensions[rr].height = alt([(ident, 30), (sit, 40), (reg_, 44), (pres, 48), (e, 24)], minimo=21.75)
    rr += 2
    band(ws, rr, "Propriedade de clientes e de fornecedores", "L", color=AMBER)
    rr += 1
    sub(ws, rr, [("B", None, "#"), ("C", None, "Item"), ("D", None, "Como é identificado"), ("E", "F", "De quem é"), ("G", "I", "Como é cuidado"), ("J", "L", "Ocorrência e comunicação ao dono")])
    t1 = rr + 1
    for j, (item, dono, ident, cuid, oc) in enumerate(ex_["prop"], 1):
        rr += 1
        num(ws, f"B{rr}", j)
        put(ws, f"C{rr}", item, f=font(10, True))
        put(ws, f"D{rr}", ident)
        put(ws, f"E{rr}", dono, h="center", merge=f"E{rr}:F{rr}")
        put(ws, f"G{rr}", cuid, merge=f"G{rr}:I{rr}")
        put(ws, f"J{rr}", oc or None, merge=f"J{rr}:L{rr}")
        ws.row_dimensions[rr].height = alt([(item, 24), (ident, 30), (cuid, 46), (oc, 66)], minimo=21.75)
    t2 = rr
    rr += 2
    band(ws, rr, "Mudanças" + (f' · registro lido em {H["mud_ref"]:%d/%m/%Y}' if H.get("mud_ref") else ""), "L", color=PURPLE)
    rr += 1
    sub(ws, rr, [("B", None, "Data"), ("C", None, "O que mudou"), ("D", None, "Motivo"), ("E", "H", "Análise antes de mudar"), ("I", None, "Quem autorizou"), ("J", "K", "Ações decorrentes"),
                 ("L", None, "Conferência")])
    m1 = rr + 1
    for d, oque, motivo, analise, aut, acoes in ex_["mud"]:
        rr += 1
        put(ws, f"B{rr}", d, h="center", fmt=DATE)
        put(ws, f"C{rr}", oque, f=font(10, True))
        put(ws, f"D{rr}", motivo)
        put(ws, f"E{rr}", analise or None, merge=f"E{rr}:H{rr}")
        put(ws, f"I{rr}", aut or None)
        put(ws, f"J{rr}", acoes or None, merge=f"J{rr}:K{rr}")
        calc(ws, f"L{rr}", f'=IF(E{rr}="","Falta a análise",IF(I{rr}="","Falta quem autorizou","OK"))', b=False, sz=9)
        ws.row_dimensions[rr].height = alt([(oque, 24), (motivo, 30), (analise, 40), (acoes, 36)], minimo=21.75)
    m2 = rr
    cf_warn(ws, f"L{m1}:L{m2}", f"L{m1}")
    rr += 2
    band(ws, rr, "Resumo automático", "L")
    for k, (text, formula, fmt) in enumerate([
        ("Controles no plano", f"=COUNTA(D{p1}:D{p2})", None),
        ("Verificações: lançadas, fora do critério", f'=COUNTA(E{g1}:E{g2})&" lançadas, "&COUNTIF(I{g1}:I{g2},"{FORA}")&" fora do critério"', None),
        ("Dentro do critério", f'=COUNTIF(I{g1}:I{g2},"{DENTRO}")/COUNTA(E{g1}:E{g2})', "0%"),
        ("Fora sem reação registrada", f"=SUM(G{q1}:G{q2})", None),
        ("Itens de terceiros com ocorrência", f"=COUNTA(J{t1}:J{t2})", None),
        ("Mudanças: registradas, a rever", f'=COUNTA(C{m1}:C{m2})&" registradas, "&(COUNTA(C{m1}:C{m2})-COUNTIF(L{m1}:L{m2},"OK"))&" a rever"', None),
    ], 1):
        label(ws, f"B{rr+k}", text, merge=f"B{rr+k}:D{rr+k}", h="right")
        calc(ws, f"E{rr+k}", formula, fmt=fmt, merge=f"E{rr+k}:L{rr+k}", h="left")
        ws.row_dimensions[rr + k].height = 21.75
    setup(ws, MUTED, f"B1:L{rr + 6}")


exemplo(wb.create_sheet("Exemplo 1 - Pizzaria"), EX1)
exemplo(wb.create_sheet("Exemplo 2 - Indústria"), EX2)

wb.active = 1
wb.save(OUT)
print("ok", OUT, "| avisos: Plano D%d, Registros E%d, Rastreabilidade D%d, Mudanças E%d, Painel D%d" % (PAV, RAV, TAV, MAV, NAV))
