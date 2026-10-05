# -*- coding: utf-8 -*-
"""Gera SWOT-modelo.xlsx no mesmo padrão visual de SIPOC-modelo.xlsx e PDCA-modelo.xlsx."""
import os
import sys
from datetime import date

# funções de formatação compartilhadas com o gerador do PDCA (bloco inicial do arquivo)
_here = os.path.dirname(os.path.abspath(__file__))
_src = open(os.path.join(_here, "build_pdca.py"), encoding="utf-8").read()
exec(_src[: _src.index("STATUS = [")])

from openpyxl.chart import BarChart, Reference  # noqa: E402
from openpyxl.chart.label import DataLabelList  # noqa: E402
from openpyxl.chart.series import DataPoint  # noqa: E402

QD = {
    "S": ("1E7B73", "D9EEEB", "Forças", "Strengths", "interno, ajuda", "O que a organização faz bem e está sob o seu controle?"),
    "W": ("A96A12", "F6E8CF", "Fraquezas", "Weaknesses", "interno, atrapalha", "O que a organização faz mal ou não tem, e poderia mudar?"),
    "O": ("2B5C8A", "DCE8F3", "Oportunidades", "Opportunities", "externo, ajuda", "O que acontece no ambiente e pode favorecer o objetivo?"),
    "T": ("B0413E", "F5DEDC", "Ameaças", "Threats", "externo, atrapalha", "O que acontece no ambiente e pode prejudicar o objetivo?"),
}
YESNO = ["Sim", "Parcial", "Não"]
YESNO_CF = [("Sim", GREEN), ("Parcial", YELLOW), ("Não", RED)]
STATUS_CF = [("Concluída", GREEN), ("Em andamento", YELLOW), ("Não iniciada", RED)]
PRIO_CF = [("Alta", RED), ("Média", YELLOW), ("Baixa", GREEN)]
N = 8  # linhas por quadrante


def note(ws, ref, text):
    c = Comment(text, "Modelo SWOT")
    c.width, c.height = 280, 110
    ws[ref].comment = c


def dv_score(ws, rng):
    dv = DataValidation(type="whole", operator="between", formula1="1", formula2="5", allow_blank=True)
    dv.promptTitle, dv.prompt = "Nota de 1 a 5", "1 = muito baixo; 5 = muito alto"
    dv.errorTitle, dv.error = "Valor inválido", "Digite um número inteiro de 1 a 5."
    dv.showErrorMessage = dv.showInputMessage = True
    ws.add_data_validation(dv)
    dv.add(rng)


SWOT_W = {"A": 2, "B": 5, "C": 20, "D": 26, "E": 10, "F": 10, "G": 9, "H": 5, "I": 20, "J": 26, "K": 10, "L": 10, "M": 9, "N": 2}
SIDE = {"L": ("B", "C", "D", "E", "F", "G"), "R": ("H", "I", "J", "K", "L", "M")}
POS = {"S": ("L", 11), "W": ("R", 11), "O": ("L", 25), "T": ("R", 25)}  # lado e primeira linha do bloco


def cells(q):
    """Colunas e linhas de entrada de um quadrante: (col fator, col impacto, col intensidade, col pontos, primeira, última)."""
    side, r0 = POS[q]
    c = SIDE[side]
    return c[1], c[3], c[4], c[5], r0 + 5, r0 + 4 + N


def quadrant(ws, q, data=None, is_ex=False):
    side, r0 = POS[q]
    cn, cf, cf2, ci, cx, cp = SIDE[side]
    dark, tint, name, eng, kind, hint = QD[q]
    put(ws, f"{cn}{r0}", q, f=font(20, True, c=WHITE), bg=dark, h="center", box=False, merge=f"{cn}{r0}:{cp}{r0}")
    put(ws, f"{cn}{r0+1}", f"{name} ({eng}) · {kind}", f=font(10, True, c=WHITE), bg=dark, h="center", merge=f"{cn}{r0+1}:{cp}{r0+1}")
    put(ws, f"{cn}{r0+2}", hint, f=font(9, i=True, c=MUTED), bg=tint, h="center", merge=f"{cn}{r0+2}:{cp}{r0+2}")
    for col, text in ((cn, "#"), (ci, "Impacto"), (cx, "Intensidade"), (cp, "Pontos")):
        put(ws, f"{col}{r0+3}", text, f=font(9 if col != cn else 10, True), bg=tint, h="center")
    put(ws, f"{cf}{r0+3}", "Fator", f=font(10, True), bg=tint, h="center", merge=f"{cf}{r0+3}:{cf2}{r0+3}")
    ws.row_dimensions[r0].height = 33.75
    ws.row_dimensions[r0 + 1].height = 21.75
    ws.row_dimensions[r0 + 2].height = 24
    ws.row_dimensions[r0 + 3].height = 21.75
    ws.row_dimensions[r0 + 4].height = 30
    if not is_ex:
        e = {"S": ("Receita própria, elogiada nas avaliações dos clientes", 5, 4),
             "W": ("Forno único, que limita a produção no horário de pico", 5, 4),
             "O": ("Novos condomínios em construção no bairro", 5, 4),
             "T": ("Aumento da taxa cobrada pelo aplicativo de delivery", 4, 4)}[q]
        ex(ws, f"{cn}{r0+4}", "Ex.", h="center")
        put(ws, f"{cf}{r0+4}", e[0], f=font(9, i=True, c=MUTED), bg=GRAY, merge=f"{cf}{r0+4}:{cf2}{r0+4}")
        ex(ws, f"{ci}{r0+4}", e[1], h="center")
        ex(ws, f"{cx}{r0+4}", e[2], h="center")
        ex(ws, f"{cp}{r0+4}", e[1] * e[2], h="center")
    first = r0 + 5
    for k in range(N):
        r = first + k
        d = data[k] if data and k < len(data) else (None, None, None)
        bg = (tint if d[0] else WHITE) if is_ex else INPUT
        put(ws, f"{cn}{r}", f"{q}{k+1}", f=font(10, True, c=MUTED), bg=GRAY, h="center")
        inp(ws, f"{cf}{r}", d[0], merge=f"{cf}{r}:{cf2}{r}", bg=bg)
        inp(ws, f"{ci}{r}", d[1], h="center", bg=bg)
        inp(ws, f"{cx}{r}", d[2], h="center", bg=bg)
        calc(ws, f"{cp}{r}", f'=IF(OR({ci}{r}="",{cx}{r}=""),"",{ci}{r}*{cx}{r})')
        ws.row_dimensions[r].height = 33 if is_ex else 30
    dv_score(ws, f"{ci}{first}:{cx}{first+N-1}")


def swot_sheet(ws, tab, data=None):
    widths(ws, SWOT_W)
    is_ex = data is not None
    title(ws, "SWOT — Matriz de análise",
          "Exemplo preenchido, para consulta. Use a aba SWOT para a sua análise." if is_ex
          else "Preencha as células em amarelo-claro. As células cinza são calculadas automaticamente.", "M")
    hb = WHITE if is_ex else INPUT
    H = data["head"] if is_ex else {}
    for r, l1, k1, l2, k2, f2 in [(4, "Objeto da análise", "objeto", "Responsável", "resp", None),
                                  (5, "Área / unidade", "area", "Data", "data", DATE),
                                  (6, "Elaborado por", "autor", "Versão", "versao", None),
                                  (8, "Horizonte de tempo", "horizonte", "Participantes", "part", None)]:
        label(ws, f"B{r}", l1, merge=f"B{r}:C{r}")
        inp(ws, f"D{r}", H.get(k1), merge=f"D{r}:G{r}", bg=hb)
        label(ws, f"H{r}", l2, merge=f"H{r}:I{r}")
        inp(ws, f"J{r}", H.get(k2), merge=f"J{r}:M{r}", bg=hb, fmt=f2)
        ws.row_dimensions[r].height = 21.75
    for r, l1, k1 in [(7, "Objetivo da análise", "objetivo"), (9, "Fora do escopo", "fora")]:
        label(ws, f"B{r}", l1, merge=f"B{r}:C{r}")
        inp(ws, f"D{r}", H.get(k1), merge=f"D{r}:M{r}", bg=hb)
        ws.row_dimensions[r].height = 30
    ws.row_dimensions[8].height = 30
    dv_date(ws, "J5")
    if not is_ex:
        note(ws, "D4", 'O que está sendo analisado: a organização, uma área, um produto ou um projeto.\nEx.: "Operação de delivery da pizzaria".')
        note(ws, "D7", 'O resultado que se quer alcançar. É ele que define o que ajuda e o que atrapalha.\nEx.: "Aumentar em 20% o faturamento do delivery".')
        note(ws, "D8", 'Período coberto pela análise.\nEx.: "12 meses".')
        note(ws, "J8", "Quem participou da oficina. Procure reunir áreas e níveis diferentes.")
        note(ws, "D9", 'O que esta análise não vai tratar.\nEx.: "Atendimento no salão".')
        note(ws, "E14", "Impacto: quanto o fator pesa no objetivo.\n1 = quase não afeta; 5 = decide o resultado.")
        note(ws, "F14", "Intensidade. Para forças e fraquezas: o quanto a característica é marcante hoje.\nPara oportunidades e ameaças: a chance de acontecer no horizonte da análise.\n1 = muito baixa; 5 = muito alta.")
    for q in "SWOT":
        quadrant(ws, q, data["fatores"][q] if is_ex else None, is_ex)
    if is_ex:
        put(ws, "B15", "Forças e fraquezas descrevem a organização como ela é hoje. Cada fator é um fato que pode ser verificado. "
            "Os pontos são o produto de impacto e intensidade, de 1 a 25.",
            f=font(9, i=True, c=MUTED), bg=GRAY, merge="B15:M15")
        put(ws, "B29", "Oportunidades e ameaças vêm do ambiente: existiriam mesmo que a organização não existisse.",
            f=font(9, i=True, c=MUTED), bg=GRAY, merge="B29:M29")

    s = 39
    band(ws, s, "Resumo automático", "M")
    label(ws, f"B{s+1}", "Quadrante", merge=f"B{s+1}:D{s+1}", h="right")
    slots = {"S": (f"E{s+1}:G{s+1}", "E"), "W": (f"H{s+1}:I{s+1}", "H"), "O": (f"J{s+1}:J{s+1}", "J"), "T": (f"K{s+1}:M{s+1}", "K")}
    for r, text in ((s + 2, "Fatores listados"), (s + 3, "Pontos"), (s + 4, "Média por fator"), (s + 5, "Aviso")):
        label(ws, f"B{r}", text, merge=f"B{r}:D{r}", h="right")
        ws.row_dimensions[r].height = 21.75 if r != s + 5 else 30
    ws.row_dimensions[s + 1].height = 21.75
    for q in "SWOT":
        rng, col = slots[q]
        a, b = rng.split(":")
        end = b[0]
        mg = (lambda r: f"{col}{r}:{end}{r}" if end != col else None)
        cfa, _, _, cpt, r1, r2 = cells(q)
        put(ws, f"{col}{s+1}", f"{q} — {QD[q][2]}", f=font(10, True, c=WHITE), bg=QD[q][0], h="center", merge=mg(s + 1))
        calc(ws, f"{col}{s+2}", f"=COUNTA({cfa}{r1}:{cfa}{r2})", fmt='0" fatores"', merge=mg(s + 2))
        calc(ws, f"{col}{s+3}", f"=SUM({cpt}{r1}:{cpt}{r2})", fmt='0" pontos"', merge=mg(s + 3))
        calc(ws, f"{col}{s+4}", f'=IF(COUNT({cpt}{r1}:{cpt}{r2})=0,"",AVERAGE({cpt}{r1}:{cpt}{r2}))', fmt='0.0', merge=mg(s + 4))
        calc(ws, f"{col}{s+5}", f'=IF({col}{s+2}=0,"Liste os fatores",IF({col}{s+2}<3,"Poucos fatores: liste ao menos 3",'
             f'IF(COUNT({cpt}{r1}:{cpt}{r2})<{col}{s+2},"Há fator sem nota","OK")))', b=False, sz=9, merge=mg(s + 5))
    cf_ok(ws, f"E{s+5}:M{s+5}")

    def wide(r, text, formula, fmt=None, height=21.75, sz=10):
        label(ws, f"B{r}", text, merge=f"B{r}:G{r}", h="right")
        calc(ws, f"H{r}", formula, fmt=fmt, merge=f"H{r}:M{r}", sz=sz)
        ws.row_dimensions[r].height = height

    med = {q: f"{slots[q][1]}{s+4}" for q in "SWOT"}
    rngs = {q: "{3}{4}:{3}{5}".format(*cells(q)) for q in "SWOT"}
    wide(s + 6, "Balanço interno (média das forças menos a das fraquezas)",
         f'=IF(OR({med["S"]}="",{med["W"]}=""),"",{med["S"]}-{med["W"]})', fmt="+0.0;-0.0;0.0")
    wide(s + 7, "Balanço externo (média das oportunidades menos a das ameaças)",
         f'=IF(OR({med["O"]}="",{med["T"]}=""),"",{med["O"]}-{med["T"]})', fmt="+0.0;-0.0;0.0")
    wide(s + 8, "Postura estratégica sugerida",
         "=IF(OR(" + ",".join(f"COUNT({rngs[q]})=0" for q in "SWOT") + '),"Dê notas aos quatro quadrantes para ver a postura",'
         f'IF({med["S"]}>={med["W"]},IF({med["O"]}>={med["T"]},"Desenvolvimento: usar as forças para aproveitar as oportunidades",'
         '"Manutenção: usar as forças para enfrentar as ameaças"),'
         f'IF({med["O"]}>={med["T"]},"Crescimento: corrigir as fraquezas para aproveitar as oportunidades",'
         '"Sobrevivência: reduzir as fraquezas e proteger-se das ameaças")))', height=30)
    wide(s + 9, "Campos do cabeçalho preenchidos", "=COUNTA(D4,J4,D5,J5,D6,J6,D7,D8,J8,D9)", fmt='0" de 10"')
    end = s + 9

    if is_ex:
        band(ws, end + 2, "Estratégias da SWOT cruzada", "M")
        r = end + 3
        put(ws, f"B{r}", "Cruzamento", f=font(10, True), bg=GRAY, h="center", merge=f"B{r}:C{r}")
        put(ws, f"D{r}", "Fatores", f=font(10, True), bg=GRAY, h="center")
        put(ws, f"E{r}", "Estratégia", f=font(10, True), bg=GRAY, h="center", merge=f"E{r}:M{r}")
        ws.row_dimensions[r].height = 21.75
        for kind, pair, text in data["estrategias"]:
            r += 1
            put(ws, f"B{r}", kind, f=font(10, True), merge=f"B{r}:C{r}")
            put(ws, f"D{r}", pair, h="center")
            put(ws, f"E{r}", text, merge=f"E{r}:M{r}")
            ws.row_dimensions[r].height = 30
        end = r
    setup(ws, tab, f"B1:M{end}", landscape=False, fit_height=True)


# ------------------------------------------------------------------ dados dos exemplos
EX1 = {
    "head": dict(objeto="Operação de delivery de uma pizzaria de bairro", resp="Gerente da loja", area="Pizzaria (loja)",
                 data=date(2026, 9, 29), autor="Equipe da loja", versao="1.0",
                 objetivo="Aumentar em 20% o faturamento do delivery.", horizonte="12 meses",
                 part="Gerente, líder da expedição, pizzaiolo, atendente e um entregador",
                 fora="Atendimento no salão e abertura de novas unidades."),
    "fatores": {
        "S": [("Receita própria e massa de fermentação longa, elogiadas nas avaliações", 5, 4),
              ("Entrega própria, com 95% dos pedidos no prazo", 4, 4),
              ("Clientela fiel no bairro, com alta recompra", 4, 3),
              ("Equipe estável, com baixa rotatividade", 3, 3)],
        "W": [("Forno único, que limita a produção no horário de pico", 5, 4),
              ("Dependência de um aplicativo para 70% dos pedidos", 4, 5),
              ("Sem controle de custo por pizza", 3, 3),
              ("Pouca presença nas redes sociais", 2, 3)],
        "O": [("Novos condomínios em construção no bairro", 5, 4),
              ("Procura crescente por opções vegetarianas e sem glúten", 3, 3),
              ("Empresas da região buscando fornecedores para eventos", 2, 3),
              ("Pedido por aplicativo de mensagens, com custo menor por venda", 4, 4)],
        "T": [("Aumento da taxa cobrada pelo aplicativo de delivery", 4, 4),
              ("Chegada de uma rede de pizzarias ao bairro", 4, 3),
              ("Alta do preço do queijo e da farinha", 3, 4),
              ("Obras na avenida principal, que atrasam as entregas", 2, 2)],
    },
    "estrategias": [
        ("SO — Desenvolvimento", "S2 × O4", "Lançar o pedido por mensagem, com a entrega própria no prazo como argumento."),
        ("SO — Desenvolvimento", "S1 × O1", "Fazer uma ação de boas-vindas nos novos condomínios, com degustação."),
        ("WO — Crescimento", "W2 × O4", "Levar 30% dos pedidos do aplicativo para o canal próprio."),
        ("WO — Crescimento", "W1 × O1", "Avaliar a compra do segundo forno antes da entrega dos condomínios."),
        ("ST — Manutenção", "S2 × T1", "Divulgar a entrega própria no prazo para levar os clientes do aplicativo ao canal próprio."),
        ("WT — Sobrevivência", "W2 × T1", "Negociar a taxa com o aplicativo e definir preços por canal, para que o aumento não consuma a margem."),
    ],
}
EX2 = {
    "head": dict(objeto="Distribuidora de materiais elétricos · Compras", resp="Gerente de Suprimentos", area="Suprimentos",
                 data=date(2026, 5, 22), autor="Equipe de compras", versao="1.0",
                 objetivo="Reduzir em 8% o custo total de aquisição, sem aumentar as faltas de material.", horizonte="12 meses",
                 part="Gerente, compradores, Controladoria, Produção e Qualidade",
                 fora="Gestão de estoques e pagamento a fornecedores."),
    "fatores": {
        "S": [("Compradores experientes, com conhecimento técnico dos itens", 4, 4),
              ("Base de fornecedores homologados para todos os itens críticos", 4, 4),
              ("Contratos de longo prazo para os principais insumos", 5, 3)],
        "W": [("Requisições chegam com especificação incompleta", 4, 4),
              ("Cotações e aprovações por e-mail, fora do sistema de compras", 4, 4),
              ("Fornecedor único contratado para três itens críticos", 5, 4),
              ("Avaliação de fornecedores sem indicadores objetivos", 3, 3)],
        "O": [("Novos fabricantes nacionais de itens hoje importados", 4, 3),
              ("Plataformas de cotação eletrônica com custo acessível", 4, 4),
              ("Outras unidades do grupo dispostas a comprar em conjunto", 3, 3)],
        "T": [("Variação do câmbio nos insumos importados", 5, 4),
              ("Aumento do frete e dos prazos logísticos", 4, 3),
              ("Exigências ambientais e sociais sobre a cadeia de fornecedores", 3, 3),
              ("Oferta de um insumo crítico concentrada em poucos fabricantes", 4, 4)],
    },
    "estrategias": [
        ("WT — Sobrevivência", "W3 × T4", "Homologar um segundo fornecedor para os três itens críticos."),
        ("WT — Sobrevivência", "W1 × T2", "Padronizar a requisição, com especificação e data de necessidade, para reduzir as compras urgentes."),
        ("ST — Manutenção", "S3 × T1", "Incluir cláusulas de reajuste e de proteção cambial nos contratos de longo prazo."),
        ("WO — Crescimento", "W2 × O2", "Implantar uma plataforma de cotação eletrônica no lugar das cotações por e-mail."),
        ("SO — Desenvolvimento", "S1 × O1", "Desenvolver fabricantes nacionais, com o apoio técnico dos compradores."),
    ],
}

# ================================================================== pasta de trabalho
wb = Workbook()
wb.properties.title = "SWOT — Modelo de matriz de análise"
wb.properties.creator = "Modelo SWOT"
wb.properties.language = "pt-BR"

# ------------------------------------------------------------------ Instruções
ws = wb.active
ws.title = "Instruções"
widths(ws, {"A": 2, "B": 24, "C": 92, "D": 2})
title(ws, "SWOT — Como usar esta planilha",
      "Modelo para analisar forças, fraquezas, oportunidades e ameaças e transformar o diagnóstico em estratégia.", "C")
r = 4


def section(text):
    global r
    band(ws, r, text, "C", sz=11)
    r += 1


def line(k, v, kbg=GRAY, vbg=None, kf=None, kh="left", height=19.5):
    global r
    put(ws, f"B{r}", k, f=kf or font(10, True), bg=kbg, h=kh)
    put(ws, f"C{r}", v, bg=vbg)
    ws.row_dimensions[r].height = height
    r += 1


section("Legenda: onde preencher")
line("Amarelo-claro", "Células de entrada. É aqui que você digita.", vbg=INPUT)
line("Cinza", "Células calculadas ou fixas (rótulos, contagens, avisos). Não altere.", vbg=GRAY)
line("Cinza em itálico", "Linha de exemplo (marcada com “Ex.”). Mostra o formato esperado e não entra nas contagens.", vbg=GRAY)
line("Listas suspensas", "Colunas de código, prioridade e status aceitam apenas as opções da lista. As notas aceitam números de 1 a 5.")
line("Comentários", "Células com triângulo vermelho no canto trazem uma dica de preenchimento. Passe o mouse sobre elas.", height=31.5)
r += 1
section("Os quatro quadrantes")
for q, text in [("S", "Forças: características internas que favorecem o objetivo. A organização controla."),
                ("W", "Fraquezas: características internas que dificultam o objetivo. A organização controla."),
                ("O", "Oportunidades: fatos ou tendências do ambiente que podem favorecer o objetivo. A organização não controla."),
                ("T", "Ameaças: fatos ou tendências do ambiente que podem prejudicar o objetivo. A organização não controla.")]:
    line(q, text, kbg=QD[q][0], vbg=QD[q][1], kf=font(10, True, c=WHITE), kh="center")
r += 1
section("Escala das notas")
for k, text in [("Impacto", "Quanto o fator pesa no objetivo. 1 = quase não afeta; 3 = afeta de forma moderada; 5 = decide o resultado."),
                ("Intensidade (S e W)", "O quanto a característica é marcante hoje. 1 = pouco perceptível; 3 = aparece com frequência; 5 = define a organização."),
                ("Intensidade (O e T)", "A chance de acontecer no horizonte da análise. 1 = improvável; 3 = pode acontecer ou não; 5 = já está acontecendo."),
                ("Pontos", "Impacto × intensidade, de 1 a 25. Prioridade alta: 15 ou mais. Média: de 8 a 14. Baixa: até 7.")]:
    line(k, text, height=31.5)
r += 1
section("Ordem recomendada de preenchimento")
for k, text in enumerate([
    "Aba SWOT: preencha o cabeçalho, com o objeto, o objetivo e o horizonte da análise.",
    "Aba SWOT, quadrantes O e T: liste as oportunidades e as ameaças. Comece pelo ambiente externo.",
    "Aba SWOT, quadrantes S e W: liste as forças e as fraquezas, com fatos que as sustentem.",
    "Aba SWOT: dê as notas de impacto e de intensidade a cada fator.",
    "Aba Priorização: veja os fatores em ordem de pontuação. Nada para digitar.",
    "Aba SWOT cruzada: combine os fatores de alta prioridade pelo código e escreva as estratégias.",
    "Aba Plano de ação: leve as estratégias escolhidas a ações, com responsável, prazo e indicador.",
    "Aba Checklist: valide a análise com o grupo e marque a data da próxima revisão.",
], 1):
    line(f"Passo {k}", text)
r += 1
section("Abas da planilha")
for k, text in [
    ("SWOT", "Modelo principal. Cabeçalho da análise, quatro quadrantes com notas e resumo automático, com a média por quadrante."),
    ("Priorização", "Fatores em ordem de pontuação, faixas de prioridade e gráfico dos pontos por quadrante."),
    ("SWOT cruzada", "Estratégias que combinam um fator interno com um fator externo."),
    ("Plano de ação", "Ações de cada estratégia, com situação calculada a partir do prazo e do status."),
    ("Checklist", "Doze verificações de qualidade da análise, com percentual de conclusão."),
    ("Exemplo 1 - Pizzaria", "Análise preenchida com postura de desenvolvimento: crescer no delivery."),
    ("Exemplo 2 - Compras", "Análise preenchida com postura de manutenção: reduzir o custo de aquisição."),
]:
    line(k, text)
r += 1
section("Regras e premissas do modelo")
for k, text in [
    ("Códigos", "Cada fator tem um código fixo, formado pela letra do quadrante e pelo número da linha: S1 a S8, W1 a W8, O1 a O8 e T1 a T8. A aba SWOT cruzada usa esses códigos para buscar o texto dos fatores."),
    ("Postura estratégica", "A postura compara a média dos pontos por fator de cada quadrante: forças com fraquezas e oportunidades com ameaças. A média, e não a soma, para que um quadrante não pese mais só por ter mais fatores. É uma leitura de apoio, que depende das notas dadas pelo grupo."),
    ("Empates", "Na aba Priorização, fatores com a mesma pontuação aparecem na ordem dos quadrantes: S, W, O e T."),
    ("Ação atrasada", "Uma ação é considerada atrasada quando o prazo é anterior à data de hoje e o status não é Concluída nem Cancelada."),
    ("Capacidade", "A planilha comporta 8 fatores por quadrante, 12 estratégias e 15 ações. Se faltar espaço, agrupe fatores parecidos: listas longas escondem o que importa."),
    ("Exemplos", "Os dados das abas de exemplo são ilustrativos, criados para este material."),
]:
    line(k, text, height=45.75 if len(text) > 100 else 31.5)
setup(ws, INK, f"B1:C{r}", landscape=False, fit_height=True)

# ------------------------------------------------------------------ SWOT
swot_sheet(wb.create_sheet("SWOT"), QD["S"][0])

# ------------------------------------------------------------------ Priorização
ws = wb.create_sheet("Priorização")
widths(ws, {"A": 2, "B": 6, "C": 9, "D": 20, "E": 52, "F": 11, "G": 12, "H": 10, "I": 14, "J": 2, "K": 2,
            "L": 9, "M": 18, "N": 40, "O": 10, "P": 12, "Q": 10, "R": 10, "S": 2})
title(ws, "Priorização dos fatores", "Esta aba é calculada a partir da aba SWOT. Não há nada para digitar aqui.", "I")
band(ws, 4, "Fatores em ordem de pontuação", "I")
for col, text in zip("BCDEFGHI", ["Pos.", "Código", "Quadrante", "Fator", "Impacto", "Intensidade", "Pontos", "Prioridade"]):
    put(ws, f"{col}5", text, f=font(10, True), bg=GRAY, h="center")
put(ws, "L4", "Apoio ao cálculo (não altere)", f=font(10, True, c=WHITE), bg=MUTED, box=False, merge="L4:R4")
for col, text in zip("LMNOPQR", ["Código", "Quadrante", "Fator", "Impacto", "Intensidade", "Pontos", "Posição"]):
    put(ws, f"{col}5", text, f=font(9, True), bg=GRAY, h="center")
ws.row_dimensions[5].height = 21.75
F1, F2 = 6, 6 + 4 * N - 1
k = 0
for q in "SWOT":
    cfa, cim, cin, cpt, r1, r2 = cells(q)
    for j in range(N):
        row = F1 + k
        src = r1 + j
        put(ws, f"L{row}", f"{q}{j+1}", f=font(9, c=MUTED), bg=GRAY, h="center")
        put(ws, f"M{row}", QD[q][2], f=font(9, c=MUTED), bg=GRAY)
        for col, c in (("N", cfa), ("O", cim), ("P", cin), ("Q", cpt)):
            put(ws, f"{col}{row}", f'=IF(SWOT!{c}{src}="","",SWOT!{c}{src})', f=font(9, c=MUTED), bg=GRAY, h="left" if col == "N" else "center")
        put(ws, f"R{row}", f'=IF(Q{row}="","",RANK(Q{row},$Q${F1}:$Q${F2})+COUNTIF($Q${F1 - 1}:Q{row - 1},Q{row}))',
            f=font(9, c=MUTED), bg=GRAY, h="center")
        k += 1
for k in range(4 * N):
    row = F1 + k
    num(ws, f"B{row}", k + 1)
    for col, src, h in (("C", "L", "center"), ("D", "M", "left"), ("E", "N", "left"), ("F", "O", "center"), ("G", "P", "center"), ("H", "Q", "center")):
        calc(ws, f"{col}{row}", f'=IFERROR(INDEX(${src}${F1}:${src}${F2},MATCH($B{row},$R${F1}:$R${F2},0)),"")',
             h=h, b=(col in "CH"))
    calc(ws, f"I{row}", f'=IF(H{row}="","",IF(H{row}>=15,"Alta",IF(H{row}>=8,"Média","Baixa")))', b=False)
    ws.row_dimensions[row].height = 21.75
cf_equal(ws, f"I{F1}:I{F2}", PRIO_CF)
for q in "SWOT":
    ws.conditional_formatting.add(f"C{F1}:D{F2}", FormulaRule(formula=[f'LEFT($C{F1},1)="{q}"'],
                                  fill=PatternFill("solid", bgColor=QD[q][1], fgColor=QD[q][1])))
s = F2 + 2
band(ws, s, "Resumo automático", "I")
summary(ws, s + 1, "Fatores com pontuação", f"=COUNT(H{F1}:H{F2})", "E", "F", val_merge=f"F{s+1}:G{s+1}")
summary(ws, s + 2, "Prioridade alta (15 pontos ou mais)", f'=COUNTIF(I{F1}:I{F2},"Alta")', "E", "F", val_merge=f"F{s+2}:G{s+2}")
summary(ws, s + 3, "Prioridade média (de 8 a 14 pontos)", f'=COUNTIF(I{F1}:I{F2},"Média")', "E", "F", val_merge=f"F{s+3}:G{s+3}")
summary(ws, s + 4, "Prioridade baixa (até 7 pontos)", f'=COUNTIF(I{F1}:I{F2},"Baixa")', "E", "F", val_merge=f"F{s+4}:G{s+4}")
t = s + 6
band(ws, t, "Pontos por quadrante", "I")
put(ws, f"B{t+1}", "Quadrante", f=font(10, True), bg=GRAY, h="center", merge=f"B{t+1}:D{t+1}")
for col, text in (("E", "Leitura"), ("F", "Fatores"), ("G", "Pontos")):
    put(ws, f"{col}{t+1}", text, f=font(10, True), bg=GRAY, h="center")
put(ws, f"H{t+1}", "Participação", f=font(10, True), bg=GRAY, h="center", merge=f"H{t+1}:I{t+1}")
ws.row_dimensions[t + 1].height = 21.75
src_col = {"S": "E", "W": "H", "O": "J", "T": "K"}
for j, q in enumerate("SWOT"):
    row = t + 2 + j
    put(ws, f"B{row}", f"{q} — {QD[q][2]}", f=font(10, True, c=WHITE), bg=QD[q][0], merge=f"B{row}:D{row}")
    put(ws, f"E{row}", QD[q][4].capitalize(), bg=QD[q][1])
    calc(ws, f"F{row}", f"=SWOT!{src_col[q]}41")
    calc(ws, f"G{row}", f"=SWOT!{src_col[q]}42")
    calc(ws, f"H{row}", f"=IF(SUM($G${t+2}:$G${t+5})=0,0,G{row}/SUM($G${t+2}:$G${t+5}))", fmt="0%", merge=f"H{row}:I{row}")
    ws.row_dimensions[row].height = 21.75
g = t + 7
band(ws, g, "Gráfico: pontos por quadrante", "I")
ch = BarChart()
ch.type = "col"
ch.height, ch.width = 8.5, 17
ch.style = 2
ch.title = None
ch.legend = None
ch.gapWidth = 120
ch.add_data(Reference(ws, min_col=7, min_row=t + 1, max_row=t + 5), titles_from_data=True)
ch.set_categories(Reference(ws, min_col=2, min_row=t + 2, max_row=t + 5))
srs = ch.series[0]
srs.graphicalProperties.solidFill = MUTED
for j, q in enumerate("SWOT"):
    pt = DataPoint(idx=j)
    pt.graphicalProperties = GraphicalProperties(solidFill=QD[q][0])
    pt.graphicalProperties.line.noFill = True
    srs.dPt.append(pt)
srs.dLbls = DataLabelList()
srs.dLbls.showVal = True
srs.dLbls.showSerName = srs.dLbls.showCatName = srs.dLbls.showLegendKey = False
ch.x_axis.delete = False
ch.y_axis.delete = False
ch.y_axis.scaling.min = 0
ch.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill="D3DBE0", w=9525))
ch.x_axis.graphicalProperties = GraphicalProperties(ln=LineProperties(solidFill="D3DBE0", w=9525))
ch.y_axis.graphicalProperties = GraphicalProperties(ln=LineProperties(noFill=True))
ws.add_chart(ch, f"B{g+1}")
for rr in range(g + 1, g + 19):
    ws.row_dimensions[rr].height = 15
setup(ws, QD["W"][0], f"B1:I{g+18}", landscape=False, fit_height=True)

# ------------------------------------------------------------------ SWOT cruzada
ws = wb.create_sheet("SWOT cruzada")
widths(ws, {"A": 2, "B": 5, "C": 11, "D": 38, "E": 11, "F": 38, "G": 22, "H": 46, "I": 13, "J": 2})
title(ws, "SWOT cruzada — estratégias",
      "Escolha o código de um fator interno e de um fator externo. O texto dos fatores vem da aba SWOT.", "I")
for col, text in zip("BCDEFGHI", ["#", "Interno", "Fator interno", "Externo", "Fator externo", "Cruzamento", "Estratégia", "Prioridade"]):
    head(ws, f"{col}4", text)
ws.row_dimensions[4].height = 21.75
ex(ws, "B5", "Ex.", h="center")
ex(ws, "C5", "W2", h="center")
ex(ws, "D5", "Dependência de um aplicativo para 70% dos pedidos")
ex(ws, "E5", "O4", h="center")
ex(ws, "F5", "Pedido por aplicativo de mensagens, com custo menor por venda")
ex(ws, "G5", "WO — Crescimento", h="center")
ex(ws, "H5", "Levar 30% dos pedidos do aplicativo para o canal próprio")
ex(ws, "I5", "Alta", h="center")
ws.row_dimensions[5].height = 31.5
C1, C2 = 6, 17


def lookup(code, qa, qb):
    fa, _, _, _, a1, a2 = cells(qa)
    fb, _, _, _, b1, b2 = cells(qb)
    n = f"VALUE(MID({code},2,1))"
    return (f'=IF({code}="","",IF(LEFT({code},1)="{qa}",INDEX(SWOT!${fa}${a1}:${fa}${a2},{n}),'
            f'INDEX(SWOT!${fb}${b1}:${fb}${b2},{n}))&"")')


for k in range(12):
    row = C1 + k
    num(ws, f"B{row}", k + 1)
    inp(ws, f"C{row}", h="center")
    calc(ws, f"D{row}", lookup(f"C{row}", "S", "W"), h="left", b=False)
    inp(ws, f"E{row}", h="center")
    calc(ws, f"F{row}", lookup(f"E{row}", "O", "T"), h="left", b=False)
    calc(ws, f"G{row}", f'=IF(OR(C{row}="",E{row}=""),"",LEFT(C{row},1)&LEFT(E{row},1)&" — "&'
         f'IF(LEFT(C{row},1)="S",IF(LEFT(E{row},1)="O","Desenvolvimento","Manutenção"),'
         f'IF(LEFT(E{row},1)="O","Crescimento","Sobrevivência")))', b=False)
    inp(ws, f"H{row}")
    inp(ws, f"I{row}", h="center")
    ws.row_dimensions[row].height = 36
dv_list(ws, f"C{C1}:C{C2}", [f"{q}{i}" for q in "SW" for i in range(1, N + 1)], "Código de uma força (S) ou de uma fraqueza (W)")
dv_list(ws, f"E{C1}:E{C2}", [f"{q}{i}" for q in "OT" for i in range(1, N + 1)], "Código de uma oportunidade (O) ou de uma ameaça (T)")
dv_list(ws, f"I{C1}:I{C2}", ["Alta", "Média", "Baixa"], "Alta, Média ou Baixa")
cf_equal(ws, f"I{C1}:I{C2}", PRIO_CF)
band(ws, 19, "Resumo automático", "I")
summary(ws, 20, "Estratégias registradas", f"=COUNTA(H{C1}:H{C2})", "F", "G")
summary(ws, 21, "SO — Desenvolvimento", f'=COUNTIF(G{C1}:G{C2},"SO*")', "F", "G")
summary(ws, 22, "WO — Crescimento", f'=COUNTIF(G{C1}:G{C2},"WO*")', "F", "G")
summary(ws, 23, "ST — Manutenção", f'=COUNTIF(G{C1}:G{C2},"ST*")', "F", "G")
summary(ws, 24, "WT — Sobrevivência", f'=COUNTIF(G{C1}:G{C2},"WT*")', "F", "G")
summary(ws, 25, "Estratégias de prioridade alta", f'=COUNTIF(I{C1}:I{C2},"Alta")', "F", "G")
summary(ws, 26, "Aviso: cruzamento",
        f'=IF(SUMPRODUCT((H{C1}:H{C2}<>"")*((C{C1}:C{C2}="")+(E{C1}:E{C2}="")>0))>0,"Há estratégia sem fator","OK")', "F", "G")
summary(ws, 27, "Aviso: fatores",
        f'=IF(SUMPRODUCT(((C{C1}:C{C2}<>"")*(D{C1}:D{C2}=""))+((E{C1}:E{C2}<>"")*(F{C1}:F{C2}="")))>0,"Há código sem fator na aba SWOT","OK")', "F", "G")
for rr in (26, 27):
    ws[f"G{rr}"].font = font(9)
cf_ok(ws, "G26:G27")
ws.freeze_panes = "C5"
setup(ws, QD["O"][0], "B1:I27", fit_height=True)

# ------------------------------------------------------------------ Plano de ação
ws = wb.create_sheet("Plano de ação")
widths(ws, {"A": 2, "B": 5, "C": 34, "D": 38, "E": 18, "F": 13, "G": 30, "H": 15, "I": 14, "J": 2})
title(ws, "Plano de ação", "Uma linha por ação. Cada ação vem de uma estratégia e tem um único responsável.", "I")
for col, text in zip("BCDEFGHI", ["#", "Estratégia ou objetivo", "Ação", "Responsável", "Prazo", "Indicador de acompanhamento", "Status", "Situação"]):
    head(ws, f"{col}4", text)
ws.row_dimensions[4].height = 21.75
ex(ws, "B5", "Ex.", h="center")
ex(ws, "C5", "Levar 30% dos pedidos do aplicativo para o canal próprio")
ex(ws, "D5", "Divulgar o pedido por mensagem na embalagem e nas redes sociais")
ex(ws, "E5", "Gerente da loja")
ex(ws, "F5", date(2026, 11, 30), h="center", fmt=DATE)
ex(ws, "G5", "% de pedidos recebidos pelo canal próprio")
ex(ws, "H5", "Em andamento", h="center")
ex(ws, "I5", "No prazo", h="center")
ws.row_dimensions[5].height = 31.5
A1, A2 = 6, 20
for k in range(15):
    row = A1 + k
    num(ws, f"B{row}", k + 1)
    for col in "CDEG":
        inp(ws, f"{col}{row}")
    inp(ws, f"F{row}", h="center", fmt=DATE)
    inp(ws, f"H{row}", h="center")
    calc(ws, f"I{row}", f'=IF(D{row}="","",IF(OR(H{row}="Concluída",H{row}="Cancelada"),H{row},'
         f'IF(F{row}="","Sem prazo",IF(F{row}<TODAY(),"Atrasada","No prazo"))))', b=False)
    ws.row_dimensions[row].height = 30
dv_list(ws, f"H{A1}:H{A2}", ["Não iniciada", "Em andamento", "Concluída", "Cancelada"], "Não iniciada, Em andamento, Concluída ou Cancelada")
dv_date(ws, f"F{A1}:F{A2}")
cf_equal(ws, f"H{A1}:H{A2}", STATUS_CF)
cf_equal(ws, f"I{A1}:I{A2}", [("Concluída", GREEN), ("No prazo", GREEN), ("Sem prazo", YELLOW), ("Atrasada", RED)])
band(ws, 22, "Resumo automático", "I")
summary(ws, 23, "Ações registradas", f"=COUNTA(D{A1}:D{A2})", "F", "G")
summary(ws, 24, "Concluídas", f'=COUNTIF(H{A1}:H{A2},"Concluída")', "F", "G")
summary(ws, 25, "Em andamento", f'=COUNTIF(H{A1}:H{A2},"Em andamento")', "F", "G")
summary(ws, 26, "Não iniciadas", f'=COUNTIF(H{A1}:H{A2},"Não iniciada")', "F", "G")
summary(ws, 27, "Atrasadas", f'=COUNTIF(I{A1}:I{A2},"Atrasada")', "F", "G")
summary(ws, 28, "Ações sem responsável ou sem prazo", f'=SUMPRODUCT((D{A1}:D{A2}<>"")*((E{A1}:E{A2}="")+(F{A1}:F{A2}="")>0))', "F", "G")
summary(ws, 29, "Percentual concluído", "=IF(G23=0,0,G24/G23)", "F", "G", fmt="0%")
ws.freeze_panes = "C5"
setup(ws, QD["T"][0], "B1:I29", fit_height=True)

# ------------------------------------------------------------------ Checklist
ws = wb.create_sheet("Checklist")
widths(ws, {"A": 2, "B": 5, "C": 70, "D": 14, "E": 44, "F": 2})
title(ws, "Checklist de validação da Matriz SWOT", "Escolha o status de cada item na lista suspensa (coluna em amarelo-claro).", "E")
for col, text in zip("BCDE", ["#", "Item de verificação", "Status", "Observações"]):
    head(ws, f"{col}4", text)
ws.row_dimensions[4].height = 21.75
for k, text in enumerate([
    "O objeto, o objetivo e o horizonte da análise estão definidos.",
    "Participaram pessoas de áreas e níveis diferentes.",
    "Cada fator é um fato ou uma tendência que pode ser verificada.",
    "Forças e fraquezas são internas: a organização controla.",
    "Oportunidades e ameaças são externas: existiriam sem a organização.",
    "Nenhum fator é uma ação ou uma solução disfarçada.",
    "Cada quadrante tem de 3 a 8 fatores.",
    "Os fatores foram pontuados pelo grupo, com critérios combinados antes.",
    "Os fatores de alta prioridade foram cruzados em estratégias.",
    "Cada estratégia escolhida virou ação, com responsável, prazo e indicador.",
    "As expectativas das partes interessadas foram consideradas.",
    "A data da próxima revisão está marcada.",
]):
    row = 5 + k
    num(ws, f"B{row}", k + 1)
    put(ws, f"C{row}", text, bg=GRAY)
    inp(ws, f"D{row}", h="center")
    inp(ws, f"E{row}")
    ws.row_dimensions[row].height = 30
dv_list(ws, "D5:D16", YESNO, "Sim, Parcial ou Não")
cf_equal(ws, "D5:D16", YESNO_CF)
band(ws, 18, "Resumo automático", "E")
summary(ws, 19, "Itens atendidos (Sim)", '=COUNTIF(D5:D16,"Sim")', "C", "D")
summary(ws, 20, "Itens parciais", '=COUNTIF(D5:D16,"Parcial")', "C", "D")
summary(ws, 21, "Itens não atendidos", '=COUNTIF(D5:D16,"Não")', "C", "D")
summary(ws, 22, "Itens sem resposta", "=COUNTBLANK(D5:D16)", "C", "D")
summary(ws, 23, "Percentual atendido", "=D19/ROWS(D5:D16)", "C", "D", fmt="0%")
summary(ws, 24, "Situação", '=IF(D19=ROWS(D5:D16),"Validado",IF(D22>0,"Em andamento","Com pendências"))', "C", "D")
setup(ws, "7A4A9A", "B1:E24", landscape=False, fit_height=True)

# ------------------------------------------------------------------ Exemplos
swot_sheet(wb.create_sheet("Exemplo 1 - Pizzaria"), MUTED, EX1)
swot_sheet(wb.create_sheet("Exemplo 2 - Compras"), MUTED, EX2)

wb.active = 1
wb.save(OUT)
print("ok", OUT)
