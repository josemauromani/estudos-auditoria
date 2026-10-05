# -*- coding: utf-8 -*-
"""Gera Calibracao-modelo.xlsx no mesmo padrão visual das planilhas anteriores da série."""
import math
import os
import sys

_here = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, _here)
# funções de formatação compartilhadas com o gerador do PDCA (bloco inicial do arquivo)
_src = open(os.path.join(_here, "build_pdca.py"), encoding="utf-8").read()
exec(_src[: _src.index("STATUS = [")])

from openpyxl.chart import BarChart, Series  # noqa: E402
from openpyxl.chart.label import DataLabelList  # noqa: E402
from cal_data import (ADEQ, APROV, CHECK, EMDIA, EX1, EX2, FORA, GROSSA, JANELA, NAO, RAZAO, REPROV, SEMCAL, SIM, SITS, TIPOS, TIPOS_REG, VENCE, VENCIDO)  # noqa: E402

BLUE, TEAL, AMBER, PURPLE, REDC = "2B5C8A", "1E7B73", "A96A12", "7A4A9A", "B0413E"
S1_T, S2_T, S3_T = "DCE8F3", "F8EBCB", "F5DEDC"
YESNO = ["Sim", "Parcial", "Não"]
YESNO_CF = [("Sim", GREEN), ("Parcial", YELLOW), ("Não", RED)]
NI, NR = 40, 80
IN, CA = "Instrumentos", "'Calibrações'"
I1, I2 = 8, 8 + NI - 1          # instrumentos
C1, C2 = 7, 7 + NR - 1          # registros de calibração e verificação
GERAL = "#,##0.###"
SIT_CF = [(EMDIA, GREEN), (VENCE, S2_T), (VENCIDO, RED), (SEMCAL, RED), (FORA, GRAY)]
SEMAV = "Falta avaliar os resultados anteriores"


def note(ws, ref, text):
    c = Comment(text, "Modelo Calibração")
    c.width, c.height = 300, 120
    ws[ref].comment = c


def cf_warn(ws, rng_, first, ok_values=("OK",)):
    cond = ",".join(f'{first}="{v}"' for v in ok_values)
    ws.conditional_formatting.add(rng_, FormulaRule(formula=[f"OR({cond})"], fill=PatternFill("solid", bgColor=GREEN, fgColor=GREEN)))
    ws.conditional_formatting.add(rng_, FormulaRule(formula=[f"AND(LEN({first})>0,NOT(OR({cond})))"],
                                  fill=PatternFill("solid", bgColor=YELLOW, fgColor=YELLOW)))


def cf_texto(ws, rng_, first, pares, resto=None):
    for t, cor in pares:
        ws.conditional_formatting.add(rng_, FormulaRule(formula=[f'{first}="{t}"'], stopIfTrue=True, fill=PatternFill("solid", bgColor=cor, fgColor=cor)))
    if resto:
        ws.conditional_formatting.add(rng_, FormulaRule(formula=[f"LEN({first})>0"], fill=PatternFill("solid", bgColor=resto, fgColor=resto)))


def hint_row(ws, row, cells, height=30):
    for ref, text in cells:
        put(ws, f"{ref}{row}", text, f=font(9, i=True, c=MUTED), bg=GRAY, h="center")
    ws.row_dimensions[row].height = height


def alt(pares, minimo=21.75, linha=12):
    n = max(max(1, math.ceil(len(str(t)) / max(c, 1))) for t, c in pares)
    return max(minimo, n * linha + 7)


# ------------------------------------------------------------------ fórmulas compartilhadas pelas abas de entrada e pelos exemplos
def f_prox(ult, interv):
    return f'IF(OR({ult}="",{interv}=""),"",EDATE({ult},{interv}))'


def f_sit(cod, emuso, ult, prox, ref):
    return (f'IF({cod}="","",IF({emuso}="{NAO}","{FORA}",IF({ult}="","{SEMCAL}",IF(AND(ISNUMBER({ref}),{prox}<{ref}),"{VENCIDO}",'
            f'IF(AND(ISNUMBER({ref}),{prox}-{ref}<MIN({JANELA},({prox}-{ult})/4)),"{VENCE}","{EMDIA}")))))')


def f_adeq(tol, res):
    return f'IF(OR({tol}="",{res}=""),"",IF({res}<={tol}/{RAZAO}+0.000000001,"{ADEQ}","{GROSSA}"))'


def f_inst(c):
    return (f'IF({c["cod"]}="","",IF({c["uso"]}="","Falta o que mede",IF({c["tipo"]}="","Falta o tipo de controle",IF({c["interv"]}="","Falta o intervalo",'
            f'IF({c["ema"]}="","Falta o erro máximo admissível",IF({c["sit"]}="{FORA}","{FORA}",IF({c["sit"]}="{VENCIDO}","Vencido: bloquear ou calibrar",'
            f'IF({c["sit"]}="{SEMCAL}","Sem calibração ou verificação",IF({c["adeq"]}="{GROSSA}","Resolução grossa para a tolerância","OK")))))))))')


def f_res(cod, erro, inc, crit):
    return f'IF(OR({cod}="",{erro}="",{crit}=""),"",IF(ABS({erro})+N({inc})<={crit}+0.000000001,"{APROV}","{REPROV}"))'


def f_cal(c, outros, codes):
    return (f'IF({c["cod"]}="",IF({outros}>0,"Falta o instrumento",""),IF(ISNA(MATCH({c["cod"]},{codes},0)),"Instrumento não cadastrado",'
            f'IF({c["crit"]}="","Falta o critério no cadastro",IF({c["reg"]}="","Falta o registro",IF({c["erro"]}="","Falta o erro encontrado",'
            f'IF(AND({c["res"]}="{REPROV}",{c["acao"]}=""),"Falta a ação",IF(AND({c["res"]}="{REPROV}",{c["imp"]}=""),"{SEMAV}","OK")))))))')


def chart_style(ch):
    ch.style = 2
    ch.title = None
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ch.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill="D3DBE0", w=9525))
    ch.x_axis.graphicalProperties = GraphicalProperties(ln=LineProperties(solidFill="D3DBE0", w=9525))
    ch.y_axis.graphicalProperties = GraphicalProperties(ln=LineProperties(noFill=True))


def colunas(ws, anchor, r1, r2, c_cat, c_val, width=16, height=7.5):
    ch = BarChart()
    ch.type = "col"
    ch.gapWidth = 70
    ch.height, ch.width = height, width
    chart_style(ch)
    ch.legend = None
    s = Series(Reference(ws, min_col=c_val, min_row=r1, max_row=r2), title="Instrumentos")
    s.graphicalProperties.solidFill = BLUE
    s.graphicalProperties.line.noFill = True
    s.dLbls = DataLabelList()
    s.dLbls.showVal = True
    s.dLbls.showSerName = s.dLbls.showCatName = s.dLbls.showLegendKey = False
    ch.append(s)
    ch.set_categories(Reference(ws, min_col=c_cat, min_row=r1, max_row=r2))
    ch.y_axis.scaling.min = 0
    ch.y_axis.majorUnit = 1
    ws.add_chart(ch, anchor)


# ================================================================== pasta de trabalho
wb = Workbook()
wb.properties.title = "Calibração e recursos de medição — Modelo"
wb.properties.creator = "Modelo Calibração"
wb.properties.language = "pt-BR"

# ------------------------------------------------------------------ Instruções
ws = wb.active
ws.title = "Instruções"
widths(ws, {"A": 2, "B": 26, "C": 90, "D": 2})
title(ws, "Calibração e recursos de medição — Como usar esta planilha",
      "Modelo para cadastrar os instrumentos, registrar as calibrações e as verificações, ler os resultados e acompanhar os vencimentos.", "C")
r = 4


def section(text):
    global r
    band(ws, r, text, "C", sz=11)
    r += 1


def line(k, v, kbg=GRAY, vbg=None, kf=None, kh="left", height=None):
    global r
    put(ws, f"B{r}", k, f=kf or font(10, True), bg=kbg, h=kh)
    put(ws, f"C{r}", v, bg=vbg)
    ws.row_dimensions[r].height = height or (19.5 if len(v) <= 96 else (31.5 if len(v) <= 192 else 45.75))
    r += 1


section("Legenda: onde preencher")
line("Amarelo-claro", "Células de entrada. É aqui que você digita.", vbg=INPUT)
line("Cinza", "Células calculadas ou fixas (próxima data, situação, adequação, resultado, conferências). Não altere.", vbg=GRAY)
line("Cinza em itálico", "Linha de exemplo (marcada com “Ex.”). Mostra o formato esperado e não entra nas contagens.", vbg=GRAY)
line("Listas suspensas", "Tipo de controle, em uso, código do instrumento, tipo do registro e checklist aceitam apenas as opções da lista.")
line("Comentários", "Células com triângulo vermelho no canto trazem uma dica de preenchimento. Passe o mouse sobre elas.")
r += 1
section("As regras de cálculo")
for k, text in [
    ("Próxima data", "A última calibração ou verificação mais o intervalo, em meses."),
    (EMDIA, f"Até a próxima data faltam mais dias do que o aviso: um quarto do intervalo, com no máximo {JANELA} dias."),
    (VENCE, f"Faltam menos dias do que o aviso. Na verificação mensal, cerca de uma semana; no intervalo de um ano, {JANELA} dias."),
    (VENCIDO, "A próxima data já passou. O instrumento deve ser bloqueado até a calibração."),
    (SEMCAL, "Não há data de calibração nem de verificação."),
    (FORA, "O instrumento está marcado como fora de uso: reprovado, quebrado, substituído."),
    ("Adequação", f"Adequado quando a resolução cabe ao menos {RAZAO} vezes na tolerância do que o instrumento mede."),
    ("Resultado", f"{APROV} quando o valor absoluto do erro, somado à incerteza, é menor ou igual ao erro máximo admissível do instrumento."),
]:
    line(k, text)
r += 1
section("Ordem recomendada de preenchimento")
for k, text in enumerate([
    "Aba Instrumentos: informe a data da leitura.",
    "Aba Instrumentos: cadastre cada instrumento que decide a conformidade, com o que mede, a tolerância, a resolução, o controle, o intervalo e o erro máximo admissível.",
    "Aba Instrumentos: informe a data da última calibração ou verificação e se o instrumento está em uso.",
    "Aba Calibrações: lance cada certificado e cada verificação, com o erro e a incerteza.",
    "Aba Calibrações: para cada reprovado, escreva a ação e a avaliação dos resultados anteriores.",
    "Aba Painel: leia os avisos. Depois, valide na aba Checklist.",
], 1):
    line(f"Passo {k}", text)
r += 1
section("Abas da planilha")
for k, text in [
    (IN, f"Até {NI} instrumentos. Calcula a próxima data, a situação, a adequação e a conferência."),
    ("Calibrações", f"Até {NR} registros. Calcula o resultado pelo critério do instrumento e a conferência."),
    ("Painel", "Os instrumentos por situação, os reprovados, o gráfico e o aviso."),
    ("Checklist", "Doze verificações do controle de instrumentos, com percentual de conclusão."),
    ("Exemplo 1 - Pizzaria", "Os instrumentos da loja, lidos em 31/03/2027."),
    ("Exemplo 2 - Indústria", "Os instrumentos da fábrica, lidos em 30/06/2027."),
]:
    line(k, text)
r += 1
section("Regras e premissas do modelo")
for k, text in [
    ("Tolerância", "A amplitude da faixa aceita: de 38 a 42 µm, 4 µm. Para um limite só, use a margem que importa: no mínimo 65 °C, com 5 °C de margem."),
    ("Erro máximo admissível", "Digite só o número, sem o sinal: 1 significa ±1. Use a mesma unidade do instrumento."),
    ("Última data", "Atualize a cada calibração ou verificação, inclusive quando o resultado foi reprovado."),
    ("Exemplos", "Os dados das abas de exemplo são ilustrativos, criados para este material."),
]:
    line(k, text)
setup(ws, INK, f"B1:C{r}", landscape=False, fit_height=True)

# ------------------------------------------------------------------ Instrumentos
ws = wb.create_sheet(IN)
widths(ws, {"A": 2, "B": 5, "C": 10, "D": 26, "E": 14, "F": 30, "G": 8, "H": 11, "I": 11, "J": 18, "K": 10, "L": 12, "M": 11, "N": 9, "O": 12, "P": 16, "Q": 16, "R": 30, "S": 2})
title(ws, "Instrumentos de medição", "Um instrumento por linha: o que mede, se é adequado, como é controlado e quando vence.", "R")
label(ws, "B4", "Data da leitura", merge="B4:D4")
inp(ws, "E4", h="center", fmt=DATE)
put(ws, "F4", "Usada para saber o que venceu e o que vence em breve. Atualize a cada leitura.", f=font(9, i=True, c=MUTED), bg=GRAY, merge="F4:R4")
dv_date(ws, "E4")
REF = "$E$4"
for col, text in zip("BCDEFGHIJKLMNOPQR", ["#", "Código", "Instrumento", "Local", "O que mede", "Unidade", "Tolerância", "Resolução", "Controle", "Intervalo (meses)", "Última", "Erro máx. (±)",
                                           "Em uso?", "Próxima", "Situação", "Adequação", "Conferência"]):
    head(ws, f"{col}6", text)
ws.row_dimensions[6].height = 33
hint_row(ws, 7, [("B", ""), ("C", "TER-01"), ("D", ""), ("E", ""), ("F", "Com o critério"), ("G", "°C, g, µm"), ("H", "Amplitude"), ("I", "Menor passo"), ("J", "Lista"), ("K", "Número"),
                 ("L", "Data"), ("M", "Número"), ("N", "Lista"), ("O", "Calculada"), ("P", "Calculada"), ("Q", "Calculada"), ("R", "Calculada")])
for k in range(NI):
    rr = I1 + k
    num(ws, f"B{rr}", k + 1)
    inp(ws, f"C{rr}", h="center")
    for col in "DEF":
        inp(ws, f"{col}{rr}")
    for col in "GJN":
        inp(ws, f"{col}{rr}", h="center")
    for col in "HIKM":
        inp(ws, f"{col}{rr}", h="center", fmt=GERAL)
    inp(ws, f"L{rr}", h="center", fmt=DATE)
    calc(ws, f"O{rr}", "=" + f_prox(f"L{rr}", f"K{rr}"), b=False, fmt=DATE)
    calc(ws, f"P{rr}", "=" + f_sit(f"C{rr}", f"N{rr}", f"L{rr}", f"O{rr}", REF), sz=9)
    calc(ws, f"Q{rr}", "=" + f_adeq(f"H{rr}", f"I{rr}"), b=False, sz=9)
    cells = dict(cod=f"C{rr}", uso=f"F{rr}", tipo=f"J{rr}", interv=f"K{rr}", ema=f"M{rr}", sit=f"P{rr}", adeq=f"Q{rr}")
    calc(ws, f"R{rr}", "=" + f_inst(cells), b=False, sz=9)
    ws.row_dimensions[rr].height = 30
dv_list(ws, f"J{I1}:J{I2}", TIPOS, "Calibração externa, por laboratório, ou verificação interna, contra um padrão próprio")
dv_list(ws, f"N{I1}:N{I2}", [SIM, NAO], "O instrumento está em uso? Não: reprovado, quebrado, substituído")
dv_number(ws, f"H{I1}:I{I2}")
dv_number(ws, f"K{I1}:K{I2}")
dv_number(ws, f"M{I1}:M{I2}")
dv_date(ws, f"L{I1}:L{I2}")
cf_texto(ws, f"P{I1}:P{I2}", f"P{I1}", SIT_CF)
cf_texto(ws, f"Q{I1}:Q{I2}", f"Q{I1}", [(ADEQ, GREEN), (GROSSA, RED)])
cf_warn(ws, f"R{I1}:R{I2}", f"R{I1}", ok_values=("OK", FORA))
note(ws, "H6", "A amplitude da faixa aceita do que o instrumento mede. De 38 a 42 µm: 4. Para um limite só, a margem que importa.")
note(ws, "M6", "O maior erro aceito no instrumento, sem o sinal: 1 significa ±1. Costuma ser uma fração da tolerância.")
s = I2 + 2
band(ws, s, "Resumo automático", "R")
P_ = f"P{I1}:P{I2}"
for k, (text, formula) in enumerate([
    ("Instrumentos cadastrados", f"=COUNTA(C{I1}:C{I2})"),
    ("Por situação", "=" + '&", "&'.join(f'COUNTIF({P_},"{x}")&" {x.lower()}"' for x in SITS)),
    ("Resolução grossa", f'=COUNTIF(Q{I1}:Q{I2},"{GROSSA}")'),
    ("Aviso", f'=IF(E{s+1}=0,"Cadastre os instrumentos",IF(E4="","Informe a data da leitura",IF(COUNTIF({P_},"{VENCIDO}")>0,"Há instrumento vencido: bloquear até calibrar",'
              f'IF(COUNTIF({P_},"{SEMCAL}")>0,"Há instrumento sem calibração ou verificação",IF(COUNTIF(Q{I1}:Q{I2},"{GROSSA}")>0,"Há instrumento com resolução grossa para a tolerância",'
              f'IF(COUNTIF({P_},"{VENCE}")>0,"Há instrumento que vence em breve: programe",IF(SUMPRODUCT((R{I1}:R{I2}<>"")*(R{I1}:R{I2}<>"OK")*(R{I1}:R{I2}<>"{FORA}"))>0,'
              f'"Há instrumento a completar: veja a coluna Conferência","OK")))))))'),
], 1):
    label(ws, f"B{s+k}", text, merge=f"B{s+k}:D{s+k}", h="right")
    calc(ws, f"E{s+k}", formula, merge=f"E{s+k}:R{s+k}", sz=9 if text in ("Aviso", "Por situação") else 10, b=text not in ("Aviso", "Por situação"), h="left")
    ws.row_dimensions[s + k].height = 21.75
IAV = s + 4
cf_warn(ws, f"E{IAV}:R{IAV}", f"E{IAV}")
ws.freeze_panes = f"E{I1}"
setup(ws, BLUE, f"B1:R{IAV}")

# ------------------------------------------------------------------ Calibrações
ws = wb.create_sheet("Calibrações")
widths(ws, {"A": 2, "B": 5, "C": 12, "D": 11, "E": 24, "F": 13, "G": 26, "H": 18, "I": 9, "J": 10, "K": 10, "L": 12, "M": 32, "N": 40, "O": 32, "P": 2})
title(ws, "Calibrações e verificações", "Um certificado ou uma verificação por linha. O resultado vem do critério do instrumento.", "O")
for col, text in zip("BCDEFGHIJKLMNO", ["#", "Data", "Código", "Instrumento", "Tipo", "Registro ou certificado", "Ponto medido", "Erro", "Incerteza", "Critério (±)", "Resultado",
                                         "Ação, se reprovado", "Avaliação dos resultados anteriores", "Conferência"]):
    head(ws, f"{col}4", text)
ws.row_dimensions[4].height = 33
hint_row(ws, 5, [("B", ""), ("C", "Data"), ("D", "Lista"), ("E", "Vem do cadastro"), ("F", "Lista"), ("G", "Número do documento"), ("H", "Valor de referência"), ("I", "Com sinal"),
                 ("J", "Número"), ("K", "Vem do cadastro"), ("L", "Calculado"), ("M", "Ajuste, reparo, troca"), ("N", "O que ele mediu desde a última boa"), ("O", "Calculada")])
ex(ws, "B6", "Ex.", h="center")
for col, v, h_, fmt in [("C", date(2027, 3, 5), "center", DATE), ("D", "TER-01", "center", None), ("E", "Termômetro de espeto", "left", None), ("F", "Verificação", "center", None),
                        ("G", "Planilha FR-07", "left", None), ("H", "Água com gelo, 0 °C", "left", None), ("I", 0.3, "center", None), ("J", 0.2, "center", None), ("K", 1, "center", None),
                        ("L", APROV, "center", None), ("M", None, "left", None), ("N", None, "left", None), ("O", "OK", "center", None)]:
    ex(ws, f"{col}6", v, h=h_, fmt=fmt)
ws.row_dimensions[6].height = 30
CODES = f"{IN}!$C${I1}:$C${I2}"
for k in range(NR):
    rr = C1 + k
    num(ws, f"B{rr}", k + 1)
    inp(ws, f"C{rr}", h="center", fmt=DATE)
    inp(ws, f"D{rr}", h="center")
    calc(ws, f"E{rr}", f'=IF(D{rr}="","",IFERROR(INDEX({IN}!$D${I1}:$D${I2},MATCH(D{rr},{CODES},0))&"",""))', h="left", b=False, sz=9)
    inp(ws, f"F{rr}", h="center")
    for col in "GHMN":
        inp(ws, f"{col}{rr}")
    inp(ws, f"I{rr}", h="center", fmt=GERAL)
    inp(ws, f"J{rr}", h="center", fmt=GERAL)
    calc(ws, f"K{rr}", f'=IF(D{rr}="","",IFERROR(IF(ISNUMBER(INDEX({IN}!$M${I1}:$M${I2},MATCH(D{rr},{CODES},0))),INDEX({IN}!$M${I1}:$M${I2},MATCH(D{rr},{CODES},0)),""),""))',
         b=False, fmt=GERAL)
    calc(ws, f"L{rr}", "=" + f_res(f"D{rr}", f"I{rr}", f"J{rr}", f"K{rr}"), sz=9)
    cells = dict(cod=f"D{rr}", crit=f"K{rr}", reg=f"G{rr}", erro=f"I{rr}", res=f"L{rr}", acao=f"M{rr}", imp=f"N{rr}")
    calc(ws, f"O{rr}", "=" + f_cal(cells, f"COUNTA(C{rr},F{rr}:J{rr},M{rr}:N{rr})", CODES), b=False, sz=9)
    ws.row_dimensions[rr].height = 30
dv_date(ws, f"C{C1}:C{C2}")
dv = DataValidation(type="list", formula1=f"={CODES}", allow_blank=True)
dv.promptTitle, dv.prompt = "Opções", "Código do instrumento, da aba Instrumentos"
dv.showInputMessage = True
ws.add_data_validation(dv)
dv.add(f"D{C1}:D{C2}")
dv_list(ws, f"F{C1}:F{C2}", TIPOS_REG, "Calibração, com certificado, ou verificação interna")
dv_number(ws, f"I{C1}:J{C2}")
cf_texto(ws, f"L{C1}:L{C2}", f"L{C1}", [(APROV, GREEN), (REPROV, RED)])
cf_texto(ws, f"O{C1}:O{C2}", f"O{C1}", [("OK", GREEN), (SEMAV, RED)], resto=YELLOW)
note(ws, "I4", "O erro encontrado, com o sinal: o valor indicado pelo instrumento menos o valor do padrão. No certificado, pode aparecer como “erro” ou como “correção”, com o sinal trocado.")
note(ws, "N4", "Quando o instrumento é reprovado, o que ele mediu desde a última calibração boa fica em dúvida. Escreva o que foi revisto, o que se encontrou e o que foi feito.")
s = C2 + 2
band(ws, s, "Resumo automático", "O")
for k, (text, formula) in enumerate([
    ("Registros", f"=COUNTA(D{C1}:D{C2})"),
    ("Aprovados e reprovados", f'=COUNTIF(L{C1}:L{C2},"{APROV}")&" aprovados, "&COUNTIF(L{C1}:L{C2},"{REPROV}")&" reprovados"'),
    ("Aviso", f'=IF(E{s+1}=0,"Lance as calibrações e as verificações",IF(COUNTIF(O{C1}:O{C2},"{SEMAV}")>0,"Há reprovado sem a avaliação dos resultados anteriores",'
              f'IF(SUMPRODUCT((O{C1}:O{C2}<>"")*(O{C1}:O{C2}<>"OK"))>0,"Há registro a completar: veja a coluna Conferência","OK")))'),
], 1):
    label(ws, f"B{s+k}", text, merge=f"B{s+k}:D{s+k}", h="right")
    calc(ws, f"E{s+k}", formula, merge=f"E{s+k}:O{s+k}", sz=9 if text == "Aviso" else 10, b=text != "Aviso", h="left")
    ws.row_dimensions[s + k].height = 21.75
CAV = s + 3
cf_warn(ws, f"E{CAV}:O{CAV}", f"E{CAV}")
ws.freeze_panes = f"F{C1}"
setup(ws, TEAL, f"B1:O{CAV}")

# ------------------------------------------------------------------ Painel
ws = wb.create_sheet("Painel")
widths(ws, {"A": 2, "B": 34, "C": 16, "D": 14, "E": 14, "F": 34, "G": 2})
title(ws, "Painel dos instrumentos", "Nada a preencher: tudo vem das outras abas.", "F")
label(ws, "B3", "Data da leitura")
calc(ws, "C3", f'=IF({IN}!E4="","",{IN}!E4)', fmt=DATE)
for col, text, m in [("B", "Indicador", None), ("C", "Resultado", None), ("D", "Como se lê", "D5:F5")]:
    put(ws, f"{col}5", text, f=font(10, True, c=WHITE), bg=INK, h="center", merge=m)
ws.row_dimensions[5].height = 21.75
Ip, Iq, Ir = (f"{IN}!{c}{I1}:{c}{I2}" for c in "PQR")
Cl, Co = (f"{CA}!{c}{C1}:{c}{C2}" for c in "LO")
IND = [
    ("Instrumentos cadastrados", f"=COUNTA({IN}!C{I1}:C{I2})", "Da aba Instrumentos."),
    (EMDIA, f'=COUNTIF({Ip},"{EMDIA}")', f"Mais longe do que o aviso: um quarto do intervalo, no máximo {JANELA} dias."),
    (VENCE, f'=COUNTIF({Ip},"{VENCE}")', "Programar a calibração ou a verificação."),
    (VENCIDO, f'=COUNTIF({Ip},"{VENCIDO}")', "Bloquear até calibrar."),
    (SEMCAL, f'=COUNTIF({Ip},"{SEMCAL}")', "Verificar antes de usar de novo."),
    (FORA, f'=COUNTIF({Ip},"{FORA}")', "Retirados do posto."),
    ("Resolução grossa", f'=COUNTIF({Iq},"{GROSSA}")', f"A resolução não cabe {RAZAO} vezes na tolerância."),
    ("Registros de calibração e verificação", f"=COUNTA({CA}!D{C1}:D{C2})", "Da aba Calibrações."),
    ("Reprovados", f'=COUNTIF({Cl},"{REPROV}")', "Erro mais incerteza fora do critério."),
    ("Reprovados sem avaliação", f'=COUNTIF({Co},"{SEMAV}")', "O que o instrumento mediu antes não foi revisto."),
    ("Registros a completar", f'=SUMPRODUCT(({Co}<>"")*({Co}<>"OK"))', "Veja a coluna Conferência da aba Calibrações."),
]
for k, (nome, formula, leit) in enumerate(IND):
    rr = 6 + k
    label(ws, f"B{rr}", nome)
    calc(ws, f"C{rr}", formula)
    put(ws, f"D{rr}", leit, f=font(9, c=MUTED), bg=GRAY, merge=f"D{rr}:F{rr}")
    ws.row_dimensions[rr].height = 21.75
IR = {n: 6 + k for k, (n, _, _) in enumerate(IND)}
s = 6 + len(IND) + 1
band(ws, s, "Resumo automático", "F")
label(ws, f"B{s+1}", "Aviso")
calc(ws, f"C{s+1}", f'=IF(C{IR["Instrumentos cadastrados"]}=0,"Cadastre os instrumentos na aba Instrumentos",IF(C3="","Informe a data da leitura na aba Instrumentos",'
     f'IF(C{IR["Reprovados sem avaliação"]}>0,"Há reprovado sem a avaliação dos resultados anteriores",IF(C{IR[VENCIDO]}>0,"Há instrumento vencido: bloquear até calibrar",'
     f'IF(C{IR[SEMCAL]}>0,"Há instrumento sem calibração ou verificação",IF(C{IR["Resolução grossa"]}>0,"Há instrumento com resolução grossa para a tolerância",'
     f'IF(C{IR[VENCE]}>0,"Há instrumento que vence em breve: programe",IF(C{IR["Registros a completar"]}>0,"Há registro a completar","OK"))))))))',
     merge=f"C{s+1}:F{s+1}", sz=9, b=False, h="left")
cf_warn(ws, f"C{s+1}:F{s+1}", f"C{s+1}")
ws.row_dimensions[s + 1].height = 21.75
NAV = s + 1
GS = s + 3
for k, sit in enumerate(SITS):
    put(ws, f"B{GS + k}", sit, f=font(9), bg=GRAY)
    calc(ws, f"C{GS + k}", f'=COUNTIF({Ip},B{GS + k})', b=False)
colunas(ws, f"D{GS}", GS, GS + len(SITS) - 1, 2, 3, width=12, height=7)
setup(ws, REDC, f"B1:F{GS + 15}", landscape=False, fit_height=True)

# ------------------------------------------------------------------ Checklist
ws = wb.create_sheet("Checklist")
widths(ws, {"A": 2, "B": 5, "C": 70, "D": 14, "E": 44, "F": 2})
title(ws, "Checklist da calibração e dos recursos de medição", "Escolha o status de cada item na lista suspensa (coluna em amarelo-claro).", "E")
for col, text in zip("BCDE", ["#", "Item de verificação", "Status", "Observações"]):
    head(ws, f"{col}4", text)
ws.row_dimensions[4].height = 21.75
for k, text in enumerate(CHECK):
    rr = 5 + k
    num(ws, f"B{rr}", k + 1)
    put(ws, f"C{rr}", text, bg=GRAY)
    inp(ws, f"D{rr}", h="center")
    inp(ws, f"E{rr}")
    ws.row_dimensions[rr].height = 30
dv_list(ws, "D5:D16", YESNO, "Sim, Parcial ou Não")
cf_equal(ws, "D5:D16", YESNO_CF)
band(ws, 18, "Resumo automático", "E")
summary(ws, 19, "Itens atendidos (Sim)", '=COUNTIF(D5:D16,"Sim")', "C", "D")
summary(ws, 20, "Itens parciais", '=COUNTIF(D5:D16,"Parcial")', "C", "D")
summary(ws, 21, "Itens não atendidos", '=COUNTIF(D5:D16,"Não")', "C", "D")
summary(ws, 22, "Itens sem resposta", "=COUNTBLANK(D5:D16)", "C", "D")
summary(ws, 23, "Percentual atendido", "=D19/ROWS(D5:D16)", "C", "D", fmt="0%")
summary(ws, 24, "Situação", '=IF(D19=ROWS(D5:D16),"Validado",IF(D22>0,"Em andamento","Com pendências"))', "C", "D")
setup(ws, PURPLE, "B1:E24", landscape=False, fit_height=True)


# ------------------------------------------------------------------ Exemplos
def sub(ws, rr, cells, height=30):
    for a, b, t in cells:
        put(ws, f"{a}{rr}", t, f=font(10, True), bg=GRAY, h="center", merge=b and f"{a}{rr}:{b}{rr}")
    ws.row_dimensions[rr].height = height


def exemplo(ws, ex_):
    H = ex_["head"]
    widths(ws, {"A": 2, "B": 9, "C": 26, "D": 14, "E": 30, "F": 7, "G": 9, "H": 9, "I": 16, "J": 8, "K": 11, "L": 9, "M": 8, "N": 11, "O": 15, "P": 15, "Q": 28, "R": 2})
    title(ws, "Calibração e recursos de medição", "Exemplo preenchido, para consulta. Use as abas Instrumentos e Calibrações para a sua organização.", "Q")
    rr = 4
    for rot, val in [("Organização", H["org"]), ("Responsável", H["resp"]), ("Padrões próprios", H["padrao"]), ("Origem", H["origem"])]:
        label(ws, f"B{rr}", rot, merge=f"B{rr}:C{rr}")
        inp(ws, f"D{rr}", val, merge=f"D{rr}:Q{rr}", bg=WHITE, h="left")
        ws.row_dimensions[rr].height = alt([(val, 170)], minimo=19.5)
        rr += 1
    label(ws, f"B{rr}", "Data da leitura", merge=f"B{rr}:C{rr}")
    inp(ws, f"D{rr}", H["ref"], bg=WHITE, h="center", fmt=DATE)
    ref = f"$D${rr}"
    rr += 2
    band(ws, rr, "Instrumentos", "Q", color=BLUE)
    rr += 1
    sub(ws, rr, [("B", None, "Código"), ("C", None, "Instrumento"), ("D", None, "Local"), ("E", None, "O que mede"), ("F", None, "Unid."), ("G", None, "Tolerância"), ("H", None, "Resolução"),
                 ("I", None, "Controle"), ("J", None, "Meses"), ("K", None, "Última"), ("L", None, "Erro máx."), ("M", None, "Em uso"), ("N", None, "Próxima"), ("O", None, "Situação"),
                 ("P", None, "Adequação"), ("Q", None, "Conferência")])
    i1 = rr + 1
    for i in ex_["insts"]:
        rr += 1
        put(ws, f"B{rr}", i["cod"], f=font(10, True, c=MUTED), h="center")
        put(ws, f"C{rr}", i["nome"], f=font(10, True))
        put(ws, f"D{rr}", i["local"])
        put(ws, f"E{rr}", i["uso"])
        put(ws, f"F{rr}", i["unid"], h="center")
        put(ws, f"G{rr}", i["tol"], h="center", fmt=GERAL)
        put(ws, f"H{rr}", i["res"], h="center", fmt=GERAL)
        put(ws, f"I{rr}", i["tipo"], h="center", f=font(9))
        put(ws, f"J{rr}", i["interv"], h="center")
        put(ws, f"K{rr}", i["ultima"], h="center", fmt=DATE)
        put(ws, f"L{rr}", i["ema"], h="center", fmt=GERAL)
        put(ws, f"M{rr}", i["emuso"], h="center")
        calc(ws, f"N{rr}", "=" + f_prox(f"K{rr}", f"J{rr}"), b=False, fmt=DATE)
        calc(ws, f"O{rr}", "=" + f_sit(f"B{rr}", f"M{rr}", f"K{rr}", f"N{rr}", ref), sz=9)
        calc(ws, f"P{rr}", "=" + f_adeq(f"G{rr}", f"H{rr}"), b=False, sz=9)
        cells = dict(cod=f"B{rr}", uso=f"E{rr}", tipo=f"I{rr}", interv=f"J{rr}", ema=f"L{rr}", sit=f"O{rr}", adeq=f"P{rr}")
        calc(ws, f"Q{rr}", "=" + f_inst(cells), b=False, sz=9)
        ws.row_dimensions[rr].height = alt([(i["uso"], 30), (i["nome"], 26), (i["local"], 14)], minimo=21.75)
    i2 = rr
    cf_texto(ws, f"O{i1}:O{i2}", f"O{i1}", SIT_CF)
    cf_texto(ws, f"P{i1}:P{i2}", f"P{i1}", [(ADEQ, GREEN), (GROSSA, RED)])
    cf_warn(ws, f"Q{i1}:Q{i2}", f"Q{i1}", ok_values=("OK", FORA))
    rr += 2
    ws.row_breaks.append(Break(id=rr - 1))
    band(ws, rr, "Calibrações e verificações", "Q", color=TEAL)
    rr += 1
    sub(ws, rr, [("B", None, "Data"), ("C", None, "Código · instrumento"), ("D", None, "Tipo"), ("E", "F", "Registro · ponto"), ("G", None, "Erro"), ("H", None, "Incerteza"),
                 ("I", None, "Critério (±)"), ("J", "K", "Resultado"), ("L", "N", "Ação"), ("O", "P", "Resultados anteriores"), ("Q", None, "Conferência")])
    c1 = rr + 1
    codes = f"$B${i1}:$B${i2}"
    nome = {i["cod"]: i["nome"] for i in ex_["insts"]}
    for c in sorted(ex_["cals"], key=lambda c: c["data"]):
        rr += 1
        put(ws, f"B{rr}", c["data"], h="center", fmt=DATE)
        put(ws, f"C{rr}", f'{c["cod"]} · {nome[c["cod"]]}')
        put(ws, f"D{rr}", c["tipo"], h="center")
        put(ws, f"E{rr}", f'{c["reg"]} · {c["ponto"]}', merge=f"E{rr}:F{rr}", f=font(9))
        put(ws, f"G{rr}", c["erro"], h="center", fmt=GERAL)
        put(ws, f"H{rr}", c["inc"], h="center", fmt=GERAL)
        calc(ws, f"I{rr}", f'=IFERROR(INDEX($L${i1}:$L${i2},MATCH(LEFT(C{rr},FIND(" ·",C{rr})-1),{codes},0)),"")', b=False, fmt=GERAL)
        calc(ws, f"J{rr}", "=" + f_res(f"C{rr}", f"G{rr}", f"H{rr}", f"I{rr}"), sz=9, merge=f"J{rr}:K{rr}")
        put(ws, f"L{rr}", c["acao"] or None, merge=f"L{rr}:N{rr}", f=font(9))
        put(ws, f"O{rr}", c["impacto"] or None, merge=f"O{rr}:P{rr}", f=font(9))
        cells = dict(cod=f'LEFT(C{rr},FIND(" ·",C{rr})-1)', crit=f"I{rr}", reg=f"E{rr}", erro=f"G{rr}", res=f"J{rr}", acao=f"L{rr}", imp=f"O{rr}")
        calc(ws, f"Q{rr}", "=" + f_cal(cells, "1", codes), b=False, sz=9)
        ws.row_dimensions[rr].height = alt([(c["reg"] + c["ponto"], 40), (c["acao"], 30), (c["impacto"], 34), (nome[c["cod"]], 20)], minimo=21.75)
    c2 = rr
    cf_texto(ws, f"J{c1}:K{c2}", f"$J{c1}", [(APROV, GREEN), (REPROV, RED)])
    cf_texto(ws, f"Q{c1}:Q{c2}", f"Q{c1}", [("OK", GREEN), (SEMAV, RED)], resto=YELLOW)
    rr += 2
    band(ws, rr, "Resumo automático", "Q")
    for k, (text, formula) in enumerate([
        ("Instrumentos por situação", "=" + '&", "&'.join(f'COUNTIF(O{i1}:O{i2},"{x}")&" {x.lower()}"' for x in SITS)),
        ("Resolução grossa", f'=COUNTIF(P{i1}:P{i2},"{GROSSA}")'),
        ("Registros: aprovados, reprovados", f'=COUNTIF(J{c1}:J{c2},"{APROV}")&" aprovados, "&COUNTIF(J{c1}:J{c2},"{REPROV}")&" reprovados"'),
        ("Registros a completar", f'=SUMPRODUCT((Q{c1}:Q{c2}<>"OK")*1)'),
    ], 1):
        label(ws, f"B{rr+k}", text, merge=f"B{rr+k}:D{rr+k}", h="right")
        calc(ws, f"E{rr+k}", formula, merge=f"E{rr+k}:Q{rr+k}", h="left")
        ws.row_dimensions[rr + k].height = 21.75
    setup(ws, MUTED, f"B1:Q{rr + 4}")


exemplo(wb.create_sheet("Exemplo 1 - Pizzaria"), EX1)
exemplo(wb.create_sheet("Exemplo 2 - Indústria"), EX2)

wb.active = 1
wb.save(OUT)
print("ok", OUT, "| avisos: Instrumentos E%d, Calibrações E%d, Painel C%d" % (IAV, CAV, NAV))
