# -*- coding: utf-8 -*-
"""Gera Indicadores-modelo.xlsx no mesmo padrão visual das planilhas anteriores da série."""
import math
import os
import sys
from datetime import date

_here = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, _here)
# funções de formatação compartilhadas com o gerador do PDCA (bloco inicial do arquivo)
_src = open(os.path.join(_here, "build_pdca.py"), encoding="utf-8").read()
exec(_src[: _src.index("STATUS = [")])

from openpyxl.chart import Series  # noqa: E402
from openpyxl.utils import get_column_letter as L  # noqa: E402
from ind_data import ATENCAO, CHECK, EX1, EX2, FORA, MAIOR, MENOR, MESES, NA_META, SEGUIDOS  # noqa: E402

BLUE, TEAL, AMBER, PURPLE, REDC = "2B5C8A", "1E7B73", "A96A12", "7A4A9A", "B0413E"
S1, S2, S3 = "2A6FB0", "D19A2E", "B0413E"
S1_T, S2_T, S3_T = "DCE8F3", "F8EBCB", "F5DEDC"
SIT_CF = [(NA_META, S1_T), (ATENCAO, S2_T), (FORA, S3_T)]
ACAO_CF = [("Manter", S1_T), ("Acompanhar", S2_T), ("Analisar na reunião", S3_T), ("Abrir análise de causa", S3_T)]
TEND_CF = [("Melhorando", S1_T), ("Piorando", S3_T)]
YESNO = ["Sim", "Parcial", "Não"]
YESNO_CF = [("Sim", GREEN), ("Parcial", YELLOW), ("Não", RED)]
ST_ACAO = ["Não iniciada", "Em andamento", "Concluída"]
ST_CF = [("Concluída", GREEN), ("Em andamento", YELLOW), ("Não iniciada", RED)]
PRAZO_CF = [("Concluída", GREEN), ("No prazo", GREEN), ("Sem prazo", YELLOW), ("Atrasada", RED)]
FREQ = ["Diária", "Semanal", "Mensal", "Trimestral", "Semestral", "Anual"]
NI, NP = 12, 12  # indicadores e períodos do modelo
GERAL = "General"


def note(ws, ref, text):
    c = Comment(text, "Modelo Indicadores")
    c.width, c.height = 300, 120
    ws[ref].comment = c


def cf_warn(ws, rng_, first, ok_values=("OK",)):
    cond = ",".join(f'{first}="{v}"' for v in ok_values)
    ws.conditional_formatting.add(rng_, FormulaRule(formula=[f"OR({cond})"], fill=PatternFill("solid", bgColor=GREEN, fgColor=GREEN)))
    ws.conditional_formatting.add(rng_, FormulaRule(formula=[f"AND(LEN({first})>0,NOT(OR({cond})))"],
                                  fill=PatternFill("solid", bgColor=YELLOW, fgColor=YELLOW)))


def hint_row(ws, row, cells, height=30):
    for ref, text, merge in cells:
        put(ws, f"{ref}{row}", text, f=font(9, i=True, c=MUTED), bg=GRAY, h="center", merge=merge and f"{ref}{row}:{merge}{row}")
    ws.row_dimensions[row].height = height


def alt(text, chars, minimo=21.75, linha=12.75):
    n = max(1, math.ceil(len(str(text)) / max(chars, 1)))
    return max(minimo, n * linha + 9)


def chart_style(ch):
    ch.style = 2
    ch.title = None
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ch.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill="D3DBE0", w=9525))
    ch.x_axis.graphicalProperties = GraphicalProperties(ln=LineProperties(solidFill="D3DBE0", w=9525))
    ch.y_axis.graphicalProperties = GraphicalProperties(ln=LineProperties(noFill=True))


def line_chart(ws, anchor, row_per, row_val, row_meta, row_lim, c1, c2, width=24, height=8):
    """Linha do resultado, com a meta e o limite de atenção. Os dados estão em linhas, de c1 a c2."""
    ch = LineChart()
    ch.height, ch.width = height, width
    chart_style(ch)
    ch.legend.position = "b"
    for row, nome, cor, dash, larg in ((row_val, "Resultado", BLUE, None, 25400), (row_meta, "Meta", MUTED, "dash", 15875),
                                        (row_lim, "Limite de atenção", MUTED, "sysDot", 15875)):
        s = Series(Reference(ws, min_col=c1, max_col=c2, min_row=row), title=nome)
        s.graphicalProperties.line.solidFill = cor
        s.graphicalProperties.line.width = larg
        if dash:
            s.graphicalProperties.line.dashStyle = dash
            s.marker.symbol = "none"
        else:
            s.marker.symbol = "circle"
            s.marker.size = 7
            s.marker.graphicalProperties = GraphicalProperties(solidFill=cor)
            s.marker.graphicalProperties.line.solidFill = WHITE
        s.smooth = False
        ch.append(s)
    ch.set_categories(Reference(ws, min_col=c1, max_col=c2, min_row=row_per))
    ws.add_chart(ch, anchor)


def f_atende(v, sentido, meta):
    return f'IF({sentido}="{MAIOR}",{v}>={meta},{v}<={meta})'


def f_situacao(v, sentido, meta, lim):
    return (f'=IF(OR({v}="",{meta}="",{lim}="",{sentido}=""),"",IF({sentido}="{MAIOR}",IF({v}>={meta},"{NA_META}",IF({v}>={lim},"{ATENCAO}","{FORA}")),'
            f'IF({v}<={meta},"{NA_META}",IF({v}<={lim},"{ATENCAO}","{FORA}"))))')


def f_tendencia(n, a, b, sentido):
    return (f'=IF(OR({n}="",{n}<6),"Sem dados",IF(ABS({b}-{a})<=0.05*ABS({a}),"Estável",IF(({b}>{a})=({sentido}="{MAIOR}"),"Melhorando","Piorando")))')


def f_acao(sit, seg):
    return (f'=IF({sit}="","",IF(N({seg})>={SEGUIDOS},"Abrir análise de causa",IF({sit}="{FORA}","Analisar na reunião",'
            f'IF({sit}="{ATENCAO}","Acompanhar","Manter"))))')


# ================================================================== pasta de trabalho
wb = Workbook()
wb.properties.title = "Indicadores — Modelo de fichas, medições, painel e análise"
wb.properties.creator = "Modelo Indicadores"
wb.properties.language = "pt-BR"

# ------------------------------------------------------------------ Instruções
ws = wb.active
ws.title = "Instruções"
widths(ws, {"A": 2, "B": 24, "C": 92, "D": 2})
title(ws, "Indicadores — Como usar esta planilha",
      "Modelo para definir os indicadores, registrar os resultados, ler a situação e registrar as decisões.", "C")
r = 4


def section(text):
    global r
    band(ws, r, text, "C", sz=11)
    r += 1


def line(k, v, kbg=GRAY, vbg=None, kf=None, kh="left", height=None):
    global r
    put(ws, f"B{r}", k, f=kf or font(10, True), bg=kbg, h=kh)
    put(ws, f"C{r}", v, bg=vbg)
    ws.row_dimensions[r].height = height or (19.5 if len(v) <= 98 else (31.5 if len(v) <= 196 else 45.75))
    r += 1


section("Legenda: onde preencher")
line("Amarelo-claro", "Células de entrada. É aqui que você digita.", vbg=INPUT)
line("Cinza", "Células calculadas ou fixas (rótulos, situações, avisos). Não altere.", vbg=GRAY)
line("Cinza em itálico", "Linha de exemplo (marcada com “Ex.”). Mostra o formato esperado e não entra nas contagens.", vbg=GRAY)
line("Listas suspensas", "Colunas de sentido, frequência e status aceitam apenas as opções da lista. Metas e resultados aceitam apenas números.")
line("Comentários", "Células com triângulo vermelho no canto trazem uma dica de preenchimento. Passe o mouse sobre elas.")
r += 1
section("As três situações")
line(NA_META, "O último resultado atende à meta.", kbg=S1, vbg=S1_T, kf=font(10, True, c=WHITE))
line(ATENCAO, "O último resultado não atende à meta, mas está dentro do limite de atenção.", kbg=S2, vbg=S2_T, kf=font(10, True))
line(FORA, "O último resultado passou do limite de atenção.", kbg=S3, vbg=S3_T, kf=font(10, True, c=WHITE))
r += 1
section("Ordem recomendada de preenchimento")
for k, text in enumerate([
    "Aba Fichas: escreva cada indicador, com o objetivo, a fórmula, a unidade, o sentido, a meta e o limite de atenção.",
    "Aba Medições: escreva o nome dos períodos na linha amarela do cabeçalho.",
    "Aba Medições: registre o resultado de cada período, da esquerda para a direita, sem pular colunas.",
    "Aba Painel: leia a situação, a tendência e a ação sugerida de cada indicador.",
    "Aba Painel: informe o número do indicador que deve aparecer no gráfico.",
    "Aba Análise: registre a decisão de cada indicador sem atingir a meta, com responsável e prazo.",
    "Aba Checklist: valide o conjunto de indicadores, pelo menos uma vez por ano.",
], 1):
    line(f"Passo {k}", text)
r += 1
section("Abas da planilha")
for k, text in [
    ("Fichas", "Definição de até 12 indicadores, com conferência dos campos."),
    ("Medições", "Resultados de cada indicador, em até 12 períodos, e contagem dos períodos seguidos sem atingir a meta."),
    ("Painel", "Leitura de cada indicador, resumo e gráfico do indicador escolhido."),
    ("Análise", "Registro das decisões tomadas nas reuniões de indicadores."),
    ("Checklist", "Doze verificações de qualidade dos indicadores, com percentual de conclusão."),
    ("Exemplo 1 - Pizzaria", "Painel de quatro indicadores de uma organização pequena."),
    ("Exemplo 2 - Compras", "Painel de quatro indicadores de um processo, com o registro da reunião de análise."),
]:
    line(k, text)
r += 1
section("Regras e premissas do modelo")
for k, text in [
    ("Situação", "É lida no último resultado registrado. A meta e o limite de atenção vêm da aba Fichas."),
    ("Limite de atenção", "Separa o desvio pequeno do grande. Quando maior é melhor, o limite fica abaixo da meta. Quando menor é melhor, fica acima."),
    ("Três períodos seguidos", "Quando os três últimos resultados não atendem à meta, a ação sugerida é abrir análise de causa, qualquer que seja a situação. A regra é uma convenção deste modelo."),
    ("Tendência", "Compara a média dos três últimos períodos com a média dos três anteriores. Diferença de até 5% é considerada estável. Pede pelo menos seis medições."),
    ("Um ano por arquivo", "O modelo guarda 12 períodos. Para o período seguinte, salve uma cópia do arquivo e limpe a aba Medições."),
    ("Exemplos", "Os dados das abas de exemplo são ilustrativos, criados para este material."),
]:
    line(k, text)
setup(ws, INK, f"B1:C{r}", landscape=False, fit_height=True)

# ------------------------------------------------------------------ Fichas
ws = wb.create_sheet("Fichas")
widths(ws, {"A": 2, "B": 5, "C": 30, "D": 30, "E": 20, "F": 40, "G": 13, "H": 17, "I": 10, "J": 11, "K": 30, "L": 13, "M": 20, "N": 28, "O": 2})
title(ws, "Fichas dos indicadores", "Um indicador por linha. Duas pessoas, com a ficha na mão, precisam chegar ao mesmo número.", "N")
for rr, l1, l2, f2 in [(4, "Organização ou processo", "Responsável pelo painel", None), (5, "Período", "Atualizado em", DATE)]:
    label(ws, f"B{rr}", l1, merge=f"B{rr}:C{rr}")
    inp(ws, f"D{rr}", merge=f"D{rr}:F{rr}")
    label(ws, f"G{rr}", l2, merge=f"G{rr}:J{rr}")
    inp(ws, f"K{rr}", merge=f"K{rr}:N{rr}", fmt=f2, h="left")
    ws.row_dimensions[rr].height = 24
dv_date(ws, "K5")
for col, text in zip("BCDEFGHIJKLMN", ["#", "Indicador", "Objetivo", "Processo", "Fórmula", "Unidade", "Sentido", "Meta", "Limite de atenção",
                                       "Fonte dos dados", "Frequência", "Responsável", "Conferência"]):
    head(ws, f"{col}7", text)
ws.row_dimensions[7].height = 30
hint_row(ws, 8, [("B", "", None), ("C", "Nome curto", None), ("D", "O que se quer alcançar", None), ("E", "Onde é medido", None),
                 ("F", "O que se divide, e por quanto", None), ("G", "%, dias, itens", None), ("H", "Escolha na lista", None), ("I", "Número", None),
                 ("J", "Número", None), ("K", "De onde vem o dado", None), ("L", "Escolha na lista", None), ("M", "Cargo ou nome", None),
                 ("N", "Campos da linha", None)])
ex(ws, "B9", "Ex.", h="center")
for col, v, h_ in [("C", "Prazo de atendimento das requisições", "left"), ("D", "Fornecer os materiais no prazo", "left"),
                   ("E", "Adquirir materiais e serviços", "left"), ("F", "Soma dos dias úteis entre a requisição e o pedido ÷ número de requisições", "left"),
                   ("G", "dias úteis", "center"), ("H", MENOR, "center"), ("I", 5, "center"), ("J", 6, "center"),
                   ("K", "Datas da requisição e do pedido, no sistema", "left"), ("L", "Mensal", "center"), ("M", "Comprador sênior", "left"),
                   ("N", "Completa", "center")]:
    ex(ws, f"{col}9", v, h=h_)
ws.row_dimensions[9].height = 36
F1 = 10
F2 = F1 + NI - 1
for k in range(NI):
    rr = F1 + k
    num(ws, f"B{rr}", k + 1)
    for col in "CDEF":
        inp(ws, f"{col}{rr}")
    inp(ws, f"G{rr}", h="center")
    inp(ws, f"H{rr}", h="center")
    inp(ws, f"I{rr}", h="center", fmt=GERAL)
    inp(ws, f"J{rr}", h="center", fmt=GERAL)
    inp(ws, f"K{rr}")
    inp(ws, f"L{rr}", h="center")
    inp(ws, f"M{rr}")
    calc(ws, f"N{rr}", f'=IF(C{rr}="","",IF(D{rr}="","Falta o objetivo",IF(F{rr}="","Falta a fórmula",IF(OR(G{rr}="",H{rr}=""),"Falta a unidade ou o sentido",'
         f'IF(I{rr}="","Falta a meta",IF(J{rr}="","Falta o limite de atenção",IF(OR(AND(H{rr}="{MAIOR}",J{rr}>=I{rr}),AND(H{rr}="{MENOR}",J{rr}<=I{rr})),'
         f'"Limite incoerente com a meta",IF(OR(K{rr}="",L{rr}="",M{rr}=""),"Falta fonte, frequência ou responsável","Completa"))))))))', b=False, sz=9)
    ws.row_dimensions[rr].height = 39
dv_list(ws, f"H{F1}:H{F2}", [MAIOR, MENOR], "Maior é melhor: entregas no prazo. Menor é melhor: reclamações, atrasos.")
dv_list(ws, f"L{F1}:L{F2}", FREQ, "Com que frequência o resultado é calculado")
dv_number(ws, f"I{F1}:J{F2}")
cf_warn(ws, f"N{F1}:N{F2}", f"N{F1}", ok_values=("Completa",))
note(ws, "F7", 'Escreva o que entra na conta, e em que ordem.\nEx.: "Pedidos entregues em até 40 minutos ÷ pedidos entregues × 100".')
note(ws, "I7", "O valor que o indicador deve atingir. Digite só o número. A unidade fica na coluna Unidade.")
note(ws, "J7", "O valor que separa o desvio pequeno do grande. Quando maior é melhor, fica abaixo da meta. Quando menor é melhor, fica acima.")
s = F2 + 2
band(ws, s, "Resumo automático", "N")
for k, (text, formula) in enumerate([
    ("Indicadores cadastrados", f"=COUNTA(C{F1}:C{F2})"),
    ("Fichas completas", f'=COUNTIF(N{F1}:N{F2},"Completa")'),
    ("Aviso", f'=IF(E{s+1}=0,"Cadastre os indicadores",IF(E{s+2}<E{s+1},"Há ficha incompleta: veja a coluna Conferência",'
              f'IF(E{s+1}>8,"São muitos indicadores: confira se todos levam a decisões","OK")))'),
], 1):
    label(ws, f"B{s+k}", text, merge=f"B{s+k}:D{s+k}", h="right")
    calc(ws, f"E{s+k}", formula, merge=f"E{s+k}:F{s+k}", sz=9 if text == "Aviso" else 10, b=text != "Aviso")
    ws.row_dimensions[s + k].height = 21.75
cf_warn(ws, f"E{s+3}:F{s+3}", f"E{s+3}")
ws.freeze_panes = "D8"
setup(ws, BLUE, f"B1:N{s+3}", fit_height=True)

# ------------------------------------------------------------------ Medições
ws = wb.create_sheet("Medições")
PC = [L(5 + k) for k in range(NP)]  # colunas E a P
CA, CZ = PC[0], PC[-1]
widths(ws, dict({"A": 2, "B": 5, "C": 34, "D": 14, "Q": 11, "R": 2}, **{c: 10 for c in PC}))
title(ws, "Medições", "Registre os resultados da esquerda para a direita, sem pular colunas. Escreva o nome de cada período na linha amarela.", "Q")
head(ws, "B5", "#")
head(ws, "C5", "Indicador")
head(ws, "D5", "Unidade")
put(ws, f"{CA}5", "Períodos", f=font(10, True, c=WHITE), bg=INK, h="center", merge=f"{CA}5:{CZ}5")
head(ws, "Q5", "Medições")
ws.row_dimensions[5].height = 21.75
for col, text in (("B", ""), ("C", "Vem da aba Fichas"), ("D", "Vem da aba Fichas"), ("Q", "Contagem")):
    put(ws, f"{col}6", text, f=font(9, i=True, c=MUTED), bg=GRAY, h="center")
for c in PC:
    inp(ws, f"{c}6", h="center")
ws.row_dimensions[6].height = 24
note(ws, f"{CA}6", 'Nome do período, como texto.\nEx.: "Jan/26", "Semana 12", "1º trimestre".')
M1 = 7
M2 = M1 + NI - 1
for k in range(NI):
    rr, fr = M1 + k, F1 + k
    num(ws, f"B{rr}", k + 1)
    calc(ws, f"C{rr}", f'=IF(Fichas!C{fr}="","",Fichas!C{fr})', h="left", b=False)
    calc(ws, f"D{rr}", f'=IF(Fichas!G{fr}="","",Fichas!G{fr})', b=False, sz=9)
    for c in PC:
        inp(ws, f"{c}{rr}", h="center", fmt=GERAL)
    calc(ws, f"Q{rr}", f"=COUNT({CA}{rr}:{CZ}{rr})", b=False)
    ws.row_dimensions[rr].height = 27
dv_number(ws, f"{CA}{M1}:{CZ}{M2}")
# células sem atingir a meta ficam com fundo cinza
ws.conditional_formatting.add(f"{CA}{M1}:{CZ}{M2}", FormulaRule(
    formula=[f'AND({CA}{M1}<>"",Fichas!$I{F1}<>"",Fichas!$H{F1}<>"",NOT(IF(Fichas!$H{F1}="{MAIOR}",{CA}{M1}>=Fichas!$I{F1},{CA}{M1}<=Fichas!$I{F1})))'],
    fill=PatternFill("solid", bgColor="D9DEE2", fgColor="D9DEE2")))
H0 = M2 + 2
band(ws, H0, "Períodos seguidos sem atingir a meta", "Q")
put(ws, f"B{H0+1}", "Contagem calculada a partir dos resultados e da meta de cada indicador. Volta a zero quando o resultado atende à meta.",
    f=font(9, i=True, c=MUTED), box=False, merge=f"B{H0+1}:Q{H0+1}")
H1 = H0 + 2
H2 = H1 + NI - 1
for k in range(NI):
    rr, dr, fr = H1 + k, M1 + k, F1 + k
    num(ws, f"B{rr}", k + 1)
    calc(ws, f"C{rr}", f'=IF(Fichas!C{fr}="","",Fichas!C{fr})', h="left", b=False)
    put(ws, f"D{rr}", "", bg=GRAY)
    for j, c in enumerate(PC):
        ant = f"N({PC[j-1]}{rr})+1" if j else "1"
        calc(ws, f"{c}{rr}", f'=IF(OR({c}{dr}="",Fichas!$I{fr}="",Fichas!$H{fr}=""),"",IF({f_atende(f"{c}{dr}", f"Fichas!$H{fr}", f"Fichas!$I{fr}")},0,{ant}))',
             b=False, fmt="0")
    calc(ws, f"Q{rr}", f'=IF(Q{dr}=0,"",N(INDEX({CA}{rr}:{CZ}{rr},Q{dr})))', fmt="0")
    ws.row_dimensions[rr].height = 19.5
ws.conditional_formatting.add(f"Q{H1}:Q{H2}", FormulaRule(formula=[f"N(Q{H1})>={SEGUIDOS}"], fill=PatternFill("solid", bgColor=S3_T, fgColor=S3_T)))
s = H2 + 2
band(ws, s, "Resumo automático", "Q")
GAP = "+".join(f'IF(Q{M1+k}=0,0,COUNTBLANK(INDEX({CA}{M1+k}:{CZ}{M1+k},1):INDEX({CA}{M1+k}:{CZ}{M1+k},Q{M1+k})))' for k in range(NI))
for k, (text, formula) in enumerate([
    ("Indicadores com medição", f'=COUNTIF(Q{M1}:Q{M2},">0")'),
    ("Períodos com nome", f"=COUNTA({CA}6:{CZ}6)"),
    ("Aviso", f'=IF(D{s+1}=0,"Registre os resultados",IF({GAP}>0,"Há coluna pulada: registre os resultados em sequência",'
              f'IF(D{s+2}<MAX(Q{M1}:Q{M2}),"Há período sem nome",IF(SUMPRODUCT((Q{M1}:Q{M2}>0)*(Fichas!I{F1}:I{F2}=""))>0,"Há indicador medido sem meta","OK"))))'),
], 1):
    label(ws, f"B{s+k}", text, merge=f"B{s+k}:C{s+k}", h="right")
    calc(ws, f"D{s+k}", formula, merge=f"D{s+k}:I{s+k}", sz=9 if text == "Aviso" else 10, b=text != "Aviso")
    ws.row_dimensions[s + k].height = 21.75
cf_warn(ws, f"D{s+3}:I{s+3}", f"D{s+3}")
ws.freeze_panes = f"{CA}7"
setup(ws, AMBER, f"B1:Q{s+3}", fit_height=True)

# ------------------------------------------------------------------ Painel
ws = wb.create_sheet("Painel")
widths(ws, {"A": 2, "B": 5, "C": 34, "D": 10, "E": 10, "F": 10, "G": 11, "H": 15, "I": 10, "J": 10, "K": 12, "L": 12, "M": 14, "N": 24, "O": 12,
            "P": 2})
title(ws, "Painel de indicadores", "Leitura do último resultado de cada indicador. Calculado a partir das abas Fichas e Medições.", "O")
label(ws, "B4", "Organização ou processo", merge="B4:C4")
calc(ws, "D4", '=IF(Fichas!D4="","",Fichas!D4)', h="left", b=False, merge="D4:H4")
label(ws, "I4", "Atualizado em", merge="I4:K4")
calc(ws, "L4", '=IF(Fichas!K5="","",Fichas!K5)', h="left", b=False, fmt=DATE, merge="L4:O4")
ws.row_dimensions[4].height = 21.75
for col, text in zip("BCDEFGHIJKLMN", ["#", "Indicador", "Meta", "Limite de atenção", "Medições", "Último resultado", "Situação", "Média",
                                       "Seguidos sem atingir a meta", "Média dos 3 anteriores", "Média dos 3 últimos", "Tendência", "Ação sugerida"]):
    head(ws, f"{col}6", text)
head(ws, "O6", "Sentido")
ws.row_dimensions[6].height = 39
P1 = 7
P2 = P1 + NI - 1
for k in range(NI):
    rr, dr, hr, fr = P1 + k, M1 + k, H1 + k, F1 + k
    row = f"Medições!${CA}${dr}:${CZ}${dr}"
    num(ws, f"B{rr}", k + 1)
    calc(ws, f"C{rr}", f'=IF(Fichas!C{fr}="","",Fichas!C{fr})', h="left", b=False)
    calc(ws, f"D{rr}", f'=IF(Fichas!I{fr}="","",Fichas!I{fr})', b=False, fmt=GERAL)
    calc(ws, f"E{rr}", f'=IF(Fichas!J{fr}="","",Fichas!J{fr})', b=False, fmt=GERAL)
    calc(ws, f"F{rr}", f'=IF(C{rr}="","",COUNT({row}))', b=False)
    calc(ws, f"G{rr}", f'=IF(N(F{rr})=0,"",INDEX({row},F{rr}))', fmt=GERAL)
    calc(ws, f"H{rr}", f_situacao(f"G{rr}", f"O{rr}", f"D{rr}", f"E{rr}"), b=False)
    calc(ws, f"I{rr}", f'=IF(N(F{rr})=0,"",AVERAGE({row}))', b=False, fmt="0.0")
    calc(ws, f"J{rr}", f'=IF(N(F{rr})=0,"",N(Medições!Q{hr}))', fmt="0")
    calc(ws, f"K{rr}", f'=IF(N(F{rr})<6,"",AVERAGE(INDEX({row},F{rr}-5):INDEX({row},F{rr}-3)))', b=False, fmt="0.0")
    calc(ws, f"L{rr}", f'=IF(N(F{rr})<6,"",AVERAGE(INDEX({row},F{rr}-2):INDEX({row},F{rr})))', b=False, fmt="0.0")
    calc(ws, f"M{rr}", f'=IF(N(F{rr})=0,"",{f_tendencia(f"F{rr}", f"K{rr}", f"L{rr}", f"O{rr}")[1:]})', b=False)
    calc(ws, f"N{rr}", f_acao(f"H{rr}", f"J{rr}"), b=False, sz=9)
    calc(ws, f"O{rr}", f'=IF(Fichas!H{fr}="","",Fichas!H{fr})', b=False, sz=8)
    ws.row_dimensions[rr].height = 27
cf_equal(ws, f"H{P1}:H{P2}", SIT_CF)
cf_equal(ws, f"M{P1}:M{P2}", TEND_CF)
cf_equal(ws, f"N{P1}:N{P2}", ACAO_CF)
ws.conditional_formatting.add(f"J{P1}:J{P2}", FormulaRule(formula=[f"N(J{P1})>={SEGUIDOS}"], fill=PatternFill("solid", bgColor=S3_T, fgColor=S3_T)))
note(ws, "J6", f"Períodos seguidos, contados a partir do último, em que o resultado não atendeu à meta. Com {SEGUIDOS} ou mais, a ação sugerida é abrir análise de causa.")
note(ws, "M6", "Compara a média dos três últimos períodos com a média dos três anteriores. Diferença de até 5% é considerada estável.")
s = P2 + 2
band(ws, s, "Resumo automático", "H")
for k, (text, formula) in enumerate([
    ("Indicadores com resultado", f'=COUNTIF(F{P1}:F{P2},">0")'),
    (NA_META, f'=COUNTIF(H{P1}:H{P2},"{NA_META}")'),
    (ATENCAO, f'=COUNTIF(H{P1}:H{P2},"{ATENCAO}")'),
    (FORA, f'=COUNTIF(H{P1}:H{P2},"{FORA}")'),
    ("Com análise de causa sugerida", f'=COUNTIF(N{P1}:N{P2},"Abrir análise de causa")'),
    ("Aviso", f'=IF(D{s+1}=0,"Registre os resultados na aba Medições",IF(Medições!D{H2+5}<>"OK",Medições!D{H2+5},'
              f'IF(D{s+5}>0,"Há indicador que pede análise de causa",IF(D{s+4}>0,"Há indicador fora da meta","OK"))))'),
], 1):
    label(ws, f"B{s+k}", text, merge=f"B{s+k}:C{s+k}", h="right")
    calc(ws, f"D{s+k}", formula, merge=f"D{s+k}:H{s+k}", sz=9 if text == "Aviso" else 10, b=text != "Aviso")
    ws.row_dimensions[s + k].height = 21.75
cf_warn(ws, f"D{s+6}:H{s+6}", f"D{s+6}")
g = s + 8
band(ws, g, "Gráfico de um indicador", "O")
label(ws, f"B{g+1}", "Indicador do gráfico, de 1 a 12", merge=f"B{g+1}:C{g+1}")
inp(ws, f"D{g+1}", 1, h="center")
calc(ws, f"E{g+1}", f'=IF(D{g+1}="","Informe o número do indicador",IF(INDEX(C{P1}:C{P2},D{g+1})="","Indicador sem cadastro",INDEX(C{P1}:C{P2},D{g+1})))',
     h="left", merge=f"E{g+1}:O{g+1}")
dvn = DataValidation(type="whole", operator="between", formula1="1", formula2=str(NI), allow_blank=True)
dvn.errorTitle, dvn.error, dvn.showErrorMessage = "Número inválido", f"Digite um número inteiro de 1 a {NI}.", True
ws.add_data_validation(dvn)
dvn.add(f"D{g+1}")
ws.row_dimensions[g + 1].height = 24
GC = [L(4 + k) for k in range(NP)]  # colunas D a O
sel = f"N($D${g+1})"
for j, (rot, tipo) in enumerate([("Período", "per"), ("Resultado", "val"), ("Meta", "meta"), ("Limite de atenção", "lim")]):
    rr = g + 2 + j
    label(ws, f"B{rr}", rot, merge=f"B{rr}:C{rr}")
    for k, c in enumerate(GC):
        mc = PC[k]
        if tipo == "per":
            f_ = f'=IF(Medições!{mc}6="","",Medições!{mc}6&"")'
        elif tipo == "val":
            f_ = f'=IF({sel}=0,NA(),IF(INDEX(Medições!{mc}{M1}:{mc}{M2},{sel})="",NA(),INDEX(Medições!{mc}{M1}:{mc}{M2},{sel})))'
        else:
            col = "D" if tipo == "meta" else "E"
            f_ = f'=IF({sel}=0,NA(),IF(OR(INDEX({col}${P1}:{col}${P2},{sel})="",ISNA({c}{g+3})),NA(),INDEX({col}${P1}:{col}${P2},{sel})))'
        calc(ws, f"{c}{rr}", f_, b=False, sz=9, fmt=GERAL)
    ws.row_dimensions[rr].height = 19.5
# os valores ausentes aparecem como #N/D para o gráfico não desenhar zeros: o texto fica da cor do fundo
ws.conditional_formatting.add(f"D{g+3}:O{g+5}", FormulaRule(formula=[f"ISNA(D{g+3})"], font=Font(name="Arial", size=9, color=GRAY)))
line_chart(ws, f"B{g+7}", g + 2, g + 3, g + 4, g + 5, 4, 15, width=27, height=8.5)
GRAF = (g + 3, g + 5)
ws.freeze_panes = "D7"
setup(ws, TEAL, f"B1:O{g+24}", fit_height=True)

# ------------------------------------------------------------------ Análise
ws = wb.create_sheet("Análise")
widths(ws, {"A": 2, "B": 5, "C": 13, "D": 11, "E": 30, "F": 38, "G": 34, "H": 36, "I": 20, "J": 13, "K": 15, "L": 14, "M": 2})
title(ws, "Registro das reuniões de análise", "Uma linha por decisão. Todo indicador sem atingir a meta tem decisão registrada, com responsável e prazo.", "L")
for col, text in zip("BCDEFGHIJKL", ["#", "Data da reunião", "Indicador", "Nome do indicador", "O que os dados mostram", "Causa provável", "Decisão",
                                      "Responsável", "Prazo", "Status", "Situação"]):
    head(ws, f"{col}4", text)
ws.row_dimensions[4].height = 30
hint_row(ws, 5, [("B", "", None), ("C", "Data", None), ("D", "Número, de 1 a 12", None), ("E", "Calculado", None), ("F", "O fato, com números", None),
                 ("G", "Hipótese, a confirmar", None), ("H", "Verbo no infinitivo", None), ("I", "Cargo ou nome", None), ("J", "Data", None),
                 ("K", "Da decisão", None), ("L", "Calculada", None)])
ex(ws, "B6", "Ex.", h="center")
for col, v, h_, f_ in [("C", date(2026, 10, 8), "center", DATE), ("D", 1, "center", None), ("E", "Prazo de atendimento das requisições", "left", None),
                       ("F", "Cinco meses seguidos acima da meta, com pico de 6,2 dias em agosto.", "left", None),
                       ("G", "Requisições incompletas voltam ao requisitante e entram de novo na fila.", "left", None),
                       ("H", "Tratar junto com o RNC 2026-31.", "left", None), ("I", "Gerente de Suprimentos", "left", None),
                       ("J", date(2026, 10, 23), "center", DATE), ("K", "Em andamento", "center", None), ("L", "No prazo", "center", None)]:
    ex(ws, f"{col}6", v, h=h_, fmt=f_)
ws.row_dimensions[6].height = 36
A1, A2 = 7, 21
PB, PC_, PN = (f"Painel!${c}${P1}:${c}${P2}" for c in "BCN")
for k in range(15):
    rr = A1 + k
    num(ws, f"B{rr}", k + 1)
    inp(ws, f"C{rr}", h="center", fmt=DATE)
    inp(ws, f"D{rr}", h="center")
    calc(ws, f"E{rr}", f'=IF(D{rr}="","",IF(INDEX({PC_},D{rr})="","Indicador sem cadastro",INDEX({PC_},D{rr})))', h="left", b=False, sz=9)
    for col in "FGHI":
        inp(ws, f"{col}{rr}")
    inp(ws, f"J{rr}", h="center", fmt=DATE)
    inp(ws, f"K{rr}", h="center")
    calc(ws, f"L{rr}", f'=IF(H{rr}="","",IF(K{rr}="Concluída","Concluída",IF(J{rr}="","Sem prazo",IF(J{rr}<TODAY(),"Atrasada","No prazo"))))', b=False, sz=9)
    ws.row_dimensions[rr].height = 39
dvn = DataValidation(type="whole", operator="between", formula1="1", formula2=str(NI), allow_blank=True)
dvn.promptTitle, dvn.prompt, dvn.showInputMessage = "Indicador", "Número do indicador, como na aba Fichas.", True
dvn.errorTitle, dvn.error, dvn.showErrorMessage = "Número inválido", f"Digite um número inteiro de 1 a {NI}.", True
ws.add_data_validation(dvn)
dvn.add(f"D{A1}:D{A2}")
dv_date(ws, f"C{A1}:C{A2}")
dv_date(ws, f"J{A1}:J{A2}")
dv_list(ws, f"K{A1}:K{A2}", ST_ACAO, "Não iniciada, Em andamento ou Concluída")
cf_equal(ws, f"K{A1}:K{A2}", ST_CF)
cf_equal(ws, f"L{A1}:L{A2}", PRAZO_CF)
note(ws, "G4", "A causa registrada aqui é uma hipótese. Quando o indicador pede análise de causa, a confirmação é feita na reunião de análise, com as pessoas do processo.")
s = A2 + 2
band(ws, s, "Resumo automático", "L")
for k, (text, formula) in enumerate([
    ("Decisões registradas", f"=COUNTA(H{A1}:H{A2})"),
    ("Concluídas", f'=COUNTIF(L{A1}:L{A2},"Concluída")'),
    ("Atrasadas", f'=COUNTIF(L{A1}:L{A2},"Atrasada")'),
    ("Indicadores que pedem análise de causa", f'=COUNTIF({PN},"Abrir análise de causa")'),
    ("Desses, sem registro nesta aba", f'=SUMPRODUCT(({PN}="Abrir análise de causa")*(COUNTIF($D${A1}:$D${A2},{PB})=0))'),
    ("Aviso", f'=IF(F{s+5}>0,"Há indicador que pede análise de causa sem decisão registrada",IF(SUMPRODUCT((H{A1}:H{A2}<>"")*(((D{A1}:D{A2}="")+(I{A1}:I{A2}="")+(J{A1}:J{A2}=""))>0))>0,'
              f'"Há decisão sem indicador, responsável ou prazo",IF(F{s+3}>0,"Há decisão atrasada",IF(F{s+1}=0,"Sem decisões registradas","OK"))))'),
], 1):
    label(ws, f"B{s+k}", text, merge=f"B{s+k}:E{s+k}", h="right")
    calc(ws, f"F{s+k}", formula, sz=9 if text == "Aviso" else 10, b=text != "Aviso")
    ws.row_dimensions[s + k].height = 30 if text == "Aviso" else 21.75
cf_warn(ws, f"F{s+6}", f"F{s+6}", ok_values=("OK", "Sem decisões registradas"))
ws.freeze_panes = "F5"
setup(ws, REDC, f"B1:L{s+6}", fit_height=True)

# ------------------------------------------------------------------ Checklist
ws = wb.create_sheet("Checklist")
widths(ws, {"A": 2, "B": 5, "C": 70, "D": 14, "E": 44, "F": 2})
title(ws, "Checklist de validação dos indicadores", "Escolha o status de cada item na lista suspensa (coluna em amarelo-claro).", "E")
for col, text in zip("BCDE", ["#", "Item de verificação", "Status", "Observações"]):
    head(ws, f"{col}4", text)
ws.row_dimensions[4].height = 21.75
for k, text in enumerate(CHECK):
    rr = 5 + k
    num(ws, f"B{rr}", k + 1)
    put(ws, f"C{rr}", text, bg=GRAY)
    inp(ws, f"D{rr}", h="center")
    inp(ws, f"E{rr}")
    ws.row_dimensions[rr].height = 30
dv_list(ws, "D5:D16", YESNO, "Sim, Parcial ou Não")
cf_equal(ws, "D5:D16", YESNO_CF)
band(ws, 18, "Resumo automático", "E")
summary(ws, 19, "Itens atendidos (Sim)", '=COUNTIF(D5:D16,"Sim")', "C", "D")
summary(ws, 20, "Itens parciais", '=COUNTIF(D5:D16,"Parcial")', "C", "D")
summary(ws, 21, "Itens não atendidos", '=COUNTIF(D5:D16,"Não")', "C", "D")
summary(ws, 22, "Itens sem resposta", "=COUNTBLANK(D5:D16)", "C", "D")
summary(ws, 23, "Percentual atendido", "=D19/ROWS(D5:D16)", "C", "D", fmt="0%")
summary(ws, 24, "Situação", '=IF(D19=ROWS(D5:D16),"Validado",IF(D22>0,"Em andamento","Com pendências"))', "C", "D")
setup(ws, PURPLE, "B1:E24", landscape=False, fit_height=True)


# ------------------------------------------------------------------ Exemplos
def example(ws, data, grafico):
    MC = [L(4 + k) for k in range(12)]  # colunas D a O, uma por mês
    widths(ws, dict({"A": 2, "B": 6, "C": 34, "P": 12, "Q": 16, "R": 20, "S": 2}, **{c: 9.5 for c in MC}))
    title(ws, "Painel de indicadores", "Exemplo preenchido, para consulta. Use as abas Fichas, Medições, Painel e Análise para os seus indicadores.", "R")
    H = data["head"]
    inds = data["inds"]
    n = len(inds)
    rr = 4
    for rot, val, fmt in [("Organização", H["org"], None), ("Período", H["periodo"], None), ("Responsável", H["por"], None),
                          ("Reunião de análise", H["reuniao"], None), ("Situação em", H["data"], DATE)]:
        label(ws, f"B{rr}", rot, merge=f"B{rr}:C{rr}")
        inp(ws, f"D{rr}", val, merge=f"D{rr}:R{rr}", bg=WHITE, fmt=fmt)
        ws.row_dimensions[rr].height = 21.75
        rr += 1

    def cab(row, cols):
        for ref, text, merge in cols:
            put(ws, f"{ref}{row}", text, f=font(10, True), bg=GRAY, h="center", merge=merge and f"{ref}{row}:{merge}{row}")
        ws.row_dimensions[row].height = 24

    # ---- fichas
    rr += 1
    band(ws, rr, "Fichas dos indicadores", "R", color=BLUE)
    rr += 1
    cab(rr, [("B", "Nº", None), ("C", "Indicador", None), ("D", "Objetivo", "G"), ("H", "Fórmula", "L"), ("M", "Unidade", "N"), ("O", "Meta", None),
             ("P", "Limite", None), ("Q", "Sentido", None), ("R", "Responsável", None)])
    f1 = rr + 1
    for i in inds:
        rr += 1
        put(ws, f"B{rr}", i["id"], f=font(10, True), bg=GRAY, h="center")
        put(ws, f"C{rr}", i["nome"], f=font(10, True))
        put(ws, f"D{rr}", i["objetivo"], merge=f"D{rr}:G{rr}")
        put(ws, f"H{rr}", i["formula"], merge=f"H{rr}:L{rr}")
        put(ws, f"M{rr}", i["unidade"], h="center", merge=f"M{rr}:N{rr}")
        put(ws, f"O{rr}", i["meta"], h="center", fmt=GERAL)
        put(ws, f"P{rr}", i["limite"], h="center", fmt=GERAL)
        put(ws, f"Q{rr}", i["sentido"], h="center", f=font(9))
        put(ws, f"R{rr}", i["resp"])
        ws.row_dimensions[rr].height = max(30, alt(i["formula"], 47 * 1.1), alt(i["objetivo"], 38 * 1.1))
    # ---- medições
    rr += 2
    band(ws, rr, "Medições", "R", color=AMBER)
    rr += 1
    per = rr
    cab(rr, [("B", "Nº", None), ("C", "Indicador", None)] + [(c, m, None) for c, m in zip(MC, MESES)] + [("P", "Média", None), ("Q", "Seguidos sem a meta", None),
                                                                                                           ("R", "", None)])
    m1 = rr + 1
    for k, i in enumerate(inds):
        rr += 1
        fr = f1 + k
        put(ws, f"B{rr}", i["id"], f=font(10, True), bg=GRAY, h="center")
        put(ws, f"C{rr}", i["nome"])
        for c, v in zip(MC, i["valores"]):
            put(ws, f"{c}{rr}", v, h="center", fmt=GERAL)
        calc(ws, f"P{rr}", f"=AVERAGE(D{rr}:O{rr})", b=False, fmt="0.0")
        put(ws, f"R{rr}", "", bg=GRAY)
        ws.row_dimensions[rr].height = 24
    m2 = rr
    ws.conditional_formatting.add(f"D{m1}:O{m2}", FormulaRule(
        formula=[f'NOT(IF($Q{f1}="{MAIOR}",D{m1}>=$O{f1},D{m1}<=$O{f1}))'], fill=PatternFill("solid", bgColor="D9DEE2", fgColor="D9DEE2")))
    # ---- contagem dos períodos seguidos
    rr += 2
    band(ws, rr, "Períodos seguidos sem atingir a meta", "R", color=AMBER)
    h1 = rr + 1
    for k, i in enumerate(inds):
        rr += 1
        dr, fr = m1 + k, f1 + k
        put(ws, f"B{rr}", i["id"], f=font(10, True), bg=GRAY, h="center")
        put(ws, f"C{rr}", i["nome"], bg=GRAY)
        for j, c in enumerate(MC):
            ant = f"N({MC[j-1]}{rr})+1" if j else "1"
            calc(ws, f"{c}{rr}", f'=IF({f_atende(f"{c}{dr}", f"$Q{fr}", f"$O{fr}")},0,{ant})', b=False, fmt="0")
        for c in "PQR":
            put(ws, f"{c}{rr}", "", bg=GRAY)
        calc(ws, f"Q{dr}", f"=O{rr}", fmt="0")
        ws.row_dimensions[rr].height = 19.5
    ws.conditional_formatting.add(f"Q{m1}:Q{m2}", FormulaRule(formula=[f"Q{m1}>={SEGUIDOS}"], fill=PatternFill("solid", bgColor=S3_T, fgColor=S3_T)))
    # ---- painel
    rr += 2
    ws.row_breaks.append(Break(id=rr - 1))
    band(ws, rr, "Painel", "R", color=TEAL)
    rr += 1
    cab(rr, [("B", "Nº", None), ("C", "Indicador", None), ("D", "Último", "E"), ("F", "Situação", "G"), ("H", "Média dos 3 anteriores", "I"),
             ("J", "Média dos 3 últimos", "K"), ("L", "Tendência", "M"), ("N", "Ação sugerida", "P"), ("Q", "Decisão", "R")])
    ws.row_dimensions[rr].height = 30
    p1 = rr + 1
    for k, i in enumerate(inds):
        rr += 1
        dr, fr = m1 + k, f1 + k
        put(ws, f"B{rr}", i["id"], f=font(10, True), bg=GRAY, h="center")
        put(ws, f"C{rr}", i["nome"], bg=GRAY)
        calc(ws, f"D{rr}", f"=O{dr}", merge=f"D{rr}:E{rr}", fmt=GERAL)
        calc(ws, f"F{rr}", f_situacao(f"D{rr}", f"$Q{fr}", f"$O{fr}", f"$P{fr}"), b=False, merge=f"F{rr}:G{rr}")
        calc(ws, f"H{rr}", f"=AVERAGE(J{dr}:L{dr})", b=False, fmt="0.0", merge=f"H{rr}:I{rr}")
        calc(ws, f"J{rr}", f"=AVERAGE(M{dr}:O{dr})", b=False, fmt="0.0", merge=f"J{rr}:K{rr}")
        calc(ws, f"L{rr}", f_tendencia("12", f"H{rr}", f"J{rr}", f"$Q{fr}"), b=False, merge=f"L{rr}:M{rr}")
        calc(ws, f"N{rr}", f_acao(f"F{rr}", f"Q{dr}"), b=False, merge=f"N{rr}:P{rr}")
        put(ws, f"Q{rr}", i["decisao"], merge=f"Q{rr}:R{rr}")
        ws.row_dimensions[rr].height = max(33, alt(i["decisao"], 36 * 1.1))
    p2 = rr
    cf_equal(ws, f"F{p1}:G{p2}", SIT_CF)
    cf_equal(ws, f"L{p1}:M{p2}", TEND_CF)
    cf_equal(ws, f"N{p1}:P{p2}", ACAO_CF)
    rr += 2
    s0 = rr
    band(ws, rr, "Resumo automático", "G")
    for k, (text, formula) in enumerate([
        (NA_META, f'=COUNTIF(F{p1}:F{p2},"{NA_META}")'),
        (ATENCAO, f'=COUNTIF(F{p1}:F{p2},"{ATENCAO}")'),
        (FORA, f'=COUNTIF(F{p1}:F{p2},"{FORA}")'),
        ("Com análise de causa sugerida", f'=COUNTIF(N{p1}:N{p2},"Abrir análise de causa")'),
    ], 1):
        label(ws, f"B{rr+k}", text, merge=f"B{rr+k}:C{rr+k}", h="right")
        calc(ws, f"D{rr+k}", formula, merge=f"D{rr+k}:G{rr+k}")
        ws.row_dimensions[rr + k].height = 21.75
    # ---- gráfico de um indicador
    gi = [i["id"] for i in inds].index(grafico)
    ind = inds[gi]
    put(ws, f"I{s0}", f'Gráfico do indicador {ind["id"]} · {ind["nome"]}, em {ind["unidade"]}', f=font(10, True, c=WHITE), bg=INK, box=False,
        merge=f"I{s0}:R{s0}")
    ga = s0 + 18
    put(ws, f"B{ga}", "Dados do gráfico", f=font(9, i=True, c=MUTED), box=False, merge=f"B{ga}:C{ga}")
    for j, (rot, f_) in enumerate([("Meta", f"=$O${f1+gi}"), ("Limite de atenção", f"=$P${f1+gi}")]):
        label(ws, f"B{ga+1+j}", rot, merge=f"B{ga+1+j}:C{ga+1+j}")
        for c in MC:
            calc(ws, f"{c}{ga+1+j}", f_, b=False, sz=9, fmt=GERAL)
        ws.row_dimensions[ga + 1 + j].height = 18
    line_chart(ws, f"I{s0+1}", per, m1 + gi, ga + 1, ga + 2, 4, 15, width=19.5, height=7.4)
    rr = ga + 3
    if "analise" in data:
        rr += 1
        ws.row_breaks.append(Break(id=rr - 1))
        band(ws, rr, "Registro da reunião de análise", "R", color=REDC)
        rr += 1
        cab(rr, [("B", "Nº", None), ("C", "Indicador", None), ("D", "O que os dados mostram", "H"), ("I", "Causa provável", "L"), ("M", "Decisão", "P"),
                 ("Q", "Responsável", None), ("R", "Prazo", None)])
        nomes = {i["id"]: i["nome"] for i in inds}
        for d, ident, mostra, causa, dec, quem, prazo in data["analise"]:
            rr += 1
            put(ws, f"B{rr}", ident, f=font(10, True), bg=GRAY, h="center")
            put(ws, f"C{rr}", nomes[ident])
            put(ws, f"D{rr}", mostra, merge=f"D{rr}:H{rr}")
            put(ws, f"I{rr}", causa, merge=f"I{rr}:L{rr}")
            put(ws, f"M{rr}", dec, merge=f"M{rr}:P{rr}")
            put(ws, f"Q{rr}", quem)
            put(ws, f"R{rr}", prazo, h="center", fmt=DATE)
            ws.row_dimensions[rr].height = max(36, alt(mostra, 47 * 1.1), alt(causa, 38 * 1.1), alt(dec, 40 * 1.1))
    setup(ws, MUTED, f"B1:R{rr}")


example(wb.create_sheet("Exemplo 1 - Pizzaria"), EX1, "P3")
example(wb.create_sheet("Exemplo 2 - Compras"), EX2, "C1")

wb.active = 1
wb.save(OUT)
print("ok", OUT)
