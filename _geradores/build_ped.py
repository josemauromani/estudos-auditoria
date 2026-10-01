# -*- coding: utf-8 -*-
"""Gera Pedidos-modelo.xlsx no mesmo padrão visual das planilhas anteriores da série."""
import math
import os
import sys

_here = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, _here)
# funções de formatação compartilhadas com o gerador do PDCA (bloco inicial do arquivo)
_src = open(os.path.join(_here, "build_pdca.py"), encoding="utf-8").read()
exec(_src[: _src.index("STATUS = [")])

from openpyxl.chart import BarChart, Series  # noqa: E402
from openpyxl.chart.label import DataLabelList  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402
from ped_data import (ACEITO, ALTERADO, CHECK, DECISOES, EMANALISE, EX1, EX2, L_INC, L_OK, L_PEND, NA, NAO, NQ, PERGUNTAS, RECUSADO, RESP, SIM)  # noqa: E402

BLUE, TEAL, AMBER, PURPLE, REDC = "2B5C8A", "1E7B73", "A96A12", "7A4A9A", "B0413E"
S1_T, S2_T, S3_T = "DCE8F3", "F8EBCB", "F5DEDC"
YESNO = ["Sim", "Parcial", "Não"]
YESNO_CF = [("Sim", GREEN), ("Parcial", YELLOW), ("Não", RED)]
NO_, NP, NM = 12, 60, 30
OF, PD, MD = "Oferta", "Pedidos", "'Mudanças'"
O1, O2 = 6, 6 + NO_ - 1
P1, P2 = 7, 7 + NP - 1
M1, M2 = 7, 7 + NM - 1
QC = [get_column_letter(10 + k) for k in range(NQ)]      # J a P: as sete perguntas, na aba Pedidos


def note(ws, ref, text):
    c = Comment(text, "Modelo Pedidos")
    c.width, c.height = 300, 120
    ws[ref].comment = c


def cf_warn(ws, rng_, first, ok_values=("OK",)):
    cond = ",".join(f'{first}="{v}"' for v in ok_values)
    ws.conditional_formatting.add(rng_, FormulaRule(formula=[f"OR({cond})"], fill=PatternFill("solid", bgColor=GREEN, fgColor=GREEN)))
    ws.conditional_formatting.add(rng_, FormulaRule(formula=[f"AND(LEN({first})>0,NOT(OR({cond})))"],
                                  fill=PatternFill("solid", bgColor=YELLOW, fgColor=YELLOW)))


def cf_texto(ws, rng_, first, pares, resto=None):
    for t, cor in pares:
        ws.conditional_formatting.add(rng_, FormulaRule(formula=[f'{first}="{t}"'], stopIfTrue=True, fill=PatternFill("solid", bgColor=cor, fgColor=cor)))
    if resto:
        ws.conditional_formatting.add(rng_, FormulaRule(formula=[f"LEN({first})>0"], fill=PatternFill("solid", bgColor=resto, fgColor=resto)))


def hint_row(ws, row, cells, height=30):
    for ref, text in cells:
        put(ws, f"{ref}{row}", text, f=font(9, i=True, c=MUTED), bg=GRAY, h="center")
    ws.row_dimensions[row].height = height


def alt(pares, minimo=21.75, linha=12):
    n = max(max(1, math.ceil(len(str(t)) / max(c, 1))) for t, c in pares)
    return max(minimo, n * linha + 7)


# ------------------------------------------------------------------ fórmulas compartilhadas pelas abas de entrada e pelos exemplos
def f_leitura(vazio, qs):
    return f'IF({vazio},"",IF(COUNTA({qs})<{NQ},"{L_INC}",IF(COUNTIF({qs},"{NAO}")>0,"{L_PEND}","{L_OK}")))'


def f_conf(vazio, qs, dec, neg, quem, ana, con):
    ac = f'OR({dec}="{ACEITO}",{dec}="{ALTERADO}")'
    return (f'IF({vazio},"",IF({dec}="","Falta a decisão",IF({quem}="","Falta quem analisou",IF({dec}="{EMANALISE}","OK",'
            f'IF(COUNTA({qs})<{NQ},"Falta responder à análise",IF({ana}="","Falta a data da análise",IF(AND({ac},{con}=""),"Falta a data da confirmação",'
            f'IF(AND({ac},{ana}>{con}),"Analisado depois de aceitar",IF(AND({dec}="{ACEITO}",COUNTIF({qs},"{NAO}")>0),"Aceito com pendência na análise",'
            f'IF(AND({dec}="{ALTERADO}",{neg}=""),"Falta o que foi negociado",IF(AND({dec}="{RECUSADO}",{neg}=""),"Falta o motivo da recusa","OK")))))))))))')


def f_mud(oque, outros, ped, analise, docs, inform, peds=None):
    existe = f'IF(COUNTIF({peds},{ped})=0,"Pedido não está na aba Pedidos",' if peds else 'IF(FALSE,"",'
    return (f'IF({oque}="",IF({outros}>0,"Falta o que mudou",""),IF({ped}="","Falta o pedido",{existe}IF({analise}="","Falta a análise",'
            f'IF({docs}<>"{SIM}","Documentos não atualizados",IF({inform}="","Falta quem foi informado","OK"))))))')


def f_oferta(item, outros, espec, onde, cap):
    return (f'IF({item}="",IF({outros}>0,"Falta o item",""),IF({espec}="","Falta a especificação",IF({onde}="","Falta onde o cliente é informado",'
            f'IF({cap}<>"{SIM}","Capacidade não confirmada","OK"))))')


DEC_CF = [(ACEITO, GREEN), (ALTERADO, S2_T), (RECUSADO, S3_T), (EMANALISE, GRAY)]
LEIT_CF = [(L_OK, GREEN), (L_PEND, S2_T), (L_INC, S3_T)]
RESP_CF = [(SIM, GREEN), (NAO, S3_T), (NA, GRAY)]
GRAVE = ["Aceito com pendência na análise", "Analisado depois de aceitar"]
PLURAL = {ACEITO: "aceitos", ALTERADO: "aceitos com alteração", RECUSADO: "recusados", EMANALISE: "em análise"}


def chart_style(ch):
    ch.style = 2
    ch.title = None
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ch.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill="D3DBE0", w=9525))
    ch.x_axis.graphicalProperties = GraphicalProperties(ln=LineProperties(solidFill="D3DBE0", w=9525))
    ch.y_axis.graphicalProperties = GraphicalProperties(ln=LineProperties(noFill=True))


def barras(ws, anchor, r1, r2, c_cat, c_val, width=20, height=8.5):
    """Barras horizontais com as respostas “não” de cada pergunta."""
    ch = BarChart()
    ch.type = "bar"
    ch.gapWidth = 60
    ch.height, ch.width = height, width
    chart_style(ch)
    ch.legend = None
    s = Series(Reference(ws, min_col=c_val, min_row=r1, max_row=r2), title="Respostas “não”")
    s.graphicalProperties.solidFill = REDC
    s.graphicalProperties.line.noFill = True
    s.dLbls = DataLabelList()
    s.dLbls.showVal = True
    s.dLbls.showSerName = s.dLbls.showCatName = s.dLbls.showLegendKey = False
    ch.append(s)
    ch.set_categories(Reference(ws, min_col=c_cat, min_row=r1, max_row=r2))
    ch.y_axis.scaling.min = 0
    ch.y_axis.majorUnit = 1
    ch.x_axis.scaling.orientation = "maxMin"
    ws.add_chart(ch, anchor)


# ================================================================== pasta de trabalho
wb = Workbook()
wb.properties.title = "Requisitos do cliente e análise de pedidos — Modelo"
wb.properties.creator = "Modelo Pedidos"
wb.properties.language = "pt-BR"

# ------------------------------------------------------------------ Instruções
ws = wb.active
ws.title = "Instruções"
widths(ws, {"A": 2, "B": 26, "C": 90, "D": 2})
title(ws, "Requisitos do cliente e análise de pedidos — Como usar esta planilha",
      "Modelo para definir a oferta, analisar cada pedido antes de aceitar e controlar as mudanças nos pedidos aceitos.", "C")
r = 4


def section(text):
    global r
    band(ws, r, text, "C", sz=11)
    r += 1


def line(k, v, kbg=GRAY, vbg=None, kf=None, kh="left", height=None):
    global r
    put(ws, f"B{r}", k, f=kf or font(10, True), bg=kbg, h=kh)
    put(ws, f"C{r}", v, bg=vbg)
    ws.row_dimensions[r].height = height or (19.5 if len(v) <= 96 else (31.5 if len(v) <= 192 else 45.75))
    r += 1


section("Legenda: onde preencher")
line("Amarelo-claro", "Células de entrada. É aqui que você digita.", vbg=INPUT)
line("Cinza", "Células calculadas ou fixas (rótulos, leituras, conferências, avisos). Não altere.", vbg=GRAY)
line("Cinza em itálico", "Linha de exemplo (marcada com “Ex.”). Mostra o formato esperado e não entra nas contagens.", vbg=GRAY)
line("Listas suspensas", "Respostas às perguntas, decisão, capacidade confirmada, documentos atualizados e checklist aceitam apenas as opções da lista.")
line("Comentários", "Células com triângulo vermelho no canto trazem uma dica de preenchimento. Passe o mouse sobre elas.")
r += 1
section("As sete perguntas da análise")
for k, (t, q) in enumerate(PERGUNTAS, 1):
    line(f"{k} · {t}", q)
r += 1
section("As regras de cálculo")
for k, text in [
    (L_OK, f"As {NQ} perguntas foram respondidas, todas com “{SIM}” ou “{NA}”."),
    (L_PEND, f"Alguma resposta é “{NAO}”. O pedido pode ser aceito com alteração negociada, recusado ou ficar em análise."),
    (L_INC, "Alguma pergunta ficou sem resposta."),
    ("Aceito com pendência", f"A decisão é “{ACEITO}”, e alguma resposta é “{NAO}”. É o aviso mais grave da planilha."),
    ("Analisado depois de aceitar", "A data da análise é posterior à data da confirmação ao cliente."),
    ("Mudança a rever", "Mudança sem análise, sem os documentos atualizados ou sem quem foi informado."),
]:
    line(k, text)
r += 1
section("Ordem recomendada de preenchimento")
for k, text in enumerate([
    "Aba Oferta: escreva cada produto ou serviço oferecido, a especificação, os requisitos legais, o que não se garante, onde o cliente é informado e se a capacidade foi confirmada.",
    "Aba Pedidos: registre cada pedido, com o que o cliente pede e os requisitos não declarados e legais identificados.",
    f"Aba Pedidos: responda às {NQ} perguntas, registre a decisão, o que foi negociado ou o motivo, quem analisou e as datas da análise e da confirmação.",
    "Aba Mudanças: registre cada alteração nos pedidos aceitos, com a análise, os documentos e quem foi informado.",
    "Aba Painel: leia os pedidos a rever, as respostas “não” por pergunta e as mudanças a rever.",
    "Aba Checklist: valide a rotina.",
], 1):
    line(f"Passo {k}", text)
r += 1
section("Abas da planilha")
for k, text in [
    (OF, f"Até {NO_} itens oferecidos. Calcula a conferência de cada item."),
    (PD, f"Até {NP} pedidos. Calcula a leitura da análise, as pendências e a conferência da decisão."),
    ("Mudanças", f"Até {NM} mudanças nos pedidos aceitos. Calcula a conferência de cada uma."),
    ("Painel", "Os indicadores, as respostas por pergunta, o gráfico e o aviso."),
    ("Checklist", "Doze verificações da rotina, com percentual de conclusão."),
    ("Exemplo 1 - Pizzaria", "As encomendas e os pedidos especiais de 22 a 28/03/2027."),
    ("Exemplo 2 - Indústria", "Os pedidos e as propostas de 01 a 15/06/2027."),
]:
    line(k, text)
r += 1
section("Regras e premissas do modelo")
for k, text in [
    ("Análise simplificada", f"Para o pedido igual ao anterior, ou pelo catálogo, as perguntas podem ser respondidas na tela de confirmação. Use “{NA}” quando a pergunta não couber."),
    ("Datas", "A data da análise deve ser igual ou anterior à da confirmação. Pedido recusado ou em análise não tem data de confirmação."),
    ("Mudanças", "O número do pedido na aba Mudanças deve ser igual ao da aba Pedidos."),
    ("Exemplos", "Os dados das abas de exemplo são ilustrativos, criados para este material."),
]:
    line(k, text)
setup(ws, INK, f"B1:C{r}", landscape=False, fit_height=True)

# ------------------------------------------------------------------ Oferta
ws = wb.create_sheet(OF)
widths(ws, {"A": 2, "B": 5, "C": 24, "D": 40, "E": 30, "F": 30, "G": 24, "H": 14, "I": 28, "J": 2})
title(ws, "A oferta", "O que a organização se dispõe a fornecer. É a base da análise: o pedido que cabe na oferta é mais simples de analisar.", "I")
for col, text in zip("BCDEFGHI", ["#", "Item", "O que se oferece", "Requisitos legais", "O que não se garante", "Onde o cliente é informado", "Capacidade confirmada?", "Conferência"]):
    head(ws, f"{col}4", text)
ws.row_dimensions[4].height = 33
hint_row(ws, 5, [("B", ""), ("C", "Produto ou serviço"), ("D", "Especificação, prazos, limites"), ("E", "Leis e regulamentos"), ("F", "Escrito, para o cliente saber"),
                 ("G", "Catálogo, site, proposta"), ("H", "Lista"), ("I", "Calculada")])
for k in range(NO_):
    rr = O1 + k
    num(ws, f"B{rr}", k + 1)
    for col in "CDEFG":
        inp(ws, f"{col}{rr}")
    inp(ws, f"H{rr}", h="center")
    calc(ws, f"I{rr}", "=" + f_oferta(f"C{rr}", f"COUNTA(D{rr}:H{rr})", f"D{rr}", f"G{rr}", f"H{rr}"), b=False, sz=9)
    ws.row_dimensions[rr].height = 33
dv_list(ws, f"H{O1}:H{O2}", [SIM, NAO], "A organização confirmou que consegue cumprir o que oferece, inclusive a capacidade?")
cf_warn(ws, f"I{O1}:I{O2}", f"I{O1}")
note(ws, "F4", "O que a organização não faz, ou não garante, dito ao cliente antes do pedido: “a cozinha é compartilhada”, “tolerância menor que ±2 µm”.")
s = O2 + 2
band(ws, s, "Resumo automático", "I")
label(ws, f"B{s+1}", "Itens na oferta", merge=f"B{s+1}:C{s+1}", h="right")
calc(ws, f"D{s+1}", f"=COUNTA(C{O1}:C{O2})", merge=f"D{s+1}:I{s+1}", h="left")
label(ws, f"B{s+2}", "Aviso", merge=f"B{s+2}:C{s+2}", h="right")
calc(ws, f"D{s+2}", f'=IF(D{s+1}=0,"Escreva a oferta",IF(COUNTIF(I{O1}:I{O2},"OK")<D{s+1},"Há item a rever: veja a coluna Conferência","OK"))', merge=f"D{s+2}:I{s+2}", sz=9, b=False, h="left")
cf_warn(ws, f"D{s+2}:I{s+2}", f"D{s+2}")
for k in (1, 2):
    ws.row_dimensions[s + k].height = 21.75
OAV = s + 2
setup(ws, BLUE, f"B1:I{OAV}", fit_height=True)

# ------------------------------------------------------------------ Pedidos
ws = wb.create_sheet(PD)
widths(ws, dict({"A": 2, "B": 5, "C": 11, "D": 11, "E": 20, "F": 12, "G": 32, "H": 30, "I": 11, "Q": 20, "R": 11, "S": 18, "T": 34, "U": 20, "V": 11, "W": 12, "X": 30, "Y": 2},
                **{c: 9 for c in QC}))
title(ws, "Análise crítica dos pedidos", "Uma linha por pedido. Responda às sete perguntas antes de confirmar ao cliente.", "X")
for col, text in zip("BCDEFGHI", ["#", "Nº do pedido", "Data", "Cliente", "Canal", "O que pede", "Requisitos não declarados e legais identificados", "Prazo pedido"]):
    head(ws, f"{col}4", text)
for k, c in enumerate(QC):
    head(ws, f"{c}4", f"{k + 1} · {PERGUNTAS[k][0]}")
for col, text in zip("QRSTUVWX", ["Leitura", "Pendências", "Decisão", "O que foi negociado, ou o motivo", "Quem analisou", "Data da análise", "Data da confirmação", "Conferência"]):
    head(ws, f"{col}4", text)
ws.row_dimensions[4].height = 45
hint_row(ws, 5, [("B", ""), ("C", "Código"), ("D", "Recebido em"), ("E", ""), ("F", "Telefone, e-mail"), ("G", "Os requisitos declarados"), ("H", "O uso, a lei, as diferenças"),
                 ("I", "Data")] + [(c, "Lista") for c in QC] + [("Q", "Calculada"), ("R", "Calculadas"), ("S", "Lista"), ("T", "Com o aceite do cliente"), ("U", "Uma função"),
                                                                ("V", "Data"), ("W", "Data"), ("X", "Calculada")])
ex(ws, "B6", "Ex.", h="center")
exrow = [("C", "E-32", "center", None), ("D", date(2027, 3, 25), "center", DATE), ("E", "Comissão de formatura", "left", None), ("F", "Telefone", "center", None),
         ("G", "50 pizzas para sábado 27/03 às 21h.", "left", None), ("H", "Cerveja pedida para turma com menores.", "left", None), ("I", date(2027, 3, 27), "center", DATE)]
exrow += [(c, v, "center", None) for c, v in zip(QC, [SIM, NAO, SIM, SIM, NAO, SIM, SIM])]
exrow += [("Q", L_PEND, "center", None), ("R", 2, "center", None), ("S", ALTERADO, "center", None), ("T", "Entrega às 18h30, sem bebida alcoólica.", "left", None),
          ("U", "Gerente da loja", "left", None), ("V", date(2027, 3, 25), "center", DATE), ("W", date(2027, 3, 25), "center", DATE), ("X", "OK", "center", None)]
for col, v, h_, fmt in exrow:
    ex(ws, f"{col}6", v, h=h_, fmt=fmt)
ws.row_dimensions[6].height = 30
for k in range(NP):
    rr = P1 + k
    num(ws, f"B{rr}", k + 1)
    for col in "CF":
        inp(ws, f"{col}{rr}", h="center")
    for col in "EGHTU":
        inp(ws, f"{col}{rr}")
    for col in "DIVW":
        inp(ws, f"{col}{rr}", h="center", fmt=DATE)
    for c in QC:
        inp(ws, f"{c}{rr}", h="center")
    inp(ws, f"S{rr}", h="center")
    qs = f"{QC[0]}{rr}:{QC[-1]}{rr}"
    vazio = f'AND(C{rr}="",E{rr}="",G{rr}="")'
    calc(ws, f"Q{rr}", "=" + f_leitura(vazio, qs), sz=9)
    calc(ws, f"R{rr}", f'=IF({vazio},"",COUNTIF({qs},"{NAO}"))', b=False)
    calc(ws, f"X{rr}", "=" + f_conf(vazio, qs, f"S{rr}", f"T{rr}", f"U{rr}", f"V{rr}", f"W{rr}"), b=False, sz=9)
    ws.row_dimensions[rr].height = 30
for col in "DIVW":
    dv_date(ws, f"{col}{P1}:{col}{P2}")
dv_list(ws, f"{QC[0]}{P1}:{QC[-1]}{P2}", RESP, "Sim, Não ou Não se aplica")
dv_list(ws, f"S{P1}:S{P2}", DECISOES, "Aceito, Aceito com alteração, Recusado ou Em análise")
cf_texto(ws, f"{QC[0]}{P1}:{QC[-1]}{P2}", f"{QC[0]}{P1}", RESP_CF)
cf_texto(ws, f"Q{P1}:Q{P2}", f"Q{P1}", LEIT_CF)
cf_texto(ws, f"S{P1}:S{P2}", f"S{P1}", DEC_CF)
cf_texto(ws, f"X{P1}:X{P2}", f"X{P1}", [("OK", GREEN)] + [(g, RED) for g in GRAVE], resto=YELLOW)
for k, c in enumerate(QC):
    note(ws, f"{c}4", PERGUNTAS[k][1])
note(ws, "H4", "O que o cliente não disse e o uso exige, os requisitos legais e as diferenças em relação à proposta. É aqui que a análise mostra o que encontrou.")
note(ws, "W4", "A data em que o aceite foi comunicado ao cliente. Deve ser igual ou posterior à data da análise.")
s = P2 + 2
band(ws, s, "Resumo automático", "X")
PBR = [("Pedidos registrados", f'=COUNTA(C{P1}:C{P2})'),
       ("Por decisão", f'=IF(E{s+1}=0,"",' + '&", "&'.join(f'COUNTIF(S{P1}:S{P2},"{d}")&" {PLURAL[d]}"' for d in DECISOES) + ')'),
       ("Pedidos a rever", f'=SUMPRODUCT((X{P1}:X{P2}<>"")*(X{P1}:X{P2}<>"OK"))'),
       ("Aviso", f'=IF(E{s+1}=0,"Registre os pedidos",IF(COUNTIF(X{P1}:X{P2},"{GRAVE[0]}")>0,"Há pedido aceito com pendência na análise",'
                 f'IF(COUNTIF(X{P1}:X{P2},"{GRAVE[1]}")>0,"Há pedido analisado depois de aceito",IF(E{s+3}>0,"Há pedido a rever: veja a coluna Conferência","OK"))))')]
for k, (text, formula) in enumerate(PBR, 1):
    label(ws, f"B{s+k}", text, merge=f"B{s+k}:D{s+k}", h="right")
    calc(ws, f"E{s+k}", formula, merge=f"E{s+k}:H{s+k}", sz=9 if text in ("Aviso", "Por decisão") else 10, b=text not in ("Aviso", "Por decisão"), h="left")
    ws.row_dimensions[s + k].height = 21.75
PAV = s + 4
cf_warn(ws, f"E{PAV}:H{PAV}", f"E{PAV}")
ws.freeze_panes = f"F{P1}"
setup(ws, TEAL, f"B1:X{PAV}")

# ------------------------------------------------------------------ Mudanças
ws = wb.create_sheet("Mudanças")
widths(ws, {"A": 2, "B": 5, "C": 12, "D": 12, "E": 36, "F": 24, "G": 36, "H": 14, "I": 26, "J": 30, "K": 2})
title(ws, "Mudanças nos pedidos aceitos", "Uma linha por alteração depois do aceite, venha do cliente ou da organização.", "J")
for col, text in zip("BCDEFGHIJ", ["#", "Data", "Pedido", "O que mudou", "Quem pediu", "Análise da mudança", "Documentos atualizados?", "Quem foi informado", "Conferência"]):
    head(ws, f"{col}4", text)
ws.row_dimensions[4].height = 33
hint_row(ws, 5, [("B", ""), ("C", "Data"), ("D", "Nº da aba Pedidos"), ("E", "De quê, para quê"), ("F", ""), ("G", "Capacidade, prazo, especificação"), ("H", "Lista"),
                 ("I", "Quem produz, entrega, compra"), ("J", "Calculada")])
ex(ws, "B6", "Ex.", h="center")
for col, v, h_, fmt in [("C", date(2027, 3, 24), "center", DATE), ("D", "E-31", "center", None), ("E", "De 30 para 35 pizzas.", "left", None), ("F", "Cliente, por telefone", "left", None),
                        ("G", "Capacidade conferida: a sexta às 19h comporta.", "left", None), ("H", SIM, "center", None), ("I", "Pizzaiolo líder", "left", None), ("J", "OK", "center", None)]:
    ex(ws, f"{col}6", v, h=h_, fmt=fmt)
ws.row_dimensions[6].height = 30
for k in range(NM):
    rr = M1 + k
    num(ws, f"B{rr}", k + 1)
    inp(ws, f"C{rr}", h="center", fmt=DATE)
    inp(ws, f"D{rr}", h="center")
    for col in "EFGI":
        inp(ws, f"{col}{rr}")
    inp(ws, f"H{rr}", h="center")
    calc(ws, f"J{rr}", "=" + f_mud(f"E{rr}", f"COUNTA(C{rr}:D{rr},F{rr}:I{rr})", f"D{rr}", f"G{rr}", f"H{rr}", f"I{rr}", peds=f"{PD}!$C${P1}:$C${P2}"), b=False, sz=9)
    ws.row_dimensions[rr].height = 30
dv_date(ws, f"C{M1}:C{M2}")
dv_list(ws, f"H{M1}:H{M2}", [SIM, NAO], "O pedido, a ordem de produção e os demais documentos foram atualizados?")
cf_warn(ws, f"J{M1}:J{M2}", f"J{M1}")
note(ws, "I4", "Quem precisa saber da mudança para agir: produção, expedição, compras, qualidade.")
s = M2 + 2
band(ws, s, "Resumo automático", "J")
for k, (text, formula) in enumerate([
    ("Mudanças registradas", f"=COUNTA(E{M1}:E{M2})"),
    ("Mudanças a rever", f'=SUMPRODUCT((J{M1}:J{M2}<>"")*(J{M1}:J{M2}<>"OK"))'),
    ("Aviso", f'=IF(E{s+1}=0,"Nenhuma mudança registrada: confirme se nenhum pedido mudou",IF(E{s+2}>0,"Há mudança a rever: veja a coluna Conferência","OK"))'),
], 1):
    label(ws, f"B{s+k}", text, merge=f"B{s+k}:D{s+k}", h="right")
    calc(ws, f"E{s+k}", formula, merge=f"E{s+k}:J{s+k}", sz=9 if text == "Aviso" else 10, b=text != "Aviso", h="left")
    ws.row_dimensions[s + k].height = 21.75
MAV = s + 3
cf_warn(ws, f"E{MAV}:J{MAV}", f"E{MAV}", ok_values=("OK", "Nenhuma mudança registrada: confirme se nenhum pedido mudou"))
ws.freeze_panes = "E7"
setup(ws, AMBER, f"B1:J{MAV}", fit_height=True)

# ------------------------------------------------------------------ Painel
ws = wb.create_sheet("Painel")
widths(ws, {"A": 2, "B": 34, "C": 14, "D": 12, "E": 12, "F": 14, "G": 12, "H": 30, "I": 2})
title(ws, "Painel dos pedidos", "Nada a preencher: tudo vem das outras abas.", "H")
for col, text, m in [("B", "Indicador", None), ("C", "Resultado", None), ("D", "Como se lê", "D4:H4")]:
    put(ws, f"{col}4", text, f=font(10, True, c=WHITE), bg=INK, h="center", merge=m)
ws.row_dimensions[4].height = 21.75
Pc, Ps, Px, Pr = (f"{PD}!{c}{P1}:{c}{P2}" for c in "CSXR")
IND = [
    ("Itens na oferta", f"=COUNTA({OF}!C{O1}:C{O2})", "Da aba Oferta."),
    ("Itens da oferta a rever", f'=SUMPRODUCT(({OF}!I{O1}:I{O2}<>"")*({OF}!I{O1}:I{O2}<>"OK"))', "Sem especificação, sem onde informar, sem a capacidade confirmada."),
    ("Pedidos registrados", f"=COUNTA({Pc})", "Da aba Pedidos."),
    (ACEITO, f'=COUNTIF({Ps},"{ACEITO}")', "Todas as respostas “sim” ou “não se aplica”."),
    (ALTERADO, f'=COUNTIF({Ps},"{ALTERADO}")', "Um “não” resolvido com o cliente."),
    (RECUSADO, f'=COUNTIF({Ps},"{RECUSADO}")', "Um “não” sem solução aceitável."),
    (EMANALISE, f'=COUNTIF({Ps},"{EMANALISE}")', "Falta informação para decidir."),
    ("Pedidos em que a análise encontrou problema", f'=COUNTIF({Pr},">0")', "Ao menos uma resposta “não”."),
    ("Pedidos a rever", f'=SUMPRODUCT(({Px}<>"")*({Px}<>"OK"))', "Veja a coluna Conferência da aba Pedidos."),
    ("Aceitos com pendência na análise", f'=COUNTIF({Px},"{GRAVE[0]}")', "O mais grave: o problema vai para a produção."),
    ("Analisados depois de aceitar", f'=COUNTIF({Px},"{GRAVE[1]}")', "O compromisso veio antes da análise."),
    ("Mudanças registradas", f"=COUNTA({MD}!E{M1}:E{M2})", "Da aba Mudanças."),
    ("Mudanças a rever", f'=SUMPRODUCT(({MD}!J{M1}:J{M2}<>"")*({MD}!J{M1}:J{M2}<>"OK"))', "Sem análise, sem documentos ou sem quem foi informado."),
]
for k, (nome, formula, leit) in enumerate(IND):
    rr = 5 + k
    label(ws, f"B{rr}", nome)
    calc(ws, f"C{rr}", formula)
    put(ws, f"D{rr}", leit, f=font(9, c=MUTED), bg=GRAY, merge=f"D{rr}:H{rr}")
    ws.row_dimensions[rr].height = 21.75
IR = {n: 5 + k for k, (n, _, _) in enumerate(IND)}
T0 = 5 + len(IND) + 1
band(ws, T0, "Respostas por pergunta", "H")
for col, text in zip("BCDEFGH", ["Pergunta", SIM, NAO, NA, "Parte com “não”", "Sem resposta", "Leitura"]):
    head(ws, f"{col}{T0+1}", text)
ws.row_dimensions[T0 + 1].height = 30
Q1 = T0 + 2
for k, (t, _) in enumerate(PERGUNTAS):
    rr = Q1 + k
    rng = f"{PD}!{QC[k]}{P1}:{QC[k]}{P2}"
    label(ws, f"B{rr}", f"{k + 1} · {t}")
    calc(ws, f"C{rr}", f'=COUNTIF({rng},"{SIM}")', b=False)
    calc(ws, f"D{rr}", f'=COUNTIF({rng},"{NAO}")')
    calc(ws, f"E{rr}", f'=COUNTIF({rng},"{NA}")', b=False)
    calc(ws, f"F{rr}", f'=IF(C{rr}+D{rr}=0,"",D{rr}/(C{rr}+D{rr}))', fmt="0%", b=False)
    calc(ws, f"G{rr}", f'=MAX(0,COUNTA({Pc})-C{rr}-D{rr}-E{rr})', b=False)
    calc(ws, f"H{rr}", f'=IF(D{rr}=0,"",IF(D{rr}=MAX(D${Q1}:D${Q1 + NQ - 1}),"A que mais encontra problema: rever a oferta",""))', b=False, sz=9, h="left")
    ws.row_dimensions[rr].height = 21.75
s = Q1 + NQ + 1
band(ws, s, "Resumo automático", "H")
label(ws, f"B{s+1}", "Aviso")
calc(ws, f"C{s+1}", f'=IF(C{IR["Pedidos registrados"]}=0,"Registre os pedidos na aba Pedidos",IF(C{IR["Aceitos com pendência na análise"]}>0,"Há pedido aceito com pendência na análise",'
     f'IF(C{IR["Analisados depois de aceitar"]}>0,"Há pedido analisado depois de aceito",IF(C{IR["Mudanças a rever"]}>0,"Há mudança sem análise, sem documentos ou sem aviso",'
     f'IF(C{IR["Itens da oferta a rever"]}>0,"Há item da oferta a rever",IF(C{IR["Pedidos a rever"]}>0,"Há pedido a rever","OK"))))))', merge=f"C{s+1}:H{s+1}", sz=9, b=False, h="left")
cf_warn(ws, f"C{s+1}:H{s+1}", f"C{s+1}")
ws.row_dimensions[s + 1].height = 21.75
NAV = s + 1
barras(ws, f"B{s + 3}", Q1, Q1 + NQ - 1, 2, 4, width=20, height=8)
setup(ws, REDC, f"B1:H{s + 20}", landscape=False, fit_height=True)

# ------------------------------------------------------------------ Checklist
ws = wb.create_sheet("Checklist")
widths(ws, {"A": 2, "B": 5, "C": 70, "D": 14, "E": 44, "F": 2})
title(ws, "Checklist dos requisitos do cliente", "Escolha o status de cada item na lista suspensa (coluna em amarelo-claro).", "E")
for col, text in zip("BCDE", ["#", "Item de verificação", "Status", "Observações"]):
    head(ws, f"{col}4", text)
ws.row_dimensions[4].height = 21.75
for k, text in enumerate(CHECK):
    rr = 5 + k
    num(ws, f"B{rr}", k + 1)
    put(ws, f"C{rr}", text, bg=GRAY)
    inp(ws, f"D{rr}", h="center")
    inp(ws, f"E{rr}")
    ws.row_dimensions[rr].height = 30
dv_list(ws, "D5:D16", YESNO, "Sim, Parcial ou Não")
cf_equal(ws, "D5:D16", YESNO_CF)
band(ws, 18, "Resumo automático", "E")
summary(ws, 19, "Itens atendidos (Sim)", '=COUNTIF(D5:D16,"Sim")', "C", "D")
summary(ws, 20, "Itens parciais", '=COUNTIF(D5:D16,"Parcial")', "C", "D")
summary(ws, 21, "Itens não atendidos", '=COUNTIF(D5:D16,"Não")', "C", "D")
summary(ws, 22, "Itens sem resposta", "=COUNTBLANK(D5:D16)", "C", "D")
summary(ws, 23, "Percentual atendido", "=D19/ROWS(D5:D16)", "C", "D", fmt="0%")
summary(ws, 24, "Situação", '=IF(D19=ROWS(D5:D16),"Validado",IF(D22>0,"Em andamento","Com pendências"))', "C", "D")
setup(ws, PURPLE, "B1:E24", landscape=False, fit_height=True)


# ------------------------------------------------------------------ Exemplos
def sub(ws, rr, cells, height=30):
    for a, b, t in cells:
        put(ws, f"{a}{rr}", t, f=font(10, True), bg=GRAY, h="center", merge=b and f"{a}{rr}:{b}{rr}")
    ws.row_dimensions[rr].height = height


EQ = [get_column_letter(7 + k) for k in range(NQ)]     # G a M: as sete perguntas, nos exemplos


def exemplo(ws, ex_):
    H = ex_["head"]
    widths(ws, dict({"A": 2, "B": 9, "C": 10, "D": 17, "E": 30, "F": 9, "N": 17, "O": 17, "P": 32, "Q": 18, "R": 10, "S": 12, "T": 28, "U": 2}, **{c: 5.5 for c in EQ}))
    title(ws, "Requisitos do cliente e análise de pedidos", "Exemplo preenchido, para consulta. Use as abas Oferta, Pedidos e Mudanças para a sua organização.", "T")
    rr = 4
    for rot, val in [("Organização", H["org"]), ("Processo", H["processo"]), ("Quem analisa", H["analisa"]), ("Registros", f'{H["periodo"]}, lidos em {H["ref"]:%d/%m/%Y}.'),
                     ("Origem", H["origem"])]:
        label(ws, f"B{rr}", rot, merge=f"B{rr}:D{rr}")
        inp(ws, f"E{rr}", val, merge=f"E{rr}:T{rr}", bg=WHITE, h="left")
        ws.row_dimensions[rr].height = alt([(val, 170)], minimo=19.5)
        rr += 1
    rr += 1
    band(ws, rr, "A oferta", "T", color=BLUE)
    rr += 1
    sub(ws, rr, [("B", "C", "Item"), ("D", "E", "O que se oferece"), ("F", "M", "Requisitos legais"), ("N", "O", "O que não se garante"), ("P", None, "Onde o cliente é informado"),
                 ("Q", "S", "Capacidade confirmada?"), ("T", None, "Conferência")])
    o1 = rr + 1
    for o in ex_["ofertas"]:
        rr += 1
        put(ws, f"B{rr}", o["item"], f=font(10, True), merge=f"B{rr}:C{rr}")
        put(ws, f"D{rr}", o["espec"], merge=f"D{rr}:E{rr}")
        put(ws, f"F{rr}", o["legal"], merge=f"F{rr}:M{rr}")
        put(ws, f"N{rr}", o["naogar"], merge=f"N{rr}:O{rr}")
        put(ws, f"P{rr}", o["onde"])
        put(ws, f"Q{rr}", o["cap"], h="center", merge=f"Q{rr}:S{rr}")
        calc(ws, f"T{rr}", "=" + f_oferta(f"B{rr}", "1", f"D{rr}", f"P{rr}", f"Q{rr}"), b=False, sz=9)
        ws.row_dimensions[rr].height = alt([(o["espec"], 44), (o["legal"], 40), (o["naogar"], 30), (o["item"], 16)], minimo=21.75)
    cf_warn(ws, f"T{o1}:T{rr}", f"T{o1}")
    rr += 2
    band(ws, rr, "Pedidos", "T", color=TEAL)
    rr += 1
    sub(ws, rr, [("B", None, "Nº"), ("C", None, "Data"), ("D", None, "Cliente"), ("E", None, "O que pede"), ("F", None, "Prazo")] + [(c, None, str(k + 1)) for k, c in enumerate(EQ)]
        + [("N", None, "Leitura"), ("O", None, "Decisão"), ("P", None, "O que foi negociado, ou o motivo"), ("Q", None, "Quem analisou"), ("R", None, "Análise"), ("S", None, "Confirmação"),
           ("T", None, "Conferência")])
    p1 = rr + 1
    for p in ex_["peds"]:
        rr += 1
        put(ws, f"B{rr}", p["num"], f=font(10, True, c=MUTED), h="center")
        put(ws, f"C{rr}", p["data"], h="center", fmt=DATE)
        put(ws, f"D{rr}", f'{p["cliente"]} · {p["canal"]}')
        put(ws, f"E{rr}", p["pede"] + (f' {p["naodecl"]}' if p["naodecl"] else ""))
        put(ws, f"F{rr}", p["prazo"], h="center", fmt="dd/mm")
        for c, a in zip(EQ, p["q"]):
            put(ws, f"{c}{rr}", a, h="center", f=font(8))
        qs = f"{EQ[0]}{rr}:{EQ[-1]}{rr}"
        calc(ws, f"N{rr}", "=" + f_leitura(f'B{rr}=""', qs), sz=9)
        put(ws, f"O{rr}", p["dec"], h="center")
        put(ws, f"P{rr}", p["neg"] or None)
        put(ws, f"Q{rr}", p["quem"])
        put(ws, f"R{rr}", p["analise"], h="center", fmt="dd/mm")
        put(ws, f"S{rr}", p["confirm"], h="center", fmt="dd/mm")
        calc(ws, f"T{rr}", "=" + f_conf(f'B{rr}=""', qs, f"O{rr}", f"P{rr}", f"Q{rr}", f"R{rr}", f"S{rr}"), b=False, sz=9)
        ws.row_dimensions[rr].height = alt([(p["pede"] + p["naodecl"], 30), (p["neg"], 34), (p["cliente"] + p["canal"], 16), (p["quem"], 17)], minimo=21.75)
    p2 = rr
    cf_texto(ws, f"{EQ[0]}{p1}:{EQ[-1]}{p2}", f"{EQ[0]}{p1}", RESP_CF)
    cf_texto(ws, f"N{p1}:N{p2}", f"N{p1}", LEIT_CF)
    cf_texto(ws, f"O{p1}:O{p2}", f"O{p1}", DEC_CF)
    cf_texto(ws, f"T{p1}:T{p2}", f"T{p1}", [("OK", GREEN)] + [(g, RED) for g in GRAVE], resto=YELLOW)
    rr += 1
    put(ws, f"B{rr}", "Perguntas: " + " · ".join(f"{k} {t.lower()}" for k, (t, _) in enumerate(PERGUNTAS, 1)) + ".", f=font(9, i=True, c=MUTED), box=False, merge=f"B{rr}:T{rr}")
    rr += 2
    ws.row_breaks.append(Break(id=rr - 1))
    band(ws, rr, "Mudanças nos pedidos aceitos", "T", color=AMBER)
    rr += 1
    sub(ws, rr, [("B", None, "Data"), ("C", None, "Pedido"), ("D", "E", "O que mudou"), ("F", "M", "Análise da mudança"), ("N", "O", "Quem pediu"), ("P", None, "Quem foi informado"),
                 ("Q", "S", "Documentos atualizados?"), ("T", None, "Conferência")])
    m1 = rr + 1
    for m in ex_["muds"]:
        rr += 1
        put(ws, f"B{rr}", m["data"], h="center", fmt=DATE)
        put(ws, f"C{rr}", m["ped"], h="center")
        put(ws, f"D{rr}", m["oque"], f=font(10, True), merge=f"D{rr}:E{rr}")
        put(ws, f"F{rr}", m["analise"] or None, merge=f"F{rr}:M{rr}")
        put(ws, f"N{rr}", m["pediu"], merge=f"N{rr}:O{rr}")
        put(ws, f"P{rr}", m["inform"] or None)
        put(ws, f"Q{rr}", m["docs"], h="center", merge=f"Q{rr}:S{rr}")
        calc(ws, f"T{rr}", "=" + f_mud(f"D{rr}", "1", f"C{rr}", f"F{rr}", f"Q{rr}", f"P{rr}", peds=f"$B${p1}:$B${p2}"), b=False, sz=9)
        ws.row_dimensions[rr].height = alt([(m["oque"], 44), (m["analise"], 42), (m["inform"], 30)], minimo=21.75)
    m2 = rr
    cf_warn(ws, f"T{m1}:T{m2}", f"T{m1}")
    rr += 2
    band(ws, rr, "Resumo automático", "T")
    for k, (text, formula) in enumerate([
        ("Pedidos por decisão", "=" + '&", "&'.join(f'COUNTIF(O{p1}:O{p2},"{d}")&" {PLURAL[d]}"' for d in DECISOES)),
        ("Pedidos em que a análise encontrou problema", "=SUMPRODUCT(((" + "+".join(f'({c}{p1}:{c}{p2}="{NAO}")' for c in EQ) + ")>0)*1)"),
        ("Pedidos a rever", f'=SUMPRODUCT((T{p1}:T{p2}<>"OK")*1)'),
        ("Respostas “não” por pergunta", "=" + '&" · "&'.join(f'"{k + 1}: "&COUNTIF({c}{p1}:{c}{p2},"{NAO}")' for k, c in enumerate(EQ))),
        ("Mudanças: registradas, a rever", f'=COUNTA(D{m1}:D{m2})&" registradas, "&SUMPRODUCT((T{m1}:T{m2}<>"OK")*1)&" a rever"'),
    ], 1):
        label(ws, f"B{rr+k}", text, merge=f"B{rr+k}:E{rr+k}", h="right")
        calc(ws, f"F{rr+k}", formula, merge=f"F{rr+k}:T{rr+k}", h="left")
        ws.row_dimensions[rr + k].height = 21.75
    setup(ws, MUTED, f"B1:T{rr + 5}")


exemplo(wb.create_sheet("Exemplo 1 - Pizzaria"), EX1)
exemplo(wb.create_sheet("Exemplo 2 - Indústria"), EX2)

wb.active = 1
wb.save(OUT)
print("ok", OUT, "| avisos: Oferta D%d, Pedidos E%d, Mudanças E%d, Painel C%d" % (OAV, PAV, MAV, NAV))
