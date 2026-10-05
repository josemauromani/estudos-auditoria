# -*- coding: utf-8 -*-
"""Gera Histograma-CEP-modelo.xlsx no mesmo padrão visual das planilhas anteriores da série."""
import os
import sys

_here = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, _here)
# funções de formatação compartilhadas com o gerador do PDCA (bloco inicial do arquivo)
_src = open(os.path.join(_here, "build_pdca.py"), encoding="utf-8").read()
exec(_src[: _src.index("STATUS = [")])

from openpyxl.chart import BarChart, Series  # noqa: E402
from cep_data import (A2, ACOMP, BASE, CAPAZ, CHECK, CPK_BOM, CPK_MIN, D2, D4, EX1, EX2, EXCL, LIMITE, N, NAOCAPAZ, NCLASSES, S_AMP, S_FORA, S_LADO,  # noqa: E402
                      S_TEND, SEP, SIM, SINAIS)

BLUE, TEAL, AMBER, REDC = "2B5C8A", "1E7B73", "A96A12", "B0413E"
S2_T = "F8EBCB"
YESNO = ["Sim", "Parcial", "Não"]
YESNO_CF = [("Sim", GREEN), ("Parcial", YELLOW), ("Não", RED)]
NS = 30
W = {"A": 2, "B": 5, "C": 12, "D": 16, "E": 9, "F": 9, "G": 9, "H": 9, "I": 9, "J": 10, "K": 34, "L": 11, "M": 11, "N": 16, "O": 10, "P": 24, "Q": 18, "R": 2}
FASE_CF = [(BASE, GREEN), (ACOMP, GRAY), (EXCL, YELLOW)]
SIG_CF = [(S_FORA, RED), (S_AMP, RED), (S_LADO, S2_T), (S_TEND, S2_T)]
NUM = "0.0##"
BR = lambda x: str(x).replace(".", ",")  # noqa: E731
BR_BOM, BR_MIN = f"{CPK_BOM:.2f}".replace(".", ","), f"{CPK_MIN:.2f}".replace(".", ",")


def note(ws, ref, text):
    c = Comment(text, "Modelo Histograma e CEP")
    c.width, c.height = 300, 120
    ws[ref].comment = c


def cf_warn(ws, rng_, first, ok_values=("OK",)):
    cond = ",".join(f'{first}="{v}"' for v in ok_values)
    ws.conditional_formatting.add(rng_, FormulaRule(formula=[f"OR({cond})"], fill=PatternFill("solid", bgColor=GREEN, fgColor=GREEN)))
    ws.conditional_formatting.add(rng_, FormulaRule(formula=[f"AND(LEN({first})>0,NOT(OR({cond})))"],
                                  fill=PatternFill("solid", bgColor=YELLOW, fgColor=YELLOW)))


def cf_texto(ws, rng_, first, pares):
    for t, cor in pares:
        ws.conditional_formatting.add(rng_, FormulaRule(formula=[f'{first}="{t}"'], stopIfTrue=True, fill=PatternFill("solid", bgColor=cor, fgColor=cor)))


def chart_style(ch):
    ch.style = 2
    ch.title = None
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ch.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill="D3DBE0", w=9525))
    ch.x_axis.graphicalProperties = GraphicalProperties(ln=LineProperties(solidFill="D3DBE0", w=9525))
    ch.y_axis.graphicalProperties = GraphicalProperties(ln=LineProperties(noFill=True))


def carta(ws, anchor, r1, r2, c_cat, series, titulo, width=26, height=7.5):
    """Gráfico de controle: a série de pontos e as linhas de limite, cada uma numa coluna."""
    ch = LineChart()
    ch.height, ch.width = height, width
    chart_style(ch)
    ch.legend.position = "b"
    ch.y_axis.title = titulo
    ch.dispBlanksAs = "gap"
    for col, nome, cor, dash, larg in series:
        s = Series(Reference(ws, min_col=col, min_row=r1, max_row=r2), title=nome)
        s.graphicalProperties.line.solidFill = cor
        s.graphicalProperties.line.width = larg
        if dash:
            s.graphicalProperties.line.dashStyle = dash
            s.marker.symbol = "none"
        else:
            s.marker.symbol = "circle"
            s.marker.size = 6
            s.marker.graphicalProperties = GraphicalProperties(solidFill=cor)
            s.marker.graphicalProperties.line.solidFill = WHITE
        s.smooth = False
        ch.append(s)
    ch.set_categories(Reference(ws, min_col=c_cat, min_row=r1, max_row=r2))
    ws.add_chart(ch, anchor)


def colunas(ws, anchor, r1, r2, c_cat, c_val, width=18, height=8):
    ch = BarChart()
    ch.type = "col"
    ch.gapWidth = 10
    ch.height, ch.width = height, width
    chart_style(ch)
    ch.legend = None
    s = Series(Reference(ws, min_col=c_val, min_row=r1, max_row=r2), title="Medições")
    s.graphicalProperties.solidFill = BLUE
    s.graphicalProperties.line.solidFill = WHITE
    ch.append(s)
    ch.set_categories(Reference(ws, min_col=c_cat, min_row=r1, max_row=r2))
    ch.y_axis.scaling.min = 0
    ws.add_chart(ch, anchor)


def conta_por(rng, valores, contem=False):
    if contem:  # a célula pode trazer mais de um valor, separados por ponto e vírgula
        return "=" + '&", "&'.join(f'SUMPRODUCT(--ISNUMBER(SEARCH("{x}",{rng})))&" {x.lower()}"' for x in valores)
    return "=" + '&", "&'.join(f'COUNTIF({rng},"{x}")&" {x.lower()}"' for x in valores)


def conta(rng, valor):
    """Quantas células da coluna Sinal trazem o sinal, sozinho ou junto de outro."""
    return f'=SUMPRODUCT(--ISNUMBER(SEARCH("{valor}",{rng})))'


# ------------------------------------------------------------------ o bloco de dados, igual na aba Dados e nos exemplos
def bloco(ws, r0, ns, ex=None, nf="General"):
    """Parâmetros, limites calculados e a tabela de subgrupos. Devolve as referências usadas pelas outras abas."""
    h = ex["head"] if ex else None
    bg = WHITE if ex else INPUT
    P = {}
    for k, (rot, key, fmt) in enumerate([("Característica medida", "carac", None), ("Unidade", "unid", None), ("Especificação: mínimo (LIE)", "lie", nf),
                                         ("Especificação: máximo (LSE)", "lse", nf), ("Subgrupos da base", "base", "0")]):
        r = r0 + k
        label(ws, f"B{r}", rot, merge=f"B{r}:D{r}")
        inp(ws, f"E{r}", h[key] if ex else None, h="left" if fmt is None else "center", fmt=fmt, merge=f"E{r}:I{r}", bg=bg)
        P[key] = f"$E${r}"
        ws.row_dimensions[r].height = 19.5
    r = r0 + 5
    label(ws, f"B{r}", "Tamanho do subgrupo", merge=f"B{r}:D{r}")
    put(ws, f"E{r}", N, bg=GRAY, h="center", merge=f"E{r}:I{r}")
    if not ex:
        dv_number(ws, f"E{r0 + 2}:E{r0 + 3}")
        note(ws, f"B{r0 + 4}", "Os primeiros subgrupos, contados de cima, que formam a base dos limites. Em branco, todos entram na base. Recomendado: 20 a 25.")
        note(ws, f"B{r0 + 5}", "Este modelo usa subgrupos de cinco medições, com as constantes de Shewhart para n = 5.")
    hr = r0 + 8
    d1, d2 = hr + 1, hr + ns
    rng = lambda c: f"${c}${d1}:${c}${d2}"  # noqa: E731
    lim_rows = [("Subgrupos na base", "nb", f'=COUNTIF({rng("N")},"{BASE}")', "0"),
                ("Linha central da média (LC)", "lc", f'=IF({{nb}}=0,"",AVERAGEIF({rng("N")},"{BASE}",{rng("L")}))', NUM),
                ("Limite superior de controle (LSC)", "lsc", f'=IF({{lc}}="","",{{lc}}+{A2}*{{rb}})', NUM),
                ("Limite inferior de controle (LIC)", "lic", f'=IF({{lc}}="","",{{lc}}-{A2}*{{rb}})', NUM),
                ("Amplitude média (R̄)", "rb", f'=IF({{nb}}=0,"",AVERAGEIF({rng("N")},"{BASE}",{rng("M")}))', NUM),
                ("Limite da amplitude (LSC de R)", "lscr", f'=IF({{rb}}="","",{D4}*{{rb}})', NUM),
                ("Desvio estimado (σ = R̄ ÷ d2)", "sigma", f'=IF({{rb}}="","",{{rb}}/{D2})', "0.0###")]
    for k, (rot, key, _, _) in enumerate(lim_rows):
        P[key] = f"$L${r0 + k}"
    for k, (rot, key, f, fmt) in enumerate(lim_rows):
        r = r0 + k
        label(ws, f"K{r}", rot)
        calc(ws, f"L{r}", f.format(**{kk: P[kk] for kk in ("nb", "lc", "rb")}), fmt=fmt, merge=f"L{r}:M{r}")
    dicas = ["Da aba: os subgrupos marcados como Base.", "Média das médias da base.", f"LC + {BR(A2)} × R̄.", f"LC − {BR(A2)} × R̄.", "Média das amplitudes da base.",
             f"{BR(D4)} × R̄. O inferior é zero.", f"R̄ ÷ {BR(D2)}. Usado na capacidade."]
    for k, t in enumerate(dicas):
        put(ws, f"N{r0 + k}", t, f=font(9, i=True, c=MUTED), bg=GRAY, merge=f"N{r0 + k}:Q{r0 + k}")
    for col, text in zip("BCDEFGHIJKLMNOPQ", ["#", "Data", "Identificação", "Medição 1", "Medição 2", "Medição 3", "Medição 4", "Medição 5", "Excluir da base",
                                            "Motivo da exclusão", "Média", "Amplitude", "Fase", "Fora da espec.", "Sinal", "Conferência"]):
        head(ws, f"{col}{hr}", text)
    ws.row_dimensions[hr].height = 33
    LC, LSC, LIC, LSCR, LIE, LSE, BS = P["lc"], P["lsc"], P["lic"], P["lscr"], P["lie"], P["lse"], P["base"]
    for k in range(ns):
        r = d1 + k
        num(ws, f"B{r}", k + 1)
        g = ex["subs"][k] if ex else None
        if ex:
            put(ws, f"C{r}", g["data"], h="center", fmt=DATE)
            put(ws, f"D{r}", g["id"])
            for j, col in enumerate("EFGHI"):
                put(ws, f"{col}{r}", g["v"][j], h="center", fmt=nf)
            put(ws, f"J{r}", g["excl"] or None, h="center")
            put(ws, f"K{r}", g["motivo"] or None, f=font(9))
        else:
            inp(ws, f"C{r}", h="center", fmt=DATE)
            inp(ws, f"D{r}")
            for col in "EFGHI":
                inp(ws, f"{col}{r}", h="center", fmt=nf)
            inp(ws, f"J{r}", h="center")
            inp(ws, f"K{r}")
        V = f"E{r}:I{r}"
        calc(ws, f"L{r}", f'=IF(COUNT({V})={N},AVERAGE({V}),"")', fmt="0.00#")
        calc(ws, f"M{r}", f'=IF(COUNT({V})={N},MAX({V})-MIN({V}),"")', b=False, fmt=NUM)
        calc(ws, f"N{r}", f'=IF(L{r}="","",IF(J{r}="{SIM}","{EXCL}",IF(OR({BS}="",B{r}<={BS}),"{BASE}","{ACOMP}")))', b=False, sz=9)
        calc(ws, f"O{r}", f'=IF(OR(COUNT({V})=0,{LIE}="",{LSE}=""),"",COUNTIF({V},"<"&{LIE})+COUNTIF({V},">"&{LSE}))', b=False)
        run = (f'OR(COUNTIF(L{r - 6}:L{r},">"&{LC})=7,COUNTIF(L{r - 6}:L{r},"<"&{LC})=7)' if k >= 6 else "FALSE")
        if k >= 5:
            up = ",".join(f"L{r - j}>L{r - j - 1}" for j in range(5))
            dn = ",".join(f"L{r - j}<L{r - j - 1}" for j in range(5))
            trend = f"AND(COUNT(L{r - 5}:L{r})=6,OR(AND({up}),AND({dn})))"
        else:
            trend = "FALSE"
        sm = f'IF(OR(L{r}>{LSC},L{r}<{LIC}),"{S_FORA}",IF({run},"{S_LADO}",IF({trend},"{S_TEND}","")))'  # o sinal da média
        calc(ws, f"P{r}", f'=IF(OR(L{r}="",{LC}=""),"",{sm}&IF(M{r}>{LSCR},IF({sm}="","","{SEP}")&"{S_AMP}",""))', b=False, sz=9)
        calc(ws, f"Q{r}", f'=IF(COUNT({V})+LEN(C{r})+LEN(D{r})+LEN(J{r})=0,"",IF(AND(COUNT({V})>0,COUNT({V})<{N}),"Faltam medições",'
                          f'IF(AND(COUNT({V})>0,C{r}=""),"Falta a data",IF(AND(J{r}="{SIM}",K{r}=""),"Falta o motivo",IF(COUNT({V})=0,"Faltam medições","OK")))))', b=False, sz=9)
        ws.row_dimensions[r].height = 19.5 if not (g and g["motivo"]) else 31.5
    cf_texto(ws, f"N{d1}:N{d2}", f"N{d1}", FASE_CF)
    for t, cor in SIG_CF:  # a célula pode trazer dois sinais: a cor segue o primeiro da lista que aparece nela
        ws.conditional_formatting.add(f"P{d1}:P{d2}", FormulaRule(formula=[f'ISNUMBER(SEARCH("{t}",P{d1}))'], stopIfTrue=True,
                                                                  fill=PatternFill("solid", bgColor=cor, fgColor=cor)))
    cf_warn(ws, f"Q{d1}:Q{d2}", f"Q{d1}")
    # medição fora da especificação, em vermelho
    ws.conditional_formatting.add(f"E{d1}:I{d2}", FormulaRule(formula=[f'AND(ISNUMBER(E{d1}),{LIE}<>"",{LSE}<>"",OR(E{d1}<{LIE},E{d1}>{LSE}))'],
                                                             fill=PatternFill("solid", bgColor=RED, fgColor=RED)))
    if not ex:
        dv_date(ws, f"C{d1}:C{d2}")
        dv_number(ws, f"E{d1}:I{d2}")
        dv_list(ws, f"J{d1}:J{d2}", [SIM], "Sim, para tirar o subgrupo do cálculo dos limites")
    P.update(hr=hr, d1=d1, d2=d2)
    return P


def resumo_rows(P, aba=""):
    """Os indicadores do resumo, comuns ao painel e aos exemplos."""
    q = lambda c: f"{aba}${c}${P['d1']}:${c}${P['d2']}"  # noqa: E731
    p = lambda k: aba + P[k]  # noqa: E731
    cp = f'IF(OR({p("sigma")}="",{p("lie")}="",{p("lse")}=""),"",IF({p("sigma")}=0,"",({p("lse")}-{p("lie")})/(6*{p("sigma")})))'
    cpk = f'IF(OR({p("sigma")}="",{p("lie")}="",{p("lse")}=""),"",IF({p("sigma")}=0,"",MIN({p("lse")}-{p("lc")},{p("lc")}-{p("lie")})/(3*{p("sigma")})))'
    return [
        ("SUBGRUPOS", None, None, None),
        ("Subgrupos completos", f'=COUNT({q("L")})', "Com as cinco medições.", "0"),
        ("Na base", f'=COUNTIF({q("N")},"{BASE}")', "Usados nos limites.", "0"),
        ("Excluídos", f'=COUNTIF({q("N")},"{EXCL}")', "Com causa conhecida.", "0"),
        ("Em acompanhamento", f'=COUNTIF({q("N")},"{ACOMP}")', "Lidos contra os limites fixos.", "0"),
        ("Linhas a completar", f'=SUMPRODUCT(({q("Q")}<>"")*({q("Q")}<>"OK"))', "Conferências diferentes de OK.", "0"),
        ("SINAIS", None, None, None),
        (S_FORA, conta(q("P"), S_FORA), "Média acima do LSC ou abaixo do LIC.", "0"),
        (S_AMP, conta(q("P"), S_AMP), "Amplitude acima do limite.", "0"),
        (S_LADO, conta(q("P"), S_LADO), "Sequência do mesmo lado da linha central.", "0"),
        (S_TEND, conta(q("P"), S_TEND), "Tendência: desgaste ou aquecimento.", "0"),
        ("Sinais na base", f'=SUMPRODUCT(({q("N")}="{BASE}")*(LEN({q("P")})>0))', "Com sinal na base, a capacidade não vale.", "0"),
        ("Sinais no acompanhamento", f'=SUMPRODUCT(({q("N")}="{ACOMP}")*(LEN({q("P")})>0))', "Cada um pede causa e ação.", "0"),
        ("Primeiro sinal no acompanhamento", f'=IFERROR(INDEX({q("C")},MATCH(1,INDEX(({q("N")}="{ACOMP}")*(LEN({q("P")})>0),0),0)),"")', "A data do primeiro aviso.", DATE),
        ("CAPACIDADE", None, None, None),
        ("Medições", f'=COUNT({aba}$E${P["d1"]}:$I${P["d2"]})', "Valores individuais.", "0"),
        ("Medições fora da especificação", f'=SUM({q("O")})', "Cada uma é uma unidade fora.", "0"),
        ("Cp", "=" + cp, "Tolerância ÷ seis desvios.", "0.00"),
        ("Cpk", "=" + cpk, f"{BR_BOM} ou mais: capaz. De {BR_MIN} a {BR_BOM}: no limite.", "0.00"),
        ("Situação da capacidade", None, "Só vale com a base sem sinal.", None),
    ]


def capac_formula(cpk_ref, base_sig_ref):
    return (f'=IF({cpk_ref}="","",IF({base_sig_ref}>0,"Base com sinal: trate antes",IF({cpk_ref}>={CPK_BOM},"{CAPAZ}",'
            f'IF({cpk_ref}>={CPK_MIN},"{LIMITE}","{NAOCAPAZ}"))))')


def escreve_resumo(ws, P, r, aba="", last="F", lab="B", val="C", leit="D"):
    IR = {}
    for nome, formula, leitura, fmt in resumo_rows(P, aba):
        if formula is None and leitura is None:
            put(ws, f"{lab}{r}", nome, f=font(9, True, c=MUTED), bg=None, box=False, merge=f"{lab}{r}:{last}{r}")
            ws.row_dimensions[r].height = 21.75
            r += 1
            continue
        IR[nome] = r
        if formula is None:
            formula = capac_formula(f"{val}{IR['Cpk']}", f"{val}{IR['Sinais na base']}")
        label(ws, f"{lab}{r}", nome)
        calc(ws, f"{val}{r}", formula, fmt=fmt, sz=9 if nome.startswith("Situação") else 10)
        put(ws, f"{leit}{r}", leitura, f=font(9, c=MUTED), bg=GRAY, merge=f"{leit}{r}:{last}{r}")
        ws.row_dimensions[r].height = 21.75
        r += 1
    cf_texto(ws, f"{val}{IR['Situação da capacidade']}", f"{val}{IR['Situação da capacidade']}",
             [(CAPAZ, GREEN), (LIMITE, S2_T), (NAOCAPAZ, RED), ("Base com sinal: trate antes", YELLOW)])
    return IR, r


def f_aviso(v, IR):
    c = lambda n: f"{v}{IR[n]}"  # noqa: E731
    return (f'=IF({c("Subgrupos completos")}=0,"Lance os subgrupos na aba Dados",IF({c("Na base")}<20,"Base com menos de 20 subgrupos: limites provisórios",'
            f'IF({c("Linhas a completar")}>0,"Há linha a completar: veja a coluna Conferência",IF({c("Sinais na base")}>0,"Há sinal na base: procure a causa e exclua o subgrupo",'
            f'IF({c("Sinais no acompanhamento")}>0,"Há sinal no acompanhamento: procure a causa e anote a ação",'
            f'IF({c("Situação da capacidade")}="{NAOCAPAZ}","Processo em controle, mas não capaz: mude o processo",'
            f'IF({c("Situação da capacidade")}="{LIMITE}","Capacidade no limite: acompanhe de perto","OK")))))))')


# ================================================================== pasta de trabalho
wb = Workbook()
wb.properties.title = "Histograma e CEP — Modelo"
wb.properties.creator = "Modelo Histograma e CEP"
wb.properties.language = "pt-BR"

# ------------------------------------------------------------------ Instruções
ws = wb.active
ws.title = "Instruções"
widths(ws, {"A": 2, "B": 26, "C": 90, "D": 2})
title(ws, "Histograma e CEP — Como usar esta planilha", "Modelo para o gráfico de média e amplitude, o histograma e a capacidade de uma característica.", "C")
r = 4


def section(text):
    global r
    band(ws, r, text, "C", sz=11)
    r += 1


def line(k, v, kbg=GRAY, vbg=None):
    global r
    put(ws, f"B{r}", k, f=font(10, True), bg=kbg)
    put(ws, f"C{r}", v, bg=vbg)
    ws.row_dimensions[r].height = 19.5 if len(v) <= 96 else (31.5 if len(v) <= 192 else 45.75)
    r += 1


section("Legenda: onde preencher")
line("Amarelo-claro", "Células de entrada. É aqui que você digita.", vbg=INPUT)
line("Cinza", "Células calculadas ou fixas (médias, amplitudes, fases, sinais, limites, capacidade). Não altere.", vbg=GRAY)
line("Vermelho na medição", "A medição está fora da especificação.")
line("Comentários", "Células com triângulo vermelho no canto trazem uma dica de preenchimento. Passe o mouse sobre elas.")
r += 1
section("As regras de cálculo")
for k, text in [
    ("Média e amplitude", f"Só com as {N} medições do subgrupo preenchidas."),
    ("Fase", f"{EXCL} quando marcado para excluir. {BASE} para os primeiros subgrupos, até o número informado. {ACOMP} para os seguintes."),
    ("Limites", f"Calculados só com os subgrupos da base: LC ± {BR(A2)} × R̄ para a média; {BR(D4)} × R̄ para a amplitude; σ = R̄ ÷ {BR(D2)}."),
    ("Sinais", f"Na média, vale a primeira regra, nesta ordem: {S_FORA.lower()}; sete médias seguidas do mesmo lado da linha central; seis médias seguidas subindo ou descendo. "
               f"A {S_AMP.lower()} é marcada junto, separada por ponto e vírgula."),
    ("Capacidade", f"Cp = (LSE − LIE) ÷ 6σ. Cpk = menor distância da LC a um limite ÷ 3σ. {CAPAZ}: {BR_BOM} ou mais; {LIMITE}: de {BR_MIN} a {BR_BOM}; abaixo: {NAOCAPAZ.lower()}."),
    ("Histograma", f"{NCLASSES} classes de largura igual a um oitavo da tolerância, a partir de dois oitavos abaixo do LIE. Cada valor entra na classe em que é maior ou igual ao início."),
]:
    line(k, text)
r += 1
section("Ordem recomendada de preenchimento")
for k, text in enumerate([
    "Aba Dados: a característica, a unidade, o LIE, o LSE e quantos subgrupos formam a base.",
    "Aba Dados: cada subgrupo com a data, a identificação e as cinco medições.",
    "Aba Dados: marque Sim em Excluir da base para os subgrupos com causa conhecida, e escreva o motivo.",
    "Abas Gráficos e Histograma: leia os pontos, os limites e a forma.",
    "Aba Painel: leia a capacidade e o aviso. Depois, valide na aba Checklist.",
], 1):
    line(f"Passo {k}", text)
r += 1
section("Abas da planilha")
for k, text in [
    ("Dados", f"Até {NS} subgrupos de {N} medições, com os limites calculados no alto."),
    ("Gráficos", "O gráfico das médias e o das amplitudes, com a linha central e os limites."),
    ("Histograma", "As classes, as contagens e o gráfico das medições individuais."),
    ("Painel", "Os números dos subgrupos, dos sinais e da capacidade, com o aviso."),
    ("Checklist", "Doze verificações do gráfico, com percentual de conclusão."),
    ("Exemplo 1 - Pizzaria", "O peso da bola de massa: em controle, mas não capaz."),
    ("Exemplo 2 - Indústria", "A espessura do filme: capaz na base, com sinais no acompanhamento."),
]:
    line(k, text)
r += 1
section("Premissas do modelo")
for k, text in [
    ("Gráfico", f"Média e amplitude (X̄-R), com subgrupos de {N} e as constantes de Shewhart para esse tamanho. Outro tamanho de subgrupo pede outras constantes."),
    ("Regras e classes", "As regras de sinal, a classificação da capacidade e as classes do histograma são uma convenção deste material, próxima das usuais."),
    ("Exemplos", "Os dados das abas de exemplo são ilustrativos, criados para este material."),
]:
    line(k, text)
setup(ws, INK, f"B1:C{r}", landscape=False, fit_height=True)

# ------------------------------------------------------------------ Dados
ws = wb.create_sheet("Dados")
widths(ws, W)
title(ws, "Dados do gráfico de controle", "Um subgrupo por linha: as cinco medições feitas juntas, nas mesmas condições.", "Q")
PD = bloco(ws, 4, NS)
ws.freeze_panes = f"E{PD['d1']}"
s = PD["d2"] + 2
band(ws, s, "Resumo automático", "Q")
label(ws, f"B{s + 1}", "Fases", merge=f"B{s + 1}:D{s + 1}", h="right")
calc(ws, f"E{s + 1}", conta_por(f"N{PD['d1']}:N{PD['d2']}", (BASE, EXCL, ACOMP)), merge=f"E{s + 1}:Q{s + 1}", h="left")
label(ws, f"B{s + 2}", "Sinais", merge=f"B{s + 2}:D{s + 2}", h="right")
calc(ws, f"E{s + 2}", conta_por(f"P{PD['d1']}:P{PD['d2']}", SINAIS, contem=True), merge=f"E{s + 2}:Q{s + 2}", h="left")
for rr in (s + 1, s + 2):
    ws.row_dimensions[rr].height = 21.75
DAV = s + 2
setup(ws, BLUE, f"B1:Q{DAV}", fit_height=True)

# ------------------------------------------------------------------ Gráficos
ws = wb.create_sheet("Gráficos")
widths(ws, {"A": 2, "B": 6, "C": 11, "D": 11, "E": 11, "F": 11, "G": 11, "H": 11, "I": 11, "J": 2})
title(ws, "Gráficos de controle", "Nada a preencher: os pontos e os limites vêm da aba Dados.", "I")
for col, text in zip("BCDEFGHI", ["#", "Média", "LSC", "LC", "LIC", "Amplitude", "LSC de R", "R̄"]):
    head(ws, f"{col}4", text)
put(ws, "B3", "Linhas sem subgrupo mostram #N/D, para o gráfico não desenhar um ponto em zero.", f=font(9, i=True, c=MUTED), bg=None, box=False, merge="B3:I3")
DA = "Dados!"
for k in range(NS):
    r, rd = 5 + k, PD["d1"] + k
    num(ws, f"B{r}", k + 1)
    have = f'{DA}L{rd}=""'
    for col, ref in (("C", f"{DA}L{rd}"), ("D", DA + PD["lsc"]), ("E", DA + PD["lc"]), ("F", DA + PD["lic"]), ("G", f"{DA}M{rd}"), ("H", DA + PD["lscr"]), ("I", DA + PD["rb"])):
        calc(ws, f"{col}{r}", f'=IF(OR({have},{ref}=""),NA(),{ref})', b=False, fmt="0.00#", sz=9)
ws.conditional_formatting.add(f"C5:I{4 + NS}", FormulaRule(formula=["ISERROR(C5)"], font=Font(color="B8C2C9")))
G2 = 4 + NS
carta(ws, "K4", 5, G2, 2, [(3, "Média", BLUE, None, 22225), (4, "LSC", REDC, "dash", 15875), (5, "LC", MUTED, "solid", 12700), (6, "LIC", REDC, "dash", 15875)], "Média")
carta(ws, "K20", 5, G2, 2, [(7, "Amplitude", TEAL, None, 22225), (8, "LSC de R", REDC, "dash", 15875), (9, "R̄", MUTED, "solid", 12700)], "Amplitude")
setup(ws, TEAL, f"B1:Z{G2}")

# ------------------------------------------------------------------ Histograma
ws = wb.create_sheet("Histograma")
widths(ws, {"A": 2, "B": 6, "C": 20, "D": 14, "E": 16, "F": 2})
title(ws, "Histograma das medições", "Nada a preencher: as classes vêm da especificação, e as medições, da aba Dados.", "E")
LIE_, LSE_ = DA + PD["lie"], DA + PD["lse"]
label(ws, "B4", "Largura da classe", merge="B4:C4")
calc(ws, "D4", f'=IF(OR({LIE_}="",{LSE_}=""),"",({LSE_}-{LIE_})/8)', fmt="0.0##")
label(ws, "B5", "Início da primeira classe", merge="B5:C5")
calc(ws, "D5", f'=IF(D4="","",{LIE_}-2*D4)', fmt="0.0##")
for col, text in zip("BCDE", ["#", "De", "Até (exclusive)", "Medições"]):
    head(ws, f"{col}7", text)
VALS = f"Dados!$E${PD['d1']}:$I${PD['d2']}"
H1 = 8
for i in range(NCLASSES):
    r = H1 + i
    num(ws, f"B{r}", i + 1)
    calc(ws, f"C{r}", f'=IF($D$4="","",{LIE_}-2*$D$4+{i}*$D$4)', b=False, fmt="0.0##")
    calc(ws, f"D{r}", f'=IF($D$4="","",{LIE_}-2*$D$4+{i + 1}*$D$4)', b=False, fmt="0.0##")
    calc(ws, f"E{r}", f'=IF(C{r}="","",COUNTIFS({VALS},">="&C{r},{VALS},"<"&D{r}))')
H2 = H1 + NCLASSES - 1
for k, (rot, f) in enumerate([("Abaixo da primeira classe", f'=IF($D$4="","",COUNTIF({VALS},"<"&C{H1}))'),
                              ("Acima da última classe", f'=IF($D$4="","",COUNTIF({VALS},">="&D{H2}))')]):
    r = H2 + 1 + k
    label(ws, f"B{r}", rot, merge=f"B{r}:D{r}", h="right")
    calc(ws, f"E{r}", f)
ws.conditional_formatting.add(f"B{H1}:E{H2}", FormulaRule(formula=[f'AND(ISNUMBER($C{H1}),OR($D{H1}<={LIE_},$C{H1}>={LSE_}))'],
                                                         fill=PatternFill("solid", bgColor=RED, fgColor=RED)))
s = H2 + 4
band(ws, s, "Resumo das medições", "E")
for k, (rot, f, fmt) in enumerate([("Medições", f"=COUNT({VALS})", "0"), ("Média geral", f'=IF(COUNT({VALS})=0,"",AVERAGE({VALS}))', "0.00#"),
                                   ("Menor", f'=IF(COUNT({VALS})=0,"",MIN({VALS}))', "0.0##"), ("Maior", f'=IF(COUNT({VALS})=0,"",MAX({VALS}))', "0.0##"),
                                   ("Fora da especificação", f'=IF(OR(COUNT({VALS})=0,{LIE_}="",{LSE_}=""),"",COUNTIF({VALS},"<"&{LIE_})+COUNTIF({VALS},">"&{LSE_}))', "0")]):
    summary(ws, s + 1 + k, rot, f, "D", "E", fmt=fmt)
HR = s + 5
colunas(ws, "G4", H1, H2, 3, 5)
put(ws, "G21", "Classes em vermelho na tabela ficam fora da especificação, inteira ou em parte.", f=font(9, i=True, c=MUTED), bg=None, box=False, merge="G21:P21")
setup(ws, AMBER, f"B1:P{HR}")

# ------------------------------------------------------------------ Painel
ws = wb.create_sheet("Painel")
widths(ws, {"A": 2, "B": 34, "C": 16, "D": 14, "E": 14, "F": 30, "G": 2})
title(ws, "Painel do processo", "Nada a preencher: tudo vem das outras abas.", "F")
label(ws, "B3", "Característica")
calc(ws, "C3", '=IF(Dados!E4="","",Dados!E4&IF(Dados!E5="",""," ("&Dados!E5&")"))', merge="C3:F3", h="left")
for col, text, m in [("B", "Indicador", None), ("C", "Resultado", None), ("D", "Como se lê", "D5:F5")]:
    put(ws, f"{col}5", text, f=font(10, True, c=WHITE), bg=INK, h="center", merge=m)
ws.row_dimensions[5].height = 21.75
IR, rr = escreve_resumo(ws, PD, 6, aba="Dados!")
s = rr + 1
band(ws, s, "Resumo automático", "F")
label(ws, f"B{s + 1}", "Aviso")
calc(ws, f"C{s + 1}", f_aviso("C", IR), merge=f"C{s + 1}:F{s + 1}", sz=9, b=False, h="left")
cf_warn(ws, f"C{s + 1}:F{s + 1}", f"C{s + 1}")
ws.row_dimensions[s + 1].height = 21.75
NAV = s + 1
setup(ws, REDC, f"B1:F{NAV}", landscape=False, fit_height=True)

# ------------------------------------------------------------------ Checklist
ws = wb.create_sheet("Checklist")
widths(ws, {"A": 2, "B": 5, "C": 70, "D": 14, "E": 44, "F": 2})
title(ws, "Checklist do CEP", "Escolha o status de cada item na lista suspensa (coluna em amarelo-claro).", "E")
for col, text in zip("BCDE", ["#", "Item de verificação", "Status", "Observações"]):
    head(ws, f"{col}4", text)
ws.row_dimensions[4].height = 21.75
for k, text in enumerate(CHECK):
    rr = 5 + k
    num(ws, f"B{rr}", k + 1)
    put(ws, f"C{rr}", text, bg=GRAY)
    inp(ws, f"D{rr}", h="center")
    inp(ws, f"E{rr}")
    ws.row_dimensions[rr].height = 30
dv_list(ws, "D5:D16", YESNO, "Sim, Parcial ou Não")
cf_equal(ws, "D5:D16", YESNO_CF)
band(ws, 18, "Resumo automático", "E")
summary(ws, 19, "Itens atendidos (Sim)", '=COUNTIF(D5:D16,"Sim")', "C", "D")
summary(ws, 20, "Itens parciais", '=COUNTIF(D5:D16,"Parcial")', "C", "D")
summary(ws, 21, "Itens não atendidos", '=COUNTIF(D5:D16,"Não")', "C", "D")
summary(ws, 22, "Itens sem resposta", "=COUNTBLANK(D5:D16)", "C", "D")
summary(ws, 23, "Percentual atendido", "=D19/ROWS(D5:D16)", "C", "D", fmt="0%")
summary(ws, 24, "Situação", '=IF(D19=ROWS(D5:D16),"Validado",IF(D22>0,"Em andamento","Com pendências"))', "C", "D")
setup(ws, PURPLE, "B1:E24", landscape=False, fit_height=True)


# ------------------------------------------------------------------ Exemplos
def exemplo(ws, ex_):
    H = ex_["head"]
    widths(ws, {**W, "R": 2, "S": 10, "T": 10, "U": 10, "V": 10, "W": 10})
    title(ws, "Histograma e CEP", "Exemplo preenchido, para consulta. Use a aba Dados para a sua característica.", "Q")
    rr = 4
    for rot, val in [("Organização", H["org"]), ("Processo", H["processo"]), ("Medição", H["instr"]), ("Subgrupo", H["subgrupo"] + ". " + H["periodo"] + "."),
                     ("Origem", H["origem"])]:
        label(ws, f"B{rr}", rot, merge=f"B{rr}:D{rr}")
        inp(ws, f"E{rr}", val, merge=f"E{rr}:Q{rr}", bg=WHITE, h="left")
        ws.row_dimensions[rr].height = 19.5 if len(val) < 140 else 31.5
        rr += 1
    P = bloco(ws, 10, len(ex_["subs"]), ex_, nf="0" if ex_ is EX1 else "0.0")
    # colunas auxiliares dos gráficos: os limites repetidos em cada linha
    for col, text in zip("STUVW", ["LSC", "LC", "LIC", "LSC de R", "R̄"]):
        head(ws, f"{col}{P['hr']}", text)
    for r in range(P["d1"], P["d2"] + 1):
        for col, key in zip("STUVW", ("lsc", "lc", "lic", "lscr", "rb")):
            calc(ws, f"{col}{r}", f"={P[key]}", b=False, fmt="0.00#", sz=9)
    g0 = P["d2"] + 2
    band(ws, g0, "Gráficos de controle", "Q", color=TEAL)
    carta(ws, f"B{g0 + 1}", P["d1"], P["d2"], 2, [(12, "Média", BLUE, None, 22225), (19, "LSC", REDC, "dash", 15875), (20, "LC", MUTED, "solid", 12700),
                                                 (21, "LIC", REDC, "dash", 15875)], "Média", width=17, height=7)
    carta(ws, f"K{g0 + 1}", P["d1"], P["d2"], 2, [(13, "Amplitude", TEAL, None, 22225), (22, "LSC de R", REDC, "dash", 15875), (23, "R̄", MUTED, "solid", 12700)],
          "Amplitude", width=17, height=7)
    s = g0 + 16
    ws.row_breaks.append(Break(id=s - 1))
    band(ws, s, "Resumo automático", "Q")
    for col, text, m in [("B", "Indicador", "B{0}:D{0}"), ("E", "Resultado", "E{0}:F{0}"), ("G", "Como se lê", "G{0}:Q{0}")]:
        put(ws, f"{col}{s + 1}", text, f=font(10, True, c=WHITE), bg=INK, h="center", merge=m.format(s + 1))
    IRx = {}
    r = s + 2
    for nome, formula, leitura, fmt in resumo_rows(P):
        if formula is None and leitura is None:
            put(ws, f"B{r}", nome, f=font(9, True, c=MUTED), bg=None, box=False, merge=f"B{r}:Q{r}")
            r += 1
            continue
        IRx[nome] = r
        if formula is None:
            formula = capac_formula(f"E{IRx['Cpk']}", f"E{IRx['Sinais na base']}")
        label(ws, f"B{r}", nome, merge=f"B{r}:D{r}")
        calc(ws, f"E{r}", formula, fmt=fmt, merge=f"E{r}:F{r}", sz=9 if nome.startswith("Situação") else 10)
        put(ws, f"G{r}", leitura, f=font(9, c=MUTED), bg=GRAY, merge=f"G{r}:Q{r}")
        ws.row_dimensions[r].height = 19.5
        r += 1
    cf_texto(ws, f"E{IRx['Situação da capacidade']}", f"E{IRx['Situação da capacidade']}",
             [(CAPAZ, GREEN), (LIMITE, S2_T), (NAOCAPAZ, RED), ("Base com sinal: trate antes", YELLOW)])
    label(ws, f"B{r}", "Aviso", merge=f"B{r}:D{r}")
    calc(ws, f"E{r}", f_aviso("E", IRx), merge=f"E{r}:Q{r}", sz=9, b=False, h="left")
    cf_warn(ws, f"E{r}:Q{r}", f"E{r}")
    ws.row_dimensions[r].height = 21.75
    setup(ws, MUTED, f"B1:Q{r}")
    return dict(P=P, IR=IRx, aviso=r)


POS = {}
POS["ex1"] = exemplo(wb.create_sheet("Exemplo 1 - Pizzaria"), EX1)
POS["ex2"] = exemplo(wb.create_sheet("Exemplo 2 - Indústria"), EX2)

wb.active = 1
wb.save(OUT)
print("ok", OUT, "| Dados", {k: PD[k] for k in ("lc", "d1", "d2")}, "resumo E%d | Histograma H%d-%d resumo até %d | Painel %s aviso C%d | exemplos %s"
      % (DAV, H1, H2, HR, IR, NAV, {k: (v["P"]["d1"], v["P"]["d2"], v["P"]["lc"], v["IR"], v["aviso"]) for k, v in POS.items()}))
