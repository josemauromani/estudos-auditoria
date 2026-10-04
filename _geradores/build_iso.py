# -*- coding: utf-8 -*-
"""Gera ISO-9001-modelo.xlsx no mesmo padrão visual das planilhas anteriores da série."""
import math
import os
import sys
from datetime import date

_here = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, _here)
# funções de formatação compartilhadas com o gerador do PDCA (bloco inicial do arquivo)
_src = open(os.path.join(_here, "build_pdca.py"), encoding="utf-8").read()
exec(_src[: _src.index("STATUS = [")])

from openpyxl.chart import BarChart, Series  # noqa: E402
from openpyxl.chart.label import DataLabelList  # noqa: E402
from iso_data import A, CHECK, DOCS, EX1, EX2, N, NUMS, P, REQ, SECOES  # noqa: E402

BLUE, TEAL, AMBER, PURPLE, REDC = "2B5C8A", "1E7B73", "A96A12", "7A4A9A", "B0413E"
BLUE_T, TEAL_T, AMBER_T, RED_T = "DCE8F3", "D9EEEB", "F6E8CF", "F5DEDC"
S1, S2, S3 = "2A6FB0", "D19A2E", "B0413E"
COR_SECAO = {4: INK, 5: INK, 6: BLUE, 7: AMBER, 8: AMBER, 9: TEAL, 10: REDC}
YESNO = ["Sim", "Parcial", "Não"]
YESNO_CF = [("Sim", GREEN), ("Parcial", YELLOW), ("Não", RED)]
SITS = [A, P, N]
SIT_CF = [(A, GREEN), (P, YELLOW), (N, RED)]
ST_ACAO = ["Não iniciada", "Em andamento", "Concluída", "Cancelada"]
ST_CF = [("Concluída", GREEN), ("Em andamento", YELLOW), ("Não iniciada", RED)]
PRAZO_CF = [("Concluída", GREEN), ("No prazo", GREEN), ("Sem prazo", YELLOW), ("Atrasada", RED)]
DOC_SIT = ["Existe", "Em elaboração", "Não existe", "Não aplicável"]
DOC_CF = [("Existe", GREEN), ("Em elaboração", YELLOW), ("Não existe", RED)]
TXT = "@"


def note(ws, ref, text):
    c = Comment(text, "Modelo ISO 9001")
    c.width, c.height = 330, 130
    ws[ref].comment = c


def cf_warn(ws, rng_, first, ok_values=("OK",)):
    """Verde para os valores aceitos; amarelo para qualquer outro texto."""
    cond = ",".join(f'{first}="{v}"' for v in ok_values)
    ws.conditional_formatting.add(rng_, FormulaRule(formula=[f"OR({cond})"], fill=PatternFill("solid", bgColor=GREEN, fgColor=GREEN)))
    ws.conditional_formatting.add(rng_, FormulaRule(formula=[f"AND(LEN({first})>0,NOT(OR({cond})))"],
                                  fill=PatternFill("solid", bgColor=YELLOW, fgColor=YELLOW)))


def hint_row(ws, row, cells, height=30):
    for ref, text, merge in cells:
        put(ws, f"{ref}{row}", text, f=font(9, i=True, c=MUTED), bg=GRAY, h="center", merge=merge and f"{ref}{row}:{merge}{row}")
    ws.row_dimensions[row].height = height


def alt(text, chars, minimo=21.75, linha=12.75):
    n = max(1, math.ceil(len(str(text)) / max(chars, 1)))
    return max(minimo, n * linha + 9)


def chart_style(ch):
    ch.style = 2
    ch.title = None
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ch.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill="D3DBE0", w=9525))
    ch.x_axis.graphicalProperties = GraphicalProperties(ln=LineProperties(solidFill="D3DBE0", w=9525))
    ch.y_axis.graphicalProperties = GraphicalProperties(ln=LineProperties(noFill=True))


def stacked(ws, anchor, r1, r2, c_cat, cols, width=24, height=8.5):
    """Barras empilhadas horizontais: uma barra por seção, com as três situações."""
    ch = BarChart()
    ch.type = "bar"
    ch.grouping = "stacked"
    ch.overlap = 100
    ch.gapWidth = 60
    ch.height, ch.width = height, width
    chart_style(ch)
    ch.legend.position = "b"
    for (col, nome), cor in zip(cols, (S1, S2, S3)):
        s = Series(Reference(ws, min_col=col, min_row=r1, max_row=r2), title=nome)
        s.graphicalProperties.solidFill = cor
        s.graphicalProperties.line.solidFill = WHITE
        ch.append(s)
    ch.set_categories(Reference(ws, min_col=c_cat, min_row=r1, max_row=r2))
    ch.y_axis.scaling.min = 0
    ch.y_axis.majorUnit = 5
    ch.x_axis.scaling.orientation = "maxMin"
    ws.add_chart(ch, anchor)


def estagio(pct, nao):
    return (f'=IF({pct}="","Sem avaliação",IF({pct}<0.5,"Inicial",IF({pct}<0.8,"Em implantação",'
            f'IF({nao}>0,"Em implantação: há requisito não atendido","Pronto para a auditoria interna"))))')


def cf_estagio(ws, ref):
    for cond, color in ((f'LEFT({ref},6)="Pronto"', GREEN), (f'{ref}="Inicial"', RED), (f"LEN({ref})>0", YELLOW)):
        ws.conditional_formatting.add(ref, FormulaRule(formula=[cond], stopIfTrue=True, fill=PatternFill("solid", bgColor=color, fgColor=color)))


def bloco_resumo(ws, top, fonte, colsec, colapl, colsit, r1, r2, last, secoes=None, rh=21.75):
    """Tabela por seção, a partir da linha 'top'. 'fonte' é o prefixo da aba com os dados ('' para a própria aba)."""
    K = f"{fonte}${colsec}${r1}:${colsec}${r2}"
    F = f"{fonte}${colsit}${r1}:${colsit}${r2}"
    apl = f'*({fonte}${colapl}${r1}:${colapl}${r2}<>"Não")' if colapl else ""
    for col, text in zip("BCDEFGHI", ["Seção", "Título", "Requisitos", "Avaliados", A, P, N, "Atendimento"]):
        head(ws, f"{col}{top}", text)
    ws.row_dimensions[top].height = 30
    rr = top
    lista = [s for s in SECOES if secoes is None or s[0] in secoes]
    for s, t, _ in lista:
        rr += 1
        put(ws, f"B{rr}", s, f=font(10, True, c=WHITE), bg=COR_SECAO[s], h="center")
        put(ws, f"C{rr}", f"{s} · {t}", f=font(10, True), bg=GRAY)
        calc(ws, f"D{rr}", f'=SUMPRODUCT(({K}={s}){apl})', b=False)
        calc(ws, f"E{rr}", f'=SUMPRODUCT(({K}={s}){apl}*({F}<>""))', b=False)
        for col, sit in zip("FGH", SITS):
            calc(ws, f"{col}{rr}", f'=SUMPRODUCT(({K}={s}){apl}*({F}="{sit}"))', b=False)
        calc(ws, f"I{rr}", f'=IF(E{rr}=0,"",(F{rr}+0.5*G{rr})/E{rr})', fmt="0%")
        ws.row_dimensions[rr].height = rh
    a, b = top + 1, rr
    rr += 1
    put(ws, f"B{rr}", "Total", f=font(10, True, c=WHITE), bg=INK, h="center", merge=f"B{rr}:C{rr}")
    for col in "DEFGH":
        calc(ws, f"{col}{rr}", f"=SUM({col}{a}:{col}{b})")
    calc(ws, f"I{rr}", f'=IF(E{rr}=0,"",(F{rr}+0.5*G{rr})/E{rr})', fmt="0%")
    ws.row_dimensions[rr].height = rh
    ws.conditional_formatting.add(f"I{a}:I{rr}", FormulaRule(formula=[f'AND(ISNUMBER(I{a}),I{a}>=0.8)'], fill=PatternFill("solid", bgColor=GREEN, fgColor=GREEN)))
    ws.conditional_formatting.add(f"I{a}:I{rr}", FormulaRule(formula=[f'AND(ISNUMBER(I{a}),I{a}<0.5)'], fill=PatternFill("solid", bgColor=RED, fgColor=RED)))
    ws.conditional_formatting.add(f"I{a}:I{rr}", FormulaRule(formula=[f'ISNUMBER(I{a})'], fill=PatternFill("solid", bgColor=YELLOW, fgColor=YELLOW)))
    return a, b, rr


# ================================================================== pasta de trabalho
wb = Workbook()
wb.properties.title = "ISO 9001 — Modelo de diagnóstico, plano de ação e informação documentada"
wb.properties.creator = "Modelo ISO 9001"
wb.properties.language = "pt-BR"

# ------------------------------------------------------------------ Instruções
ws = wb.active
ws.title = "Instruções"
widths(ws, {"A": 2, "B": 24, "C": 92, "D": 2})
title(ws, "ISO 9001 — Como usar esta planilha",
      "Modelo para avaliar o atendimento a cada requisito, planejar as ações e conferir a informação documentada.", "C")
r = 4


def section(text):
    global r
    band(ws, r, text, "C", sz=11)
    r += 1


def line(k, v, kbg=GRAY, vbg=None, kf=None, kh="left", height=None):
    global r
    put(ws, f"B{r}", k, f=kf or font(10, True), bg=kbg, h=kh)
    put(ws, f"C{r}", v, bg=vbg)
    ws.row_dimensions[r].height = height or (19.5 if len(v) <= 98 else (31.5 if len(v) <= 196 else 45.75))
    r += 1


section("Legenda: onde preencher")
line("Amarelo-claro", "Células de entrada. É aqui que você digita.", vbg=INPUT)
line("Cinza", "Células calculadas ou fixas (requisitos, perguntas, contagens, avisos). Não altere.", vbg=GRAY)
line("Listas suspensas", "Colunas de aplicabilidade, situação e status aceitam apenas as opções da lista. Datas são conferidas na digitação.")
line("Comentários", "Na aba Diagnóstico, o título de cada requisito traz um comentário com o resumo do que a norma pede. Passe o mouse sobre ele.")
r += 1
section("As três situações")
line(A, "A evidência mostra que o requisito é atendido, de forma consistente. Vale 1 ponto.", kbg=S1, vbg=BLUE_T, kf=font(10, True, c=WHITE))
line(P, "Parte do requisito é atendida, ou ele é atendido em parte dos casos. Vale meio ponto.", kbg=S2, vbg=AMBER_T, kf=font(10, True))
line(N, "Não há evidência de atendimento. Não vale ponto.", kbg=S3, vbg=RED_T, kf=font(10, True, c=WHITE))
r += 1
section("Ordem recomendada de preenchimento")
for k, text in enumerate([
    "Aba Diagnóstico: preencha o cabeçalho e marque como não aplicáveis os requisitos que não valem para a organização, com a justificativa.",
    "Aba Diagnóstico: para cada requisito, leia a pergunta, procure a evidência e registre a situação.",
    "Aba Diagnóstico: escreva a evidência encontrada e, quando houver lacuna, o que falta.",
    "Aba Resumo: leia o percentual por seção e o estágio do sistema.",
    "Aba Plano de ação: escreva uma ação para cada lacuna, com responsável e prazo.",
    "Aba Documentos: confira se a informação documentada exigida pela norma existe.",
    "Aba Checklist: valide o sistema antes da auditoria interna.",
], 1):
    line(f"Passo {k}", text)
r += 1
section("Abas da planilha")
for k, text in [
    ("Diagnóstico", "Os 45 requisitos do estudo, com pergunta, aplicabilidade, situação, evidência e lacuna."),
    ("Resumo", "Resultado por seção, percentual de atendimento, estágio do sistema e gráfico."),
    ("Plano de ação", "Ações para as lacunas, com requisito, responsável, prazo e status."),
    ("Documentos", "A informação documentada exigida pela norma, com o documento ou o registro da organização."),
    ("Checklist", "Doze verificações do sistema, com percentual de conclusão."),
    ("Exemplo 1 - Pizzaria", "Diagnóstico completo de uma organização pequena, nas sete seções."),
    ("Exemplo 2 - Compras", "Diagnóstico de um processo, com os requisitos ligados a ele."),
]:
    line(k, text)
r += 1
section("Regras e premissas do modelo")
for k, text in [
    ("Requisitos", "A lista tem 45 itens. Os requisitos 7.1.1 e 7.1.2 foram reunidos. Os subitens de 4.4, 5.2, 6.1, 6.2, 7.1.5, 7.5, 8.2.3, 8.3, 8.7, 9.2, 9.3 e 10.2 foram tratados em conjunto, no requisito de que fazem parte. As perguntas e os resumos foram escritos com palavras próprias e não substituem o texto da norma."),
    ("Percentual", "Atendimento = (requisitos atendidos + metade dos atendidos em parte) ÷ requisitos avaliados. Os requisitos não aplicáveis ficam fora da conta. A pontuação é uma convenção deste modelo."),
    ("Estágio", "Abaixo de 50%: inicial. De 50% a 79%: em implantação. A partir de 80%, sem requisito não atendido: pronto para a auditoria interna. As faixas são uma convenção deste modelo."),
    ("Não aplicável", "Todo requisito marcado como não aplicável pede justificativa na coluna de evidência. A justificativa também vai para o escopo do sistema."),
    ("Diagnóstico e auditoria", "O diagnóstico é uma autoavaliação. Ele orienta a implantação e não substitui a auditoria interna, que pede auditor independente."),
    ("Exemplos", "Os dados das abas de exemplo são ilustrativos, criados para este material."),
]:
    line(k, text)
setup(ws, INK, f"B1:C{r}", landscape=False, fit_height=True)

# ------------------------------------------------------------------ Diagnóstico
ws = wb.create_sheet("Diagnóstico")
widths(ws, {"A": 2, "B": 10, "C": 36, "D": 40, "E": 12, "F": 17, "G": 44, "H": 40, "I": 9, "J": 24, "K": 8, "L": 2})
title(ws, "Diagnóstico de atendimento à ISO 9001:2015", "Registre a situação de cada requisito, com a evidência encontrada e o que falta.", "J")
for rr, l1, l2, f2 in [(4, "Organização ou processo", "Data do diagnóstico", DATE), (5, "Escopo avaliado", "Responsável", None)]:
    label(ws, f"B{rr}", l1, merge=f"B{rr}:C{rr}")
    inp(ws, f"D{rr}", merge=f"D{rr}:F{rr}")
    label(ws, f"G{rr}", l2)
    inp(ws, f"H{rr}", merge=f"H{rr}:J{rr}", fmt=f2)
    ws.row_dimensions[rr].height = 24
dv_date(ws, "H4")
for col, text in zip("BCDEFGHIJK", ["Requisito", "Título", "Pergunta de diagnóstico", "Aplicável?", "Situação", "Evidência encontrada", "O que falta",
                                     "Pontos", "Conferência", "Seção"]):
    head(ws, f"{col}7", text)
ws.row_dimensions[7].height = 24
hint_row(ws, 8, [("B", "Da norma", None), ("C", "Passe o mouse para ler o resumo", None), ("D", "Responda com fatos", None), ("E", "Sim ou Não", None),
                 ("F", "Escolha na lista", None), ("G", "Documento, registro, local e data", None), ("H", "Só quando houver lacuna", None),
                 ("I", "Calculado", None), ("J", "Calculada", None), ("K", "", None)])
rr = 8
DROWS = {}
for s, t, q in SECOES:
    rr += 1
    put(ws, f"B{rr}", f"Seção {s} · {t}", f=font(10, True, c=WHITE), bg=COR_SECAO[s], box=False, merge=f"B{rr}:J{rr}")
    ws.row_dimensions[rr].height = 21.75
    for q_ in [x for x in REQ if x["secao"] == s]:
        rr += 1
        DROWS[q_["num"]] = rr
        put(ws, f"B{rr}", q_["num"], f=font(10, True), bg=GRAY, h="center", fmt=TXT)
        put(ws, f"C{rr}", q_["titulo"], f=font(10, True), bg=GRAY)
        note(ws, f"C{rr}", "O que pede, em resumo:\n" + q_["pede"])
        put(ws, f"D{rr}", q_["pergunta"], f=font(10, i=True), bg=GRAY)
        inp(ws, f"E{rr}", "Sim", h="center")
        inp(ws, f"F{rr}", h="center")
        inp(ws, f"G{rr}")
        inp(ws, f"H{rr}")
        calc(ws, f"I{rr}", f'=IF(OR(E{rr}="Não",F{rr}=""),"",IF(F{rr}="{A}",1,IF(F{rr}="{P}",0.5,0)))', b=False, fmt="0.0")
        calc(ws, f"J{rr}", f'=IF(E{rr}="Não",IF(G{rr}="","Justifique a não aplicabilidade","Não aplicável"),IF(F{rr}="","Não avaliado",'
             f'IF(G{rr}="","Falta a evidência",IF(AND(F{rr}<>"{A}",H{rr}=""),"Falta a lacuna","Completo"))))', b=False, sz=9)
        put(ws, f"K{rr}", s, f=font(9, c=MUTED), bg=GRAY, h="center")
        ws.row_dimensions[rr].height = max(36, alt(q_["titulo"], 36 * 1.05), alt(q_["pergunta"], 40 * 1.05))
D1, D2 = 9, rr
assert len(DROWS) == 45 and D2 == 60, (len(DROWS), D2)
dv_list(ws, f"E{D1}:E{D2}", ["Sim", "Não"], "O requisito vale para a organização?")
dv_list(ws, f"F{D1}:F{D2}", SITS, "Atende, Atende em parte ou Não atende")
cf_equal(ws, f"F{D1}:F{D2}", SIT_CF)
cf_equal(ws, f"E{D1}:E{D2}", [("Não", "D9DEE2")])
for cond, color in ((f'OR(J{D1}="Completo",J{D1}="Não aplicável")', GREEN), (f'J{D1}="Não avaliado"', "E3E8EB"), (f"LEN(J{D1})>0", YELLOW)):
    ws.conditional_formatting.add(f"J{D1}:J{D2}", FormulaRule(formula=[cond], stopIfTrue=True, fill=PatternFill("solid", bgColor=color, fgColor=color)))
note(ws, "E7", "Marque Não só quando a atividade não existe na organização e a falta do requisito não afeta o produto, o serviço nem o cliente. Escreva a justificativa na coluna de evidência.")
note(ws, "G7", "A evidência precisa poder ser encontrada por outra pessoa: nome do documento ou do registro, local, data e amostra.")
ws.column_dimensions["K"].hidden = True
ws.freeze_panes = "D9"
setup(ws, BLUE, f"B1:J{D2}")
ws.print_title_rows = "7:7"

# ------------------------------------------------------------------ Resumo
ws = wb.create_sheet("Resumo")
widths(ws, {"A": 2, "B": 8, "C": 36, "D": 13, "E": 13, "F": 13, "G": 17, "H": 13, "I": 14, "J": 2})
title(ws, "Resumo do diagnóstico", "Resultado por seção da norma. Esta aba é calculada a partir da aba Diagnóstico.", "I")
label(ws, "B4", "Organização ou processo", merge="B4:C4")
calc(ws, "D4", '=IF(Diagnóstico!D4="","",Diagnóstico!D4)', h="left", b=False, merge="D4:F4")
label(ws, "G4", "Data do diagnóstico")
calc(ws, "H4", '=IF(Diagnóstico!H4="","",Diagnóstico!H4)', h="left", b=False, fmt=DATE, merge="H4:I4")
ws.row_dimensions[4].height = 21.75
a, b, tot = bloco_resumo(ws, 6, "Diagnóstico!", "K", "E", "F", D1, D2, "I")
put(ws, "D6", "Aplicáveis", f=font(10, True, c=WHITE), bg=INK, h="center")
s = tot + 2
band(ws, s, "Resumo automático", "I")
DG = f"Diagnóstico!$J${D1}:$J${D2}"
for k, (text, formula, fmt) in enumerate([
    ("Requisitos não aplicáveis", f'=COUNTIF(Diagnóstico!$E${D1}:$E${D2},"Não")', None),
    ("Requisitos ainda não avaliados", f"=D{tot}-E{tot}", None),
    ("Percentual de atendimento", f"=I{tot}", "0%"),
    ("Seção com o menor atendimento", f'=IF(COUNT(I{a}:I{b})=0,"",INDEX(C{a}:C{b},MATCH(MIN(I{a}:I{b}),I{a}:I{b},0)))', None),
    ("Estágio do sistema", estagio(f"I{tot}", f"H{tot}"), None),
    ("Aviso", f'=IF(E{tot}=0,"Preencha a aba Diagnóstico",IF(COUNTIF({DG},"Justifique*")>0,"Há requisito não aplicável sem justificativa",'
              f'IF(D{tot}-E{tot}>0,"Há requisito sem avaliação: o percentual é parcial",IF(COUNTIF({DG},"Falta*")>0,'
              f'"Há requisito sem evidência ou sem lacuna registrada","OK"))))', None),
], 1):
    label(ws, f"B{s+k}", text, merge=f"B{s+k}:D{s+k}", h="right")
    calc(ws, f"E{s+k}", formula, fmt=fmt, merge=f"E{s+k}:I{s+k}", sz=9 if text == "Aviso" else 10, b=text != "Aviso")
    ws.row_dimensions[s + k].height = 24
cf_estagio(ws, f"E{s+5}")
cf_warn(ws, f"E{s+6}:I{s+6}", f"E{s+6}")
stacked(ws, f"B{s+8}", a, b, 3, [(6, A), (7, P), (8, N)], width=24.5, height=9)
setup(ws, TEAL, f"B1:I{s+27}", fit_height=True)

# ------------------------------------------------------------------ Plano de ação
ws = wb.create_sheet("Plano de ação")
widths(ws, {"A": 2, "B": 5, "C": 11, "D": 32, "E": 36, "F": 40, "G": 20, "H": 13, "I": 15, "J": 13, "K": 14, "L": 2})
title(ws, "Plano de ação do diagnóstico", "Uma ação para cada lacuna. O título e a lacuna vêm da aba Diagnóstico.", "K")
for col, text in zip("BCDEFGHIJK", ["#", "Requisito", "Título", "O que falta", "Ação", "Responsável", "Prazo", "Status", "Concluída em", "Situação"]):
    head(ws, f"{col}4", text)
ws.row_dimensions[4].height = 24
hint_row(ws, 5, [("B", "", None), ("C", "Escolha na lista", None), ("D", "Calculado", None), ("E", "Calculado", None),
                 ("F", "Verbo no infinitivo e resultado esperado", None), ("G", "Cargo ou nome", None), ("H", "Data", None), ("I", "Da ação", None),
                 ("J", "Data", None), ("K", "Calculada", None)])
ex(ws, "B6", "Ex.", h="center")
for col, v, h_, f_ in [("C", "4.3", "center", TXT), ("D", "Determinando o escopo do sistema de gestão da qualidade", "left", None),
                       ("E", "Escrever o escopo: produtos, loja, canais de venda e requisitos não aplicáveis.", "left", None),
                       ("F", "Escrever a declaração de escopo e aprová-la com o dono.", "left", None), ("G", "Gerente da loja", "left", None),
                       ("H", date(2026, 10, 30), "center", DATE), ("I", "Em andamento", "center", None), ("J", "", "center", None),
                       ("K", "No prazo", "center", None)]:
    ex(ws, f"{col}6", v, h=h_, fmt=f_)
ws.row_dimensions[6].height = 36
P1, P2 = 7, 26
DB, DC, DH = (f"Diagnóstico!${c}${D1}:${c}${D2}" for c in "BCH")
for k in range(20):
    rr = P1 + k
    num(ws, f"B{rr}", k + 1)
    inp(ws, f"C{rr}", h="center", fmt=TXT)
    calc(ws, f"D{rr}", f'=IF(C{rr}="","",IF(ISNUMBER(MATCH(C{rr}&"",{DB},0)),INDEX({DC},MATCH(C{rr}&"",{DB},0)),"Requisito não encontrado"))', h="left", b=False, sz=9)
    calc(ws, f"E{rr}", f'=IF(C{rr}="","",IF(ISNUMBER(MATCH(C{rr}&"",{DB},0)),INDEX({DH},MATCH(C{rr}&"",{DB},0))&"",""))', h="left", b=False, sz=9)
    inp(ws, f"F{rr}")
    inp(ws, f"G{rr}")
    inp(ws, f"H{rr}", h="center", fmt=DATE)
    inp(ws, f"I{rr}", h="center")
    inp(ws, f"J{rr}", h="center", fmt=DATE)
    calc(ws, f"K{rr}", f'=IF(F{rr}="","",IF(I{rr}="Concluída","Concluída",IF(I{rr}="Cancelada","Cancelada",'
         f'IF(H{rr}="","Sem prazo",IF(H{rr}<TODAY(),"Atrasada","No prazo")))))', b=False, sz=9)
    ws.row_dimensions[rr].height = 36
dv_list(ws, f"C{P1}:C{P2}", NUMS, "Número do requisito, como na aba Diagnóstico")
dv_list(ws, f"I{P1}:I{P2}", ST_ACAO, "Não iniciada, Em andamento, Concluída ou Cancelada")
dv_date(ws, f"H{P1}:H{P2}")
dv_date(ws, f"J{P1}:J{P2}")
cf_equal(ws, f"I{P1}:I{P2}", ST_CF)
cf_equal(ws, f"K{P1}:K{P2}", PRAZO_CF)
s = P2 + 2
band(ws, s, "Resumo automático", "K")
DE, DF = (f"Diagnóstico!${c}${D1}:${c}${D2}" for c in "EF")
for k, (text, formula, fmt) in enumerate([
    ("Ações registradas", f"=COUNTA(F{P1}:F{P2})", None),
    ("Concluídas", f'=COUNTIF(K{P1}:K{P2},"Concluída")', None),
    ("Atrasadas", f'=COUNTIF(K{P1}:K{P2},"Atrasada")', None),
    ("Percentual concluído", f'=IF(E{s+1}-COUNTIF(K{P1}:K{P2},"Cancelada")<=0,0,E{s+2}/(E{s+1}-COUNTIF(K{P1}:K{P2},"Cancelada")))', "0%"),
    ("Lacunas do diagnóstico", f'=SUMPRODUCT(({DE}<>"Não")*(({DF}="{P}")+({DF}="{N}")))', None),
    ("Lacunas sem ação no plano", f'=SUMPRODUCT(({DE}<>"Não")*(({DF}="{P}")+({DF}="{N}"))*(COUNTIF($C${P1}:$C${P2},{DB})=0))', None),
    ("Aviso", f'=IF(E{s+5}=0,"Preencha a aba Diagnóstico",IF(E{s+1}=0,"Registre as ações",IF(SUMPRODUCT((F{P1}:F{P2}<>"")*(((C{P1}:C{P2}="")+(G{P1}:G{P2}="")+(H{P1}:H{P2}=""))>0))>0,'
              f'"Há ação sem requisito, responsável ou prazo",IF(E{s+6}>0,"Há lacuna sem ação",IF(E{s+3}>0,"Há ação atrasada","OK")))))', None),
], 1):
    label(ws, f"B{s+k}", text, merge=f"B{s+k}:D{s+k}", h="right")
    calc(ws, f"E{s+k}", formula, fmt=fmt, sz=9 if text == "Aviso" else 10, b=text != "Aviso")
    ws.row_dimensions[s + k].height = 30 if text == "Aviso" else 21.75
cf_warn(ws, f"E{s+7}", f"E{s+7}")
ws.freeze_panes = "D5"
setup(ws, AMBER, f"B1:K{s+7}", fit_height=True)

# ------------------------------------------------------------------ Documentos
ws = wb.create_sheet("Documentos")
widths(ws, {"A": 2, "B": 5, "C": 11, "D": 16, "E": 52, "F": 40, "G": 30, "H": 16, "I": 2})
title(ws, "Informação documentada exigida pela norma", "Informe o documento ou o registro que atende a cada item, e onde ele fica.", "H")
for col, text in zip("BCDEFGH", ["#", "Requisito", "Verbo", "O que precisa existir", "Documento ou registro da organização", "Onde fica, ou quem responde",
                                  "Situação"]):
    head(ws, f"{col}4", text)
ws.row_dimensions[4].height = 24
ex(ws, "B5", "Ex.", h="center")
for col, v, h_, f_ in [("C", "7.2", "center", TXT), ("D", "Reter", "center", None), ("E", "Evidência de competência", "left", None),
                       ("F", "Lista de presença e ficha de treinamento", "left", None), ("G", "Pasta de pessoal, com o gerente", "left", None),
                       ("H", "Existe", "center", None)]:
    ex(ws, f"{col}5", v, h=h_, fmt=f_)
ws.row_dimensions[5].height = 24
O1 = 6
for k, (req, tipo, texto, exemplo) in enumerate(DOCS):
    rr = O1 + k
    num(ws, f"B{rr}", k + 1)
    put(ws, f"C{rr}", req, f=font(10, True), bg=GRAY, h="center", fmt=TXT)
    put(ws, f"D{rr}", tipo, f=font(10, True), bg=BLUE_T if tipo == "Manter" else (TEAL_T if tipo == "Reter" else AMBER_T), h="center")
    put(ws, f"E{rr}", texto, bg=GRAY)
    note(ws, f"E{rr}", "Exemplo de como atender:\n" + exemplo)
    inp(ws, f"F{rr}")
    inp(ws, f"G{rr}")
    inp(ws, f"H{rr}", h="center")
    ws.row_dimensions[rr].height = max(30, alt(texto, 52 * 1.05))
O2 = O1 + len(DOCS) - 1
dv_list(ws, f"H{O1}:H{O2}", DOC_SIT, "Existe, Em elaboração, Não existe ou Não aplicável")
cf_equal(ws, f"H{O1}:H{O2}", DOC_CF + [("Não aplicável", "D9DEE2")])
s = O2 + 2
band(ws, s, "Resumo automático", "H")
H_ = f"H{O1}:H{O2}"
for k, (text, formula, fmt) in enumerate([
    ("Itens da lista", f"=ROWS({H_})", None),
    ("Existem", f'=COUNTIF({H_},"Existe")', None),
    ("Em elaboração", f'=COUNTIF({H_},"Em elaboração")', None),
    ("Não existem", f'=COUNTIF({H_},"Não existe")', None),
    ("Não aplicáveis", f'=COUNTIF({H_},"Não aplicável")', None),
    ("Percentual existente", f"=IF(F{s+1}-F{s+5}<=0,0,F{s+2}/(F{s+1}-F{s+5}))", "0%"),
    ("Aviso", f'=IF(COUNTA({H_})=0,"Informe a situação de cada item",IF(COUNTBLANK({H_})>0,"Há item sem situação",'
              f'IF(SUMPRODUCT(({H_}="Existe")*(F{O1}:F{O2}=""))>0,"Há item existente sem o nome do documento",IF(F{s+4}>0,"Há informação documentada exigida que não existe","OK"))))', None),
], 1):
    label(ws, f"B{s+k}", text, merge=f"B{s+k}:E{s+k}", h="right")
    calc(ws, f"F{s+k}", formula, fmt=fmt, sz=9 if text == "Aviso" else 10, b=text != "Aviso")
    ws.row_dimensions[s + k].height = 21.75
cf_warn(ws, f"F{s+7}", f"F{s+7}")
ws.freeze_panes = "F5"
setup(ws, PURPLE, f"B1:H{s+7}")
ws.print_title_rows = "4:4"

# ------------------------------------------------------------------ Checklist
ws = wb.create_sheet("Checklist")
widths(ws, {"A": 2, "B": 5, "C": 70, "D": 14, "E": 44, "F": 2})
title(ws, "Checklist de validação do sistema", "Escolha o status de cada item na lista suspensa (coluna em amarelo-claro).", "E")
for col, text in zip("BCDE", ["#", "Item de verificação", "Status", "Observações"]):
    head(ws, f"{col}4", text)
ws.row_dimensions[4].height = 21.75
for k, text in enumerate(CHECK):
    rr = 5 + k
    num(ws, f"B{rr}", k + 1)
    put(ws, f"C{rr}", text, bg=GRAY)
    inp(ws, f"D{rr}", h="center")
    inp(ws, f"E{rr}")
    ws.row_dimensions[rr].height = 30
dv_list(ws, "D5:D16", YESNO, "Sim, Parcial ou Não")
cf_equal(ws, "D5:D16", YESNO_CF)
band(ws, 18, "Resumo automático", "E")
summary(ws, 19, "Itens atendidos (Sim)", '=COUNTIF(D5:D16,"Sim")', "C", "D")
summary(ws, 20, "Itens parciais", '=COUNTIF(D5:D16,"Parcial")', "C", "D")
summary(ws, 21, "Itens não atendidos", '=COUNTIF(D5:D16,"Não")', "C", "D")
summary(ws, 22, "Itens sem resposta", "=COUNTBLANK(D5:D16)", "C", "D")
summary(ws, 23, "Percentual atendido", "=D19/ROWS(D5:D16)", "C", "D", fmt="0%")
summary(ws, 24, "Situação", '=IF(D19=ROWS(D5:D16),"Validado",IF(D22>0,"Em andamento","Com pendências"))', "C", "D")
setup(ws, PURPLE, "B1:E24", landscape=False, fit_height=True)


# ------------------------------------------------------------------ Exemplos
def example(ws, data, subtitulo, com_estagio=True):
    widths(ws, {"A": 2, "B": 10, "C": 38, "D": 17, "E": 17, "F": 17, "G": 20, "H": 17, "I": 16, "J": 8, "K": 2})
    title(ws, "Diagnóstico de atendimento à ISO 9001:2015", subtitulo, "I")
    H = data["head"]
    rr = 4
    for rot, val, fmt in [("Organização", H["org"], None), ("Data", H["data"], DATE), ("Responsável", H["por"], None), ("Escopo avaliado", H["escopo"], None),
                          ("Critério", H["criterio"], None), ("Método", H["metodo"], None)]:
        label(ws, f"B{rr}", rot, merge=f"B{rr}:C{rr}")
        inp(ws, f"D{rr}", val, merge=f"D{rr}:I{rr}", bg=WHITE, fmt=fmt)
        ws.row_dimensions[rr].height = 21.75
        rr += 1
    rr += 1
    hd = rr
    put(ws, f"B{rr}", "Requisito", f=font(10, True, c=WHITE), bg=INK, h="center")
    put(ws, f"C{rr}", "Título", f=font(10, True, c=WHITE), bg=INK, h="center")
    put(ws, f"D{rr}", "Situação", f=font(10, True, c=WHITE), bg=INK, h="center")
    put(ws, f"E{rr}", "Evidência encontrada", f=font(10, True, c=WHITE), bg=INK, h="center", merge=f"E{rr}:F{rr}")
    put(ws, f"G{rr}", "O que falta", f=font(10, True, c=WHITE), bg=INK, h="center", merge=f"G{rr}:H{rr}")
    put(ws, f"I{rr}", "Pontos", f=font(10, True, c=WHITE), bg=INK, h="center")
    put(ws, f"J{rr}", "Seção", f=font(10, True, c=WHITE), bg=INK, h="center")
    ws.row_dimensions[rr].height = 24
    r1 = rr + 1
    secs = []
    for s, t, _ in SECOES:
        itens = [x for x in REQ if x["secao"] == s and x["num"] in data["itens"]]
        if not itens:
            continue
        secs.append(s)
        rr += 1
        put(ws, f"B{rr}", f"Seção {s} · {t}", f=font(10, True, c=WHITE), bg=COR_SECAO[s], box=False, merge=f"B{rr}:I{rr}")
        ws.row_dimensions[rr].height = 21.75
        for q_ in itens:
            sit, ev, falta = data["itens"][q_["num"]]
            rr += 1
            put(ws, f"B{rr}", q_["num"], f=font(10, True), bg=GRAY, h="center", fmt=TXT)
            put(ws, f"C{rr}", q_["titulo"], f=font(10, True), bg=GRAY)
            put(ws, f"D{rr}", sit, h="center")
            put(ws, f"E{rr}", ev, merge=f"E{rr}:F{rr}")
            put(ws, f"G{rr}", falta, merge=f"G{rr}:H{rr}")
            calc(ws, f"I{rr}", f'=IF(D{rr}="{A}",1,IF(D{rr}="{P}",0.5,0))', b=False, fmt="0.0")
            put(ws, f"J{rr}", s, f=font(9, c=MUTED), bg=GRAY, h="center")
            ws.row_dimensions[rr].height = max(30, alt(q_["titulo"], 38 * 1.05), alt(ev, 34 * 1.1), alt(falta, 37 * 1.1))
    r2 = rr
    cf_equal(ws, f"D{r1}:D{r2}", SIT_CF)
    ws.column_dimensions["J"].hidden = True
    rr += 2
    ws.row_breaks.append(Break(id=rr - 1))
    band(ws, rr, "Resultado por seção", "I")
    a, b, tot = bloco_resumo(ws, rr + 1, "", "J", None, "D", r1, r2, "I", secoes=secs, rh=19.5)
    ws.row_dimensions[rr + 1].height = 24
    put(ws, f"D{rr+1}", "Requisitos", f=font(10, True, c=WHITE), bg=INK, h="center")
    s0 = tot + 2
    band(ws, s0, "Resumo automático", "I")
    for k, (text, formula, fmt) in enumerate([
        ("Percentual de atendimento", f"=I{tot}", "0%"),
        ("Seção com o menor atendimento", f"=INDEX(C{a}:C{b},MATCH(MIN(I{a}:I{b}),I{a}:I{b},0))", None),
        ("Estágio do sistema", estagio(f"I{tot}", f"H{tot}"), None) if com_estagio else
        ("Requisitos com lacuna", f"=G{tot}+H{tot}", None),
    ], 1):
        label(ws, f"B{s0+k}", text, merge=f"B{s0+k}:D{s0+k}", h="right")
        calc(ws, f"E{s0+k}", formula, fmt=fmt, merge=f"E{s0+k}:I{s0+k}")
        ws.row_dimensions[s0 + k].height = 21
    if com_estagio:
        cf_estagio(ws, f"E{s0+3}")
    stacked(ws, f"B{s0+5}", a, b, 3, [(6, A), (7, P), (8, N)], width=26, height=5.2)
    for k in range(4, 17):
        ws.row_dimensions[s0 + k].height = 12.75
    setup(ws, MUTED, f"B1:I{s0+16}")
    ws.print_title_rows = f"{hd}:{hd}"


example(wb.create_sheet("Exemplo 1 - Pizzaria"), EX1, "Exemplo preenchido, para consulta. Diagnóstico completo de uma organização pequena.")
example(wb.create_sheet("Exemplo 2 - Compras"), EX2, "Exemplo preenchido, para consulta. Requisitos da norma ligados a um processo.", com_estagio=False)

wb.active = 1
wb.save(OUT)
print("ok", OUT)
