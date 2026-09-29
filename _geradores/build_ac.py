# -*- coding: utf-8 -*-
"""Gera Analise-Critica-modelo.xlsx no mesmo padrão visual das planilhas anteriores da série."""
import math
import os
import sys
from datetime import date

_here = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, _here)
# funções de formatação compartilhadas com o gerador do PDCA (bloco inicial do arquivo)
_src = open(os.path.join(_here, "build_pdca.py"), encoding="utf-8").read()
exec(_src[: _src.index("STATUS = [")])

from openpyxl.chart import BarChart, Series  # noqa: E402
from openpyxl.chart.label import DataLabelList  # noqa: E402
from ac_data import (A_ATRASO, A_PRAZO, A_VENCIDA, ATE, CHECK, CONC, CRI, CRITERIOS, ENTRADAS, EX1, EX2, FAV, MEL, MUD, NAO, NOME, PARCIAL, PIOR,  # noqa: E402
                     REC, SIM, SIT_ACAO, SITS, STATUS, TENDS, TIPOS)

BLUE, TEAL, AMBER, PURPLE, REDC = "2B5C8A", "1E7B73", "A96A12", "7A4A9A", "B0413E"
BLUE_T, AMBER_T, RED_T = "DCE8F3", "F8EBCB", "F5DEDC"
SIT_CF = [(FAV, BLUE_T), (ATE, AMBER_T), (CRI, RED_T)]
YESNO = [SIM, PARCIAL, NAO]
YESNO_CF = [(SIM, GREEN), (PARCIAL, YELLOW), (NAO, RED)]
ACAO_CF = [(A_PRAZO, GREEN), (A_ATRASO, YELLOW), (A_VENCIDA, RED)]
PRESENCA = ["Presente", "Ausente", "Justificado"]
ITENS = [e[0] for e in ENTRADAS]
NPART, NACAO, NSAI = 10, 15, 15
MOEDA = '"R$" #,##0'
# linhas fixas, usadas nas fórmulas entre abas
P1, P2 = 16, 16 + NPART - 1          # participantes, na aba Reunião
C1 = P2 + 4                           # critérios da avaliação do sistema
A1, A2 = 9, 9 + NACAO - 1            # aba Ações anteriores
E1, E2 = 6, 6 + len(ENTRADAS) - 1    # aba Entradas
S1, S2 = 7, 7 + NSAI - 1             # aba Saídas
ANT = "'Ações anteriores'"


def note(ws, ref, text):
    c = Comment(text, "Modelo Análise crítica")
    c.width, c.height = 300, 120
    ws[ref].comment = c


def cf_warn(ws, rng_, first, ok_values=("OK",)):
    cond = ",".join(f'{first}="{v}"' for v in ok_values)
    ws.conditional_formatting.add(rng_, FormulaRule(formula=[f"OR({cond})"], fill=PatternFill("solid", bgColor=GREEN, fgColor=GREEN)))
    ws.conditional_formatting.add(rng_, FormulaRule(formula=[f"AND(LEN({first})>0,NOT(OR({cond})))"],
                                  fill=PatternFill("solid", bgColor=YELLOW, fgColor=YELLOW)))


def cf_falta(ws, rng_, first, ok="Completa", neutro=()):
    """Verde para a conferência completa, amarelo para o que falta, sem cor para os textos neutros."""
    ws.conditional_formatting.add(rng_, FormulaRule(formula=[f'{first}="{ok}"'], stopIfTrue=True, fill=PatternFill("solid", bgColor=GREEN, fgColor=GREEN)))
    for t in neutro:
        ws.conditional_formatting.add(rng_, FormulaRule(formula=[f'{first}="{t}"'], stopIfTrue=True, fill=PatternFill("solid", bgColor=GRAY, fgColor=GRAY)))
    ws.conditional_formatting.add(rng_, FormulaRule(formula=[f"LEN({first})>0"], fill=PatternFill("solid", bgColor=YELLOW, fgColor=YELLOW)))


def hint_row(ws, row, cells, height=30):
    for ref, text, merge in cells:
        put(ws, f"{ref}{row}", text, f=font(9, i=True, c=MUTED), bg=GRAY, h="center", merge=merge and f"{ref}{row}:{merge}{row}")
    ws.row_dimensions[row].height = height


def alt(pares, minimo=21.75, linha=12):
    """Altura da linha para o maior texto: pares de (texto, largura da coluna em caracteres)."""
    n = max(max(1, math.ceil(len(str(t)) / max(c, 1))) for t, c in pares)
    return max(minimo, n * linha + 7)


def chart_style(ch):
    ch.style = 2
    ch.title = None
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ch.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill="D3DBE0", w=9525))
    ch.x_axis.graphicalProperties = GraphicalProperties(ln=LineProperties(solidFill="D3DBE0", w=9525))
    ch.y_axis.graphicalProperties = GraphicalProperties(ln=LineProperties(noFill=True))


def barras(ws, anchor, r1, r2, c_cat, c_val, width=16, height=6.5):
    """Barras horizontais: número de ações em cada situação."""
    ch = BarChart()
    ch.type = "bar"
    ch.gapWidth = 60
    ch.height, ch.width = height, width
    chart_style(ch)
    ch.legend = None
    s = Series(Reference(ws, min_col=c_val, min_row=r1, max_row=r2), title="Ações")
    s.graphicalProperties.solidFill = BLUE
    s.graphicalProperties.line.noFill = True
    s.dLbls = DataLabelList()
    s.dLbls.showVal = True
    s.dLbls.showSerName = s.dLbls.showCatName = s.dLbls.showLegendKey = False
    ch.append(s)
    ch.set_categories(Reference(ws, min_col=c_cat, min_row=r1, max_row=r2))
    ch.y_axis.scaling.min = 0
    ch.y_axis.majorUnit = 1
    ch.y_axis.number_format = "0"
    ch.x_axis.scaling.orientation = "maxMin"
    ws.add_chart(ch, anchor)


def f_situacao(r, ref, c_acao="D", c_prazo="F", c_status="G", c_feito="H"):
    return (f'=IF({c_acao}{r}="","",IF({c_prazo}{r}="","Falta o prazo",IF({c_status}{r}="","Falta o status",IF({c_status}{r}="Cancelada","Cancelada",'
            f'IF({c_status}{r}="{CONC}",IF({c_feito}{r}="","Falta a data de conclusão",IF({c_feito}{r}<={c_prazo}{r},"{A_PRAZO}","{A_ATRASO}")),'
            f'IF({c_prazo}{r}<{ref},"{A_VENCIDA}","No prazo"))))))')


def f_conf_entrada(r, res="G", tend="H", sit="I", conc="J", dec="K", primeira="E"):
    return (f'=IF(COUNTA({primeira}{r}:{conc}{r})=0,"Não analisada",IF({res}{r}="","Falta o resultado",IF({tend}{r}="","Falta a tendência",IF({sit}{r}="","Falta a situação",'
            f'IF({conc}{r}="","Falta a conclusão",IF(AND({sit}{r}="{CRI}",{dec}{r}=0),"Falta a decisão",IF(AND({sit}{r}="{ATE}",{dec}{r}=0),"Atenção sem decisão","Completa")))))))')


# ================================================================== pasta de trabalho
wb = Workbook()
wb.properties.title = "Análise crítica pela direção — Modelo"
wb.properties.creator = "Modelo Análise crítica"
wb.properties.language = "pt-BR"

# ------------------------------------------------------------------ Instruções
ws = wb.active
ws.title = "Instruções"
widths(ws, {"A": 2, "B": 24, "C": 92, "D": 2})
title(ws, "Análise crítica pela direção — Como usar esta planilha",
      "Modelo para preparar a pauta, registrar as entradas e as decisões, e acompanhar as ações das análises anteriores.", "C")
r = 4


def section(text):
    global r
    band(ws, r, text, "C", sz=11)
    r += 1


def line(k, v, kbg=GRAY, vbg=None, kf=None, kh="left", height=None):
    global r
    put(ws, f"B{r}", k, f=kf or font(10, True), bg=kbg, h=kh)
    put(ws, f"C{r}", v, bg=vbg)
    ws.row_dimensions[r].height = height or (19.5 if len(v) <= 98 else (31.5 if len(v) <= 196 else 45.75))
    r += 1


section("Legenda: onde preencher")
line("Amarelo-claro", "Células de entrada. É aqui que você digita.", vbg=INPUT)
line("Cinza", "Células calculadas ou fixas (rótulos, entradas da norma, contagens, avisos). Não altere.", vbg=GRAY)
line("Cinza em itálico", "Linha de exemplo (marcada com “Ex.”). Mostra o formato esperado e não entra nas contagens.", vbg=GRAY)
line("Listas suspensas", "Colunas de tendência, situação, tipo, status e avaliação aceitam apenas as opções da lista.")
line("Comentários", "Células com triângulo vermelho no canto trazem uma dica de preenchimento. Passe o mouse sobre elas.")
r += 1
section("As três situações de uma entrada")
line(FAV, "O resultado atende ao planejado, e a tendência não preocupa. A direção registra a conclusão.", kbg=BLUE, vbg=BLUE_T, kf=font(10, True, c=WHITE))
line(ATE, "O resultado não atende em parte, ou a tendência é de piora. A direção decide, ou registra por que não vai agir agora.", kbg="D19A2E", vbg=AMBER_T,
     kf=font(10, True))
line(CRI, "O resultado compromete um objetivo, um cliente ou a operação. A direção decide na reunião.", kbg=REDC, vbg=RED_T, kf=font(10, True, c=WHITE))
r += 1
section("Ordem recomendada de preenchimento")
for k, text in enumerate([
    "Aba Reunião: escreva a data, o período analisado e os participantes.",
    "Aba Ações anteriores: liste as ações decididas nas análises anteriores, com prazo e status.",
    "Aba Entradas: escreva o resultado, a tendência e a situação de cada entrada. Use “nada a relatar” quando for o caso.",
    "Aba Entradas: na reunião, escreva a conclusão da direção sobre cada entrada.",
    "Aba Saídas: registre cada decisão, com a entrada de origem, o tipo, o responsável, o prazo e o recurso.",
    "Aba Reunião: responda às quatro perguntas da avaliação do sistema, com justificativa.",
    "Aba Checklist: valide a análise crítica.",
], 1):
    line(f"Passo {k}", text)
r += 1
section("Abas da planilha")
for k, text in [
    ("Reunião", "Dados da reunião, até 10 participantes, avaliação do sistema e situação da ata."),
    ("Ações anteriores", "Até 15 ações das análises anteriores, com a situação de cada uma na data da reunião, e gráfico."),
    ("Entradas", "As doze entradas, com resultado, tendência, situação, conclusão e número de decisões."),
    ("Saídas", "Até 15 decisões, com conferência dos campos, contagem por tipo e valor total."),
    ("Checklist", "Doze verificações de qualidade da análise crítica, com percentual de conclusão."),
    ("Exemplo 1 - Pizzaria", "A primeira análise crítica de uma loja: entradas, decisões e avaliação do sistema."),
    ("Exemplo 2 - Indústria", "Uma análise crítica semestral, com as ações das análises anteriores."),
]:
    line(k, text)
r += 1
section("Regras e premissas do modelo")
for k, text in [
    ("As doze entradas", "São os itens do requisito 9.3.2 da ISO 9001:2015, resumidos com palavras próprias. O item “c” tem sete partes, de c1 a c7."),
    ("Situação da entrada", "As três situações são uma convenção deste modelo. A norma não classifica as entradas."),
    ("Entrada crítica", "Pede pelo menos uma decisão na aba Saídas. Sem ela, a conferência mostra “Falta a decisão”."),
    ("Entrada em atenção", "Pode ficar sem decisão, e o motivo vai na conclusão. A conferência mostra “Atenção sem decisão”, como lembrete."),
    ("Situação da ação", "É calculada pela data da reunião, informada na aba Reunião. Se a data estiver em branco, a planilha usa a data de hoje."),
    ("Primeira análise", "Na primeira análise crítica, a aba Ações anteriores fica em branco, e o aviso dela informa que não há ações."),
    ("Exemplos", "Os dados das abas de exemplo são ilustrativos, criados para este material."),
]:
    line(k, text)
setup(ws, INK, f"B1:C{r}", landscape=False, fit_height=True)

# ------------------------------------------------------------------ Reunião
ws = wb.create_sheet("Reunião")
widths(ws, {"A": 2, "B": 5, "C": 28, "D": 26, "E": 15, "F": 16, "G": 44, "H": 2})
title(ws, "Reunião de análise crítica", "Dados da reunião, participantes e avaliação do sistema. A situação da ata é calculada no fim da página.", "G")
for k, (rot, fmt) in enumerate([("Organização", None), ("Data da reunião", DATE), ("Período analisado", None), ("Local", None), ("Quem conduz", None),
                                ("Quem redige a ata", None), ("Data da análise anterior", DATE), ("Data prevista da próxima análise", DATE)]):
    rr = 4 + k
    label(ws, f"B{rr}", rot, merge=f"B{rr}:C{rr}")
    inp(ws, f"D{rr}", merge=f"D{rr}:G{rr}", fmt=fmt, h="left")
    ws.row_dimensions[rr].height = 21.75
    if fmt:
        dv_date(ws, f"D{rr}")
note(ws, "B8", "A análise crítica é conduzida pela alta direção: dono, sócio, diretor ou quem dirige a parte da empresa coberta pelo sistema.")
note(ws, "B10", "Deixe em branco na primeira análise crítica.")
band(ws, 13, "Participantes", "G")
for col, text in zip("BCDEFG", ["#", "Nome", "Função", "Alta direção?", "Presença", "Observação"]):
    head(ws, f"{col}14", text)
ws.row_dimensions[14].height = 21.75
ex(ws, "B15", "Ex.", h="center")
for col, v, h_ in [("C", "Maria Souza", "left"), ("D", "Diretora geral", "left"), ("E", SIM, "center"), ("F", "Presente", "center"), ("G", "Conduz a reunião", "left")]:
    ex(ws, f"{col}15", v, h=h_)
ws.row_dimensions[15].height = 21.75
assert P1 == 16
for k in range(NPART):
    rr = P1 + k
    num(ws, f"B{rr}", k + 1)
    inp(ws, f"C{rr}")
    inp(ws, f"D{rr}")
    inp(ws, f"E{rr}", h="center")
    inp(ws, f"F{rr}", h="center")
    inp(ws, f"G{rr}")
    ws.row_dimensions[rr].height = 21.75
dv_list(ws, f"E{P1}:E{P2}", [SIM, NAO], "Sim ou Não")
dv_list(ws, f"F{P1}:F{P2}", PRESENCA, "Presente, Ausente ou Justificado")
band(ws, C1 - 2, "Avaliação do sistema", "G")
put(ws, f"B{C1-1}", "Critério", f=font(10, True, c=WHITE), bg=INK, h="center", merge=f"B{C1-1}:C{C1-1}")
put(ws, f"D{C1-1}", "Pergunta", f=font(10, True, c=WHITE), bg=INK, h="center", merge=f"D{C1-1}:E{C1-1}")
head(ws, f"F{C1-1}", "Resposta")
head(ws, f"G{C1-1}", "Justificativa")
ws.row_dimensions[C1 - 1].height = 21.75
for k, (nome, perg, desc) in enumerate(CRITERIOS):
    rr = C1 + k
    label(ws, f"B{rr}", nome, merge=f"B{rr}:C{rr}")
    put(ws, f"D{rr}", perg, bg=GRAY, merge=f"D{rr}:E{rr}")
    inp(ws, f"F{rr}", h="center")
    inp(ws, f"G{rr}")
    note(ws, f"B{rr}", desc)
    ws.row_dimensions[rr].height = 33
C2 = C1 + 3
dv_list(ws, f"F{C1}:F{C2}", YESNO, "Sim, Parcial ou Não")
cf_equal(ws, f"F{C1}:F{C2}", YESNO_CF)
s = C2 + 2
band(ws, s, "Resumo automático", "G")
# posições dos resumos das outras abas (definidas adiante com as mesmas contas)
ENT_S = E2 + 2   # faixa do resumo na aba Entradas
SAI_S = S2 + 2   # faixa do resumo na aba Saídas
ANT_S = A2 + 2   # faixa do resumo na aba Ações anteriores
for k, (text, formula) in enumerate([
    ("Participantes presentes", f'=COUNTIF(F{P1}:F{P2},"Presente")'),
    ("Alta direção presente", f'=COUNTIFS(E{P1}:E{P2},"{SIM}",F{P1}:F{P2},"Presente")'),
    ("Meses desde a análise anterior", '=IF(OR(D5="",D10=""),"",ROUND((D5-D10)/30.4375,0))'),
    ("Entradas analisadas, de 12", f"=Entradas!E{ENT_S+1}"),
    ("Decisões registradas", f"=Saídas!E{SAI_S+1}"),
    ("Ações anteriores atrasadas", f"={ANT}!E{ANT_S+4}"),
    ("Situação da ata", f'=IF(D5="","Informe a data da reunião",IF(E{s+2}=0,"Falta a presença da alta direção",IF(E{s+4}<12,"Há entrada sem análise: veja a aba Entradas",'
                        f'IF(Entradas!E{ENT_S+6}>0,"Há entrada crítica sem decisão",IF(COUNTIF(Saídas!K{S1}:K{S2},"Falta*")+COUNTIF(Saídas!K{S1}:K{S2},"Prazo*")>0,'
                        f'"Há decisão com o registro incompleto",IF(COUNTA(F{C1}:F{C2})<4,"Falta a avaliação do sistema",IF(COUNTA(G{C1}:G{C2})<4,"Falta a justificativa da avaliação",'
                        f'IF(D11="","Marque a data da próxima análise",IF(D11<=D5,"A próxima análise tem data anterior à reunião","OK")))))))))'),
], 1):
    label(ws, f"B{s+k}", text, merge=f"B{s+k}:D{s+k}", h="right")
    calc(ws, f"E{s+k}", formula, merge=f"E{s+k}:G{s+k}", sz=9 if k == 7 else 10, b=k != 7)
    ws.row_dimensions[s + k].height = 21.75
cf_warn(ws, f"E{s+7}:G{s+7}", f"E{s+7}")
setup(ws, BLUE, f"B1:G{s+7}", landscape=False, fit_height=True)

# ------------------------------------------------------------------ Ações anteriores
ws = wb.create_sheet("Ações anteriores")
widths(ws, {"A": 2, "B": 5, "C": 14, "D": 46, "E": 22, "F": 13, "G": 16, "H": 14, "I": 11, "J": 38, "K": 24, "L": 2})
title(ws, "Ações das análises anteriores", "Entrada “a”. Uma ação por linha, com o que foi decidido nas análises anteriores. Na primeira análise, deixe em branco.", "K")
label(ws, "B4", "Data de referência", merge="B4:C4")
calc(ws, "D4", '=IF(Reunião!D5="","",Reunião!D5)', fmt=DATE, h="left")
put(ws, "E4", "É a data da reunião. Se ela estiver em branco, a planilha usa a data de hoje.", f=font(9, i=True, c=MUTED), bg=GRAY, merge="E4:K4")
ws.row_dimensions[4].height = 21.75
for col, text in zip("BCDEFGHIJK", ["#", "Data da análise", "Ação decidida", "Responsável", "Prazo", "Status", "Concluída em", "Eficaz?", "Resultado ou observação",
                                     "Situação"]):
    head(ws, f"{col}6", text)
ws.row_dimensions[6].height = 24
hint_row(ws, 7, [("B", "", None), ("C", "Da reunião que decidiu", None), ("D", "Como está na ata", None), ("E", "Nome ou cargo", None), ("F", "Combinado na ata", None),
                 ("G", "Escolha na lista", None), ("H", "Só para as concluídas", None), ("I", "Funcionou?", None), ("J", "O que mudou depois da ação", None),
                 ("K", "Calculada", None)])
ex(ws, "B8", "Ex.", h="center")
for col, v, h_, fmt in [("C", date(2026, 2, 12), "center", DATE), ("D", "Contratar um analista para o Laboratório.", "left", None), ("E", "Gerente de RH", "left", None),
                        ("F", date(2026, 4, 30), "center", DATE), ("G", CONC, "center", None), ("H", date(2026, 5, 15), "center", DATE), ("I", SIM, "center", None),
                        ("J", "Admissão em 15/05.", "left", None), ("K", A_ATRASO, "center", None)]:
    ex(ws, f"{col}8", v, h=h_, fmt=fmt)
ws.row_dimensions[8].height = 24
for k in range(NACAO):
    rr = A1 + k
    num(ws, f"B{rr}", k + 1)
    inp(ws, f"C{rr}", h="center", fmt=DATE)
    inp(ws, f"D{rr}")
    inp(ws, f"E{rr}")
    inp(ws, f"F{rr}", h="center", fmt=DATE)
    inp(ws, f"G{rr}", h="center")
    inp(ws, f"H{rr}", h="center", fmt=DATE)
    inp(ws, f"I{rr}", h="center")
    inp(ws, f"J{rr}")
    calc(ws, f"K{rr}", f_situacao(rr, 'IF($D$4="",TODAY(),$D$4)'), b=False, sz=9)
    ws.row_dimensions[rr].height = 30
for col in "CFH":
    dv_date(ws, f"{col}{A1}:{col}{A2}")
dv_list(ws, f"G{A1}:G{A2}", STATUS, "Concluída, Em andamento, Não iniciada ou Cancelada")
dv_list(ws, f"I{A1}:I{A2}", YESNO, "Sim, Parcial ou Não")
cf_equal(ws, f"I{A1}:I{A2}", YESNO_CF)
cf_equal(ws, f"K{A1}:K{A2}", ACAO_CF)
ws.conditional_formatting.add(f"K{A1}:K{A2}", FormulaRule(formula=[f'LEFT(K{A1},5)="Falta"'], fill=PatternFill("solid", bgColor=YELLOW, fgColor=YELLOW)))
note(ws, "I6", "Concluir não é o mesmo que resolver. Diga se a ação trouxe o resultado esperado.")
s = A2 + 2
assert s == ANT_S
band(ws, s, "Resumo automático", "F")
put(ws, f"H{s}", "Ações por situação", f=font(10, True, c=WHITE), bg=INK, box=False, merge=f"H{s}:K{s}")
for k, (text, formula, fmt) in enumerate([
    ("Ações registradas", f"=COUNTA(D{A1}:D{A2})", None),
    ("Concluídas", f'=COUNTIF(K{A1}:K{A2},"{CONC}*")', None),
    ("Percentual concluído, sem as canceladas", f'=IF(E{s+1}-COUNTIF(K{A1}:K{A2},"Cancelada")=0,"",E{s+2}/(E{s+1}-COUNTIF(K{A1}:K{A2},"Cancelada")))', "0%"),
    ("Atrasadas", f'=COUNTIF(K{A1}:K{A2},"{A_VENCIDA}")', None),
    ("Concluídas e eficazes", f'=COUNTIFS(G{A1}:G{A2},"{CONC}",I{A1}:I{A2},"{SIM}")', None),
    ("Aviso", f'=IF(E{s+1}=0,"Sem ações anteriores",IF(COUNTIF(K{A1}:K{A2},"Falta*")>0,"Há ação com o registro incompleto: veja a coluna Situação",'
              f'IF(COUNTIFS(G{A1}:G{A2},"{CONC}",I{A1}:I{A2},"")>0,"Há ação concluída sem avaliação da eficácia",'
              f'IF(E{s+4}>0,"Há ação atrasada: leve à reunião","OK"))))', None),
], 1):
    label(ws, f"B{s+k}", text, merge=f"B{s+k}:D{s+k}", h="right")
    calc(ws, f"E{s+k}", formula, fmt=fmt, merge=f"E{s+k}:F{s+k}", sz=9 if text == "Aviso" else 10, b=text != "Aviso")
    ws.row_dimensions[s + k].height = 21.75
cf_warn(ws, f"E{s+6}:F{s+6}", f"E{s+6}")
for k, t in enumerate(SIT_ACAO, 1):
    label(ws, f"H{s+k}", t, merge=f"H{s+k}:J{s+k}")
    calc(ws, f"K{s+k}", f'=COUNTIF(K{A1}:K{A2},"{t}")')
barras(ws, f"H{s+7}", s + 1, s + 5, 8, 11, width=17, height=6.2)
ws.freeze_panes = "D7"
setup(ws, TEAL, f"B1:K{s+19}", fit_height=True)

# ------------------------------------------------------------------ Entradas
ws = wb.create_sheet("Entradas")
widths(ws, {"A": 2, "B": 7, "C": 28, "D": 36, "E": 20, "F": 24, "G": 48, "H": 14, "I": 13, "J": 40, "K": 10, "L": 22, "M": 2})
title(ws, "Entradas da análise crítica", "As doze entradas que a direção precisa considerar. Preencha antes da reunião. A conclusão é escrita na reunião.", "L")
for col, text in zip("BCDEFGHIJKL", ["Item", "Entrada", "O que a direção quer saber", "Quem apresenta", "Fonte", "Resultado no período", "Tendência", "Situação",
                                      "Conclusão da direção", "Decisões", "Conferência"]):
    head(ws, f"{col}4", text)
ws.row_dimensions[4].height = 24
hint_row(ws, 5, [("B", "9.3.2", None), ("C", "Fixa", None), ("D", "Fixo", None), ("E", "Nome ou cargo", None), ("F", "Documento ou sistema", None),
                 ("G", "Fatos e números. Se não houver, escreva “nada a relatar”", None), ("H", "Escolha na lista", None), ("I", "Escolha na lista", None),
                 ("J", "Uma frase, escrita na reunião", None), ("K", "Da aba Saídas", None), ("L", "Calculada", None)])
for k, (item, nome, olhar, fonte, _) in enumerate(ENTRADAS):
    rr = E1 + k
    put(ws, f"B{rr}", item, f=font(10, True), bg=GRAY, h="center", fmt="@")
    put(ws, f"C{rr}", nome, f=font(10, True), bg=GRAY)
    put(ws, f"D{rr}", olhar, f=font(9, c=MUTED), bg=GRAY)
    inp(ws, f"E{rr}")
    inp(ws, f"F{rr}")
    inp(ws, f"G{rr}")
    inp(ws, f"H{rr}", h="center")
    inp(ws, f"I{rr}", h="center")
    inp(ws, f"J{rr}")
    calc(ws, f"K{rr}", f"=COUNTIF(Saídas!$C${S1}:$C${S2},B{rr})", b=False)
    calc(ws, f"L{rr}", f_conf_entrada(rr), b=False, sz=9)
    note(ws, f"C{rr}", f"Onde buscar: {fonte[0].lower()}{fonte[1:]}.")
    ws.row_dimensions[rr].height = 45
dv_list(ws, f"H{E1}:H{E2}", TENDS, "Melhorando, Estável, Piorando ou Sem histórico")
dv_list(ws, f"I{E1}:I{E2}", SITS, "Favorável, Atenção ou Crítica")
cf_equal(ws, f"I{E1}:I{E2}", SIT_CF)
cf_falta(ws, f"L{E1}:L{E2}", f"L{E1}", neutro=("Não analisada",))
s = E2 + 2
assert s == ENT_S
band(ws, s, "Resumo automático", "L")
for k, (text, formula) in enumerate([
    ("Entradas analisadas, de 12", f'=SUMPRODUCT((G{E1}:G{E2}<>"")*(I{E1}:I{E2}<>"")*(J{E1}:J{E2}<>""))'),
    ("Favoráveis", f'=COUNTIF(I{E1}:I{E2},"{FAV}")'),
    ("Em atenção", f'=COUNTIF(I{E1}:I{E2},"{ATE}")'),
    ("Críticas", f'=COUNTIF(I{E1}:I{E2},"{CRI}")'),
    ("Com tendência de piora", f'=COUNTIF(H{E1}:H{E2},"{PIOR}")'),
    ("Críticas sem decisão", f'=COUNTIFS(I{E1}:I{E2},"{CRI}",K{E1}:K{E2},0)'),
    ("Aviso", f'=IF(COUNTIF(L{E1}:L{E2},"Não analisada")=12,"Preencha as entradas",IF(COUNTIF(L{E1}:L{E2},"Não analisada")>0,"Há entrada sem análise",'
              f'IF(COUNTIF(L{E1}:L{E2},"Falta a decisão")>0,"Há entrada crítica sem decisão: registre na aba Saídas",'
              f'IF(COUNTIF(L{E1}:L{E2},"Falta*")>0,"Há entrada com o registro incompleto: veja a coluna Conferência",'
              f'IF(COUNTIF(L{E1}:L{E2},"Atenção sem decisão")>0,"Há entrada em atenção sem decisão: confira se o motivo está na conclusão","OK")))))'),
], 1):
    label(ws, f"B{s+k}", text, merge=f"B{s+k}:D{s+k}", h="right")
    calc(ws, f"E{s+k}", formula, merge=f"E{s+k}:G{s+k}", sz=9 if text == "Aviso" else 10, b=text != "Aviso")
    ws.row_dimensions[s + k].height = 21.75
cf_warn(ws, f"E{s+7}:G{s+7}", f"E{s+7}")
ws.freeze_panes = "D6"
setup(ws, AMBER, f"B1:L{s+7}", fit_height=True)

# ------------------------------------------------------------------ Saídas
ws = wb.create_sheet("Saídas")
widths(ws, {"A": 2, "B": 5, "C": 10, "D": 28, "E": 22, "F": 50, "G": 24, "H": 13, "I": 32, "J": 15, "K": 28, "L": 2})
title(ws, "Saídas: decisões e ações", "Uma decisão por linha. Comece com verbo, e informe de que entrada a decisão veio.", "K")
for col, text in zip("BCDEFGHIJK", ["#", "Entrada", "Nome da entrada", "Tipo", "Decisão", "Responsável", "Prazo", "Recurso necessário", "Valor", "Conferência"]):
    head(ws, f"{col}4", text)
ws.row_dimensions[4].height = 24
hint_row(ws, 5, [("B", "", None), ("C", "Item, de a até f", None), ("D", "Da aba Entradas", None), ("E", "Escolha na lista", None),
                 ("F", "O que será feito", None), ("G", "Nome ou cargo", None), ("H", "Depois da reunião", None), ("I", "Pessoas, equipamento, serviço", None),
                 ("J", "Em reais", None), ("K", "Calculada", None)])
ex(ws, "B6", "Ex.", h="center")
for col, v, h_, fmt in [("C", "c3", "center", "@"), ("D", NOME["c3"], "left", None), ("E", REC, "center", None),
                        ("F", "Aprovar a compra do medidor de espessura em linha para as extrusoras 3 e 4.", "left", None), ("G", "Diretor geral", "left", None),
                        ("H", date(2027, 6, 30), "center", DATE), ("I", "Dois medidores, com instalação", "left", None), ("J", 96000, "center", MOEDA),
                        ("K", "Completa", "center", None)]:
    ex(ws, f"{col}6", v, h=h_, fmt=fmt)
ws.row_dimensions[6].height = 30
assert S1 == 7
for k in range(NSAI):
    rr = S1 + k
    num(ws, f"B{rr}", k + 1)
    inp(ws, f"C{rr}", h="center", fmt="@")
    calc(ws, f"D{rr}", f'=IF(C{rr}="","",IFERROR(VLOOKUP(C{rr},Entradas!$B${E1}:$C${E2},2,FALSE),"Entrada desconhecida"))', h="left", b=False, sz=9)
    inp(ws, f"E{rr}", h="center")
    inp(ws, f"F{rr}")
    inp(ws, f"G{rr}")
    inp(ws, f"H{rr}", h="center", fmt=DATE)
    inp(ws, f"I{rr}")
    inp(ws, f"J{rr}", h="center", fmt=MOEDA)
    calc(ws, f"K{rr}", f'=IF(COUNTA(C{rr},E{rr}:J{rr})=0,"",IF(F{rr}="","Falta a decisão",IF(C{rr}="","Falta a entrada",IF(D{rr}="Entrada desconhecida","Entrada desconhecida",IF(E{rr}="","Falta o tipo",'
         f'IF(G{rr}="","Falta o responsável",IF(H{rr}="","Falta o prazo",IF(AND(Reunião!$D$5<>"",H{rr}<Reunião!$D$5),"Prazo anterior à reunião",'
         f'IF(AND(E{rr}="{REC}",I{rr}="",J{rr}=""),"Falta descrever o recurso","Completa")))))))))', b=False, sz=9)
    ws.row_dimensions[rr].height = 33
dv_list(ws, f"C{S1}:C{S2}", ITENS, "Item da entrada: a, b, c1 a c7, d, e ou f")
dv_list(ws, f"E{S1}:E{S2}", TIPOS, "Melhoria, Mudança no sistema ou Recursos")
dv_date(ws, f"H{S1}:H{S2}")
dv_number(ws, f"J{S1}:J{S2}")
cf_falta(ws, f"K{S1}:K{S2}", f"K{S1}")
note(ws, "E4", "Melhoria: o que será feito melhor.\nMudança no sistema: o que muda em processos, documentos, objetivos ou indicadores.\nRecursos: pessoas, equipamentos ou dinheiro.")
note(ws, "F4", 'Comece com verbo: aprovar, contratar, rever, auditar.\nEvite: "melhorar a comunicação", "verificar a possibilidade".')
s = S2 + 2
assert s == SAI_S
band(ws, s, "Resumo automático", "K")
for k, (text, formula, fmt) in enumerate([
    ("Decisões registradas", f"=COUNTA(F{S1}:F{S2})", None),
    ("De melhoria", f'=COUNTIF(E{S1}:E{S2},"{MEL}")', None),
    ("De mudança no sistema", f'=COUNTIF(E{S1}:E{S2},"{MUD}")', None),
    ("De recursos", f'=COUNTIF(E{S1}:E{S2},"{REC}")', None),
    ("Valor total aprovado", f"=SUM(J{S1}:J{S2})", MOEDA),
    ("Aviso", f'=IF(E{s+1}=0,"Registre as decisões",IF(COUNTIF(K{S1}:K{S2},"Falta*")+COUNTIF(K{S1}:K{S2},"Prazo*")>0,'
              f'"Há decisão com o registro incompleto: veja a coluna Conferência",IF(COUNTIF(D{S1}:D{S2},"Entrada desconhecida")>0,"Há decisão com entrada desconhecida: use os itens de a até f",'
              f'IF(Entradas!E{ENT_S+6}>0,"Há entrada crítica sem decisão","OK"))))', None),
], 1):
    label(ws, f"B{s+k}", text, merge=f"B{s+k}:D{s+k}", h="right")
    calc(ws, f"E{s+k}", formula, fmt=fmt, merge=f"E{s+k}:F{s+k}", sz=9 if text == "Aviso" else 10, b=text != "Aviso")
    ws.row_dimensions[s + k].height = 21.75
cf_warn(ws, f"E{s+6}:F{s+6}", f"E{s+6}")
ws.freeze_panes = "E6"
setup(ws, REDC, f"B1:K{s+6}", fit_height=True)

# ------------------------------------------------------------------ Checklist
ws = wb.create_sheet("Checklist")
widths(ws, {"A": 2, "B": 5, "C": 70, "D": 14, "E": 44, "F": 2})
title(ws, "Checklist da análise crítica", "Escolha o status de cada item na lista suspensa (coluna em amarelo-claro).", "E")
for col, text in zip("BCDE", ["#", "Item de verificação", "Status", "Observações"]):
    head(ws, f"{col}4", text)
ws.row_dimensions[4].height = 21.75
for k, text in enumerate(CHECK):
    rr = 5 + k
    num(ws, f"B{rr}", k + 1)
    put(ws, f"C{rr}", text, bg=GRAY)
    inp(ws, f"D{rr}", h="center")
    inp(ws, f"E{rr}")
    ws.row_dimensions[rr].height = 30
dv_list(ws, "D5:D16", YESNO, "Sim, Parcial ou Não")
cf_equal(ws, "D5:D16", YESNO_CF)
band(ws, 18, "Resumo automático", "E")
summary(ws, 19, "Itens atendidos (Sim)", '=COUNTIF(D5:D16,"Sim")', "C", "D")
summary(ws, 20, "Itens parciais", '=COUNTIF(D5:D16,"Parcial")', "C", "D")
summary(ws, 21, "Itens não atendidos", '=COUNTIF(D5:D16,"Não")', "C", "D")
summary(ws, 22, "Itens sem resposta", "=COUNTBLANK(D5:D16)", "C", "D")
summary(ws, 23, "Percentual atendido", "=D19/ROWS(D5:D16)", "C", "D", fmt="0%")
summary(ws, 24, "Situação", '=IF(D19=ROWS(D5:D16),"Validado",IF(D22>0,"Em andamento","Com pendências"))', "C", "D")
setup(ws, PURPLE, "B1:E24", landscape=False, fit_height=True)


# ------------------------------------------------------------------ Exemplos
def sub(ws, rr, cells):
    """Linha de cabeçalho cinza de uma tabela do exemplo: (coluna inicial, coluna final ou None, texto)."""
    for a, b, t in cells:
        put(ws, f"{a}{rr}", t, f=font(10, True), bg=GRAY, h="center", merge=b and f"{a}{rr}:{b}{rr}")
    ws.row_dimensions[rr].height = 24


def exemplo(ws, data):
    widths(ws, {"A": 2, "B": 7, "C": 26, "D": 22, "E": 48, "F": 14, "G": 14, "H": 40, "I": 13, "J": 22, "K": 2})
    title(ws, "Análise crítica pela direção", "Exemplo preenchido, para consulta. Use as abas Reunião, Ações anteriores, Entradas e Saídas para a sua análise.", "J")
    H = data["head"]
    pres = [p[0] for p in data["part"] if p[3] == "Presente"]
    rr = 4
    for rot, val, fmt in [("Organização", H["org"], None), ("Data da reunião", H["data"], DATE), ("Período analisado", H["periodo"], None),
                          ("Participantes", ", ".join(pres) + ".", None), ("Análise anterior", H["anterior"] or "Não houve: é a primeira análise crítica.", DATE if H["anterior"] else None),
                          ("Origem", H["origem"], None)]:
        label(ws, f"B{rr}", rot, merge=f"B{rr}:C{rr}")
        inp(ws, f"D{rr}", val, merge=f"D{rr}:J{rr}", bg=WHITE, fmt=fmt, h="left")
        ws.row_dimensions[rr].height = 19.5
        rr += 1
    REF = "$D$5"
    rr += 1

    # ---- ações das análises anteriores
    a1 = a2 = None
    if data["anteriores"]:
        band(ws, rr, "Entrada “a” · ações das análises anteriores", "J", color=TEAL)
        rr += 1
        sub(ws, rr, [("B", None, "Nº"), ("C", None, "Data da análise"), ("D", None, "Responsável"), ("E", None, "Ação decidida"), ("F", None, "Prazo"),
                     ("G", None, "Concluída em"), ("H", None, "Eficácia e resultado"), ("I", None, "Status"), ("J", None, "Situação")])
        a1 = rr + 1
        for a in data["anteriores"]:
            rr += 1
            res = (f'Eficaz: {a["eficaz"].lower()}. ' if a["eficaz"] else "") + a["obs"]
            put(ws, f"B{rr}", a["n"], f=font(10, True), bg=GRAY, h="center")
            put(ws, f"C{rr}", a["origem"], h="center", fmt=DATE)
            put(ws, f"D{rr}", a["quem"])
            put(ws, f"E{rr}", a["acao"])
            put(ws, f"F{rr}", a["prazo"], h="center", fmt=DATE)
            put(ws, f"G{rr}", a["feito"], h="center", fmt=DATE)
            put(ws, f"H{rr}", res.strip())
            put(ws, f"I{rr}", a["status"], h="center")
            calc(ws, f"J{rr}", f_situacao(rr, REF, c_acao="E", c_prazo="F", c_status="I", c_feito="G"), b=False, sz=9)
            ws.row_dimensions[rr].height = alt([(a["acao"], 46), (res, 38), (a["quem"], 20)], minimo=24)
        a2 = rr
        cf_equal(ws, f"J{a1}:J{a2}", ACAO_CF)
        rr += 2
        band(ws, rr, "Ações por situação", "E")
        c0 = rr
        for k, t in enumerate(SIT_ACAO, 1):
            label(ws, f"B{rr+k}", t, merge=f"B{rr+k}:D{rr+k}", h="right")
            calc(ws, f"E{rr+k}", f'=COUNTIF(J{a1}:J{a2},"{t}")')
            ws.row_dimensions[rr + k].height = 21.75
        label(ws, f"B{rr+6}", "Percentual concluído, sem as canceladas", merge=f"B{rr+6}:D{rr+6}", h="right")
        calc(ws, f"E{rr+6}", f"=(E{rr+1}+E{rr+2})/(SUM(E{rr+1}:E{rr+5})-E{rr+5})", fmt="0%")
        ws.row_dimensions[rr + 6].height = 21.75
        barras(ws, f"F{c0}", c0 + 1, c0 + 5, 2, 5, width=17, height=5.6)
        rr += 12
        ws.row_breaks.append(Break(id=rr - 1))

    # ---- entradas
    band(ws, rr, "Entradas analisadas", "J", color=AMBER)
    rr += 1
    sub(ws, rr, [("B", None, "Item"), ("C", None, "Entrada"), ("D", None, "Quem apresenta"), ("E", None, "Resultado no período"), ("F", None, "Tendência"),
                 ("G", None, "Situação"), ("H", None, "Conclusão da direção"), ("I", None, "Decisões"), ("J", None, "Conferência")])
    e1 = rr + 1
    e2 = e1 + len(data["entradas"]) - 1
    d1 = e2 + 4              # primeira linha das decisões: uma linha em branco, a faixa e o cabeçalho antes
    d2 = d1 + len(data["saidas"]) - 1
    for e in data["entradas"]:
        rr += 1
        put(ws, f"B{rr}", e["item"], f=font(10, True), bg=GRAY, h="center", fmt="@")
        put(ws, f"C{rr}", NOME[e["item"]], f=font(10, True))
        put(ws, f"D{rr}", e["quem"])
        put(ws, f"E{rr}", e["resultado"])
        put(ws, f"F{rr}", e["tend"], h="center")
        put(ws, f"G{rr}", e["sit"], h="center")
        put(ws, f"H{rr}", e["conclusao"])
        calc(ws, f"I{rr}", f'=COUNTIF($C${d1}:$C${d2},B{rr}&" ·*")', b=False)
        calc(ws, f"J{rr}", f_conf_entrada(rr, res="E", tend="F", sit="G", conc="H", dec="I", primeira="D"), b=False, sz=9)
        ws.row_dimensions[rr].height = alt([(e["resultado"], 50), (e["conclusao"], 42), (NOME[e["item"]], 24)], minimo=31)
    assert rr == e2
    cf_equal(ws, f"G{e1}:G{e2}", SIT_CF)
    cf_falta(ws, f"J{e1}:J{e2}", f"J{e1}", neutro=("Não analisada",))
    rr += 2
    ws.row_breaks.append(Break(id=rr - 1))

    # ---- saídas
    band(ws, rr, "Saídas: decisões e ações", "J", color=REDC)
    rr += 1
    sub(ws, rr, [("B", None, "Nº"), ("C", None, "Entrada"), ("D", None, "Tipo"), ("E", None, "Decisão"), ("F", "G", "Responsável"), ("H", None, "Recurso necessário"),
                 ("I", None, "Prazo"), ("J", None, "Valor")])
    assert rr + 1 == d1, (rr, d1)
    for s_ in data["saidas"]:
        rr += 1
        put(ws, f"B{rr}", s_["n"], f=font(10, True), bg=GRAY, h="center")
        put(ws, f"C{rr}", f'{s_["item"]} · {NOME[s_["item"]]}')
        put(ws, f"D{rr}", s_["tipo"], h="center")
        put(ws, f"E{rr}", s_["decisao"])
        put(ws, f"F{rr}", s_["quem"], merge=f"F{rr}:G{rr}")
        put(ws, f"H{rr}", s_["recurso"])
        put(ws, f"I{rr}", s_["prazo"], h="center", fmt=DATE)
        put(ws, f"J{rr}", s_["valor"], h="center", fmt=MOEDA)
        ws.row_dimensions[rr].height = alt([(s_["decisao"], 48), (NOME[s_["item"]], 22), (s_["recurso"], 40)], minimo=30)
    assert rr == d2
    rr += 2

    # ---- avaliação do sistema
    band(ws, rr, "Avaliação do sistema", "J", color=BLUE)
    rr += 1
    sub(ws, rr, [("B", "C", "Critério"), ("D", "E", "Pergunta"), ("F", None, "Resposta"), ("G", "J", "Justificativa")])
    v1 = rr + 1
    for (nome, perg, _), (resp, just) in zip(CRITERIOS, data["sistema"]):
        rr += 1
        label(ws, f"B{rr}", nome, merge=f"B{rr}:C{rr}")
        put(ws, f"D{rr}", perg, merge=f"D{rr}:E{rr}")
        put(ws, f"F{rr}", resp, h="center")
        put(ws, f"G{rr}", just, merge=f"G{rr}:J{rr}")
        ws.row_dimensions[rr].height = 24
    cf_equal(ws, f"F{v1}:F{rr}", YESNO_CF)
    rr += 2

    # ---- resumo
    band(ws, rr, "Resumo automático", "J")
    itens = [("Entradas analisadas, de 12", f'=SUMPRODUCT((E{e1}:E{e2}<>"")*(G{e1}:G{e2}<>"")*(H{e1}:H{e2}<>""))', None),
             ("Favoráveis, em atenção e críticas", f'=COUNTIF(G{e1}:G{e2},"{FAV}")&", "&COUNTIF(G{e1}:G{e2},"{ATE}")&" e "&COUNTIF(G{e1}:G{e2},"{CRI}")', None),
             ("Críticas sem decisão", f'=COUNTIFS(G{e1}:G{e2},"{CRI}",I{e1}:I{e2},0)', None),
             ("Decisões registradas", f"=COUNTA(E{d1}:E{d2})", None),
             ("De melhoria, de mudança no sistema e de recursos", f'=COUNTIF(D{d1}:D{d2},"{MEL}")&", "&COUNTIF(D{d1}:D{d2},"{MUD}")&" e "&COUNTIF(D{d1}:D{d2},"{REC}")', None),
             ("Valor total aprovado", f"=SUM(J{d1}:J{d2})", MOEDA)]
    if a1:
        itens.append(("Ações anteriores atrasadas", f'=COUNTIF(J{a1}:J{a2},"{A_VENCIDA}")', None))
    for k, (text, formula, fmt) in enumerate(itens, 1):
        label(ws, f"B{rr+k}", text, merge=f"B{rr+k}:D{rr+k}", h="right")
        calc(ws, f"E{rr+k}", formula, fmt=fmt)
        ws.row_dimensions[rr + k].height = 21.75
    setup(ws, MUTED, f"B1:J{rr + len(itens)}")


exemplo(wb.create_sheet("Exemplo 1 - Pizzaria"), EX1)
exemplo(wb.create_sheet("Exemplo 2 - Indústria"), EX2)

wb.active = 1
wb.save(OUT)
print("ok", OUT)
